"""Compare a class-per-sheet review export with a populated SPEC reference.

Coverage candidates are not engineering matches or reference-code approvals.
Source and expected workbooks remain unchanged. No API or database is used.
"""
import argparse
from collections import Counter
from fractions import Fraction
import hashlib
import json
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


FAMILIES = {
    'PIPE': 'Piping', '45 DEG LR ELBOW': '45 Degree Direction Change',
    '45 DEG ELBOW': '45 Degree Direction Change', '90 DEG ELBOW': '90 Degree Direction Change',
    '90 DEG LR ELBOW': '90 Degree Direction Change', '90 DEG SR ELBOW': '90 Degree SR Elbow',
    'CON. REDUCER': 'Concentric Size Change', 'CONCENTRIC REDUCER': 'Concentric Size Change',
    'CONCENTRIC SWAGE': 'Concentric Size Change', 'ECC. REDUCER': 'Eccentric Size Change',
    'ECCENTRIC REDUCER': 'Eccentric Size Change', 'ECCENTRIC SWAGE': 'Eccentric Size Change',
    'TEE': 'Tee', 'EQUAL TEE': 'Tee', 'REDUCING TEE': 'Reducing Tee', 'CAP': 'Cap',
    'FLANGE': 'Flange', 'BLIND FLANGE': 'Blind Flange', 'ORIFICE FLANGE': 'Orifice Flange',
    'GATE VALVE': 'Gate Valve', 'CHECK VALVE': 'Check Valve', 'GLOBE VALVE': 'Globe Valve',
    'BALL VALVE': 'Ball Valve', 'NEEDLE VALVE': 'Needle Valve', 'COUPLING': 'Coupling',
    'PIPE NIPPLE': 'Nipple', 'NIPPLE': 'Nipple', 'PLUG': 'Plug', 'RPAD': 'Reinforcing pad',
    'SOCKOLET': 'Sockolet', 'WELDOLET': 'Weldolet', 'ELBOLET': 'Elbolet',
}
SUSPECT_CODES = {'AND', 'BRANCH', 'DATA', 'DATE', 'DEFINITION', 'DESIGNATION', 'FOR',
                 'FROM', 'HAS', 'IAL', 'IALS', 'IFICATION', 'INDEX', 'IS', 'ON',
                 'PRESSURE', 'REFERENCE', 'SHALL', 'SHEETS', 'TO', 'VALVE'}


def text(value):
    return '' if value is None else str(value).strip()


def size(value, unit=''):
    raw = text(value)
    unit = text(unit).lower()
    if raw.upper().startswith('DN') or re.search(r'\bmm\b', raw, re.I):
        unit = 'mm'
    elif raw.upper().startswith('NPS') or any(char in raw for char in '\"“”″') or re.search(r'\bin(?:ches)?\b', raw, re.I):
        unit = 'in'
    raw = re.sub(r'^(DN|NPS)\s*', '', raw, flags=re.I)
    raw = re.sub(r'\b(?:mm|in|inch|inches)\b', '', raw, flags=re.I)
    raw = raw.translate(str.maketrans({'“': '', '”': '', '″': '', '"': '', '½': ' 1/2', '¼': ' 1/4', '¾': ' 3/4'})).strip()
    try:
        parts = raw.split()
        number = sum((Fraction(part) for part in parts), Fraction()) if parts else None
        return number, unit
    except (ValueError, ZeroDivisionError):
        return None, unit


def marker_data(sheet):
    headers = []
    active = False
    result = []
    for index, row in enumerate(sheet.iter_rows(values_only=True), 1):
        marker = text(row[0]) if row else ''
        if marker == 'Head':
            headers = [text(value) for value in row]
        elif marker == 'Start':
            active = True
        elif marker == 'End':
            active = False
        elif active and any(text(value) for value in row[1:]):
            result.append({'Excel row': index, **dict(zip(headers, row))})
    return headers, result


def review_data(sheet):
    metadata, headers, rows = {}, [], []
    for index, row in enumerate(sheet.iter_rows(values_only=True), 1):
        if not headers and len(row) >= 2 and row[0] is not None:
            metadata[text(row[0])] = row[1]
        if row and text(row[0]) == 'Type' and 'Size From' in row:
            headers = [text(value) for value in row]
        elif headers and any(text(value) for value in row):
            rows.append({'Excel row': index, **dict(zip(headers, row))})
    return metadata, rows


def write_sheet(book, name, headers, rows):
    sheet = book.create_sheet(name)
    for values in [headers, *rows]:
        sheet.append([json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value for value in values])
        # Text from source workbooks is kept literal, not re-emitted as formulas.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='17365D')
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(wrap_text=True)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(index)].width = min(65, max(16, len(header) + 3))
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    return sheet


def compare(extracted, expected, output, default_template=None, legacy_template=None):
    if output.resolve() in {extracted.resolve(), expected.resolve()}:
        raise ValueError('Comparison output must not overwrite either source workbook.')
    raw = load_workbook(extracted, read_only=True, data_only=True)
    reference = load_workbook(expected, read_only=True, data_only=True)
    current = load_workbook(default_template, read_only=True, data_only=True) if default_template else None
    legacy = load_workbook(legacy_template, read_only=True, data_only=True) if legacy_template else None
    _, reference_classes = marker_data(reference['PipingMaterialsClassData'])
    names = {text(row.get('SPECNAME')) for row in reference_classes if text(row.get('SPECNAME'))}
    if len(names) != 1:
        raise ValueError('The expected workbook must identify one explicit SPECNAME for this comparison.')
    class_name = names.pop()
    if class_name not in raw.sheetnames:
        raise ValueError(f'The extracted review workbook has no exact class sheet {class_name}.')
    metadata, components = review_data(raw[class_name])
    expected_headers, expected_rows = marker_data(reference['PipingCommodityFilter'])
    candidate_rows = []
    expected_without_candidates = []
    reverse_candidates = {component['Excel row']: [] for component in components}
    for row in expected_rows:
        family = text(row.get('SHORTCODE'))
        low, unit = size(row.get('FIRSTSIZEFROM'), row.get('FIRSTSIZEUNITS'))
        high, other_unit = size(row.get('FIRSTSIZETO'), row.get('FIRSTSIZEUNITS'))
        family_rows = [component for component in components if FAMILIES.get(text(component.get('Sub-Type')).upper()) == family]
        candidates = []
        for component in family_rows:
            first, first_unit = size(component.get('Size From'))
            last, last_unit = size(component.get('Size To'))
            first_unit, last_unit = first_unit or last_unit, last_unit or first_unit
            if None not in (first, last, low, high) and unit and first_unit == last_unit == unit == other_unit and first <= low <= high <= last:
                candidates.append(component)
                reverse_candidates[component['Excel row']].append(row['Excel row'])
        status = 'Coverage candidate; engineering fields unverified' if candidates else 'No comparable first-size coverage' if family_rows else 'No mapped component family'
        secondary = bool(text(row.get('SECONDSIZEFROM')) or text(row.get('SECONDSIZETO')))
        candidate_rows.append([row['Excel row'], family, row.get('FIRSTSIZEFROM'), row.get('FIRSTSIZETO'), row.get('FIRSTSIZEUNITS'), row.get('SECONDSIZEFROM'), row.get('SECONDSIZETO'), row.get('COMMODITYCODE'), status, len(candidates), ', '.join(str(item['Excel row']) for item in candidates), 'Unverified' if secondary else 'Not stated', 'Not compared: raw export has no explicit commodity-code field'])
        if not candidates:
            expected_without_candidates.append(candidate_rows[-1])
    inventory = []
    for name in reference.sheetnames:
        headers, rows = marker_data(reference[name])
        new_headers = marker_data(current[name])[0] if current and name in current.sheetnames else []
        old_headers = marker_data(legacy[name])[0] if legacy and name in legacy.sheetnames else []
        inventory.append([name, len(rows), len(headers), name in raw.sheetnames,
                          bool(current and name in current.sheetnames), headers == new_headers if headers else 'Metadata sheet',
                          bool(legacy and name in legacy.sheetnames), ', '.join(value for value in headers if value and value not in old_headers), ', '.join(headers)])
    class_inventory = []
    for sheet in raw:
        info, rows = review_data(sheet)
        code = text(info.get('Class Code')) or sheet.title
        class_inventory.append([sheet.title, code, len(rows), 'Review: possible prose/header token' if code in SUSPECT_CODES else '', 'Header only' if not rows else 'Has component candidates'])
    source_gaps = []
    for component in components:
        missing = [name for name in ('Material Std', 'End Conn.', 'Schedule/Rating', 'Description') if not text(component.get(name))]
        family = FAMILIES.get(text(component.get('Sub-Type')).upper(), '')
        notes = text(component.get('Notes'))
        source_gaps.append([component['Excel row'], component.get('Type'), component.get('Sub-Type'), family, component.get('Size From'), component.get('Size To'), ', '.join(missing), 'Header row was extracted as a component' if text(component.get('Sub-Type')).upper() == 'PART NAME' else 'Not in this commodity-filter mapping' if not family else '', re.search(r'Source page\s+(\d+)', notes).group(1) if re.search(r'Source page\s+(\d+)', notes) else '', ', '.join(map(str, reverse_candidates[component['Excel row']])), notes])
    stats = {'reference_class': class_name, 'extracted_sheets': len(raw.sheetnames), 'expected_sheets': len(reference.sheetnames),
             'raw_sheet_name_matches': len(set(raw.sheetnames) & set(reference.sheetnames)),
             'new_default_sheet_order_matches': bool(current and current.sheetnames == reference.sheetnames),
             'new_default_header_mismatches': sum(item[5] is False for item in inventory) if current else None,
             'legacy_sheets': len(legacy.sheetnames) if legacy else None,
             'extracted_class_rows': len(components), 'expected_commodity_rows': len(expected_rows),
             'expected_rows_with_family_first_size_candidates': sum(bool(item[9]) for item in candidate_rows),
             'expected_rows_without_candidates': len(expected_without_candidates),
             'expected_rows_with_secondary_sizes': sum(item[11] == 'Unverified' for item in candidate_rows),
             'source_rows_with_material_field': sum(bool(text(item.get('Material Std'))) for item in components),
             'source_rows_with_end_connection_field': sum(bool(text(item.get('End Conn.'))) for item in components),
             'source_header_rows_misclassified': sum(text(item.get('Sub-Type')).upper() == 'PART NAME' for item in components),
             'source_rows_without_commodity_filter_candidates': sum(not values for values in reverse_candidates.values()),
             'fully_verified_engineering_matches': 0,
             'class_header_only_sheets': sum(item[4] == 'Header only' for item in class_inventory),
             'possible_prose_class_sheets': sum(bool(item[3]) for item in class_inventory),
             'source_sha256': hashlib.sha256(extracted.read_bytes()).hexdigest(),
             'expected_sha256': hashlib.sha256(expected.read_bytes()).hexdigest()}
    book = Workbook(); book.remove(book.active)
    summary = [[key.replace('_', ' '), value] for key, value in stats.items()]
    summary += [['Scope', f'Only {class_name} source rows compared with the populated expected class. Other source classes are not expected to match one AA1B workbook.'],
                ['Coverage scope', 'Candidate analysis targets PipingCommodityFilter only. Gasket/bolt selection and other rule sheets require their own structured fields and reference configuration.'],
                ['Interpretation', 'Family/first-size coverage candidates are not verified matches. No count subtraction is an extraction-accuracy score.'],
                ['Reference authority', 'The supplied workbook is a comparison baseline; document/revision equivalence and engineering approval are not established.'],
                ['Identifiers', 'Source ME numbers in Notes are not automatically SmartPlant commodity codes. No expected code was assigned to a source row.'],
                ['Structured field gaps', 'Blank material/end-connection fields do not prove the source descriptions contain no material or end details; inspect the literal descriptions and PDF evidence.'],
                ['Verification count', 'Zero fully verified matches means this report has not completed engineering matching; it is not a zero-percent accuracy score.'],
                ['Default change', '50-sheet schema only; example engineering values and source-machine metadata are removed.'],
                ['Release state', 'Review report only. Existing matching, partial-processing and engineering-release gates remain.']]
    write_sheet(book, 'Summary', ['Measure', 'Value'], summary)
    write_sheet(book, 'Sheet comparison', ['Expected sheet', 'Expected data rows', 'Columns', 'In raw export', 'In new default', 'New headers match', 'In old default', 'New column spelling vs old', 'Expected headers'], inventory)
    write_sheet(book, 'AA1B class fields', ['Extracted field', 'Extracted value', 'Expected field / value', 'Assessment'], [
        ['Class Code', metadata.get('Class Code'), f'SPECNAME={class_name}', 'Same class code'],
        ['Material', metadata.get('Material'), reference_classes[0].get('MATERIALSDESCRIPTION'), 'Missing source class metadata' if not metadata.get('Material') else 'Review description equivalence'],
        ['Services', metadata.get('Services'), reference_classes[0].get('FLUIDSERVICE'), 'Missing source class metadata' if not metadata.get('Services') else 'Review service equivalence'],
        ['Rating', metadata.get('Rating'), reference_classes[0].get('MATERIALSDESCRIPTION'), 'Expected rating is in reference description; source field missing' if not metadata.get('Rating') else 'Review'],
        ['Facing', metadata.get('Facing'), reference_classes[0].get('MATERIALSDESCRIPTION'), 'Expected facing is in reference description; source field missing' if not metadata.get('Facing') else 'Review'],
        ['Corrosion Allowance', metadata.get('Corrosion Allowance'), marker_data(reference['CorrosionAllowance'])[1], 'Review explicit source/reference units and values'],
    ])
    raw_headers = ['Excel row', *[key for key in components[0] if key != 'Excel row']] if components else ['Excel row']
    write_sheet(book, 'AA1B extracted', raw_headers, [[component.get(key) for key in raw_headers] for component in components])
    write_sheet(book, 'AA1B expected', ['Excel row', *expected_headers], [[row.get(key) for key in ['Excel row', *expected_headers]] for row in expected_rows])
    comparison_headers = ['Expected Excel row', 'Expected ShortCode', 'First size from', 'First size to', 'Unit', 'Second size from', 'Second size to', 'Expected commodity code', 'Coverage result', 'Source candidate count', 'Source Excel rows', 'Second-size check', 'Identifier check']
    write_sheet(book, 'Coverage candidates', comparison_headers, candidate_rows)
    write_sheet(book, 'Expected without candidates', comparison_headers, expected_without_candidates)
    gap_headers = ['Source Excel row', 'Type', 'Sub-type', 'Candidate expected family', 'Size from', 'Size to', 'Missing source fields', 'Layout/classification issue', 'Source PDF page', 'Expected candidate rows', 'Source evidence']
    write_sheet(book, 'Extraction field gaps', gap_headers, source_gaps)
    write_sheet(book, 'Source without PCF candidates', gap_headers, [row for row in source_gaps if not reverse_candidates[row[0]]])
    fields = []
    common_mapping = {'SPECNAME': ('Class Code', 'Exact class-code comparison'),
                      'SHORTCODE': ('Sub-Type', 'Explicit family mapping; candidate only'),
                      'FIRSTSIZEFROM': ('Size From / Size To', 'First-size containment comparison'),
                      'FIRSTSIZETO': ('Size From / Size To', 'First-size containment comparison'),
                      'FIRSTSIZEUNITS': ('Unit stated in size text', 'No unit conversion'),
                      'SECONDSIZEFROM': ('No separate structured field', 'Unverified'),
                      'SECONDSIZETO': ('No separate structured field', 'Unverified'),
                      'COMMODITYCODE': ('No explicit commodity-code field', 'Reference matching/review required'),
                      'MATERIALSDESCRIPTION': ('Class Material; component descriptions', 'Class material is blank; descriptions require separate interpretation'),
                      'FLUIDSERVICE': ('Services', 'Class service field comparison'),
                      'CORROSIONALLOWANCE': ('Corrosion Allowance', 'Class metadata comparison; missing source value'),
                      'FIRSTSIZESCHEDULE': ('Schedule/Rating', 'Mixed source field; schedule vs rating not automatically resolved')}
    for name in reference.sheetnames:
        headers, records = marker_data(reference[name])
        for header in headers:
            if not header: continue
            if header == 'Head':
                source_field, status = 'Workbook control marker', 'Structure comparison only'
            else:
                source_field, status = common_mapping.get(header.upper(), ('No direct review-export column', 'Not verified; rule/reference configuration or additional extraction needed'))
            fields.append([name, header, sum(bool(text(record.get(header))) for record in records), source_field, status])
    write_sheet(book, 'Expected field mapping', ['Expected sheet', 'Expected column', 'Populated reference cells', 'Potential source field', 'Comparison scope / gap'], fields)
    write_sheet(book, 'Class inventory', ['Source sheet', 'Class code', 'Component rows', 'Class-heading concern', 'Extraction state'], class_inventory)
    book.properties.title = 'RADAI AA1B workbook comparison - engineering review'
    book.save(output);book.close()
    output.with_suffix('.json').write_text(json.dumps(stats, indent=2) + '\n', encoding='utf-8')
    for opened in (raw, reference, current, legacy):
        if opened: opened.close()
    return stats


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--extracted', required=True, type=Path)
    parser.add_argument('--expected', required=True, type=Path)
    parser.add_argument('--output', required=True, type=Path)
    parser.add_argument('--default-template', type=Path)
    parser.add_argument('--legacy-template', type=Path)
    args = parser.parse_args()
    print(json.dumps(compare(args.extracted, args.expected, args.output, args.default_template, args.legacy_template), indent=2))
