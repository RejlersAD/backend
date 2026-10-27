"""Prepare a schema-only SPEC template from a user-selected reference workbook.

Workbook cells are data, never instructions. Reference engineering rows and
origin-machine metadata are removed; original source workbooks are not changed.
"""
import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


def headers(book, name):
    if name not in book.sheetnames:
        return []
    for row in book[name].iter_rows(min_row=1, max_row=min(book[name].max_row, 60)):
        if str(row[0].value or '').strip() == 'Head':
            return [str(cell.value or '').strip() for cell in row]
    # The retained legacy PipingCommodityFilter has an unmarked header row.
    for row in book[name].iter_rows(min_row=1, max_row=min(book[name].max_row, 10)):
        if row[0].value in (None, '') and sum(isinstance(cell.value, str) and bool(cell.value.strip()) for cell in row[1:]) >= 3:
            return [str(cell.value or '').strip() for cell in row]
    return []


def prepare(reference, output, legacy):
    if reference.resolve() == output.resolve():
        raise ValueError('The reference workbook must remain unchanged.')
    old = load_workbook(legacy, read_only=True, data_only=True)
    book = load_workbook(reference, keep_links=False)
    profile = {'name': 'SmartPlant SPEC - 50-sheet reference structure',
               'source_filename': reference.name,
               'source_sha256': hashlib.sha256(reference.read_bytes()).hexdigest(),
               'sheet_names': book.sheetnames, 'canonical_headers': {},
               'index_sheet': 'Index', 'index_name_column': 2, 'index_status_column': 3,
               'revision_sheet': 'Revision History', 'revision_title_cell': 'A2'}
    for sheet in book:
        actual = headers(book, sheet.title)
        canonical = {value.casefold(): value for value in headers(old, sheet.title) if value}
        profile['canonical_headers'][sheet.title] = {
            value: canonical.get(value.casefold(), value) for value in actual if value}
        markers = {str(sheet.cell(row, 1).value or '').strip(): row
                   for row in range(1, sheet.max_row + 1)
                   if str(sheet.cell(row, 1).value or '').strip() in {'Start', 'End'}}
        if 'Start' in markers and 'End' in markers:
            start, end = markers['Start'], markers['End']
            first = start + 1
            if end > first:
                # Retain one blank, styled exemplar for generated data rows.
                if end > first + 1:
                    sheet.delete_rows(first + 1, end - first - 1)
                for cell in sheet[first]:
                    if isinstance(cell, MergedCell):
                        continue
                    cell.value = None
                    cell.comment = None
                    cell.hyperlink = None
            else:
                sheet.insert_rows(first)
        if sheet.title == profile['revision_sheet']:
            for row in sheet.iter_rows(min_col=2):
                for cell in row:
                    if isinstance(cell, MergedCell):
                        continue
                    cell.value = None
                    cell.comment = None
                    cell.hyperlink = None
        if sheet.title == profile['index_sheet']:
            for row in sheet.iter_rows(min_row=5):
                if row[1].value in book.sheetnames:
                    row[2].value = 'Not populated'
    old.close()
    book.properties.creator = 'RADAI'
    book.properties.lastModifiedBy = 'RADAI'
    book.properties.title = 'RADAI SPEC template - structure only'
    book.properties.subject = 'Blank template; engineering values require extraction and review'
    book.save(output)
    book.close()
    profile_path = output.with_name('smartplant_spec_profile.json')
    profile_path.write_text(json.dumps(profile, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
    return profile


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--reference', required=True, type=Path)
    parser.add_argument('--output', required=True, type=Path)
    parser.add_argument('--legacy', required=True, type=Path)
    args = parser.parse_args()
    result = prepare(args.reference, args.output, args.legacy)
    print(json.dumps({'sheets': len(result['sheet_names']), 'source_sha256': result['source_sha256']}))
