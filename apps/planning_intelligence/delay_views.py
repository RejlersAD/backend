from io import BytesIO
import json

from django.core.exceptions import ValidationError as ModelValidationError
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from rest_framework import serializers
from rest_framework.permissions import IsAuthenticated
from rest_framework.negotiation import DefaultContentNegotiation
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.rbac.action_policy import module_action_allowed
from .access import accessible_projects
from .delay_serializers import DelayCommandSerializer
from .services.delay_analysis import DelayAnalysisError
from .services.delay_cases import delay_command, delay_state
from .services.schedule_approval import ScheduleApprovalError


class DelayAnalysisView(APIView):
    permission_classes = [IsAuthenticated]
    business_approval_actions = {'DelayAnalysisView'}

    @property
    def permission_action(self):
        request = getattr(self, 'request', None)
        if not request or request.method in {'GET', 'HEAD', 'OPTIONS'}:
            return 'read'
        command = request.data.get('action') if isinstance(request.data, dict) else None
        if command in {'approve_case', 'reject_case'}:
            return 'approve'
        if command == 'return_case' and not module_action_allowed(request.user, 'planning_package', 'update'):
            return 'approve'
        return 'update'

    def get(self, request, project_id):
        project = get_object_or_404(accessible_projects(request.user), pk=project_id)
        identifiers = {key: serializers.IntegerField(min_value=1).run_validation(request.query_params[key])
                       for key in ('baseline_id', 'case_id') if key in request.query_params}
        try:
            return Response(delay_state(project, request.user, **identifiers))
        except ScheduleApprovalError as exc:
            return Response(exc.payload, status=exc.status_code)

    def post(self, request, project_id):
        project = get_object_or_404(accessible_projects(request.user), pk=project_id)
        serializer = DelayCommandSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        try:
            return Response(delay_command(project, request.user, serializer.validated_data))
        except ScheduleApprovalError as exc:
            return Response(exc.payload, status=exc.status_code)
        except DelayAnalysisError as exc:
            return Response({'detail': str(exc), 'code': exc.code, 'issues': exc.issues}, status=400)
        except ModelValidationError as exc:
            raise serializers.ValidationError(getattr(exc, 'message_dict', exc.messages)) from exc


def _cell(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, allow_nan=False)
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub(lambda match: f'\\u{ord(match.group()):04x}', value)
        if len(value) > 30000:
            value = value[:29900] + ' [continued in the complete JSON package sheet]'
        if value.lstrip().startswith(('=', '+', '-', '@')):
            return "'" + value
    return value


def _sheet(book, title, rows, fields):
    sheet = book.create_sheet(title)
    sheet.append(fields)
    for row in rows:
        sheet.append([_cell(row.get(key)) for key in fields])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 28


class ReviewExportNegotiation(DefaultContentNegotiation):
    def filter_renderers(self, renderers, format):
        return renderers if format in {'json', 'xlsx'} else super().filter_renderers(renderers, format)


class DelayCaseExportView(APIView):
    permission_classes = [IsAuthenticated]
    content_negotiation_class = ReviewExportNegotiation
    # DRF otherwise treats ?format=xlsx as a renderer override and returns 404.
    def get_format_suffix(self, **kwargs):
        return None

    def get(self, request, project_id, case_id):
        project = get_object_or_404(accessible_projects(request.user), pk=project_id)
        kind = request.query_params.get('format', 'json')
        if kind not in {'json', 'xlsx'}:
            raise serializers.ValidationError('Choose JSON or XLSX for the review package.')
        state = delay_state(project, request.user, case_id=case_id)
        case = state['case']
        if not case.get('run'):
            raise serializers.ValidationError('Calculate a case before exporting its review package.')
        package = {'schema': 'radai.delay-review/1.0', 'project_id': project.pk,
                   'baseline': state['baseline'], 'method_boundary': state['method_boundary'], 'case': case}
        if kind == 'json':
            content = json.dumps(package, ensure_ascii=False, allow_nan=False, indent=2).encode('utf-8')
            mime = 'application/json'
        else:
            book = Workbook()
            book.remove(book.active)
            run, result = case['run'], case['run']['result']
            metadata = {'project_id': project.pk, 'baseline_id': case['baseline_id'],
                'case_id': case['id'], 'case_name': case['name'], 'revision': case['revision'],
                'status': case['status'], 'source_stale': case['source_stale'],
                'run_id': run['id'], 'run_fingerprint': run['fingerprint'],
                'reference_report_id': case['reference_report_id'], 'data_date': case.get('reference_data_date'),
                'method_boundary': state['method_boundary'], 'issues': case['issues'] + result.get('issues', []),
                'recommendation': case['recommendation'], 'assessment': case.get('recommendation_assessment'),
                'reviewed_by_id': case.get('reviewed_by_id'), 'reviewed_at': case.get('reviewed_at'),
                'limitations': result.get('limitations', [])}
            _sheet(book, 'Review', [{'field': key, 'value': value} for key, value in metadata.items()], ['field', 'value'])
            _sheet(book, 'Events', run.get('events', []), ['id', 'revision', 'title', 'description', 'start_date', 'end_date',
                'activity_ids', 'evidence', 'governance_source', 'risk_source'])
            changes = [dict(row, scenario='Impact') for row in case['changes']]
            for scenario in case['scenarios']:
                changes.extend(dict(row, scenario=scenario['id']) for row in scenario['changes'])
            _sheet(book, 'Changes', changes, ['scenario', 'event_id', 'field', 'activity_id', 'predecessor_id',
                'successor_id', 'expected_before', 'value', 'evidence', 'reason'])
            scenarios = [dict(result.get('impact') or {}, id='impact', name='Combined impact')] + result.get('scenarios', [])
            _sheet(book, 'Scenarios', scenarios, ['id', 'name', 'status', 'forecast_finish',
                'net_finish_shift_calendar_days', 'recovered_calendar_days', 'contract_overrun_calendar_days'])
            for title, key in [('Activities', 'affected_activities'), ('Milestones', 'affected_milestones')]:
                rows = [dict(row, scenario=scenario['id']) for scenario in scenarios for row in scenario.get(key, [])]
                _sheet(book, title, rows, ['scenario', 'activity_id', 'external_id', 'name', 'reference_start',
                    'reference_finish', 'scenario_start', 'scenario_finish', 'start_shift_calendar_days',
                    'finish_shift_calendar_days', 'reference_float', 'scenario_float'])
            _sheet(book, 'Paths', [{'scenario': scenario['id'], 'paths': scenario.get('paths', {})} for scenario in scenarios], ['scenario', 'paths'])
            # Excel cells have a 32,767-character limit. Retain the complete
            # review package in ordered chunks as well as the readable sheets.
            complete = json.dumps(package, ensure_ascii=True, allow_nan=False, separators=(',', ':'))
            archive = book.create_sheet('Complete JSON package')
            archive.append(['part', 'json'])
            for index in range(0, len(complete), 29000):
                archive.append([index // 29000 + 1, complete[index:index + 29000]])
                archive.cell(row=archive.max_row, column=2).data_type = 's'
            stream = BytesIO()
            book.save(stream)
            content = stream.getvalue()
            mime = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        response = HttpResponse(content, content_type=mime)
        response['Content-Disposition'] = f'attachment; filename="delay-case-{case_id}-run-{case["run"]["id"]}.{kind}"'
        response['Cache-Control'] = 'private, no-store'
        return response
