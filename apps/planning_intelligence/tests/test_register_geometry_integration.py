"""Geometry joins existing extraction without rewriting source or approvals."""
from copy import deepcopy
import hashlib
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from django.core.files.base import ContentFile
from django.db import connection
from django.test import SimpleTestCase
from django.test.utils import CaptureQueriesContext

from ..intelligence_serializers import DocumentProfileSerializer
from ..services.document_intelligence import profile_document, run_document_intelligence
from ..services.document_plan import build_document_plan, source_files
from ..services.extraction_coverage import file_coverage, summarize_coverage
from ..services.parsers import extract_text_with_coverage
from ..services.pdf_register_geometry import GEOMETRY_VERSION, extract_pdf_register_rows
from ..services.planning_package_sources import build_planning_package_sources
from ..services.register_geometry_cache import RegisterGeometrySourceChanged, cached_register_geometry, ensure_register_geometry
from ..services.register_rows import extract_register_rows, register_rows_for_file
from .test_document_intelligence import DocumentIntelligenceFixture
from .test_pdf_register_geometry import make_register_pdf


def source_pdf():
    return make_register_pdf([
        {'code': '3.1.1', 'title': 'Reliability Availability and\nMaintenance Study Report', 'merged': True},
        {'code': '3.1.2', 'title': 'Emergency Response Plan', 'remarks': 'Included in safety report'},
        {'code': '3.1.3', 'title': 'Conditional Design Report', 'remarks': 'Only if equipment changes'},
        {'code': '3.1.4', 'title': 'Excluded Design Report', 'remarks': 'Not required'},
        {'code': '3.1.5', 'title': 'Unmarked Design Report', 'marks': []},
    ])


def parsed_source():
    stream = source_pdf()
    stream.name = 'synthetic/register.pdf'
    text, _confidence, coverage = extract_text_with_coverage(stream, 'register.pdf')
    return stream, text, coverage


def mixed_project_sources(project_id, first_file_id):
    """Distinct real formats and vocabulary, without relying on a saved SOW."""
    from docx import Document
    from openpyxl import Workbook

    document = Document()
    document.add_paragraph('Environmental monitoring handover')
    table = document.add_table(rows=1, cols=3)
    for cell, label in zip(table.rows[0].cells, ['Item', 'Department', 'Document Title']):
        cell.text = label
    for values in [('1', 'Environment', 'Sediment sample register'),
                   ('2', 'Operations', 'Valve service | readiness "B"')]:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    word_stream = BytesIO()
    document.save(word_stream)

    workbook = Workbook()
    workbook.active.title = 'Site A'
    workbook.active.append(['Document Number', 'Department', 'Document Title'])
    workbook.active.append(['ENV-A-1', 'Ecology', 'Wetland habitat\nsurvey dossier'])
    second = workbook.create_sheet('Site B')
    second.append(['Document Title', 'Document Number', 'Department'])
    second.append(['Shared readiness record', 'ENV-B-2', 'Operations'])
    workbook_stream = BytesIO()
    workbook.save(workbook_stream)
    workbook.close()

    pdf_stream = make_register_pdf([{'code': '9.4.7', 'discipline': 'Water',
                                    'title': 'Aquifer monitoring\nacceptance dossier'}])
    sources = []
    for offset, (stream, filename, category) in enumerate([
        (word_stream, 'field-brief.docx', 'other'),
        (workbook_stream, 'ecology-register.xlsx', 'mdr'),
        (pdf_stream, 'groundwater-plan.PDF', 'other'),
    ]):
        stream.name = f'isolated/{project_id}/{filename}'
        text, _confidence, coverage = extract_text_with_coverage(stream, filename)
        sources.append({'id': first_file_id + offset, 'project_id': project_id,
                        'filename': filename, 'category': category, 'parse_status': 'done',
                        'text': text, 'structured_evidence': coverage.get('structured_evidence') or {}})
    return sources


class RegisterGeometryReadTests(SimpleTestCase):
    def test_duplicate_register_numbers_keep_uniquely_located_titles_eligible(self):
        stream = make_register_pdf([
            {'code': '1', 'title': 'Cooling Water Design Criteria'},
            {'code': '1', 'title': 'Circulation Pump Specification'},
        ])
        stream.name = 'synthetic/numbered-register.pdf'
        text, _, coverage = extract_text_with_coverage(stream, 'numbered-register.pdf')
        rows = extract_register_rows(text, structured_evidence=coverage['structured_evidence'])
        self.assertEqual(len(rows), 2)
        self.assertNotEqual(rows[0]['start'], rows[1]['start'])
        source = {'id': 318, 'project_id': 917, 'filename': 'numbered-register.pdf',
                  'category': 'other', 'parse_status': 'done', 'text': text,
                  'structured_evidence': coverage['structured_evidence']}
        package = build_planning_package_sources([source], [])
        self.assertEqual({row['title'] for row in package['deliverables']},
                         {'Cooling Water Design Criteria', 'Circulation Pump Specification'})
        self.assertEqual(package['excluded_inventory'], [])
        for row in rows:
            self.assertEqual(row['source_locator']['quote'], text[row['start']:row['end']])

    def test_ambiguous_duplicate_geometry_stays_inventory_without_false_quotes_or_fallback(self):
        stream = make_register_pdf([
            {'code': '1', 'title': 'Cooling Water Design Criteria'},
            {'code': '1', 'title': 'Cooling Water Design Criteria'},
        ])
        stream.name = 'synthetic/repeated-register.pdf'
        text, _, coverage = extract_text_with_coverage(stream, 'repeated-register.pdf')
        rows = extract_register_rows(text, structured_evidence=coverage['structured_evidence'])
        self.assertEqual(len(rows), 2)
        self.assertEqual({row['source_layout'] for row in rows}, {'pdf_geometry_register'})
        self.assertNotEqual(rows[0]['source_locator']['bbox'], rows[1]['source_locator']['bbox'])
        for row in rows:
            self.assertEqual(row['title_boundary_status'], 'ambiguous')
            self.assertIsNone(row['start'])
            self.assertIsNone(row['end'])
            self.assertEqual(row['source_excerpt'], '')
            self.assertNotIn('quote', row['source_locator'])
            self.assertNotIn('character_start', row['source_locator'])
            self.assertIn('Cooling Water Design Criteria', row['source_locator']['literal_cells']['title'])
        source = {'id': 318, 'project_id': 917, 'filename': 'repeated-register.pdf',
                  'category': 'other', 'parse_status': 'done', 'text': text,
                  'structured_evidence': coverage['structured_evidence']}
        plan = build_document_plan([source], project_id=917)
        self.assertEqual(plan['activities'], [])
        self.assertEqual(len(plan['register_inventory']), 2)
        package = build_planning_package_sources([source], [])
        self.assertEqual(package['deliverables'], [])
        self.assertEqual(len(package['excluded_inventory']), 2)
        self.assertEqual(len({row['id'] for row in package['excluded_inventory']}), 2)
        self.assertTrue(all('character_start' not in row['source_references'][0]['locator'] for row in package['excluded_inventory']))

    def test_ambiguous_geometry_suppresses_matching_positive_text_fallback_rows(self):
        text = ('Table 9: Applicable Deliverables for Work Packages\n'
                'S. No. Discipline Document / Deliverable Description Work Packages Remarks\n'
                '9.4 General\n9.4.7 General Cooling Water Design Criteria X\n'
                '9.4.7 General Cooling Water Design Criteria X\n')
        fallback = extract_register_rows(text)
        self.assertEqual(len(fallback), 2)
        self.assertTrue(all(row['title_boundary_status'] == 'explicit_marked_row' for row in fallback))
        from ..services.register_geometry_cache import SCHEMA_VERSION
        geometry = []
        for index, source in enumerate(fallback):
            row = deepcopy(source)
            row.update(start=None, end=None, source_layout='pdf_geometry_register',
                       literal_cells={'title': row['original_title']})
            row['source_locator'] = {'page': 1, 'bbox': [40, 100 + index * 50, 600, 145 + index * 50],
                                     'text_row_status': 'ambiguous_repeated_item'}
            geometry.append(row)
        evidence = {'register_geometry': {'schema_version': SCHEMA_VERSION, 'extractor_version': GEOMETRY_VERSION, 'status': 'parsed',
                    'text_sha256': hashlib.sha256(text.encode()).hexdigest(), 'rows': geometry}}
        rows = extract_register_rows(text, structured_evidence=evidence)
        self.assertEqual(len(rows), 2)
        self.assertTrue(all(row['title_boundary_status'] == 'ambiguous' for row in rows))
        self.assertTrue(all(row['source_line'] is None for row in rows))
        source = {'id': 318, 'project_id': 917, 'filename': 'repeated-register.pdf',
                  'category': 'other', 'parse_status': 'done', 'text': text, 'structured_evidence': evidence}
        package = build_planning_package_sources([source], [])
        self.assertEqual(package['deliverables'], [])
        self.assertEqual(len(package['excluded_inventory']), 2)

    def test_actual_word_workbook_and_pdf_sources_share_the_package_catalog(self):
        with patch('apps.planning_intelligence.services.project_ai.call_project_ai',
                   side_effect=AssertionError('Mixed-format extraction called a provider')):
            sources = mixed_project_sources(project_id=812, first_file_id=700)
            result = build_planning_package_sources(sources, [])
        self.assertEqual({row['title'] for row in result['deliverables']}, {
            'Sediment sample register', 'Valve service | readiness "B"',
            'Wetland habitat\nsurvey dossier', 'Shared readiness record',
            'Aquifer monitoring acceptance dossier',
        })
        self.assertEqual(len(result['deliverables']), 5)
        self.assertEqual(result['excluded_inventory'], [])
        self.assertEqual({ref['file_id'] for row in result['deliverables'] for ref in row['source_references']}, {700, 701, 702})
        self.assertEqual({ref['project_id'] for row in result['deliverables'] for ref in row['source_references']}, {812})
        self.assertNotIn('register_geometry', sources[0]['structured_evidence'])
        self.assertNotIn('register_geometry', sources[1]['structured_evidence'])
        pdf = next(row for row in result['deliverables'] if row['title'] == 'Aquifer monitoring acceptance dossier')
        self.assertEqual(pdf['source_references'][0]['locator']['register_item'], '9.4.7')
        self.assertEqual(pdf['source_references'][0]['locator']['page'], 1)

    def test_two_projects_and_multiple_files_keep_distinct_scope_and_ignore_foreign_claims(self):
        first = mixed_project_sources(project_id=812, first_file_id=700)
        second = mixed_project_sources(project_id=913, first_file_id=900)
        quote = 'Additional statutory sampling dossier'
        for sources, project_id, file_id in ((first, 812, 703), (second, 913, 903)):
            sources.append({'id': file_id, 'project_id': project_id, 'filename': 'scope-notes.txt',
                            'category': 'other', 'parse_status': 'done', 'text': quote})
        claim = {'id': 37, 'fact_type': 'deliverable', 'status': 'detected', 'extraction_method': 'ai',
                 'source_file_id': 703, 'value': {'name': quote}, 'source_excerpt': quote,
                 'source_locator': {'character_start': 0, 'character_end': len(quote), 'quote': quote}}
        first_catalog = build_planning_package_sources(first, [claim])
        foreign_catalog = build_planning_package_sources(second, [claim])
        second_catalog = build_planning_package_sources(second, [{**claim, 'id': 58, 'source_file_id': 903}])
        self.assertEqual(len(first_catalog['deliverables']), 6)
        self.assertEqual(len(foreign_catalog['deliverables']), 5)
        self.assertEqual(len(second_catalog['deliverables']), 6)
        self.assertNotIn(quote, {row['title'] for row in foreign_catalog['deliverables']})
        self.assertTrue({row['id'] for row in first_catalog['deliverables']}.isdisjoint(
            {row['id'] for row in second_catalog['deliverables']}))
        for project_id, catalog in ((812, first_catalog), (913, second_catalog)):
            self.assertEqual({ref['project_id'] for row in catalog['deliverables'] for ref in row['source_references']}, {project_id})
        duplicated = {**deepcopy(first[1]), 'id': 799}
        repeated = build_planning_package_sources([first[1], duplicated], [])['deliverables']
        same_title = [row for row in repeated if row['title'] == 'Shared readiness record']
        self.assertEqual(len(same_title), 2)
        self.assertEqual(len({row['id'] for row in same_title}), 2)
        self.assertEqual({row['source_references'][0]['file_id'] for row in same_title}, {701, 799})

    def test_pdf_cache_isolation_and_renamed_source_preserve_content_only(self):
        stream = make_register_pdf([{'code': '9.4.7', 'discipline': 'Water', 'title': 'Aquifer\nmonitoring dossier'}])
        stream.name = 'uploads/project-a/report.pdf'
        text, _confidence, coverage = extract_text_with_coverage(stream, 'report.pdf')
        uploaded = SimpleNamespace(file=stream, original_filename='report.pdf', extracted_text=text,
                                   size_bytes=len(stream.getvalue()), document_profile=SimpleNamespace(extraction_coverage=coverage))
        self.assertIsNotNone(cached_register_geometry(uploaded))
        copied = BytesIO(stream.getvalue())
        copied.name = 'uploads/project-b/report.pdf'
        other = SimpleNamespace(file=copied, original_filename='report.pdf', extracted_text=text,
                                size_bytes=len(copied.getvalue()), document_profile=SimpleNamespace(extraction_coverage=deepcopy(coverage)))
        self.assertIsNone(cached_register_geometry(other))
        copied.name = 'uploads/project-b/renamed-environmental-handover.PDF'
        other.original_filename = 'renamed-environmental-handover.PDF'
        renamed_text, _, renamed_coverage = extract_text_with_coverage(copied, other.original_filename)
        other.document_profile.extraction_coverage = renamed_coverage
        self.assertEqual(renamed_text, text)
        old_rows, new_rows = register_rows_for_file(uploaded), register_rows_for_file(other)
        self.assertEqual([(row['register_item'], row['original_title'], row['applicability_status']) for row in old_rows],
                         [(row['register_item'], row['original_title'], row['applicability_status']) for row in new_rows])
        self.assertEqual(cached_register_geometry(uploaded)['file_sha256'], cached_register_geometry(other)['file_sha256'])
        self.assertNotEqual(cached_register_geometry(uploaded)['source_storage_name'], cached_register_geometry(other)['source_storage_name'])

    def test_non_pdf_and_unsaved_sources_do_not_query_for_a_geometry_profile(self):
        from ..models import PlanningFile
        for filename, adding in [('notes.txt', False), ('register.pdf', True)]:
            with self.subTest(filename=filename):
                source = PlanningFile(pk=42, original_filename=filename, file=filename)
                source._state.adding = adding
                self.assertIsNone(cached_register_geometry(source))

    def test_parse_preserves_source_text_and_keeps_full_titles_out_of_remarks(self):
        stream, text, coverage = parsed_source()
        with patch('apps.planning_intelligence.services.register_geometry_cache.build_register_geometry', return_value={}):
            without_geometry, _, _ = extract_text_with_coverage(stream, 'register.pdf')
        self.assertEqual(text, without_geometry)
        geometry = coverage['structured_evidence']['register_geometry']
        self.assertEqual(geometry['file_sha256'], hashlib.sha256(stream.getvalue()).hexdigest())
        rows = extract_register_rows(text, structured_evidence=coverage['structured_evidence'])
        self.assertEqual(len(rows), 5)
        self.assertEqual(rows[0]['name'], 'Reliability Availability and Maintenance Study Report')
        self.assertEqual([row['applicability_status'] for row in rows],
                         ['marked', 'bundled', 'conditional', 'not_required', 'not_marked'])
        for row in rows:
            locator = row['source_locator']
            self.assertEqual(locator['quote'], text[row['start']:row['end']])
            self.assertEqual(locator['page'], 1)
            self.assertEqual(len(locator['bbox']), 4)

    def test_source_and_package_projection_preserve_review_only_inventory(self):
        _stream, text, coverage = parsed_source()
        source = {'id': 1, 'project_id': 2, 'filename': 'register.pdf', 'category': 'sow',
                  'parse_status': 'done', 'text': text, 'structured_evidence': coverage['structured_evidence']}
        plan = build_document_plan([source], project_id=2)
        self.assertEqual(len(plan['register_inventory']), 5)
        self.assertEqual([row['name'] for row in plan['activities']],
                         ['Reliability Availability and Maintenance Study Report'])
        package = build_planning_package_sources([source], [])
        self.assertEqual(len(package['deliverables']), 1)
        self.assertEqual(len(package['excluded_inventory']), 4)
        self.assertIsNone(plan['activities'][0]['duration_days'])
        self.assertFalse(plan['ready_for_calculation'])

    def test_cached_rows_replace_only_matching_page_and_item(self):
        text = 'Serial No | Discipline | Document Title\n1 | Process | First title\n\fSerial No | Discipline | Document Title\n1 | Process | Second title\n'
        fallback = extract_register_rows(text)
        replacement = deepcopy(fallback[0])
        replacement.update(name='Full first title', original_title='Full first title', source_layout='pdf_geometry_register')
        replacement['source_locator'].update(page=1, bbox=[0, 0, 10, 10])
        from ..services.register_geometry_cache import SCHEMA_VERSION
        evidence = {'register_geometry': {'schema_version': SCHEMA_VERSION, 'extractor_version': GEOMETRY_VERSION, 'status': 'parsed',
                    'text_sha256': hashlib.sha256(text.encode()).hexdigest(), 'rows': [replacement]}}
        rows = extract_register_rows(text, structured_evidence=evidence)
        self.assertEqual([row['name'] for row in rows], ['Full first title', 'Second title'])
        evidence['register_geometry']['text_sha256'] = 'stale'
        self.assertEqual(extract_register_rows(text, structured_evidence=evidence), fallback)

    def test_old_schema_and_extractor_caches_fall_back_without_cross_version_reuse(self):
        stream, text, coverage = parsed_source()
        fallback = extract_register_rows(text)
        source = SimpleNamespace(file=stream, original_filename='register.pdf', extracted_text=text,
                                 size_bytes=len(stream.getvalue()))
        for key in ('schema_version', 'extractor_version'):
            with self.subTest(key=key):
                stale_coverage = deepcopy(coverage)
                stale_coverage['structured_evidence']['register_geometry'][key] = 'obsolete-version'
                source.document_profile = SimpleNamespace(extraction_coverage=stale_coverage)
                self.assertIsNone(cached_register_geometry(source))
                self.assertEqual(extract_register_rows(text, structured_evidence=stale_coverage['structured_evidence']), fallback)

    def test_short_ai_title_cannot_bypass_review_only_geometry_cell(self):
        _stream, text, coverage = parsed_source()
        source = {'id': 1, 'project_id': 2, 'filename': 'register.pdf', 'category': 'sow',
                  'parse_status': 'done', 'text': text, 'structured_evidence': coverage['structured_evidence']}
        quote = 'Emergency Response Plan'
        start = text.index(quote)
        fact = {'id': 9, 'source_file_id': 1, 'fact_type': 'deliverable', 'status': 'detected',
                'extraction_method': 'ai', 'value': {'name': 'Emergency Response'},
                'source_excerpt': quote, 'source_locator': {'character_start': start,
                    'character_end': start + len(quote), 'quote': quote}}
        result = build_planning_package_sources([source], [fact])
        self.assertFalse(any(row['title'] == 'Emergency Response' for row in result['deliverables']))
        self.assertTrue(any(row['reason'] == 'ai_claim_requires_register_review' for row in result['excluded_inventory']))

    def test_source_storage_text_and_size_mismatch_reject_cached_geometry(self):
        stream, text, coverage = parsed_source()
        uploaded = SimpleNamespace(file=stream, extracted_text=text, size_bytes=len(stream.getvalue()),
                                   document_profile=SimpleNamespace(extraction_coverage=coverage))
        self.assertIsNotNone(cached_register_geometry(uploaded))
        uploaded.extracted_text += 'changed'
        self.assertIsNone(cached_register_geometry(uploaded))
        uploaded.extracted_text = text
        stream.name = 'replacement/register.pdf'
        self.assertIsNone(cached_register_geometry(uploaded))
        stream.name = 'synthetic/register.pdf'
        uploaded.size_bytes += 1
        self.assertIsNone(cached_register_geometry(uploaded))

    def test_optional_geometry_library_failure_does_not_erase_extracted_text(self):
        stream, text, _coverage = parsed_source()
        from pdfminer.pdfparser import PDFSyntaxError
        with patch('apps.planning_intelligence.services.register_geometry_cache.build_register_geometry',
                   side_effect=PDFSyntaxError('Synthetic unreadable geometry')):
            fallback, _confidence, coverage = extract_text_with_coverage(stream, 'register.pdf')
        self.assertEqual(fallback, text)
        self.assertNotIn('register_geometry', coverage.get('structured_evidence') or {})


class RegisterGeometryAnalysisTests(DocumentIntelligenceFixture):
    def test_ambiguous_geometry_facts_persist_cell_provenance_without_false_character_offsets(self):
        stream = make_register_pdf([
            {'code': '1', 'title': 'Cooling Water Design Criteria'},
            {'code': '1', 'title': 'Cooling Water Design Criteria'},
        ])
        text, _, coverage = extract_text_with_coverage(stream, 'repeated-register.pdf')
        file_obj = self.source('repeated-register.pdf', 'other', text)
        file_obj.file.save('repeated-register.pdf', ContentFile(stream.getvalue()), save=False)
        file_obj.size_bytes = len(stream.getvalue())
        file_obj.save()
        coverage['structured_evidence']['register_geometry']['source_storage_name'] = file_obj.file.name
        profile_document(file_obj, extraction_coverage=coverage)
        run, intelligence = run_document_intelligence(self.project, user=self.owner, allow_ai=False)
        facts = list(run.facts.filter(fact_type='deliverable', extraction_method='deterministic'))
        self.assertEqual(len(facts), 2)
        self.assertEqual(intelligence['register_summary']['row_count'], 2)
        summary_rows = [row for discipline in intelligence['disciplines'].values() for row in discipline['register_rows']]
        self.assertEqual(len(summary_rows), 2)
        self.assertTrue(all(row['title_boundary_status'] == 'ambiguous' for row in summary_rows))
        for fact in facts:
            self.assertEqual(fact.status, 'detected')
            self.assertEqual(fact.value['title_boundary_status'], 'ambiguous')
            self.assertEqual(fact.source_excerpt, '')
            self.assertNotIn('character_start', fact.source_locator)
            self.assertNotIn('character_end', fact.source_locator)
            self.assertNotIn('quote', fact.source_locator)
            self.assertEqual(fact.source_locator['page'], 1)
            self.assertEqual(len(fact.source_locator['bbox']), 4)
        self.assertNotEqual(facts[0].source_locator['bbox'], facts[1].source_locator['bbox'])
        from ..services.evidence_graph import refresh_evidence_graph
        from .test_scheduling_engine import grant_planning_test_actions
        grant_planning_test_actions((self.owner,), ('read', 'create', 'update'))
        graph = refresh_evidence_graph(self.project, self.owner)
        identities = list(graph.nodes.filter(current=True, kind='fact', property='identity'))
        self.assertEqual(len(identities), 2)
        for node in identities:
            self.assertEqual(node.status, 'detected')
            self.assertFalse(node.validation['quote_verified'])
            self.assertFalse(node.sources[0]['quote_verified'])
            self.assertEqual(node.sources[0]['verbatim'], '')
            self.assertNotIn('character_start', node.sources[0]['locator'])
            self.assertEqual(node.sources[0]['locator']['literal_cells']['title'], 'Cooling Water Design Criteria')
        self.assertNotEqual(identities[0].sources[0]['locator']['bbox'], identities[1].sources[0]['locator']['bbox'])

    def saved_pdf(self, *, retain_cache=False):
        stream, text, coverage = parsed_source()
        file_obj = self.source('register.pdf', 'sow', text)
        file_obj.file.save('register.pdf', ContentFile(stream.getvalue()), save=False)
        file_obj.content_type = 'application/pdf'
        file_obj.size_bytes = len(stream.getvalue())
        file_obj.save()
        geometry = coverage['structured_evidence']['register_geometry']
        geometry['source_storage_name'] = file_obj.file.name
        if not retain_cache:
            coverage.pop('structured_evidence')
        profile_document(file_obj, extraction_coverage=coverage)
        return file_obj

    def test_analysis_refreshes_once_and_resume_retains_text_revision_and_prior_facts(self):
        file_obj = self.saved_pdf()
        before = (file_obj.extracted_text, file_obj.updated_at)
        with patch('apps.planning_intelligence.services.pdf_register_geometry.extract_pdf_register_rows', wraps=extract_pdf_register_rows) as extractor:
            run, intelligence = run_document_intelligence(self.project, user=self.owner, allow_ai=False)
        self.assertEqual(extractor.call_count, 1)
        facts = list(run.facts.filter(fact_type='deliverable').order_by('pk').values('id', 'value', 'status', 'source_locator'))
        self.assertEqual(len(facts), 5)
        self.assertEqual(facts[0]['value']['name'], 'Reliability Availability and Maintenance Study Report')
        for fact in facts:
            locator = fact['source_locator']
            self.assertEqual(locator['quote'], file_obj.extracted_text[locator['character_start']:locator['character_end']])
            self.assertEqual(fact['status'], 'detected')
        self.assertEqual(intelligence['register_summary']['row_count'], 5)
        run.summary['ai_checkpoint'] = {'saved': 'unchanged checkpoint identity'}
        run.save(update_fields=['summary'])
        from ..services.document_intelligence import analyze_project
        with patch('apps.planning_intelligence.services.pdf_register_geometry.extract_pdf_register_rows', wraps=extract_pdf_register_rows) as extractor, \
                patch('apps.planning_intelligence.services.document_intelligence.analyze_project', wraps=analyze_project) as analyze:
            resumed, _ = run_document_intelligence(self.project, user=self.owner, allow_ai=False, resume_run=run)
        self.assertEqual(extractor.call_count, 0)
        self.assertEqual(analyze.call_args.kwargs['resume_state'], run.summary['ai_checkpoint'])
        self.assertNotEqual(resumed.pk, run.pk)
        self.assertEqual(facts, list(run.facts.filter(fact_type='deliverable').order_by('pk').values('id', 'value', 'status', 'source_locator')))
        file_obj.refresh_from_db()
        self.assertEqual((file_obj.extracted_text, file_obj.updated_at), before)

    def test_get_projections_do_not_read_storage_or_write_and_coverage_does_not_leak_rows(self):
        file_obj = self.saved_pdf(retain_cache=True)
        with patch('django.db.models.fields.files.FieldFile.open', side_effect=AssertionError('GET opened original storage')):
            with CaptureQueriesContext(connection) as queries:
                sources = source_files(self.project)
                rows = register_rows_for_file(file_obj)
                profile = file_obj.document_profile
                serialized = DocumentProfileSerializer(profile).data
                coverage = summarize_coverage([file_obj], analyzed_file_ids=[file_obj.pk])
        self.assertEqual(len(rows), 5)
        self.assertEqual(len(sources[0]['structured_evidence']['register_geometry']['rows']), 5)
        self.assertFalse(any(query['sql'].lstrip().upper().startswith(('UPDATE ', 'INSERT ', 'DELETE ')) for query in queries))
        for public in (serialized['extraction_coverage'], coverage['files'][0], file_coverage(file_obj, profile.extraction_coverage)):
            self.assertNotIn('structured_evidence', public)
            self.assertEqual(public['structured_evidence_summary']['register_geometry']['row_count'], 5)
        profile_document(file_obj)
        profile.refresh_from_db()
        self.assertEqual(len(profile.extraction_coverage['structured_evidence']['register_geometry']['rows']), 5)

    def test_explicit_refresh_rehashes_binary_even_when_storage_name_and_text_match(self):
        file_obj = self.saved_pdf(retain_cache=True)
        cached = cached_register_geometry(file_obj)
        altered = make_register_pdf([{'code': '3.2.1', 'title': 'Other register'}]).getvalue()
        # Deliberately simulate storage replacement without a record revision.
        # The new bytes cannot reuse the old cached geometry by name alone.
        with file_obj.file.storage.open(file_obj.file.name, 'wb') as target:
            target.write(altered)
        with patch('apps.planning_intelligence.services.pdf_register_geometry.extract_pdf_register_rows', wraps=extract_pdf_register_rows) as extractor:
            with self.assertRaises(RegisterGeometrySourceChanged):
                ensure_register_geometry(file_obj)
        self.assertEqual(extractor.call_count, 0)
        file_obj.document_profile.refresh_from_db()
        self.assertEqual(file_obj.document_profile.extraction_coverage['structured_evidence']['register_geometry'], cached)

    def test_valid_cache_refresh_changes_generation_identity_without_source_revision(self):
        file_obj = self.saved_pdf()
        run, _ = run_document_intelligence(self.project, user=self.owner, allow_ai=False)
        profile = file_obj.document_profile
        profile.refresh_from_db()
        geometry = profile.extraction_coverage['structured_evidence'].pop('register_geometry')
        profile.save(update_fields=['extraction_coverage'])
        from ..services.operational_jobs import generation_fingerprint
        request = {'generation_options': {'mode': 'planning_package', 'intelligence_run_id': run.pk}}
        before = generation_fingerprint(self.project, request)
        source_revision = file_obj.updated_at
        ensure_register_geometry(file_obj)
        after = generation_fingerprint(self.project, request)
        self.assertNotEqual(before, after)
        file_obj.refresh_from_db()
        self.assertEqual(file_obj.updated_at, source_revision)
        self.assertEqual(cached_register_geometry(file_obj)['file_sha256'], geometry['file_sha256'])

    def test_analysis_job_identity_changes_with_geometry_version_but_ai_checkpoint_does_not(self):
        from ..services.document_intelligence import _extraction_source_manifest
        from ..services.intelligence import _augment_with_ai
        from ..services.operational_jobs import operation_fingerprint
        from ..services.preview_confirmation import source_fingerprint

        file_obj = self.saved_pdf(retain_cache=True)
        source_state = (file_obj.extracted_text, file_obj.updated_at, file_obj.size_bytes, file_obj.file.name)
        manifest = _extraction_source_manifest([file_obj])
        source_identity = source_fingerprint(self.project)
        initial_ai, revised_ai = {'notes': []}, {'notes': []}
        with patch('apps.planning_intelligence.services.project_ai.get_project_ai_config',
                   return_value={'provider': 'claude', 'model': 'isolated-test-model'}), \
                patch('apps.planning_intelligence.services.project_ai.project_provider', return_value='claude'), \
                patch('apps.planning_intelligence.services.project_ai.call_project_ai',
                      side_effect=AssertionError('Fingerprint verification must never call a provider')), \
                CaptureQueriesContext(connection) as queries:
            original = operation_fingerprint(self.project, 'analyze', {})
            self.assertEqual(operation_fingerprint(self.project, 'analyze', {}), original)
            _augment_with_ai(initial_ai, [file_obj], self.project, self.owner, allow_requests=False)
            with patch('apps.planning_intelligence.services.pdf_register_geometry.GEOMETRY_VERSION',
                       GEOMETRY_VERSION + '-next-test-version'):
                changed = operation_fingerprint(self.project, 'analyze', {})
                self.assertNotEqual(changed, original)
                self.assertEqual(operation_fingerprint(self.project, 'analyze', {}), changed)
                _augment_with_ai(revised_ai, [file_obj], self.project, self.owner,
                                 allow_requests=False, resume_state=initial_ai['ai_checkpoint'])
                self.assertEqual(source_fingerprint(self.project), source_identity)
                self.assertEqual(_extraction_source_manifest([file_obj]), manifest)
            self.assertEqual(operation_fingerprint(self.project, 'analyze', {}), original)
        self.assertEqual(revised_ai['ai_checkpoint'], initial_ai['ai_checkpoint'])
        self.assertGreater(revised_ai['ai_processing_coverage']['chunks_total'], 0)
        self.assertEqual(revised_ai['ai_processing_coverage']['calls_this_pass'], 0)
        self.assertFalse(any(query['sql'].lstrip().upper().startswith(('UPDATE ', 'INSERT ', 'DELETE ')) for query in queries))
        file_obj.refresh_from_db()
        self.assertEqual((file_obj.extracted_text, file_obj.updated_at, file_obj.size_bytes, file_obj.file.name), source_state)

    def test_unavailable_pdf_geometry_preserves_existing_text_fallback(self):
        file_obj = self.source('absent.pdf', 'sow', 'Serial No | Discipline | Document Title\n1 | Process | Existing title\n')
        self.assertIsNone(ensure_register_geometry(file_obj))
        run, _ = run_document_intelligence(self.project, user=self.owner, allow_ai=False)
        self.assertEqual(run.facts.get(fact_type='deliverable').value['name'], 'Existing title')
