"""Generic extraction must expose omissions and never invent missing scope."""
import io
import json
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from django.test import SimpleTestCase, TestCase

from ..models import DocumentIntelligenceRun, IntelligenceFact, PlanningFile, PlanningProject
from ..services.document_intelligence import _extract_file_facts, _persist_ai_facts, run_document_intelligence
from ..services.extraction_coverage import file_coverage
from ..services.intelligence import analyze_project, source_chunks
from ..services.parsers import extract_text_with_coverage
from ..services.register_rows import extract_register_rows
from ..services.schedule_basis import _deliverable_rows
from ..services.simple_schedule_proposal import source_constraints
from django.utils import timezone


class GenericParserCoverageTests(SimpleTestCase):
    def test_csv_quoted_title_and_empty_columns_survive_parsing(self):
        source = b'Serial No,Discipline,Document No,Document Title,Revision\n1,PROCUREMENT,,"Vendor, \"\"A\"\"\nDossier",\n'
        text, _, coverage = extract_text_with_coverage(io.BytesIO(source), 'purchase-register.csv')
        rows = extract_register_rows(text)
        self.assertEqual(rows[0]['original_title'], 'Vendor, "A"\nDossier')
        self.assertEqual(rows[0]['document_number'], '')
        self.assertEqual(coverage['status'], 'complete')
        self.assertFalse(coverage['semantic_coverage_verified'])

    def test_utf16_text_is_not_silently_discarded(self):
        expected = 'Contractor shall issue final safety notes.\n\u0645\u0631\u0627\u062c\u0639\u0629'
        text, _, coverage = extract_text_with_coverage(io.BytesIO(expected.encode('utf-16')), 'requirements.txt')
        self.assertEqual(text, expected)
        self.assertEqual(coverage['status'], 'complete')

    def test_existing_markdown_upload_format_keeps_source_text(self):
        text, _, coverage = extract_text_with_coverage(io.BytesIO(b'# Requirements\nSupplier must provide test report.'), 'scope.md')
        self.assertIn('Supplier must provide test report.', text)
        self.assertEqual(coverage['status'], 'complete')

    def test_invalid_utf8_reports_replaced_characters(self):
        text, _, coverage = extract_text_with_coverage(io.BytesIO(b'Issue the report \xff.'), 'notes.txt')
        self.assertIn('\ufffd', text)
        self.assertEqual(coverage['status'], 'partial')
        self.assertIn('text_decode_replacement', [item['code'] for item in coverage['issues']])

    def test_unknown_binary_and_legacy_xls_do_not_become_text(self):
        for filename in ('source.bin', 'book.xls', 'plan.mpp'):
            with self.subTest(filename=filename):
                text, confidence, coverage = extract_text_with_coverage(io.BytesIO(b'\x00Binary scope'), filename)
                self.assertEqual((text, confidence), ('', 0))
                self.assertEqual(coverage['status'], 'unsupported')

    def test_text_retention_limit_is_explicit(self):
        with patch('apps.planning_intelligence.services.parsers.MAX_EXTRACTED_CHARS', 20):
            text, _, coverage = extract_text_with_coverage(io.BytesIO(b'A' * 40), 'large.txt')
        self.assertTrue(text.endswith('...[truncated]'))
        self.assertEqual(coverage['characters_extracted'], 40)
        self.assertTrue(coverage['text_truncated'])
        self.assertEqual(coverage['status'], 'partial')

    def test_all_workbook_sheets_including_hidden_are_covered(self):
        import openpyxl
        stream = io.BytesIO()
        workbook = openpyxl.Workbook()
        workbook.active.title = 'Procurement'
        workbook.active.append(['Document Title', 'Department'])
        workbook.active.append(['Supplier assessment', 'Commercial'])
        sheet = workbook.create_sheet('Hidden constraints')
        sheet.sheet_state = 'hidden'
        sheet.append(['Must obtain approval before purchase'])
        workbook.save(stream)
        text, _, coverage = extract_text_with_coverage(stream, 'multi-sheet.xlsx')
        self.assertIn('Supplier assessment', text)
        self.assertIn('Must obtain approval', text)
        self.assertEqual(coverage['units_total'], 2)
        self.assertEqual(coverage['units_processed'], 2)
        self.assertEqual(coverage['units'][1]['sheet'], 'Hidden constraints')
        self.assertTrue(coverage['units'][1]['hidden'])
        self.assertEqual(coverage['status'], 'partial')  # Non-cell content is not certified.

    def test_docx_keeps_body_order_and_blank_table_cells(self):
        import docx
        document = docx.Document()
        document.add_paragraph('Start of required work')
        table = document.add_table(rows=2, cols=3)
        table.cell(0, 0).text = 'Document No'
        table.cell(0, 1).text = 'Document Title'
        table.cell(0, 2).text = 'Revision'
        table.cell(1, 1).text = 'Constructability | review'
        document.add_paragraph('End constraint: complete after review')
        stream = io.BytesIO()
        document.save(stream)
        text, _, coverage = extract_text_with_coverage(stream, 'construction.docx')
        self.assertLess(text.index('Start of required work'), text.index('--- Table: 1 ---'))
        self.assertLess(text.index('--- Table: 1 ---'), text.index('End constraint'))
        self.assertIn('|"Constructability | review"|', text)
        self.assertEqual(coverage['units_total'], 3)
        self.assertEqual(coverage['status'], 'partial')

    def test_mixed_pdf_ocr_runs_on_scanned_page_not_only_empty_document(self):
        pdf = MagicMock()
        pdf.__enter__.return_value.pages = [
            SimpleNamespace(extract_text=lambda: 'A readable engineering requirement.'),
            SimpleNamespace(extract_text=lambda: ''),
            SimpleNamespace(extract_text=lambda: 'Final visible constraint.'),
        ]
        with patch('pdfplumber.open', return_value=pdf), patch('apps.planning_intelligence.services.parsers._ocr_pdf_page', return_value='Scanned instruction') as ocr:
            text, _, coverage = extract_text_with_coverage(io.BytesIO(b'%PDF'), 'mixed.pdf')
        ocr.assert_called_once()
        self.assertEqual(ocr.call_args.args[1], 1)
        self.assertIn('Scanned instruction', text)
        self.assertEqual(coverage['units_total'], 3)
        self.assertEqual(coverage['units_processed'], 3)
        self.assertEqual(text.count('\f'), 2)

    def test_ocr_limit_marks_each_unprocessed_page(self):
        pdf = MagicMock()
        pdf.__enter__.return_value.pages = [SimpleNamespace(extract_text=lambda: '') for _ in range(3)]
        with patch('pdfplumber.open', return_value=pdf), patch('apps.planning_intelligence.services.parsers.MAX_OCR_PAGES', 1), patch('apps.planning_intelligence.services.parsers._ocr_pdf_page', return_value='Scanned source'):
            _, _, coverage = extract_text_with_coverage(io.BytesIO(b'%PDF'), 'scan.pdf')
        self.assertEqual(coverage['units_total'], 3)
        self.assertEqual(coverage['units_processed'], 1)
        self.assertEqual(coverage['units_skipped'], 2)
        self.assertEqual(coverage['status'], 'partial')

    def test_legacy_text_does_not_claim_complete_extraction(self):
        source = PlanningFile(id=1, original_filename='old.pdf', parse_status='done', extracted_text='Stored text')
        self.assertEqual(file_coverage(source)['status'], 'unknown')

    def test_multipage_tiff_covers_each_frame(self):
        from PIL import Image
        stream = io.BytesIO()
        first = Image.new('RGB', (20, 20), 'white')
        second = Image.new('RGB', (20, 20), 'white')
        first.save(stream, format='TIFF', save_all=True, append_images=[second])
        with patch('pytesseract.image_to_string', side_effect=['First page requirement', 'Second page constraint']):
            text, _, coverage = extract_text_with_coverage(stream, 'scanned-contract.tiff')
        self.assertIn('Second page constraint', text)
        self.assertEqual(coverage['units_total'], 2)
        self.assertEqual(coverage['units_processed'], 2)


class GenericIntelligenceEvidenceTests(SimpleTestCase):
    def source(self, text, category='other', pk=1):
        return PlanningFile(id=pk, original_filename='test.txt', category=category, parse_status='done', extracted_text=text)

    def test_absent_scope_does_not_invent_catalogue_deliverables_or_hse(self):
        result = analyze_project([self.source('Project Name: Internal Operations')], allow_ai=False)
        self.assertFalse(any(group['in_scope'] for group in result['disciplines'].values()))
        self.assertFalse(any(group['deliverables'] for group in result['disciplines'].values()))
        self.assertEqual(result['hse_studies'], [])
        self.assertIsNone(result['detected_duration_months'])

    def test_topic_mentions_never_create_catalogue_scope_or_required_studies(self):
        for text in (
            'Training topics: HAZOP, ENVID, SIL and Piping Material Specification.',
            'Reference index: P&ID, Process Flow Diagram, Equipment List and Electrical Engineering.',
            'Examples only: Civil Engineering foundation drawings and procurement vendor data.',
            'HAZOP Study is excluded. Piping Material Specification is not required.',
        ):
            with self.subTest(text=text):
                source = self.source(text, category='sow')
                result = analyze_project([source], allow_ai=False)
                self.assertEqual(result['disciplines'], {})
                self.assertEqual(result['hse_studies'], [])
                self.assertEqual(result['available_hse_studies'], [])
                rows = {'run': DocumentIntelligenceRun(id=1), 'facts': [], '_seen': set()}
                _extract_file_facts(rows, source)
                self.assertFalse(any(fact.fact_type in {'deliverable', 'hse_study', 'discipline'} for fact in rows['facts']))

    def test_explicit_narrative_requirement_remains_verbatim_without_catalogue_expansion(self):
        text = 'The Contractor shall prepare the P&ID and HAZOP Study.'
        source = self.source(text, category='sow')
        rows = {'run': DocumentIntelligenceRun(id=1), 'facts': [], '_seen': set()}
        _extract_file_facts(rows, source)
        requirement = next(fact for fact in rows['facts'] if fact.fact_type == 'requirement')
        self.assertEqual(requirement.value, text)
        self.assertIn(text, requirement.source_excerpt)
        self.assertEqual(requirement.source_locator['line'], 1)
        self.assertFalse(any(fact.fact_type in {'deliverable', 'hse_study'} for fact in rows['facts']))

    def test_explicit_register_scope_is_retained_beside_unstructured_topic_mentions(self):
        register = self.source('Document Number|Document Title|Department\nQA-1|Readiness assessment|Assurance\n', category='mdr')
        topics = self.source('Background: HAZOP, P&ID, Piping Material Specification.', pk=2)
        result = analyze_project([register, topics], allow_ai=False)
        self.assertEqual(result['deliverable_source'], 'register')
        self.assertEqual(result['disciplines']['assurance']['deliverables'], ['Readiness assessment'])
        self.assertEqual(result['hse_studies'], [])
        self.assertEqual(result['register_summary']['row_count'], 1)

    def test_warranty_is_not_project_duration(self):
        result = analyze_project([self.source('Warranty duration is 24 months.\nWarranty shall remain valid for 24 months.')], allow_ai=False)
        self.assertIsNone(result['detected_duration_months'])

    def test_register_content_is_recognized_without_filename_or_category_inference(self):
        source = self.source('Document Title|Department\nSupplier Assessment|Commercial\n', category='other')
        result = analyze_project([source], allow_ai=False)
        self.assertEqual(result['deliverable_source'], 'register')
        rows = {'run': DocumentIntelligenceRun(id=1), 'facts': [], '_seen': set()}
        _extract_file_facts(rows, source, include_catalogue_deliverables=False)
        facts = [item for item in rows['facts'] if item.fact_type == 'deliverable']
        self.assertEqual(facts[0].value['original_title'], 'Supplier Assessment')
        self.assertTrue(facts[0].value['source_register'])

    def test_requirements_after_first_hundred_are_retained(self):
        source = self.source('\n'.join(f'The contractor shall deliver requirement {index}.' for index in range(165)))
        rows = {'run': DocumentIntelligenceRun(id=1), 'facts': [], '_seen': set()}
        _extract_file_facts(rows, source)
        requirements = [fact for fact in rows['facts'] if fact.fact_type == 'requirement']
        self.assertEqual(len(requirements), 165)
        self.assertIn('requirement 164', requirements[-1].value)

    def test_constraints_never_borrow_nearby_numbers_or_invent_survey_links(self):
        result = source_constraints([{'id': 1, 'filename': 'anything.txt', 'category': 'other', 'text': (
            'Company review\nUnrelated access period 25 working days\n'
            'Review\n36 working days\n'
            'Weeks from award\n1. 24\n'
            'Site survey documents must be submitted prior to commencement of design.\n'
            'Client review requires 10 working days.\n'
            'Delivery is 8 weeks after purchase order receipt.'
        )}])
        self.assertEqual([item['value'] for item in result if item['kind'] == 'review_days'], [10])
        self.assertEqual([item['value'] for item in result if item['kind'] == 'relative_weeks'], [8])
        self.assertFalse(any(item['kind'] == 'sequence' for item in result))
        self.assertTrue(all(not item['executable'] for item in result))
        self.assertTrue(any('Site survey' in item['value'] for item in result if item['kind'] == 'constraint_candidate'))

    def test_wrapped_explicit_review_sentence_keeps_its_quote(self):
        text = ('Allow two (02) weeks (10 working days) time duration required\n'
                'by COMPANY for the review of submitted technical documents.')
        result = source_constraints([{'id': 1, 'filename': 'clauses.txt', 'category': 'other', 'text': text}])
        review = next(item for item in result if item['kind'] == 'review_days')
        self.assertEqual(review['value'], 10)
        reference = review['source_references'][0]
        locator = reference['locator']
        self.assertEqual(reference['excerpt'], text[locator['character_start']:locator['character_end']])
        self.assertFalse(review['applicability_verified'])

    def test_chunks_cover_every_character_and_file_without_overlap_or_omissions(self):
        files = [self.source('A' * 27 + '\n' + 'B' * 44), self.source('Last source requirement', pk=2)]
        chunks = list(source_chunks(files, max_chars=17))
        for source in files:
            selected = [chunk for chunk in chunks if chunk['source_file_id'] == source.pk]
            self.assertEqual(''.join(chunk['text'] for chunk in selected), source.extracted_text)
            self.assertEqual(selected[-1]['character_end'], len(source.extracted_text))

    @patch('apps.planning_intelligence.services.claude_client.get_claude_config', return_value={'configured': True})
    @patch('apps.planning_intelligence.services.claude_client.call_claude')
    def test_ai_claims_require_real_file_quote_and_value(self, call, _config):
        source = self.source('Project Name: Site Extension\nContractor shall prepare Permit Matrix.')
        call.return_value = {'text': json.dumps({'facts': [
            {'type': 'project_name', 'value': 'Invented name', 'source_file_id': 1, 'quote': 'Project Name: Site Extension'},
            {'type': 'deliverable', 'value': 'Permit Matrix', 'source_file_id': 99, 'quote': 'Contractor shall prepare Permit Matrix.'},
            {'type': 'duration_months', 'value': 6, 'source_file_id': 1, 'quote': 'Project duration: 6 months'},
            {'type': 'deliverable', 'value': 'Permit Matrix', 'source_file_id': 1, 'quote': 'Contractor shall prepare Permit Matrix.', 'discipline': 'Construction'},
        ]})}
        result = analyze_project([source])
        self.assertEqual(result['ai_processing_coverage']['rejected_claim_count'], 3)
        self.assertEqual(len(result['ai_evidence_facts']), 1)
        self.assertIsNone(result['ai_evidence_facts'][0]['discipline'])
        self.assertIn('Permit Matrix', result['disciplines']['not_specified']['deliverables'])
        rows = {'run': DocumentIntelligenceRun(id=1), 'facts': [], '_seen': set()}
        _persist_ai_facts(rows, result, [source])
        self.assertEqual(rows['facts'][0].source_file_id, 1)
        self.assertIn('Permit Matrix', rows['facts'][0].source_excerpt)

    @patch.dict('os.environ', {'PLANNING_AI_MAX_CHUNKS': '2'})
    @patch('apps.planning_intelligence.services.intelligence.CLAUDE_MAX_INPUT_CHARS', 20)
    @patch('apps.planning_intelligence.services.claude_client.get_claude_config', return_value={'configured': True})
    @patch('apps.planning_intelligence.services.claude_client.call_claude', return_value={'text': '{"facts": []}'})
    def test_ai_budget_never_claims_complete_coverage(self, call, _config):
        result = analyze_project([self.source('X' * 75)])
        coverage = result['ai_processing_coverage']
        self.assertEqual(call.call_count, 2)
        self.assertEqual(coverage['chunks_total'], 4)
        self.assertEqual(coverage['chunks_skipped'], 2)
        self.assertEqual(coverage['status'], 'partial')
        self.assertEqual(coverage['characters_processed'], 40)

    @patch('apps.planning_intelligence.services.claude_client.get_claude_config', return_value={'configured': True})
    @patch('apps.planning_intelligence.services.claude_client.call_claude', return_value={'text': '{"facts": "not an array"}'})
    def test_invalid_ai_shape_is_reported_failed(self, _call, _config):
        result = analyze_project([self.source('Required project work')])
        self.assertEqual(result['ai_processing_coverage']['chunks_failed'], 1)
        self.assertEqual(result['ai_processing_coverage']['status'], 'partial')

    @patch('apps.planning_intelligence.services.claude_client.get_claude_config', return_value={'configured': True})
    @patch('apps.planning_intelligence.services.claude_client.call_claude', return_value={'text': '{"facts": []}', 'stop_reason': 'max_tokens'})
    def test_ai_output_limit_is_partial_even_with_valid_json(self, _call, _config):
        result = analyze_project([self.source('Required project work')])
        coverage = result['ai_processing_coverage']
        self.assertEqual(coverage['chunks_partial'], 1)
        self.assertEqual(coverage['chunks_processed'], 0)
        self.assertEqual(coverage['status'], 'partial')


class PersistedExtractionCoverageTests(TestCase):
    def test_analysis_records_unparsed_uploads_and_unknown_legacy_coverage(self):
        project = PlanningProject.objects.create(name='Coverage project')
        first = PlanningFile.objects.create(project=project, file='scope.txt', original_filename='scope.txt', parse_status='done', extracted_text='Project Name: Coverage project')
        second = PlanningFile.objects.create(project=project, file='pending.pdf', original_filename='pending.pdf', parse_status='pending')
        run, result = run_document_intelligence(project, allow_ai=False)
        coverage = result['processing_coverage']
        self.assertEqual(coverage['file_count'], 2)
        self.assertEqual(coverage['analyzed_file_count'], 1)
        self.assertEqual(coverage['status'], 'partial')
        rows = {row['file_id']: row for row in coverage['files']}
        self.assertEqual(rows[first.pk]['status'], 'unknown')
        self.assertFalse(rows[second.pk]['included_in_analysis'])
        self.assertEqual(run.summary['processing_coverage'], coverage)

    def test_generic_basis_preserves_similar_titles_and_source_identity(self):
        project = PlanningProject.objects.create(name='Generic scope')
        source = PlanningFile.objects.create(project=project, file='notes.txt', original_filename='notes.txt', parse_status='done')
        run = DocumentIntelligenceRun.objects.create(project=project, status='succeeded', started_at=timezone.now(), summary={'base_intelligence': {'document_driven': True}})
        for index, name in enumerate(['Foundation layout Area A', 'Foundation layout Area B', 'P&IDs']):
            IntelligenceFact.objects.create(run=run, source_file=source, fact_type='deliverable', key=str(index), value={'name': name, 'original_title': name, 'discipline': 'civil'}, normalized_value=name.lower(), source_locator={'line': index + 1}, status='confirmed')
        rows = _deliverable_rows(run)
        self.assertEqual([row['canonical_name'] for row in rows], ['Foundation layout Area A', 'Foundation layout Area B', 'P&IDs'])
        self.assertEqual(rows[2]['references'][0]['locator'], {'line': 3})
        self.assertEqual(len({row['source_identity'] for row in rows}), 3)

    def test_confirmed_grounded_ai_deliverable_survives_register_path(self):
        project = PlanningProject.objects.create(name='Generic mixed sources')
        source = PlanningFile.objects.create(project=project, file='notes.txt', original_filename='notes.txt', parse_status='done')
        run = DocumentIntelligenceRun.objects.create(project=project, status='succeeded', started_at=timezone.now(), summary={'base_intelligence': {'document_driven': True, 'deliverable_source': 'register'}})
        fact = IntelligenceFact.objects.create(run=run, source_file=source, fact_type='deliverable', extraction_method='ai', key='permit', value={'name': 'Permit Matrix', 'original_title': 'Permit Matrix', 'discipline': 'not_specified'}, normalized_value='permit matrix', source_locator={'line': 7}, source_excerpt='Contractor shall prepare Permit Matrix.', status='confirmed')
        IntelligenceFact.objects.create(run=run, fact_type='deliverable', extraction_method='ai', key='invented', value={'name': 'Invented work'}, normalized_value='invented work', status='confirmed')
        rows = _deliverable_rows(run)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['original_title'], 'Permit Matrix')
        self.assertEqual(rows[0]['fact_ids'], [fact.pk])
