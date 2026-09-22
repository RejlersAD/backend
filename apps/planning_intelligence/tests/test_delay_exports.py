"""Lossless review-package exports without database or business-data access."""
from copy import deepcopy
from io import BytesIO
import json
from types import SimpleNamespace
from unittest.mock import patch

from django.test import SimpleTestCase
from openpyxl import load_workbook
from rest_framework.test import APIRequestFactory, force_authenticate

from ..delay_views import DelayCaseExportView, _cell


def review_state():
    return {
        'baseline': {'id': 11, 'name': 'Approved baseline'},
        'method_boundary': 'Prospective sensitivity; no contractual entitlement.',
        'case': {
            'id': 3, 'baseline_id': 11, 'name': 'Delay review', 'revision': 2,
            'status': 'calculated', 'source_stale': False,
            'reference_report_id': 17, 'reference_data_date': '2026-09-21',
            'issues': [], 'recommendation': {}, 'recommendation_assessment': {},
            'changes': [], 'scenarios': [],
            'run': {
                'id': 8, 'fingerprint': 'a' * 64, 'events': [],
                'result': {
                    'impact': {'status': 'available', 'forecast_finish': '2026-12-20',
                               'paths': {'nodes': [], 'edges': [], 'witnesses': []},
                               'affected_activities': [], 'affected_milestones': []},
                    'scenarios': [], 'issues': [], 'limitations': ['No baseline changes.'],
                },
            },
        },
    }


class DelayExportTests(SimpleTestCase):
    def export(self, state, kind):
        actor = SimpleNamespace(pk=7, is_authenticated=True)
        project = SimpleNamespace(pk=5)
        request = APIRequestFactory().get('/delay-analysis/cases/3/export/', {'format': kind})
        force_authenticate(request, user=actor)
        with patch('apps.planning_intelligence.delay_views.accessible_projects', return_value=object()), \
             patch('apps.planning_intelligence.delay_views.get_object_or_404', return_value=project), \
             patch('apps.planning_intelligence.delay_views.delay_state', return_value=state):
            return DelayCaseExportView.as_view()(request, project_id=project.pk, case_id=3)

    def archive_package(self, workbook):
        archive = workbook['Complete JSON package']
        chunks = []
        for index, row in enumerate(archive.iter_rows(min_row=2), 1):
            self.assertEqual(row[0].value, index)
            self.assertEqual(row[1].data_type, 's')
            self.assertLessEqual(len(row[1].value), 29000)
            chunks.append(row[1].value)
        return json.loads(''.join(chunks))

    def test_large_paths_and_event_evidence_roundtrip_exactly(self):
        state = review_state()
        paths = state['case']['run']['result']['impact']['paths']
        paths['nodes'] = [
            {'activity_id': index, 'external_id': f'ACT-{index}',
             'name': 'Discipline deliverable / source trace ' * 4}
            for index in range(1, 1101)
        ]
        paths['edges'] = [
            {'predecessor_id': index, 'successor_id': index + 1, 'type': 'FS', 'lag_days': '0'}
            for index in range(1, 1100)
        ]
        paths['witnesses'] = [{'activity_ids': list(range(1, 1101)), 'type': 'driving'}]
        evidence = [{'reference': f'Source {index}: ' + 'Quoted passage ' * 280,
                     'fact': {'sources': [{'locator': {'page': index + 1}, 'quote': 'Exact \u0394 / \u5de5\u7a0b text'}]}}
                    for index in range(50)]
        state['case']['run']['events'] = [{'id': 1, 'revision': 1, 'title': 'Recorded event', 'evidence': evidence}]
        original = deepcopy(state)

        response = self.export(state, 'xlsx')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Cache-Control'], 'private, no-store')
        self.assertTrue(response['Content-Disposition'].endswith('delay-case-3-run-8.xlsx"'))
        book = load_workbook(BytesIO(response.content))
        package = self.archive_package(book)
        expected = json.loads(self.export(state, 'json').content)
        self.assertEqual(package, expected)
        self.assertEqual(len(package['case']['run']['result']['impact']['paths']['nodes']), 1100)
        self.assertEqual(package['case']['run']['events'][0]['evidence'], evidence)
        self.assertIn('complete JSON package sheet', book['Paths']['B2'].value)
        event_headers = [cell.value for cell in book['Events'][1]]
        evidence_cell = book['Events'].cell(row=2, column=event_headers.index('evidence') + 1)
        self.assertIn('complete JSON package sheet', evidence_cell.value)
        self.assertEqual(state, original)

    def test_formulas_and_control_characters_are_safe_and_lossless(self):
        state = review_state()
        state['case']['name'] = '=HYPERLINK("https://example.invalid","untrusted")'
        state['case']['run']['events'] = [
            {'id': 1, 'revision': 1, 'title': '\t+SUM(1,2)',
             'description': 'Exact NUL\x00 and SOH\x01 and VT\x0b, normal\nline and \u5de5\u7a0b.',
             'evidence': [{'reference': '@SUM(1,2)\x02 signed source'}]},
        ]
        state['case']['changes'] = [{'event_id': 1, 'reason': '  -1+2', 'evidence': '=1+1'}]
        response = self.export(state, 'xlsx')
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.content))
        expected = json.loads(self.export(state, 'json').content)
        self.assertEqual(self.archive_package(book), expected)
        for sheet in book:
            for row in sheet:
                for cell in row:
                    self.assertNotEqual(cell.data_type, 'f', f'Executable formula in {sheet.title}!{cell.coordinate}')
        event_headers = [cell.value for cell in book['Events'][1]]
        title = book['Events'].cell(row=2, column=event_headers.index('title') + 1)
        description = book['Events'].cell(row=2, column=event_headers.index('description') + 1)
        self.assertTrue(title.value.startswith("'"))
        self.assertIn('NUL\\u0000', description.value)
        self.assertIn('SOH\\u0001', description.value)
        self.assertIn('normal\nline', description.value)

    def test_readable_cell_formula_guard_handles_leading_whitespace(self):
        for value in ['=1+1', ' +SUM(1,2)', '\t-1+2', '\n@SUM(1,2)']:
            with self.subTest(value=value):
                self.assertEqual(_cell(value), "'" + value)
        self.assertEqual(_cell(-7), -7)
        self.assertEqual(_cell('Scope & review'), 'Scope & review')

    def test_export_requires_a_calculated_case_and_supported_format(self):
        state = review_state()
        state['case']['run'] = None
        self.assertEqual(self.export(state, 'xlsx').status_code, 400)
        # DRF rejects unsupported renderer suffixes before the view handler.
        self.assertEqual(self.export(review_state(), 'csv').status_code, 404)
