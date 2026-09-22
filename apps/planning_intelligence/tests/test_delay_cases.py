"""Delay scenarios retain published evidence and require independent review."""
from copy import deepcopy
from datetime import date
from io import BytesIO
import json
from uuid import uuid4

from django.core.exceptions import ValidationError as ModelValidationError
from django.db import DatabaseError, connection, transaction
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIRequestFactory, force_authenticate

from apps.core.project_models import ProjectMember
from apps.project_control.models import CostLedgerEntry
from apps.rbac.action_policy import module_action_allowed
from apps.rbac.route_guard import ModuleActionGuardMixin
from apps.users.models import User
from ..delay_views import DelayAnalysisView
from ..models import (DelayAnalysisCase, DelayAnalysisRun, DelayEvent, EvidenceDocumentVersion, EvidenceGraph,
    EvidenceNode, GovernanceItem, OperationalControlReport, PlanningFile, PlanningProject, PlanningRiskRecord,
    ScheduleBaseline)
from . import test_operational_controls as operational_fixture
from .test_scheduling_engine import grant_planning_test_actions


class DelayCaseTests(TestCase):
    def setUp(self):
        # Compose the existing approved-report fixture within this test's own DB
        # transaction; do not inherit and rerun all of its unrelated test methods.
        self.ops = operational_fixture.OperationalControlTests(methodName='runTest')
        self.ops.setUp()
        self.client = self.ops.client
        self.owner, self.manager, self.outsider = self.ops.owner, self.ops.manager, self.ops.outsider
        grant_planning_test_actions((self.owner, self.outsider), ('export',))
        self.project, self.baseline = self.ops.project, self.ops.baseline
        self.a, self.b = self.ops.a, self.ops.b
        self.url = f'/api/v1/planning-intelligence/projects/{self.project.pk}/delay-analysis/'
        self.reference = None

    def command(self, action, expected=200, **values):
        response = self.client.post(self.url, {'action': action, **values}, format='json')
        self.assertEqual(response.status_code, expected, getattr(response, 'data', response.content))
        return response.data

    def published(self):
        if self.reference is None:
            self.reference = self.ops.published()
        self.client.force_authenticate(self.owner)
        return self.reference

    def event(self, **values):
        state = self.command('create_event', **({'baseline_id': self.baseline.pk,
            'title': 'Late approved package information', 'activity_ids': [self.a.pk],
            'start_date': '2026-08-27', 'end_date': '2026-08-28',
            'evidence': [{'reference': 'Signed event log EVENT-001, revision 1, page 2'}]} | values))
        return state['events'][0]

    def change(self, event, before='3', value='5', **values):
        return {'event_id': event['id'], 'activity_id': self.a.pk, 'field': 'remaining_duration_days',
            'expected_before': before, 'value': value, 'evidence': 'Approved planner estimate EST-001',
            'reason': 'Explicit remaining-work proposal for technical sensitivity'} | values

    def case(self, event=None, *, calculate=True, **values):
        reference = self.published()
        event = event or self.event()
        case = self.command('create_case', baseline_id=self.baseline.pk, reference_report_id=reference['id'],
            name='Event sensitivity', event_ids=[event['id']])['case']
        case = self.command('save_case', case_id=case['id'], revision=case['revision'],
            **({'changes': [self.change(event)]} | values))['case']
        if calculate:
            case = self.command('calculate_case', case_id=case['id'], revision=case['revision'])['case']
        return case

    def submit(self, case):
        return self.command('submit_case', case_id=case['id'], revision=case['revision'],
            source_fingerprint=case['source_fingerprint'], reason='Ready for independent technical review')['case']

    def approve(self, case):
        self.client.force_authenticate(self.manager)
        return self.command('approve_case', case_id=case['id'], revision=case['revision'],
            source_fingerprint=case['source_fingerprint'], reason='Technical sensitivity reviewed, no entitlement awarded')['case']

    def source(self, project=None):
        project = project or self.project
        graph, _ = EvidenceGraph.objects.get_or_create(project=project, defaults={
            'schema_version': 'test', 'rule_version': 'test', 'source_fingerprint': 'a' * 64})
        file = PlanningFile.objects.create(project=project, original_filename='Event evidence.pdf',
            file='planning/test-event-evidence.pdf', category='other', extracted_text='Confirmed event evidence.')
        document = EvidenceDocumentVersion.objects.create(id=uuid4(), graph=graph, source_file=file,
            filename=file.original_filename, storage_name=file.file.name, file_sha256='b' * 64,
            text_sha256='c' * 64, extracted_text=file.extracted_text, integrity_status='verified')
        sources = [{'document_version': str(document.pk), 'text_sha256': document.text_sha256,
            'locator': {'page': 2, 'line': 4}, 'verbatim': file.extracted_text, 'quote_verified': True}]
        fact = EvidenceNode.objects.create(id=uuid4(), graph=graph, document_version=document,
            kind='fact', entity_id='event:exact-source', property='requirement', value=file.extracted_text,
            provenance_type='document_evidence', sources=sources, status='accepted', current=True)
        return document, fact

    def test_only_published_exact_baseline_reports_can_be_selected(self):
        draft = self.ops.report()
        event = self.event()
        self.command('create_case', expected=404, baseline_id=self.baseline.pk, reference_report_id=draft['id'],
            name='Draft reference forbidden', event_ids=[event['id']])
        other = ScheduleBaseline.objects.create(schedule=self.ops.schedule, source_version=self.ops.version,
            name='Different frozen scope', snapshot={'activities': [{'id': self.b.pk}]},
            approved_by=self.manager, approved_at=timezone.now())
        self.command('create_event', expected=400, baseline_id=other.pk, title='Wrong frozen activity',
            activity_ids=[self.a.pk], evidence=[{'reference': 'Exact scope must match'}])
        self.assertFalse(DelayAnalysisCase.objects.exists())

    def test_outsider_and_read_only_member_cannot_change_cases(self):
        self.client.force_authenticate(self.outsider)
        self.assertEqual(self.client.get(self.url).status_code, 404)
        self.command('create_event', expected=404, baseline_id=self.baseline.pk, title='Not accessible',
            activity_ids=[self.a.pk], evidence=[{'reference': 'Outside project'}])
        reviewer = User.objects.create_user(username='delay-reader', email='delay-reader@example.test')
        ProjectMember.objects.create(project=self.ops.enterprise, user=reviewer, role='reviewer')
        grant_planning_test_actions((reviewer,), ('read', 'create', 'update', 'approve'))
        self.client.force_authenticate(reviewer)
        self.assertEqual(self.client.get(self.url).status_code, 200)
        self.command('create_event', expected=403, baseline_id=self.baseline.pk, title='Read only',
            activity_ids=[self.a.pk], evidence=[{'reference': 'Read access is not write authority'}])
        self.assertFalse(DelayEvent.objects.exists())

    def test_document_and_fact_associations_are_exact_and_project_scoped(self):
        foreign = PlanningProject.objects.create(name='Other source project', created_by=self.owner)
        document, fact = self.source(foreign)
        self.command('create_event', expected=404, baseline_id=self.baseline.pk, title='Wrong document project',
            activity_ids=[self.a.pk], evidence=[{'reference': 'Foreign source', 'document_version_id': str(document.pk)}])
        local_document, local_fact = self.source()
        another_document, _ = self.source()
        self.command('create_event', expected=400, baseline_id=self.baseline.pk, title='Wrong fact version',
            activity_ids=[self.a.pk], evidence=[{'reference': 'Mismatched exact source',
                'document_version_id': str(another_document.pk), 'fact_id': str(local_fact.pk)}])
        event = self.event(evidence=[{'reference': 'Verified quoted passage',
            'document_version_id': str(local_document.pk), 'fact_id': str(local_fact.pk)}])
        case = self.case(event)
        captured = case['run']['events'][0]['evidence'][0]
        self.assertEqual(captured['document']['text_sha256'], local_document.text_sha256)
        self.assertEqual(captured['fact']['sources'][0]['locator'], {'page': 2, 'line': 4})
        self.assertEqual(captured['fact_id'], str(local_fact.pk))

    def test_impact_and_recovery_are_reviewed_without_mutating_business_schedule_or_money(self):
        self.published()
        event = self.event()
        before_baseline = deepcopy(self.baseline.snapshot)
        before_report = deepcopy(OperationalControlReport.objects.get(pk=self.reference['id']).publication)
        ledger_count = CostLedgerEntry.objects.count()
        case = self.case(event, scenarios=[{'id': 'recover', 'name': 'Reviewed remaining effort',
            'changes': [self.change(event, before='5', value='3')]}])
        self.assertEqual(case['run']['result']['impact']['net_finish_shift_calendar_days'], 2, case['run']['result'])
        self.assertEqual(case['run']['result']['scenarios'][0]['net_finish_shift_calendar_days'], 0)
        case = self.approve(self.submit(case))
        self.assertEqual(case['status'], 'approved')
        self.assertEqual(case['recommendation_assessment']['entitlement_status'], 'not_determined')
        self.assertIsNone(case['recommendation_assessment']['approved_extension_calendar_days'])
        self.baseline.refresh_from_db()
        self.assertEqual(self.baseline.snapshot, before_baseline)
        self.assertEqual(OperationalControlReport.objects.get(pk=self.reference['id']).publication, before_report)
        self.assertEqual(CostLedgerEntry.objects.count(), ledger_count)
        self.ops.period.refresh_from_db()
        self.assertEqual(self.ops.period.status, 'open')

    def test_event_revision_and_input_fingerprint_guard_new_calculation_and_submission(self):
        case = self.case()
        old_run = DelayAnalysisRun.objects.get(pk=case['run']['id'])
        old_input = deepcopy(old_run.input_snapshot)
        event = DelayEvent.objects.get(pk=case['event_ids'][0])
        self.command('update_event', expected=409, event_id=event.pk, revision=event.revision + 1,
            title='Stale edit', reason='Stale browser revision')
        self.command('update_event', event_id=event.pk, revision=event.revision,
            evidence=[{'reference': 'Corrected signed event log revision 2'}], reason='Corrected source evidence')
        fresh = self.client.get(self.url, {'case_id': case['id']}).data['case']
        self.assertTrue(fresh['source_stale'])
        self.assertFalse(fresh['permissions']['can_submit'])
        self.command('submit_case', expected=409, case_id=case['id'], revision=case['revision'],
            source_fingerprint=case['source_fingerprint'], reason='Old capture must not submit')
        recalculated = self.command('calculate_case', case_id=case['id'], revision=case['revision'])['case']
        self.assertNotEqual(recalculated['run']['id'], old_run.pk)
        self.assertNotEqual(recalculated['source_fingerprint'], case['source_fingerprint'])
        old_run.refresh_from_db()
        self.assertEqual(old_run.input_snapshot, old_input)

    def test_fact_acceptance_change_invalidates_the_unreviewed_source_snapshot(self):
        document, fact = self.source()
        event = self.event(evidence=[{'reference': 'Accepted original statement',
            'document_version_id': str(document.pk), 'fact_id': str(fact.pk)}])
        case = self.case(event)
        EvidenceNode.objects.filter(pk=fact.pk).update(current=False, status='superseded')
        state = self.client.get(self.url, {'case_id': case['id']}).data['case']
        self.assertTrue(state['source_stale'])
        self.assertNotEqual(state['source_fingerprint'], case['source_fingerprint'])

    def test_fact_correction_stales_pending_run_and_preserves_immutable_original_value(self):
        document, fact = self.source()
        event = self.event(evidence=[{'reference': 'Accepted event fact',
            'document_version_id': str(document.pk), 'fact_id': str(fact.pk)}])
        case = self.case(event)
        old_run = DelayAnalysisRun.objects.get(pk=case['run']['id'])
        old_input = deepcopy(old_run.input_snapshot)
        original_value = deepcopy(fact.value)

        corrected_value = 'Corrected event fact: information arrived later.'
        if connection.vendor == 'postgresql':
            with self.assertRaises(DatabaseError), transaction.atomic():
                EvidenceNode.objects.filter(pk=fact.pk).update(value=corrected_value)
        fact.refresh_from_db()
        self.assertEqual(fact.value, original_value)
        self.assertEqual(fact.status, 'accepted')
        self.assertTrue(fact.current)

        corrected = EvidenceNode.objects.create(id=uuid4(), graph=fact.graph, document_version=document,
            kind=fact.kind, entity_id=fact.entity_id, property=fact.property, value=corrected_value,
            unit=fact.unit, provenance_type='planner_input', status='accepted', current=True,
            validation={'corrects_fact_id': str(fact.pk), 'reason': 'Reviewed correction to the event statement'})
        self.command('update_event', event_id=event['id'], revision=event['revision'],
            evidence=[{'reference': 'Accepted corrected event statement',
                'document_version_id': str(document.pk), 'fact_id': str(corrected.pk)}],
            reason='Explicitly select the corrected fact without overwriting the original')
        state = self.client.get(self.url, {'case_id': case['id']}).data['case']
        self.assertTrue(state['source_stale'])
        self.assertFalse(state['permissions']['can_submit'])
        self.assertNotEqual(state['source_fingerprint'], case['source_fingerprint'])
        self.command('submit_case', expected=409, case_id=case['id'], revision=case['revision'],
            source_fingerprint=case['source_fingerprint'], reason='A changed fact needs a new calculation')

        recalculated = self.command('calculate_case', case_id=case['id'], revision=case['revision'])['case']
        self.assertNotEqual(recalculated['run']['id'], old_run.pk)
        captured = recalculated['run']['events'][0]['evidence'][0]
        self.assertEqual(captured['fact_id'], str(corrected.pk))
        self.assertEqual(captured['fact']['value'], corrected_value)
        self.assertEqual(captured['fact']['status'], fact.status)
        old_run.refresh_from_db()
        self.assertEqual(old_run.input_snapshot, old_input)
        self.assertEqual(old_run.input_snapshot['events'][0]['evidence'][0]['fact']['value'], original_value)

    def test_fact_only_reference_captures_its_exact_document_without_claiming_rechecked_bytes(self):
        document, fact = self.source()
        event = self.event(evidence=[{'reference': 'Exact cited fact, document resolved from that fact',
            'fact_id': str(fact.pk)}])
        case = self.case(event)
        captured = case['run']['events'][0]['evidence'][0]
        self.assertEqual(captured['kind'], 'fact_reference')
        self.assertEqual(captured['document_version_id'], str(document.pk))
        self.assertEqual(captured['document']['filename'], document.filename)
        self.assertEqual(captured['document']['file_sha256'], document.file_sha256)
        self.assertEqual(captured['document']['text_sha256'], document.text_sha256)
        self.assertEqual(captured['document']['recorded_integrity_status'], 'verified')
        self.assertEqual(captured['document']['storage_verification'], 'not_rechecked')
        self.assertEqual(captured['fact']['document_version_id'], str(document.pk))
        for field in ('kind', 'value', 'unit', 'provenance_type'):
            self.assertEqual(captured['fact'][field], getattr(fact, field))
        for issues in (case['issues'], case['run']['result']['issues']):
            warning = next(row for row in issues if row['code'] == 'document_bytes_not_rechecked')
            self.assertEqual(warning['severity'], 'warning')
        self.assertFalse(case['source_stale'])

    def test_approval_only_project_manager_can_review_through_the_module_guard_but_cannot_edit(self):
        case = self.case()
        reviewer = User.objects.create_user(username='delay-approval-only', email='delay-approval-only@example.test')
        ProjectMember.objects.create(project=self.ops.enterprise, user=reviewer, role='project_manager')
        grant_planning_test_actions((reviewer,), ('read', 'approve'))
        for action in ('read', 'approve'):
            self.assertTrue(module_action_allowed(reviewer, 'planning_package', action))
        for action in ('create', 'update'):
            self.assertFalse(module_action_allowed(reviewer, 'planning_package', action))

        # Match secure_module_endpoints' wrapper, including its original class
        # name: the production guard checks that explicit business route name.
        guarded = type(DelayAnalysisView.__name__, (ModuleActionGuardMixin, DelayAnalysisView),
            {'__module__': DelayAnalysisView.__module__}).as_view()
        factory = APIRequestFactory()

        def request(method, payload):
            value = getattr(factory, method)(self.url, payload, format='json')
            force_authenticate(value, user=reviewer)
            return guarded(value, project_id=self.project.pk)

        denied = request('post', {'action': 'save_case', 'case_id': case['id'],
            'revision': case['revision'], 'name': 'Approval authority cannot edit the proposal'})
        self.assertEqual(denied.status_code, 403, denied.data)
        self.assertIn('update permission', str(denied.data['detail']))
        self.assertEqual(DelayAnalysisCase.objects.get(pk=case['id']).name, case['name'])

        case = self.submit(case)
        detail = request('get', {'case_id': case['id']})
        self.assertEqual(detail.status_code, 200, detail.data)
        self.assertTrue(detail.data['case']['permissions']['can_approve'])
        self.assertFalse(detail.data['permissions']['can_write'])
        approved = request('post', {'action': 'approve_case', 'case_id': case['id'],
            'revision': case['revision'], 'source_fingerprint': case['source_fingerprint'],
            'reason': 'Independent technical review by the assigned project manager'})
        self.assertEqual(approved.status_code, 200, approved.data)
        self.assertEqual(approved.data['case']['status'], 'approved')
        self.assertEqual(approved.data['case']['reviewed_by_id'], reviewer.pk)

    def test_every_contributor_is_excluded_from_approving_the_case(self):
        case = self.case()
        self.client.force_authenticate(self.manager)
        case = self.command('save_case', case_id=case['id'], revision=case['revision'], name='Manager reviewed draft')['case']
        self.client.force_authenticate(self.owner)
        case = self.command('save_case', case_id=case['id'], revision=case['revision'], name='Owner submitted draft')['case']
        case = self.command('calculate_case', case_id=case['id'], revision=case['revision'])['case']
        case = self.submit(case)
        for actor in (self.owner, self.manager):
            self.client.force_authenticate(actor)
            self.command('approve_case', expected=403, case_id=case['id'], revision=case['revision'],
                source_fingerprint=case['source_fingerprint'], reason='Contributor cannot independently approve')
        self.assertEqual(DelayAnalysisCase.objects.get(pk=case['id']).status, 'submitted')

    def test_reviewed_case_and_runs_are_immutable_and_revision_retains_history(self):
        case = self.approve(self.submit(self.case()))
        record = DelayAnalysisCase.objects.get(pk=case['id'])
        run = record.current_run
        frozen = deepcopy(run.input_snapshot)
        self.client.force_authenticate(self.owner)
        self.command('save_case', expected=409, case_id=record.pk, revision=record.revision, name='Cannot overwrite')
        revised = self.command('revise_case', case_id=record.pk, reason='New event information for a separate review')['case']
        self.assertEqual(revised['supersedes_id'], record.pk)
        self.assertEqual(revised['status'], 'draft')
        self.assertIsNone(revised['current_run_id'])
        record.name = 'Forbidden model edit'
        with self.assertRaises(ModelValidationError):
            record.save()
        with self.assertRaises(ModelValidationError):
            run.save()
        if connection.vendor == 'postgresql':
            for queryset, values in [(DelayAnalysisCase.objects.filter(pk=record.pk), {'name': 'Forbidden SQL edit'}),
                                     (DelayAnalysisRun.objects.filter(pk=run.pk), {'result': {}})]:
                with self.assertRaises(DatabaseError), transaction.atomic():
                    queryset.update(**values)
                with self.assertRaises(DatabaseError), transaction.atomic():
                    queryset.delete()
        run.refresh_from_db()
        self.assertEqual(run.input_snapshot, frozen)

    def test_reviewed_result_remains_frozen_after_its_event_changes(self):
        case = self.approve(self.submit(self.case()))
        stored_run = deepcopy(case['run'])
        event = DelayEvent.objects.get(pk=case['event_ids'][0])
        self.client.force_authenticate(self.owner)
        self.command('update_event', event_id=event.pk, revision=event.revision,
            title='Later event description', evidence=[{'reference': 'Subsequent evidence'}], reason='Historical correction')
        current = self.client.get(self.url, {'case_id': case['id']}).data['case']
        self.assertEqual(current['run'], stored_run)
        self.assertEqual(current['source_fingerprint'], case['source_fingerprint'])
        self.assertFalse(current['source_stale'])

    def test_extension_recommendation_requires_recorded_contract_notice_and_assessments(self):
        case = self.case(recommendation={'requested_extension_calendar_days': 2})
        error = self.command('submit_case', expected=400, case_id=case['id'], revision=case['revision'],
            source_fingerprint=case['source_fingerprint'], reason='Insufficient contract basis')
        self.assertIn('contract_clause_reference', error)
        self.assertIn('notice_reference', error)
        self.assertIn('concurrency_assessment', error)
        self.assertEqual(DelayAnalysisCase.objects.get(pk=case['id']).status, 'calculated')

    def test_deleted_governance_reference_can_be_removed_without_losing_event_history(self):
        item = GovernanceItem.objects.create(version=self.ops.version, activity=self.a,
            item_type='issue', title='Recorded prerequisite issue', raised_by=self.owner)
        event = self.event(governance_item_id=item.pk)
        item.soft_delete()
        state = self.command('update_event', event_id=event['id'], revision=event['revision'],
            governance_item_id=None, reason='Retire a superseded governance reference')
        self.assertIsNone(next(row for row in state['events'] if row['id'] == event['id'])['governance_item_id'])
        self.assertTrue(GovernanceItem.objects.filter(pk=item.pk, is_deleted=True).exists())

    def test_governance_activity_reassignment_cannot_silently_change_event_scope(self):
        item = GovernanceItem.objects.create(version=self.ops.version, activity=self.a,
            item_type='issue', title='Original explicit activity', raised_by=self.owner)
        event = self.event(governance_item_id=item.pk)
        case = self.case(event)
        GovernanceItem.objects.filter(pk=item.pk).update(activity=self.b)
        self.command('calculate_case', expected=409, case_id=case['id'], revision=case['revision'])
        self.assertEqual(DelayEvent.objects.get(pk=event['id']).activity_ids, [self.a.pk])

    def test_exports_retain_frozen_trace_and_exclude_commercial_actuals(self):
        self.ops.source_cost()  # Included in the published report's private source manifest.
        case = self.approve(self.submit(self.case()))
        self.client.force_authenticate(self.owner)
        state = self.client.get(self.url, {'case_id': case['id']})
        self.assertEqual(state.status_code, 200, state.data)
        self.assertNotIn('999.17', json.dumps(state.data, default=str))
        self.assertNotIn('Sensitive cost', json.dumps(state.data, default=str))
        export_url = f'{self.url}cases/{case["id"]}/export/'
        response = self.client.get(export_url, {'format': 'json'})
        self.assertEqual(response.status_code, 200, getattr(response, 'data', response.content))
        data = json.loads(response.content)
        text = json.dumps(data, default=str)
        self.assertIn(case['source_fingerprint'], text)
        self.assertNotIn('999.17', text)
        self.assertNotIn('source_actuals', text)
        self.assertNotIn('Sensitive cost', text)
        response = self.client.get(export_url, {'format': 'xlsx'})
        self.assertEqual(response.status_code, 200, getattr(response, 'data', response.content))
        workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=False)
        values = '\n'.join(str(value) for sheet in workbook for row in sheet.iter_rows(values_only=True) for value in row if value is not None)
        self.assertIn(case['source_fingerprint'], values)
        self.assertNotIn('999.17', values)
        self.assertNotIn('Sensitive cost', values)
        self.client.force_authenticate(self.outsider)
        self.assertEqual(self.client.get(export_url, {'format': 'json'}).status_code, 404)
