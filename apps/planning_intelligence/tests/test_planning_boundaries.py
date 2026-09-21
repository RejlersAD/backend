"""Evidence gates cannot be bypassed by calculation, approval or native exports."""
import datetime as dt
import io
import json
from decimal import Decimal
from unittest.mock import patch

from django.utils import timezone
from openpyxl import load_workbook

from ..models import CalendarException, ScheduleBaseline
from ..services.cpm import SchedulingError, calculate_schedule_version
from ..services.planning_boundaries import freeze_schedule_inputs, calculation_inputs_current, accepted_input_validation
from ..services.operational_jobs import schedule_state_fingerprint
from ..services.schedule_approval import ScheduleApprovalError, require_accepted_schedule_inputs
from ..services.schedule_export_contract import ScheduleExportError, export_capabilities
from ..services.schedule_exports import generate_schedule_export
from .test_scheduling_engine import ScheduleFixture


class PlanningBoundaryTests(ScheduleFixture):
    def document_activity(self, code='A', duration=2, **metadata):
        return self.activity(code, duration, metadata={'evidence_policy': 'document_driven', **metadata})

    def test_missing_calendar_is_not_replaced_with_five_day_week(self):
        row = self.activity('A', 2)
        self.schedule.default_calendar = None
        self.schedule.save(update_fields=['default_calendar'])
        with self.assertRaises(SchedulingError) as error:
            calculate_schedule_version(self.version)
        self.assertEqual(error.exception.code, 'planning_inputs_not_accepted')
        self.assertIn('calendar_not_specified', {item['code'] for item in error.exception.issues})
        row.refresh_from_db()
        self.assertIsNone(row.planned_start)

    def test_document_input_gate_prevents_calculation_and_approval(self):
        row = self.document_activity()
        blocking = [{'code': 'duration_not_accepted', 'message': 'Review the two source assertions.',
                     'blocks': ['calculation', 'approval', 'export']}]
        with patch('apps.planning_intelligence.services.evidence_graph.validate_schedule_inputs', return_value=blocking):
            with self.assertRaises(SchedulingError):
                calculate_schedule_version(self.version)
            with self.assertRaises(ScheduleApprovalError):
                require_accepted_schedule_inputs(self.version)
        row.refresh_from_db()
        self.assertIsNone(row.planned_start)
        self.assertIsNone(row.total_float_days)

    def test_document_durations_are_not_rounded(self):
        self.document_activity(duration=Decimal('1.25'))
        with patch('apps.planning_intelligence.services.evidence_graph.validate_schedule_inputs', return_value=[]):
            with self.assertRaises(SchedulingError) as error:
                calculate_schedule_version(self.version)
        self.assertIn('duration_resolution_unsupported', {item['code'] for item in error.exception.issues})

    def test_missing_duration_placeholder_is_not_calculated_as_zero(self):
        row = self.activity('A', 0, metadata={'duration_pending': True})
        with self.assertRaises(SchedulingError):
            calculate_schedule_version(self.version)
        row.refresh_from_db()
        self.assertIsNone(row.planned_start)

    def test_approved_version_cannot_be_recalculated_by_service(self):
        row = self.activity('A', 2)
        calculate_schedule_version(self.version)
        row.refresh_from_db()
        previous = (row.planned_start, row.planned_finish, row.total_float_days)
        self.version.status = 'baselined'
        self.version.save(update_fields=['status'])
        with self.assertRaises(SchedulingError) as error:
            calculate_schedule_version(self.version)
        self.assertEqual(error.exception.code, 'immutable_schedule')
        row.refresh_from_db()
        self.assertEqual(previous, (row.planned_start, row.planned_finish, row.total_float_days))

    def test_accepted_calculation_is_reproducible_and_preserves_calendar_exception(self):
        first = self.document_activity('A', 2)
        second = self.document_activity('B', 3)
        self.link(first, second, kind='FS', lag=1)
        CalendarException.objects.create(calendar=self.calendar, date=dt.date(2026, 8, 25), is_working=False)
        with patch('apps.planning_intelligence.services.evidence_graph.validate_schedule_inputs', return_value=[]):
            calculate_schedule_version(self.version)
            before = list(self.version.activities.values_list('external_id', 'planned_start', 'planned_finish', 'total_float_days'))
            calculate_schedule_version(self.version)
        after = list(self.version.activities.values_list('external_id', 'planned_start', 'planned_finish', 'total_float_days'))
        self.assertEqual(before, after)
        self.assertEqual(before[0][2], dt.date(2026, 8, 26))
        self.assertEqual(before[1][1], dt.date(2026, 8, 28))

    def test_calendar_exception_changes_job_identity_and_approval_freshness(self):
        self.document_activity()
        with patch('apps.planning_intelligence.services.evidence_graph.validate_schedule_inputs', return_value=[]):
            calculate_schedule_version(self.version)
            self.assertTrue(calculation_inputs_current(self.version))
            before = schedule_state_fingerprint(self.version)
            CalendarException.objects.create(calendar=self.calendar, date=dt.date(2026, 8, 25), is_working=False)
            self.assertNotEqual(before, schedule_state_fingerprint(self.version))
            self.assertFalse(calculation_inputs_current(self.version))
            with self.assertRaises(ScheduleApprovalError) as error:
                require_accepted_schedule_inputs(self.version)
        self.assertEqual(error.exception.payload['code'], 'calculation_inputs_stale')

    def test_source_planned_dates_are_compared_without_becoming_constraints(self):
        activity = self.document_activity(evidence_entity_id='task:A')
        evidence = {'accepted_inputs': {'task:A': {'start_date': '2026-08-27', 'finish_date': '2026-08-28'}},
                    'facts': [{'id': 'source-start-fact', 'entity_id': 'task:A', 'property': 'start_date', 'status': 'accepted'},
                              {'id': 'source-finish-fact', 'entity_id': 'task:A', 'property': 'finish_date', 'status': 'accepted'}]}
        with patch('apps.planning_intelligence.services.evidence_graph.validate_schedule_inputs', return_value=[]), \
                patch('apps.planning_intelligence.services.evidence_graph.evidence_graph_snapshot', return_value=evidence):
            run = calculate_schedule_version(self.version)
            activity.refresh_from_db()
            self.assertEqual(activity.planned_start, dt.date(2026, 8, 24))
            self.assertEqual(activity.planned_finish, dt.date(2026, 8, 25))
            self.assertEqual(activity.constraint_type, 'none')
            self.assertIsNone(activity.constraint_date)
            mismatches = [item for item in run.issues if item['code'] == 'source_planned_date_mismatch']
            self.assertEqual(len(mismatches), 2)
            self.assertEqual(mismatches[0]['source_fact_id'], 'source-start-fact')
            self.assertEqual(mismatches[0]['source_date'], '2026-08-27')
            self.assertEqual(mismatches[0]['calculated_date'], '2026-08-24')
            readiness = accepted_input_validation(self.version)
            self.assertTrue(readiness['ready_for_calculation'])
            self.assertFalse(readiness['ready_for_approval'])
            self.assertFalse(readiness['ready_for_export'])
            with self.assertRaises(ScheduleApprovalError):
                require_accepted_schedule_inputs(self.version)
            content, _, _ = generate_schedule_export(self.version, 'json')
            self.assertEqual(json.loads(content)['export_state'], 'structured_draft')
            self.assertFalse(json.loads(content)['readiness']['ready_for_export'])

    def test_nonworking_upper_constraints_use_previous_working_day(self):
        self.project.planned_end_date = dt.date(2026, 9, 4)
        self.project.save(update_fields=['planned_end_date'])
        start_bound = self.activity('START-BOUND', 2, constraint_type='start_no_later', constraint_date=dt.date(2026, 8, 30))
        finish_bound = self.activity('FINISH-BOUND', 2, constraint_type='finish_no_later', constraint_date=dt.date(2026, 8, 30))
        calculate_schedule_version(self.version)
        start_bound.refresh_from_db()
        finish_bound.refresh_from_db()
        self.assertEqual(start_bound.late_start, dt.date(2026, 8, 28))
        self.assertEqual(finish_bound.late_start, dt.date(2026, 8, 27))
        self.assertEqual(finish_bound.late_finish, dt.date(2026, 8, 28))
        self.assertEqual(start_bound.constraint_date, dt.date(2026, 8, 30))
        self.project.refresh_from_db()
        self.assertEqual(self.project.planned_end_date, dt.date(2026, 9, 4))

    def test_exact_nonworking_constraint_is_not_silently_shifted_to_monday(self):
        activity = self.activity('EXACT', 2, constraint_type='must_start', constraint_date=dt.date(2026, 8, 30))
        for constraint_type in ('must_start', 'must_finish'):
            with self.subTest(constraint_type=constraint_type):
                activity.constraint_type = constraint_type
                activity.save(update_fields=['constraint_type'])
                with self.assertRaises(SchedulingError) as error:
                    calculate_schedule_version(self.version)
                self.assertEqual(error.exception.code, 'constraint_nonworking_date')
                self.assertEqual(error.exception.issues[0]['constraint_date'], '2026-08-30')
                activity.refresh_from_db()
                self.assertIsNone(activity.planned_start)
                self.assertIsNone(activity.planned_finish)
                self.assertEqual(activity.constraint_date, dt.date(2026, 8, 30))


class ScheduleExportBoundaryTests(ScheduleFixture):
    def test_capabilities_do_not_claim_unimplemented_xml_or_validated_xer(self):
        adapters = {row['format']: row for row in export_capabilities()}
        self.assertEqual(adapters['primavera_xml']['status'], 'unavailable')
        self.assertEqual(adapters['mspdi']['status'], 'implemented_subset')
        self.assertEqual(adapters['mspdi']['verification']['vendor_application_roundtrip'], 'not_tested')
        self.assertEqual(adapters['xer']['status'], 'legacy_unvalidated')
        with self.assertRaises(ScheduleExportError):
            generate_schedule_export(self.version, 'primavera_xml')

    def test_document_plan_cannot_use_legacy_xer_serializer(self):
        self.activity('A', 2, metadata={'evidence_policy': 'document_driven'})
        with self.assertRaises(ScheduleExportError):
            generate_schedule_export(self.version, 'xer')

    def test_json_keeps_draft_missing_values_and_provenance(self):
        self.activity('A', 0, metadata={'duration_pending': True, 'source_references': [{'file_id': 17, 'locator': {'row': 4}}]})
        content, _, _ = generate_schedule_export(self.version, 'json')
        snapshot = json.loads(content)
        self.assertEqual(snapshot['export_state'], 'structured_draft')
        self.assertFalse(snapshot['readiness']['ready_for_calculation'])
        self.assertIsNone(snapshot['activities'][0]['planned_start'])
        self.assertEqual(snapshot['activities'][0]['metadata']['source_references'][0]['locator']['row'], 4)
        self.assertIn('sha256', snapshot['traceability'])

    def test_excel_preserves_long_evidence_and_never_executes_document_formula(self):
        evidence = 'Verbatim evidence: ' + ('source data ' * 4000)
        activity = self.activity('A', 2, metadata={'source_evidence': evidence})
        activity.name = '=HYPERLINK("https://invalid.example")'
        activity.save(update_fields=['name'])
        content, _, _ = generate_schedule_export(self.version, 'xlsx')
        workbook = load_workbook(io.BytesIO(content))
        headers = [cell.value for cell in workbook['Activities'][1]]
        name_cell = workbook['Activities'].cell(2, headers.index('name') + 1)
        self.assertEqual(name_cell.data_type, 's')
        chunks = ''.join(row[1] for row in list(workbook['Traceability'].values)[1:])
        self.assertEqual(json.loads(chunks)['activities'][0]['metadata']['source_evidence'], evidence)

    def test_baseline_export_uses_frozen_calendar_and_schedule_after_source_changes(self):
        row = self.activity('A', 2)
        calculate_schedule_version(self.version)
        row.refresh_from_db()
        from ..schedule_serializers import ScheduleActivitySerializer, ScheduleVersionSerializer
        inputs = freeze_schedule_inputs(self.version)
        frozen = {'version': ScheduleVersionSerializer(self.version).data,
                  'activities': ScheduleActivitySerializer(self.version.activities.all(), many=True).data,
                  'wbs': [], 'relationships': [], 'accepted_inputs': inputs}
        ScheduleBaseline.objects.create(schedule=self.schedule, source_version=self.version, name='Accepted',
                                        approved_by=self.owner, approved_at=timezone.now(), snapshot=frozen)
        self.version.status = 'baselined'
        self.version.save(update_fields=['status'])
        self.calendar.working_weekdays = [0, 1, 2, 3, 4, 5, 6]
        self.calendar.save(update_fields=['working_weekdays'])
        self.project.name = 'Renamed after baseline'
        self.project.save(update_fields=['name'])
        self.schedule.name, self.schedule.code = 'Renamed schedule', 'RENAMED'
        self.schedule.save(update_fields=['name', 'code'])
        content, _, filename = generate_schedule_export(self.version, 'json')
        snapshot = json.loads(content)
        self.assertEqual(snapshot['export_state'], 'approved_baseline')
        self.assertEqual(snapshot['project']['name'], 'FEED Schedule')
        self.assertEqual(snapshot['schedule']['name'], 'Master')
        self.assertEqual(snapshot['schedule']['code'], 'MASTER')
        self.assertEqual(filename, 'FEED_Schedule_MASTER_v1_Schedule.json')
        self.assertEqual(snapshot['calendar']['working_weekdays'], [0, 1, 2, 3, 4])
        self.assertEqual(snapshot['traceability']['sha256'], inputs['sha256'])
        self.assertEqual(snapshot['activities'][0]['planned_finish'], row.planned_finish.isoformat())
