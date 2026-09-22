"""Source-preservation tests; these do not claim Microsoft Project app testing."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZipFile

from django.test import SimpleTestCase
from lxml import etree
from openpyxl import load_workbook

from ..services.ms_project_export import NS, build_mspdi, mspdi_bundle, verify_mspdi
from ..services.schedule_export_contract import ScheduleExportError, export_capabilities
from ..services.schedule_exports import generate_schedule_export


def exchange_snapshot():
    shifts = [{'from': '08:30:00', 'to': '12:30:00'}, {'from': '13:15:00', 'to': '17:15:00'}]
    calendar = {'id': 31, 'name': 'Site & office', 'working_weekdays': [0, 1, 2, 3, 4],
                'hours_per_day': '8.00', 'timezone': 'Asia/Dubai',
                'working_times': {str(day): deepcopy(shifts) for day in range(5)},
                'exceptions': [{'date': '2026-10-07', 'is_working': False, 'working_hours': '0',
                                'working_times': [], 'name': 'Company holiday'},
                               {'date': '2026-10-11', 'is_working': True, 'working_hours': '4.00',
                                'working_times': [{'from': '09:00:00', 'to': '13:00:00'}], 'name': 'Review & issue'}]}
    dates = ['2026-10-05', '2026-10-06', '2026-10-08', '2026-10-09', '2026-10-12']
    activities = [{'id': i + 101, 'external_id': f'ACT-{i + 1}', 'name': f'Activity {i + 1} — Design & review <issued>',
                   'activity_type': 'task', 'duration_days': '1.00', 'wbs_node': 71, 'calendar': 31,
                   'planned_start': day, 'planned_finish': day, 'total_float_days': '-1.25', 'free_float_days': '0.50',
                   'is_critical': True, 'constraint_type': 'none', 'constraint_date': None,
                   'metadata': {'evidence': {'document': 'Scope & specification.pdf', 'page': 2}}}
                  for i, day in enumerate(dates)]
    activities[4].update(activity_type='finish_milestone', duration_days='0.00')
    return {'schema_version': '2.0', 'export_state': 'structured_draft',
            'project': {'id': 1, 'name': 'Plant <A> & expansion'},
            'schedule': {'id': 2, 'name': 'Approved timing', 'code': 'SCH-A', 'planned_start': dates[0]},
            'version': {'id': 3, 'status': 'calculated', 'version': 2},
            'calendar': calendar, 'calendars': [calendar],
            'wbs': [{'id': 70, 'parent': None, 'name': 'Engineering', 'code': 'ENG'},
                    {'id': 71, 'parent': 70, 'name': 'Mechanical & piping', 'code': 'ENG.MECH'}],
            'activities': activities,
            'relationships': [{'id': i + 20, 'predecessor': 101, 'successor': i + 102,
                               'relationship_type': kind, 'lag_days': lag}
                              for i, (kind, lag) in enumerate([('FS', '0'), ('SS', '-0.5'), ('FF', '1.25'), ('SF', '-1')])],
            'resources': [{'id': 12, 'code': 'EMP-7', 'name': 'Engineer <Lead>', 'resource_type': 'labor',
                           'unit': 'hour', 'unit_cost': '125.25', 'capacity_units_per_day': '8'}],
            'assignments': [{'id': 55, 'activity': 101, 'resource': 12, 'budgeted_hours': '7.50',
                             'planned_units': '0.75', 'budgeted_cost': '939.375'}],
            'traceability': {'sha256': 'original-frozen-hash', 'planning_build': {'accepted_plan': 'exact snapshot'}},
            'readiness': {'ready_for_calculation': True, 'ready_for_export': True},
            'risk_register': [{'id': 10, 'name': 'Vendor response', 'source': 'Contract clause 8'}],
            'governance': {'items': [], 'reviews': []}, 'controls': None, 'baselines': []}


class MicrosoftProjectExchangeTests(SimpleTestCase):
    def test_schema_and_independent_parser_preserve_typed_logic_signed_lag_and_shifts(self):
        snapshot = exchange_snapshot()
        original = deepcopy(snapshot)
        content, provenance, report = build_mspdi(snapshot)
        root = etree.fromstring(content)
        ns = {'p': NS}
        self.assertEqual(report['xml_well_formed'], 'passed')
        self.assertEqual(report['vendor_application_roundtrip'], 'not_tested')
        self.assertFalse(report['native_mpp'])
        self.assertEqual(len(root.findall('p:Tasks/p:Task', ns)), 7)
        links = root.findall('p:Tasks/p:Task/p:PredecessorLink', ns)
        self.assertEqual({item.findtext('p:Type', namespaces=ns) for item in links}, {'0', '1', '2', '3'})
        self.assertEqual([item.findtext('p:LinkLag', namespaces=ns) for item in links], ['0', '-2400', '6000', '-4800'])
        self.assertIn(b'&amp;', content)
        self.assertIn(b'&lt;', content)
        self.assertEqual(provenance['snapshot'], original)
        self.assertEqual(snapshot, original)
        first = root.findall('p:Tasks/p:Task', ns)[2]
        self.assertEqual(first.findtext('p:Start', namespaces=ns), '2026-10-05T08:30:00')
        self.assertEqual(first.findtext('p:Finish', namespaces=ns), '2026-10-05T17:15:00')
        self.assertEqual(first.findtext('p:TotalSlack', namespaces=ns), '-6000')

    def test_zip_keeps_exact_baseline_evidence_resources_risk_and_sha_binding(self):
        snapshot = exchange_snapshot()
        snapshot.update(export_state='approved_baseline', baseline_id=999, baseline_approved_at='2026-10-01T10:00:00Z')
        with ZipFile(BytesIO(mspdi_bundle(snapshot))) as bundle:
            self.assertEqual(set(bundle.namelist()), {'schedule.xml', 'radai-provenance.json', 'verification.json'})
            provenance = json.loads(bundle.read('radai-provenance.json'))
            report = json.loads(bundle.read('verification.json'))
            self.assertEqual(provenance['snapshot'], snapshot)
            self.assertEqual(provenance['snapshot']['traceability']['sha256'], 'original-frozen-hash')
            self.assertEqual(provenance['xml_sha256'], sha256(bundle.read('schedule.xml')).hexdigest())
            self.assertEqual(report['provenance_sha256'], sha256(bundle.read('radai-provenance.json')).hexdigest())
            root = etree.fromstring(bundle.read('schedule.xml'))
            self.assertEqual(len(root.findall(f'.//{{{NS}}}Baseline')), 5)

    def test_date_only_full_day_duration_crosses_nonworking_exception_exactly(self):
        snapshot = exchange_snapshot()
        snapshot['activities'][0].update(planned_finish='2026-10-08', duration_days='3')
        content, _, _ = build_mspdi(snapshot)
        self.assertIn(b'<Duration>PT86400S</Duration>', content)

    def test_start_and_finish_milestones_keep_their_distinct_working_boundaries(self):
        snapshot = exchange_snapshot()
        snapshot['activities'][0].update(activity_type='start_milestone', duration_days='0')
        content, _, _ = build_mspdi(snapshot)
        root = etree.fromstring(content)
        tasks = root.findall(f'.//{{{NS}}}Task')
        self.assertEqual(tasks[2].findtext(f'{{{NS}}}Start'), '2026-10-05T08:30:00')
        self.assertEqual(tasks[2].findtext(f'{{{NS}}}Finish'), '2026-10-05T08:30:00')
        self.assertEqual(tasks[-1].findtext(f'{{{NS}}}Start'), '2026-10-12T17:15:00')

    def test_all_supported_explicit_constraints_preserved(self):
        for kind, numeric in [('must_start', '2'), ('must_finish', '3'), ('start_no_earlier', '4'),
                              ('start_no_later', '5'), ('finish_no_later', '7')]:
            with self.subTest(kind=kind):
                snapshot = exchange_snapshot()
                snapshot['activities'][0].update(constraint_type=kind, constraint_date='2026-10-05')
                content, _, _ = build_mspdi(snapshot)
                first = etree.fromstring(content).findall(f'.//{{{NS}}}Task')[2]
                self.assertEqual(first.findtext(f'{{{NS}}}ConstraintType'), numeric)
                self.assertEqual(first.findtext(f'{{{NS}}}ConstraintDate')[:10], '2026-10-05')

    def assert_rejected(self, snapshot, code):
        with self.assertRaises(ScheduleExportError) as error:
            build_mspdi(snapshot)
        self.assertIn(code, {item['code'] for item in error.exception.payload['issues']})

    def test_missing_shifts_are_not_replaced_with_office_hours(self):
        snapshot = exchange_snapshot()
        snapshot['calendars'][0]['working_times'] = {}
        self.assert_rejected(snapshot, 'mspdi_working_times_required')

    def test_overlap_mismatched_hours_and_missing_working_exception_shifts_rejected(self):
        for mutation, code in [
                (lambda s: s['calendar']['working_times']['0'][1].update({'from': '12:00:00'}), 'mspdi_calendar_interval'),
                (lambda s: s['calendar'].update(hours_per_day='7'), 'mspdi_calendar_hours_mismatch'),
                (lambda s: s['calendar']['exceptions'][1].update(working_times=[]), 'mspdi_working_times_required')]:
            snapshot = exchange_snapshot()
            mutation(snapshot)
            self.assert_rejected(snapshot, code)

    def test_missing_dates_unknown_duration_and_nonworking_endpoint_are_not_repaired(self):
        for changes, code in [({'planned_start': None}, 'mspdi_date_required'),
                              ({'duration_days': None}, 'mspdi_numeric_input'),
                              ({'planned_start': '2026-10-07', 'planned_finish': '2026-10-07'}, 'mspdi_nonworking_endpoint'),
                              ({'duration_days': '0.5'}, 'mspdi_duration_date_mismatch')]:
            snapshot = exchange_snapshot()
            snapshot['activities'][0].update(changes)
            self.assert_rejected(snapshot, code)

    def test_unrepresentable_lag_precision_and_nonfinite_float_are_rejected(self):
        snapshot = exchange_snapshot()
        snapshot['relationships'][0]['lag_days'] = '0.00001'
        self.assert_rejected(snapshot, 'mspdi_precision_unsupported')
        snapshot = exchange_snapshot()
        snapshot['activities'][0]['total_float_days'] = 'NaN'
        self.assert_rejected(snapshot, 'mspdi_numeric_input')

    def test_unsupported_loe_material_assignments_and_bad_wbs_are_blocked(self):
        snapshot = exchange_snapshot()
        snapshot['activities'][0]['activity_type'] = 'level_of_effort'
        self.assert_rejected(snapshot, 'mspdi_activity_type_unsupported')
        snapshot = exchange_snapshot()
        snapshot['resources'][0]['resource_type'] = 'material'
        self.assert_rejected(snapshot, 'mspdi_material_assignment_unsupported')
        snapshot = exchange_snapshot()
        snapshot['wbs'][0]['parent'] = 71
        self.assert_rejected(snapshot, 'mspdi_wbs_hierarchy')

    def test_xml_control_chars_never_silently_stripped(self):
        snapshot = exchange_snapshot()
        snapshot['activities'][0]['name'] += '\x01'
        self.assert_rejected(snapshot, 'mspdi_invalid_xml_text')

    def test_cycles_and_orphan_assignments_are_rejected_without_repair(self):
        snapshot = exchange_snapshot()
        snapshot['relationships'].append({'id': 99, 'predecessor': 102, 'successor': 101,
                                          'relationship_type': 'FS', 'lag_days': '0'})
        self.assert_rejected(snapshot, 'mspdi_relationship_cycle')
        snapshot = exchange_snapshot()
        snapshot['assignments'][0]['resource'] = 99999
        self.assert_rejected(snapshot, 'mspdi_assignment_reference')

    def test_unreferenced_calendar_is_retained_in_sidecar_without_inventing_shifts(self):
        snapshot = exchange_snapshot()
        snapshot['calendars'].append({'id': 900, 'name': 'Unrelated project calendar', 'working_times': {}})
        content, provenance, _ = build_mspdi(snapshot)
        self.assertEqual(len(etree.fromstring(content).findall(f'.//{{{NS}}}Calendar')), 1)
        self.assertEqual(len(provenance['snapshot']['calendars']), 2)

    def test_json_excel_and_xml_bundle_retain_same_source_data(self):
        snapshot = exchange_snapshot()
        version = SimpleNamespace(version=2, status='calculated')
        with patch('apps.planning_intelligence.services.schedule_exports.schedule_snapshot', side_effect=lambda v: deepcopy(snapshot)), \
             patch('apps.planning_intelligence.services.schedule_exports.calculation_inputs_current', return_value=True):
            data = json.loads(generate_schedule_export(version, 'json')[0])
            workbook = load_workbook(BytesIO(generate_schedule_export(version, 'xlsx')[0]))
            chunks = [row[1] for row in workbook['Traceability'].iter_rows(min_row=2, values_only=True)]
            excel_snapshot = json.loads(''.join(chunks))
            with ZipFile(BytesIO(generate_schedule_export(version, 'mspdi_zip')[0])) as bundle:
                native_snapshot = json.loads(bundle.read('radai-provenance.json'))['snapshot']
        for key in ['activities', 'wbs', 'calendars', 'relationships', 'resources', 'assignments', 'traceability', 'risk_register']:
            self.assertEqual(data[key], snapshot[key])
            self.assertEqual(excel_snapshot[key], data[key])
            self.assertEqual(native_snapshot[key], data[key])
        self.assertIn('Risk register', workbook.sheetnames)

    def test_draft_or_blocked_input_cannot_export_as_native_calculated_schedule(self):
        snapshot = exchange_snapshot()
        snapshot['version']['status'] = 'draft'
        self.assert_rejected(snapshot, 'mspdi_schedule_not_ready')
        snapshot = exchange_snapshot()
        snapshot['readiness']['ready_for_export'] = False
        self.assert_rejected(snapshot, 'mspdi_schedule_not_ready')

    def test_independent_verifier_detects_name_shift_and_relationship_corruption(self):
        snapshot = exchange_snapshot()
        content, provenance, _ = build_mspdi(snapshot)
        for before, after in [(b'ACT-1</WBS>', b'WRONG</WBS>'),
                              (b'2026-10-05T08:30:00</Start>', b'2026-10-05T09:30:00</Start>'),
                              (b'<LinkLag>-2400</LinkLag>', b'<LinkLag>-2401</LinkLag>')]:
            self.assertIn(before, content)
            with self.assertRaises(ScheduleExportError) as error:
                verify_mspdi(content.replace(before, after), snapshot, provenance['identity_mapping'])
            self.assertEqual(error.exception.payload['code'], 'mspdi_preservation_failed')

    def test_schema_rejects_wrong_duration_datatype_before_field_comparison(self):
        snapshot = exchange_snapshot()
        content, provenance, _ = build_mspdi(snapshot)
        with self.assertRaises(ScheduleExportError) as error:
            verify_mspdi(content.replace(b'<Duration>PT28800S</Duration>', b'<Duration>eight hours</Duration>'),
                         snapshot, provenance['identity_mapping'])
        self.assertEqual(error.exception.payload['code'], 'mspdi_validation_failed')

    def test_export_service_uses_real_xml_and_zip_extensions_and_staleness_guard(self):
        version = SimpleNamespace(version=2, status='calculated')
        with patch('apps.planning_intelligence.services.schedule_exports.schedule_snapshot', side_effect=lambda v: exchange_snapshot()), \
             patch('apps.planning_intelligence.services.schedule_exports.calculation_inputs_current', return_value=True):
            for format, mime, extension in [('mspdi', 'application/xml; charset=utf-8', '.xml'),
                                             ('mspdi_zip', 'application/zip', '.zip')]:
                content, content_type, name = generate_schedule_export(version, format)
                self.assertEqual(content_type, mime)
                self.assertTrue(name.endswith(extension))
                self.assertTrue(content)
        with patch('apps.planning_intelligence.services.schedule_exports.calculation_inputs_current', return_value=False):
            with self.assertRaises(ScheduleExportError) as error:
                generate_schedule_export(version, 'mspdi')
            self.assertEqual(error.exception.payload['code'], 'mspdi_calculation_stale')

    def test_capabilities_never_claim_vendor_application_or_native_mpp_validation(self):
        formats = {row['format']: row for row in export_capabilities()}
        for name in ['mspdi', 'mspdi_zip']:
            self.assertEqual(formats[name]['status'], 'implemented_subset')
            self.assertEqual(formats[name]['verification']['vendor_application_roundtrip'], 'not_tested')
        self.assertEqual(formats['primavera_xml']['status'], 'unavailable')
        self.assertEqual(formats['xer']['status'], 'legacy_unvalidated')
