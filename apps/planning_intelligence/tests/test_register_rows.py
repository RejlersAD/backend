"""Source registers retain their own titles and rows, independent of catalogues."""
import io
from unittest import TestCase

from openpyxl import Workbook

from ..services.register_rows import (
    extract_legacy_register_rows, extract_register_rows, extract_workbook_register_rows,
    normalize_register_discipline,
    register_row_requires_review,
)


class RegisterRowExtractionTests(TestCase):
    MATRIX_HEADER = ('Table 8: Applicable Deliverables for Work Packages\n'
                     'Work Packages\nDocument / Deliverable\n'
                     'S. NO. Discipline Remarks\nDescription 1 2 3\n')

    def test_matrix_retains_applicability_and_remarks_without_assigning_package_columns(self):
        text = self.MATRIX_HEADER + ('4.2 General\n'
                '4.2.1 General Access dossier X x Common report\n'
                '4.2.2 General Optional drawing\n'
                '4.2.3 General Shelter report X If a new building is required\n'
                '4.2.4 General Existing study X Already covered in another report\n'
                '4.2.5 General Release register X\n')
        rows = extract_register_rows(text)
        self.assertEqual(len(rows), 5)
        self.assertEqual([row['register_item'] for row in rows], ['4.2.1', '4.2.2', '4.2.3', '4.2.4', '4.2.5'])
        self.assertEqual([row['applicability_status'] for row in rows], ['marked', 'not_marked', 'conditional', 'conditional', 'marked'])
        self.assertEqual([row['name'] for row in rows if not register_row_requires_review(row)], ['Access dossier', 'Release register'])
        self.assertEqual(rows[0]['source_remarks'], 'Common report')
        self.assertEqual(rows[0]['applicability_marks'], 'X x')
        self.assertEqual(rows[0]['explicit_dimensions'], {'discipline': {'value': 'General', 'header': 'Discipline'}})
        for row in rows:
            self.assertEqual(row['package_columns_status'], 'not_resolved')
            self.assertNotIn('package', row['explicit_dimensions'])
            self.assertEqual(text[row['start']:row['end']].strip(), row['source_excerpt'])
            self.assertEqual(row['source_locator']['register_item'], row['register_item'])

    def test_matrix_requires_caption_and_complete_column_header(self):
        body = '4.2 General\n4.2.1 General Access dossier X\n'
        for header in (self.MATRIX_HEADER.replace('Applicable Deliverables', 'Reference Documents'),
                       self.MATRIX_HEADER.replace('Remarks', ''),
                       self.MATRIX_HEADER.replace('Discipline', ''),
                       self.MATRIX_HEADER.split('\n', 1)[1]):
            self.assertEqual(extract_register_rows(header + body), [])

    def test_matrix_interleaved_text_quarantines_both_possible_title_boundaries(self):
        text = self.MATRIX_HEADER + ('4.2 General\n4.2.1 General Clear report X\n'
                'possibly a wrapped title or remark\n4.2.2 General Other report X\n'
                '4.2.3 General Last report X\n')
        rows = extract_register_rows(text)
        self.assertEqual([row['title_boundary_status'] for row in rows], ['ambiguous', 'ambiguous', 'explicit_marked_row'])
        self.assertEqual([row['name'] for row in rows if not register_row_requires_review(row)], ['Last report'])

    def test_matrix_retains_serial_only_and_wrapped_discipline_rows_as_inventory(self):
        text = self.MATRIX_HEADER + ('4.2 General\n4.2.1 General Clear report X\n'
                '4.3 Specialist Team\nSpecialist\n4.3.1 Wrapped report X\nTeam\n'
                '4.3.2\nAnother wrapped title\n')
        rows = extract_register_rows(text)
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[1]['discipline'], 'not_specified')
        self.assertEqual(rows[1]['applicability_status'], 'ambiguous')
        self.assertTrue(register_row_requires_review(rows[2]))
        self.assertEqual(rows[2]['original_title'], '4.3.2')

    def test_matrix_page_furniture_preserves_rows_but_next_appendix_is_not_scope(self):
        text = self.MATRIX_HEADER + ('4.2 General\n4.2.1 General First report X\n'
                'Register.docx 1 / 2\fExample Classification: Internal\n'
                '4.2.2 General Second report X\nAppendix 9: References\n'
                '4.2.3 General Existing drawing X\n')
        rows = extract_register_rows(text)
        self.assertEqual([row['name'] for row in rows], ['First report', 'Second report'])
        self.assertTrue(all(not register_row_requires_review(row) for row in rows))

    def test_matrix_multiple_marker_runs_cannot_silently_truncate_a_title_containing_x(self):
        rows = extract_register_rows(self.MATRIX_HEADER + '4.2 General\n4.2.1 General Equipment X study X\n')
        self.assertTrue(register_row_requires_review(rows[0]))

    def test_captioned_pdf_register_retains_groups_without_inventing_disciplines(self):
        text = ('APPENDIX 2 - PROJECT DELIVERABLES\nTable 4: Project Deliverables\n'
                'S. No. Description\nGeneral\n1 Site access dossier\n2 Survey record (if required)\n'
                'Community Liaison\n1 Stakeholder briefing pack\n')
        rows = extract_register_rows(text)
        self.assertEqual([row['name'] for row in rows], ['Site access dossier', 'Survey record (if required)', 'Stakeholder briefing pack'])
        self.assertEqual([row['source_group'] for row in rows], ['General', 'General', 'Community Liaison'])
        self.assertEqual({row['discipline'] for row in rows}, {'not_specified'})
        self.assertTrue(all(row['explicit_dimensions'] == {} for row in rows))
        for row in rows:
            self.assertEqual(text[row['start']:row['end']].strip(), row['source_excerpt'])
            self.assertEqual(row['source_locator']['source_group'], row['source_group'])

    def test_caption_and_serial_header_are_both_required(self):
        for text in ('S. No. Description\nGeneral\n1 Site access dossier\n',
                     'Table 4: Project Deliverables\n1 Site access dossier\n',
                     'The contractor shall provide deliverables\nS. No. Description\n1 Site access dossier\n',
                     'Table 4: Project Deliverables\nnot a table\nstill prose\nmore prose\nS. No. Description\n1 Dossier\n'):
            self.assertEqual(extract_register_rows(text), [], text)

    def test_repeated_pdf_headers_keep_distinct_duplicate_serials_and_skip_furniture(self):
        furniture = 'Repeated confidential source document header\n'
        text = (furniture + 'Table 4: Engineering Deliverables\nS. No. Description\nGeneral\n'
                '12 Interface report\n12 Interface drawing\n' + furniture + 'Page: 8\n'
                'S. No. Description\n13 Release package\n' + furniture + 'Page: 9\n')
        rows = extract_register_rows(text)
        self.assertEqual([row['name'] for row in rows], ['Interface report', 'Interface drawing', 'Release package'])
        self.assertEqual([row['register_item'] for row in rows], [12, 12, 13])
        self.assertEqual(len({row['source_line'] for row in rows}), 3)
        self.assertTrue(all(row['source_group'] == 'General' for row in rows))

    def test_wrapped_interleaved_serial_preserves_quote_and_marks_ambiguous_boundary(self):
        text = ('Table 2: Deliverable Register\nSerial No Description\nSupport Services\n'
                '1 Clean boundary drawing\nCondition report for retained\n2\nstructures\n'
                '3 Subsequent clear report\n')
        rows = extract_register_rows(text)
        self.assertEqual([row['name'] for row in rows], ['Clean boundary drawing', 'Condition report for retained structures', 'Subsequent clear report'])
        self.assertEqual(rows[1]['title_boundary_status'], 'ambiguous')
        self.assertIn('retained\n2\nstructures', rows[1]['source_excerpt'])
        self.assertEqual(rows[0]['title_boundary_status'], 'explicit_numbered_row')

    def test_unknown_wrapped_text_is_not_claimed_as_explicit_columns(self):
        rows = extract_register_rows('Table 3: Deliverables\nS. No. Description\n1 A long title\npossibly a footer\n2 Clear title\n')
        self.assertEqual(rows[0]['title_boundary_status'], 'ambiguous')
        self.assertEqual(rows[1]['name'], 'Clear title')

    def test_new_appendices_software_reference_lists_and_notes_are_excluded(self):
        text = ('Table 4: FEED Deliverables\nS. No. Description\n1 Intended report\n'
                'Notes:\n1. The list requires review.\n2. Terms are subject to agreement.\n'
                'APPENDIX 4 - SOFTWARE\nS. No. Description\n1 Spreadsheet application\n'
                'Table 9: Reference Documents\nS. No. Description\n1 Existing source drawing\n')
        rows = extract_register_rows(text)
        self.assertEqual([row['name'] for row in rows], ['Intended report'])

    def test_bare_note_labels_end_table_before_numbered_prose(self):
        for label in ('Notes', 'Note', 'NOTES', 'Note.', 'Notes:'):
            text = ('Table 1: Deliverables\nS. No. Description\nGeneral\n'
                    f'1 Scope dossier\n2 Scope report\n{label}\n1 Review assumptions\n2 Check access\n')
            self.assertEqual([row['name'] for row in extract_register_rows(text)], ['Scope dossier', 'Scope report'], label)

    def test_possible_title_continuation_before_serial_reset_remains_ambiguous(self):
        for title, continuation in (('Design criteria for', 'existing facilities'),
                                    ('Design criteria for', 'Existing Facilities'),
                                    ('Design criteria', 'for existing facilities'),
                                    ('Survey records /', 'Supporting Details')):
            text = ('Table 1: Deliverables\nS. No. Description\nGeneral\n'
                    f'1 {title}\n{continuation}\n1 Inspection report\n')
            rows = extract_register_rows(text)
            self.assertEqual([row['name'] for row in rows], [f'{title} {continuation}', 'Inspection report'])
            self.assertEqual(rows[0]['title_boundary_status'], 'ambiguous')
            self.assertEqual(rows[1]['source_group'], 'General')
            self.assertEqual(rows[1]['title_boundary_status'], 'explicit_numbered_row')

    def test_captioned_and_delimited_tables_preserve_each_source_occurrence(self):
        text = ('Table 3: Deliverables\nS. No. Description\n1 Captioned report\n'
                'APPENDIX 4 - DOCUMENT REGISTER\nSerial No|Discipline|Document Title\n'
                '1|General|Delimited report\n')
        self.assertEqual([row['name'] for row in extract_register_rows(text)], ['Captioned report', 'Delimited report'])

    def workbook(self, rows, sheet_name='MDR'):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name
        for row in rows:
            sheet.append(row)
        stream = io.BytesIO()
        workbook.save(stream)
        stream.seek(0)
        self.addCleanup(stream.close)
        return stream

    def test_workbook_preserves_original_titles_discipline_and_true_row(self):
        stream = self.workbook([
            [], [], [None, 'SL. NO.', 'DISCIPLINE', 'DRAWING / DOCUMENT (TITLE)'],
            [None, None, 'GENERAL', None],
            [None, 1, 'GENERAL', 'MASTER DELIVERABLE REGISTER'],
            [None, None, 'HVAC', None],
            [None, 219, 'HVAC', 'HVAC ADEQUEACY REPORT -FAR-0'],
            [None, 220, 'HVAC', 'HVAC ADEQUEACY REPORT -FAR-6'],
        ], 'Sheet1')

        rows = extract_workbook_register_rows(stream)

        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[1]['name'], 'HVAC ADEQUEACY REPORT -FAR-0')
        self.assertEqual(rows[1]['original_title'], rows[1]['name'])
        self.assertEqual(rows[1]['discipline'], 'hvac')
        self.assertEqual(rows[1]['discipline_label'], 'HVAC')
        self.assertEqual(rows[1]['register_item'], 219)
        self.assertEqual(rows[1]['source_locator'], {'sheet': 'Sheet1', 'row': 7})
        self.assertEqual(stream.tell(), 0)

    def test_legacy_flattened_excel_is_not_labelled_as_worksheet_rows(self):
        text = (
            '--- Sheet: Sheet1 ---\nSL. NO. | DISCIPLINE | DRAWING / DOCUMENT (TITLE)\n'
            'GENERAL\n1 | GENERAL | MASTER DELIVERABLE REGISTER\n'
            'HVAC\n219 | HVAC | HVAC ADEQUEACY REPORT -FAR-0\n'
        )
        rows = extract_register_rows(text)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1]['source_locator'], {'sheet': 'Sheet1', 'line': 6})
        self.assertIsNone(rows[1]['row_number'])
        self.assertEqual(text[rows[1]['start']:rows[1]['end']].strip(), rows[1]['source_excerpt'])

    def test_220_row_register_roundtrips_all_six_groups_without_catalogue_names(self):
        # Match the observed workbook layout and size using generated test data;
        # no project documents or their confidential contents enter the fixture.
        from ..services.parsers import _extract_xlsx

        contents = [[], [], [None, 'SL. NO.', 'DISCIPLINE', 'DRAWING / DOCUMENT (TITLE)']]
        expected = []
        for label, count in (
            ('GENERAL', 34), ('HSE', 81), ('INSTRUMENTATION', 41),
            ('ELECTRICAL', 24), ('CIVIL', 38), ('HVAC', 2),
        ):
            contents.append([None, None, label, None])
            for index in range(count):
                serial = len(expected) + 1
                title = f'{label} Deliverable {index + 1} (+/- 15%) -FAR-{serial % 2}'
                expected.append((serial, label, title))
                contents.append([None, serial, label, title])
        stream = self.workbook(contents, 'Sheet1')

        workbook_rows = extract_workbook_register_rows(stream)
        text_rows = extract_register_rows(_extract_xlsx(stream))

        for rows in (workbook_rows, text_rows):
            self.assertEqual(len(rows), 220)
            self.assertEqual(
                [(row['register_item'], row['discipline_label'], row['name']) for row in rows],
                expected,
            )
        self.assertEqual(workbook_rows[-1]['row_number'], 229)

    def test_quoted_csv_titles_keep_commas_quotes_and_newlines(self):
        rows = extract_register_rows(
            'Serial No,Discipline,Document Title,Document No,Revision\n'
            '1,CIVIL,"FOUNDATION, \"\"A\"\"\nREPORT",CV-001,02\n'
            '2,CIVIL,"FOUNDATION, B",CV-002,01\n'
        )
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]['name'], 'FOUNDATION, "A"\nREPORT')
        self.assertEqual(rows[0]['document_number'], 'CV-001')
        self.assertEqual(rows[0]['document_revision'], '02')
        self.assertEqual(rows[1]['source_line'], 4)

    def test_tsv_with_empty_columns_does_not_shift_titles(self):
        rows = extract_register_rows(
            '\tDocument No\tDepartment\tDocument Title\tRev\n'
            '\t\tElectrical\tLOAD LIST - FAR-0\t\n'
        )
        self.assertEqual(rows[0]['name'], 'LOAD LIST - FAR-0')
        self.assertEqual(rows[0]['document_number'], '')
        self.assertEqual(rows[0]['discipline'], 'electrical')

    def test_excel_parse_preserves_blank_optional_columns_and_quoted_title_content(self):
        from ..services.parsers import _extract_xlsx

        title = 'FOUNDATION | "SPECIAL"\nDRAWING'
        stream = self.workbook([
            [None, 'SL. NO.', 'DISCIPLINE', 'DOCUMENT NO', 'DOCUMENT TITLE', 'REV'],
            [None, 1, 'CIVIL', None, title, None],
            [None, 2, 'CIVIL', 'CV-002', 'DETAIL DRAWING', 'A'],
        ])
        text = _extract_xlsx(stream)
        rows = extract_register_rows(text)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]['name'], title)
        self.assertEqual(rows[0]['document_number'], '')
        self.assertEqual(rows[0]['document_revision'], '')
        self.assertEqual(rows[1]['name'], 'DETAIL DRAWING')
        self.assertEqual(rows[1]['document_number'], 'CV-002')
        self.assertEqual(rows[1]['document_revision'], 'A')

    def test_excel_rows_and_legacy_pdf_rows_can_form_one_exact_register(self):
        spreadsheet = extract_register_rows(
            'SL. NO. | DISCIPLINE | DOCUMENT TITLE\n1 | GENERAL | PROJECT PLAN\n'
        )
        pdf = extract_legacy_register_rows(
            '1 CIVIL ABC-CV-001 Foundation drawing North Island NEW 1 B\n'
            '2 HVAC ABC-HV-002 Ventilation layout North Island EXISTING 2 C\n'
        )
        self.assertEqual([row['name'] for row in spreadsheet + pdf], [
            'PROJECT PLAN', 'Foundation drawing North Island', 'Ventilation layout North Island',
        ])
        self.assertEqual(pdf[1]['discipline'], 'hvac')
        self.assertEqual(pdf[1]['document_number'], 'ABC-HV-002')
        self.assertEqual(pdf[1]['document_revision'], 'C')
        self.assertEqual(pdf[1]['source_line'], 2)
        self.assertIn('North Island', pdf[1]['source_excerpt'])
        self.assertEqual(pdf[1]['title_boundary_status'], 'ambiguous')

    def test_legacy_pdf_shared_document_title_endings_are_retained(self):
        rows = extract_legacy_register_rows(
            '1 CIVIL ABC-CV-001 Foundation Inspection Report NEW 1 B\n'
            '2 CIVIL ABC-CV-002 Concrete Inspection Report NEW 1 B\n'
        )
        self.assertEqual([row['name'] for row in rows], [
            'Foundation Inspection Report', 'Concrete Inspection Report',
        ])

    def test_missing_merged_discipline_is_carried_forward(self):
        stream = self.workbook([
            ['Serial Number', 'Discipline', 'Document Title'],
            [1, 'INSTRUMENTATION', 'SYSTEM ARCHITECTURE DIAGRAM FOR FAR 0'],
            [2, None, 'HSSD ADEQUACY REPORT-FAR0'],
        ])
        rows = extract_workbook_register_rows(stream)
        self.assertEqual([row['discipline'] for row in rows], ['instrumentation'] * 2)
        legacy = extract_register_rows(
            'Serial Number | Discipline | Document Title\n'
            '1 | INSTRUMENTATION | SYSTEM ARCHITECTURE DIAGRAM FOR FAR 0\n'
            '2 | HSSD ADEQUACY REPORT-FAR0\n'
        )
        self.assertEqual(legacy[1]['discipline_label'], 'INSTRUMENTATION')
        self.assertEqual(legacy[1]['name'], 'HSSD ADEQUACY REPORT-FAR0')

    def test_repeated_headers_blanks_and_sections_are_not_deliverables(self):
        rows = extract_register_rows(
            'SL. NO. | DISCIPLINE | DRAWING / DOCUMENT (TITLE)\n'
            'HSE\n1 | HSE | HAZOP REPORT\n\n'
            'SL. NO. | DISCIPLINE | DRAWING / DOCUMENT (TITLE)\n'
            '2 | HSE | HAZOP REPORT\nTotal | HSE | 2\n'
        )
        self.assertEqual([row['register_item'] for row in rows], [1, 2])
        # Two source entries with the same title are still two deliverables.
        self.assertEqual([row['name'] for row in rows], ['HAZOP REPORT', 'HAZOP REPORT'])

    def test_each_sheet_uses_its_own_columns_and_discipline(self):
        stream = self.workbook([
            ['Serial No', 'Discipline', 'Document Title'],
            [1, 'HSE', 'HAZOP REPORT'],
        ], 'Studies')
        import openpyxl
        workbook = openpyxl.load_workbook(stream)
        other = workbook.create_sheet('Other')
        other.append(['Document No', 'Document Title'])
        other.append(['PKG-1', 'EPC PACKAGE'])
        stream.seek(0)
        workbook.save(stream)
        stream.seek(0)
        rows = extract_workbook_register_rows(stream)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1]['discipline'], 'not_specified')
        self.assertEqual(rows[1]['source_locator'], {'sheet': 'Other', 'row': 2})
        self.assertIsNone(rows[1]['register_item'])

    def test_unknown_department_name_is_preserved_without_guessing(self):
        rows = extract_register_rows(
            'Serial No;Department;Deliverable Name\n1;People & Culture;Recruitment plan\n'
        )
        self.assertEqual(rows[0]['discipline_label'], 'People & Culture')
        self.assertEqual(rows[0]['discipline'], 'people_culture')
        self.assertEqual(rows[0]['name'], 'Recruitment plan')

    def test_unstructured_prose_does_not_become_register_rows(self):
        self.assertEqual(extract_register_rows('The contractor shall prepare a material specification.'), [])
        self.assertEqual(extract_register_rows('Title | Note\nElectrical Design | Proposed'), [])

    def test_known_disciplines_are_matched_as_labels_not_substrings(self):
        self.assertEqual(normalize_register_discipline('Civil / Structural'), 'civil')
        self.assertEqual(normalize_register_discipline('Instrumentation & Control'), 'instrumentation')
        self.assertEqual(normalize_register_discipline('HVAC'), 'hvac')
        self.assertEqual(normalize_register_discipline('GENERAL'), 'general')
        self.assertEqual(normalize_register_discipline('Civilian Engagement'), 'civilian_engagement')

    def test_corrupt_workbook_restores_stream_position(self):
        stream = io.BytesIO(b'not a workbook')
        stream.seek(2)
        with self.assertRaises(Exception):
            extract_workbook_register_rows(stream)
        self.assertEqual(stream.tell(), 2)
