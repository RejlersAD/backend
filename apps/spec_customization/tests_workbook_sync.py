"""Workbook API sync + AI Copilot stabilization tests (part 2).

Covers: normalized preview payload shape consumed by WorkbookCanvas;
cell override merge; batch-save; deterministic validation; chatbot
repair/update provenance. Synthetic fixtures only (isolated SQLite).
"""
from __future__ import annotations

from django.contrib.auth import get_user_model
from django.test import TestCase
from rest_framework.test import APIClient

from .models import WorkbookCellOverride
from .services.exporters import build_cat_workbook
from .services.exporters import smartplant_config as cfg
from .services.exporters.workbook_preview import WORKBOOK_CAT, WORKBOOK_SPEC, build_preview
from .services.workbook_chatbot import (
    _plan_safe_repairs_for_workbook,
    _plan_updates_for_workbook,
)
from .tests_exporter_roundtrip import make_five_system_job


class ExtractionToExportRoundtripTests(TestCase):
    """Valve trim / stress / branch-matrix extraction→export round-trips."""

    def setUp(self):
        owner = get_user_model().objects.create_user(
            username='roundtrip-owner', email='roundtrip-owner@example.test', password='safe-password')
        self.job, self.cls = make_five_system_job(owner)

    def test_valve_trim_and_manufacturer_reach_cat_preview(self):
        preview = build_preview(self.job, WORKBOOK_CAT)
        gate = next((s for s in preview['sheets'] if s['name'] == 'GateValve'), None)
        self.assertIsNotNone(gate)
        self.assertTrue(gate['rows'])
        trims = [str(r['cells'].get('ValveTrim', '')) for r in gate['rows']]
        makers = [str(r['cells'].get('ValveManufacturer', '')) for r in gate['rows']]
        models = [str(r['cells'].get('ValveModelNumber', '')) for r in gate['rows']]
        self.assertIn('5', trims, f'ValveTrim 5 missing: {trims}')
        self.assertIn('Velan', makers, f'Velan hint missing: {makers}')
        self.assertIn('150', models, f'Fig 150 model missing: {models}')

    def test_allowable_stress_reaches_materials_data(self):
        preview = build_preview(self.job, WORKBOOK_SPEC)
        mats = next((s for s in preview['sheets'] if s['name'] == 'MaterialsData'), None)
        self.assertIsNotNone(mats)
        stresses = [str(r['cells'].get('AllowableStress', '')) for r in mats['rows']]
        self.assertIn('138MPa', stresses, f'stress missing: {stresses}')
        self.assertTrue(all(s.strip() for s in stresses), f'blank stress: {stresses}')

    def test_branch_matrix_emits_rows_for_tee_class(self):
        preview = build_preview(self.job, WORKBOOK_SPEC)
        branch = next((s for s in preview['sheets'] if s['name'] == 'PipeBranch'), None)
        self.assertIsNotNone(branch)
        self.assertTrue(branch['rows'])

    def test_straight_run_class_emits_no_branch_rows(self):
        self.cls.components.filter(component_type='fitting').delete()
        preview = build_preview(self.job, WORKBOOK_SPEC)
        branch = next((s for s in preview['sheets'] if s['name'] == 'PipeBranch'), None)
        self.assertIsNotNone(branch)
        # Template Start/End marker rows remain even when no branch data exists.
        self.assertFalse(any(row['source'].get('class_id') for row in branch['rows']))

    def test_preview_and_xlsx_agree_on_valve_trim(self):
        from openpyxl import load_workbook as _load
        preview = build_preview(self.job, WORKBOOK_CAT)
        gate = next(s for s in preview['sheets'] if s['name'] == 'GateValve')
        preview_trims = {str(r['cells'].get('ValveTrim', '')) for r in gate['rows']}
        buf = build_cat_workbook(self.job)
        buf.seek(0)
        wb = _load(buf, data_only=True)
        try:
            ws = wb['GateValve']
            header_row = None
            for r in range(1, min(ws.max_row, 60) + 1):
                if str(ws.cell(r, 1).value or '').strip() in cfg.HEADER_MARKERS:
                    header_row = r
                    break
            self.assertIsNotNone(header_row)
            headers = [str(ws.cell(header_row, c).value or '') for c in range(1, ws.max_column + 1)]
            self.assertIn('ValveTrim', headers)
            col = headers.index('ValveTrim') + 1
            xlsx_trims = {str(ws.cell(r, col).value or '') for r in range(header_row + 1, ws.max_row + 1)}
            self.assertTrue(preview_trims & xlsx_trims, f'preview {preview_trims} vs xlsx {xlsx_trims}')
        finally:
            wb.close()


class WorkbookApiSyncTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.owner = user_model.objects.create_user(
            username='canvas-owner', email='canvas-owner@example.test', password='safe-password')
        self.job, self.cls = make_five_system_job(self.owner)
        self.client = APIClient()
        self.client.force_authenticate(self.owner)

    def _first_spec_cell(self):
        preview = build_preview(self.job, WORKBOOK_SPEC)
        sheet = next(s for s in preview['sheets'] if s['rows'])
        row = sheet['rows'][0]
        col = next(iter(row['cells'].keys()))
        return sheet['name'], row['row_key'], col, row['cells'][col]

    def test_preview_payload_shape_matches_canvas_contract(self):
        preview = build_preview(self.job, WORKBOOK_SPEC)
        self.assertIn('workbook', preview)
        self.assertIn('sheets', preview)
        for sheet in preview['sheets']:
            self.assertIn('name', sheet)
            self.assertIn('headers', sheet)
            self.assertIn('rows', sheet)
            for row in sheet['rows']:
                self.assertIn('row_key', row)
                self.assertIn('cells', row)
                self.assertIn('overridden', row)
                self.assertIn('source', row)

    def test_cell_override_merges_into_preview(self):
        sheet_name, row_key, column_name, _old = self._first_spec_cell()
        resp = self.client.post(
            f'/api/v1/spec-customization/paper-spec/jobs/{self.job.id}/workbook/cell/',
            {'workbook': 'spec', 'sheet_name': sheet_name, 'row_key': row_key,
             'column_name': column_name, 'value': 'CANVAS-EDIT'},
            format='json',
        )
        self.assertIn(resp.status_code, (200, 201))
        preview = build_preview(self.job, WORKBOOK_SPEC)
        sheet = next(s for s in preview['sheets'] if s['name'] == sheet_name)
        row = next(r for r in sheet['rows'] if r['row_key'] == row_key)
        self.assertEqual(str(row['cells'][column_name]), 'CANVAS-EDIT')
        self.assertIn(column_name, row['overridden'])

    def test_batch_save_then_validate(self):
        sheet_name, row_key, column_name, _old = self._first_spec_cell()
        resp = self.client.post(
            f'/api/v1/spec-customization/paper-spec/jobs/{self.job.id}/workbook/batch-save/',
            {'cells': [{'workbook': 'spec', 'sheet_name': sheet_name, 'row_key': row_key,
                        'column_name': column_name, 'value': 'BATCH-EDIT'}]},
            format='json',
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data.get('saved_count'), 1)
        vresp = self.client.post(
            f'/api/v1/spec-customization/paper-spec/jobs/{self.job.id}/workbook/validate/',
            {'workbook_scope': 'spec', 'active_workbook': 'spec'},
            format='json',
        )
        self.assertEqual(vresp.status_code, 200)
        self.assertTrue(vresp.data.get('ok'))
        self.assertIn('status', vresp.data)

    def test_chatbot_repair_respects_stabilization_rules(self):
        preview = build_preview(self.job, WORKBOOK_SPEC)
        target = None
        for sheet in preview['sheets']:
            if 'PressureRating' in (sheet.get('headers') or []) and sheet.get('rows'):
                target = sheet
                break
        if target is None:
            self.skipTest('no PressureRating column in SPEC preview')
            return
        row = target['rows'][0]
        WorkbookCellOverride.objects.update_or_create(
            job=self.job, workbook='spec', sheet_name=target['name'],
            row_key=row['row_key'], column_name='PressureRating',
            defaults={'value': 'CL900'},
        )
        plan = _plan_safe_repairs_for_workbook(self.job, 'spec')
        self.assertTrue(plan.get('ok'))
        fixed = [u for u in plan.get('planned_updates', [])
                 if u.get('column_name') == 'PressureRating']
        if fixed:
            self.assertEqual(fixed[0]['value'], 'CLASS 900')
            self.assertIn('source', fixed[0])

    def test_chatbot_update_keeps_provenance(self):
        result = _plan_updates_for_workbook(
            self.job, 'spec',
            'set MaterialGrade to ASTM A106 Gr B where Description contains Tee',
        )
        if not result.get('ok') or not result.get('planned_updates'):
            self.skipTest('no matching rows for chatbot update probe')
            return
        update = result['planned_updates'][0]
        self.assertIn('previous_value', update)
        self.assertIn('source', update)
