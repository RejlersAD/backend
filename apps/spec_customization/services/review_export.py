"""Visible coverage notices for review downloads; never engineering approval."""
from openpyxl.styles import Font, PatternFill, Alignment

from .completeness import extraction_completeness


def stamp_review_notice(book, job, workbook):
    from ..models import PipingClassComponent
    from .exporters import smartplant_config as cfg

    coverage = extraction_completeness(job)
    revision = cfg.SPEC_TEMPLATE_PROFILE['revision_sheet']
    sheet = book[revision] if workbook == 'spec' and revision in book.sheetnames else book.create_sheet('Review Notice', 0)
    sheet['A2'] = 'RADAI REVIEW ONLY - NOT FOR SMARTPLANT IMPORT'
    sheet['A2'].font = Font(bold=True, color='9C0006', size=13)
    sheet['A3'] = 'Unverified engineering data. Resolve extraction gaps and matching/release checks before use.'
    sheet['A3'].font = Font(bold=True, color='9C0006')
    failed = coverage['failed_chunks']
    ranges = '; '.join(f"{item.get('page_from', '?')}-{item.get('page_to', '?')}: {item.get('reason', 'processing/review gap')}" for item in failed)
    count = PipingClassComponent.objects.filter(piping_class__job=job).count()
    metadata = [
        ('Source document', job.document.original_filename),
        ('Extraction job', str(job.pk)),
        ('Extraction status', job.status),
        ('Source coverage', coverage['state']),
        ('Component candidates', str(count)),
        ('Chunks processed / total', f'{job.chunks_done} / {job.chunks_total}'),
        ('Failed/review page ranges', ranges[:30000] or 'No recorded page gaps; engineering correctness remains unverified.'),
        ('Release state', 'REVIEW ONLY - no import/release approval'),
    ]
    for row, (label, value) in enumerate(metadata, 5):
        sheet.cell(row, 1, label)
        cell = sheet.cell(row, 2, str(value).replace('\x00', ''))
        cell.data_type = 's'
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    sheet['B8'].fill = PatternFill('solid', fgColor='FFF2CC')
    if workbook != 'spec':
        sheet.column_dimensions['A'].width = 30
        sheet.column_dimensions['B'].width = 90
    book.properties.title = f'RADAI REVIEW ONLY {workbook.upper()} workbook'
    book.properties.subject = f'Unverified review data; source coverage: {coverage["state"]}; not for import'
    book.properties.description = 'Review download only. Existing engineering validation and frozen-release requirements still apply.'
    return coverage
