"""The user-selected SPEC schema must not introduce sample engineering facts."""
from pathlib import Path
import tempfile

from django.test import SimpleTestCase, TestCase
from django.contrib.auth import get_user_model
from openpyxl import Workbook, load_workbook

from .services.exporters import smartplant_config as cfg
from .services.exporters.smartplant_exporter import build_spec_workbook, _write_row
from .services.exporters.workbook_preview import build_preview, WORKBOOK_SPEC
from .tests_exporter_roundtrip import make_five_system_job


class TemplateStructureTests(SimpleTestCase):
    def test_default_has_fifty_sheets_and_no_reference_rows(self):
        book = load_workbook(cfg.SPEC_TEMPLATE_PATH, read_only=True, data_only=True)
        try:
            self.assertEqual(len(book.sheetnames), 50)
            self.assertEqual(book.sheetnames[:4], ['Revision History', 'Index', 'PipingCommodityFilter', 'PipingMaterialsClassData'])
            self.assertEqual(book.sheetnames[-1], 'WeldGapRule')
            for sheet in book:
                active = False
                for row in sheet.iter_rows(values_only=True):
                    marker = str(row[0] or '').strip()
                    if marker == 'Start': active = True
                    elif marker == 'End': active = False
                    elif active:
                        self.assertFalse(any(value not in (None, '') for value in row), sheet.title)
            self.assertIsNone(book['Revision History']['B5'].value)
            self.assertIsNone(book['Revision History']['B6'].value)
        finally:
            book.close()

    def test_header_case_does_not_drop_values_or_create_extra_columns(self):
        book = Workbook(); sheet = book.active
        _write_row(sheet, 2, {'SPECNAME': 2, 'COMMODITYCODE': 3}, {'SpecName': 'A1', 'CommodityCode': 'REVIEWED-CODE'})
        self.assertEqual((sheet['B2'].value, sheet['C2'].value), ('A1', 'REVIEWED-CODE'))
        self.assertEqual(sheet.max_column, 3)

    def test_reference_preparation_preserves_input_and_removes_merged_metadata(self):
        from scripts.prepare_s3d_spec_template import prepare
        with tempfile.TemporaryDirectory() as directory:
            reference = Path(directory) / 'reference.xlsx'; legacy = Path(directory) / 'legacy.xlsx'; output = Path(directory) / 'output.xlsx'
            book = Workbook(); sheet = book.active; sheet.title = 'PipingMaterialsClassData'
            sheet.append(['Head', 'SPECNAME']); sheet.append(['Start']); sheet.append([None, 'SAMPLE-ONLY']); sheet.append(['End'])
            metadata = book.create_sheet('Revision History'); metadata.merge_cells('B5:D5'); metadata['B5'] = 'SOURCE-MACHINE'
            book.save(reference); book.save(legacy); book.close()
            original = reference.read_bytes()
            prepare(reference, output, legacy)
            self.assertEqual(original, reference.read_bytes())
            book = load_workbook(output)
            try:
                self.assertIsNone(book['PipingMaterialsClassData']['B3'].value)
                self.assertIsNone(book['Revision History']['B5'].value)
            finally:
                book.close()

    def test_comparison_normalizes_explicit_fractions_without_converting_units(self):
        from scripts.compare_s3d_workbooks import size
        self.assertEqual(size('½”'), size('0.50', 'in'))
        self.assertEqual(size('1 1/2"'), size('1.5', 'in'))
        self.assertNotEqual(size('DN 15')[1], size('0.50', 'in')[1])
        self.assertIsNone(size('unsure')[0])

    def test_comparison_marks_size_candidates_without_approving_codes_or_secondary_sizes(self):
        from scripts.compare_s3d_workbooks import compare
        with tempfile.TemporaryDirectory() as directory:
            raw = Path(directory) / 'raw.xlsx'; reference = Path(directory) / 'expected.xlsx'; report = Path(directory) / 'report.xlsx'
            book = Workbook(); sheet = book.active; sheet.title = 'AA1B'
            sheet.append(['Class Code', 'AA1B'])
            sheet.append(['Type', 'Sub-Type', 'Size From', 'Size To', 'Material Std', 'End Conn.', 'Description', 'Notes'])
            sheet.append(['pipe', 'PIPE', '1/2"', '2"', None, None, 'Literal source description', 'Source page 3'])
            book.save(raw); book.close()
            book = Workbook(); sheet = book.active; sheet.title = 'PipingMaterialsClassData'
            sheet.append(['Head', 'SPECNAME', 'MATERIALSDESCRIPTION', 'FLUIDSERVICE']); sheet.append(['Start']); sheet.append([None, 'AA1B', 'Reference description', 'Reference service']); sheet.append(['End'])
            sheet = book.create_sheet('PipingCommodityFilter')
            sheet.append(['Head', 'SPECNAME', 'SHORTCODE', 'FIRSTSIZEFROM', 'FIRSTSIZETO', 'FIRSTSIZEUNITS', 'SECONDSIZEFROM', 'SECONDSIZETO', 'COMMODITYCODE']); sheet.append(['Start']); sheet.append([None, 'AA1B', 'Piping', '1', '1', 'in', '0.5', '0.5', 'REFERENCE-CODE']); sheet.append(['End'])
            sheet = book.create_sheet('CorrosionAllowance'); sheet.append(['Head', 'CORROSIONALLOWANCE']); sheet.append(['Start']); sheet.append(['End'])
            book.save(reference); book.close()
            original = (raw.read_bytes(), reference.read_bytes())
            result = compare(raw, reference, report)
            self.assertEqual(result['expected_rows_with_family_first_size_candidates'], 1)
            self.assertEqual(result['expected_rows_with_secondary_sizes'], 1)
            self.assertEqual(result['fully_verified_engineering_matches'], 0)
            self.assertEqual(original, (raw.read_bytes(), reference.read_bytes()))
            with self.assertRaises(ValueError):
                compare(raw, reference, raw)


class TemplateExportTests(TestCase):
    def setUp(self):
        owner = get_user_model().objects.create_user(username='profile-owner', email='profile@example.test')
        self.job, _ = make_five_system_job(owner)

    def test_preview_uses_existing_field_keys_without_legacy_extra_sheets(self):
        preview = build_preview(self.job, WORKBOOK_SPEC)
        sheets = {sheet['name']: sheet for sheet in preview['sheets']}
        self.assertNotIn('MaterialsData', sheets)
        self.assertIn('MaterialsData', preview['excluded_legacy_sections'])
        self.assertIn('SpecName', sheets['PipingMaterialsClassData']['headers'])
        self.assertEqual(sheets['PipingMaterialsClassData']['rows'][0]['cells']['SpecName'], 'A1')
        self.assertIn('SpecName', sheets['PipingCommodityFilter']['headers'])
        self.assertNotIn('SPECNAME', sheets['PipingCommodityFilter']['headers'])
        for sheet in sheets.values():
            keys = {key.casefold(): key for row in sheet['rows'] for key in row['cells']}
            for header in sheet['headers']:
                if header.casefold() in keys:
                    self.assertEqual(header, keys[header.casefold()], sheet['name'])

    def test_export_retains_exact_schema_and_updates_index_without_sample_facts(self):
        book = load_workbook(build_spec_workbook(self.job), data_only=True)
        try:
            self.assertEqual(book.sheetnames, cfg.SPEC_TEMPLATE_PROFILE['sheet_names'])
            self.assertEqual(book['PipingCommodityFilter']['B4'].value, 'SPECNAME')
            self.assertEqual(book['PipingMaterialsClassData']['B6'].value, 'A1')
            self.assertIn('DRAFT', book['Revision History']['A2'].value)
            statuses = [row[2] for row in book['Index'].iter_rows(min_row=5, values_only=True)]
            self.assertNotIn('Success', statuses)
        finally:
            book.close()
