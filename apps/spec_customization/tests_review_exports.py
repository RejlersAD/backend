"""Explicit review attachments preserve partial coverage and release guards."""
import io

from django.contrib.auth import get_user_model
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIRequestFactory, force_authenticate

from .services.exporters import smartplant_config as cfg
from .release_models import SpecAuditEvent
from .tests_exporter_roundtrip import make_five_system_job
from .views import export_smartplant_spec, export_smartplant_cat


class ReviewWorkbookDownloadsTests(TestCase):
    def setUp(self):
        self.owner = get_user_model().objects.create_user(username='review-export-owner', email='review@example.test')
        self.job, self.cls = make_five_system_job(self.owner)
        self.job.chunks_total = 2
        self.job.chunks_done = 1
        self.job.config_snapshot = {'extraction_outcome': {'failed_chunks': [
            {'page_from': 11, 'page_to': 20, 'reason': 'timeout'}]}}
        self.job.save()

    def request(self, kind='spec', query=None, user=None, anonymous=False):
        request = APIRequestFactory().get('/review-export/', query if query is not None else {'draft': '1', 'review': '1'})
        if not anonymous:
            force_authenticate(request, user=user or self.owner)
        view = export_smartplant_spec if kind == 'spec' else export_smartplant_cat
        return view(request, job_id=self.job.pk)

    def test_partial_review_spec_retains_fifty_sheets_data_and_visible_coverage(self):
        response = self.request()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['X-Spec-Export-State'], 'review')
        self.assertEqual(response['X-Spec-Coverage'], 'partial')
        self.assertIn('REVIEW_', response['Content-Disposition'])
        event = SpecAuditEvent.objects.get(job_identity=self.job.pk, event='review_export')
        self.assertEqual(event.details, {'workbook': 'spec', 'coverage': 'partial'})
        book = load_workbook(io.BytesIO(response.content), read_only=True)
        try:
            self.assertEqual(book.sheetnames, cfg.SPEC_TEMPLATE_PROFILE['sheet_names'])
            notice = book['Revision History']
            self.assertIn('NOT FOR SMARTPLANT IMPORT', notice['A2'].value)
            self.assertEqual(notice['B8'].value, 'partial')
            self.assertIn('11-20: timeout', notice['B11'].value)
            self.assertIn('A1', [cell.value for row in book['PipingCommodityFilter'] for cell in row])
        finally:
            book.close()

    def test_partial_review_cat_retains_catalog_and_notice(self):
        response = self.request('cat')
        self.assertEqual(response.status_code, 200)
        book = load_workbook(io.BytesIO(response.content), read_only=True)
        try:
            self.assertEqual(book.sheetnames[0], 'Review Notice')
            self.assertGreater(len(book.sheetnames), 1)
            self.assertEqual(book['Review Notice']['B8'].value, 'partial')
            self.assertIn('Unverified', book.properties.subject)
        finally:
            book.close()

    def test_normal_draft_and_release_still_reject_partial_jobs(self):
        for kind in ('spec', 'cat'):
            for query in ({}, {'draft': '1'}):
                with self.subTest(kind=kind, query=query):
                    self.assertEqual(self.request(kind, query).status_code, 409)

    def test_review_mode_cannot_select_release_or_omit_draft(self):
        for query in ({'review': '1'}, {'draft': '1', 'review': '1', 'release_id': 'a-release'}):
            self.assertEqual(self.request(query=query).status_code, 400)

    def test_running_and_empty_jobs_have_clear_errors(self):
        self.job.status = 'processing'
        self.job.save(update_fields=['status'])
        response = self.request()
        self.assertEqual(response.status_code, 409)
        self.assertIn('finish', response.data['error'])
        self.job.status = 'completed'
        self.job.save(update_fields=['status'])
        self.cls.components.all().delete()
        response = self.request()
        self.assertEqual(response.status_code, 409)
        self.assertIn('No component rows', response.data['error'])

    def test_ownership_and_authentication_are_required(self):
        other = get_user_model().objects.create_user(username='other-review-user', email='other@example.test')
        self.assertIn(self.request(user=other).status_code, (403, 404))
        self.assertIn(self.request(anonymous=True).status_code, (401, 403))
