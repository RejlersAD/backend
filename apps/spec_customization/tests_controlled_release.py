"""Synthetic Phase 3/4 lifecycle, reproducibility and access boundaries."""
from io import BytesIO
from unittest.mock import patch
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.files.base import ContentFile
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from .matching_models import MatchingWorkbookSet, ComponentMatchDecision
from .models import PaperSpecDocument, PaperSpecExtractionJob, PipingClass, PipingClassComponent, WorkbookCellOverride
from .project_models import SpecProject
from .release_models import SpecRelease, SpecAuditEvent, ValidationRun, ImportAttempt
from .services.controlled_release import digest, canonical_workbook
from .services.component_matcher import parse_match_xlsx
from .services.reference_intelligence import file_sha256
from .tests_matching_pipeline import RemoteMemoryStorage, xlsx, spec_xlsx


class ControlledReleaseTests(TestCase):
    def setUp(self):
        import uuid
        suffix = uuid.uuid4().hex[:8]
        self.owner = get_user_model().objects.create_user(username=f'release-owner-{suffix}', email=f'owner-{suffix}@example.test')
        self.other = get_user_model().objects.create_user(username=f'release-other-{suffix}', email=f'other-{suffix}@example.test')
        self.client = APIClient()
        self.client.force_authenticate(self.owner)
        self.project = SpecProject.objects.create(name='Release probe', created_by=self.owner)
        self.document = PaperSpecDocument.objects.create(file=ContentFile(b'%PDF-synthetic', name='source.pdf'),
            original_filename='source.pdf', sha256_hash=digest(b'%PDF-synthetic'), uploaded_by=self.owner, project_id=str(self.project.pk))
        self.job = PaperSpecExtractionJob.objects.create(document=self.document, created_by=self.owner, status='completed',
            config_snapshot={'extraction_outcome': {'failed_chunks': []}})
        self.pc = PipingClass.objects.create(job=self.job, class_code='A', material_grade='A105', pressure_rating='150')
        self.component = PipingClassComponent.objects.create(piping_class=self.pc, component_type='pipe', size_from='2', description='Pipe')
        storage = RemoteMemoryStorage()
        for kind in ('match', 'spec', 'cat'):
            self.enterContext(patch.object(MatchingWorkbookSet._meta.get_field(f'{kind}_file'), 'storage', storage))
        self.reference = MatchingWorkbookSet.objects.create(project=self.project, schema_valid=True, is_parsed=True,
            standard_edition='Pilot edition', units='NPS/class',
            match_file=ContentFile(xlsx('Mappings', ['PDF', 'B', 'C', 'Catalog'], [['PIPE', '', '', 'Straight Pipe']]), name='Match.xlsx'),
            spec_file=ContentFile(spec_xlsx(), name='SPEC.xlsx'),
            cat_file=ContentFile(xlsx('StraightPipe', ['IndustryCommodityCode', 'Npd[1]', 'PressureRating[1]', 'MaterialGrade'],
                [['CODE-001', 2, 150, 'A105']]), name='CAT.xlsx'))
        parse_match_xlsx(self.reference.match_file, self.reference)
        for kind in ('match', 'spec', 'cat'):
            setattr(self.reference, f'{kind}_file_hash', file_sha256(getattr(self.reference, f'{kind}_file')))
        self.reference.save()
        self.assertEqual(self.client.post('/api/v1/spec-customization/matching-set-freeze/', {'workbook_set_id': str(self.reference.pk)}).status_code, 200)
        self.job.reference_workbook_set = self.reference
        self.job.save(update_fields=['reference_workbook_set'])
        self.decision = ComponentMatchDecision.objects.create(job=self.job, workbook_set=self.reference, pdf_component_name='PIPE',
            nominal_size='2', pressure_rating='150', material_grade='A105', state='unique_match',
            matched_commodity_code='CODE-001', review_decision='approved', final_commodity_code='CODE-001',
            reviewed_by=self.owner, candidates=[{'commodity_code': 'CODE-001'}])
        self.enterContext(override_settings(S3D_APPROVED_REFERENCE_SET_IDS=[str(self.reference.pk)],
            S3D_RELEASE_AUTHORITY_USER_IDS=[str(self.owner.pk)], S3D_DISPOSABLE_PILOT_TARGETS=[
                {'id': 'disposable-test', 'label': 'Synthetic target', 'disposable': True, 'sp3d_version': 'Agreed test version', 'import_profile': 'Piping pilot'}]))
        # Exercise command boundaries with a small deterministic workbook;
        # a separate test uses the actual builders/preview/export integration.
        def preview(job, kind):
            if kind == 'spec':
                cells = {'CommodityCode': 'CODE-001'}
            else:
                cells = {'IndustryCommodityCode': 'CODE-001'}
            return {'sheets': [{'name': 'Synthetic', 'headers': list(cells), 'rows': [{'row_key': 'synthetic', 'cells': cells,
                'source': {'class_id': str(self.pc.pk), 'component_id': str(self.component.pk)}, 'match_decision': {'approved': True, 'commodity_code': 'CODE-001'}}]}], 'workbook': kind}
        self.preview_patch = self.enterContext(patch('apps.spec_customization.services.controlled_release.build_preview', side_effect=preview))
        def workbook(*args, **kwargs):
            out = BytesIO(); wb = Workbook(); wb.active.append(['CODE-001']); wb.save(out); out.seek(0); return out
        from .services.exporters.smartplant_exporter import build_spec_workbook, build_cat_workbook
        self.real_spec_exporter, self.real_cat_exporter = build_spec_workbook, build_cat_workbook
        self.enterContext(patch('apps.spec_customization.services.exporters.smartplant_exporter.build_spec_workbook', side_effect=workbook))
        self.enterContext(patch('apps.spec_customization.services.exporters.smartplant_exporter.build_cat_workbook', side_effect=workbook))
        self.root = f'/api/v1/spec-customization/paper-spec/jobs/{self.job.pk}/'

    def frozen(self):
        validation = self.client.post(self.root + 'releases/validate/')
        self.assertEqual(validation.status_code, 201, validation.data)
        draft = self.client.post(self.root + 'releases/')
        response = self.client.post(self.root + f"releases/{draft.data['id']}/freeze/")
        self.assertEqual(response.status_code, 200, response.data)
        return SpecRelease.objects.get(pk=response.data['id'])

    def test_validate_freeze_and_released_bytes_are_reproducible(self):
        release = self.frozen()
        first = self.client.get(self.root + 'export-spec/')
        second = self.client.get(self.root + 'export-spec/')
        self.assertEqual(first.status_code, 200)
        self.assertEqual(first.content, second.content)
        self.assertEqual(digest(first.content), release.manifest['spec_sha256'])
        self.assertEqual(first['X-Spec-Export-State'], 'frozen')
        self.assertTrue(SpecAuditEvent.objects.filter(job_identity=self.job.pk, event='released_export').exists())
        self.assertEqual(ValidationRun.objects.filter(job=self.job).count(), 1)

    def test_unknown_completeness_validation_errors_and_authority_block_freeze(self):
        with override_settings(S3D_RELEASE_AUTHORITY_USER_IDS=[]):
            draft = self.client.post(self.root + 'releases/')
            response = self.client.post(self.root + f"releases/{draft.data['id']}/freeze/")
            self.assertEqual(response.status_code, 409)
            self.assertIn('Release authority is not configured.', response.data['blockers'])
        self.job.config_snapshot = {}
        self.job.save(update_fields=['config_snapshot'])
        state = self.client.get(self.root + 'releases/')
        self.assertTrue(any('complete coverage' in value for value in state.data['blockers']))
        self.job.config_snapshot = {'extraction_outcome': {'failed_chunks': []}}
        self.job.save(update_fields=['config_snapshot'])
        def invalid_preview(job, kind):
            return {'workbook': kind, 'sheets': [{'name': 'Synthetic', 'headers': ['CommodityCode'], 'rows': [{'row_key': 'broken', 'source': {'class_id': str(self.pc.pk)}, 'cells': {'CommodityCode': ''}}]}]}
        self.preview_patch.side_effect = invalid_preview
        validation = self.client.post(self.root + 'releases/validate/')
        self.assertGreater(validation.data['error_count'], 0)
        draft = self.client.post(self.root + 'releases/')
        response = self.client.post(self.root + f"releases/{draft.data['id']}/freeze/")
        self.assertEqual(response.status_code, 409)
        self.assertTrue(any('zero errors' in value for value in response.data['blockers']))

    def test_cell_override_invalidates_without_mutating_historical_release(self):
        release = self.frozen()
        original = release.manifest.copy()
        response = self.client.post(self.root + 'workbook/cell/', {'workbook': 'spec', 'sheet_name': 'Synthetic',
            'row_key': 'synthetic', 'column_name': 'Description', 'value': 'Reviewed correction'})
        self.assertEqual(response.status_code, 201)
        self.assertEqual(self.client.get(self.root + 'export-spec/').status_code, 409)
        release.refresh_from_db()
        self.assertEqual(release.state, 'frozen')
        self.assertEqual(release.manifest, original)
        self.assertTrue(release.successors.filter(state='invalidated').exists())
        self.assertTrue(SpecAuditEvent.objects.filter(job_identity=self.job.pk, event='cell_override').exists())

    def test_direct_data_changes_cannot_bypass_stale_export_guard(self):
        self.frozen()
        self.pc.material_grade = 'Changed'
        self.pc.save(update_fields=['material_grade'])
        self.assertEqual(self.client.get(self.root + 'export-cat/').status_code, 409)

    def test_unauthorized_release_and_pilot_access(self):
        self.client.force_authenticate(self.other)
        for route in ('releases/', 'release-history/', 'import-attempts/'):
            self.assertEqual(self.client.get(self.root + route).status_code, 403)
        self.assertEqual(self.client.post(self.root + 'releases/validate/').status_code, 403)

    def test_evidence_cannot_be_updated_or_deleted(self):
        release = self.frozen()
        release.state = 'invalidated'
        with self.assertRaises(ValidationError): release.save()
        with self.assertRaises(ValidationError): release.delete()
        with self.assertRaises(ValidationError): SpecAuditEvent.objects.all().delete()
        with self.assertRaises(ValidationError): ValidationRun.objects.all().update(error_count=0)

    def test_stale_draft_requires_new_cycle(self):
        self.client.post(self.root + 'releases/validate/')
        draft = self.client.post(self.root + 'releases/')
        self.client.post(self.root + 'workbook/cell/', {'workbook': 'cat', 'sheet_name': 'Synthetic', 'row_key': 'synthetic', 'column_name': 'Description', 'value': 'new'})
        response = self.client.post(self.root + f"releases/{draft.data['id']}/freeze/")
        self.assertEqual(response.status_code, 409)
        self.assertTrue(any('stale' in value for value in response.data['blockers']))

    def test_import_report_is_persisted_with_ground_truth_and_reconciliation(self):
        release = self.frozen()
        response = self.client.post(self.root + 'import-attempts/', {'release_id': str(release.pk), 'target_id': 'disposable-test',
            'status': 'failed', 'rejection_count': 2, 'report': SimpleUploadedFile('rejections.txt', b'Synthetic rejection report'),
            'measurements': '{"true_positive": 8, "false_positive": 2, "false_negative": 1, "ground_truth_reference": "reviewed sample A"}'}, format='multipart')
        self.assertEqual(response.status_code, 201, response.data)
        first = ImportAttempt.objects.get(pk=response.data['id'])
        self.assertEqual(first.report_sha256, digest(b'Synthetic rejection report'))
        replay = self.client.post(self.root + 'import-attempts/', {'release_id': str(release.pk), 'target_id': 'disposable-test', 'status': 'failed', 'rejection_count': 2, 'report': SimpleUploadedFile('rejections.txt', b'Synthetic rejection report'), 'measurements': '{"true_positive": 8, "false_positive": 2, "false_negative": 1, "ground_truth_reference": "reviewed sample A"}'}, format='multipart')
        self.assertEqual(replay.status_code, 200, replay.data)
        self.assertEqual(replay.data['id'], str(first.pk))
        download = self.client.get(self.root + f'import-attempts/{first.pk}/report/')
        self.assertEqual(download.status_code, 200)
        self.assertEqual(download.content, b'Synthetic rejection report')
        response = self.client.post(self.root + 'import-attempts/', {'release_id': str(release.pk), 'parent_id': str(first.pk),
            'target_id': 'disposable-test', 'status': 'reconciled', 'report': SimpleUploadedFile('reconciled.txt', b'Synthetic reconciliation')}, format='multipart')
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(ImportAttempt.objects.filter(release=release).count(), 2)

    def test_pilot_rejects_unconfigured_target_and_missing_report(self):
        release = self.frozen()
        payload = {'release_id': str(release.pk), 'target_id': 'unapproved', 'status': 'planned'}
        self.assertEqual(self.client.post(self.root + 'import-attempts/', payload).status_code, 409)
        payload.update(status='succeeded', target_id='disposable-test')
        self.assertEqual(self.client.post(self.root + 'import-attempts/', payload).status_code, 400)

    def test_zip_timestamps_do_not_change_canonical_workbook_hash(self):
        wb = Workbook(); wb.active.append(['same input'])
        first = BytesIO(); wb.save(first)
        second = BytesIO(); wb.save(second)
        self.assertEqual(canonical_workbook(first.getvalue()), canonical_workbook(second.getvalue()))

    def test_reviewed_codes_reach_real_preview_and_export(self):
        from .services.exporters.workbook_preview import build_preview
        # Use the real preview and exporter implementations for this bridge.
        cat = build_preview(self.job, 'cat')
        spec = build_preview(self.job, 'spec')
        cat_rows = [row for sheet in cat['sheets'] for row in sheet['rows'] if row.get('source', {}).get('component_id') == str(self.component.pk)]
        self.assertTrue(cat_rows)
        self.assertTrue(any(row['cells'].get('IndustryCommodityCode') == 'CODE-001' for row in cat_rows))
        self.assertTrue(any(row['cells'].get('CommodityCode') == 'CODE-001' for sheet in spec['sheets'] for row in sheet['rows']), [(sheet['name'], row.get('source'), row['cells'].get('CommodityCode'), row['cells'].get('IndustryCommodityCode')) for sheet in spec['sheets'] for row in sheet['rows'] if row.get('source', {}).get('class_id')])
        self.assertTrue(any(row.get('provenance', {}).get('IndustryCommodityCode') == 'reference' for row in cat_rows))
        output = self.real_cat_exporter(self.job, preview_data=cat)
        workbook = load_workbook(output, read_only=True)
        self.assertTrue(any('CODE-001' in row for sheet in workbook for row in sheet.iter_rows(values_only=True)))
        self.assertEqual(cat['reference_set_id'], str(self.reference.pk))
