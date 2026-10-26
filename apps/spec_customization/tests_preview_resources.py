"""Resource and ownership regressions for large-source upload/preview recovery."""
from pathlib import Path
import tempfile
from types import SimpleNamespace
from unittest.mock import Mock, patch

from django.contrib.auth import get_user_model
from django.test import SimpleTestCase, TestCase
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIRequestFactory, force_authenticate

from .services.exporters import workbook_preview as preview
from .models import PaperSpecDocument, PaperSpecExtractionJob


class PreviewResourceTests(SimpleTestCase):
    def setUp(self):
        directory = self.enterContext(tempfile.TemporaryDirectory())
        self.path = str(Path(directory) / 'template.xlsx')
        book = Workbook()
        book.active.title = 'Generated'
        for name in ('Generated', 'Static'):
            sheet = book[name] if name in book.sheetnames else book.create_sheet(name)
            sheet.append(['Head', 'Description'])
            if name == 'Static':
                sheet.append(['', 'Static source text'])
            sheet.append(['End', ''])
        book.save(self.path)
        book.close()
        self.books = []

        def tracked(*args, **kwargs):
            result = load_workbook(*args, **kwargs)
            result.close = Mock(wraps=result.close)
            self.books.append(result)
            return result

        preview._HEADERS_CACHE.clear()
        self.enterContext(patch.object(preview, 'load_workbook', side_effect=tracked))
        self.enterContext(patch.object(preview.cfg, 'SPEC_TEMPLATE_PATH', self.path))
        self.enterContext(patch.object(preview, '_enumerate_spec_rows', side_effect=lambda job: iter([
            ('Generated', 'source:1', {}, {'Head': '', 'Description': 'Literal source text'})])))
        self.overrides = self.enterContext(patch.object(preview, '_load_overrides', return_value={}))
        self.bindings = self.enterContext(patch('apps.spec_customization.services.release_provenance.apply_reference_bindings',
                                               side_effect=lambda job, result: result))

    def test_one_template_open_for_all_headers_and_static_rows_and_resource_closed(self):
        data = preview.build_preview(SimpleNamespace(id='synthetic'), 'spec')
        self.assertEqual(len(self.books), 1)
        self.books[0].close.assert_called_once()
        sheets = {sheet['name']: sheet for sheet in data['sheets']}
        self.assertEqual(sheets['Generated']['rows'][0]['cells']['Description'], 'Literal source text')
        self.assertEqual(sheets['Static']['rows'][0]['cells']['Description'], 'Static source text')
        self.assertIsNone(preview._TEMPLATE_CONTEXT.get())

    def test_failure_also_closes_resource_and_restores_request_context(self):
        self.bindings.side_effect = RuntimeError('Synthetic final-stage failure')
        with self.assertRaises(RuntimeError):
            preview.build_preview(SimpleNamespace(id='synthetic'), 'spec')
        self.books[0].close.assert_called_once()
        self.assertIsNone(preview._TEMPLATE_CONTEXT.get())

    def test_fresh_build_observes_saved_override_without_retaining_previous_result(self):
        first = preview.build_preview(SimpleNamespace(id='synthetic'), 'spec')
        self.overrides.return_value = {('Generated', 'source:1'): {'Description': 'Reviewed edit'}}
        second = preview.build_preview(SimpleNamespace(id='synthetic'), 'spec')
        self.assertEqual(first['sheets'][0]['rows'][0]['cells']['Description'], 'Literal source text')
        self.assertEqual(second['sheets'][0]['rows'][0]['cells']['Description'], 'Reviewed edit')
        self.assertEqual(second['sheets'][0]['rows'][0]['provenance']['Description'], 'override')
        self.assertEqual(len(self.books), 2)


class OwnUploadRecoveryTests(TestCase):
    def setUp(self):
        self.owner = get_user_model().objects.create_user(username='recovery-owner', email='owner@example.test', is_superuser=True)
        other = get_user_model().objects.create_user(username='recovery-other', email='other@example.test')
        document = PaperSpecDocument.objects.create(original_filename='source.pdf', sha256_hash='9' * 64)
        self.own = PaperSpecExtractionJob.objects.create(document=document, created_by=self.owner)
        PaperSpecExtractionJob.objects.create(document=document, created_by=other)

    def request(self, query, authenticated=True):
        from .views import list_jobs
        request = APIRequestFactory().get('/api/v1/spec-customization/paper-spec/jobs/', query)
        if authenticated:
            force_authenticate(request, user=self.owner)
        return list_jobs(request)

    def test_admin_can_recover_only_own_uploads_with_mine_filter(self):
        response = self.request({'mine': '1'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual([row['id'] for row in response.data], [str(self.own.pk)])
        self.assertEqual(len(self.request({}).data), 2)

    def test_own_upload_recovery_requires_authentication(self):
        self.assertIn(self.request({'mine': '1'}, authenticated=False).status_code, (401, 403))

    def test_compression_survives_production_route_guard_wrapping(self):
        import gzip
        import json
        from django.urls import path
        from apps.rbac.route_guard import secure_module_endpoints
        from apps.rbac.models import Module, Permission, Organization, UserProfile
        from .views import workbook_preview
        organization, _ = Organization.objects.get_or_create(code='PREVIEW', defaults={'name': 'Preview test'})
        profile, _ = UserProfile.objects.get_or_create(user=self.owner, defaults={'organization': organization})
        profile.status = 'active'
        profile.save(update_fields=['status'])
        self.owner.refresh_from_db()
        module, _ = Module.objects.get_or_create(code='spec_customization', defaults={'name': 'Spec Customization'})
        Permission.objects.get_or_create(code='spec_customization.read',
                                         defaults={'name': 'Read specs', 'module': module, 'action': 'read'})
        request = APIRequestFactory().get(
            f'/api/v1/spec-customization/paper-spec/jobs/{self.own.pk}/workbook/',
            HTTP_ACCEPT_ENCODING='gzip')
        force_authenticate(request, user=self.owner)
        patterns = [path('api/v1/spec-customization/paper-spec/jobs/<uuid:job_id>/workbook/', workbook_preview)]
        secure_module_endpoints(patterns)
        expected = {'sheets': [], 'literal_source': 'source evidence ' * 1000}
        with patch('apps.spec_customization.views.build_preview', return_value=expected):
            response = patterns[0].callback(request, job_id=self.own.pk)
        response.render()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Encoding'], 'gzip')
        self.assertEqual(json.loads(gzip.decompress(response.content)), expected)
        anonymous = APIRequestFactory().get(request.path, HTTP_ACCEPT_ENCODING='gzip')
        denied = patterns[0].callback(anonymous, job_id=self.own.pk)
        self.assertIn(denied.status_code, (401, 403))
