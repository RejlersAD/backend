"""Project-scoped extraction API for the deterministic P&ID → IO List flow
(alignment spec §8). Registered under the existing instrument-io-workflow
URL tree so the module RBAC ('instrument_io_list') applies unchanged.

Endpoints (all project-scoped):
  POST /projects/<pid>/extraction-files/            multipart upload (auto kind)
  GET  /projects/<pid>/extraction-files/            slot summaries
  PUT  /projects/<pid>/extraction-legend/           { legendId, rules }
  PUT  /projects/<pid>/extraction-doc-info/         { docNo, revision, ... }
  POST /projects/<pid>/extract/                     { useAi } → { extractionId, aiEstimate }
  GET  /projects/<pid>/extractions/<eid>/           ?tab=io|service|lines|setpoints|untagged|vendor|conflicts&sheet=
  GET  /projects/<pid>/extractions/<eid>/status/    job polling
  POST /projects/<pid>/extractions/<eid>/decisions/ ServiceDecision[]
  POST /projects/<pid>/extractions/<eid>/ai/<loop_key>/  BYOK suggestion
  GET  /projects/<pid>/extractions/<eid>/export/?format=iolist|review
"""
import logging
import os
import re
import threading

from django.db import close_old_connections
from django.http import FileResponse
from rest_framework import status
from rest_framework.decorators import api_view, parser_classes, permission_classes
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from .models import IOListDocument, IOListProject
from .models_extraction import (IOListAISuggestion, IOListExtraction,
                                IOListServiceDecision)

logger = logging.getLogger(__name__)


def _project_or_404(request, project_id):
    return IOListProject.objects.filter(id=project_id).first()


def _detect_kind(filename, file_bytes):
    """pdf → P&ID or I/O List by content; xlsx → vendor DB or IO list baseline
    by sheet names."""
    name = filename.lower()
    if name.endswith('.pdf'):
        import fitz
        doc = fitz.open(stream=file_bytes, filetype='pdf')
        try:
            text = '\n'.join(doc.load_page(i).get_text('text') or ''
                             for i in range(min(doc.page_count, 2)))
            has_smart = 'IN#' in text
            has_table = 'TAG_Number' in text or 'Loop_Name' in text
            return 'iolist' if (has_table and not has_smart) else 'pid'
        finally:
            doc.close()
    if name.endswith(('.xlsx', '.xlsm')):
        import openpyxl
        import io as _io
        wb = openpyxl.load_workbook(_io.BytesIO(file_bytes), read_only=True)
        try:
            if 'IO_DATABASE' in wb.sheetnames:
                return 'vendor'
            return 'iolist-xlsx'
        finally:
            wb.close()
    return None


def _document_summary(doc):
    return {
        'id': str(doc.id), 'name': doc.document_number or (doc.pdf_file.name.split('/')[-1]),
        'documentType': doc.document_type,
        'vendorSystem': (doc.extraction_stats or {}).get('vendor_system', ''),
        'uploadedAt': doc.created_at.isoformat() if doc.created_at else None,
    }


@api_view(['POST'])
@parser_classes([MultiPartParser, FormParser])
@permission_classes([IsAuthenticated])
def upload_extraction_file(request, project_id):
    """§3 slot upload. Auto-detects P&ID / IO-list / vendor and stores the
    file as an IOListDocument with the matching document_type."""
    project = _project_or_404(request, project_id)
    if project is None:
        return Response({'error': 'Project not found.'}, status=status.HTTP_404_NOT_FOUND)
    upload = request.FILES.get('file')
    if not upload:
        return Response({'error': 'No file supplied.'}, status=status.HTTP_400_BAD_REQUEST)
    data = upload.read()
    kind = _detect_kind(upload.name, data)
    if kind is None:
        return Response({'error': 'Unsupported file type — upload PDF or .xlsx.'},
                        status=status.HTTP_400_BAD_REQUEST)

    doc_type = {'pid': 'pid_drawing', 'iolist': 'io_list', 'iolist-xlsx': 'io_list',
                'vendor': 'vendor_db'}[kind]
    import hashlib
    from django.core.files.base import ContentFile
    doc = IOListDocument(
        project=project, uploaded_by=request.user, document_type=doc_type,
        document_number=upload.name, status='uploaded',
        pdf_sha256=hashlib.sha256(data).hexdigest(),
    )
    doc.pdf_file.save(upload.name, ContentFile(data))
    doc.save()

    summary = {'id': str(doc.id), 'kind': kind, 'name': upload.name}
    if kind == 'pid':
        from .services.pid_io import spans
        import fitz
        import tempfile, os
        tmp = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False)
        try:
            tmp.write(data)
            tmp.close()
            pages = spans.read_document(tmp.name)
            smart_in = sum(len([s for s in p.smart_tags if s.text.startswith('IN#')]) for p in pages)
            smart_eq = sum(len([s for s in p.smart_tags if s.text.startswith('EQ#')]) for p in pages)
            drawings = []
            for p in pages:
                m = re.findall(r'\d{2}-PP-152-\d{5}(?:-P2)?', p.text)
                if m:
                    drawings.append(m[-1])
            summary.update({
                'sheets': len(pages), 'drawings': drawings,
                'pdfType': 'Vector — text layer + smart-tag layer' if smart_in else 'Vector — text layer',
                'smartTags': {'instrument': smart_in, 'equipment': smart_eq},
            })
        finally:
            os.unlink(tmp.name)
    elif kind == 'iolist-xlsx':
        import tempfile, os
        from .services.pid_io import baseline
        tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        try:
            tmp.write(data)
            tmp.close()
            info = baseline.read_baseline(tmp.name)
            if 'error' in info:
                summary['error'] = info['error']
            else:
                summary.update({
                    'document': info['doc_no'], 'revision': info['revision'],
                    'sheet': info['sheet_name'], 'rows': len(info['rows']),
                    'columns': len(info['headers']), 'systems': info['systems'],
                })
        finally:
            os.unlink(tmp.name)
    elif kind == 'vendor':
        import tempfile, os
        from .services.pid_io import vendor as vendor_mod
        tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        try:
            tmp.write(data)
            tmp.close()
            db = vendor_mod.read_vendor_db(tmp.name)
            if isinstance(db, dict) and 'error' in db:
                summary['error'] = db['error']
            else:
                # DCS and ESD are two distinct, both-required databases; the
                # detected system rides on the document for ordering.
                doc.extraction_stats = {'vendor_system': db.system}
                doc.save(update_fields=['extraction_stats', 'updated_at'])
                summary.update({'document': db.name, 'liveTags': db.live_count,
                                'vendorSystem': db.system})
        finally:
            os.unlink(tmp.name)
    return Response(summary, status=status.HTTP_201_CREATED)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def list_extraction_files(request, project_id):
    project = _project_or_404(request, project_id)
    if project is None:
        return Response({'error': 'Project not found.'}, status=status.HTTP_404_NOT_FOUND)
    docs = project.documents.filter(
        document_type__in=('pid_drawing', 'io_list', 'vendor_db'))
    return Response({'files': [_document_summary(d) for d in docs]})


def _latest_extraction(project):
    return project.extractions.order_by('-created_at').first()


@api_view(['PUT'])
@permission_classes([IsAuthenticated])
def put_legend(request, project_id):
    project = _project_or_404(request, project_id)
    if project is None:
        return Response({'error': 'Project not found.'}, status=status.HTTP_404_NOT_FOUND)
    ext = _latest_extraction(project) or IOListExtraction(project=project, created_by=request.user)
    ext.stage_state = {**(ext.stage_state or {}), '2': 'complete',
                       'legendId': request.data.get('legendId')}
    ext.legend_rules = request.data.get('rules', {})
    ext.save()
    return Response({'ok': True, 'stageState': ext.stage_state})


@api_view(['PUT'])
@permission_classes([IsAuthenticated])
def put_doc_info(request, project_id):
    project = _project_or_404(request, project_id)
    if project is None:
        return Response({'error': 'Project not found.'}, status=status.HTTP_404_NOT_FOUND)
    ext = _latest_extraction(project) or IOListExtraction(project=project, created_by=request.user)
    ext.doc_number = str(request.data.get('docNo', ext.doc_number))[:128]
    ext.revision = str(request.data.get('revision', ext.revision))[:32]
    ext.rejlers_number = str(request.data.get('rejlersNumber', ext.rejlers_number))[:128]
    ext.stage_state = {**(ext.stage_state or {}), '3': 'complete'}
    ext.save()
    return Response({'ok': True, 'docNumber': ext.doc_number, 'revision': ext.revision})


def _file_path(doc):
    """Return a local file path for a stored document. S3-backed storage
    doesn't support absolute paths (.path raises), so stream the bytes into
    a NamedTemporaryFile and hand the engine a real local path."""
    if not doc or not doc.pdf_file:
        return None
    import tempfile
    suffix = os.path.splitext(doc.pdf_file.name)[1] or '.bin'
    tmp = tempfile.NamedTemporaryFile(suffix=suffix, delete=False)
    try:
        with doc.pdf_file.open('rb') as fh:
            for chunk in fh.chunks():
                tmp.write(chunk)
        tmp.close()
        return tmp.name
    except Exception:
        try:
            os.unlink(tmp.name)
        except OSError:
            pass
        raise


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def start_extraction(request, project_id):
    """Runs the deterministic engine in a worker thread; poll
    …/extractions/<id>/status/ for progress. useAi only gates the BYOK step
    (the drawing rules always run first and cost nothing)."""
    project = _project_or_404(request, project_id)
    if project is None:
        return Response({'error': 'Project not found.'}, status=status.HTTP_404_NOT_FOUND)
    docs = list(project.documents.filter(
        document_type__in=('pid_drawing', 'io_list', 'vendor_db')).order_by('-created_at'))
    pid_doc = next((d for d in docs if d.document_type == 'pid_drawing'), None)
    if pid_doc is None:
        return Response({'error': 'Upload a P&ID drawing first (Stage 1).'},
                        status=status.HTTP_400_BAD_REQUEST)
    # EVERY uploaded P&ID is processed (multi-file upload in Stage 1)
    pid_docs = [d for d in docs if d.document_type == 'pid_drawing']
    iolist_doc = next((d for d in docs if d.document_type == 'io_list'
                       and d.pdf_file.name.lower().endswith(('.xlsx', '.xlsm'))), None)
    # DCS first, then ESD — GDS rows live in the DCS database (spec §6.8).
    vendor_docs = sorted(
        (d for d in docs if d.document_type == 'vendor_db'),
        key=lambda d: 0 if (d.extraction_stats or {}).get('vendor_system') == 'DCS' else 1)

    # mark the previous extraction outdated (§2 re-check rule)
    project.extractions.filter(status=IOListExtraction.STATUS_COMPLETED).update(is_outdated=True)

    ext = IOListExtraction.objects.create(
        project=project, pid_document=pid_doc, iolist_document=iolist_doc,
        vendor_documents=[str(d.id) for d in vendor_docs],
        use_ai=bool(request.data.get('useAi')), created_by=request.user,
        status=IOListExtraction.STATUS_RUNNING,
    )

    def _run(extraction_id):
        close_old_connections()
        try:
            from .services.pid_io import engine
            extraction = IOListExtraction.objects.get(id=extraction_id)
            # every P&ID uploaded to the project is processed
            pid_paths = [_file_path(d) for d in IOListDocument.objects.filter(
                project=extraction.project, document_type='pid_drawing')]
            result = engine.run_extraction(
                pid_paths,
                baseline_path=_file_path(extraction.iolist_document),
                vendor_paths=[_file_path(d) for d in
                              IOListDocument.objects.filter(id__in=extraction.vendor_documents)],
            )
            extraction.summary = result['summary']
            extraction.tags_payload = result['tags']
            extraction.lines_payload = result['lines']
            extraction.untagged_payload = result['untagged']
            extraction.setpoints_payload = result['setpoints']
            extraction.vendor_payload = result['vendor']
            extraction.cause_notes_payload = result.get('cause_notes', [])
            extraction.alarmtrip_payload = result.get('alarmtrip', {})
            extraction.warnings = result['warnings']
            extraction.status = IOListExtraction.STATUS_COMPLETED
            extraction.stage_state = {**(extraction.stage_state or {}), '1': 'complete'}
            extraction.save()
        except Exception as exc:
            logger.exception('io-list extraction failed')
            extraction = IOListExtraction.objects.get(id=extraction_id)
            extraction.status = IOListExtraction.STATUS_FAILED
            extraction.error = str(exc)
            extraction.save(update_fields=['status', 'error', 'updated_at'])
        finally:
            close_old_connections()

    threading.Thread(target=_run, args=(ext.id,), daemon=True).start()
    return Response({'extractionId': str(ext.id), 'aiEstimate': None},
                    status=status.HTTP_202_ACCEPTED)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def extraction_status(request, project_id, extraction_id):
    ext = IOListExtraction.objects.filter(id=extraction_id, project_id=project_id).first()
    if not ext:
        return Response({'error': 'Extraction not found.'}, status=status.HTTP_404_NOT_FOUND)
    return Response({'status': ext.status, 'error': ext.error, 'summary': ext.summary,
                     'isOutdated': ext.is_outdated})


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def list_extractions(request, project_id):
    """Extraction history for the project — newest first; past results stay
    reopenable, replaced ones are flagged outdated."""
    project = _project_or_404(request, project_id)
    if project is None:
        return Response({'error': 'Project not found.'}, status=status.HTTP_404_NOT_FOUND)
    out = []
    for ext in project.extractions.all()[:25]:
        s = ext.summary or {}
        out.append({
            'id': str(ext.id), 'status': ext.status, 'isOutdated': ext.is_outdated,
            'createdAt': ext.created_at.isoformat(),
            'createdBy': getattr(ext.created_by, 'email', '') or '',
            'docNumber': ext.doc_number, 'revision': ext.revision,
            'useAi': ext.use_ai,
            'summary': {k: s.get(k) for k in (
                'tags', 'fieldIO', 'derived', 'withLine', 'setpoints', 'untagged',
                'vendorMatched', 'vendorCells', 'vendorConflicts', 'aiTags', 'elapsed_s')},
            'error': ext.error[:300] if ext.error else '',
        })
    return Response({'extractions': out})


@api_view(['PATCH'])
@permission_classes([IsAuthenticated])
def patch_extraction_row(request, project_id, extraction_id, tag):
    """Verify/edit one extracted row. Only allowlisted fields are writable;
    every edit is stamped (who/when) on the row itself. Confirming a row is a
    decision, stored via the decisions endpoint — not here."""
    ext = IOListExtraction.objects.filter(id=extraction_id, project_id=project_id).first()
    if not ext:
        return Response({'error': 'Extraction not found.'}, status=status.HTTP_404_NOT_FOUND)
    allowed = {'service', 'line', 'ioType', 'setpoint'}
    updates = {k: v for k, v in request.data.items() if k in allowed}
    if not updates:
        return Response({'error': f'Allowed fields: {sorted(allowed)}'},
                        status=status.HTTP_400_BAD_REQUEST)

    from django.utils import timezone
    stamp = {'editedBy': getattr(request.user, 'email', ''), 'editedAt': timezone.now().isoformat()}
    row = None
    for t in ext.tags_payload:
        if t.get('tag') == tag:
            row = t
            break
    if row is None:
        return Response({'error': 'Row not found.'}, status=status.HTTP_404_NOT_FOUND)
    for field in ('service', 'line', 'ioType'):
        if field in updates:
            row[field] = str(updates[field])[:255]
    if 'setpoint' in updates and isinstance(updates['setpoint'], dict):
        sp = updates['setpoint']
        point = str(sp.get('point') or '')[:16]
        for existing in ext.setpoints_payload:
            if existing.get('tag') == tag and existing.get('point') == point:
                existing['value'] = str(sp.get('value') or '')[:32]
                existing['editedBy'] = stamp['editedBy']
                existing['editedAt'] = stamp['editedAt']
                break
        else:
            ext.setpoints_payload.append({
                'tag': tag, 'point': point, 'value': str(sp.get('value') or '')[:32],
                'unit': str(sp.get('unit') or ''), 'label': 'manual entry',
                'column': str(sp.get('column') or 'Set point'), 'sheet': row.get('page', 0) + 1,
                'evidence': 'Edited by reviewer', **stamp,
            })
    row.update(stamp)
    ext.save(update_fields=['tags_payload', 'setpoints_payload', 'updated_at'])
    return Response({'ok': True, 'row': row})


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_extraction(request, project_id, extraction_id):
    """Tabbed result payload (§7.3). ?sheet= filters by sheet number."""
    ext = IOListExtraction.objects.filter(id=extraction_id, project_id=project_id).first()
    if not ext:
        return Response({'error': 'Extraction not found.'}, status=status.HTTP_404_NOT_FOUND)
    if ext.status != IOListExtraction.STATUS_COMPLETED:
        return Response({'status': ext.status, 'error': ext.error}, status=status.HTTP_202_ACCEPTED)

    tab = request.query_params.get('tab', 'io')
    sheet = request.query_params.get('sheet')
    decisions = {d.tag: d.decision for d in ext.decisions.all()}
    suggestions = {s.loop_key: {'suggestion': s.suggestion, 'status': s.status}
                   for s in ext.ai_suggestions.all()}

    def _sheet(rows):
        if not sheet:
            return rows
        return [r for r in rows if str(r.get('sheet') or (r.get('page', 0) + 1)) == str(sheet)]

    tags = _sheet(ext.tags_payload)
    if tab == 'io':
        rows = [t for t in tags if t.get('ioClass') in ('AI', 'AO', 'DI', 'DO') or t.get('derived')]
        for t in rows:
            t['decision'] = decisions.get(t['tag'])
            t['ai'] = suggestions.get(t['loopKey'])
        return Response({'tab': tab, 'summary': ext.summary, 'rows': rows})
    if tab == 'service':
        rows = [t for t in tags if not t.get('derived')]
        for t in rows:
            t['decision'] = decisions.get(t['tag'])
            t['ai'] = suggestions.get(t['loopKey'])
        return Response({'tab': tab, 'rows': rows})
    if tab == 'lines':
        return Response({'tab': tab, 'rows': _sheet(ext.lines_payload)})
    if tab == 'setpoints':
        return Response({'tab': tab, 'rows': _sheet(ext.setpoints_payload)})
    if tab == 'untagged':
        return Response({'tab': tab, 'rows': _sheet(ext.untagged_payload)})
    if tab == 'vendor':
        return Response({'tab': tab, **ext.vendor_payload})
    if tab == 'conflicts':
        return Response({'tab': tab, 'rows': ext.vendor_payload.get('conflicts', [])})
    if tab == 'alarmtrip':
        from .services.pid_io import alarmtrip
        scope = request.query_params.get('scope', 'pid')
        payload = alarmtrip.for_extraction(ext, scope)
        return Response({'tab': tab, **payload})
    return Response({'error': f'Unknown tab {tab!r}.'}, status=status.HTTP_400_BAD_REQUEST)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def post_decisions(request, project_id, extraction_id):
    ext = IOListExtraction.objects.filter(id=extraction_id, project_id=project_id).first()
    if not ext:
        return Response({'error': 'Extraction not found.'}, status=status.HTTP_404_NOT_FOUND)
    items = request.data if isinstance(request.data, list) else request.data.get('decisions', [])
    saved = 0
    for item in items:
        tag = str(item.get('tag') or '').strip()
        decision = str(item.get('decision') or '').strip()
        if not tag or decision not in ('keep-io-list', 'use-generated', 'ask-ai', 'confirm'):
            continue
        IOListServiceDecision.objects.update_or_create(
            extraction=ext, tag=tag,
            defaults={'decision': decision, 'decided_by': request.user})
        saved += 1
    return Response({'ok': True, 'saved': saved})


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def ai_suggest(request, project_id, extraction_id, loop_key):
    """§7.5 — one BYOK call per low-confidence loop. The suggestion stays
    'suggested' until the user confirms; it never overwrites anything."""
    ext = IOListExtraction.objects.filter(id=extraction_id, project_id=project_id).first()
    if not ext:
        return Response({'error': 'Extraction not found.'}, status=status.HTTP_404_NOT_FOUND)
    members = [t for t in ext.tags_payload if t.get('loopKey') == loop_key]
    if not members:
        return Response({'error': 'Loop not found in this extraction.'},
                        status=status.HTTP_404_NOT_FOUND)

    from apps.core.ai_consumer_clients import (AIProviderOperationError,
                                               lazy_provider_client)
    import json as _json

    context = {
        'loop': loop_key,
        'members': [{'tag': t['tag'], 'system': t['system'], 'ioType': t['ioType'],
                     'line': t.get('line') or '', 'equipment': t.get('equipment') or '',
                     'serviceEvidence': t.get('serviceEvidence') or ''} for t in members],
    }
    prompt = (
        'You are an instrument engineer completing an IO list. Given the loop members '
        'and rule-based evidence from a P&ID, propose the SERVICE text and LINE NUMBER. '
        'Return ONLY JSON: {"service": str, "line": str, "reason": str}.\n'
        f'Loop data: {_json.dumps(context)}')
    try:
        client = lazy_provider_client('anthropic', __import__('anthropic').Anthropic,
                                      api_key=lambda: '', timeout=60)
        resp = client.messages.create(model='claude-haiku-4-5', max_tokens=512,
                                      messages=[{'role': 'user', 'content': prompt}])
        text = ''.join(b.text for b in resp.content if getattr(b, 'type', None) == 'text')
        suggestion = _json.loads(text[text.find('{'):text.rfind('}') + 1] or '{}')
    except AIProviderOperationError as exc:
        return Response({'error': str(exc)}, status=status.HTTP_502_BAD_GATEWAY)
    except Exception as exc:
        logger.exception('io-list AI suggestion failed')
        return Response({'error': f'AI suggestion failed: {exc}'},
                        status=status.HTTP_502_BAD_GATEWAY)

    obj, _ = IOListAISuggestion.objects.update_or_create(
        extraction=ext, loop_key=loop_key,
        defaults={'suggestion': suggestion, 'created_by': request.user})
    return Response({'loopKey': loop_key, 'suggestion': suggestion, 'status': obj.status})


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def export_extraction(request, project_id, extraction_id):
    ext = IOListExtraction.objects.filter(id=extraction_id, project_id=project_id).first()
    if not ext:
        return Response({'error': 'Extraction not found.'}, status=status.HTTP_404_NOT_FOUND)
    fmt = request.query_params.get('fmt') or request.query_params.get('format') or 'iolist'
    try:
        from .services.pid_io import export as export_mod
        path = export_mod.build_workbook(ext, fmt)
        # Read into memory: the temp file must not be tied to the response's
        # streaming lifecycle (test clients and some servers close it early).
        import io as _io
        import os as _os
        with open(path, 'rb') as fh:
            payload = _io.BytesIO(fh.read())
        _os.unlink(path)
        if fmt == 'review':
            filename = f'{ext.doc_number or "io-list"}-review.xlsx'
        else:
            doc = ext.doc_number or 'io-list'
            rev = ext.revision or ''
            filename = f'{doc}_IO_List_{rev}.xlsx' if rev else f'{doc}_IO_List.xlsx'
        return FileResponse(payload, as_attachment=True, filename=filename)
    except Exception as exc:
        logger.exception('io-list export failed')
        return Response({'error': f'Export failed: {exc}'},
                        status=status.HTTP_500_INTERNAL_SERVER_ERROR)
