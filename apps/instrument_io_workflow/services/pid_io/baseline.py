"""IO list baseline workbook reader (§3 slot 2).

The reference baseline (13-IE-363-70001) stores the IO list in the
'Honeywell - H2 Point Database' sheet: row 1 carries column numbers, row 2
the 53 headers, data starts at row 3. The sheet is located by the row
containing `TAG_Number`, never by name.
"""
import re

import openpyxl

TAG_HEADER = 'TAG_Number'


def _norm_header(h):
    return re.sub(r'\s+', ' ', str(h or '').replace('\n', ' ')).strip()


def find_io_sheet(wb):
    """The sheet whose first rows contain a TAG_Number header cell."""
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=6, values_only=True):
            if any(_norm_header(c) == TAG_HEADER for c in row if c):
                return ws
    return None


def read_baseline(path):
    """Returns {headers, rows, doc_no, revision, sheet_name, systems} — rows
    are dicts keyed by the normalised headers, tags normalised (uppercase,
    spaces removed) into '_tag_norm'."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = find_io_sheet(wb)
        if ws is None:
            return {'error': 'No sheet with a TAG_Number header row found.'}
        rows_iter = ws.iter_rows(values_only=True)
        header_row = None
        headers = []
        raw = []
        for row in rows_iter:
            if header_row is None:
                if any(_norm_header(c) == TAG_HEADER for c in row if c):
                    headers = [_norm_header(c) for c in row]
                    header_row = True
                continue
            raw.append(row)
        records = []
        for row in raw:
            rec = {headers[i]: (row[i] if i < len(row) else None) for i in range(len(headers)) if headers[i]}
            tag = rec.get(TAG_HEADER)
            if tag is None or str(tag).strip() == '':
                continue
            rec['_tag_norm'] = re.sub(r'\s+', '', str(tag)).upper()
            records.append(rec)
        systems = {}
        for rec in records:
            sysname = str(rec.get('System') or '').strip() or '?'
            systems[sysname] = systems.get(sysname, 0) + 1
        # document number / revision from the file name (13-IE-363-70001_Rev.Q1.xlsx)
        m = re.match(r'([0-9A-Z-]+?)(?:[_ ]Rev\.?([A-Z0-9]+))?\.xlsx$',
                     path.replace('\\', '/').split('/')[-1], re.IGNORECASE)
        return {
            'headers': headers, 'rows': records, 'sheet_name': ws.title,
            'doc_no': m.group(1) if m else '', 'revision': m.group(2) if m and m.group(2) else '',
            'systems': systems,
        }
    finally:
        wb.close()
