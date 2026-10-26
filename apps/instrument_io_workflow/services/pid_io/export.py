"""§7.6 — workbook exports for an IOListExtraction.

`iolist`: the 53-column baseline layout, vendor-filled cells in green,
conflict cells in amber, plus `Set point` (after High_High_Alarm_Value),
`Vendor source` and `Conflicts` columns.

`review`: README, PID_TAGS, LINE_TAGS, SETPOINTS, SERVICE_CHECK,
IOL_NOT_ON_PID, UNTAGGED, IO_LIST_AUTOFILLED, VENDOR_CONFLICTS,
NOT_IN_VENDOR.
"""
import tempfile

import openpyxl
from openpyxl.styles import PatternFill

_GREEN = PatternFill('solid', start_color='E8F5E9')   # vendor-filled
_AMBER = PatternFill('solid', start_color='FFF3E0')   # conflict

_ALARM_COLS = ('Low_Low_Alarm_Value', 'Low_Alarm_Value',
               'High_Alarm_Value', 'High_High_Alarm_Value')


def _baseline_rows(extraction):
    vendor = extraction.vendor_payload or {}
    return vendor.get('baseline', {}), vendor.get('fills', {}), vendor.get('conflicts', []), vendor.get('baselineRows', [])


def build_workbook(extraction, fmt):
    wb = openpyxl.Workbook()
    if fmt == 'review':
        _review_sheets(wb, extraction)
    else:
        _iolist_sheet(wb, extraction)
    tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
    wb.save(tmp.name)
    return tmp.name


def _iolist_sheet(wb, extraction):
    ws = wb.active
    ws.title = 'IO List'
    baseline, fills, conflicts, rows = _baseline_rows(extraction)
    headers = list(baseline.get('headers') or [])
    if 'Set point' not in headers:
        anchor = 'High_High_Alarm_Value' if 'High_High_Alarm_Value' in headers else None
        headers.insert(headers.index(anchor) + 1 if anchor else len(headers), 'Set point')
    headers += ['Vendor source', 'Conflicts']
    ws.append(headers)

    # set points per tag → column value
    sp_by_tag = {}
    for sp in extraction.setpoints_payload:
        if sp.get('value'):
            sp_by_tag.setdefault(sp['tag'], []).append(sp)

    conflict_by_tag = {}
    for c in conflicts:
        conflict_by_tag.setdefault(c['tag'], []).append(c)

    for rec in rows:
        tag = rec.get('_tag_norm', '')
        fill = fills.get(tag, {})
        cells = fill.get('cells', {})
        row = []
        for h in headers:
            if h == 'Set point':
                sps = sp_by_tag.get(tag, [])
                row.append('; '.join(f"{s['point']} = {s['value']} {s['unit']}"
                                     for s in sps))
            elif h == 'Vendor source':
                row.append(fill.get('source', ''))
            elif h == 'Conflicts':
                row.append('; '.join(f"{c['field']}: {c['ioList'] if 'ioList' in c else c.get('io_list')} vs {c['vendor']}"
                                     for c in conflict_by_tag.get(tag, []) if c['type'] == 'engineering'))
            elif h in cells:
                row.append(cells[h])
            else:
                v = rec.get(h)
                row.append('' if v is None else v)
        ws.append(row)
        r = ws.max_row
        for idx, h in enumerate(headers, start=1):
            if h in cells:
                ws.cell(row=r, column=idx).fill = _GREEN
            if any(c['field'] == h for c in conflict_by_tag.get(tag, []) if c['type'] == 'engineering'):
                ws.cell(row=r, column=idx).fill = _AMBER


def _review_sheets(wb, extraction):
    s = extraction.summary or {}
    ws = wb.active
    ws.title = 'README'
    ws.append(['IO List extraction — review workbook'])
    for k, v in s.items():
        ws.append([k, v])

    def _sheet(name, headers, rows):
        w = wb.create_sheet(name)
        w.append(headers)
        for r in rows:
            w.append(r)

    tags = extraction.tags_payload
    _sheet('PID_TAGS',
           ['Tag', 'Sheet', 'System', 'Bubble', 'IO class', 'Line', 'Line source',
            'Service', 'Confidence', 'Evidence', 'Confirmed by smart tag'],
           [[t['tag'], t['page'] + 1, t['system'], t['bubble'], t.get('ioType', ''),
             t.get('line', ''), t.get('lineSource', ''), t.get('service', ''),
             t.get('serviceConfidence', 0), t.get('serviceEvidence', ''),
             'yes' if t.get('confirmedBySmartTag') else 'no'] for t in tags])
    _sheet('LINE_TAGS',
           ['Sheet', 'Line number', 'Anchored by', 'Tags tapped on the line',
            'Loop members (inherited)'],
           [[l['sheet'], l['line'], l['anchoredBy'], ', '.join(l['tagsOwn']),
             ', '.join(l['tagsLoop'])] for l in extraction.lines_payload])
    _sheet('SETPOINTS',
           ['Sheet', 'Tag', 'Point', 'Set point (P&ID)', 'Scale label', 'IO list column'],
           [[sp['sheet'], sp['tag'], sp['point'],
             f"{sp['value']} {sp['unit']}" if sp['value'] else 'Letter only — C&E',
             sp['label'], sp['column']] for sp in extraction.setpoints_payload])
    _sheet('SERVICE_CHECK',
           ['Tag', 'Generated from drawing', 'Confidence', 'Why'],
           [[t['tag'], t.get('service', ''), t.get('serviceConfidence', 0),
             t.get('serviceEvidence', '')] for t in tags if t.get('service')])
    vendor = extraction.vendor_payload or {}
    on_pid = {t['norm'] for t in tags}
    baseline_rows = vendor.get('baselineRows') or []
    _sheet('IOL_NOT_ON_PID', ['Tag'], [[r.get('_tag_norm')] for r in baseline_rows
                                       if r.get('_tag_norm') and r['_tag_norm'] not in on_pid])
    _sheet('UNTAGGED', ['Sheet', 'Text in bubble', 'Classified as', 'Nearest tag', 'Distance pt'],
           [[u['sheet'], u['text'], u['kind'], u['nearestTag'], u['distancePt']]
            for u in extraction.untagged_payload])
    fills = vendor.get('fills') or {}
    _sheet('IO_LIST_AUTOFILLED', ['Tag', 'Column', 'Value', 'Source'],
           [[tag, col, val, f.get('source', '')] for tag, f in fills.items()
            for col, val in f.get('cells', {}).items()])
    _sheet('VENDOR_CONFLICTS', ['Tag', 'Field', 'IO List Rev Q1', 'Vendor database', 'Type'],
           [[c['tag'], c['field'], c.get('io_list', ''), c['vendor'], c['type']]
            for c in vendor.get('conflicts', [])])
    _sheet('NOT_IN_VENDOR', ['Tag'], [[t] for t in vendor.get('noRecord', [])])
