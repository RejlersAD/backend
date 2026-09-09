"""
Loader for project-specific Pipe Class Conversation workbook records.

Soft-coded strategy:
- No hardcoded sheet schemas.
- Every non-empty row is stored as an ordered JSON cell list.
- Supports dry-run and idempotent reload for safe production updates.

Usage:
    python manage.py load_pipe_class_conversation --source /path/to/Pipe_Class_Conversation.xlsx --dry-run
    python manage.py load_pipe_class_conversation --source /path/to/Pipe_Class_Conversation.xlsx --reload
"""
from __future__ import annotations

from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from apps.valve_standards import constants as c
from apps.valve_standards.models import PipeClassConversationRecord, Standard


def _norm_cell(v):
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        return s if s else None
    return str(v)


class Command(BaseCommand):
    help = 'Load Pipe Class Conversation workbook rows into Valve Standards reference dataset.'

    def add_arguments(self, parser):
        parser.add_argument('--source', type=str, required=True, help='Path to Pipe_Class_Conversation.xlsx')
        parser.add_argument('--reload', action='store_true', help='Delete existing dataset rows before import')
        parser.add_argument('--dry-run', action='store_true', help='Parse and report only, without database writes')

    def handle(self, *args, **options):
        source = Path(options['source'])
        if not source.exists():
            raise CommandError(f'Source file not found: {source}')

        wb = load_workbook(source, data_only=True)

        counts = {'total_rows': 0, 'index_entry': 0, 'sheet_row': 0, 'branch_row': 0}
        records = []

        for ws in wb.worksheets:
            sheet_name = ws.title.strip()
            for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                cells = [_norm_cell(v) for v in row]
                if not any(v is not None for v in cells):
                    continue

                record_type = self._infer_record_type(sheet_name, row_idx, cells)
                primary_key_text = self._primary_key_text(cells)
                records.append({
                    'source_sheet': sheet_name[:64],
                    'source_row': row_idx,
                    'record_type': record_type,
                    'primary_key_text': primary_key_text[:128],
                    'cells': cells,
                })
                counts['total_rows'] += 1
                counts[record_type] = counts.get(record_type, 0) + 1

        if options['dry_run']:
            self.stdout.write(self.style.WARNING('DRY RUN: no rows written.'))
            self._print_summary(source, counts)
            return

        standard, _ = Standard.objects.get_or_create(
            code=c.B16_5_STANDARD_CODE,
            defaults={
                'title': c.B16_5_STANDARD_TITLE,
                'edition_year': c.B16_5_STANDARD_EDITION_YEAR,
            },
        )

        with transaction.atomic():
            if options['reload']:
                deleted, _ = PipeClassConversationRecord.objects.filter(
                    standard=standard,
                ).delete()
                self.stdout.write(f'Reload enabled: deleted {deleted} existing row(s).')

            objs = [
                PipeClassConversationRecord(
                    standard=standard,
                    source_sheet=r['source_sheet'],
                    source_row=r['source_row'],
                    record_type=r['record_type'],
                    primary_key_text=r['primary_key_text'],
                    cells=r['cells'],
                )
                for r in records
            ]
            PipeClassConversationRecord.objects.bulk_create(objs, batch_size=2000)

        self._print_summary(source, counts)
        self.stdout.write(self.style.SUCCESS(f'Inserted {len(records)} row(s) into {standard.code}.'))

    def _infer_record_type(self, sheet_name, row_idx, cells):
        first = (cells[0] or '').lower() if cells else ''
        if sheet_name.upper() == 'INDEX':
            if row_idx <= 4:
                return 'sheet_row'
            if first.startswith('source:'):
                return 'sheet_row'
            if first == 'piping class':
                return 'sheet_row'
            return 'index_entry'
        if sheet_name.upper() == 'BRANCH TABLES':
            return 'branch_row'
        return 'sheet_row'

    def _primary_key_text(self, cells):
        for val in cells[:3]:
            if val is not None:
                return str(val)
        return ''

    def _print_summary(self, source, counts):
        self.stdout.write(
            f'Parsed workbook: {source}\n'
            f"  total_rows: {counts['total_rows']}\n"
            f"  index_entry: {counts.get('index_entry', 0)}\n"
            f"  sheet_row: {counts.get('sheet_row', 0)}\n"
            f"  branch_row: {counts.get('branch_row', 0)}"
        )
