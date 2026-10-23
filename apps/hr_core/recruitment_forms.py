"""Scoped external forms and explicit HR invitation/review commands (HR-REC-02)."""
import hashlib
import io
import secrets
from datetime import timedelta
from urllib.parse import urlsplit

from django.conf import settings
from django.core.mail import EmailMessage, get_connection
from django.db import transaction
from django.http import HttpResponse
from django.utils import timezone
from rest_framework.exceptions import APIException, NotFound, PermissionDenied, ValidationError
from rest_framework.permissions import AllowAny, IsAuthenticated
from rest_framework.response import Response
from rest_framework.throttling import SimpleRateThrottle
from rest_framework.views import APIView

from .governance import audit
from .models import RecruitmentApplicationLink, RecruitmentCandidate, RecruitmentFormInvitation, RecruitmentFormSubmission, RecruitmentOpening
from .recruitment import RecruitmentHRPermission
from .recruitment_form_schema import DECLARATION, VERSION, schema_for, validate_form, display_sections


def digest(value):
    return hashlib.sha256(value.encode()).hexdigest()


def no_store(response):
    response['Cache-Control'] = 'no-store, private'
    response['Referrer-Policy'] = 'no-referrer'
    response['X-Robots-Tag'] = 'noindex, nofollow'
    return response


class PrivateResponseMixin:
    def finalize_response(self, request, response, *args, **kwargs):
        return no_store(super().finalize_response(request, response, *args, **kwargs))


class ClosedCVLink(APIException):
    status_code = 409
    default_detail = {'detail': 'This CV link has already received a submission and is now closed.', 'code': 'cv_link_closed', 'kind': 'cv'}


def require_active(invitation):
    if invitation.revoked_at or invitation.expires_at <= timezone.now():
        raise PermissionDenied('This invitation has expired or was revoked. Please contact HR.')
    if invitation.opening_id and not invitation.opening.is_open:
        raise PermissionDenied('This position is closed. Please contact HR.')
    if invitation.application_link_id:
        parent = RecruitmentApplicationLink.objects.select_for_update().get(pk=invitation.application_link_id)
        if parent.revoked_at or parent.expires_at <= timezone.now():
            raise PermissionDenied('This application link has expired or was revoked. Please contact HR.')
        if invitation.kind == 'cv' and invitation.state not in ('submitted', 'verified') and RecruitmentFormSubmission.objects.filter(invitation__application_link=parent).exists():
            raise ClosedCVLink()
    if invitation.kind == 'pre_onboarding' and not invitation.application_link_id and (not invitation.candidate_id or invitation.candidate.interview_outcome != 'suitable'):
        raise PermissionDenied('Pre-onboarding is not currently available. Please contact HR.')


def summary(invitation):
    return dict(id=invitation.pk, kind=invitation.kind, title=invitation.get_kind_display(),
                candidate=invitation.candidate_id, opening=invitation.opening_id,
                position=invitation.opening.title if invitation.opening_id else '', recipient=invitation.recipient,
                expires_at=invitation.expires_at, revoked=bool(invitation.revoked_at),
                state=invitation.state, delivery_status=invitation.delivery_status,
                correction_note=invitation.correction_note, draft_revision=invitation.draft_revision,
                application_link=invitation.application_link_id,
                versions=list(invitation.submissions.order_by('-version').values('id', 'version', 'created_at')))


def event(request, invitation, label):
    audit(actor=request.user, action='recruitment.form', object_type='recruitment_candidate',
          object_id=invitation.candidate_id, metadata={'label': label, 'invitation_id': invitation.pk, 'form_kind': invitation.kind}, request=request)


def send_email(recipient, subject, body):
    connection = get_connection(timeout=10)
    sent = EmailMessage(subject, body, settings.DEFAULT_FROM_EMAIL, [recipient], connection=connection).send(fail_silently=False)
    if sent != 1:
        raise RuntimeError('Email transport did not accept the message.')


def invitation_link(token):
    origin = getattr(settings, 'RECRUITMENT_FRONTEND_URL', settings.FRONTEND_URL).rstrip('/')
    parsed = urlsplit(origin)
    if parsed.scheme not in ('http', 'https') or not parsed.netloc or parsed.query or parsed.fragment:
        raise ValidationError({'detail': 'HR candidate form origin is not configured correctly.'})
    return f'{origin}/candidate-forms#invite={token}'


class HRFormsView(PrivateResponseMixin, APIView):
    permission_classes = [IsAuthenticated, RecruitmentHRPermission]


class RecruitmentFormSchemaView(HRFormsView):
    permission_action = 'read'

    def get(self, request, kind):
        if kind not in ('cv', 'pre_onboarding'):
            raise NotFound()
        return Response(schema_for(kind))


class RecruitmentInvitationsView(HRFormsView):
    def get(self, request):
        try:
            candidate = int(request.query_params.get('candidate', ''))
        except (ValueError, TypeError):
            raise ValidationError({'candidate': 'Choose a candidate.'})
        invitations = RecruitmentFormInvitation.objects.filter(candidate_id=candidate).select_related('opening', 'candidate').order_by('-created_at')
        return Response([summary(item) for item in invitations])

    @transaction.atomic
    def post(self, request):
        kind = request.data.get('kind')
        if kind not in ('cv', 'pre_onboarding'):
            raise ValidationError({'kind': 'Choose CV or pre-onboarding.'})
        try:
            candidate = RecruitmentCandidate.objects.select_for_update().get(pk=request.data.get('candidate'))
            opening = RecruitmentOpening.objects.select_for_update().get(pk=request.data.get('opening'), is_open=True)
        except (RecruitmentCandidate.DoesNotExist, RecruitmentOpening.DoesNotExist, ValueError, TypeError):
            raise ValidationError({'detail': 'Choose an existing candidate and an open position.'})
        if kind == 'pre_onboarding' and candidate.interview_outcome != 'suitable':
            raise ValidationError({'detail': 'Record a suitable interview outcome before inviting pre-onboarding.'})
        now = timezone.now()
        RecruitmentFormInvitation.objects.filter(candidate=candidate, opening=opening, kind=kind, revoked_at__isnull=True).update(revoked_at=now, session_hash='')
        token = secrets.token_urlsafe(32)
        link = invitation_link(token)
        invitation = RecruitmentFormInvitation.objects.create(candidate=candidate, opening=opening, owner=request.user, kind=kind,
            recipient=candidate.email, token_hash=digest(token), expires_at=now + timedelta(days=3))
        event(request, invitation, 'CV invitation created' if kind == 'cv' else 'Pre-onboarding invitation created')
        return Response({**summary(invitation), 'link': link, 'token': token}, status=201)


class RecruitmentInvitationCommandView(HRFormsView):
    permission_action = 'update'

    def post(self, request, pk):
        action = request.data.get('action')
        if action not in ('revoke', 'send'):
            raise ValidationError({'action': 'Choose send or revoke.'})
        with transaction.atomic():
            try:
                invitation = RecruitmentFormInvitation.objects.select_for_update(of=('self',)).select_related('candidate', 'opening').get(pk=pk)
            except RecruitmentFormInvitation.DoesNotExist:
                raise NotFound()
            if action == 'revoke':
                invitation.revoked_at = invitation.revoked_at or timezone.now()
                invitation.session_hash = ''
                invitation.save(update_fields=['revoked_at', 'session_hash', 'updated_at'])
                event(request, invitation, 'Form invitation revoked')
                return Response(summary(invitation))
            require_active(invitation)
            if not invitation.candidate_id or not invitation.recipient:
                raise ValidationError({'detail': 'This applicant has not submitted contact details yet.'})
            token = request.data.get('token', '')
            if not isinstance(token, str) or not secrets.compare_digest(digest(token), invitation.token_hash):
                raise ValidationError({'detail': 'Create a new link before sending this invitation.'})
            if invitation.delivery_status == 'sending':
                return Response({'detail': 'This invitation is already being sent.'}, status=409)
            if invitation.delivery_status == 'sent':
                return Response(summary(invitation))
            invitation.delivery_status = 'sending'
            invitation.save(update_fields=['delivery_status', 'updated_at'])
        # The attempt is durable before contacting an external email transport.
        try:
            send_email(invitation.recipient, f'Rejlers - {invitation.get_kind_display()}',
                f'Hello {invitation.candidate.name},\n\nHR invites you to complete the {invitation.get_kind_display()} for {invitation.opening.title if invitation.opening_id else "General application"}.\n\n{invitation_link(token)}\n\nThe link expires on {invitation.expires_at.isoformat()}. Please keep it private.\n\nRejlers HR')
        except Exception:
            RecruitmentFormInvitation.objects.filter(pk=pk, delivery_status='sending').update(delivery_status='failed')
            return Response({'detail': 'Email could not be sent. Check the configured HR email transport and retry; the private link remains available.'}, status=502)
        RecruitmentFormInvitation.objects.filter(pk=pk).update(delivery_status='sent')
        invitation.refresh_from_db()
        event(request, invitation, 'Form invitation email accepted by transport')
        return Response(summary(invitation))


class RecruitmentFormReviewView(HRFormsView):
    business_approval_actions = ('RecruitmentFormReviewView',)

    @property
    def permission_action(self):
        if self.request.method in ('GET', 'HEAD'):
            return 'read'
        return 'approve' if self.request.data.get('action') == 'verify' else 'update'

    def get(self, request, pk):
        # GET must use the read grant, irrespective of review authority.
        try:
            invitation = RecruitmentFormInvitation.objects.select_related('candidate', 'opening').get(pk=pk)
        except RecruitmentFormInvitation.DoesNotExist:
            raise NotFound()
        from .recruitment_cv import attachment_summary
        submissions = []
        for item in invitation.submissions.order_by('-version').prefetch_related('attachments'):
            submissions.append({key: getattr(item, key) for key in ('id', 'version', 'data', 'declaration_text', 'schema_version', 'created_at')})
            submissions[-1]['attachments'] = [attachment_summary(file) for file in item.attachments.all()]
        return Response({**summary(invitation), 'schema': schema_for(invitation.kind), 'submissions': submissions})

    @transaction.atomic
    def post(self, request, pk):
        try:
            invitation = RecruitmentFormInvitation.objects.select_for_update(of=('self',)).select_related('candidate', 'opening').get(pk=pk)
        except RecruitmentFormInvitation.DoesNotExist:
            raise NotFound()
        latest = invitation.submissions.order_by('-version').first()
        if not latest or request.data.get('version') != latest.version or invitation.state not in ('submitted', 'verified'):
            return Response({'detail': 'The submission changed. Reload before reviewing.'}, status=409)
        action = request.data.get('action')
        if action == 'verify':
            if (invitation.opening_id and not invitation.opening.is_open) or (invitation.kind == 'pre_onboarding' and invitation.candidate.interview_outcome != 'suitable'):
                raise ValidationError({'detail': 'The position or interview outcome is no longer eligible for verification.'})
            invitation.state = 'verified'
            invitation.correction_note = ''
        elif action == 'correction':
            note = request.data.get('note', '')
            if not isinstance(note, str) or not note.strip() or len(note) > 4000:
                raise ValidationError({'note': 'Provide a correction note under 4000 characters.'})
            invitation.state = 'correction_requested'
            invitation.correction_note = note.strip()
        else:
            raise ValidationError({'action': 'Choose verify or correction.'})
        invitation.save(update_fields=['state', 'correction_note', 'updated_at'])
        event(request, invitation, 'Form verified by HR' if action == 'verify' else 'Form correction requested by HR')
        return Response(summary(invitation))


class RecruitmentFormExportView(HRFormsView):
    permission_action = 'export'

    def get(self, request, pk, version):
        try:
            submission = RecruitmentFormSubmission.objects.select_related('invitation__candidate', 'invitation__opening').get(invitation_id=pk, version=version)
        except RecruitmentFormSubmission.DoesNotExist:
            raise NotFound()
        from docx import Document
        invitation = submission.invitation
        document = Document()
        document.add_heading('REJLERS', 0)
        document.add_heading(invitation.get_kind_display(), 1)
        photo = submission.attachments.filter(kind='photo').first()
        if photo:
            from docx.shared import Inches
            document.add_picture(io.BytesIO(bytes(photo.content)), width=Inches(1.2))
        document.add_paragraph(f'Candidate: {invitation.candidate.name}\nPosition: {invitation.opening.title if invitation.opening_id else "General application"}\nVersion: {version}\nSubmitted: {submission.created_at.isoformat()}')
        for section in display_sections(invitation.kind):
            document.add_heading(section['title'], 2)
            value = submission.data.get(section['key'], [] if section['repeatable'] else {})
            for index, row in enumerate(value if section['repeatable'] else [value]):
                if section['repeatable']: document.add_paragraph(f'Entry {index + 1}', style='Heading 3')
                for spec in section['fields']:
                    content = row.get(spec['key'], '')
                    document.add_paragraph(f"{spec['label']}: {'Yes' if content is True else 'No' if content is False else content}")
        if submission.declaration_text:
            document.add_heading('Company-template declaration', 2)
            document.add_paragraph(submission.declaration_text)
        document.add_paragraph('Typed names are acknowledgements, not verified digital signatures.')
        output = io.BytesIO()
        document.save(output)
        response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        response['Content-Disposition'] = f'attachment; filename="RAD-{invitation.kind}-CAN-{invitation.candidate_id:04d}-v{version}.docx"'
        event(request, invitation, 'Submitted form exported by HR')
        return no_store(response)


class FormThrottle(SimpleRateThrottle):
    rate = '20/min'
    scope = 'recruitment_forms'

    def get_cache_key(self, request, view):
        return self.cache_format % {'scope': self.scope, 'ident': self.get_ident(request)}


class PublicFormView(PrivateResponseMixin, APIView):
    authentication_classes = []
    permission_classes = [AllowAny]
    throttle_classes = [FormThrottle]

    def initial(self, request, *args, **kwargs):
        super().initial(request, *args, **kwargs)
        if request.method == 'POST' and not isinstance(request.data, dict):
            raise ValidationError({'detail': 'A JSON object is required.'})


class CandidateFormTemplateView(PublicFormView):
    """Blank company template only; never resolves an invitation or personal data."""

    def get(self, request, kind):
        if kind not in ('cv', 'pre_onboarding'):
            raise NotFound()
        return Response(schema_for(kind))


def pre_onboarding_verification(invitation, *, send_code=False, code='', email=None):
    now = timezone.now()
    if send_code:
        if invitation.otp_sent_at and now - invitation.otp_sent_at < timedelta(seconds=60):
            return Response({'detail': 'Please wait one minute before requesting another code.'}, status=429)
        from rest_framework.serializers import EmailField
        try:
            recipient = EmailField().run_validation(email if email is not None else invitation.recipient)
        except ValidationError:
            raise ValidationError({'detail': 'Enter a valid personal email address in the form.'})
        verification = f'{secrets.randbelow(1000000):06d}'
        try:
            send_email(recipient, 'Rejlers - pre-onboarding verification code', f'Your one-time verification code is {verification}. It expires in 10 minutes. Do not share this code.')
        except Exception:
            return Response({'detail': 'The verification email could not be sent. Please contact HR.'}, status=502)
        invitation.recipient = recipient
        invitation.verified_email = ''
        invitation.otp_hash = digest(verification)
        invitation.otp_expires_at = now + timedelta(minutes=10)
        invitation.otp_sent_at = now
        invitation.otp_attempts = 0
        invitation.save(update_fields=['recipient', 'verified_email', 'otp_hash', 'otp_expires_at', 'otp_sent_at', 'otp_attempts', 'updated_at'])
        return Response({'verification_required': True, 'code_sent': True})
    if not code:
        return Response({'verification_required': True})
    if not isinstance(code, str) or not invitation.otp_hash or not invitation.otp_expires_at or invitation.otp_expires_at <= now or invitation.otp_attempts >= 5:
        raise PermissionDenied('The verification code is expired or unavailable. Request a new code.')
    invitation.otp_attempts += 1
    invitation.save(update_fields=['otp_attempts', 'updated_at'])
    if not secrets.compare_digest(digest(code), invitation.otp_hash):
        return Response({'detail': 'Incorrect verification code.'}, status=403)
    invitation.otp_hash = ''
    invitation.verified_email = invitation.recipient
    invitation.save(update_fields=['otp_hash', 'verified_email', 'updated_at'])
    return None


class CandidateFormAccessView(PublicFormView):
    @transaction.atomic
    def post(self, request):
        token = request.data.get('token', '')
        if not isinstance(token, str) or not 32 <= len(token) <= 100:
            raise PermissionDenied('Invalid invitation. Please contact HR.')
        try:
            invitation = RecruitmentFormInvitation.objects.select_for_update(of=('self',)).select_related('candidate', 'opening').get(token_hash=digest(token))
        except RecruitmentFormInvitation.DoesNotExist:
            raise PermissionDenied('Invalid invitation. Please contact HR.')
        require_active(invitation)
        if invitation.kind == 'cv' and invitation.submissions.exists():
            return Response({'closed': True, 'kind': 'cv', 'state': 'submitted'})
        if invitation.kind == 'pre_onboarding' and (invitation.recipient or invitation.candidate_id or not invitation.application_link_id):
            result = pre_onboarding_verification(invitation, send_code=request.data.get('send_code') is True, code=request.data.get('code', ''))
            if result is not None:
                return result
        session = secrets.token_urlsafe(32)
        invitation.session_hash = digest(session)
        invitation.session_expires_at = min(timezone.now() + timedelta(hours=1), invitation.expires_at)
        invitation.save(update_fields=['session_hash', 'session_expires_at', 'otp_hash', 'updated_at'])
        return Response({'session': session})


def session_invitation(request):
    header = request.headers.get('Authorization', '')
    token = header[7:] if header.startswith('Bearer ') else ''
    if not 32 <= len(token) <= 100:
        raise PermissionDenied('Your form session has expired. Reopen your invitation link.')
    try:
        invitation = RecruitmentFormInvitation.objects.select_for_update(of=('self',)).select_related('candidate', 'opening').get(session_hash=digest(token), session_expires_at__gt=timezone.now())
    except RecruitmentFormInvitation.DoesNotExist:
        raise PermissionDenied('Your form session has expired. Reopen your invitation link.')
    require_active(invitation)
    if request.method != 'GET' and invitation.kind == 'cv' and invitation.submissions.exists():
        raise ClosedCVLink()
    return invitation


class CandidateFormSessionView(PublicFormView):
    class DraftThrottle(FormThrottle):
        scope = 'candidate_drafts'
        rate = '90/min'

    throttle_classes = [DraftThrottle]

    @transaction.atomic
    def get(self, request):
        invitation = session_invitation(request)
        if invitation.kind == 'cv' and invitation.submissions.exists():
            return Response({'closed': True, 'kind': 'cv', 'state': 'submitted'})
        from .recruitment_cv import latest_attachments, attachment_summary
        draft = dict(invitation.draft)
        if not draft and invitation.candidate_id:
            draft = {'candidate': {'name': invitation.candidate.name, 'email': invitation.recipient, 'phone': invitation.candidate.mobile}} if invitation.kind == 'cv' else {'contact': {'email': invitation.recipient, 'phone': invitation.candidate.mobile}, 'identity': {'nationality': invitation.candidate.nationality}, 'beneficiary': {'candidate_name': invitation.candidate.name, 'date': timezone.localdate().isoformat()}}
        if invitation.kind == 'cv':
            draft['candidate'] = dict(draft.get('candidate', {}))
            draft['candidate'].setdefault('position_applied', invitation.opening.title if invitation.opening_id else (invitation.candidate.position if invitation.candidate_id else ''))
            draft['candidate'].setdefault('nationality', invitation.candidate.nationality if invitation.candidate_id else '')
        return Response({'title': invitation.get_kind_display(), 'kind': invitation.kind, 'position': invitation.opening.title if invitation.opening_id else '',
            'expires_at': invitation.expires_at, 'location': invitation.opening.location if invitation.opening_id else '', 'public_application': bool(invitation.application_link_id),
            'department': invitation.opening.department if invitation.opening_id else '', 'schema': schema_for(invitation.kind), 'data': draft,
            'revision': invitation.draft_revision, 'state': invitation.state, 'correction_note': invitation.correction_note,
            'verified_email': invitation.verified_email, 'version': invitation.submissions.count(), 'attachments': [attachment_summary(item) for item in latest_attachments(invitation)]})

    @transaction.atomic
    def post(self, request):
        invitation = session_invitation(request)
        action = request.data.get('action')
        if action in ('send_code', 'verify_email'):
            if invitation.kind != 'pre_onboarding' or invitation.state in ('submitted', 'verified'):
                raise ValidationError({'detail': 'Email confirmation is available on an editable pre-onboarding form.'})
            if action == 'verify_email' and str(request.data.get('email', '')).strip().casefold() != invitation.recipient.casefold():
                raise ValidationError({'detail': 'Your email changed. Request a new verification code.'})
            result = pre_onboarding_verification(invitation, send_code=action == 'send_code', code=request.data.get('code', ''), email=request.data.get('email'))
            return result if result is not None else Response({'verified_email': invitation.verified_email, 'email_verified': True})
        if action not in ('save', 'submit'):
            raise ValidationError({'action': 'Choose save or submit.'})
        if invitation.state in ('submitted', 'verified'):
            return Response({'detail': 'This form has already been submitted.', 'state': invitation.state}, status=409)
        if request.data.get('revision') != invitation.draft_revision:
            return Response({'detail': 'This draft changed in another session. Reload before saving.'}, status=409)
        data = validate_form(invitation.kind, request.data.get('data'), final=action == 'submit')
        if action == 'submit' and invitation.kind == 'pre_onboarding':
            email = data['contact']['email']
            if not invitation.verified_email or email.casefold() != invitation.verified_email.casefold():
                raise ValidationError({'detail': 'Confirm your personal email address before submitting the form.', 'verification_required': True})
        from .recruitment_cv import latest_attachments
        attachments = latest_attachments(invitation)
        if action == 'submit' and invitation.kind == 'cv' and not any(item.kind == 'photo' for item in attachments):
            raise ValidationError({'detail': 'Upload your professional photo before submitting your CV.'})
        if action == 'submit' and invitation.application_link_id and not invitation.candidate_id:
            if invitation.kind == 'cv':
                identity = data['candidate']
                name, email, mobile, nationality = identity['name'], identity['email'], identity['phone'], identity.get('nationality', '')
            else:
                identity, contact = data['identity'], data['contact']
                name = f"{identity['first_name']} {identity['family_name']}".strip()
                email, mobile, nationality = contact['email'], contact['phone'], identity['nationality']
                if email.casefold() != invitation.recipient.casefold():
                    raise ValidationError({'detail': 'Use the email address verified for this application.'})
            if len(name) > 200 or len(email) > 254 or len(nationality) > 100:
                raise ValidationError({'detail': 'Candidate name, email or nationality exceeds the supported length.'})
            # A shared link never merges or exposes an existing person's record.
            from .recruitment_projection import cv_projection
            invitation.recipient = email
            invitation.candidate = RecruitmentCandidate.objects.create(name=name, email=email,
                mobile=mobile if len(mobile) <= 40 else '', nationality=nationality,
                position=(data['candidate'].get('position_applied') or (invitation.opening.title if invitation.opening_id else '')) if invitation.kind == 'cv' else invitation.opening.title if invitation.opening_id else '', department=invitation.opening.department if invitation.opening_id else '', status='Applied',
                **(cv_projection(data) if invitation.kind == 'cv' else {}))
        invitation.draft = data
        invitation.draft_revision += 1
        if action == 'submit':
            if invitation.kind == 'cv' and invitation.candidate_id:
                changed = []
                for source, target in [('position_applied', 'position'), ('nationality', 'nationality')]:
                    if data['candidate'].get(source):
                        setattr(invitation.candidate, target, data['candidate'][source])
                        changed.append(target)
                if changed:
                    invitation.candidate.save(update_fields=[*changed, 'updated_at'])
            declaration = DECLARATION.format(**data['beneficiary']) if invitation.kind == 'pre_onboarding' else ''
            submission = RecruitmentFormSubmission.objects.create(invitation=invitation, version=invitation.submissions.count() + 1, data=data, schema_version=VERSION, declaration_text=declaration)
            submission.attachments.set(attachments)
            invitation.state = 'submitted'
            event(request, invitation, 'Candidate CV submitted' if invitation.kind == 'cv' else 'Pre-onboarding form submitted')
            from apps.notifications.services import NotificationService
            transaction.on_commit(lambda: NotificationService.create_notification(
                recipient=invitation.owner, title='Candidate form submitted',
                message=f'{invitation.get_kind_display()} is ready for HR review (CAN-{invitation.candidate_id:04d}).',
                category='INFO', priority='NORMAL', force_in_app=True, send_email=False,
                action_url='/hr/talent-acquisition', action_label='Review candidate forms',
                metadata={'invitation_id': invitation.pk, 'candidate_id': invitation.candidate_id}))
        invitation.save(update_fields=['candidate', 'recipient', 'draft', 'draft_revision', 'state', 'updated_at'])
        return Response({'revision': invitation.draft_revision, 'state': invitation.state, 'version': invitation.submissions.count()})
