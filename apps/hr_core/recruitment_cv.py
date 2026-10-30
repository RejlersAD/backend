"""Bounded private CV uploads and candidate-reviewed extraction (HR-REC-04)."""
import base64
import hashlib
import io
import zipfile

from django.db import transaction
from django.http import HttpResponse
from django.utils.text import get_valid_filename
from rest_framework.exceptions import NotFound, PermissionDenied, ValidationError
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.response import Response

from .models import RecruitmentFormAttachment
from .recruitment_forms import HRFormsView, PublicFormView, FormThrottle, no_store, session_invitation


class UploadThrottle(FormThrottle):
    scope = 'candidate_uploads'
    rate = '12/min'


class ExtractionThrottle(FormThrottle):
    scope = 'candidate_cv_extraction'
    rate = '3/min'


def attachment_summary(attachment):
    return {'id': attachment.pk, 'kind': attachment.kind, 'name': attachment.name,
            'size': attachment.size, 'content_type': attachment.content_type,
            'created_at': attachment.created_at, 'extracted': bool(attachment.extraction)}


def latest_attachments(invitation):
    result = {}
    for attachment in invitation.attachments.defer('content', 'extraction').order_by('-created_at', '-pk'):
        result.setdefault(attachment.kind, attachment)
    return list(result.values())


def editable(invitation):
    if invitation.kind != 'cv':
        raise ValidationError({'detail': 'Uploads are available for the CV form only.'})
    if invitation.state in ('submitted', 'verified'):
        raise PermissionDenied('This form is submitted. Ask HR to request a correction before replacing files.')


def photo_bytes(content):
    from PIL import Image, ImageOps, UnidentifiedImageError
    try:
        with Image.open(io.BytesIO(content)) as image:
            if image.format not in ('JPEG', 'PNG', 'WEBP') or image.width * image.height > 20000000:
                raise ValueError()
            image.load()
            image = ImageOps.exif_transpose(image).convert('RGB')
            image.thumbnail((1200, 1200))
            output = io.BytesIO()
            image.save(output, format='JPEG', quality=88, optimize=True)
            return output.getvalue()
    except (UnidentifiedImageError, OSError, ValueError, Image.DecompressionBombError):
        raise ValidationError({'detail': 'Upload a valid JPG, PNG or WebP photo under 5 MB and 20 megapixels.'})


def cv_source(content, name, include_images=False):
    """No macros, remote fetches, scripts, embedded files or unlimited expansion."""
    if name.lower().endswith('.pdf') and content.startswith(b'%PDF-'):
        import pymupdf
        try:
            with pymupdf.open(stream=content, filetype='pdf') as document:
                if document.is_encrypted or not 1 <= document.page_count <= 20 or document.embfile_count():
                    raise ValueError()
                text, images, scanned_pages = [], [], 0
                for page in document:
                    page_text = page.get_text().strip()
                    text.append(page_text)
                    if len(page_text) < 30:
                        scanned_pages += 1
                        if scanned_pages > 10:
                            raise ValueError()
                    if include_images and len(page_text) < 30:
                        if len(images) >= 10:
                            raise ValueError()
                        scale = min(2.5, 2000 / max(page.rect.width, page.rect.height, 1))
                        pixmap = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
                        if pixmap.width * pixmap.height > 5000000:
                            raise ValueError()
                        images.append(base64.b64encode(pixmap.tobytes('jpeg')).decode())
                combined = '\n'.join(text)
                if len(combined) > 60000:
                    raise ValueError()
                return combined, images, 'application/pdf'
        except Exception:
            raise ValidationError({'detail': 'Use an unlocked PDF of at most 20 pages (scanned CVs: at most 10 pages), with no embedded files.'})
    if name.lower().endswith('.docx') and content.startswith(b'PK'):
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 1500 or sum(item.file_size for item in entries) > 25000000:
                    raise ValueError()
                if any('vbaproject' in item.filename.lower() or '/embeddings/' in item.filename.lower() for item in entries):
                    raise ValueError()
            from docx import Document
            document = Document(io.BytesIO(content))
            text = '\n'.join([p.text for p in document.paragraphs] + [
                ' | '.join(dict.fromkeys(cell.text for cell in row.cells)) for table in document.tables for row in table.rows])
            if not text.strip() or len(text) > 60000:
                raise ValueError()
            return text, [], 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        except Exception:
            raise ValidationError({'detail': 'Use a readable DOCX CV without macros or embedded files (maximum 60,000 characters).'})
    raise ValidationError({'detail': 'Upload a PDF or DOCX CV. Other file types are not supported.'})


class CandidateUploadView(PublicFormView):
    parser_classes = [MultiPartParser, FormParser]
    throttle_classes = [UploadThrottle]

    def post(self, request):
        upload, kind = request.FILES.get('file'), request.data.get('kind')
        if not upload or kind not in ('photo', 'cv'):
            raise ValidationError({'detail': 'Choose a professional photo or CV file.'})
        limit = (5 if kind == 'photo' else 8) * 1024 * 1024
        if not 0 < upload.size <= limit:
            raise ValidationError({'detail': f'The file must be under {5 if kind == "photo" else 8} MB.'})
        # Validate session before doing parsing work; recheck after validation.
        with transaction.atomic():
            editable(session_invitation(request))
        content = upload.read(limit + 1)
        name = get_valid_filename(upload.name.replace('\\', '/').split('/')[-1])[:190]
        if kind == 'photo':
            content, name, content_type = photo_bytes(content), 'professional-photo.jpg', 'image/jpeg'
        else:
            _, _, content_type = cv_source(content, name)
        with transaction.atomic():
            invitation = session_invitation(request)
            editable(invitation)
            # Bound retained revisions; submitted snapshots are never overwritten.
            sizes = list(invitation.attachments.values_list('size', flat=True))
            if len(sizes) >= 12 or sum(sizes) + len(content) > 24 * 1024 * 1024:
                raise ValidationError({'detail': 'This form has reached its upload limit. Contact HR for a new invitation.'})
            attachment = RecruitmentFormAttachment.objects.create(invitation=invitation, kind=kind,
                name=name, content_type=content_type, content=content, size=len(content), sha256=hashlib.sha256(content).hexdigest())
        return Response(attachment_summary(attachment), status=201)


def file_response(attachment):
    response = HttpResponse(bytes(attachment.content), content_type=attachment.content_type)
    response['Content-Disposition'] = f'attachment; filename="{attachment.name}"'
    response['X-Content-Type-Options'] = 'nosniff'
    response['Content-Security-Policy'] = "default-src 'none'; sandbox"
    return no_store(response)


class CandidateAttachmentView(PublicFormView):
    @transaction.atomic
    def get(self, request, pk):
        invitation = session_invitation(request)
        try:
            attachment = invitation.attachments.get(pk=pk)
        except RecruitmentFormAttachment.DoesNotExist:
            raise NotFound()
        return file_response(attachment)


class RecruitmentAttachmentView(HRFormsView):
    permission_action = 'export'

    def get(self, request, pk, attachment):
        try:
            document = RecruitmentFormAttachment.objects.filter(pk=attachment, invitation_id=pk,
                recruitmentformsubmission__isnull=False).distinct().get()
        except RecruitmentFormAttachment.DoesNotExist:
            raise NotFound()
        return file_response(document)


class RecruitmentCandidatePhotoView(HRFormsView):
    permission_action = 'read'

    def get(self, request, pk):
        photo = RecruitmentFormAttachment.objects.filter(invitation__candidate_id=pk, kind='photo',
            recruitmentformsubmission__isnull=False).order_by('-recruitmentformsubmission__created_at', '-pk').first()
        if photo is None:
            raise NotFound()
        return file_response(photo)


def extraction_failure(error):
    from .recruitment_cv_local import CVTemplateError
    if isinstance(error, CVTemplateError):
        return 'cv_template_only', 'Your file is saved, but it contains an unfilled RAD CV template. Upload your completed CV or enter your details below. Your existing entries are kept.'
    if isinstance(error, (TimeoutError, RuntimeError)):
        return 'ocr_unavailable', 'Your CV is saved. Local OCR could not finish. Retry or upload a readable Word or text PDF CV. Your existing entries are kept.'
    if isinstance(error, (ValidationError, ValueError, KeyError, IndexError)):
        return 'cv_mapping_failed', 'Your CV is saved. Auto-fill could not identify the CV sections. Upload a readable CV or enter your details below. Your existing entries are kept.'
    return 'extraction_unavailable', 'Your CV is saved. Auto-fill could not complete. You can retry or complete the form manually.'


def extract_cv(attachment):
    from .recruitment_cv_local import extract_local
    return extract_local(attachment)


class CandidateCVExtractView(PublicFormView):
    throttle_classes = [ExtractionThrottle]

    def post(self, request):
        with transaction.atomic():
            invitation = session_invitation(request)
            editable(invitation)
            try:
                attachment = invitation.attachments.get(pk=request.data.get('attachment'), kind='cv')
            except (RecruitmentFormAttachment.DoesNotExist, ValueError, TypeError):
                raise NotFound()
        try:
            proposed = extract_cv(attachment)
        except Exception as error:
            code, detail = extraction_failure(error)
            return Response({'detail': detail, 'error_code': code, 'file_saved': True, 'manual_entry_available': True}, status=503)
        with transaction.atomic():
            invitation = session_invitation(request)
            editable(invitation)
            if invitation.attachments.filter(kind='cv').order_by('-created_at', '-pk').values_list('pk', flat=True).first() != attachment.pk:
                return Response({'detail': 'A newer CV was uploaded. Retry auto-fill for the latest CV.'}, status=409)
            attachment.extraction = proposed
            attachment.save(update_fields=['extraction', 'updated_at'])
        return Response({'data': proposed, 'review_required': True,
            'detail': 'Review and correct the extracted information before submitting. Missing details are left blank.'})
