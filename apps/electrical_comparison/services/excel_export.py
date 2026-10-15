"""
Electrical Comparison — Excel exporter.

When `source` is given (the frontend's per-tab Export Excel button),
one sheet with just that comparison's rows — unchanged from before.
When `source` is omitted (the history list's "Download" button, which
exports everything), one sheet PER comparison tab that actually has
data, instead of dumping every source into a single sheet where rows
from different comparisons (different status vocabularies, e.g. a
'matched' equipment_list row next to a 'fully_matched' combined row)
sat side by side.
openpyxl only, no third-party templates — same convention as
apps.valve_mto.excel_export / apps.instrument_io_workflow.excel_export.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

_HEADER_FILL = PatternFill(start_color='003366', end_color='003366', fill_type='solid')
_HEADER_FONT = Font(color='FFFFFF', bold=True)
_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)

# Single-sheet, one-source export (source= given) — keeps the Source
# column since, historically, this could also be called with no source
# (now split out below), where Source was the only thing telling rows
# from different comparisons apart.
_COLUMNS = [
    ('tag_number', 'Tag Number'),
    ('equipment_type', 'Equipment Type'),
    ('description', 'Description'),
    ('status', 'Status'),
    ('source', 'Source'),
    ('remarks', 'Remarks'),
]

# Per-tab sheet columns (no Source column — the sheet itself, and its
# name, already say which comparison these rows belong to).
_TAB_COLUMNS = [
    ('tag_number', 'Tag'),
    ('equipment_type', 'Equipment Type'),
    ('description', 'Description'),
    ('status', 'Status'),
    ('remarks', 'Remarks'),
]

# Sheet name + row order for the full (no source filter) export — same
# comparison-tab order the frontend displays (TAB_ORDER in
# SingleLineDiagram.jsx). Only a source with at least one saved row gets
# its own sheet.
_TAB_SHEET_NAMES = [
    ('equipment_list', 'P&ID vs Equipment List'),
    ('load_list', 'P&ID vs Load List'),
    ('equipment_vs_loadlist', 'Equipment vs Load List'),
    ('combined', 'Full Comparison'),
]


def _write_sheet(ws, rows, columns):
    for idx, (_, label) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=idx, value=label)
        cell.fill, cell.font, cell.alignment = _HEADER_FILL, _HEADER_FONT, _CENTER

    for r_idx, row in enumerate(rows, start=2):
        for c_idx, (field, _) in enumerate(columns, start=1):
            ws.cell(row=r_idx, column=c_idx, value=getattr(row, field))

    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(length + 2, 12), 60)


def export_job_to_xlsx(job, source: str | None = None) -> bytes:
    wb = Workbook()

    if source:
        ws = wb.active
        ws.title = 'Electrical Comparison'
        _write_sheet(ws, job.results.filter(source=source), _COLUMNS)
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()

    all_rows = list(job.results.all())
    wb.remove(wb.active)
    any_sheet = False
    for src, sheet_name in _TAB_SHEET_NAMES:
        rows = [r for r in all_rows if r.source == src]
        if not rows:
            continue
        _write_sheet(wb.create_sheet(title=sheet_name), rows, _TAB_COLUMNS)
        any_sheet = True
    if not any_sheet:
        # No saved rows at all (e.g. a P&ID-only job with nothing to
        # compare) — a workbook needs at least one sheet to open cleanly.
        _write_sheet(wb.create_sheet(title='Electrical Comparison'), [], _TAB_COLUMNS)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
