"""SLD analysis exports (spec §6.5). Headers say SLD, never P&ID.

- format=workbook: Summary sheet + one sheet per result tab.
- format=drafting: sheet · location · issue · severity only.
"""
from __future__ import annotations

import io

import openpyxl
from openpyxl.styles import Font

BOLD = Font(bold=True)

TAB_COLUMNS = {
    'equipment': ['Tag', 'Description', 'Category', 'List sheet', 'On SLD', 'Status', 'Reason'],
    'load': ['Tag', 'Description (list)', 'Board (list)', 'Board (SLD)', 'kW (list)', 'kW (SLD)', 'Status', 'Issue'],
    'eqload': ['Tag', 'Description (Load List)', 'Board', 'Rated kW', 'Status', 'Equipment List description', 'List sheet'],
    'drafting': ['Check', 'Sheet', 'Location', 'Result', 'Severity'],
    'interfaces': ['Tag', 'Where on SLD', 'Meaning'],
    'cables': ['Cable tag', 'Size', 'Sheet'],
}
TAB_TITLES = {
    'equipment': 'SLD vs Equipment List',
    'load': 'SLD vs Load List',
    'eqload': 'Equipment vs Load List',
    'drafting': 'Drafting check',
    'interfaces': 'Interfaces',
    'cables': 'Cables',
}


def _row(tab, item):
    if tab == 'equipment':
        return [item.get('tag'), item.get('description'), item.get('category'),
                item.get('listSheet'), item.get('onSld'), item.get('status'), item.get('reason')]
    if tab == 'load':
        return [item.get('tag'), item.get('description'), item.get('boardList'), item.get('boardSld'),
                item.get('kwList'), item.get('kwSld'), item.get('status'), item.get('issue')]
    if tab == 'eqload':
        return [item.get('tag'), item.get('description'), item.get('board'), item.get('ratedKw'),
                item.get('status'), item.get('equipmentDescription'), item.get('equipmentSheet')]
    if tab == 'drafting':
        return [item.get('check'), item.get('sheet'), item.get('location'),
                item.get('result'), item.get('severity')]
    if tab == 'interfaces':
        return [item.get('tag'), item.get('where'), item.get('meaning')]
    if tab == 'cables':
        return [item.get('tag'), item.get('size'), item.get('sheet')]
    return []


def build_export(analysis, fmt='workbook'):
    result = analysis.result or {}
    wb = openpyxl.Workbook()

    if fmt == 'drafting':
        ws = wb.active
        ws.title = 'Drafting comments'
        ws.append(['Sheet', 'Location', 'Issue', 'Severity'])
        for c in ws[1]:
            c.font = BOLD
        for f in result.get('drafting', []):
            ws.append([f.get('sheet'), f.get('location'),
                       f"{f.get('check')} — {f.get('result')}", f.get('severity')])
    else:
        ws = wb.active
        ws.title = 'Summary'
        h = result.get('headline', {})
        drawing = result.get('drawing', {})
        ws.append(['RADAI SLD Comparison — Summary'])
        ws['A1'].font = BOLD
        ws.append(['SLD drawing', f"{drawing.get('drawing_no', '')} Rev {drawing.get('rev', '')}"])
        ws.append(['Analysed', str(analysis.created_at)])
        for label, key in [
            ('Boards on SLD', 'boards'), ('Transformers', 'transformers'), ('Generators', 'generators'),
            ('Matched with Equipment List', None), ('Load List conflicts', None),
            ('Drafting issues', None), ('Interfaces', 'interfaces'),
        ]:
            if key:
                ws.append([label, h.get(key, 0)])
        ws.append(['Matched with Equipment List', f"{h.get('matched', 0)} of {h.get('matchedTotal', 0)}"])
        ws.append(['Load List conflicts', f"{h.get('loadConflicts', 0)} of {h.get('loadTotal', 0)}"])
        ws.append(['Drafting issues', f"{h.get('draftingTotal', 0)} ({h.get('draftingHigh', 0)} high · {h.get('draftingMedium', 0)} medium)"])

        for tab, title in TAB_TITLES.items():
            ws = wb.create_sheet(title[:31])
            ws.append(TAB_COLUMNS[tab])
            for c in ws[1]:
                c.font = BOLD
            for item in result.get(tab, []):
                ws.append(_row(tab, item))

    buf = io.BytesIO()
    wb.save(buf)
    name = f"SLD_drafting_comments_{analysis.id}.xlsx" if fmt == 'drafting' else f"SLD_comparison_{analysis.id}.xlsx"
    return buf.getvalue(), name
