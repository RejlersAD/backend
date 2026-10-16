"""Legend Workbook Importer — one multi-sheet .xlsx → per-section legends.

Companion to legend_pack_splitter.py (which handles PDF/image packs via AI).
This module handles structured Excel workbooks where each SHEET is a legend
artefact: code tables (Fluid Codes, Insulation Codes, …), numbering-system
sheets (Line Designation, Valve Numbering, …) and SYMBOL sheets with images
embedded in the SYMBOL column (openpyxl ``ws._images``, anchored per row).

Pipeline:
  1. ``preview_workbook`` — per-sheet kind + predicted section (soft-coded
     SHEET_RULES exact map, then KEYWORD_RULES fuzzy fallback) so the user
     can confirm/override before anything is written.
  2. ``import_workbook`` — applies a user-confirmed {sheet: section|null}
     map: code tables become lookup dicts overlaid onto the section's default
     template (same overlay philosophy as the pack splitter), numbering and
     notes sheets become notes-style definitions, and pictures from EVERY
     mapped sheet create LegendSymbolImage rows. Original cell positions,
     merged ranges and picture associations live in definition.imported_sheets
     for aligned previews and extraction reference context. One active
     PidCheckerV2LegendSheet per (project, section) — existing project
     legend updated in place when present, siblings deactivated.

Everything is soft-coded at module level; a failing sheet never aborts the
import — it lands in ``warnings`` instead.
"""
from __future__ import annotations

import copy
import io
import logging
import re
from typing import Any, Dict, List, Optional

from django.db import transaction

from ..legend_defaults import (
    SECTIONS, SECTION_LABELS, DEFAULT_TEMPLATES,
    SECTION_LINE_LIST, SECTION_EQUIPMENT_LIST, SECTION_INSTRUMENT_INDEX,
    SECTION_VALVE, SECTION_PIPING, SECTION_CONTROL_VALVE_REGULATOR,
    SECTION_INSTRUMENT_SIGNAL, SECTION_EQUIPMENT_SYMBOLS,
    SECTION_INSTRUMENT_FUNCTION, SECTION_INSTRUMENT_TYPICAL_LETTER,
    SECTION_SCOPE_SYMBOLS, SECTION_MISCELLANEOUS, SECTION_LIMIT_LINE,
    SECTION_OTHER_SPECIALTIES, SECTION_SPECIAL_PIPING, SECTION_DRAWING_CONT,
    SECTION_GENERAL_INSTRUMENT,
)
from ..models import PidCheckerV2LegendSheet, LegendSymbolImage

logger = logging.getLogger(__name__)

# Soft-coded: S3/disk prefix for the structured import-record archive.
WORKBOOK_ARCHIVE_PREFIX = 'legend_workbook_imports'

ACCEPTED_EXTENSIONS = ('.xlsx', '.xlsm')
HEADER_SCAN_DEPTH   = 10      # rows scanned at the top of each sheet
MIN_HEADER_HITS     = 2       # alias hits needed to accept a header row
MAX_SAMPLE_ROWS     = 3
MAX_NOTES_LINES     = 400     # cap per notes-style sheet

# Header aliases (lower-cased) → logical column role.
HEADER_ALIASES = {
    'code':  {'code', 'codes', 'symbol', 'segment', 'abbreviation', 'abbr', 'letter', 'letters', 'type',
              'piping class', 'area code number'},
    'desc':  {'description', 'desc', 'name', 'meaning', 'service', 'remarks', 'remark', 'piping material'},
    'skip':  {'s.no.', 's.no', 'sl.no', 'sl no', 'no', 'no.', 'sno', 'sr no', 'item'},
}

# ─── Soft-coded sheet → section rules ────────────────────────────────────
# Exact sheet-name map (lower-cased). Value: dict(section, kind,
# overlay=field_key|None). `overlay` = the field of the section's DEFAULT
# template whose `lookup` gets replaced by this sheet's code table (same
# overlay approach as legend_pack_splitter.SECTION_LOOKUP_OVERLAYS). When
# None, a standalone notes/lookup definition is built instead.
SHEET_RULES: Dict[str, Dict[str, Any]] = {
    # ── Consolidated legend workbooks (e.g. Legends_Full.xlsx) ──────────────
    'line numbering':               {'section': SECTION_LINE_LIST,       'kind': 'numbering'},
    'equipment numbering':          {'section': SECTION_EQUIPMENT_LIST,  'kind': 'numbering'},
    'p&id numbering':               {'section': SECTION_DRAWING_CONT,    'kind': 'numbering'},
    'notes & drawing details':      {'section': SECTION_MISCELLANEOUS,   'kind': 'notes'},
    'symbols':                      {'section': SECTION_MISCELLANEOUS,   'kind': 'symbols'},
    'all codes':                    {'section': SECTION_LINE_LIST,       'kind': 'numbering'},
    'fluid codes':                  {'section': SECTION_LINE_LIST,       'kind': 'codes', 'overlay': 'service'},
    'line designation':             {'section': SECTION_LINE_LIST,       'kind': 'numbering'},
    'piping material codes':        {'section': SECTION_LINE_LIST,       'kind': 'codes', 'overlay': 'spec'},
    'insulation codes':             {'section': SECTION_LINE_LIST,       'kind': 'codes', 'overlay': 'insulation'},
    'existing area codes':          {'section': SECTION_SCOPE_SYMBOLS,   'kind': 'codes'},
    'new area codes':               {'section': SECTION_SCOPE_SYMBOLS,   'kind': 'codes'},
    'drawing number designation':   {'section': SECTION_DRAWING_CONT,    'kind': 'numbering'},
    'valve numbering':              {'section': SECTION_VALVE,           'kind': 'numbering'},
    'valve symbols':                {'section': SECTION_VALVE,           'kind': 'symbols'},
    'esd valve types':              {'section': SECTION_VALVE,           'kind': 'symbols'},
    'line symbols':                 {'section': SECTION_LIMIT_LINE,      'kind': 'symbols'},
    'control valve symbols':        {'section': SECTION_CONTROL_VALVE_REGULATOR, 'kind': 'symbols'},
    'motor & valve control types':  {'section': SECTION_CONTROL_VALVE_REGULATOR, 'kind': 'symbols'},
    'pressure-safety symbols':      {'section': SECTION_OTHER_SPECIALTIES,       'kind': 'symbols'},
    'general symbols':              {'section': SECTION_MISCELLANEOUS,   'kind': 'symbols'},
    'safety & fire water':          {'section': SECTION_OTHER_SPECIALTIES,       'kind': 'symbols'},
    'typical assemblies':           {'section': SECTION_SPECIAL_PIPING,  'kind': 'symbols'},
    'analyser sample connection':   {'section': SECTION_INSTRUMENT_INDEX,        'kind': 'symbols'},
    'p&id notes & legend':          {'section': SECTION_MISCELLANEOUS,   'kind': 'notes'},
    'abbreviations':                {'section': SECTION_MISCELLANEOUS,   'kind': 'codes'},
    'equipment codes':              {'section': SECTION_EQUIPMENT_LIST,  'kind': 'codes'},
    'equipment symbols':            {'section': SECTION_EQUIPMENT_SYMBOLS,       'kind': 'symbols'},
    'instrument id table':          {'section': SECTION_INSTRUMENT_INDEX,        'kind': 'codes'},
    'instrument letter positions':  {'section': SECTION_INSTRUMENT_TYPICAL_LETTER, 'kind': 'numbering'},
    'instrument tagging & esd':     {'section': SECTION_INSTRUMENT_INDEX,        'kind': 'numbering'},
    'instrument symbols':           {'section': SECTION_GENERAL_INSTRUMENT,      'kind': 'symbols'},
    'instrument connections':       {'section': SECTION_INSTRUMENT_SIGNAL,       'kind': 'symbols'},
    'level instrument connections': {'section': SECTION_INSTRUMENT_SIGNAL,       'kind': 'symbols'},
    'level instrument types':       {'section': SECTION_GENERAL_INSTRUMENT,      'kind': 'symbols'},
}

# Fuzzy fallback — first matching keyword wins (order matters).
# (keyword, section, kind, overlay_field|None)
KEYWORD_RULES = [
    ('fluid',           SECTION_LINE_LIST,       'codes',     'service'),
    ('insulation',      SECTION_LINE_LIST,       'codes',     'insulation'),
    ('piping material', SECTION_LINE_LIST,       'codes',     'spec'),
    ('line designation', SECTION_LINE_LIST,      'numbering', None),
    ('line symbol',     SECTION_LIMIT_LINE,      'symbols',   None),
    ('valve numbering', SECTION_VALVE,           'numbering', None),
    ('valve symbol',    SECTION_VALVE,           'symbols',   None),
    ('valve',           SECTION_VALVE,           'codes',     None),
    ('control valve',   SECTION_CONTROL_VALVE_REGULATOR, 'symbols', None),
    ('equipment symbol', SECTION_EQUIPMENT_SYMBOLS, 'symbols', None),
    ('equipment',       SECTION_EQUIPMENT_LIST,  'codes',     None),
    ('instrument letter', SECTION_INSTRUMENT_TYPICAL_LETTER, 'numbering', None),
    ('instrument',      SECTION_INSTRUMENT_INDEX, 'codes',    None),
    ('drawing number',  SECTION_DRAWING_CONT,    'numbering', None),
    ('area code',       SECTION_SCOPE_SYMBOLS,   'codes',     None),
    ('piping',          SECTION_PIPING,          'codes',     None),
    ('abbreviation',    SECTION_MISCELLANEOUS,   'codes',     None),
    ('notes',           SECTION_MISCELLANEOUS,   'notes',     None),
    ('legend',          SECTION_MISCELLANEOUS,   'notes',     None),
    ('symbol',          SECTION_MISCELLANEOUS,   'symbols',   None),
    ('code',            SECTION_MISCELLANEOUS,   'codes',     None),
]

VALID_KINDS = ('codes', 'numbering', 'symbols', 'notes', 'fields', 'unknown')


# ─── Small helpers ────────────────────────────────────────────────────────
def _norm(v) -> str:
    return re.sub(r'\s+', ' ', str(v or '')).strip()


def _predict_sheet(sheet_name: str, image_count: int) -> Dict[str, Any]:
    """Soft-coded prediction: exact map first, then keyword scoring."""
    key = sheet_name.strip().lower()
    rule = SHEET_RULES.get(key)
    if rule:
        return {'section': rule['section'], 'kind': rule['kind'],
                'overlay': rule.get('overlay'), 'confidence': 'high'}
    for kw, section, kind, overlay in KEYWORD_RULES:
        if kw in key:
            return {'section': section, 'kind': kind, 'overlay': overlay,
                    'confidence': 'medium'}
    # Heuristic fallback: embedded pictures ⇒ symbol sheet.
    if image_count:
        return {'section': SECTION_MISCELLANEOUS, 'kind': 'symbols',
                'overlay': None, 'confidence': 'low'}
    return {'section': None, 'kind': 'unknown', 'overlay': None, 'confidence': 'low'}


def _detect_table(ws) -> Dict[str, Any]:
    """Find the header row (first HEADER_SCAN_DEPTH rows) and map columns.
    Returns {header_row (0-based|None), code_col, desc_col, headers}."""
    headers: List[str] = []
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=HEADER_SCAN_DEPTH, values_only=True)):
        cells = [_norm(c).lower() for c in row]
        code_col = desc_col = None
        hits = 0
        for c_idx, cell in enumerate(cells):
            if not cell:
                continue
            if cell in HEADER_ALIASES['skip']:
                hits += 1
            elif cell in HEADER_ALIASES['code'] and code_col is None:
                code_col = c_idx
                hits += 1
            elif cell in HEADER_ALIASES['desc'] and desc_col is None:
                desc_col = c_idx
                hits += 1
        if hits >= MIN_HEADER_HITS:
            headers = [_norm(c) for c in row if _norm(c)]
            return {'header_row': r_idx, 'code_col': code_col, 'desc_col': desc_col, 'headers': headers}
    return {'header_row': None, 'code_col': None, 'desc_col': None, 'headers': []}


def _iter_data_rows(ws, header_row: int):
    start = (header_row if header_row is not None else 0) + 2  # 1-based min_row
    for row in ws.iter_rows(min_row=start, values_only=True):
        yield row


def _extract_codes(ws, table) -> Dict[str, str]:
    lookup: Dict[str, str] = {}
    if table['code_col'] is None or table['desc_col'] is None:
        return lookup
    for row in _iter_data_rows(ws, table['header_row']):
        code = _norm(row[table['code_col']] if table['code_col'] < len(row) else '')
        desc = _norm(row[table['desc_col']] if table['desc_col'] < len(row) else '')
        if not code or not desc or code in ('—', '-'):
            continue
        lookup[code.upper()] = desc
    return lookup


def _extract_notes(ws, table) -> List[str]:
    """Numbering/notes sheets → readable text lines (format line, segment
    explanations, free-text notes)."""
    lines: List[str] = []
    # Title / format lines above the header
    for row in ws.iter_rows(min_row=1, max_row=(table['header_row'] or HEADER_SCAN_DEPTH), values_only=True):
        txt = ' '.join(_norm(c) for c in row if _norm(c))
        if txt:
            lines.append(txt)
    # Data rows
    for row in _iter_data_rows(ws, table['header_row']):
        cells = [_norm(c) for c in row if _norm(c)]
        if not cells:
            continue
        if len(cells) >= 2:
            lines.append(f'{cells[0]} = {cells[1]}' + (f' ({cells[2]})' if len(cells) > 2 else ''))
        else:
            lines.append(cells[0])
        if len(lines) >= MAX_NOTES_LINES:
            break
    return lines[:MAX_NOTES_LINES]


def _sample_rows(ws, table) -> List[List[str]]:
    out: List[List[str]] = []
    for row in _iter_data_rows(ws, table['header_row']):
        cells = [_norm(c) for c in row]
        if any(cells):
            out.append([c for c in cells if c][:6])
        if len(out) >= MAX_SAMPLE_ROWS:
            break
    return out


def _build_codes_definition(section: str, title: str, lookup: Dict[str, str],
                            overlay_field: Optional[str], notes: Optional[List[str]] = None) -> dict:
    """Compile-safe definition: deep-copy the section default template and
    overlay the code table onto the target field's lookup when configured;
    otherwise a standalone lookup/notes definition."""
    tpl = DEFAULT_TEMPLATES.get(section)
    if tpl and overlay_field:
        definition = copy.deepcopy(tpl['definition'])
        for field in definition.get('fields') or []:
            if field.get('key') == overlay_field:
                field['lookup'] = lookup
                break
        else:
            overlay_field = None
    if not tpl or not overlay_field:
        definition = {
            'separator': '-',
            'fields': [{
                'key': 'code',
                'label': title,
                'regex': r'[A-Z0-9]{1,8}',
                'optional': True,
                'lookup': lookup,
                'notes': '\n'.join(notes or []),
            }],
        }
    if notes:
        definition.setdefault('notes', '\n'.join(notes))
    return definition


def _build_notes_definition(title: str, lines: List[str]) -> dict:
    return {
        'separator': '-',
        'fields': [{
            'key': 'note',
            'label': title,
            'regex': r'.*',
            'optional': True,
            'notes': '\n'.join(lines),
        }],
    }


def _extract_symbol_images(ws, table) -> List[Dict[str, Any]]:
    """Embedded pictures anchored in the SYMBOL column → (symbol_name, bytes).

    openpyxl anchors are 0-based; the anchor row matches the data row whose
    DESCRIPTION cell names the symbol.
    """
    out: List[Dict[str, Any]] = []
    desc_col = table['desc_col'] if table['desc_col'] is not None else 2
    for img in getattr(ws, '_images', []):
        try:
            anchor_row = img.anchor._from.row          # 0-based
            anchor_col = img.anchor._from.col
            values = next(ws.iter_rows(min_row=anchor_row + 1, max_row=anchor_row + 1, values_only=True))
            name = _norm(values[desc_col]) if desc_col < len(values) else ''
            if not name:
                name = ' | '.join(_norm(c) for c in values if _norm(c))
            # Multi-column sheets contain several different pictures per row.
            # Keep their column context, not just the shared row description.
            row_images = [other for other in ws._images if other.anchor._from.row == anchor_row]
            if len(row_images) > 1:
                header_row = (table['header_row'] + 1) if table['header_row'] is not None else 4
                column_label = _norm(ws.cell(header_row, anchor_col + 1).value)
                name = f'{name or ws.title} | {column_label or f"Column {anchor_col + 1}"}'
            data = img._data()
            if data:
                out.append({
                    'name': name or f'{ws.title} row {anchor_row + 1}',
                    'data': data, 'row': anchor_row + 1, 'column': anchor_col + 1,
                })
        except Exception as exc:
            logger.warning('[LegendWorkbook] image extraction failed on %s: %s', ws.title, exc)
    return out


def _sheet_content(ws, kind: str) -> Dict[str, Any]:
    """Keep column positions and line breaks, including non-tabular sheets."""
    return {
        'name': ws.title, 'kind': kind,
        'rows': [
            [str(cell) if cell is not None else '' for cell in row]
            for row in ws.iter_rows(
                max_row=max([ws.max_row] + [image.anchor._from.row + 1 for image in ws._images]),
                max_col=max([ws.max_column] + [image.anchor._from.col + 1 for image in ws._images]),
                values_only=True,
            )
        ],
        'merged_ranges': [str(item) for item in ws.merged_cells.ranges],
        'images': [],
    }


def _extract_fields_definition(ws) -> Optional[dict]:
    """Keep the existing fields-only Excel import format supported."""
    for header_row, row in enumerate(ws.iter_rows(max_row=HEADER_SCAN_DEPTH, values_only=True), start=1):
        columns = {_norm(value).lower(): index for index, value in enumerate(row) if _norm(value)}
        if 'key' not in columns or 'regex' not in columns:
            continue
        definition = {'separator': '-', 'fields': []}
        for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
            field = {key: values[index] for key, index in columns.items() if index < len(values)}
            if not field.get('key'):
                continue
            item = {
                'key': str(field['key']), 'label': str(field.get('label') or field['key']),
                'regex': str(field.get('regex') or '[A-Z0-9]+'),
            }
            for key in ('suffix', 'notes'):
                if field.get(key):
                    item[key] = str(field[key])
            if _norm(field.get('optional')).lower() in ('yes', 'true', '1', 'y'):
                item['optional'] = True
            if field.get('separator'):
                definition['separator'] = str(field['separator'])
            if field.get('lookup'):
                pairs = [pair.partition('=') for pair in re.split(r'\||\r?\n|;', str(field['lookup']))]
                item['lookup'] = {key.strip(): value.strip() for key, separator, value in pairs if separator and key.strip()}
            definition['fields'].append(item)
        return definition if definition['fields'] else None
    return None


def _dedupe_name(base: str, used: set) -> str:
    max_length = LegendSymbolImage._meta.get_field('symbol_name').max_length
    base = base[:max_length]
    name = base
    i = 2
    while name in used:
        suffix = f' ({i})'
        name = f'{base[:max_length - len(suffix)]}{suffix}'
        i += 1
    used.add(name)
    return name


# ─── Public API ───────────────────────────────────────────────────────────
def preview_workbook(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """Per-sheet prediction table for the confirm UI. Never raises on a
    single bad sheet — it lands in the sheet's own warning."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheets: List[Dict[str, Any]] = []
    for ws in wb.worksheets:
        try:
            image_count = len(getattr(ws, '_images', []))
            table = _detect_table(ws)
            pred = _predict_sheet(ws.title, image_count)
            fields_definition = _extract_fields_definition(ws)
            data_rows = max(0, (ws.max_row or 1) - ((table['header_row'] or 0) + 1))
            sheets.append({
                'name': ws.title,
                'predicted_section': pred['section'],
                'confidence': pred['confidence'],
                'kind': 'fields' if fields_definition else pred['kind'],
                'headers': table['headers'],
                'row_count': data_rows,
                'image_count': image_count,
                'sample_rows': _sample_rows(ws, table),
            })
        except Exception as exc:                         # pragma: no cover
            sheets.append({'name': ws.title, 'predicted_section': None, 'confidence': 'low',
                           'kind': 'unknown', 'headers': [], 'row_count': 0,
                           'image_count': 0, 'sample_rows': [], 'warning': str(exc)})
    return {'filename': filename, 'sheet_count': len(sheets), 'sheets': sheets}


def import_workbook(file_bytes: bytes, filename: str, *, user, project, sheet_map: Dict[str, Optional[str]],
                    valve_mto: bool = False) -> Dict[str, Any]:
    """Apply the confirmed sheet→section map. ``sheet_map`` value None/'' =
    skip that sheet. Returns a summary dict; per-sheet failures are collected
    in ``warnings`` and never abort the import."""
    import openpyxl
    from django.core.files.base import ContentFile
    if valve_mto:
        from apps.valve_mto.models import ValveMTOLegend, VALVE_MTO_LEGEND_SECTIONS
        if any(section and section not in VALVE_MTO_LEGEND_SECTIONS for section in sheet_map.values()):
            raise ValueError('Map worksheets only to Valve MTO legend sections.')

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    pack_name = (filename or 'Legend workbook').rsplit('.', 1)[0]
    if not any(sheet_map.get(ws.title) in SECTIONS for ws in wb.worksheets):
        raise ValueError('Map at least one worksheet to a valid legend section.')

    warnings: List[str] = []
    skipped: List[str] = []
    # section → aggregated content
    section_lookups: Dict[str, Dict[str, Dict[str, str]]] = {}   # section → overlay_field|'_standalone' → lookup
    section_notes: Dict[str, List[str]] = {}
    images_per_section: Dict[str, int] = {}
    # Per-section detail records for the archive (legend refs added after persist).
    section_details: Dict[str, Dict[str, Any]] = {}
    section_sheets: Dict[str, List[Dict[str, Any]]] = {}
    section_images: Dict[str, List[Dict[str, Any]]] = {}
    section_definitions: Dict[str, dict] = {}

    for ws in wb.worksheets:
        section = sheet_map.get(ws.title)
        if not section:
            skipped.append(ws.title)
            continue
        if section not in SECTIONS:
            warnings.append(f'{ws.title}: unknown section {section!r} — skipped')
            continue
        try:
            rule = SHEET_RULES.get(ws.title.strip().lower(), {})
            kind = rule.get('kind') or _predict_sheet(ws.title, len(getattr(ws, '_images', [])))['kind']
            table = _detect_table(ws)
            fields_definition = _extract_fields_definition(ws)
            if fields_definition:
                kind = 'fields'
            content = _sheet_content(ws, kind)
            section_sheets.setdefault(section, []).append(content)
            # Pictures occur on notes/numbering sheets too, not only sheets
            # predicted as "symbols" (e.g. Instrument Tagging & ESD).
            imgs = _extract_symbol_images(ws, table)
            for image in imgs:
                image['sheet_content'] = content
            section_images.setdefault(section, []).extend(imgs)
            if len(imgs) != len(ws._images):
                warnings.append(f'{ws.title}: extracted {len(imgs)} of {len(ws._images)} embedded pictures')

            if kind == 'fields':
                if section in section_definitions:
                    warnings.append(f'{ws.title}: multiple field definitions mapped to {section}; first rule set retained, all source cells preserved')
                else:
                    section_definitions[section] = fields_definition
            elif kind == 'symbols':
                if not imgs:
                    warnings.append(f'{ws.title}: symbol sheet but no extractable images')
                # A symbol sheet may also carry a code/description table — keep it as notes.
                notes = _extract_notes(ws, table)
                if notes:
                    section_notes.setdefault(section, []).extend(notes[:50])
            elif kind == 'codes':
                lookup = _extract_codes(ws, table)
                if lookup:
                    overlay = rule.get('overlay') or _predict_sheet(ws.title, 0).get('overlay')
                    section_lookups.setdefault(section, {}).setdefault(
                        overlay or f'_standalone:{ws.title}', {},
                    ).update(lookup)
                    det = section_details.setdefault(section, {'sheets': [], 'rows_imported': 0, 'images': []})
                    det['sheets'].append(ws.title)
                    det['rows_imported'] += len(lookup)
                else:
                    section_notes.setdefault(section, []).extend(_extract_notes(ws, table))
                    if ws.title.strip().lower() != 'instrument id table':
                        warnings.append(f'{ws.title}: no code/description rows detected; original cells preserved')
            elif kind in ('numbering', 'notes'):
                notes = _extract_notes(ws, table)
                if notes:
                    section_notes.setdefault(section, []).extend(notes)
                    det = section_details.setdefault(section, {'sheets': [], 'rows_imported': 0, 'images': []})
                    det['sheets'].append(ws.title)
                    det['rows_imported'] += len(notes)
                else:
                    warnings.append(f'{ws.title}: no text content detected')
            else:
                section_notes.setdefault(section, []).extend(_extract_notes(ws, table))
                warnings.append(f'{ws.title}: unrecognised table layout; preserved as reference text')
        except Exception as exc:
            logger.exception('[LegendWorkbook] sheet %s failed', ws.title)
            warnings.append(f'{ws.title}: {exc}')

    # ── Persist one legend per section ──────────────────────────────────
    created: List[Dict[str, Any]] = []
    updated: List[Dict[str, Any]] = []
    sections_touched = set(section_sheets)
    with transaction.atomic():
        for section in sorted(sections_touched):
            imgs = section_images.get(section, [])
            saved = (_inline_symbol_images(imgs, section) if valve_mto
                     else _save_symbol_images(imgs, section=section, project=project))
            images_per_section[section] = len(saved)
            for image, record in zip(imgs, saved):
                image['sheet_content']['images'].append({
                    'name': record['symbol_name'], 'path': record['path'],
                    'row': image['row'], 'column': image['column'],
                    **({'data_url': record['data_url']} if valve_mto else {}),
                })
            det = section_details.setdefault(section, {'sheets': [], 'rows_imported': 0, 'images': []})
            det['images'] = saved
            lookups = section_lookups.get(section, {})
            notes = section_notes.get(section, [])
            label = SECTION_LABELS.get(section, section)

            # Build the definition — overlay EVERY mapped code table onto its
            # field of the section default template (fluid codes → service,
            # piping material codes → spec, insulation codes → insulation, …).
            # Standalone tables (no mapped field) ride along as a lookup field
            # or notes text so nothing is lost.
            definition = copy.deepcopy(section_definitions.get(section))
            overlay_map = {k: v for k, v in lookups.items() if not k.startswith('_standalone:')}
            standalone_tables = {k.split(':', 1)[1]: v for k, v in lookups.items() if k.startswith('_standalone:')}
            standalone = {
                code: description
                for lookup in standalone_tables.values()
                for code, description in lookup.items()
            } or None
            extra_note_lines: List[str] = list(notes or [])
            if overlay_map and DEFAULT_TEMPLATES.get(section):
                if definition is None:
                    definition = copy.deepcopy(DEFAULT_TEMPLATES[section]['definition'])
                for field_key, lookup in overlay_map.items():
                    for field in definition.get('fields') or []:
                        if field.get('key') == field_key:
                            field['lookup'] = lookup
                            break
                    else:
                        # Template has no such field — keep the table as notes.
                        extra_note_lines.extend(f'{code} = {desc}' for code, desc in lookup.items())
            if definition is None and standalone:
                definition = _build_codes_definition(section, label, standalone, None, None)
            if definition is None and extra_note_lines:
                definition = _build_notes_definition(label, extra_note_lines)
            if definition is None:
                definition = _build_notes_definition(label, [])
            definition['imported_sheets'] = section_sheets[section]
            if definition is not None and extra_note_lines:
                definition.setdefault('notes', '')
                definition['notes'] = (definition['notes'] + '\n' if definition['notes'] else '') + '\n'.join(extra_note_lines[:MAX_NOTES_LINES])

            name = f'{pack_name} — {label}'
            if valve_mto:
                name = name[:200]
            # Deactivate siblings in this project scope, then update the
            # existing active project legend in place or create a new one.
            qs = (ValveMTOLegend.objects.filter(created_by=user, section=section) if valve_mto
                  else PidCheckerV2LegendSheet.objects.filter(project=project, section=section))
            existing = qs.filter(is_active=True).first() or qs.order_by('-updated_at').first()
            qs.update(is_active=False)
            if existing:
                existing.name = name
                existing.description = f'Imported from legend workbook "{pack_name}".'
                existing.definition = definition
                existing.is_active = True
                if user and not existing.created_by_id:
                    existing.created_by = user
                existing.save()
                updated.append({'legend_id': str(existing.pk if valve_mto else existing.legend_id), 'section': section,
                                'section_label': label, 'name': existing.name})
            else:
                model = ValveMTOLegend if valve_mto else PidCheckerV2LegendSheet
                obj = model.objects.create(
                    created_by=user,
                    **({} if valve_mto else {'project': project}),
                    section=section,
                    name=name,
                    description=f'Imported from legend workbook "{pack_name}".',
                    definition=definition,
                    is_active=True,
                )
                created.append({'legend_id': str(obj.pk if valve_mto else obj.legend_id), 'section': section,
                                'section_label': label, 'name': obj.name})

    logger.info('[LegendWorkbook] %s → %d created, %d updated, images=%s, skipped=%d',
                pack_name, len(created), len(updated), images_per_section, len(skipped))
    summary = {
        'created': created,
        'updated': updated,
        'sections': len(created) + len(updated),
        'images': images_per_section,
        'image_count': sum(images_per_section.values()),
        'skipped': skipped,
        'warnings': warnings,
        'file': filename,
        'pack_name': pack_name,
    }

    # Merge legend refs into the per-section archive detail.
    for entry in created + updated:
        det = section_details.setdefault(entry['section'], {'sheets': [], 'rows_imported': 0, 'images': []})
        det['legend_id'] = entry['legend_id']
        det['legend_name'] = entry['name']

    # Archive the import record (workbook bytes + JSON summary) to the default
    # storage (S3 when enabled). Best-effort — never breaks the import.
    summary['archive'] = {} if valve_mto else _archive_import(
        file_bytes, filename, user=user, project=project,
        sheet_map=sheet_map, summary=summary, section_details=section_details,
        warnings=warnings,
    )
    return summary


def _inline_symbol_images(images: List[Dict[str, Any]], section: str) -> List[Dict[str, str]]:
    """Keep small workbook references inside the user's isolated JSON definition."""
    import base64
    from PIL import Image

    saved = []
    used = set()
    for index, item in enumerate(images):
        name = _dedupe_name(_norm(item['name']).upper() or 'SYMBOL', used)
        with Image.open(io.BytesIO(item['data'])) as picture:
            rgba = picture.convert('RGBA')
            background = Image.new('RGB', rgba.size, 'white')
            background.paste(rgba, mask=rgba.getchannel('A'))
            buffer = io.BytesIO()
            background.save(buffer, format='PNG')
        saved.append({
            'symbol_name': name, 'path': f'inline:{section}:{index}',
            'data_url': 'data:image/png;base64,' + base64.b64encode(buffer.getvalue()).decode('ascii'),
        })
    return saved


# Soft-coded: parallel storage writes for legend symbol images. With S3 as
# the default storage (production), each image used to be its own serial
# HTTPS PUT inside the request — a 100+ symbol workbook blew past the
# client's 120 s timeout in production. Storage PUTs are idempotent and
# independent, so they run in a small thread pool; DB rows are written
# serially afterwards (fast, local). Storage operations don't participate
# in the surrounding transaction.atomic() anyway (S3 is not transactional),
# so behaviour is unchanged — just faster.
LEGEND_IMAGE_UPLOAD_WORKERS = 6


def _save_symbol_images(imgs: List[Dict[str, Any]], *, section: str, project) -> List[Dict[str, str]]:
    """Create/replace LegendSymbolImage rows. Returns one record per saved
    image: {symbol_name, path} (storage path — S3 key when S3 is the default
    storage). LegendSymbolImage.project is a pid_verification.PIDVProject FK
    (NOT the organizer project used for legend sheets) — resolve the matching
    PIDVProject by project_id, falling back to project=None (still usable via
    the cross-project fallback tier in SymbolImagesListView)."""
    from apps.pid_verification.models import PIDVProject
    from django.core.files.base import ContentFile
    from django.core.files.storage import default_storage
    from concurrent.futures import ThreadPoolExecutor
    import uuid as _uuid

    pidv_project = None
    if project is not None:
        pidv_project = PIDVProject.objects.filter(project_id=project.project_id).first()

    from ..views import _normalize_symbol_image

    # Phase 0 (serial, fast): dedupe names + resolve existing rows (DB reads).
    prepared = []
    used: set = set()
    for item in imgs:
        name = _dedupe_name(_norm(item['name']).upper() or 'SYMBOL', used)
        data, content_type = _normalize_symbol_image(item['data'], '.png')
        obj = LegendSymbolImage.objects.filter(
            project=pidv_project, section=section, symbol_name=name,
        ).first()
        old_path = obj.image_file.name if (obj is not None and obj.image_file) else None
        prepared.append({
            'symbol_name': name, 'data': data, 'content_type': content_type,
            'obj': obj, 'old_path': old_path,
        })

    # Phase 1 (parallel): storage deletes+PUTs only — no DB access in threads.
    def _store(p):
        if p['old_path']:
            try:
                default_storage.delete(p['old_path'])
            except Exception:
                logger.warning('[LegendWorkbook] could not delete old symbol image %s', p['old_path'])
        p['path'] = default_storage.save(f'{_uuid.uuid4().hex}.png', ContentFile(p['data']))
        return p

    if len(prepared) > 1:
        with ThreadPoolExecutor(max_workers=LEGEND_IMAGE_UPLOAD_WORKERS) as pool:
            prepared = list(pool.map(_store, prepared))
    elif prepared:
        prepared = [_store(prepared[0])]

    # Phase 2 (serial): DB rows with the already-stored paths (no storage I/O).
    saved: List[Dict[str, str]] = []
    for p in prepared:
        obj = p['obj']
        if obj is None:
            obj = LegendSymbolImage(project=pidv_project, section=section, symbol_name=p['symbol_name'])
        obj.content_type = p['content_type']
        obj.image_file.name = p['path']
        obj.save()
        saved.append({'symbol_name': p['symbol_name'], 'path': p['path']})
    return saved


def _archive_import(file_bytes: bytes, filename: str, *, user, project,
                    sheet_map: Dict[str, Optional[str]], summary: Dict[str, Any],
                    section_details: Dict[str, Any], warnings: List[str]) -> Dict[str, str]:
    """Archive the structured import record through Django's default_storage
    (storage-agnostic: S3 when USE_S3 is on, local disk otherwise).

    Saves the original workbook bytes + a JSON summary under
    WORKBOOK_ARCHIVE_PREFIX/<project_id|global>/. NEVER raises — an archiving
    failure appends a warning and returns empty paths instead of breaking the
    import.
    """
    try:
        from django.core.files.base import ContentFile
        from django.core.files.storage import default_storage
        from django.utils import timezone

        ts = timezone.now().strftime('%Y%m%d_%H%M%S')
        scope = str(project.project_id) if project is not None else 'global'
        safe_name = re.sub(r'[^A-Za-z0-9._-]+', '_', filename or 'workbook.xlsx')
        wb_path = f'{WORKBOOK_ARCHIVE_PREFIX}/{scope}/{ts}_{safe_name}'
        json_path = f'{WORKBOOK_ARCHIVE_PREFIX}/{scope}/{ts}_summary.json'

        default_storage.save(wb_path, ContentFile(file_bytes))

        import json as _json
        record = {
            'imported_at': timezone.now().isoformat(),
            'user_id': getattr(user, 'pk', None),
            'filename': filename,
            'sheet_map': sheet_map,
            'sections': section_details,
            'created': summary.get('created', []),
            'updated': summary.get('updated', []),
            'skipped': summary.get('skipped', []),
            'warnings': summary.get('warnings', []),
        }
        default_storage.save(json_path, ContentFile(
            _json.dumps(record, indent=2, default=str).encode('utf-8'),
        ))
        logger.info('[LegendWorkbook] archived import → %s , %s', wb_path, json_path)
        return {'workbook': wb_path, 'summary': json_path}
    except Exception as exc:                             # pragma: no cover
        logger.warning('[LegendWorkbook] archive failed (import unaffected): %s', exc)
        warnings.append(f'Archive to storage failed: {exc}')
        return {'workbook': '', 'summary': ''}
