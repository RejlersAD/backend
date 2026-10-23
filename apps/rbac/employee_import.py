"""HR-IMP-24: reviewed employee register import, using existing identity stores."""
import csv
import io
import re
from collections import defaultdict
from zipfile import BadZipFile

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from apps.hr_core.services import EmployeeService
from apps.rbac.models import Organization, Role, UserProfile, UserRole

UPDATE_FIELDS = ('first_name', 'last_name', 'department', 'job_title', 'phone', 'employee_id', 'location')
FIELD_LABELS = ('First Name', 'Last Name', 'Department', 'Job Title', 'Phone', 'Employee ID', 'Location')


def parse_register(upload):
    extension = upload.name.rsplit('.', 1)[-1].lower()
    if extension not in {'csv', 'xlsx', 'xlsm'}:
        raise ValueError('Upload a CSV or Excel (.xlsx/.xlsm) file.')
    if upload.size > 5 * 1024 * 1024:
        raise ValueError('File exceeds the 5 MB limit.')
    def header(value):
        return re.sub(r'[^a-z0-9]+', '_', str(value or '').strip().lower()).strip('_')
    if extension == 'csv':
        data = list(csv.reader(io.StringIO(upload.read().decode('utf-8-sig'))))
    else:
        try:
            workbook = load_workbook(upload, read_only=True, data_only=True)
        except (BadZipFile, InvalidFileException, KeyError) as error:
            raise ValueError('This is not a readable Excel workbook. Save it as XLSX or CSV and try again.') from error
        try:
            data = list(workbook.active.iter_rows(values_only=True))
        finally:
            workbook.close()
    if not data:
        raise ValueError('The file is empty.')
    headers = [header(value) for value in data[0]]
    named = [value for value in headers if value]
    if len(named) != len(set(named)):
        raise ValueError('The file contains duplicate column headings.')
    if not {'email', 'employee_id', 'username'}.intersection(headers):
        raise ValueError('Include Email, Employee ID or Username to identify employees.')
    rows, ignored = [], 0
    for number, values in enumerate(data[1:], 2):
        if not any(str(value).strip() for value in values if value is not None):
            ignored += 1
            continue
        row = {key: str(value).strip() if value is not None else '' for key, value in zip(headers, values) if key}
        rows.append((number, row))
    return rows, ignored


def import_register(upload, *, actor, preview, mode, update_fields, organization_id=None, role_id=None, audit=None):
    if mode not in {'update_existing', 'upsert'}:
        raise ValueError('Choose update_existing or upsert import mode.')
    if not isinstance(update_fields, list) or any(field not in UPDATE_FIELDS for field in update_fields):
        raise ValueError('Select supported employee fields to update.')
    fields = set(update_fields)
    rows, ignored = parse_register(upload)
    organization = role = None
    if mode == 'upsert':
        organization = Organization.objects.filter(pk=organization_id).first() if organization_id else None
        role = Role.objects.filter(pk=role_id, is_active=True).first() if role_id else None
        if not organization or not role:
            raise ValueError('Select a valid organization and active role for new employees.')

    User = get_user_model()
    results = {'success': [], 'updated': [], 'unchanged': [], 'failed': [], 'skipped': []}
    preview_rows = []
    seen = defaultdict(set)
    seen_profiles = set()
    with transaction.atomic():
        profiles_query = UserProfile.objects.select_related('user', 'organization', 'canonical_employee')
        users_query = User.objects.all()
        if not preview:
            # Serializes overlapping imports against the same identity records.
            list(users_query.select_for_update().values_list('pk', flat=True))
            list(UserProfile.objects.select_for_update().values_list('pk', flat=True))
        profiles = list(profiles_query)
        indexes = {field: defaultdict(list) for field in ('email', 'employee_id', 'username')}
        for profile in profiles:
            for field in indexes:
                value = getattr(profile, field) if field == 'employee_id' else getattr(profile.user, field)
                if value:
                    indexes[field][value.strip().casefold()].append(profile)
        users_by_email, users_by_username = defaultdict(list), defaultdict(list)
        for user in users_query:
            if user.email:
                users_by_email[user.email.strip().casefold()].append(user)
            users_by_username[user.username.strip().casefold()].append(user)

        for number, row in rows:
            report = {'row': number, 'email': row.get('email', ''), 'employee_id': row.get('employee_id', ''),
                      'name': f"{row.get('first_name', '')} {row.get('last_name', '')}".strip()}
            try:
                candidates_by_id = {}
                candidate_sets = []
                matched_by = []
                for field, index in indexes.items():
                    value = row.get(field, '').casefold()
                    if not value:
                        continue
                    if value in seen[field]:
                        raise ValueError(f'Duplicate {field.replace("_", " ")} in this file; each employee must appear once.')
                    candidates = index.get(value, [])
                    if candidates:
                        candidates_by_id.update({candidate.pk: candidate for candidate in candidates})
                        candidate_sets.append({candidate.pk for candidate in candidates})
                        matched_by.append(field)
                matching_ids = set.intersection(*candidate_sets) if candidate_sets else set()
                if candidate_sets and not matching_ids:
                    raise ValueError('Identity conflict: email, employee ID or username point to different employees.')
                if len(matching_ids) > 1:
                    raise ValueError('Ambiguous identity: more than one employee matches all supplied identifiers.')
                profile = candidates_by_id[next(iter(matching_ids))] if matching_ids else None
                if profile:
                    report['profile_id'] = str(profile.pk)
                if profile and profile.is_deleted:
                    raise ValueError('This employee is archived; restore the original record before importing.')
                if profile and profile.canonical_employee_id and profile.canonical_employee.user_id not in (None, profile.user_id):
                    raise ValueError('Canonical employee identity belongs to another account; reconcile it before importing.')
                if profile and profile.pk in seen_profiles:
                    raise ValueError('Duplicate employee match in this file; each employee must appear once.')
                if profile:
                    seen_profiles.add(profile.pk)
                # Reserve identities only after resolving conflicts. No duplicate creation within a file.
                for field in indexes:
                    if row.get(field):
                        seen[field].add(row[field].casefold())
                if profile is None and mode == 'update_existing':
                    report.update(state='skipped', reason='No existing employee matches; creation is disabled.')
                    results['skipped'].append(report); preview_rows.append(report)
                    continue
                email = row.get('email', '')
                if email:
                    try:
                        validate_email(email)
                    except ValidationError as error:
                        raise ValueError('Invalid email address.') from error
                    if len(email) > User._meta.get_field('email').max_length:
                        raise ValueError('Email address is too long.')
                elif profile is None:
                    raise ValueError('Email is missing. A real email is required to create a new employee.')
                if profile is None:
                    username = row.get('username') or email
                    if len(username) > User._meta.get_field('username').max_length:
                        raise ValueError('Username is too long.')
                    if users_by_email.get(email.casefold()) or users_by_username.get(username.casefold()):
                        raise ValueError('An account already uses this email or username without a matching employee profile; reconcile that account first.')
                    if not row.get('username'):
                        if email.casefold() in seen['username']:
                            raise ValueError('Duplicate generated username in this file.')
                        seen['username'].add(email.casefold())
                else:
                    # A stale email/username must never identify somebody else's orphan account.
                    for field, owners in [('email', users_by_email), ('username', users_by_username)]:
                        field_owners = owners.get(row.get(field, '').casefold(), [])
                        if field_owners and all(user.pk != profile.user_id for user in field_owners):
                            raise ValueError(f'Identity conflict: {field} belongs to a different account.')
                changes = {}
                for field in (fields if profile else UPDATE_FIELDS):
                    value = row.get(field, '')
                    target = profile.user if profile and field in {'first_name', 'last_name'} else profile
                    current = getattr(target, field, '') if target else ''
                    model = User if field in {'first_name', 'last_name'} else UserProfile
                    model_field = model._meta.get_field(field)
                    if value and model_field.max_length and len(value) > model_field.max_length:
                        raise ValueError(f'{field.replace("_", " ").title()} exceeds {model_field.max_length} characters.')
                    if value and value != current:
                        changes[field] = value
                report.update(state='existing' if profile else 'new', matched_by=matched_by,
                              changed_fields=list(changes), incoming={field: row.get(field, '') for field in UPDATE_FIELDS})
                if preview:
                    preview_rows.append(report)
                    continue
                if profile:
                    # Keep the lookup snapshot immutable if this row rolls back.
                    profile = UserProfile.objects.select_related('user', 'canonical_employee').get(pk=profile.pk)
                # New and updated employee/account/master changes commit together for each row.
                with transaction.atomic():
                    if profile:
                        if not changes:
                            report.update(reason='Selected fields already match.', state='unchanged')
                            results['unchanged'].append(report)
                            continue
                        user_fields, profile_fields = [], []
                        for field, value in changes.items():
                            target = profile.user if field in {'first_name', 'last_name'} else profile
                            setattr(target, field, value)
                            (user_fields if target is profile.user else profile_fields).append(field)
                        if user_fields:
                            # Saving User fires login bookkeeping signals. Import only
                            # the selected identity fields; sync HR explicitly below.
                            User.objects.filter(pk=profile.user_id).update(**{field: getattr(profile.user, field) for field in user_fields})
                        if profile_fields:
                            from django.utils import timezone
                            # Match the existing bulk import's bounded writes. Profile
                            # save signals synchronize biometric records as well; that
                            # integration is outside a master-register data correction.
                            UserProfile.objects.filter(pk=profile.pk).update(
                                **{field: getattr(profile, field) for field in profile_fields}, updated_at=timezone.now())
                        reviewed_user = {field: getattr(profile.user, field) for field in
                            ('first_name', 'last_name', 'email', 'username', 'password', 'is_active', 'last_login')}
                        reviewed_profile = {field: getattr(profile, field) for field in
                            ('department', 'job_title', 'phone', 'employee_id', 'location', 'manager_id', 'organization_id', 'status')}
                        EmployeeService.sync_from_rbac_profile(profile, set(changes))
                        # Canonical compatibility signals mirror a full record back
                        # to RBAC, including old master codes/blank contact values.
                        # Keep the reviewed import values and unselected account
                        # settings intact within the same row transaction.
                        User.objects.filter(pk=profile.user_id).update(**reviewed_user)
                        UserProfile.objects.filter(pk=profile.pk).update(**reviewed_profile)
                        report['state'] = 'updated'
                        bucket = 'updated'
                    else:
                        username = row.get('username') or email
                        if len(username) > User._meta.get_field('username').max_length:
                            raise ValueError('Username is too long.')
                        user = User.objects.create_user(username=username, email=email,
                            first_name=row.get('first_name', ''), last_name=row.get('last_name', ''), password=None)
                        # No spreadsheet-driven access status or credential changes.
                        profile = UserProfile.objects.create(user=user, organization=organization,
                            **{field: row.get(field, '') for field in UPDATE_FIELDS if field not in {'first_name', 'last_name'}})
                        # Profile signals may assign Default automatically. New employees
                        # receive exactly the role explicitly selected for this import.
                        UserRole.objects.filter(user_profile=profile).exclude(role=role).delete()
                        UserRole.objects.update_or_create(user_profile=profile, role=role,
                            defaults={'assigned_by': actor, 'is_primary': True})
                        EmployeeService.sync_from_rbac_profile(profile)
                        report.update(state='created', email_sent=False, email_queued=False)
                        report['profile_id'] = str(profile.pk)
                        bucket = 'success'
                    if audit:
                        audit(user=actor, action='bulk_upload', resource_type='UserProfile', resource_id=str(profile.pk),
                              resource_repr='Employee register import', metadata={'row': number, 'state': report['state'], 'fields': list(changes)})
                # Keep the initial identity snapshot unchanged throughout this file.
                # Earlier updates must not make a row rejected in preview importable.
                results[bucket].append(report)
            except (ValueError, ValidationError) as error:
                report.update(state='invalid', error=str(error)); results['failed'].append(report); preview_rows.append(report)
            except Exception:
                report.update(state='invalid', error='This row could not be saved. No changes from this row were committed.')
                results['failed'].append(report); preview_rows.append(report)
        if preview:
            counts = {state: sum(row['state'] == state for row in preview_rows) for state in ('new', 'existing', 'invalid', 'skipped')}
            return {'preview': True, 'summary': {'total': len(rows), 'ignored_blank': ignored, **counts},
                    'columns': [{'id': key, 'label': label} for key, label in zip(UPDATE_FIELDS, FIELD_LABELS)], 'rows': preview_rows}
    return {'message': 'Employee register import completed.', 'summary': {'total_processed': len(rows),
            'ignored_blank': ignored, 'created': len(results['success']), 'successful': len(results['success']),
            'updated': len(results['updated']), 'unchanged': len(results['unchanged']),
            'failed': len(results['failed']), 'skipped': len(results['skipped']), 'emails_sent': 0, 'emails_queued': 0}, 'details': results}
