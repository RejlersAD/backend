import base64
from io import BytesIO
from datetime import datetime, timezone
from types import SimpleNamespace
from unittest import TestCase
from unittest.mock import patch

import fitz
from docx import Document
from docx.oxml.ns import qn
from PIL import Image as PILImage
from PyPDF2 import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas

from apps.procurement.services.purchase_order_exports import (
    _approval_display,
    _html_blocks,
    build_purchase_order_docx,
    build_purchase_order_pdf,
)


class PurchaseOrderExportTests(TestCase):
    def test_empty_introduction_and_editor_markup_omit_the_entire_scope_page(self):
        empty_narratives = (
            '', ' \n\t\u00a0 ', '<p><br></p>',
            '<div><p> &nbsp;\t</p><p><br /></p></div>', '&nbsp;',
            '<p>\u200b\u200c\u200d&#xfeff;&nbsp;<br></p>',
            '&lt;p&gt;&lt;br&gt;&lt;/p&gt;',
            '<div data-po-page-break="true">Page Break</div>', '<table></table>',
        )
        for narrative in empty_narratives:
            for show_heading in (False, True):
                with self.subTest(narrative=narrative, show_heading=show_heading):
                    order = self._order()
                    order.description = narrative
                    order.contact_persons = {'order_introduction': ' \t\u00a0 ',
                                             'show_scope_heading': show_heading}
                    content, warnings = build_purchase_order_pdf(order)
                    self.assertEqual(warnings, [])
                    with fitz.open(stream=content, filetype='pdf') as pdf:
                        self.assertEqual(len(pdf), 2)
                        self.assertIn('SUMMARY OF PRICES', pdf[1].get_text())
                        text = '\n'.join(page.get_text() for page in pdf)
                        self.assertNotIn('PURCHASE ORDER:', text)
                        self.assertNotIn('PO DESCRIPTION & SCOPE', text)
                        self.assertNotIn(order.title, pdf[1].get_text())
                        self.assertIn('Order Confirmation:', pdf[0].get_text())
                    word = Document(BytesIO(build_purchase_order_docx(order)))
                    paragraphs = [paragraph.text for paragraph in word.paragraphs]
                    self.assertNotIn(order.title, paragraphs)
                    self.assertFalse(any(text.startswith('PURCHASE ORDER:') for text in paragraphs))
                    self.assertNotIn('PO DESCRIPTION & SCOPE', paragraphs)
                    self.assertEqual([paragraph.text for paragraph in word.paragraphs
                                      if paragraph.paragraph_format.page_break_before], ['SUMMARY OF PRICES'])

    def test_intro_only_body_only_and_legacy_intro_keep_scope_without_title_as_body(self):
        cases = (
            ({'order_introduction': 'Agreed introduction only.'}, '<p><br></p>', 'Agreed introduction only.'),
            ({'order_introduction': ''}, '<p>Recorded scope only.</p>', 'Recorded scope only.'),
            ({}, '', 'issue this purchase order to'),
        )
        for contacts, narrative, expected in cases:
            with self.subTest(contacts=contacts, narrative=narrative):
                order = self._order()
                order.description = narrative
                order.contact_persons = {**contacts, 'show_scope_heading': False}
                content, warnings = build_purchase_order_pdf(order)
                self.assertEqual(warnings, [])
                with fitz.open(stream=content, filetype='pdf') as pdf:
                    self.assertEqual(len(pdf), 3)
                    scope = pdf[1].get_text()
                    self.assertIn('PURCHASE ORDER:', scope)
                    self.assertIn(expected, scope)
                    self.assertEqual(scope.count(order.title), 1)
                    self.assertNotIn('PO DESCRIPTION & SCOPE', scope)
                    self.assertIn('SUMMARY OF PRICES', pdf[2].get_text())
                word = Document(BytesIO(build_purchase_order_docx(order)))
                paragraphs = [paragraph.text for paragraph in word.paragraphs]
                self.assertNotIn(order.title, paragraphs)
                self.assertEqual(sum(text.startswith('PURCHASE ORDER:') for text in paragraphs), 1)
                self.assertIn(expected, '\n'.join(paragraphs))
                self.assertNotIn('PO DESCRIPTION & SCOPE', paragraphs)

    def test_image_only_scope_survives_without_introduction_or_heading(self):
        image = BytesIO()
        PILImage.new('RGB', (19, 11), 'navy').save(image, 'PNG')
        encoded = base64.b64encode(image.getvalue()).decode()
        order = self._order()
        order.contact_persons = {'order_introduction': '', 'show_scope_heading': False}
        order.description = f'<p><br></p><img src="data:image/png;base64,{encoded}" width="100"><p>&nbsp;</p>'
        content, warnings = build_purchase_order_pdf(order)
        self.assertEqual(warnings, [])
        with fitz.open(stream=content, filetype='pdf') as pdf:
            self.assertEqual(len(pdf), 3)
            self.assertTrue(any(image[2:4] == (19, 11) for image in pdf[1].get_images()))
            self.assertEqual(pdf[1].get_text().count(order.title), 1)
        word = Document(BytesIO(build_purchase_order_docx(order)))
        self.assertTrue(word.inline_shapes)
        self.assertNotIn(order.title, [paragraph.text for paragraph in word.paragraphs])

    def test_table_only_scope_preserves_authored_cells_and_empty_table_grid(self):
        for first_cell in ('Recorded service deliverable', '&nbsp;'):
            with self.subTest(first_cell=first_cell):
                order = self._order()
                order.contact_persons = {'order_introduction': '', 'show_scope_heading': False}
                order.description = f'<table><tr><td>{first_cell}</td><td><br></td></tr></table>'
                content, warnings = build_purchase_order_pdf(order)
                self.assertEqual(warnings, [])
                with fitz.open(stream=content, filetype='pdf') as pdf:
                    self.assertEqual(len(pdf), 3)
                    self.assertTrue(pdf[1].get_drawings())
                    if first_cell != '&nbsp;':
                        self.assertIn(first_cell, pdf[1].get_text())
                word = Document(BytesIO(build_purchase_order_docx(order)))
                tables = [table for table in word.tables if table.style.name == 'Table Grid'
                          and len(table.rows) == 1 and len(table.columns) == 2]
                self.assertEqual(len(tables), 1)
                self.assertEqual(tables[0].cell(0, 0).text.strip(), '' if first_cell == '&nbsp;' else first_cell)

    def test_pending_po_shows_ceo_identity_without_inventing_a_completed_approval(self):
        order = self._order()
        order.status, order.approved_at = 'draft', None
        order.approved_by_name = 'Preselected Approver'
        order.approval_log = [{'user_id': 'po-reviewer', 'stage': 'Final Management Sign-off', 'status': 'Pending'}]
        content, _ = build_purchase_order_pdf(order)
        with fitz.open(stream=content, filetype='pdf') as pdf:
            text = '\n'.join(page.get_text() for page in pdf)
        self.assertIn('PO status: Draft', text)
        self.assertIn('Approval pending:', text)
        self.assertIn('Jarmo Suominen', text)
        self.assertIn('CEO, Rejlers Abu Dhabi', text)
        self.assertNotIn('Preselected Approver', text)
        self.assertNotIn('Approved by:', text)
        document = Document(BytesIO(build_purchase_order_docx(order)))
        text = '\n'.join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        self.assertIn('Approval pending:', text)
        self.assertIn('Jarmo Suominen', text)
        self.assertIn('CEO, Rejlers Abu Dhabi', text)
        self.assertNotIn('Preselected Approver', text)
        self.assertNotIn('Approved by:', text)

    def test_unassigned_draft_pdf_and_word_do_not_claim_approval_was_requested(self):
        order = self._order()
        order.status, order.approved_at = 'draft', None
        order.approved_by_name = 'Preselected Approver'
        order.approval_log = []
        content, warnings = build_purchase_order_pdf(order)
        self.assertEqual(warnings, [])
        with fitz.open(stream=content, filetype='pdf') as pdf:
            self.assertEqual(len(pdf), 3)
            text = '\n'.join(page.get_text() for page in pdf)
            self.assertTrue(pdf[0].search_for('Order Confirmation:'))
            self.assertIn('Approval not requested:', text)
            self.assertIn('Jarmo Suominen', text)
            self.assertIn('CEO, Rejlers Abu Dhabi', text)
            self.assertNotIn('Approval pending:', text)
            self.assertNotIn('Preselected Approver', text)
            self.assertNotIn('Approved by:', text)
        document = Document(BytesIO(build_purchase_order_docx(order)))
        cell = next(cell for table in document.tables for row in table.rows for cell in row.cells
                    if 'Approval not requested:' in cell.text)
        self.assertEqual(cell.vertical_alignment, 0)
        self.assertIn('Jarmo Suominen', cell.text)
        self.assertIn('CEO, Rejlers Abu Dhabi', cell.text)
        self.assertNotIn('Preselected Approver', cell.text)

    def test_linked_requisition_approval_and_source_rows_do_not_supply_po_approval(self):
        order = self._order()
        order.status, order.approved_at, order.approved_by_name = 'draft', None, ''
        order.pr_reference = SimpleNamespace(
            pr_number='PR-APPROVED', status='approved', approved_by_name='PR Signer',
            approved_at=datetime(2026, 9, 3, tzinfo=timezone.utc),
        )
        order.approval_log = [{
            'source': 'signed_purchase_requisition_pdf', 'external': True,
            'user_id': 'pr-reviewer', 'approver': 'PR Signer', 'status': 'Approved',
            'evidence_document_id': 'retained-pr-source',
        }]
        content, _ = build_purchase_order_pdf(order)
        with fitz.open(stream=content, filetype='pdf') as pdf:
            text = '\n'.join(page.get_text() for page in pdf)
        document = Document(BytesIO(build_purchase_order_docx(order)))
        word_text = '\n'.join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        for rendered in (text, word_text):
            self.assertIn('Approval not requested:', rendered)
            self.assertIn('Jarmo Suominen', rendered)
            self.assertIn('CEO, Rejlers Abu Dhabi', rendered)
            self.assertNotIn('PR Signer', rendered)
            self.assertNotIn('Approved by:', rendered)
            self.assertNotIn('Approval pending:', rendered)

    def test_pending_status_needs_an_assigned_internal_stage_and_stopped_routes_are_not_pending(self):
        order = self._order()
        order.approved_at, order.approved_by_name = None, ''
        for status in ('draft', 'sent', 'in_progress'):
            with self.subTest(status=status):
                order.status = status
                order.approval_log = [{
                    'approver_email': 'signer@example.test', 'stage': 'Final Management Sign-off',
                    'status': 'in_review',
                }]
                display = _approval_display(order)
                self.assertEqual(display['heading'], 'Approval pending:')
                self.assertEqual(display['name'], 'Jarmo Suominen')
                self.assertIn('CEO, Rejlers Abu Dhabi', display['title'])
        order.status = 'draft'
        for row in (
            {'status': 'pending', 'stage': 'Final Management Sign-off'},
            {'status': 'pending', 'user_id': 'pr-reviewer', 'source': 'signed_purchase_requisition_pdf'},
            {'status': 'pending', 'user_id': 'source-reviewer', 'evidence_document_id': 'source'},
        ):
            with self.subTest(row=row):
                order.approval_log = [row]
                self.assertEqual(_approval_display(order)['heading'], 'Approval not requested:')
        order.approval_log = [
            {'status': 'Rejected', 'user_id': 'first-reviewer'},
            {'status': 'Pending', 'user_id': 'next-reviewer'},
        ]
        self.assertEqual(_approval_display(order)['heading'], 'Approval record:')

    def test_pdf_and_docx_flag_mismatched_approval_instead_of_printing_signature(self):
        order = self._order()
        order.approved_by_name = 'Assigned Approver'
        order.approved_by_id = 'assigned'
        signature = BytesIO()
        PILImage.new('RGB', (181, 51), 'navy').save(signature, format='PNG')
        order.approval_signature = 'data:image/png;base64,' + base64.b64encode(signature.getvalue()).decode()
        order.approval_log = [{
            'user_id': 'assigned', 'approved_by_id': 'different-person',
            'status': 'Approved', 'signature': order.approval_signature,
        }]
        content, warnings = build_purchase_order_pdf(order)
        self.assertEqual(warnings, [])
        with fitz.open(stream=content, filetype='pdf') as pdf:
            self.assertIn('Approval review is required', ' '.join(page.get_text() for page in pdf))
            self.assertFalse(any(image[2:4] == (181, 51) for page in pdf for image in page.get_images()))
        document = Document(BytesIO(build_purchase_order_docx(order)))
        self.assertIn('Approval review is required', ' '.join(
            paragraph.text for table in document.tables for row in table.rows
            for cell in row.cells for paragraph in cell.paragraphs
        ))
        self.assertNotIn(signature.getvalue(), [part.blob for part in document.part.package.parts])

    def test_later_lifecycle_with_missing_approval_details_is_unknown_not_pending(self):
        for status in ('sent', 'completed', 'approved'):
            with self.subTest(status=status):
                order = self._order()
                order.status, order.approved_at, order.approved_by_name = status, None, ''
                content, _ = build_purchase_order_pdf(order)
                with fitz.open(stream=content, filetype='pdf') as pdf:
                    text = '\n'.join(page.get_text() for page in pdf)
                self.assertIn('Approval record:', text)
                self.assertIn('Not recorded', text)
                self.assertNotIn('Not yet approved', text)
                self.assertNotIn('Approval pending:', text)
                document = Document(BytesIO(build_purchase_order_docx(order)))
                text = '\n'.join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
                self.assertIn('Approval record:', text)
                self.assertIn('Not recorded', text)

    def test_recorded_timestamp_supplies_date_when_separate_approval_date_is_missing(self):
        order = self._order()
        self.assertIsNone(order.approved_date)
        content, _ = build_purchase_order_pdf(order)
        with fitz.open(stream=content, filetype='pdf') as pdf:
            text = pdf[0].get_text()
        self.assertIn('Approved by:', text)
        self.assertIn('2026-09-03', text)
        document = Document(BytesIO(build_purchase_order_docx(order)))
        text = '\n'.join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        self.assertIn('Date: 2026-09-03', text)

    def _order(self, attachments=None):
        return SimpleNamespace(
            vendor=SimpleNamespace(name='Test Vendor'),
            po_number='RAD-PRJ-PUR-0001_2026',
            po_date='2026-09-03',
            items=[{'description': 'Test item', 'quantity': 2, 'unit_price': 50, 'uom': 'EA'}],
            total_amount=105,
            tax_amount=5,
            discount_amount=0,
            vat_percentage=5,
            currency='USD',
            title='Test Purchase Order',
            description='<p>First scope&nbsp;paragraph</p><p>Second scope paragraph</p><ul><li>Required document</li></ul>',
            seller_reference='Vendor Contact',
            quote_ref='',
            seller_license_no='',
            seller_address='Vendor City',
            seller_contact_person='',
            seller_phone='+971 1 234 5678',
            seller_fax='+971 1 234 5679',
            seller_email='vendor@example.com',
            invoicing_attn='Accounts Payable',
            invoicing_emails=['aneef.thadikkarantavida@rejlers.ae'],
            company_fax='+971 2 639 7448',
            buyer_reference_pm='Test Buyer',
            buyer_reference_email='richahannah.thomas@rejlers.ae',
            contact_persons={},
            project_number='590001',
            rad_project_no='',
            payment_terms='30 days',
            payment_mode='Bank Transfer',
            delivery_terms='DAP',
            expected_delivery='2026-09-30',
            marking='RAD-PRJ-PUR-0001_2026',
            form_note='(PO no. to be used in all documents)',
            status='sent',
            approved_by_name='Jarmo Suominen',
            approved_by_title='Sr. Vice President, Middle East\nCEO, Rejlers Abu Dhabi',
            approved_at=datetime(2026, 9, 3, tzinfo=timezone.utc),
            approved_date=None,
            confirmation_date=None,
            attachments=attachments or [],
        )

    def _one_page_pdf(self):
        output = BytesIO()
        pdf = canvas.Canvas(output, pagesize=A4)
        pdf.drawString(50, 780, 'Supporting document')
        pdf.save()
        return output.getvalue()

    def test_pdf_adds_one_cover_and_source_per_attachment_in_order(self):
        attachments = [
            {
                'title': f'Attachment {index}',
                'description': f'Description {index}',
                'filename': f'attachment-{index}.pdf',
                's3_key': str(index),
                'content_type': 'application/pdf',
            }
            for index in range(1, 4)
        ]
        with patch(
            'apps.procurement.services.purchase_order_exports._download_attachment',
            return_value=self._one_page_pdf(),
        ):
            content, warnings = build_purchase_order_pdf(self._order(attachments))

        exported = PdfReader(BytesIO(content))
        self.assertEqual(warnings, [])
        # Cover/details, narrative, and price summary, then one cover and one
        # source page for each of the three attachments.
        self.assertEqual(len(exported.pages), 9)
        narrative_text = exported.pages[1].extract_text()
        self.assertIn('PURCHASE ORDER', narrative_text)
        # Header and footer use the official Rejlers image asset rather than a
        # substitute text/vector wordmark.
        xobjects = exported.pages[1]['/Resources'].get('/XObject', {})
        self.assertGreaterEqual(len(xobjects), 2)
        self.assertIn('Rejlers International Engineering Solutions', narrative_text)
        self.assertIn('Page 2', narrative_text)
        self.assertIn('First scope paragraph', narrative_text)
        self.assertIn('Second scope paragraph', narrative_text)
        self.assertNotIn('&nbsp;', narrative_text)
        first_page_text = exported.pages[0].extract_text()
        self.assertIn('Seller Address:', first_page_text)
        self.assertIn('Seller Name:', first_page_text)
        self.assertIn('Seller Ref. no:', first_page_text)
        self.assertIn('Contact Person:', first_page_text)
        self.assertRegex(
            first_page_text,
            r'Buyer\s+Reference:\s+Test Buyer\s+Procurement Manager\s+richahannah\.thomas@rejlers\.ae',
        )
        self.assertIn('aneef.thadikkarantavida@rejlers.ae', first_page_text)
        self.assertIn('Phone Number:', first_page_text)
        self.assertIn('Fax:', first_page_text)
        self.assertIn('Email:', first_page_text)
        self.assertNotIn('Phone / Email:', first_page_text)
        self.assertRegex(first_page_text, r'Seller Ref\. no:\s+—')
        self.assertRegex(first_page_text, r'Contact Person:\s+Phone Number:')
        price_summary_text = exported.pages[2].extract_text()
        self.assertIn('SUMMARY OF PRICES', price_summary_text)
        self.assertRegex(price_summary_text, r'Total Price:\s+USD 100\.00')
        self.assertRegex(price_summary_text, r'VAT \(5%\):\s+USD 5\.00')
        self.assertRegex(price_summary_text, r'Total Sum:\s+USD 105\.00')
        self.assertNotIn('Grand Total', price_summary_text)
        self.assertNotIn('AED', price_summary_text)
        first_cover_text = exported.pages[3].extract_text()
        self.assertIn('PURCHASE ORDER', first_cover_text)
        self.assertIn('Page 4', first_cover_text)
        self.assertIn('Description 1', first_cover_text)
        self.assertNotIn('attachment-1.pdf', first_cover_text)

    def test_word_keeps_native_price_summary_before_supporting_cover_pages(self):
        content, warnings = build_purchase_order_docx(self._order([{
            'title': 'Attachment 1',
            'description': 'Retained supporting cover',
            'filename': 'support.pdf',
        }]), with_warnings=True)
        document = Document(BytesIO(content))
        self.assertEqual(warnings, ['support.pdf: file could not be downloaded'])
        page_images = document.element.xpath('//wp:anchor/wp:docPr/@descr')
        self.assertEqual(page_images, ['Original PDF page 4'])
        self.assertEqual(len(document.sections), 2)
        rendered_text = '\n'.join(document.element.xpath('//w:t/text()'))
        header_text = '\n'.join(
            cell.text
            for table in document.sections[0].header.tables
            for row in table.rows
            for cell in row.cells
        )
        footer_text = '\n'.join(
            cell.text
            for table in document.sections[0].footer.tables
            for row in table.rows
            for cell in row.cells
        )

        self.assertIn('SUMMARY OF PRICES', rendered_text)
        self.assertRegex(rendered_text, r'Total Price:\s+USD 100\.00')
        self.assertRegex(rendered_text, r'VAT \(5%\):\s+USD 5\.00')
        self.assertRegex(rendered_text, r'Total Sum:\s+USD 105\.00')
        self.assertNotIn('Grand Total', rendered_text)
        self.assertNotIn('AED', rendered_text)
        self.assertIn('First scope\u00a0paragraph', rendered_text)
        self.assertIn('Second scope paragraph', rendered_text)
        self.assertNotIn('&nbsp;', rendered_text)
        self.assertIn('PURCHASE ORDER', header_text)
        self.assertIn('RAD-PRJ-PUR-0001_2026', header_text)
        self.assertIn('HOME OF THE', header_text)
        self.assertIn('Rejlers International Engineering Solutions', footer_text)
        self.assertIn('Page ', footer_text)

    def test_long_purchase_summary_stays_on_first_page_without_displacing_approval(self):
        order = self._order()
        order.title = (
            'Supply of Smart Interop Publisher License — May 2026 (PO Form 29) '
            '(1 No) for RFIN XLPE Project = USD 1,755.40 1 No: USD 1,755.40 '
            'equally shared in RAB Projects 5900863 (H2 Extraction) and 5901055 (ADOC)'
        )

        content, warnings = build_purchase_order_pdf(order)
        exported = PdfReader(BytesIO(content))
        first_page_text = exported.pages[0].extract_text()

        self.assertEqual(warnings, [])
        self.assertEqual(len(exported.pages), 3)
        self.assertIn('5901055', first_page_text)
        self.assertIn('(ADOC)', first_page_text)
        self.assertIn('Approved by:', first_page_text)
        self.assertIn('Order Confirmation:', first_page_text)

    def test_word_company_form_preserves_editable_cover_fields_and_branding(self):
        order = self._order()
        order.contact_persons['buyer_references'] = [{
            'name': 'Buyer & Contract Manager', 'designation': 'Project Manager',
            'email': 'buyer@example.test',
        }]
        order.seller_reference = 'SELLER-REF-726'
        order.seller_contact_person = 'Supplier Contact'
        document = Document(BytesIO(build_purchase_order_docx(order)))
        text = '\n'.join(document.element.xpath('//w:t/text()'))
        for expected in (
            'Seller:', 'Seller Address:', 'Invoicing Address:',
            'Accounts Payable', 'PO Box 39317', 'Fax: +971 2 639 7448',
            'Seller Reference:', 'Quote Ref.:', 'License No.:',
            'Buyer & Contract Manager', 'Project Manager', 'buyer@example.test',
            'Seller Name:', 'Seller Ref. no:', 'SELLER-REF-726',
            'Contact Person:', 'Supplier Contact', 'Phone Number:',
            'Fax:', '+971 1 234 5679', 'Email:', 'vendor@example.com',
        ):
            self.assertIn(expected, text)
        self.assertNotIn('<br', text)
        self.assertNotIn('<b>', text)
        self.assertNotIn('Phone / Email:', text)
        self.assertNotIn('Seller information', text)
        # Native table cells and paragraphs remain editable; the body is not
        # a screenshot of the PDF. Brand marks live in repeating page furniture.
        self.assertGreater(len(document.tables), 4)
        self.assertFalse(document.element.xpath('//w:drawing'))
        section = document.sections[0]
        self.assertAlmostEqual(section.page_width.mm, 210, delta=.1)
        self.assertAlmostEqual(section.page_height.mm, 297, delta=.1)
        header_title = section.header.tables[0].cell(0, 0).paragraphs[0]
        self.assertEqual(header_title.text, 'PURCHASE ORDER')
        self.assertEqual(next(run for run in header_title.runs if run.text).font.size.pt, 15)
        self.assertTrue(section.header._element.xpath('.//w:drawing'))
        self.assertEqual(len(section.footer._element.xpath('.//w:drawing')), 3)
        self.assertTrue(section.footer._element.xpath('.//w:instrText[text()=" PAGE "]'))
        self.assertEqual(len(document.element.xpath('//w:pPr/w:pageBreakBefore')), 2)

    def test_word_prices_match_pdf_columns_borders_and_totals(self):
        order = self._order()
        order.items_table_headers = {
            '__column_order': ['description', 'quantity', 'total_price'],
            'description': 'Agreed service', 'quantity': 'Quantity',
            'total_price': 'Amount',
        }
        document = Document(BytesIO(build_purchase_order_docx(order)))
        table = next(table for table in document.tables if table.cell(0, 0).text == 'Agreed service')
        self.assertEqual([cell.text for cell in table.rows[1].cells], ['Test item', '2', 'USD 100.00'])
        self.assertAlmostEqual(table.columns[0].width / table.columns[1].width, 3, delta=.01)
        self.assertIsNotNone(table.rows[0]._tr.find(qn('w:trPr')).find(qn('w:tblHeader')))
        totals = document.tables[-1]
        self.assertEqual([[cell.text for cell in row.cells] for row in totals.rows], [
            ['Total Price:', 'USD 100.00'], ['VAT (5%):', 'USD 5.00'], ['Total Sum:', 'USD 105.00'],
        ])
        self.assertEqual(totals.cell(2, 0)._tc.find('.//' + qn('w:shd')).get(qn('w:fill')), 'E2E8F0')

    def test_custom_opening_statement_matches_pdf_and_word_as_literal_text(self):
        order = self._order()
        custom = 'The Buyer & Seller agree to <specific> scope.\nDelivery follows the approved schedule.'
        order.contact_persons['order_introduction'] = custom
        word = Document(BytesIO(build_purchase_order_docx(order)))
        self.assertTrue(any(paragraph.text == custom for paragraph in word.paragraphs))
        content, _ = build_purchase_order_pdf(order)
        with fitz.open(stream=content, filetype='pdf') as pdf:
            text = pdf[1].get_text()
            for line in custom.splitlines():
                self.assertIn(line, text)
            self.assertNotIn('We, Rejlers', text)
        order.contact_persons['order_introduction'] = ''
        word = Document(BytesIO(build_purchase_order_docx(order)))
        self.assertFalse(any('We, Rejlers' in p.text for p in word.paragraphs))
        content, _ = build_purchase_order_pdf(order)
        with fitz.open(stream=content, filetype='pdf') as pdf:
            self.assertNotIn('We, Rejlers', pdf[1].get_text())
        del order.contact_persons['order_introduction']
        word = Document(BytesIO(build_purchase_order_docx(order)))
        self.assertTrue(any('We, Rejlers International Engineering Solutions (Buyer)' in p.text for p in word.paragraphs))

    def test_rich_text_normalization_decodes_entities_and_keeps_blocks(self):
        self.assertEqual(
            _html_blocks('&lt;p&gt;Alpha&amp;nbsp;Beta&lt;/p&gt;<div>Gamma<br>Delta</div>'),
            ['Alpha Beta', 'Gamma', 'Delta'],
        )

    def _assert_body_clear_of_footer(self, pdf):
        # Branded page furniture is blue/white. Business text is slate/black;
        # inspect actual glyph rectangles rather than page counts alone.
        furniture_colors = {0x3275B6, 0x0870AA, 0xFFFFFF}
        for page in pdf:
            for block in page.get_text('dict')['blocks']:
                for line in block.get('lines', []):
                    for span in line['spans']:
                        if span['color'] in furniture_colors or span['bbox'][1] < 34 * mm:
                            continue
                        self.assertLessEqual(
                            span['bbox'][3], page.rect.height - 40 * mm,
                            f'Page {page.number + 1} enters the footer: {span["text"]}',
                        )

    def _long_order(self):
        order = self._order()
        order.vendor.name = 'Synthetic International Instrumentation and Technical Services Limited'
        order.seller_address = '\n'.join(f'Seller address level {index:02d}' for index in range(8))
        order.invoicing_attn = 'Invoice contact: Synthetic Accounts and Commercial Administration'
        order.invoicing_emails = [
            f'accounting.department.region{index}.contract.invoices@example.invalid'
            for index in range(4)
        ]
        order.company_fax = 'INVOICEFAX9988'
        order.contact_persons = {'buyer_references': [
            {'name': f'Synthetic Buyer {index}', 'designation': 'Senior Commercial Administrator',
             'email': f'buyer{index}.regional.contracts@example.invalid'}
            for index in range(4)
        ]}
        order.payment_terms = 'Payment follows receipt of a fully checked invoice and accepted deliverables.'
        order.seller_contact_person = 'Synthetic Supplier Contact for Regional Commercial Administration'
        order.seller_email = 'supplier.commercial.department.contract.administration@example.invalid'
        order.approved_by_name = 'Synthetic Authorised Approver'
        order.approved_by_title = 'Director of Commercial Operations\nRegional Engineering Division'
        order.approved_date = '2026-09-12'
        order.confirmation_date = '2026-09-13'
        return order

    def _realistic_long_contact_order(self):
        # Match the reported cover's density: two buyers, three invoice email
        # addresses and the same long supplier contact in both cover columns.
        order = self._order()
        order.vendor.name = 'SYNTHETIC SURVEYS WORK MEASUREMENT & SPACE L.L.C.'
        order.title = 'Provision of Piping Design engineer for 2 Months'
        order.seller_address = 'Al Example Tower, Office M04, Abu Dhabi, United Arab Emirates'
        order.seller_reference = (
            'Mr. Synthetic Contact\ninfo@syntheticsurveys.com\n'
            'supplier@syntheticsurveys.com'
        )
        order.quote_ref = 'CESPR092026016 R0 & Email Dated 14.09.2026'
        order.seller_license_no = 'CN-4349991'
        order.invoicing_attn = 'Attn. Mr. Synthetic Accounts Contact'
        order.invoicing_emails = [
            'accounts.payable.contact@example.ae',
            'cc. uae.finance@example.ae',
            'uae.procurement@example.ae',
        ]
        order.company_fax = 'INVOICEFAX9988'
        order.contact_persons = {'buyer_references': [
            {'name': 'Synthetic Primary Buyer', 'designation': 'Procurement Manager',
             'email': 'primarybuyer.thomas@example.ae'},
            {'name': 'Synthetic Second Buyer', 'designation': 'Procurement Engineer',
             'email': 'secondarybuyer.ravichandran@example.ae'},
        ]}
        order.payment_terms = '45 days net for agreed payment milestones'
        order.delivery_terms = 'Services completed and accepted'
        order.marking = 'RAD-PRJ-PUR-0126_SEP2026'
        return order

    def test_realistic_cover_keeps_every_field_legible_on_first_page(self):
        order = self._realistic_long_contact_order()
        content, warnings = build_purchase_order_pdf(order)
        self.assertEqual(warnings, [])
        with fitz.open(stream=content, filetype='pdf') as pdf:
            self.assertEqual(len(pdf), 3)
            self._assert_body_clear_of_footer(pdf)
            page = pdf[0]
            text = ''.join(page.get_text().split())
            for label in (
                'Seller:', 'Seller Address:', 'Invoicing Address:', 'Seller Reference:',
                'Buyer Reference:', 'Payment Terms:', 'Payment Mode:', 'Project:',
                'Delivery terms:', 'Delivery date:', 'Marking:', 'Purchase Summary:',
                'Total Purchase Price:', 'Total Sum:', 'Approved by:', 'Order Confirmation:',
                'Seller Signature:', 'Seller Name:', 'Seller Ref. no:', 'Contact Person:',
                'Phone Number:', 'Fax:', 'Email:',
            ):
                self.assertIn(''.join(label.split()), text)
            for value in (
                order.vendor.name, order.seller_address, order.seller_reference,
                order.quote_ref, order.seller_license_no, order.invoicing_attn,
                *order.invoicing_emails, order.payment_terms, order.delivery_terms,
                order.marking, order.title, order.seller_phone, order.seller_fax,
                order.seller_email,
            ):
                self.assertIn(''.join(value.split()), text)
            for reference in order.contact_persons['buyer_references']:
                for value in reference.values():
                    self.assertIn(''.join(value.split()), text)
            # Keep the identity/date roughly 40mm above the previous bottom
            # anchor, grouped with the heading and clear of the footer.
            heading = page.search_for('Approved by:')[0]
            approver = page.search_for('Jarmo Suominen')[0]
            date_line = page.search_for('2026-09-03')[0]
            self.assertAlmostEqual(approver.x0, heading.x0, delta=1)
            self.assertGreater(approver.y0, page.rect.height - 120 * mm)
            self.assertLess(approver.y0, page.rect.height - 95 * mm)
            self.assertGreater(approver.y0, heading.y1)
            self.assertLess(date_line.x1, page.rect.width / 2)
            self.assertGreater(date_line.y0, approver.y1)
            self.assertGreater(date_line.y1, page.rect.height - 88 * mm)
            self.assertLess(date_line.y1, page.rect.height - 82 * mm)
            self.assertLess(
                page.search_for('INVOICEFAX9988')[0].y1,
                min(rect.y0 for rect in page.search_for('Payment Terms:')),
            )
            # Evaluate the actual exported glyph sizes after any fit scaling.
            for block in page.get_text('dict')['blocks']:
                for line in block.get('lines', []):
                    for span in line['spans']:
                        if (span['bbox'][1] >= 34 * mm
                                and span['bbox'][3] <= page.rect.height - 40 * mm):
                            self.assertGreaterEqual(span['size'], 8,
                                f'Cover text became too small: {span["text"]}')
            self.assertIn('PO DESCRIPTION & SCOPE', pdf[1].get_text())
            self.assertNotIn('Order Confirmation:', pdf[1].get_text())
            self.assertIn('SUMMARY OF PRICES', pdf[2].get_text())

    def test_long_addresses_and_contacts_fit_on_first_page_without_footer_overlap(self):
        order = self._long_order()
        content, warnings = build_purchase_order_pdf(order)
        self.assertEqual(warnings, [])
        with fitz.open(stream=content, filetype='pdf') as pdf:
            self._assert_body_clear_of_footer(pdf)
            text = pdf[0].get_text()
            compact_text = ''.join(text.split())
            for value in [
                order.vendor.name, order.seller_address, *order.invoicing_emails,
                order.seller_contact_person, order.seller_email,
                order.approved_by_name, order.approved_by_title,
                order.payment_terms, order.approved_date,
            ]:
                self.assertIn(''.join(value.split()), compact_text)
            for reference in order.contact_persons['buyer_references']:
                self.assertIn(''.join(reference['email'].split()), compact_text)
            fax = [(page.number, rect.y1) for page in pdf for rect in page.search_for('INVOICEFAX9988')]
            payment = [(page.number, rect.y0) for page in pdf for rect in page.search_for('Payment Terms:')]
            self.assertEqual(len(fax), 1)
            self.assertEqual(len({page_number for page_number, _ in payment}), 1)
            self.assertLess(fax[0], min(payment))
            self.assertEqual(fax[0][0], 0)
            self.assertEqual(min(payment)[0], 0)
            self.assertIn('Approved by:', text)
            self.assertIn('Order Confirmation:', text)
            self.assertEqual(len(pdf), 3)

    def test_single_address_taller_than_page_preserves_every_line_on_first_page(self):
        order = self._order()
        address_lines = [f'ADDRESSLINE{index:03d} Synthetic site location' for index in range(125)]
        order.seller_address = '\n'.join(address_lines)
        content, warnings = build_purchase_order_pdf(order)
        self.assertEqual(warnings, [])
        with fitz.open(stream=content, filetype='pdf') as pdf:
            self._assert_body_clear_of_footer(pdf)
            text = pdf[0].get_text()
            for index in range(125):
                self.assertEqual(text.count(f'ADDRESSLINE{index:03d}'), 1)
            for label in ('Payment Terms:', 'Order Confirmation:', 'Phone Number:'):
                self.assertIn(''.join(label.split()), ''.join(text.split()))
            self.assertIn('105.00USD', ''.join(text.split()))
            self.assertEqual(len(pdf), 3)
            self.assertIn('PO DESCRIPTION & SCOPE', pdf[1].get_text())
            self.assertIn('SUMMARY OF PRICES', pdf[2].get_text())

    def test_signature_approver_and_confirmation_stay_together_on_first_page(self):
        order = self._realistic_long_contact_order()
        order.approved_by_name = 'Synthetic Authorised Approver'
        order.approved_date = '2026-09-12'
        signature = BytesIO()
        PILImage.new('RGB', (180, 50), 'navy').save(signature, format='PNG')
        order.approval_signature = 'data:image/png;base64,' + base64.b64encode(signature.getvalue()).decode()
        content, warnings = build_purchase_order_pdf(order)
        self.assertEqual(warnings, [])
        with fitz.open(stream=content, filetype='pdf') as pdf:
            self._assert_body_clear_of_footer(pdf)
            approval_pages = [page for page in pdf if page.search_for('Approved by:')]
            self.assertEqual(len(approval_pages), 1)
            page = approval_pages[0]
            self.assertEqual(page.number, 0)
            self.assertEqual(len(pdf), 3)
            heading = page.search_for('Approved by:')[0]
            approver = page.search_for(order.approved_by_name)[0]
            self.assertTrue(page.search_for('Order Confirmation:'))
            signature_images = [image for image in page.get_images() if image[2:4] == (180, 50)]
            self.assertEqual(len(signature_images), 1)
            signature_rect = page.get_image_rects(signature_images[0][0])[0]
            self.assertLess(heading.y1, signature_rect.y0)
            self.assertLess(signature_rect.y1, approver.y0)
            self.assertLess(approver.y0 - signature_rect.y1, 5 * mm)
            self.assertAlmostEqual(signature_rect.x0, approver.x0, delta=1)
            approval_date = page.search_for(order.approved_date)[0]
            self.assertGreater(approver.y0, page.rect.height - 120 * mm)
            self.assertLess(approver.y0, page.rect.height - 95 * mm)
            self.assertGreater(approval_date.y0, approver.y1)
            self.assertLess(approval_date.x1, page.rect.width / 2)
            self.assertGreater(approval_date.y1, page.rect.height - 88 * mm)
            # The seal now sits beside the signature/name. Remaining identity
            # lines flow beneath it, still well clear of the printed footer.
            self.assertLess(approval_date.y1, page.rect.height - 75 * mm)
            self.assertFalse(page.search_for('__________________________'))

    def test_pending_approval_identity_is_higher_and_word_cell_stays_top_aligned(self):
        order = self._realistic_long_contact_order()
        order.status, order.approved_at, order.approved_date = 'draft', None, None
        order.approval_log = [{'user_id': 'po-reviewer', 'stage': 'Final Management Sign-off', 'status': 'Pending'}]
        content, warnings = build_purchase_order_pdf(order)
        self.assertFalse(warnings)
        with fitz.open(stream=content, filetype='pdf') as pdf:
            self.assertEqual(len(pdf), 3)
            self._assert_body_clear_of_footer(pdf)
            page = pdf[0]
            heading = page.search_for('Approval pending:')[0]
            identity = page.search_for('Jarmo Suominen')[0]
            self.assertGreater(identity.y0, heading.y1)
            self.assertAlmostEqual(identity.x0, heading.x0, delta=1)
            self.assertGreater(identity.y0, page.rect.height - 115 * mm)
            self.assertLess(identity.y0, page.rect.height - 95 * mm)
            self.assertNotIn('Approved by:', page.get_text())
            self.assertTrue(page.search_for('Order Confirmation:'))
        word = Document(BytesIO(build_purchase_order_docx(order)))
        cell = next(cell for table in word.tables for row in table.rows for cell in row.cells if 'Approval pending:' in cell.text)
        self.assertEqual(cell.vertical_alignment, 0)  # Word TOP alignment.
        self.assertIn('Jarmo Suominen', cell.text)
        self.assertIn('CEO, Rejlers Abu Dhabi', cell.text)

    def test_pending_cover_shows_all_projects_and_blank_contact_without_signing_artwork(self):
        order = self._order()
        order.status, order.approved_at, order.approved_date = 'draft', None, None
        order.approval_log = [{'user_id': 'po-reviewer', 'stage': 'Final Management Sign-off', 'status': 'Pending'}]
        order.project_number = '590001, 590002, 590003'
        order.pr_reference = SimpleNamespace(project_details=[
            {'project_number': '590001'}, {'project_number': '590002'}, {'project_number': '590003'},
        ])
        image = BytesIO()
        PILImage.new('RGB', (181, 51), 'navy').save(image, format='PNG')
        order.approval_signature = 'data:image/png;base64,' + base64.b64encode(image.getvalue()).decode()
        with patch('apps.procurement.services.purchase_order_exports.completed_jarmo_profile_artwork') as artwork:
            content, warnings = build_purchase_order_pdf(order)
            word_content = build_purchase_order_docx(order)
        artwork.assert_not_called()
        self.assertFalse(warnings)
        with fitz.open(stream=content, filetype='pdf') as pdf:
            cover = pdf[0].get_text()
            self.assertIn('590001, 590002, 590003', cover)
            self.assertIn('Jarmo Suominen', cover)
            self.assertIn('CEO, Rejlers Abu Dhabi', cover)
            self.assertIn('Approval pending:', cover)
            self.assertNotIn('Approved by:', cover)
            self.assertEqual(cover.count('Vendor Contact'), 1)  # Seller Reference only.
            self.assertRegex(cover, r'Contact Person:\s+Phone Number:')
            self.assertFalse(any(entry[2:4] == (181, 51) for page in pdf for entry in page.get_images()))
        word = Document(BytesIO(word_content))
        text = '\n'.join(word.element.xpath('//w:t/text()'))
        self.assertIn('590001, 590002, 590003', text)
        self.assertIn('Jarmo Suominen', text)
        self.assertIn('CEO, Rejlers Abu Dhabi', text)
        self.assertNotIn('Approved by:', text)
        self.assertRegex(text, r'Contact Person:\s+Phone Number:')
        self.assertNotIn(image.getvalue(), [part.blob for part in word.part.package.parts])
        self.assertIsNone(order.approved_at)
        self.assertIsNone(order.approved_date)

    def test_project_expansion_keeps_explicit_overrides_and_approved_recorded_terms(self):
        from apps.procurement.serializers import PurchaseOrderSerializer
        from apps.procurement.services.purchase_order_project_display import purchase_order_project_reference

        order = self._order()
        order.status, order.approved_at = 'draft', None
        order.project = SimpleNamespace(project_number='590001', project_name='First project')
        order.pr_reference = SimpleNamespace(project_details=[
            {'project_number': ' 590001 '}, {'project_code': '590002'}, {'code': '590003'},
            {'project_number': '590001'}, {'project_name': 'A description is not a project code'},
        ])
        self.assertEqual(purchase_order_project_reference(order), '590001, 590002, 590003')
        self.assertEqual(PurchaseOrderSerializer().get_project_display(order), '590001, 590002, 590003')
        order.project_number = 'OVERRIDE-42'
        self.assertEqual(purchase_order_project_reference(order), 'OVERRIDE-42')
        self.assertEqual(PurchaseOrderSerializer().get_project_display(order), 'OVERRIDE-42')
        order.project_number = '590001'
        order.approval_log = [{'status': 'Approved'}]
        self.assertEqual(purchase_order_project_reference(order), '590001')
        content, _ = build_purchase_order_pdf(order)
        with fitz.open(stream=content, filetype='pdf') as pdf:
            self.assertNotIn('590002', pdf[0].get_text())
        self.assertEqual(order.project_number, '590001')

    def test_unlocked_historical_project_labels_render_all_three_numbers(self):
        order = self._order()
        order.status, order.approved_at = 'draft', None
        order.project_number = '5901142'
        order.pr_reference = SimpleNamespace(project_details=[
            {'source': 'historical', 'value': '5901142-SARB PRODUCED WATER TREATMENT PROJECT'},
            {'source': 'historical', 'label': 'C & F CED.FWA T31 Plant Modifications (MOCs) for Upper Zakum Package 5901086'},
            {'source': 'historical', 'value': 'Detailed Engineering for (NEB) (10522 & 10523), 5901056'},
        ])
        content, _ = build_purchase_order_pdf(order)
        with fitz.open(stream=content, filetype='pdf') as pdf:
            self.assertIn('5901142, 5901086, 5901056', pdf[0].get_text())
        word = Document(BytesIO(build_purchase_order_docx(order)))
        text = '\n'.join(cell.text for table in word.tables for row in table.rows for cell in row.cells)
        self.assertIn('5901142, 5901086, 5901056', text)
        self.assertEqual(order.project_number, '5901142')

    def test_explicit_single_or_empty_project_selection_overrides_linked_pr_and_stale_primary(self):
        from apps.procurement.serializers import PurchaseOrderSerializer
        from apps.procurement.services.purchase_order_project_display import purchase_order_project_reference

        order = self._order()
        order.status, order.approved_at = 'draft', None
        order.project = SimpleNamespace(project_number='5901142', project_name='Old primary project')
        order.rad_project_no = '5901142'
        order.pr_reference = SimpleNamespace(project_details=[
            {'project_number': '5901142'}, {'project_number': '5901086'}, {'project_number': '5901056'},
        ])
        for number, selections in (
            ('5901086', [{'project_number': '5901086', 'project_name': 'Selected project'}]),
            ('', []),
        ):
            with self.subTest(number=number):
                order.project_number = number
                order.contact_persons = {'project_selections': selections}
                self.assertEqual(purchase_order_project_reference(order), number)
                self.assertEqual(PurchaseOrderSerializer().get_project_display(order), number or None)
                content, _ = build_purchase_order_pdf(order)
                with fitz.open(stream=content, filetype='pdf') as pdf:
                    cover = pdf[0].get_text()
                word = Document(BytesIO(build_purchase_order_docx(order)))
                word_text = '\n'.join(cell.text for table in word.tables for row in table.rows for cell in row.cells)
                for text in (cover, word_text):
                    self.assertNotIn('5901142', text)
                    self.assertNotIn('5901056', text)
                    if number:
                        self.assertIn(number, text)
                    else:
                        self.assertNotIn('5901086', text)
                        self.assertNotIn('Multiple Projects', text)
                        self.assertRegex(text, r'Project:?\s+—')
