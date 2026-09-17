"""Canonical PO exports show real approval artwork without inventing signatures."""

import base64
from io import BytesIO
from pathlib import Path
from unittest import TestCase
from unittest.mock import patch

import fitz
from docx import Document
from docx.oxml.ns import qn
from PIL import Image as PILImage
from PyPDF2 import PdfReader
from PyPDF2.generic import ContentStream
from reportlab.lib.units import mm

from apps.procurement.services.purchase_order_approval_artwork import DEFAULT_APPROVAL_STAMP_REFERENCE
from apps.procurement.services.purchase_order_exports import build_purchase_order_docx, build_purchase_order_pdf
from . import test_purchase_order_exports as export_fixtures


STAMP_PATH = Path(__file__).resolve().parents[1] / 'assets' / 'commercial-license-stamp.png'
SIGNATURE_SIZE = (181, 53)


class PurchaseOrderStampTests(TestCase):
    _order = export_fixtures.PurchaseOrderExportTests._order

    def setUp(self):
        self.stamp_bytes = STAMP_PATH.read_bytes()
        with PILImage.open(BytesIO(self.stamp_bytes)) as stamp:
            self.stamp_size = stamp.size
        stream = BytesIO()
        PILImage.new('RGB', SIGNATURE_SIZE, (35, 20, 145)).save(stream, format='PNG')
        self.signature_bytes = stream.getvalue()

    def recorded_order(self, *, dense=True):
        order = (export_fixtures.PurchaseOrderExportTests._realistic_long_contact_order(self)
                 if dense else self._order())
        order.status = 'completed'
        order.approved_by_id = 'recorded-ceo'
        order.approved_by_name = 'Synthetic Recorded CEO'
        order.approved_by_title = 'Chief Executive Officer'
        order.approved_date = '2026-09-12'
        order.approval_signature = 'data:image/png;base64,' + base64.b64encode(self.signature_bytes).decode()
        order.approval_stamp = DEFAULT_APPROVAL_STAMP_REFERENCE
        order.approval_log = [{
            'level': 5, 'stage': 'Final Management Sign-off',
            'user_id': order.approved_by_id, 'approved_by_id': order.approved_by_id,
            'signature_user_id': order.approved_by_id,
            'status': 'Approved', 'date': '2026-09-12T10:00:00Z',
            'signature': order.approval_signature,
        }]
        return order

    @staticmethod
    def image_parts(document):
        result = {}
        for part in document.part.package.parts:
            if not part.content_type.startswith('image/'):
                continue
            with PILImage.open(BytesIO(part.blob)) as image:
                result.setdefault(image.size, []).append(part)
        return result

    def assert_artwork_absent(self, order):
        content, _warnings = build_purchase_order_pdf(order)
        with fitz.open(stream=content, filetype='pdf') as pdf:
            dimensions = [image[2:4] for page in pdf for image in page.get_images()]
            self.assertNotIn(SIGNATURE_SIZE, dimensions)
            self.assertNotIn(self.stamp_size, dimensions)
        document = Document(BytesIO(build_purchase_order_docx(order)))
        parts = self.image_parts(document)
        self.assertNotIn(SIGNATURE_SIZE, parts)
        self.assertNotIn(self.stamp_size, parts)

    def test_completed_verified_po_pdf_has_visible_full_opacity_signature_and_stamp_on_cover(self):
        order = self.recorded_order()
        content, warnings = build_purchase_order_pdf(order)
        self.assertFalse(warnings)
        with fitz.open(stream=content, filetype='pdf') as pdf:
            self.assertEqual(len(pdf), 3)
            export_fixtures.PurchaseOrderExportTests._assert_body_clear_of_footer(self, pdf)
            page = pdf[0]
            images = page.get_images()
            signature = [item for item in images if item[2:4] == SIGNATURE_SIZE]
            stamps = [item for item in images if item[2:4] == self.stamp_size]
            self.assertEqual(len(signature), 1)
            self.assertEqual(len(stamps), 1)
            signature_rect = page.get_image_rects(signature[0][0])[0]
            stamp_rect = page.get_image_rects(stamps[0][0])[0]
            self.assertLessEqual(signature_rect.x1, stamp_rect.x0 + 1)
            self.assertLessEqual(signature_rect.width, 42 * mm + 1)
            self.assertLessEqual(signature_rect.height, 18 * mm + 1)
            self.assertLessEqual(stamp_rect.width, 30 * mm + 1)
            self.assertGreater(stamp_rect.width, 25 * mm)
            name_rect = page.search_for(order.approved_by_name)[0]
            # The signature sits immediately above the name, and the seal is
            # alongside it (not stranded under the "Approved by" heading).
            self.assertGreater(name_rect.y0, signature_rect.y1)
            self.assertLess(name_rect.y0 - signature_rect.y1, 5 * mm)
            self.assertLess(abs((signature_rect.y0 + signature_rect.y1) / 2
                                - (stamp_rect.y0 + stamp_rect.y1) / 2), mm)
            self.assertGreater(stamp_rect.y1, name_rect.y0)
            self.assertFalse(stamp_rect.intersects(name_rect))
            self.assertLess(stamp_rect.y1, page.search_for(order.approved_by_title)[0].y0)
            for rectangle in (signature_rect, stamp_rect):
                self.assertGreater(rectangle.y0, page.search_for('Approved by:')[0].y1)
                self.assertLess(rectangle.y1, page.rect.height - 40 * mm)
            # Inspect rendered pixels too: an embedded but transparent/covered
            # stamp is not a visible stamp.
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), clip=stamp_rect, colorspace=fitz.csRGB, alpha=False)
            rendered = PILImage.frombytes('RGB', (pixmap.width, pixmap.height), pixmap.samples)
            blue_ink = sum(1 for red, green, blue in rendered.getdata()
                           if blue > red + 20 and blue > green + 10 and min(red, green) < 190)
            self.assertGreater(blue_ink, 100)
            artwork_names = {'/' + signature[0][7], '/' + stamps[0][7]}

        reader = PdfReader(BytesIO(content))
        page = reader.pages[0]
        resources = page['/Resources'].get_object()
        graphics_states = resources.get('/ExtGState', {})
        graphics_states = graphics_states.get_object() if hasattr(graphics_states, 'get_object') else graphics_states
        alpha, stack, painted = 1.0, [], []
        for operands, operator in ContentStream(page.get_contents(), reader).operations:
            if operator == b'q':
                stack.append(alpha)
            elif operator == b'Q':
                alpha = stack.pop()
            elif operator == b'gs':
                state = graphics_states[operands[0]].get_object()
                alpha = float(state.get('/ca', alpha))
            elif operator == b'Do' and str(operands[0]) in artwork_names:
                painted.append(alpha)
        self.assertEqual(len(painted), 2)
        self.assertEqual(painted, [1.0, 1.0])

    def test_compact_cover_keeps_signature_and_seal_at_signer_name(self):
        order = self.recorded_order(dense=False)
        content, _warnings = build_purchase_order_pdf(order)
        with fitz.open(stream=content, filetype='pdf') as pdf:
            page = pdf[0]
            name = page.search_for(order.approved_by_name)[0]
            signature = next(image for image in page.get_images() if image[2:4] == SIGNATURE_SIZE)
            stamp = next(image for image in page.get_images() if image[2:4] == self.stamp_size)
            signature_rect = page.get_image_rects(signature[0])[0]
            stamp_rect = page.get_image_rects(stamp[0])[0]
            self.assertGreater(name.y0, signature_rect.y1)
            self.assertLess(name.y0 - signature_rect.y1, 5 * mm)
            self.assertLess(abs((signature_rect.y0 + signature_rect.y1) / 2
                                - (stamp_rect.y0 + stamp_rect.y1) / 2), mm)
            self.assertFalse(stamp_rect.intersects(name))
            self.assertLess(stamp_rect.y1, page.search_for(order.approved_by_title)[0].y0)
            self.assertGreater(name.y0, page.rect.height - 120 * mm)
            self.assertLess(name.y0, page.rect.height - 95 * mm)
            export_fixtures.PurchaseOrderExportTests._assert_body_clear_of_footer(self, pdf)

    def test_word_keeps_stamp_and_recorded_signature_as_images_on_same_approval_row(self):
        order = self.recorded_order()
        document = Document(BytesIO(build_purchase_order_docx(order)))
        parts = self.image_parts(document)
        self.assertEqual(len(parts.get(self.stamp_size, [])), 1)
        self.assertEqual(len(parts.get(SIGNATURE_SIZE, [])), 1)
        self.assertEqual(parts[self.stamp_size][0].blob, self.stamp_bytes)
        self.assertEqual(parts[SIGNATURE_SIZE][0].blob, self.signature_bytes)
        rows = []
        for shape in document.inline_shapes:
            blips = shape._inline.xpath('.//a:blip')
            if not blips:
                continue
            part = document.part.related_parts[blips[0].get(qn('r:embed'))]
            if part not in parts[self.stamp_size] + parts[SIGNATURE_SIZE]:
                continue
            row = next(parent for parent in shape._inline.iterancestors() if parent.tag == qn('w:tr'))
            rows.append(row)
            for opacity in shape._inline.xpath('.//a:alpha'):
                self.assertEqual(opacity.get('val'), '100000')
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0], rows[1])
        self.assertIn(order.approved_by_name, ''.join(document.element.itertext()))

    def test_completed_status_without_signature_never_gets_a_stamp_or_default_ceo_signature(self):
        order = self.recorded_order()
        order.approval_signature = ''
        order.approval_log[0]['signature'] = ''
        self.assert_artwork_absent(order)

    def test_legacy_completed_valid_signature_uses_default_seal_without_rewriting_record(self):
        order = self.recorded_order()
        order.approval_stamp = ''
        content, _warnings = build_purchase_order_pdf(order)
        with fitz.open(stream=content, filetype='pdf') as pdf:
            self.assertIn(self.stamp_size, [image[2:4] for image in pdf[0].get_images()])
        document = Document(BytesIO(build_purchase_order_docx(order)))
        self.assertIn(self.stamp_size, self.image_parts(document))
        self.assertEqual(order.approval_stamp, '')

    def test_saved_custom_stamp_takes_precedence_over_the_default_company_seal(self):
        order = self.recorded_order()
        custom_size = (127, 113)
        stream = BytesIO()
        PILImage.new('RGB', custom_size, (145, 30, 40)).save(stream, format='PNG')
        custom_bytes = stream.getvalue()
        order.approval_stamp = 'data:image/png;base64,' + base64.b64encode(custom_bytes).decode()
        content, _warnings = build_purchase_order_pdf(order)
        with fitz.open(stream=content, filetype='pdf') as pdf:
            dimensions = [image[2:4] for page in pdf for image in page.get_images()]
            self.assertIn(custom_size, dimensions)
            self.assertNotIn(self.stamp_size, dimensions)
        document = Document(BytesIO(build_purchase_order_docx(order)))
        parts = self.image_parts(document)
        self.assertEqual(parts[custom_size][0].blob, custom_bytes)
        self.assertNotIn(self.stamp_size, parts)

    def test_mismatched_recorded_signer_hides_both_signature_and_stamp(self):
        order = self.recorded_order()
        order.approval_log[0]['approved_by_id'] = 'different-person'
        self.assert_artwork_absent(order)

    def test_original_pdf_evidence_url_does_not_become_a_synthetic_signature_or_stamp(self):
        order = self.recorded_order()
        order.approval_signature = 'https://files.example.invalid/original-signed-po.pdf#page=1'
        order.approval_stamp = order.approval_signature
        order.approval_log = [{
            'status': 'Approved', 'external': True, 'evidence_document_id': 'original-po',
            'signature_verified': True, 'stamp_verified': True,
        }]
        with patch('urllib.request.urlopen') as external_fetch:
            self.assert_artwork_absent(order)
        external_fetch.assert_not_called()

    def test_recorded_final_employee_is_not_replaced_by_a_default_ceo(self):
        order = self.recorded_order()
        order.approved_by_name = 'Synthetic Procurement Approver'
        order.approved_by_title = 'Procurement Manager'
        content, _warnings = build_purchase_order_pdf(order)
        with fitz.open(stream=content, filetype='pdf') as pdf:
            text = '\n'.join(page.get_text() for page in pdf)
            self.assertIn(''.join(order.approved_by_name.split()), ''.join(text.split()))
        self.assertNotIn('Jarmo Suominen', text)
        document = Document(BytesIO(build_purchase_order_docx(order)))
        text = ''.join(document.element.itertext())
        self.assertIn(order.approved_by_name, text)
        self.assertNotIn('Jarmo Suominen', text)
