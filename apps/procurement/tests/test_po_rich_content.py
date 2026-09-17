import base64
import gc
from io import BytesIO
from unittest import TestCase
from unittest.mock import patch

import fitz
from docx import Document
from docx.oxml.ns import qn
from PIL import Image
from PyPDF2 import PdfReader
from reportlab.pdfgen import canvas

from apps.procurement.services.po_rich_content import parse_rich_content
from apps.procurement.services.purchase_order_exports import build_purchase_order_docx, build_purchase_order_pdf
from apps.procurement.tests import test_purchase_order_exports as export_tests


class PurchaseOrderRichContentTests(TestCase):
    def order(self, narrative):
        order = export_tests.PurchaseOrderExportTests()._order()
        order.description = narrative
        return order

    def render(self, narrative):
        order = self.order(narrative)
        content, warnings = build_purchase_order_pdf(order)
        self.assertEqual(warnings, [])
        return fitz.open(stream=content, filetype='pdf'), Document(BytesIO(build_purchase_order_docx(order)))

    def test_paragraphs_keep_inline_styles_alignment_and_explicit_spacing(self):
        pdf, word = self.render('<h2 style="text-align:center">Scope Heading</h2><p style="text-align:justify;line-height:1.8;margin-bottom:14pt">Normal <strong>Bold Fee</strong> <em>Italic</em> <u>Underlined</u> <s>Removed</s> <span style="font-family:Times New Roman;font-size:18pt;color:#cc1122;background-color:#ffee33">4,386 USD</span></p>')
        spans = [span for page in pdf for block in page.get_text('dict')['blocks'] if 'lines' in block for line in block['lines'] for span in line['spans']]
        price = next(span for span in spans if span['text'] == '4,386 USD')
        self.assertAlmostEqual(price['size'], 18)
        self.assertEqual(price['color'], 0xcc1122)
        self.assertIn('Times', price['font'])
        self.assertTrue(any('Bold' in span['font'] for span in spans if 'Bold Fee' in span['text']))
        paragraph = next(paragraph for paragraph in word.paragraphs if 'Bold Fee' in paragraph.text)
        self.assertEqual(paragraph.alignment, 3)
        self.assertAlmostEqual(paragraph.paragraph_format.line_spacing, 1.8)
        self.assertAlmostEqual(paragraph.paragraph_format.space_after.pt, 14)
        price_run = next(run for run in paragraph.runs if run.text == '4,386 USD')
        self.assertEqual(price_run.font.name, 'Times New Roman')
        self.assertEqual(price_run.font.size.pt, 18)
        self.assertIn('ffee33', price_run._r.xml)
        self.assertTrue(next(run for run in paragraph.runs if run.text == 'Removed').font.strike)
        self.assertTrue(next(run for run in paragraph.runs if run.text == 'Underlined').underline)
        self.assertEqual(next(paragraph for paragraph in word.paragraphs if paragraph.text == 'Scope Heading').style.name, 'Heading 2')

    def test_inline_newlines_do_not_create_new_paragraphs_or_drop_literal_entities(self):
        blocks = parse_rich_content('<p>First\n wrapped line &amp; value &lt;script&gt;literal&lt;/script&gt;<br>Next line</p>')
        self.assertEqual(len(blocks), 1)
        self.assertEqual(''.join(run.text for run in blocks[0].runs), 'First wrapped line & value <script>literal</script>\nNext line')
        self.assertEqual(len(parse_rich_content('&lt;p&gt;Escaped&lt;/p&gt;')), 1)

    def test_legacy_plain_text_newlines_and_malformed_links_remain_readable(self):
        pdf, word = self.render('First legacy paragraph\nSecond legacy paragraph')
        self.assertTrue(any(p.text == 'First legacy paragraph' for p in word.paragraphs))
        self.assertTrue(any(p.text == 'Second legacy paragraph' for p in word.paragraphs))
        text = '\n'.join(page.get_text() for page in pdf)
        self.assertIn('First legacy paragraph\nSecond legacy paragraph', text)
        self.assertTrue(parse_rich_content('<p><a href="http://[">Visible malformed link</a></p>'))

    def test_nested_lists_start_override_and_editable_numbering(self):
        narrative = '<ol start="3"><li>Third clause<ul><li>Nested bullet</li></ul></li><li value="7"><p>Seventh clause</p></li><li>Eighth clause</li></ol>'
        blocks = parse_rich_content(narrative)
        self.assertEqual([block.list_info['marker'] for block in blocks], ['3.', '•', '7.', '8.'])
        self.assertEqual([block.list_info['depth'] for block in blocks], [0, 1, 0, 0])
        pdf, word = self.render(narrative)
        text = '\n'.join(page.get_text() for page in pdf)
        for expected in ('3.', '7.', '8.', 'Nested bullet'):
            self.assertIn(expected, text)
        paragraphs = [p for p in word.paragraphs if p.text in {'Third clause', 'Nested bullet', 'Seventh clause', 'Eighth clause'}]
        self.assertEqual(len(paragraphs), 4)
        self.assertTrue(all(p._p.pPr.numPr is not None for p in paragraphs))
        self.assertGreater(paragraphs[1].paragraph_format.left_indent, paragraphs[0].paragraph_format.left_indent)
        self.assertIn('w:val="7"', word.part.numbering_part.element.xml)

    def test_six_column_table_stays_a_table_and_merged_cells_remain_editable(self):
        narrative = '<table><thead><tr><th>SN</th><th>Discipline</th><th>Period</th><th>Manhours</th><th>Unit Price USD</th><th>Total Price USD</th></tr></thead><tbody><tr><td rowspan="2">1</td><td>Engineering</td><td>Sept</td><td>40</td><td>109.65</td><td>4,386</td></tr><tr><td colspan="5" style="background-color:#eeeeee"><strong>Fee includes all deliverables</strong></td></tr></tbody></table>'
        pdf, word = self.render(narrative)
        text = '\n'.join(page.get_text() for page in pdf)
        self.assertIn('Engineering', text)
        self.assertIn('4,386', text)
        table = next(table for table in word.tables if table.cell(0, 0).text == 'SN')
        self.assertEqual(len(table.columns), 6)
        self.assertEqual(table.cell(1, 0)._tc, table.cell(2, 0)._tc)
        self.assertEqual(table.cell(2, 1)._tc, table.cell(2, 5)._tc)
        self.assertIn('w:tblHeader', table.rows[0]._tr.xml)
        self.assertIn('w:gridSpan', table.rows[2]._tr.xml)

    def test_long_table_splits_across_pages_with_repeated_header(self):
        narrative = '<table><thead><tr><th>Unique Narrative Column</th><th>Quantity</th></tr></thead><tbody>' + ''.join(f'<tr><td>Deliverable {index}</td><td>{index}</td></tr>' for index in range(75)) + '</tbody></table>'
        pdf, word = self.render(narrative)
        pages = [page.get_text() for page in pdf]
        self.assertGreater(sum('Unique Narrative Column' in page for page in pages), 1)
        self.assertIn('Deliverable 74', '\n'.join(pages))
        table = next(table for table in word.tables if 'Unique Narrative Column' in table.cell(0, 0).text)
        self.assertEqual(len(table.rows), 76)

    def test_table_column_widths_and_all_header_table_remain_renderable(self):
        narrative = '<table style="width:100%"><colgroup><col style="width:25%"><col style="width:75%"></colgroup><tr><th>Short code</th><th>Detailed scope text</th></tr><tr><th>A</th><th>Engineering specification</th></tr></table>'
        pdf, word = self.render(narrative)
        self.assertIn('Engineering specification', '\n'.join(page.get_text() for page in pdf))
        table = next(table for table in word.tables if table.cell(0, 0).text == 'Short code')
        self.assertAlmostEqual(table.columns[1].width / table.columns[0].width, 3, places=2)

    def test_long_table_of_header_cells_can_continue_without_an_oversized_repeating_header(self):
        narrative = '<table>' + ''.join(f'<tr><th>Header-cell deliverable {index}</th><th>{index}</th></tr>' for index in range(75)) + '</table>'
        order = self.order(narrative)
        content, _ = build_purchase_order_pdf(order)
        with fitz.open(stream=content, filetype='pdf') as document:
            text = '\n'.join(page.get_text() for page in document)
            self.assertIn('Header-cell deliverable 74', text)
            self.assertGreater(len(document), 3)
        word = Document(BytesIO(build_purchase_order_docx(order)))
        table = next(table for table in word.tables if table.cell(0, 0).text == 'Header-cell deliverable 0')
        self.assertNotIn('w:tblHeader', table._tbl.xml)

    def test_manual_page_break_is_real_and_editor_label_is_not_exported(self):
        pdf, word = self.render('<p>Before manual boundary</p><div data-po-page-break="true">Page Break</div><p>After manual boundary</p>')
        pages = [page.get_text() for page in pdf]
        before = next(i for i, text in enumerate(pages) if 'Before manual boundary' in text)
        after = next(i for i, text in enumerate(pages) if 'After manual boundary' in text)
        self.assertEqual(after, before + 1)
        self.assertNotIn('Page Break', '\n'.join(pages))
        self.assertNotIn('Page Break', '\n'.join(p.text for p in word.paragraphs))
        paragraphs = word.paragraphs
        index = next(i for i, p in enumerate(paragraphs) if p.text == 'Before manual boundary')
        self.assertIn('w:type="page"', paragraphs[index + 1]._p.xml)

    def test_inline_images_and_safe_links_render_without_remote_requests(self):
        image = BytesIO()
        Image.new('RGB', (130, 70), 'navy').save(image, 'PNG')
        data = base64.b64encode(image.getvalue()).decode()
        narrative = f'<p><a href="https://example.com/quote"><strong>Quotation link</strong></a></p><img src="data:image/png;base64,{data}" style="width:120pt"><img src="http://169.254.169.254/private" alt="External image"><script>dangerous text</script><p><a href="javascript:alert(1)">Unsafe link text</a></p>'
        with patch('urllib.request.urlopen', side_effect=AssertionError('No network access')):
            pdf, word = self.render(narrative)
        self.assertTrue(any(image[2:4] == (130, 70) for page in pdf for image in page.get_images()))
        self.assertTrue(any(link.get('uri') == 'https://example.com/quote' for page in pdf for link in page.get_links()))
        self.assertIn('https://example.com/quote', '\n'.join(str(rel.target_ref) for rel in word.part.rels.values()))
        self.assertNotIn('javascript:', word.element.xml)
        text = '\n'.join(page.get_text() for page in pdf)
        self.assertIn('External image', text)
        self.assertNotIn('dangerous text', text)

    def test_price_columns_respect_saved_order_labels_and_specification(self):
        order = self.order('<p>Scope</p>')
        order.items[0].update(specification='SPEC-AX7', custom_measure='Measured value')
        order.items_table_headers = {'__column_order': ['specification', 'custom_measure', 'description'], 'description': 'Scope Item', 'specification': 'Agreed Specification', 'custom_measure': 'Measure'}
        pdf, _ = build_purchase_order_pdf(order)
        with fitz.open(stream=pdf, filetype='pdf') as document:
            text = document[-1].get_text()
            self.assertLess(text.index('Agreed'), text.index('Measure'))
            self.assertIn('SPEC-AX7', text)
            self.assertIn('Measured value', text)
            self.assertNotIn('Discount', text)
        word = Document(BytesIO(build_purchase_order_docx(order)))
        table = word.tables[-1]
        self.assertEqual([cell.text for cell in table.rows[0].cells], ['Agreed Specification', 'Measure', 'Scope Item'])
        self.assertEqual(table.cell(1, 0).text, 'SPEC-AX7')

    def test_persisted_purchase_summary_matches_preview_and_explicit_blank_uses_title(self):
        order = self.order('<p>Scope</p>')
        order.summary = 'Legacy transient summary'
        order.contact_persons = {'purchase_summary': 'Saved engineering summary'}
        pdf, _ = build_purchase_order_pdf(order)
        with fitz.open(stream=pdf, filetype='pdf') as document:
            self.assertIn('Saved engineering summary', document[0].get_text())
            self.assertNotIn('Legacy transient summary', document[0].get_text())
        word = Document(BytesIO(build_purchase_order_docx(order)))
        self.assertTrue(any('Saved engineering summary' in cell.text for table in word.tables for row in table.rows for cell in row.cells))
        order.contact_persons['purchase_summary'] = ''
        pdf, _ = build_purchase_order_pdf(order)
        with fitz.open(stream=pdf, filetype='pdf') as document:
            self.assertIn(order.title, document[0].get_text())
            self.assertNotIn('Legacy transient summary', document[0].get_text())

    def test_server_supplied_preview_attachment_bytes_use_same_pdf_path_without_storage(self):
        order = self.order('<p>Scope</p>')
        order.attachments = [{'filename': 'pending.pdf', 'title': 'Reviewed appendix title', 'description': 'Supporting scope', '_preview_content': export_tests.PurchaseOrderExportTests()._one_page_pdf()}]
        with patch('django.core.files.storage.default_storage.open', side_effect=AssertionError('No storage read for uploaded bytes')):
            content, warnings = build_purchase_order_pdf(order)
        self.assertEqual(warnings, [])
        with fitz.open(stream=content, filetype='pdf') as document:
            self.assertIn('Supporting document', document[-1].get_text())
            self.assertIn('Reviewed appendix title', document[-2].get_text())
            self.assertIn('Supporting scope', document[-2].get_text())

    def test_distinct_attachment_sources_survive_collection_between_pdf_merges(self):
        order = self.order('<p>Main source</p>')
        order.attachments = []
        for index in range(8):
            source = BytesIO()
            page = canvas.Canvas(source)
            page.drawString(40, 800, f'UNIQUE APPENDIX CONTENT {index}')
            page.save()
            order.attachments.append({'filename': f'source-{index}.pdf', '_preview_content': source.getvalue()})

        def read_source(*args, **kwargs):
            gc.collect()
            return PdfReader(*args, **kwargs)

        with patch('apps.procurement.services.purchase_order_exports.PdfReader', side_effect=read_source):
            content, warnings = build_purchase_order_pdf(order)
        self.assertEqual(warnings, [])
        with fitz.open(stream=content, filetype='pdf') as document:
            self.assertEqual(len(document), 3 + 2 * len(order.attachments))
            for index in range(8):
                self.assertIn(f'UNIQUE APPENDIX CONTENT {index}', document[4 + 2 * index].get_text())
