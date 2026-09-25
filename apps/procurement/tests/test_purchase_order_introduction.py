"""Editable PO opening text survives saves and respects approved terms."""

from copy import deepcopy
from io import BytesIO

import fitz
from docx import Document
from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from rest_framework.exceptions import ValidationError

from apps.procurement.models import PurchaseOrder
from apps.procurement.services.purchase_order_document_preview import document_preview_order
from apps.procurement.services.purchase_order_exports import build_purchase_order_docx, build_purchase_order_pdf
from apps.procurement.services.purchase_order_introduction import purchase_order_introduction
from . import test_po_approved_content as fixtures


@override_settings(ROOT_URLCONF=fixtures.__name__, TEAMS_APPROVAL_WEBHOOK_URL='', WEB_PUSH_VAPID_PRIVATE_KEY='')
class PurchaseOrderIntroductionTests(TestCase):
    setUp = fixtures.PurchaseOrderApprovedContentTests.setUp
    stage = staticmethod(fixtures.PurchaseOrderApprovedContentTests.stage)
    approve = fixtures.PurchaseOrderApprovedContentTests.approve

    def patch_statement(self, value):
        contacts = deepcopy(self.order.contact_persons or {})
        contacts['order_introduction'] = value
        return self.client.patch(self.url, {'contact_persons': contacts}, format='json')

    def test_saved_statement_round_trip_and_reset_preserve_other_contacts(self):
        contacts = {'buyer_references': [{'name': 'Example buyer'}], 'purchase_summary': 'Kept summary'}
        self.order.contact_persons = contacts
        self.order.save()
        response = self.patch_statement('We issue this order under the agreed framework.\nBuyer and Seller confirm the scope.')
        self.assertEqual(response.status_code, 200, response.data)
        self.order.refresh_from_db()
        read = self.client.get(self.url)
        self.assertEqual(read.status_code, 200)
        self.assertEqual(read.data['contact_persons'], self.order.contact_persons)
        self.assertEqual(self.order.contact_persons['purchase_summary'], 'Kept summary')
        self.assertEqual(self.order.contact_persons['buyer_references'], contacts['buyer_references'])
        self.assertIn('agreed framework', purchase_order_introduction(self.order))
        reset = self.client.patch(self.url, {'contact_persons': contacts}, format='json')
        self.assertEqual(reset.status_code, 200, reset.data)
        self.order.refresh_from_db()
        self.assertNotIn('order_introduction', self.order.contact_persons)
        self.assertEqual(purchase_order_introduction(self.order),
                         'We, Rejlers International Engineering Solutions (Buyer), issue this purchase order to Content supplier (Seller).')

    def test_invalid_saved_and_preview_statements_are_denied_without_changes(self):
        before = PurchaseOrder.objects.values().get(pk=self.order.pk)
        for value in (None, 4, {}, [], 'x' * 10001, 'Invalid\x00text', 'Invalid\ufffetext'):
            with self.subTest(value_type=type(value).__name__):
                response = self.patch_statement(value)
                self.assertEqual(response.status_code, 400, response.data)
                self.assertIsInstance(response.data['contact_persons'], list)
                with self.assertRaises(ValidationError):
                    document_preview_order({'contact_persons': {'order_introduction': value}}, base=self.order)
                self.assertEqual(PurchaseOrder.objects.values().get(pk=self.order.pk), before)

    def test_approved_statement_is_locked_but_unchanged_value_remains_allowed(self):
        self.order.contact_persons = {'order_introduction': 'Reviewed opening statement.'}
        self.order.save()
        self.approve()
        before = PurchaseOrder.objects.values().get(pk=self.order.pk)
        for value in ('Changed seller obligations.', ''):
            response = self.patch_statement(value)
            self.assertEqual(response.status_code, 400, response.data)
            self.assertIn('revised purchase order', str(response.data))
            self.assertEqual(PurchaseOrder.objects.values().get(pk=self.order.pk), before)
        unchanged = self.patch_statement('Reviewed opening statement.')
        self.assertEqual(unchanged.status_code, 200, unchanged.data)
        self.order.refresh_from_db()
        self.assertEqual(self.order.approval_log, before['approval_log'])
        self.assertEqual(self.order.approval_signature, before['approval_signature'])

    def test_update_permission_is_required(self):
        outsider = get_user_model().objects.create_user('statement-reader', email='reader@example.test')
        self.client.force_authenticate(outsider)
        response = self.patch_statement('Unauthorized replacement')
        self.assertEqual(response.status_code, 403)
        self.order.refresh_from_db()
        self.assertNotIn('order_introduction', self.order.contact_persons)

    def test_pdf_word_and_unsaved_preview_share_literal_multiline_statement(self):
        custom = 'We, Example Buyer (Buyer), order from Example Seller (Seller).\r\nTerms <reviewed> & agreed.'
        self.assertEqual(self.patch_statement(custom).status_code, 200)
        self.order.refresh_from_db()
        before = PurchaseOrder.objects.values().get(pk=self.order.pk)
        pdf, warnings = build_purchase_order_pdf(self.order)
        self.assertFalse(warnings)
        with fitz.open(stream=pdf, filetype='pdf') as document:
            pdf_text = '\n'.join(page.get_text() for page in document)
        word = Document(BytesIO(build_purchase_order_docx(self.order)))
        word_text = '\n'.join(paragraph.text for paragraph in word.paragraphs)
        for text in (pdf_text, word_text):
            self.assertIn('Example Buyer (Buyer)', text)
            self.assertIn('Terms <reviewed> & agreed.', text)
            self.assertNotIn('issue this purchase order to Content supplier', text)
        snapshot = document_preview_order({'contact_persons': {'order_introduction': 'Unsaved scope agreement.'}}, base=self.order)
        self.assertEqual(purchase_order_introduction(snapshot), 'Unsaved scope agreement.')
        self.assertEqual(PurchaseOrder.objects.values().get(pk=self.order.pk), before)

    def test_untouched_legacy_order_keeps_default_without_injection(self):
        from apps.procurement.services.purchase_order_content import purchase_order_content_fingerprint
        fingerprint = purchase_order_content_fingerprint(self.order)
        self.assertNotIn('order_introduction', self.order.contact_persons)
        self.assertIn('issue this purchase order to Content supplier', purchase_order_introduction(self.order))
        self.assertNotIn('order_introduction', self.order.contact_persons)
        self.assertNotIn('show_order_introduction', self.order.contact_persons)
        self.assertEqual(purchase_order_content_fingerprint(self.order), fingerprint)

    def test_introduction_visibility_round_trip_retains_wording_and_other_metadata(self):
        text = 'Preserved Buyer and Seller agreement.'
        original_contacts = {'order_introduction': text, 'purchase_summary': 'Preserved summary',
                             'buyer_references': [{'name': 'Preserved reference'}]}
        self.order.contact_persons = original_contacts
        self.order.description = '<p>Preserved scope narrative.</p>'
        self.order.save()
        for visible in (False, True):
            with self.subTest(visible=visible):
                contacts = {**original_contacts, 'show_order_introduction': visible}
                before = PurchaseOrder.objects.values().get(pk=self.order.pk)
                snapshot = document_preview_order({'contact_persons': contacts}, base=self.order)
                pdf, warnings = build_purchase_order_pdf(snapshot)
                self.assertFalse(warnings)
                with fitz.open(stream=pdf, filetype='pdf') as document:
                    pdf_text = '\n'.join(page.get_text() for page in document)
                self.assertEqual(text in pdf_text, visible)
                self.assertIn('Preserved scope narrative.', pdf_text)
                self.assertEqual(PurchaseOrder.objects.values().get(pk=self.order.pk), before)
                response = self.client.patch(self.url, {'contact_persons': contacts}, format='json')
                self.assertEqual(response.status_code, 200, response.data)
                self.order.refresh_from_db()
                self.assertEqual(self.client.get(self.url).data['contact_persons'], contacts)
                word = Document(BytesIO(build_purchase_order_docx(self.order)))
                word_text = '\n'.join(paragraph.text for paragraph in word.paragraphs)
                self.assertEqual(text in word_text, visible)
                self.assertIn('Preserved scope narrative.', word_text)

    def test_hidden_introduction_does_not_reintroduce_standard_text_when_override_is_reset(self):
        for custom in (None, 'Retained hidden statement.'):
            for visible in (False, True):
                with self.subTest(custom=custom, visible=visible):
                    contacts = {'show_order_introduction': visible}
                    if custom is not None:
                        contacts['order_introduction'] = custom
                    snapshot = document_preview_order({'contact_persons': contacts, 'description': ''}, base=self.order)
                    self.assertEqual(bool(purchase_order_introduction(snapshot)), visible)
                    pdf, warnings = build_purchase_order_pdf(snapshot)
                    self.assertFalse(warnings)
                    with fitz.open(stream=pdf, filetype='pdf') as document:
                        self.assertEqual(len(document), 3 if visible else 2)
                        pdf_text = '\n'.join(page.get_text() for page in document)
                    word = Document(BytesIO(build_purchase_order_docx(snapshot)))
                    word_text = '\n'.join(paragraph.text for paragraph in word.paragraphs)
                    for text in (pdf_text, word_text):
                        self.assertEqual((custom or 'issue this purchase order to') in text, visible)
                    self.assertEqual(snapshot.contact_persons, contacts)

    def test_introduction_visibility_requires_boolean_and_preserves_approved_lock(self):
        before = PurchaseOrder.objects.values().get(pk=self.order.pk)
        for value in ('false', 0, 1, None, [], {}):
            contacts = {**self.order.contact_persons, 'show_order_introduction': value}
            response = self.client.patch(self.url, {'contact_persons': contacts}, format='json')
            self.assertEqual(response.status_code, 400, response.data)
            with self.assertRaises(ValidationError):
                document_preview_order({'contact_persons': contacts}, base=self.order)
            self.assertEqual(PurchaseOrder.objects.values().get(pk=self.order.pk), before)
        self.approve()
        approved = PurchaseOrder.objects.values().get(pk=self.order.pk)
        response = self.client.patch(self.url, {'contact_persons': {
            **self.order.contact_persons, 'show_order_introduction': False,
        }}, format='json')
        self.assertEqual(response.status_code, 400, response.data)
        self.assertIn('revised purchase order', str(response.data))
        self.assertEqual(PurchaseOrder.objects.values().get(pk=self.order.pk), approved)

    def test_cleared_statement_stays_blank_after_save_and_is_omitted_from_pdf_and_word(self):
        self.order.description = '<p>Scope retained after clearing the introduction.</p>'
        self.order.save()
        for blank in ('', '   ', '\r\n\t'):
            with self.subTest(blank=repr(blank)):
                self.assertEqual(self.patch_statement(blank).status_code, 200)
                self.order.refresh_from_db()
                self.assertEqual(self.client.get(self.url).data['contact_persons']['order_introduction'], blank)
                self.assertEqual(purchase_order_introduction(self.order), '')
                before = PurchaseOrder.objects.values().get(pk=self.order.pk)
                pdf, warnings = build_purchase_order_pdf(self.order)
                self.assertFalse(warnings)
                with fitz.open(stream=pdf, filetype='pdf') as document:
                    text = '\n'.join(page.get_text() for page in document)
                    self.assertIn('PO DESCRIPTION & SCOPE', text)
                    self.assertIn('Scope retained after clearing the introduction.', text)
                    self.assertNotIn('issue this purchase order to', text)
                    self.assertNotIn('(Seller).', text)
                word = Document(BytesIO(build_purchase_order_docx(self.order)))
                paragraphs = [paragraph.text for paragraph in word.paragraphs]
                scope_index = paragraphs.index('PO DESCRIPTION & SCOPE')
                self.assertIn('Scope retained after clearing the introduction.', paragraphs)
                # There is no empty introduction paragraph left above scope.
                self.assertTrue(paragraphs[scope_index - 1].startswith('PURCHASE ORDER:'))
                self.assertNotIn('issue this purchase order to', '\n'.join(paragraphs))
                self.assertEqual(PurchaseOrder.objects.values().get(pk=self.order.pk), before)

    def test_cleared_unsaved_preview_omits_statement_without_changing_saved_order(self):
        before = PurchaseOrder.objects.values().get(pk=self.order.pk)
        snapshot = document_preview_order({'contact_persons': {'order_introduction': ''}}, base=self.order)
        self.assertEqual(purchase_order_introduction(snapshot), '')
        pdf, _ = build_purchase_order_pdf(snapshot)
        with fitz.open(stream=pdf, filetype='pdf') as document:
            self.assertNotIn('issue this purchase order to', '\n'.join(page.get_text() for page in document))
        self.assertEqual(PurchaseOrder.objects.values().get(pk=self.order.pk), before)

    def test_scope_heading_toggle_preserves_narrative_in_saved_and_preview_documents(self):
        self.order.description = '<p>Retained engineering scope narrative.</p>'
        self.order.contact_persons = {'order_introduction': 'Agreed opening statement.',
                                      'purchase_summary': 'Retained purchase summary'}
        self.order.save()
        for visible in (None, False, True):
            with self.subTest(visible=visible):
                contacts = deepcopy(self.order.contact_persons)
                if visible is not None:
                    contacts['show_scope_heading'] = visible
                before = PurchaseOrder.objects.values().get(pk=self.order.pk)
                snapshot = document_preview_order({'contact_persons': contacts}, base=self.order)
                pdf, warnings = build_purchase_order_pdf(snapshot)
                self.assertFalse(warnings)
                with fitz.open(stream=pdf, filetype='pdf') as document:
                    text = '\n'.join(page.get_text() for page in document)
                self.assertEqual('PO DESCRIPTION & SCOPE' in text, visible is not False)
                self.assertIn('Retained engineering scope narrative.', text)
                self.assertIn('Agreed opening statement.', text)
                self.assertEqual(PurchaseOrder.objects.values().get(pk=self.order.pk), before)
                response = self.client.patch(self.url, {'contact_persons': contacts}, format='json')
                self.assertEqual(response.status_code, 200, response.data)
                self.order.refresh_from_db()
                self.assertEqual(self.client.get(self.url).data['contact_persons'], contacts)
                word = Document(BytesIO(build_purchase_order_docx(self.order)))
                word_text = '\n'.join(paragraph.text for paragraph in word.paragraphs)
                self.assertEqual('PO DESCRIPTION & SCOPE' in word_text, visible is not False)
                self.assertIn('Retained engineering scope narrative.', word_text)
                self.assertIn('Agreed opening statement.', word_text)

    def test_scope_heading_rejects_invalid_values_and_preserves_approved_content_lock(self):
        before = PurchaseOrder.objects.values().get(pk=self.order.pk)
        for value in ('false', 0, 1, None, [], {}):
            contacts = {**self.order.contact_persons, 'show_scope_heading': value}
            response = self.client.patch(self.url, {'contact_persons': contacts}, format='json')
            self.assertEqual(response.status_code, 400, response.data)
            with self.assertRaises(ValidationError):
                document_preview_order({'contact_persons': contacts}, base=self.order)
            self.assertEqual(PurchaseOrder.objects.values().get(pk=self.order.pk), before)
        self.approve()
        approved = PurchaseOrder.objects.values().get(pk=self.order.pk)
        response = self.client.patch(self.url, {'contact_persons': {
            **self.order.contact_persons, 'show_scope_heading': False,
        }}, format='json')
        self.assertEqual(response.status_code, 400, response.data)
        self.assertIn('revised purchase order', str(response.data))
        self.assertEqual(PurchaseOrder.objects.values().get(pk=self.order.pk), approved)

    def test_empty_current_scope_omits_the_page_without_changing_saved_content_or_approval(self):
        self.order.description = '<p>Previously approved scope.</p>'
        self.order.save()
        self.approve()
        before = PurchaseOrder.objects.values().get(pk=self.order.pk)
        contacts = {**self.order.contact_persons, 'order_introduction': '', 'show_scope_heading': False}
        snapshot = document_preview_order({
            'contact_persons': contacts, 'description': '<p>&nbsp;<br></p>',
        }, base=self.order)
        self.assertFalse(snapshot.approval_log)
        content, warnings = build_purchase_order_pdf(snapshot)
        self.assertFalse(warnings)
        with fitz.open(stream=content, filetype='pdf') as pdf:
            self.assertEqual(len(pdf), 2)
            self.assertIn('SUMMARY OF PRICES', pdf[1].get_text())
            self.assertNotIn('PURCHASE ORDER:', '\n'.join(page.get_text() for page in pdf))
        word = Document(BytesIO(build_purchase_order_docx(snapshot)))
        self.assertFalse(any(paragraph.text.startswith('PURCHASE ORDER:') for paragraph in word.paragraphs))
        self.assertEqual(PurchaseOrder.objects.values().get(pk=self.order.pk), before)
