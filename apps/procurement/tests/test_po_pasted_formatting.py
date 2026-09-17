"""Word/browser markup emitted by paste and the editor formatting commands."""
from io import BytesIO
from pathlib import Path
from unittest import TestCase

import fitz
from docx import Document
from docx.oxml.ns import qn
from reportlab.lib.units import mm
from reportlab.lib.pagesizes import A4

from apps.procurement.services.po_rich_content import parse_rich_content, pdf_rich_flowables
from apps.procurement.services.purchase_order_exports import _pdf_styles, build_purchase_order_pdf, build_purchase_order_docx
from apps.procurement.tests import test_purchase_order_exports as export_tests


class PurchaseOrderPastedFormattingTests(TestCase):
    def fixture(self):
        return (Path(__file__).parent / 'fixtures' / 'po_word_paste.html').read_text(encoding='utf-8')

    def render(self, description):
        order = export_tests.PurchaseOrderExportTests()._order()
        order.description = description
        content, warnings = build_purchase_order_pdf(order)
        self.assertFalse(warnings)
        return fitz.open(stream=content, filetype='pdf'), Document(BytesIO(build_purchase_order_docx(order)))

    def test_word_pasted_tabs_use_same_label_value_stops_in_pdf_and_editable_word(self):
        pdf, word = self.render(self.fixture())
        page = pdf[1]
        label = page.search_for('Supplier:')[0]
        supplier = page.search_for('Example Engineering LLC')[0]
        reference = page.search_for('RAD-PRJ-PUR-0126_SEP2026')[0]
        hours = page.search_for('Hours:')[0]
        self.assertAlmostEqual(supplier.x0 - label.x0, 96, delta=.5)
        self.assertAlmostEqual(reference.x0, supplier.x0, delta=.5)
        self.assertAlmostEqual(hours.x0 - label.x0, 210, delta=.5)
        paragraph = next(p for p in word.paragraphs if p.text.startswith('Supplier:'))
        self.assertEqual(paragraph.text, 'Supplier:\tExample Engineering LLC')
        self.assertEqual([round(stop.position.pt) for stop in paragraph.paragraph_format.tab_stops], [96, 210, 324])
        self.assertIn('<w:tab/>', paragraph._p.xml)

    def test_nbsp_and_word_spaceruns_are_not_collapsed(self):
        source = '<p>Label A&nbsp;&nbsp;&nbsp;&nbsp;B.</p><p>Code:<span style="mso-spacerun:yes">    </span>Value</p>'
        blocks = parse_rich_content(source)
        self.assertIn('A\u00a0\u00a0\u00a0\u00a0B', ''.join(run.text for run in blocks[0].runs))
        self.assertEqual(''.join(run.text for run in blocks[1].runs), 'Code:    Value')
        pdf, word = self.render(source)
        text = pdf[1].get_text()
        self.assertIn('A    B', text.replace('\u00a0', ' '))
        self.assertIn('Code:    Value', text.replace('\u00a0', ' '))
        self.assertTrue(any('A\u00a0\u00a0\u00a0\u00a0B' in p.text for p in word.paragraphs))

    def test_browser_execcommand_font_tags_and_link_defaults_keep_style(self):
        source = '<p><font face="Times New Roman" size="4" color="#c00000">Edited red font</font> <a href="https://example.test/plain">Editor hyperlink</a> <a href="https://example.test/explicit" style="color:#008800;text-decoration:none">Explicit green link</a></p>'
        pdf, word = self.render(source)
        spans = [span for block in pdf[1].get_text('dict')['blocks'] if 'lines' in block for line in block['lines'] for span in line['spans']]
        red = next(span for span in spans if span['text'] == 'Edited red font')
        self.assertEqual(red['color'], 0xc00000)
        self.assertAlmostEqual(red['size'], 13.5)
        self.assertIn('Times', red['font'])
        link = next(span for span in spans if 'Editor hyperlink' in span['text'])
        self.assertEqual(link['color'], 0x1d4ed8)
        paragraph = next(p for p in word.paragraphs if 'Edited red font' in p.text)
        hyperlinks = paragraph._p.findall(qn('w:hyperlink'))
        self.assertEqual(hyperlinks[0].find('.//' + qn('w:color')).get(qn('w:val')).lower(), '1d4ed8')
        self.assertEqual(hyperlinks[0].find('.//' + qn('w:u')).get(qn('w:val')), 'single')
        self.assertEqual(hyperlinks[1].find('.//' + qn('w:color')).get(qn('w:val')).lower(), '008800')

    def test_word_table_widths_padding_row_height_and_cell_alignment_survive(self):
        source = self.fixture()
        table_block = next(block for block in parse_rich_content(source) if block.kind == 'table')
        flowable = pdf_rich_flowables([table_block], 176 * mm, _pdf_styles()['body'])[0]
        flowable.wrap(176 * mm, 1000)
        self.assertAlmostEqual(flowable._colWidths[1] / sum(flowable._colWidths), .30)
        self.assertGreaterEqual(flowable._rowHeights[1], 48)
        self.assertEqual(flowable._cellStyles[1][0].leftPadding, 2)
        self.assertEqual(flowable._cellStyles[1][0].rightPadding, 9)
        self.assertEqual(flowable._cellStyles[1][0].topPadding, 3)
        self.assertEqual(flowable._cellStyles[1][0].bottomPadding, 6)
        self.assertEqual(flowable._cellStyles[1][2].valign, 'MIDDLE')
        self.assertEqual(flowable._cellStyles[1][4].valign, 'BOTTOM')
        _, word = self.render(source)
        table = next(table for table in word.tables if table.cell(0, 0).text == 'SN')
        self.assertEqual(table.rows[1].height.pt, 48)
        margins = table.cell(1, 0)._tc.find('.//' + qn('w:tcMar'))
        self.assertEqual({child.tag.split('}')[1]: int(child.get(qn('w:w'))) for child in margins}, {'top': 60, 'right': 180, 'bottom': 120, 'left': 40})
        self.assertEqual(table.cell(1, 2).vertical_alignment, 1)
        self.assertEqual(table.cell(1, 4).vertical_alignment, 3)

    def test_editor_insert_table_padding_overrides_default_cell_padding(self):
        source = '<table style="border-collapse:collapse;width:100%"><tbody><tr><th style="border:1px solid #64748b;padding:6px;background:#f1f5f9;font-weight:700;">Heading 1</th><th>Default cell</th></tr><tr><td style="border:1px solid #64748b;padding:6px;">Cell</td><td>Default cell</td></tr></tbody></table>'
        table = next(block for block in parse_rich_content(source) if block.kind == 'table')
        flowable = pdf_rich_flowables([table], 176 * mm, _pdf_styles()['body'])[0]
        flowable.wrap(176 * mm, 1000)
        self.assertEqual(flowable._cellStyles[1][0].leftPadding, 4.5)
        self.assertEqual(flowable._cellStyles[1][1].leftPadding, 6)
        self.assertEqual(flowable._cellvalues[1][0][0].style.leftIndent, 0)

    def test_real_browser_toolbar_output_preserves_combined_font_commands_and_table(self):
        # Captured from the real editor after Bold, Underline, colour and the
        # native size=4 dropdown actions, rather than hand-normalized HTML.
        source = (Path(__file__).parent / 'fixtures' / 'po_browser_formatting.html').read_text(encoding='utf-8')
        pdf, word = self.render(source)
        spans = [span for block in pdf[1].get_text('dict')['blocks'] if 'lines' in block for line in block['lines'] for span in line['spans']]
        formatted = next(span for span in spans if span['text'] == 'Current narrative text')
        self.assertEqual(formatted['color'], 0xc00000)
        self.assertIn('Bold', formatted['font'])
        self.assertAlmostEqual(formatted['size'], 13.5)
        authored = next(span for span in spans if span['text'] == 'Authored amount')
        self.assertEqual(authored['color'], 0xc00000)
        self.assertIn('Times-Bold', authored['font'])
        self.assertAlmostEqual(authored['size'], 14)
        run = next(p for p in word.paragraphs if p.text == 'Current narrative text').runs[0]
        self.assertTrue(run.bold)
        self.assertTrue(run.underline)
        self.assertEqual(run.font.size.pt, 13.5)
        self.assertEqual(str(run.font.color.rgb).lower(), 'c00000')
        table = next(block for block in parse_rich_content(source) if block.kind == 'table')
        rendered = pdf_rich_flowables([table], 176 * mm, _pdf_styles()['body'])[0]
        rendered.wrap(176 * mm, 1000)
        self.assertAlmostEqual(sum(rendered._colWidths), 176 * mm * .8)
        self.assertAlmostEqual(rendered._colWidths[0] / sum(rendered._colWidths), .3)
        self.assertGreaterEqual(rendered._rowHeights[0], 48)
        vertical_lines = [item for drawing in pdf[1].get_drawings() for item in drawing['items']
                          if item[0] == 'l' and abs(item[1].x - item[2].x) < .1 and abs(item[1].y - item[2].y) >= 47]
        self.assertTrue(vertical_lines)
        left, right = min(item[1].x for item in vertical_lines), max(item[1].x for item in vertical_lines)
        self.assertAlmostEqual(right - left, (A4[0] - 32 * mm - 12) * .8, delta=.5)
