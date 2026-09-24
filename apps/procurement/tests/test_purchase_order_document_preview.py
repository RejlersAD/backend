"""Editor previews use real exports without mutating records or source files."""
from datetime import date
from io import BytesIO
import json
from unittest.mock import patch

import fitz
from docx import Document
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.utils import timezone
from reportlab.pdfgen import canvas

from apps.procurement.models import PurchaseOrder, PODocument, Vendor
from apps.procurement.services.purchase_order_document_preview import document_preview_order
from apps.procurement.services.purchase_order_exports import build_purchase_order_pdf
from . import test_signed_po_originating_pr as originating


@override_settings(ROOT_URLCONF=originating.__name__)
class PurchaseOrderDocumentPreviewTests(TestCase):
    setUp = originating.SignedPOOriginatingPRTests.setUp
    revoke = originating.SignedPOOriginatingPRTests.revoke

    def preview(self, snapshot=None, **options):
        payload = {'snapshot': json.dumps(snapshot or {
            'po_number': 'RAD-PRJ-PUR-0126_SEP2026', 'po_date': '2026-09-15',
            'vendor': str(self.vendor.pk), 'title': 'Current editor title',
            'description': '<p>Current unsaved narrative <b>Important term</b></p>',
            'total_amount': '4486.68', 'tax_amount': '224.33', 'currency': 'USD',
            'items': [{'description': 'Piping engineer', 'quantity': 396, 'unit_price': 11}],
        }), 'format': 'pdf', **options}
        return self.client.post('/api/v1/procurement/orders/preview-document/', payload, format='multipart')

    def order(self, **values):
        return PurchaseOrder.objects.create(
            vendor=self.vendor, title='Saved title', description='<p>Saved narrative</p>',
            po_number='RAD-PRJ-PUR-0126_SEP2026', total_amount=100, created_by=self.user, **values)

    def text(self, response):
        self.assertEqual(response.status_code, 200, getattr(response, 'data', response.content[:500]))
        with fitz.open(stream=response.content, filetype='pdf') as document:
            return '\n'.join(page.get_text() for page in document)

    def test_new_draft_pdf_and_word_use_current_editor_content_without_any_save(self):
        with patch('django.core.files.storage.default_storage.save') as storage_save:
            pdf = self.preview()
            self.assertIn('Current unsaved narrative', self.text(pdf))
            word = self.preview(format='word')
            self.assertEqual(word.status_code, 200)
            document = Document(BytesIO(word.content))
            self.assertIn('Current unsaved narrative', '\n'.join(p.text for p in document.paragraphs))
            self.assertFalse(storage_save.called)
        self.assertFalse(PurchaseOrder.objects.exists())
        self.assertFalse(PODocument.objects.exists())
        self.assertEqual(Vendor.objects.count(), 1)
        self.assertEqual(pdf['Cache-Control'], 'no-store')
        self.assertIn('.pdf', pdf['Content-Disposition'])
        self.assertIn('.docx', word['Content-Disposition'])

    def test_saved_snapshot_exports_edits_without_updating_the_saved_record(self):
        order = self.order()
        before = PurchaseOrder.objects.values().get(pk=order.pk)
        response = self.preview({'description': '<p>Unsaved replacement scope</p>'}, order_id=str(order.pk))
        text = self.text(response)
        self.assertIn('Unsaved replacement scope', text)
        self.assertNotIn('Saved narrative', text)
        self.assertEqual(PurchaseOrder.objects.values().get(pk=order.pk), before)

    def test_new_editor_cover_keeps_contact_blank_and_shows_projects_and_unsigned_ceo(self):
        snapshot = {
            'po_number': 'RAD-PRJ-PUR-0126_SEP2026', 'vendor': str(self.vendor.pk),
            'title': 'Three-project engineering scope', 'total_amount': '100.00',
            'seller_reference': 'Seller Reference Contact', 'seller_contact_person': '',
            'seller_phone': '+971500001234',
            'project_number': '590001, 590002, 590003',
            'approved_by_name': 'Untrusted client signer', 'approved_date': '2026-09-15',
            'approval_signature': 'untrusted-client-signature',
        }
        with patch('apps.procurement.services.purchase_order_exports.completed_jarmo_profile_artwork') as artwork:
            response = self.preview(snapshot)
            self.assertEqual(response.status_code, 200)
            word_response = self.preview(snapshot, format='word')
            self.assertEqual(word_response.status_code, 200)
        artwork.assert_not_called()
        with fitz.open(stream=response.content, filetype='pdf') as document:
            cover = document[0].get_text()
            self.assertIn('590001, 590002, 590003', cover)
            self.assertIn('Jarmo Suominen', cover)
            self.assertIn('CEO, Rejlers Abu Dhabi', cover)
            self.assertIn('Approval not requested:', cover)
            self.assertNotIn('Approval pending:', cover)
            self.assertNotIn('Approved by:', cover)
            self.assertNotIn('Untrusted client signer', cover)
            self.assertEqual(cover.count('Seller Reference Contact'), 1)
            self.assertRegex(cover, r'Contact Person:\s+Phone Number:')
        word = Document(BytesIO(word_response.content))
        # Company cover fields are native nested Word tables.
        text = '\n'.join(word.element.xpath('//w:t/text()'))
        self.assertIn('590001, 590002, 590003', text)
        self.assertIn('Jarmo Suominen', text)
        self.assertIn('CEO, Rejlers Abu Dhabi', text)
        self.assertIn('Approval not requested:', text)
        self.assertNotIn('Approval pending:', text)
        self.assertNotIn('Approved by:', text)
        self.assertNotIn('Untrusted client signer', text)
        self.assertRegex(text, r'Contact Person:\s+Phone Number:')
        self.assertIn('Fax:', text)
        self.assertIn('Email:', text)
        self.assertFalse(PurchaseOrder.objects.exists())
        self.assertFalse(PODocument.objects.exists())

    def test_unchanged_snapshot_has_identical_page_rendering_to_saved_export(self):
        order = self.order(status='sent', approved_by_name='Recorded signer', approved_date=date(2026, 9, 15),
                           approved_at=timezone.now())
        saved, _ = build_purchase_order_pdf(order)
        response = self.preview({'description': order.description}, order_id=str(order.pk))
        self.assertEqual(response.status_code, 200)
        with fitz.open(stream=saved, filetype='pdf') as first, fitz.open(stream=response.content, filetype='pdf') as second:
            self.assertEqual(len(first), len(second))
            self.assertEqual([page.get_pixmap().samples for page in first], [page.get_pixmap().samples for page in second])

    def test_client_cannot_forge_approval_and_saved_approval_does_not_approve_changed_narrative(self):
        fake = {'description': '<p>New draft</p>', 'status': 'sent', 'approved_by_name': 'Forged signer',
                'approved_date': '2026-09-15', 'approved_at': '2026-09-15T09:00:00Z',
                'approval_log': [{'status': 'Approved', 'approver': 'Forged signer'}]}
        self.assertNotIn('Forged signer', self.text(self.preview(fake)))
        order = self.order(status='sent', approved_by_name='Recorded signer', approved_date=date(2026, 9, 15),
                           approved_at=timezone.now())
        before = PurchaseOrder.objects.values().get(pk=order.pk)
        text = self.text(self.preview(fake, order_id=str(order.pk)))
        self.assertNotIn('Recorded signer', text)
        self.assertIn('Approval not requested', text)
        self.assertIn('Jarmo Suominen', text)
        self.assertEqual(PurchaseOrder.objects.values().get(pk=order.pk), before)

    def test_uploaded_attachment_is_merged_in_memory_and_cannot_supply_storage_keys(self):
        stream = BytesIO()
        source = canvas.Canvas(stream)
        source.drawString(40, 800, 'NEW UNSAVED APPENDIX')
        source.save()
        upload = SimpleUploadedFile('appendix.pdf', stream.getvalue(), content_type='application/pdf')
        metadata = [{'new_file_index': 0, 'title': 'Current attachment title',
                     's3_key': 'foreign-secret', 'url': 'https://example.invalid/foreign'}]
        with patch('django.core.files.storage.default_storage.save') as storage_save, \
                patch('django.core.files.storage.default_storage.open') as storage_open:
            response = self.preview(attachments=upload, attachment_metadata=json.dumps(metadata))
            text = self.text(response)
            self.assertIn('Current attachment title', text)
            self.assertIn('NEW UNSAVED APPENDIX', text)
            storage_save.assert_not_called()
            storage_open.assert_not_called()
        self.assertFalse(PODocument.objects.exists())

    def test_saved_attachment_selection_uses_only_base_indices_and_preserves_order(self):
        order = self.order(attachments=[{'filename': 'first.pdf', 's3_key': 'owned-first'},
                                       {'filename': 'second.pdf', 's3_key': 'owned-second'}])
        result = document_preview_order({}, base=order, attachment_metadata=[
            {'existing_attachment_index': 1, 'title': 'Second first', 's3_key': 'foreign'},
            {'existing_attachment_index': 0, 'title': 'First second'},
        ])
        self.assertEqual([row['s3_key'] for row in result.attachments], ['owned-second', 'owned-first'])
        self.assertEqual(order.attachments[0]['s3_key'], 'owned-first')
        response = self.preview(order_id=str(order.pk), attachment_metadata=json.dumps([{'existing_attachment_index': 2}]))
        self.assertEqual(response.status_code, 400)
        result = document_preview_order({}, base=order, attachment_metadata=[])
        self.assertEqual(result.attachments, [])

    def test_render_permissions_match_new_draft_and_saved_edit_actions(self):
        self.revoke('read')
        self.assertEqual(self.preview().status_code, 200)  # create-only editor still works
        self.revoke('create')
        self.assertEqual(self.preview().status_code, 403)
        order = self.order()
        self.assertEqual(self.preview(order_id=str(order.pk)).status_code, 200)
        self.revoke('update')
        self.assertEqual(self.preview(order_id=str(order.pk)).status_code, 403)

    def test_invalid_snapshot_data_returns_field_errors_without_orphans(self):
        for snapshot in ({'total_amount': 'NaN'}, {'items': ['bad row']}, {'po_date': 'not a date'},
                         {'vendor': 'not a vendor'}, {'description': {'script': 'bad'}},
                         {'items': [{'quantity': 'garbage'}]}):
            with self.subTest(snapshot=snapshot):
                self.assertEqual(self.preview(snapshot).status_code, 400)
        self.assertEqual(self.preview(format='html').status_code, 400)
        self.assertFalse(PurchaseOrder.objects.exists())

    def test_empty_new_form_still_previews_without_registration_requirements(self):
        response = self.preview({'vendor': '', 'title': '', 'description': '', 'po_number': '',
                                 'total_amount': '', 'net_amount': '', 'items': [], 'vat_percentage': '',
                                 'vat_basis': 'unconfirmed', 'summary': '', 'contact_persons': {}})
        self.assertIn('Vendor not selected', self.text(response))
        self.assertFalse(PurchaseOrder.objects.exists())

    def test_unsupported_nested_paste_returns_actionable_error_without_saving(self):
        response = self.preview({'description': '<div>' * 90 + 'Nested clause' + '</div>' * 90})
        self.assertEqual(response.status_code, 400)
        self.assertIn('description', response.data)
        self.assertFalse(PurchaseOrder.objects.exists())

    def test_purchase_summary_survives_normal_save_and_matches_snapshot_export(self):
        order = self.order()
        contacts = {'commercial': [{'name': 'Existing contact'}], 'purchase_summary': 'Short vendor summary'}
        response = self.client.patch(f'/api/v1/procurement/orders/{order.pk}/',
                                     {'contact_persons': contacts}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        order.refresh_from_db()
        self.assertEqual(order.contact_persons, contacts)
        saved, _ = build_purchase_order_pdf(order)
        preview = self.preview({'summary': 'Short vendor summary', 'contact_persons': contacts}, order_id=str(order.pk))
        self.assertIn('Short vendor summary', self.text(preview))
        with fitz.open(stream=saved, filetype='pdf') as first, fitz.open(stream=preview.content, filetype='pdf') as second:
            self.assertEqual([p.get_pixmap().samples for p in first], [p.get_pixmap().samples for p in second])

    def test_explicit_project_removals_survive_save_reload_and_pdf_preview(self):
        self.pr.project_details = [
            {'project_number': '5901142'}, {'project_number': '5901086'}, {'project_number': '5901056'},
        ]
        self.pr.save(update_fields=['project_details'])
        order = self.order(pr_reference=self.pr, project_number='5901142, 5901086, 5901056', rad_project_no='5901142')
        endpoint = f'/api/v1/procurement/orders/{order.pk}/'
        for number, selections in (
            ('5901086', [{'project_number': '5901086', 'project_name': 'Selected project'}]),
            ('', []),
        ):
            with self.subTest(number=number):
                contacts = {'commercial': [{'name': 'Existing contact'}], 'project_selections': selections}
                snapshot = {'project_number': number, 'contact_persons': contacts}
                response = self.client.patch(endpoint, snapshot, format='json')
                self.assertEqual(response.status_code, 200, response.data)
                order.refresh_from_db()
                self.assertEqual(order.project_number, number)
                self.assertEqual(order.contact_persons, contacts)
                reloaded = self.client.get(endpoint)
                self.assertEqual(reloaded.status_code, 200, reloaded.data)
                self.assertEqual(reloaded.data['project_number'], number)
                self.assertEqual(reloaded.data['project_display'], number or None)
                self.assertEqual(reloaded.data['contact_persons']['project_selections'], selections)
                saved, _ = build_purchase_order_pdf(order)
                with fitz.open(stream=saved, filetype='pdf') as pdf:
                    saved_text = pdf[0].get_text()
                preview_text = self.text(self.preview(snapshot, order_id=str(order.pk)))
                for text in (saved_text, preview_text):
                    self.assertNotIn('5901142', text)
                    self.assertNotIn('5901056', text)
                    if number:
                        self.assertIn(number, text)
                    else:
                        self.assertNotIn('5901086', text)
                        self.assertNotIn('Multiple Projects', text)
                        self.assertRegex(text, r'Project:\s+—')
        order.approval_log = [{'stage': 'CEO', 'status': 'Approved'}]
        order.save(update_fields=['approval_log'])
        response = self.client.patch(endpoint, {
            'project_number': '5901142',
            'contact_persons': {'project_selections': [{'project_number': '5901142'}]},
        }, format='json')
        self.assertEqual(response.status_code, 400, response.data)
        order.refresh_from_db()
        self.assertEqual(order.project_number, '')
        self.assertEqual(order.contact_persons['project_selections'], [])
