from io import BytesIO
from unittest import TestCase

import fitz
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.oxml.ns import qn
from PIL import Image

from apps.procurement.services.purchase_order_word_pages import (
    MAX_PAGE_IMAGE_EDGE,
    MAX_WORD_PAGE_POINTS,
    append_pdf_pages,
)


def supporting_pdf():
    """Synthetic cover, portrait evidence, and a rotated landscape drawing."""
    with fitz.open() as source:
        for index, (width, height, color, rotation) in enumerate((
            (595, 842, (1, 0, 0), 0),
            (420, 595, (0, 1, 0), 0),
            (420, 595, (0, 0, 1), 90),
        )):
            page = source.new_page(width=width, height=height)
            page.draw_rect(page.rect, color=None, fill=color)
            page.insert_text((30, 40), f'Synthetic supporting page {index + 1}')
            page.set_rotation(rotation)
        return source.tobytes()


def image_parts(document):
    for anchor in document._element.body.xpath('.//wp:anchor'):
        relationship_id = anchor.xpath('.//a:blip')[0].get(qn('r:embed'))
        yield document.part.related_parts[relationship_id]


class PurchaseOrderWordPageTests(TestCase):
    def test_editable_body_and_branding_survive_and_supporting_pages_have_no_overlay(self):
        document = Document()
        document.add_paragraph('Editable purchase order body')
        document.sections[0].header.paragraphs[0].text = 'Purchase order header'
        document.sections[0].footer.paragraphs[0].text = 'Purchase order footer'
        self.assertEqual(append_pdf_pages(document, supporting_pdf()), 3)
        stream = BytesIO()
        document.save(stream)
        reopened = Document(stream)
        self.assertEqual(len(reopened.sections), 4)
        self.assertEqual(reopened.paragraphs[0].text, 'Editable purchase order body')
        self.assertEqual(reopened.sections[0].header.paragraphs[0].text, 'Purchase order header')
        self.assertEqual(reopened.sections[0].footer.paragraphs[0].text, 'Purchase order footer')
        for section in reopened.sections[1:]:
            self.assertEqual(section.start_type, WD_SECTION_START.NEW_PAGE)
            for name in ('header', 'first_page_header', 'even_page_header',
                         'footer', 'first_page_footer', 'even_page_footer'):
                part = getattr(section, name)
                self.assertFalse(part.is_linked_to_previous)
                self.assertEqual(''.join(paragraph.text for paragraph in part.paragraphs), '')

    def test_image_pages_preserve_color_order_source_dimensions_and_rotation(self):
        document = Document()
        append_pdf_pages(document, supporting_pdf())
        self.assertEqual(len(document.sections), 3)
        for section, size in zip(document.sections, ((595, 842), (420, 595), (595, 420))):
            self.assertAlmostEqual(section.page_width.pt, size[0], delta=.05)
            self.assertAlmostEqual(section.page_height.pt, size[1], delta=.05)
        self.assertEqual(document.sections[-1].orientation, WD_ORIENT.LANDSCAPE)
        for part, expected in zip(image_parts(document), ((255, 0, 0), (0, 255, 0), (0, 0, 255))):
            with Image.open(BytesIO(part.blob)) as image:
                self.assertEqual(image.getpixel((image.width // 2, image.height // 2)), expected)
                self.assertLessEqual(max(image.size), MAX_PAGE_IMAGE_EDGE)
        self.assertEqual(len(list(image_parts(document))), 3)

    def test_floating_pages_use_absolute_page_origin_and_no_inline_or_page_breaks(self):
        document = Document()
        append_pdf_pages(document, supporting_pdf())
        self.assertEqual(len(document._element.body.xpath('.//wp:inline')), 0)
        self.assertEqual(len(document._element.body.xpath('.//w:br')), 0)
        anchors = document._element.body.xpath('.//wp:anchor')
        self.assertEqual(len(anchors), 3)
        for anchor, section in zip(anchors, document.sections):
            for name in ('positionH', 'positionV'):
                position = anchor.find(qn(f'wp:{name}'))
                self.assertEqual(position.get('relativeFrom'), 'page')
                self.assertEqual(position.find(qn('wp:posOffset')).text, '0')
            extent = anchor.find(qn('wp:extent'))
            self.assertEqual(int(extent.get('cx')), section.page_width)
            self.assertEqual(int(extent.get('cy')), section.page_height)

    def test_skipped_body_does_not_get_duplicated_and_appending_no_pages_is_unchanged(self):
        document = Document()
        document.add_paragraph('Editable purchase order')
        self.assertEqual(append_pdf_pages(document, supporting_pdf(), start_page=1), 2)
        self.assertEqual(len(document.sections), 3)
        self.assertEqual(len(list(image_parts(document))), 2)
        self.assertEqual(document._element.body.xpath('.//wp:docPr')[0].get('descr'), 'Original PDF page 2')
        original = document._element.xml
        self.assertEqual(append_pdf_pages(document, supporting_pdf(), start_page=3), 0)
        self.assertEqual(document._element.xml, original)

    def test_empty_paragraphs_do_not_create_a_leading_blank_page(self):
        document = Document()
        document.add_paragraph()
        append_pdf_pages(document, supporting_pdf(), start_page=2)
        self.assertEqual(len(document.sections), 1)
        self.assertEqual(len(list(image_parts(document))), 1)

    def test_oversize_drawing_is_bounded_without_changing_its_aspect_ratio(self):
        with fitz.open() as source:
            source.new_page(width=4000, height=2000)
            content = source.tobytes()
        document = Document()
        append_pdf_pages(document, content)
        section = document.sections[0]
        self.assertEqual(section.page_width.pt, MAX_WORD_PAGE_POINTS)
        self.assertEqual(section.page_width / section.page_height, 2)
        with Image.open(BytesIO(next(image_parts(document)).blob)) as image:
            self.assertLessEqual(max(image.size), MAX_PAGE_IMAGE_EDGE)
            self.assertEqual(image.width / image.height, 2)

    def test_bad_pdf_or_page_selection_fails_without_changing_the_document(self):
        document = Document()
        document.add_paragraph('Saved content')
        original = document._element.xml
        for start_page in (-1, True, 1.5, '1', 4):
            with self.subTest(start_page=start_page), self.assertRaises(ValueError):
                append_pdf_pages(document, supporting_pdf(), start_page=start_page)
            self.assertEqual(document._element.xml, original)
        with self.assertRaises(Exception):
            append_pdf_pages(document, b'not a PDF')
        self.assertEqual(document._element.xml, original)

    def test_password_protected_pdf_is_not_silently_omitted(self):
        with fitz.open() as source:
            source.new_page()
            content = source.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256,
                                    owner_pw='owner-test', user_pw='reader-test')
        document = Document()
        with self.assertRaisesRegex(ValueError, 'password'):
            append_pdf_pages(document, content)
        self.assertEqual(len(list(image_parts(document))), 0)
