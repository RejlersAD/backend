"""Word retains the PDF packet, canonical source bindings and export warnings."""

from copy import deepcopy
from io import BytesIO
import json
from unittest.mock import patch

import fitz
from docx import Document
from docx.oxml.ns import qn
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase, override_settings
from PIL import Image, ImageChops

from apps.procurement.models import PODocument, PurchaseOrder
from apps.procurement.services.purchase_order_exports import (
    _main_pdf,
    build_purchase_order_docx,
    build_purchase_order_pdf,
)
from apps.procurement.services.purchase_order_word_pages import MAX_PAGE_IMAGE_EDGE, PDF_PAGE_DPI
from apps.rbac.models import Permission, RolePermission
from . import test_purchase_order_document_preview as preview_fixtures
from . import test_purchase_order_exports as export_fixtures
from . import test_purchase_order_word_pages as page_fixtures
from . import test_signed_po_originating_pr as originating


EXPORTS = 'apps.procurement.services.purchase_order_exports'


def page_images(document):
    for anchor in document._element.body.xpath('.//wp:anchor'):
        label = anchor.find(qn('wp:docPr')).get('descr', '')
        if not label.startswith('Original PDF page '):
            continue
        relationship_id = anchor.xpath('.//a:blip')[0].get(qn('r:embed'))
        yield document.part.related_parts[relationship_id].blob


class PurchaseOrderWordPacketTests(SimpleTestCase):
    def order(self, attachments=None):
        order = export_fixtures.PurchaseOrderExportTests()._order(attachments)
        order.status = 'draft'
        order.approved_at = None
        order.approval_log = []
        order.contact_persons = {'order_introduction': 'Editable Buyer and Seller statement.'}
        return order

    def assert_pdf_tail_matches(self, order, word, pdf):
        images = list(page_images(word))
        with fitz.open(stream=_main_pdf(order), filetype='pdf') as main:
            body_count = len(main)
        with fitz.open(stream=pdf, filetype='pdf') as canonical:
            self.assertEqual(len(images), len(canonical) - body_count)
            for image_bytes, page in zip(images, list(canonical)[body_count:]):
                scale = min(PDF_PAGE_DPI / 72,
                            MAX_PAGE_IMAGE_EDGE / max(page.rect.width, page.rect.height))
                rendered = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
                expected = Image.frombytes('RGB', (rendered.width, rendered.height), rendered.samples)
                with Image.open(BytesIO(image_bytes)) as embedded:
                    self.assertEqual(embedded.size, expected.size)
                    self.assertIsNone(ImageChops.difference(embedded.convert('RGB'), expected).getbbox())

    def test_editable_body_and_ordered_attachment_covers_match_the_canonical_pdf_pixels(self):
        image = BytesIO()
        Image.new('RGB', (120, 80), 'orange').save(image, format='PNG')
        order = self.order([
            {'title': 'First evidence', 'filename': 'three-pages.pdf',
             'content_type': 'application/pdf', '_preview_content': page_fixtures.supporting_pdf()},
            {'title': 'Second evidence', 'filename': 'photo.png',
             'content_type': 'image/png', '_preview_content': image.getvalue()},
        ])
        original = deepcopy(order.__dict__)
        word_bytes, warnings = build_purchase_order_docx(order, with_warnings=True)
        canonical, pdf_warnings = build_purchase_order_pdf(order)
        self.assertEqual(warnings, pdf_warnings)
        self.assertFalse(warnings)
        word = Document(BytesIO(word_bytes))
        self.assertIn('Editable Buyer and Seller statement.', '\n'.join(word.element.xpath('//w:t/text()')))
        self.assertEqual(len(list(page_images(word))), 6)  # Two covers, three PDF pages, one image.
        self.assert_pdf_tail_matches(order, word, canonical)
        self.assertEqual(order.__dict__, original)

    def test_missing_and_unsupported_sources_keep_their_covers_and_the_same_warnings(self):
        order = self.order([
            {'title': 'Unavailable evidence', 'filename': 'missing.pdf'},
            {'title': 'Unsupported evidence', 'filename': 'drawing.dwg', '_preview_content': b'opaque'},
            {'title': 'Readable evidence', 'filename': 'readable.pdf',
             '_preview_content': page_fixtures.supporting_pdf()},
        ])
        word_bytes, warnings = build_purchase_order_docx(order, with_warnings=True)
        canonical, pdf_warnings = build_purchase_order_pdf(order)
        self.assertEqual(warnings, pdf_warnings)
        self.assertEqual(len(warnings), 2)
        self.assertIn('missing.pdf', warnings[0])
        self.assertIn('drawing.dwg', warnings[1])
        word = Document(BytesIO(word_bytes))
        self.assertEqual(len(list(page_images(word))), 6)  # All three covers remain, then three readable pages.
        self.assert_pdf_tail_matches(order, word, canonical)

    def test_signed_original_uses_confirmed_document_binding_instead_of_attachment_url_or_key(self):
        order = self.order([{
            'type': 'signed_purchase_order_pdf', 'document_id': 'confirmed-original',
            'filename': 'original.pdf', 's3_key': 'untrusted/other-order.pdf',
            'url': 'https://external.invalid/private.pdf',
        }])
        order.pk = 'synthetic-saved-order'
        sources = [{'id': 'confirmed-original', 'storage_key': 'procurement/orders/authorized/original.pdf'}]
        original = deepcopy(order.attachments)
        with patch(f'{EXPORTS}.uploaded_purchase_order_sources', return_value=sources), \
             patch('django.core.files.storage.default_storage.open',
                   side_effect=lambda *_args: BytesIO(page_fixtures.supporting_pdf())) as opened, \
             patch('requests.get') as external_get:
            content, warnings = build_purchase_order_docx(order, with_warnings=True)
        opened.assert_called_once_with('procurement/orders/authorized/original.pdf', 'rb')
        external_get.assert_not_called()
        self.assertFalse(warnings)
        self.assertEqual(len(list(page_images(Document(BytesIO(content))))), 4)
        self.assertEqual(order.attachments, original)

    def test_url_only_sources_are_not_fetched_and_legacy_byte_return_remains_supported(self):
        order = self.order([{'filename': 'remote.pdf', 'url': 'https://external.invalid/remote.pdf'}])
        with patch('django.core.files.storage.default_storage.open') as opened, \
             patch('requests.get') as external_get:
            content = build_purchase_order_docx(order)
        self.assertIsInstance(content, bytes)
        self.assertEqual(len(list(page_images(Document(BytesIO(content))))), 1)  # Cover remains.
        opened.assert_not_called()
        external_get.assert_not_called()

    def test_orders_without_attachments_keep_native_text_without_image_pages(self):
        content, warnings = build_purchase_order_docx(self.order(), with_warnings=True)
        self.assertFalse(warnings)
        self.assertFalse(list(page_images(Document(BytesIO(content)))))


@override_settings(ROOT_URLCONF=originating.__name__)
class PurchaseOrderWordPacketAPITests(TestCase):
    preview = preview_fixtures.PurchaseOrderDocumentPreviewTests.preview
    order = preview_fixtures.PurchaseOrderDocumentPreviewTests.order

    def setUp(self):
        originating.SignedPOOriginatingPRTests.setUp(self)
        permission = Permission.objects.get(module__code='procurement_orders', action='export', is_active=True)
        RolePermission.objects.get_or_create(role=self.role, permission=permission)
        cache.clear()

    def test_saved_word_download_reports_missing_attachment_without_saving_any_record_or_file(self):
        order = self.order(attachments=[{'filename': 'missing.pdf', 'title': 'Missing source'}])
        before = PurchaseOrder.objects.values().get(pk=order.pk)
        with patch('django.core.files.storage.default_storage.save') as storage_save:
            response = self.client.get(f'/api/v1/procurement/orders/{order.pk}/export-word/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['X-PO-Attachment-Warnings'], '1')
        self.assertEqual(response['Cache-Control'], 'no-store')
        self.assertIn('.docx', response['Content-Disposition'])
        self.assertEqual(len(list(page_images(Document(BytesIO(response.content))))), 1)
        self.assertEqual(PurchaseOrder.objects.values().get(pk=order.pk), before)
        self.assertFalse(PODocument.objects.exists())
        storage_save.assert_not_called()

    def test_word_preview_ignores_client_storage_override_and_reports_the_saved_source_warning(self):
        order = self.order(attachments=[{'filename': 'saved.pdf', 's3_key': 'procurement/saved/source.pdf'}])
        before = PurchaseOrder.objects.values().get(pk=order.pk)
        metadata = [{'existing_attachment_index': 0, 's3_key': 'private/unrelated.pdf',
                     'url': 'https://external.invalid/private.pdf', '_preview_content': 'client content'}]
        with patch('django.core.files.storage.default_storage.open', side_effect=FileNotFoundError) as opened, \
             patch('django.core.files.storage.default_storage.save') as storage_save:
            response = self.preview({'description': '<p>Unsaved replacement scope</p>'},
                                    order_id=str(order.pk), format='word',
                                    attachment_metadata=json.dumps(metadata))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['X-PO-Attachment-Warnings'], '1')
        opened.assert_called_once_with('procurement/saved/source.pdf', 'rb')
        storage_save.assert_not_called()
        word = Document(BytesIO(response.content))
        self.assertIn('Unsaved replacement scope', '\n'.join(word.element.xpath('//w:t/text()')))
        self.assertEqual(len(list(page_images(word))), 1)
        self.assertEqual(PurchaseOrder.objects.values().get(pk=order.pk), before)
        self.assertFalse(PODocument.objects.exists())

    def test_new_word_preview_includes_uploaded_pages_without_storage_or_database_writes(self):
        upload = SimpleUploadedFile('supporting.pdf', page_fixtures.supporting_pdf(),
                                    content_type='application/pdf')
        with patch('django.core.files.storage.default_storage.save') as storage_save, \
             patch('django.core.files.storage.default_storage.open') as storage_open:
            response = self.preview(format='word',
                                    attachment_metadata=json.dumps([{'new_file_index': 0}]),
                                    attachments=[upload])
        self.assertEqual(response.status_code, 200)
        self.assertNotIn('X-PO-Attachment-Warnings', response)
        self.assertEqual(len(list(page_images(Document(BytesIO(response.content))))), 4)
        storage_save.assert_not_called()
        storage_open.assert_not_called()
        self.assertFalse(PurchaseOrder.objects.exists())
        self.assertFalse(PODocument.objects.exists())
