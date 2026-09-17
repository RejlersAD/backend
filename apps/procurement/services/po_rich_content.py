"""One safe, editable narrative model for the PO PDF and Word renderers.

The editor's HTML is data: no stylesheet, script, network image or local path is
executed/read. Paragraphs, runs, lists, cells and explicit page breaks retain
their order instead of being flattened into a paragraph for every visual line.
"""
from __future__ import annotations

import base64
import html
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from io import BytesIO
from urllib.parse import urlsplit
from xml.sax.saxutils import escape, quoteattr

from PIL import Image as PILImage
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.platypus import Image, PageBreak, Paragraph, Spacer, Table, TableStyle


@dataclass
class Node:
    tag: str
    attrs: dict = field(default_factory=dict)
    children: list = field(default_factory=list)


@dataclass
class Run:
    text: str
    style: dict = field(default_factory=dict)


@dataclass
class Block:
    kind: str
    style: dict = field(default_factory=dict)
    runs: list = field(default_factory=list)
    rows: list = field(default_factory=list)
    image: bytes | None = None
    width: float | None = None
    height: float | None = None
    list_info: dict | None = None
    columns: list = field(default_factory=list)
    row_styles: list = field(default_factory=list)


@dataclass
class Cell:
    blocks: list
    colspan: int = 1
    rowspan: int = 1
    heading: bool = False
    style: dict = field(default_factory=dict)


_VOID = {'br', 'img', 'hr', 'meta', 'link', 'input', 'wbr', 'col'}
_UNSAFE = {'script', 'style', 'iframe', 'object', 'embed', 'head', 'svg', 'math', 'template'}
_BLOCKS = {'p', 'div', 'section', 'article', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'pre', 'blockquote', 'ul', 'ol', 'li', 'table', 'hr'}
_INHERITED = {'font_size', 'font_family', 'bold', 'italic', 'underline', 'strike', 'color', 'background', 'align', 'line_height', 'pre', 'vertical', 'href', 'list_depth', 'tab_stops', 'tab_interval'}


class _Tree(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node('root')
        self.stack = [self.root]

    def handle_starttag(self, tag, attrs):
        # Deep malformed HTML must not exhaust the renderer's recursion stack.
        if len(self.stack) > 80:
            raise ValueError('The narrative contains excessively nested markup.')
        node = Node(tag, dict(attrs))
        self.stack[-1].children.append(node)
        if tag not in _VOID:
            self.stack.append(node)

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        if tag not in _VOID:
            self.handle_endtag(tag)

    def handle_endtag(self, tag):
        for index in range(len(self.stack) - 1, 0, -1):
            if self.stack[index].tag == tag:
                del self.stack[index:]
                break

    def handle_data(self, data):
        self.stack[-1].children.append(data)


def _length(value, base=10.5):
    match = re.fullmatch(r'\s*(-?\d+(?:\.\d+)?)\s*(pt|px|em|rem|%|in|cm|mm)?\s*', str(value or ''), re.I)
    if not match:
        return None
    amount, unit = float(match[1]), (match[2] or 'px').lower()
    return amount * {'pt': 1, 'px': .75, 'em': base, 'rem': 12, '%': base / 100, 'in': 72, 'cm': 72 / 2.54, 'mm': 72 / 25.4}[unit]


def _color(value):
    value = str(value or '').strip()
    if not value or value in {'transparent', 'inherit', 'initial'}:
        return None
    rgb = re.fullmatch(r'rgba?\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)(?:\s*,\s*[\d.]+)?\s*\)', value)
    if rgb:
        return '#%02x%02x%02x' % tuple(min(255, int(part)) for part in rgb.groups())
    if re.fullmatch(r'#[a-fA-F0-9]{3}(?:[a-fA-F0-9]{3})?', value):
        return colors.HexColor(value if len(value) == 7 else '#' + ''.join(char * 2 for char in value[1:])).hexval().replace('0x', '#')
    named = colors.getAllNamedColors().get(value.lower())
    return named.hexval().replace('0x', '#') if named is not None else None


def _style(node, parent):
    result = {key: value for key, value in parent.items() if key in _INHERITED}
    tag, attrs = node.tag, node.attrs
    if tag in {'b', 'strong', 'th'}:
        result['bold'] = True
    if tag == 'th':
        result['align'] = 'center'
    if tag in {'i', 'em'}:
        result['italic'] = True
    if tag == 'u':
        result['underline'] = True
    if tag in {'s', 'strike', 'del'}:
        result['strike'] = True
    if tag == 'a' and _safe_link(attrs.get('href')):
        result.update(color='#1d4ed8', underline=True)
    if tag in {'sub', 'sup'}:
        result['vertical'] = tag
    if tag in {'pre', 'code'}:
        result['font_family'] = 'Courier New'
    if tag == 'pre':
        result['pre'] = True
    if re.fullmatch('h[1-6]', tag):
        result.update(font_size={1: 24, 2: 20, 3: 16, 4: 14, 5: 12, 6: 10.5}[int(tag[1])], bold=True, heading=int(tag[1]), before=8, after=5)
    if tag == 'blockquote':
        result.update(indent=18, italic=True)
    css = dict(pair.split(':', 1) for pair in (part.strip() for part in attrs.get('style', '').split(';')) if ':' in pair)
    css = {key.strip().lower(): value.strip() for key, value in css.items()}
    for name in ('margin', 'padding'):
        values = css.get(name, '').split()
        if 1 <= len(values) <= 4:
            top, right, bottom, left = (values * 4)[:4] if len(values) == 1 else (values[0], values[1], values[0], values[1]) if len(values) == 2 else (values[0], values[1], values[2], values[1]) if len(values) == 3 else values
            for side, value in zip(('top', 'right', 'bottom', 'left'), (top, right, bottom, left)):
                css.setdefault(f'{name}-{side}', value)
    align = css.get('text-align', attrs.get('align', '')).lower()
    if align in {'left', 'right', 'center', 'justify'}:
        result['align'] = align
    font = css.get('font-family', attrs.get('face'))
    if font:
        result['font_family'] = font.split(',')[0].strip(' "\'')[:100]
    size = _length(css.get('font-size'), result.get('font_size', 10.5))
    if size is None and tag == 'font' and attrs.get('size'):
        # Chrome's native execCommand('fontSize') uses legacy HTML sizes in
        # CSS pixels (10/13/16/18/24/32/48), converted here to physical points.
        size = {'1': 7.5, '2': 9.75, '3': 12, '4': 13.5, '5': 18, '6': 24, '7': 36}.get(attrs['size'])
    if size is not None:
        result['font_size'] = min(72, max(5, size))
    if 'font-weight' in css:
        result['bold'] = css['font-weight'] in {'bold', 'bolder', '600', '700', '800', '900'}
    if 'font-style' in css:
        result['italic'] = css['font-style'] in {'italic', 'oblique'}
    decoration = css.get('text-decoration', css.get('text-decoration-line', ''))
    if decoration:
        result.update(underline='underline' in decoration, strike='line-through' in decoration)
    for field_name, css_name in (('color', 'color'), ('background', 'background-color')):
        value = _color(css.get(css_name, attrs.get('color') if field_name == 'color' else css.get('background')))
        if value:
            result[field_name] = value
    line = css.get('line-height', '')
    if re.fullmatch(r'\d+(\.\d+)?', line):
        result['line_height'] = ('multiple', min(5, max(.5, float(line))))
    elif line and line != 'normal':
        value = _length(line, result.get('font_size', 10.5))
        if value:
            result['line_height'] = ('points', min(200, max(5, value)))
    for field_name, css_names in (
        ('before', ('margin-top',)), ('after', ('margin-bottom',)),
        ('indent', ('margin-left', 'padding-left')), ('right_indent', ('margin-right', 'padding-right')),
        ('first_indent', ('text-indent',)),
    ):
        values = [_length(css.get(name), result.get('font_size', 10.5)) for name in css_names if not (tag in {'td', 'th'} and name.startswith('padding-'))]
        if any(value is not None for value in values):
            result[field_name] = max(-36 if field_name == 'first_indent' else 0, min(180, sum(value or 0 for value in values)))
    if css.get('white-space') in {'pre', 'pre-wrap', 'break-spaces'}:
        result['pre'] = True
    if css.get('mso-spacerun') == 'yes':
        result['pre'] = True
    if css.get('mso-tab-count', '').isdigit():
        result['tab_count'] = min(20, int(css['mso-tab-count']))
    stops = css.get('tab-stops', css.get('mso-tab-stops', ''))
    if stops:
        result['tab_stops'] = sorted({value for token in stops.split() if (value := _length(token)) is not None and 0 < value < 1500})
    interval = _length(css.get('mso-default-tab-stop'))
    if interval and interval > 0:
        result['tab_interval'] = interval
    if css.get('page-break-before') == 'always' or css.get('break-before') == 'page':
        result['break_before'] = True
    if css.get('page-break-after') == 'always' or css.get('break-after') == 'page':
        result['break_after'] = True
    if css.get('page-break-inside') == 'avoid' or css.get('break-inside') == 'avoid':
        result['keep'] = True
    width = css.get('width', attrs.get('width', ''))
    if str(width).strip().endswith('%'):
        try:
            result['width'] = ('fraction', min(1, max(.01, float(width[:-1]) / 100)))
        except ValueError:
            pass
    elif _length(width) is not None:
        result['width'] = ('points', max(1, _length(width)))
    padding = {}
    for side in ('top', 'right', 'bottom', 'left'):
        value = _length(css.get('padding-' + side))
        if value is not None:
            padding[side] = min(72, max(0, value))
    if padding:
        result['cell_padding'] = padding
    height = _length(css.get('height', attrs.get('height', css.get('min-height'))))
    if height is not None:
        result['min_height'] = max(0, min(600, height))
    valign = css.get('vertical-align', attrs.get('valign', '')).lower()
    if valign in {'top', 'middle', 'bottom'}:
        result['valign'] = valign
    border = css.get('border', '')
    border_width = re.search(r'(\d+(?:\.\d+)?(?:pt|px))', border)
    if border_width:
        result['border_width'] = min(5, _length(border_width[1]))
    if border in {'0', 'none', '0px'}:
        result['border_width'] = 0
    border_color = _color(css.get('border-color', ''))
    if not border_color:
        border_color = next((_color(part) for part in border.split() if _color(part)), None)
    if border_color:
        result['border_color'] = border_color
    return result


def _safe_link(value):
    value = str(value or '').strip()
    try:
        return value if urlsplit(value).scheme.lower() in {'https', 'http', 'mailto', 'tel'} else ''
    except ValueError:
        return ''


def _image(node, style):
    source = str(node.attrs.get('src') or '')
    if not re.match(r'^data:image/(?:png|jpeg|jpg|gif|webp);base64,', source, re.I) or len(source) > 8_000_000:
        return Block('paragraph', style, [Run(node.attrs.get('alt') or '[Image unavailable]', style)])
    try:
        data = base64.b64decode(source.split(',', 1)[1], validate=True)
        with PILImage.open(BytesIO(data)) as image:
            if image.width * image.height > 20_000_000:
                raise ValueError('Image too large')
            width, height = image.size
            # A common safe format also supports GIF/WebP in editable Word.
            output = BytesIO()
            image.convert('RGBA').save(output, format='PNG')
            data = output.getvalue()
    except (ValueError, OSError, PILImage.DecompressionBombError):
        return Block('paragraph', style, [Run(node.attrs.get('alt') or '[Image unavailable]', style)])
    css = dict(pair.split(':', 1) for pair in node.attrs.get('style', '').split(';') if ':' in pair)
    css = {key.strip(): value.strip() for key, value in css.items()}
    width_text = css.get('width', node.attrs.get('width', ''))
    if str(width_text).strip().endswith('%'):
        try:
            style['width_fraction'] = min(1, max(.01, float(width_text[:-1]) / 100))
        except ValueError:
            pass
        requested_width = None
    else:
        requested_width = _length(width_text)
    requested_height = _length(css.get('height', node.attrs.get('height')))
    requested_width = requested_width if requested_width and requested_width > 0 else None
    requested_height = requested_height if requested_height and requested_height > 0 else None
    if requested_width and not requested_height:
        requested_height = requested_width * height / width
    elif requested_height and not requested_width:
        requested_width = requested_height * width / height
    return Block('image', style, image=data, width=requested_width or width * .75, height=requested_height or height * .75)


class _Builder:
    def __init__(self):
        self.sequence = 0

    def content(self, nodes, parent=None):
        parent = parent or {}
        blocks, runs = [], []

        def flush(force=False):
            nonlocal runs
            if runs and any(run.text.strip() for run in runs) or force:
                if runs and not parent.get('pre'):
                    runs[0].text = runs[0].text.lstrip(' ')
                    runs[-1].text = runs[-1].text.rstrip(' ')
                blocks.append(Block('paragraph', dict(parent), runs))
            runs = []

        def visit(node, inherited):
            if isinstance(node, str):
                # NBSPs and tabs are intentional layout from Word/browser
                # editing; only collapsible HTML whitespace becomes a space.
                value = node if inherited.get('pre') else re.sub(r'[ \r\n\f\v]+', ' ', node)
                if value:
                    runs.append(Run(value, dict(inherited)))
                return
            if node.tag in _UNSAFE:
                return
            style = _style(node, inherited)
            if style.get('tab_count'):
                runs.append(Run('\t' * style['tab_count'], style))
                return
            if node.attrs.get('data-po-page-break') == 'true':
                flush()
                blocks.append(Block('break'))
                return
            if node.tag == 'br':
                runs.append(Run('\n', dict(style)))
            elif node.tag == 'img':
                flush()
                blocks.append(_image(node, style))
            elif node.tag in _BLOCKS:
                flush()
                if style.get('break_before'):
                    blocks.append(Block('break'))
                if node.tag in {'ol', 'ul'}:
                    blocks.extend(self.list(node, style, inherited.get('list_depth', 0)))
                elif node.tag == 'table':
                    blocks.append(self.table(node, style))
                elif node.tag == 'hr':
                    blocks.append(Block('paragraph', style, [Run('____________________________', style)]))
                else:
                    children = self.content(node.children, style)
                    blocks.extend(children or [Block('paragraph', style, [Run('\u00a0', style)])])
                if style.get('break_after'):
                    blocks.append(Block('break'))
            else:
                if node.tag == 'a':
                    style['href'] = _safe_link(node.attrs.get('href'))
                for child in node.children:
                    visit(child, style)

        for node in nodes:
            visit(node, parent)
        flush()
        return blocks

    def list(self, node, style, depth):
        ordered = node.tag == 'ol'
        try:
            number = int(node.attrs.get('start', 1))
        except (TypeError, ValueError):
            number = 1
        self.sequence += 1
        sequence = self.sequence
        result = []
        for child in node.children:
            if not isinstance(child, Node) or child.tag != 'li':
                continue
            try:
                if ordered and child.attrs.get('value') is not None:
                    number = int(child.attrs['value'])
                    self.sequence += 1
                    sequence = self.sequence
            except (TypeError, ValueError):
                pass
            item_style = {**_style(child, style), 'list_depth': depth + 1}
            items = self.content(child.children, item_style)
            if not items or items[0].kind != 'paragraph':
                items.insert(0, Block('paragraph', item_style, [Run('\u00a0', item_style)]))
            marker = f'{number}.' if ordered else '\u2022'
            items[0].list_info = {'ordered': ordered, 'number': number, 'depth': depth, 'sequence': sequence, 'marker': marker}
            for item in items:
                if not item.list_info:
                    item.style['indent'] = item.style.get('indent', 0) + (depth + 1) * 18
            result.extend(items)
            number += 1
        return result

    def table(self, node, style):
        rows, columns, row_styles = [], [], []

        def visit(children, header=False):
            for child in children:
                if not isinstance(child, Node):
                    continue
                if child.tag in {'thead', 'tbody', 'tfoot'}:
                    visit(child.children, child.tag == 'thead')
                elif child.tag == 'colgroup':
                    for column in child.children:
                        if isinstance(column, Node) and column.tag == 'col':
                            columns.append(_style(column, {}).get('width'))
                elif child.tag == 'tr':
                    cells = []
                    row_style = _style(child, style)
                    for cell in child.children:
                        if not isinstance(cell, Node) or cell.tag not in {'td', 'th'}:
                            continue
                        cell_style = _style(cell, row_style)
                        cell_style.setdefault('after', 0)
                        if 'cell_padding' not in cell_style and node.attrs.get('cellpadding'):
                            padding = _length(node.attrs['cellpadding'])
                            if padding is not None:
                                cell_style['cell_padding'] = {side: max(0, min(72, padding)) for side in ('top', 'right', 'bottom', 'left')}
                        spans = []
                        for key in ('colspan', 'rowspan'):
                            try:
                                spans.append(min(64, max(1, int(cell.attrs.get(key, 1)))))
                            except (TypeError, ValueError):
                                spans.append(1)
                        cells.append(Cell(self.content(cell.children, cell_style), *spans, header or cell.tag == 'th', cell_style))
                    if cells:
                        rows.append(cells)
                        row_styles.append(row_style)
        visit(node.children)
        return Block('table', style, rows=rows, columns=columns, row_styles=row_styles)


def parse_rich_content(value):
    source = str(value or '')
    # Decode whole escaped editor documents, not ordinary entities inside HTML
    # (e.g. a literal &lt;script&gt; typed in a paragraph).
    for _ in range(3):
        if re.search(r'<[a-zA-Z][^>]*>', source) or not re.search(r'&(?:amp;)*lt;', source):
            break
        source = html.unescape(source)
    if not re.search(r'<[a-zA-Z][^>]*>', source):
        return [Block('paragraph', runs=[Run(html.unescape(line))]) for line in source.splitlines() if line.strip()]
    tree = _Tree()
    tree.feed(source)
    return _Builder().content(tree.root.children)


def _pdf_font(style):
    family = style.get('font_family', '').lower()
    base = 'Courier' if any(word in family for word in ('courier', 'mono', 'consolas')) else 'Times' if any(word in family for word in ('times', 'serif', 'georgia')) and 'sans' not in family else 'Helvetica'
    bold, italic = style.get('bold'), style.get('italic')
    if base == 'Times':
        return 'Times-' + ('BoldItalic' if bold and italic else 'Bold' if bold else 'Italic' if italic else 'Roman')
    return base + ('-BoldOblique' if bold and italic else '-Bold' if bold else '-Oblique' if italic else '')


def _pdf_markup(runs):
    parts = []
    for run in runs:
        style = run.style
        value = escape(run.text).replace('\n', '<br/>')
        if style.get('pre'):
            value = value.replace(' ', '&#160;')
        attrs = f'name={quoteattr(_pdf_font(style))}'
        for key, attr in (('font_size', 'size'), ('color', 'color'), ('background', 'backColor')):
            if style.get(key):
                attrs += f' {attr}={quoteattr(str(style[key]))}'
        value = f'<font {attrs}>{value}</font>'
        for key, tag in (('underline', 'u'), ('strike', 'strike')):
            if style.get(key):
                value = f'<{tag}>{value}</{tag}>'
        if style.get('vertical') in {'sub', 'sup'}:
            value = f'<{style["vertical"]}>{value}</{style["vertical"]}>'
        if style.get('href'):
            value = f'<link href={quoteattr(style["href"])}>{value}</link>'
        parts.append(value)
    return ''.join(parts) or '&#160;'


def _table_grid(block):
    positions, occupied, columns = [], set(), 0
    for row, cells in enumerate(block.rows):
        column = 0
        for cell in cells:
            while (row, column) in occupied:
                column += 1
            span = min(cell.rowspan, len(block.rows) - row)
            positions.append((row, column, span, cell))
            for y in range(row, row + span):
                for x in range(column, column + cell.colspan):
                    occupied.add((y, x))
            column += cell.colspan
            columns = max(columns, column)
    if columns > 64:
        raise ValueError('The narrative table contains too many columns.')
    return positions, columns


def _table_widths(block, positions, count, available):
    requested = block.style.get('width')
    width = min(available, requested[1] * available if requested and requested[0] == 'fraction' else requested[1] if requested else available)
    widths = [None] * count
    for index, value in enumerate(block.columns[:count]):
        if value:
            widths[index] = width * value[1] if value[0] == 'fraction' else value[1]
    for _, column, _, cell in positions:
        requested = cell.style.get('width')
        if not requested:
            continue
        value = width * requested[1] if requested[0] == 'fraction' else requested[1]
        for index in range(column, column + cell.colspan):
            if widths[index] is None:
                widths[index] = value / cell.colspan
    allocated = sum(value or 0 for value in widths)
    missing = sum(value is None for value in widths)
    widths = [value if value is not None else max(width * .03, (width - allocated) / missing) for value in widths]
    # The printable area is authoritative even for pasted fixed-width tables.
    scale = width / sum(widths)
    return [value * scale for value in widths]


def _header_rows(block):
    count = 0
    for row in block.rows:
        if not all(cell.heading for cell in row):
            break
        count += 1
    # All-TH pasted tables have no distinct repeatable header section.
    return count if count < len(block.rows) else 0


class _RichTable(Table):
    def wrap(self, availWidth, availHeight):
        self.width, self.height = super().wrap(availWidth, availHeight)
        return self.width, self.height


def _cell_padding(cell):
    return {**dict.fromkeys(('top', 'right', 'bottom', 'left'), 6), **cell.style.get('cell_padding', {})}


def _row_heights(block):
    return [max([row_style.get('min_height', 0)] + [cell.style.get('min_height', 0) for cell in row if cell.rowspan == 1])
            for row, row_style in zip(block.rows, block.row_styles or [{} for _ in block.rows])]


def _pdf_tabbed_paragraph(block, width, paragraph_style):
    """Lay out Word tab fields at measured stops, keeping each field editable.

    ReportLab Paragraph has no tab-stop support. Borderless row segments retain
    styled runs and wrap the value column without approximating tabs with spaces.
    """
    width = max(12, width)
    rows, cells, runs = [], [], []
    start, cursor = 0, 0
    stops = block.style.get('tab_stops', [])
    interval = block.style.get('tab_interval', 36)

    def flush_cell():
        nonlocal runs
        cells.append((start, runs))
        runs = []

    def flush_row():
        nonlocal cells, start, cursor
        flush_cell()
        rows.append(cells)
        cells, start, cursor = [], 0, 0

    for run in block.runs:
        for part in re.split(r'(\t|\n)', run.text):
            if part == '\t':
                flush_cell()
                target = next((position for position in stops if position > cursor + .1), None)
                if target is None:
                    target = (int(cursor / interval) + 1) * interval
                start = cursor = min(max(0, width - 24), target)
            elif part == '\n':
                flush_row()
            elif part:
                runs.append(Run(part, run.style))
                cursor += stringWidth(part, _pdf_font(run.style), run.style.get('font_size', paragraph_style.fontSize))
    if runs or cells:
        flush_row()
    result = []
    for row_index, fields in enumerate(rows):
        # Consecutive tabs may make empty fields; their widths still carry the
        # original tab positions, while the final field gets remaining space.
        starts = [position for position, _ in fields]
        values, widths, pending = [], [], []
        for index, (position, content) in enumerate(fields):
            end = starts[index + 1] if index + 1 < len(fields) else width
            if end <= position:
                if content:
                    if values:
                        values[-1].extend(content)
                    else:
                        pending.extend(content)
                continue
            widths.append(end - position)
            values.append(pending + list(content))
            pending = []
        if not values:
            continue
        cell_style = ParagraphStyle('POTabField', parent=paragraph_style, leftIndent=0, rightIndent=0, firstLineIndent=0, spaceBefore=0, spaceAfter=0, alignment=TA_LEFT)
        rendered = [Paragraph(_pdf_markup(value), cell_style) for value in values]
        if paragraph_style.leftIndent:
            widths.insert(0, paragraph_style.leftIndent)
            marker = block.list_info['marker'] if block.list_info and row_index == 0 else ''
            rendered.insert(0, Paragraph(escape(marker), cell_style))
        result.append(_RichTable([rendered], colWidths=widths,
                                 hAlign='LEFT', splitByRow=1, splitInRow=1,
                                 spaceBefore=paragraph_style.spaceBefore if row_index == 0 else 0,
                                 spaceAfter=paragraph_style.spaceAfter if row_index == len(rows) - 1 else 0,
                                 style=TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 0), ('RIGHTPADDING', (0, 0), (-1, -1), 0), ('TOPPADDING', (0, 0), (-1, -1), 0), ('BOTTOMPADDING', (0, 0), (-1, -1), 0)])))
    return result


def pdf_rich_flowables(blocks, width, base_style):
    result = []
    for block in blocks:
        style = block.style
        if block.kind == 'break':
            result.append(PageBreak())
        elif block.kind == 'paragraph':
            size = style.get('font_size', base_style.fontSize)
            max_size = max([size] + [run.style.get('font_size', size) for run in block.runs])
            line = style.get('line_height', ('multiple', 1.3))
            leading = max_size * line[1] if line[0] == 'multiple' else line[1]
            indent = style.get('indent', 0)
            bullet = None
            if block.list_info:
                indent += 18 * (block.list_info['depth'] + 1)
                bullet = block.list_info['marker']
            paragraph_style = ParagraphStyle(
                'PORich', parent=base_style, fontName=_pdf_font(style), fontSize=size,
                leading=leading, alignment={'left': TA_LEFT, 'center': TA_CENTER, 'right': TA_RIGHT, 'justify': TA_JUSTIFY}.get(style.get('align'), TA_LEFT),
                spaceBefore=style.get('before', 0), spaceAfter=style.get('after', 4),
                leftIndent=indent, rightIndent=style.get('right_indent', 0),
                firstLineIndent=style.get('first_indent', 0), bulletIndent=max(0, indent - 14),
                bulletFontName=_pdf_font(style), bulletFontSize=size,
                textColor=colors.toColor(style.get('color', '#000000')),
                keepWithNext=bool(style.get('heading')), allowWidows=0, allowOrphans=0,
            )
            if any('\t' in run.text for run in block.runs):
                result.extend(_pdf_tabbed_paragraph(block, width - indent - style.get('right_indent', 0), paragraph_style))
            else:
                result.append(Paragraph(_pdf_markup(block.runs), paragraph_style, bulletText=bullet))
        elif block.kind == 'image':
            factor = min(width * style.get('width_fraction', 1) / block.width, 500 / block.height)
            if not style.get('width_fraction'):
                factor = min(1, factor)
            rendered = Image(BytesIO(block.image), width=block.width * factor, height=block.height * factor)
            rendered.hAlign = style.get('align', 'left').upper()
            result.extend((rendered, Spacer(1, 4)))
        elif block.kind == 'table' and block.rows:
            positions, columns = _table_grid(block)
            widths = _table_widths(block, positions, columns, width)
            grid = [['' for _ in range(columns)] for _ in block.rows]
            commands = [('GRID', (0, 0), (-1, -1), .5, colors.HexColor('#64748b')), ('VALIGN', (0, 0), (-1, -1), 'TOP')]
            for row, column, rowspan, cell in positions:
                padding = _cell_padding(cell)
                grid[row][column] = pdf_rich_flowables(cell.blocks, max(12, sum(widths[column:column + cell.colspan]) - padding['left'] - padding['right']), base_style) or [Paragraph('&#160;', base_style)]
                if cell.colspan > 1 or rowspan > 1:
                    commands.append(('SPAN', (column, row), (column + cell.colspan - 1, row + rowspan - 1)))
                if cell.style.get('background'):
                    commands.append(('BACKGROUND', (column, row), (column + cell.colspan - 1, row + rowspan - 1), colors.toColor(cell.style['background'])))
                for edge, value in padding.items():
                    commands.append((edge.upper() + 'PADDING', (column, row), (column + cell.colspan - 1, row + rowspan - 1), value))
                commands.append(('VALIGN', (column, row), (column + cell.colspan - 1, row + rowspan - 1), {'middle': 'MIDDLE', 'bottom': 'BOTTOM'}.get(cell.style.get('valign'), 'TOP')))
                if 'border_width' in cell.style or 'border_color' in cell.style:
                    commands.append(('BOX', (column, row), (column + cell.colspan - 1, row + rowspan - 1), cell.style.get('border_width', .5), colors.toColor(cell.style.get('border_color', '#64748b'))))
            result.extend((_RichTable(grid, colWidths=widths, minRowHeights=_row_heights(block), hAlign=style.get('align', 'left').upper(), repeatRows=_header_rows(block), splitByRow=1, splitInRow=1, style=TableStyle(commands)), Spacer(1, 6)))
    return result


def _docx_number(document, info, cache):
    key = info['sequence']
    if key in cache:
        return cache[key]
    numbering = document.part.numbering_part.element
    abstract_id = max([int(element.get(qn('w:abstractNumId'))) for element in numbering.findall(qn('w:abstractNum'))] + [-1]) + 1
    number_id = max([int(element.get(qn('w:numId'))) for element in numbering.findall(qn('w:num'))] + [0]) + 1
    abstract = OxmlElement('w:abstractNum')
    abstract.set(qn('w:abstractNumId'), str(abstract_id))
    level = OxmlElement('w:lvl')
    level.set(qn('w:ilvl'), '0')
    for tag, value in (('start', info['number'] if info['ordered'] else 1), ('numFmt', 'decimal' if info['ordered'] else 'bullet'), ('lvlText', '%1.' if info['ordered'] else '\u2022')):
        element = OxmlElement('w:' + tag)
        element.set(qn('w:val'), str(value))
        level.append(element)
    abstract.append(level)
    numbering.append(abstract)
    number = OxmlElement('w:num')
    number.set(qn('w:numId'), str(number_id))
    reference = OxmlElement('w:abstractNumId')
    reference.set(qn('w:val'), str(abstract_id))
    number.append(reference)
    numbering.append(number)
    cache[key] = number_id
    return number_id


def _docx_run(paragraph, run):
    style = run.style
    rendered = paragraph.add_run(run.text)
    rendered.bold, rendered.italic = bool(style.get('bold')), bool(style.get('italic'))
    rendered.underline, rendered.font.strike = bool(style.get('underline')), bool(style.get('strike'))
    rendered.font.name = style.get('font_family', 'Arial')
    rendered.font.size = Pt(style.get('font_size', 10.5))
    if style.get('color'):
        rendered.font.color.rgb = RGBColor.from_string(style['color'].lstrip('#'))
    if style.get('background'):
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), style['background'].lstrip('#'))
        rendered._r.get_or_add_rPr().append(shading)
    rendered.font.superscript = style.get('vertical') == 'sup'
    rendered.font.subscript = style.get('vertical') == 'sub'
    if style.get('href'):
        from docx.opc.constants import RELATIONSHIP_TYPE
        relationship = paragraph.part.relate_to(style['href'], RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        hyperlink = OxmlElement('w:hyperlink')
        hyperlink.set(qn('r:id'), relationship)
        hyperlink.append(rendered._r)
        paragraph._p.append(hyperlink)


def append_docx_rich_content(container, blocks, width=504, *, document=None, numbering=None):
    document = document or container
    numbering = {} if numbering is None else numbering
    for block in blocks:
        style = block.style
        if block.kind == 'break':
            container.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif block.kind == 'paragraph':
            paragraph = container.add_paragraph()
            if style.get('heading'):
                paragraph.style = f'Heading {style["heading"]}'
            paragraph.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'right': WD_ALIGN_PARAGRAPH.RIGHT, 'center': WD_ALIGN_PARAGRAPH.CENTER, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}.get(style.get('align'), WD_ALIGN_PARAGRAPH.LEFT)
            formatting = paragraph.paragraph_format
            formatting.space_before, formatting.space_after = Pt(style.get('before', 0)), Pt(style.get('after', 4))
            formatting.left_indent, formatting.right_indent = Pt(style.get('indent', 0)), Pt(style.get('right_indent', 0))
            formatting.first_line_indent = Pt(style.get('first_indent', 0))
            line = style.get('line_height', ('multiple', 1.3))
            formatting.line_spacing = line[1] if line[0] == 'multiple' else Pt(line[1])
            formatting.keep_with_next = bool(style.get('heading'))
            for stop in style.get('tab_stops', []):
                formatting.tab_stops.add_tab_stop(Pt(stop))
            if any('\t' in run.text for run in block.runs) and not style.get('tab_stops'):
                interval = style.get('tab_interval', 36)
                for multiple in range(1, int(width / interval) + 1):
                    formatting.tab_stops.add_tab_stop(Pt(interval * multiple))
            if block.list_info:
                info = block.list_info
                formatting.left_indent = Pt(style.get('indent', 0) + 18 * (info['depth'] + 1))
                formatting.first_line_indent = Pt(-14)
                properties = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                properties.get_or_add_ilvl().val = 0
                properties.get_or_add_numId().val = _docx_number(document, info, numbering)
            for run in block.runs:
                _docx_run(paragraph, run)
        elif block.kind == 'image':
            factor = min(width * style.get('width_fraction', 1) / block.width, 500 / block.height)
            if not style.get('width_fraction'):
                factor = min(1, factor)
            paragraph = container.add_paragraph()
            paragraph.alignment = {'center': WD_ALIGN_PARAGRAPH.CENTER, 'right': WD_ALIGN_PARAGRAPH.RIGHT}.get(style.get('align'), WD_ALIGN_PARAGRAPH.LEFT)
            paragraph.add_run().add_picture(BytesIO(block.image), width=Pt(block.width * factor), height=Pt(block.height * factor))
        elif block.kind == 'table' and block.rows:
            positions, columns = _table_grid(block)
            widths = _table_widths(block, positions, columns, width)
            if hasattr(container, 'sections'):
                table = container.add_table(rows=len(block.rows), cols=columns)
            else:
                table = container.add_table(rows=len(block.rows), cols=columns, width=Pt(width))
            table.style = 'Table Grid'
            table.autofit = False
            for column, column_width in zip(table.columns, widths):
                column.width = Pt(column_width)
            for row in table.rows:
                for cell, column_width in zip(row.cells, widths):
                    cell.width = Pt(column_width)
            for row, height in zip(table.rows, _row_heights(block)):
                if height:
                    row.height = Pt(height)
                    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            for row, column, rowspan, cell in positions:
                rendered = table.cell(row, column)
                if cell.colspan > 1 or rowspan > 1:
                    rendered = rendered.merge(table.cell(row + rowspan - 1, column + cell.colspan - 1))
                initial = rendered.paragraphs[0]._p
                padding = _cell_padding(cell)
                append_docx_rich_content(rendered, cell.blocks, max(12, sum(widths[column:column + cell.colspan]) - padding['left'] - padding['right']), document=document, numbering=numbering)
                rendered.vertical_alignment = {'middle': WD_CELL_VERTICAL_ALIGNMENT.CENTER, 'bottom': WD_CELL_VERTICAL_ALIGNMENT.BOTTOM}.get(cell.style.get('valign'), WD_CELL_VERTICAL_ALIGNMENT.TOP)
                margins = OxmlElement('w:tcMar')
                for edge, value in padding.items():
                    element = OxmlElement('w:' + edge)
                    element.set(qn('w:w'), str(round(value * 20)))
                    element.set(qn('w:type'), 'dxa')
                    margins.append(element)
                rendered._tc.get_or_add_tcPr().append(margins)
                if cell.blocks:
                    initial.getparent().remove(initial)
                if not rendered.paragraphs:
                    rendered.add_paragraph()
                if cell.style.get('background'):
                    shading = OxmlElement('w:shd')
                    shading.set(qn('w:fill'), cell.style['background'].lstrip('#'))
                    rendered._tc.get_or_add_tcPr().append(shading)
                if 'border_width' in cell.style or 'border_color' in cell.style:
                    borders = OxmlElement('w:tcBorders')
                    for edge in ('top', 'left', 'bottom', 'right'):
                        element = OxmlElement('w:' + edge)
                        element.set(qn('w:val'), 'single' if cell.style.get('border_width', .5) else 'nil')
                        element.set(qn('w:sz'), str(round(cell.style.get('border_width', .5) * 8)))
                        element.set(qn('w:color'), cell.style.get('border_color', '#64748b').lstrip('#'))
                        borders.append(element)
                    rendered._tc.get_or_add_tcPr().append(borders)
            for index in range(_header_rows(block)):
                header = OxmlElement('w:tblHeader')
                table.rows[index]._tr.get_or_add_trPr().append(header)
