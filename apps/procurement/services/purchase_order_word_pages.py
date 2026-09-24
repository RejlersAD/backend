"""Preserve canonical supporting PDF pages in the Purchase Order Word export.

Only already-authorized PDF bytes enter this helper. It never looks up records,
downloads files, or changes their contents. Supporting pages are visual evidence,
while the preceding generated Purchase Order body remains editable Word content.
"""

from __future__ import annotations

from io import BytesIO

import fitz
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


PDF_PAGE_DPI = 180
MAX_PAGE_IMAGE_EDGE = 3000
MAX_WORD_PAGE_POINTS = 22 * 72


def _page_paragraph(paragraph):
    """A floating page needs an anchor, not a page-sized line of text."""
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_before = Pt(0)
    paragraph_format.space_after = Pt(0)
    paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph_format.line_spacing = Pt(1)
    paragraph_format.keep_with_next = False
    paragraph_format.keep_together = False
    paragraph_format.page_break_before = False
    paragraph_format.widow_control = False
    # Formatting the paragraph mark also avoids inherited heading line heights.
    properties = paragraph._p.get_or_add_pPr()
    run_properties = properties.find(qn('w:rPr'))
    if run_properties is None:
        run_properties = OxmlElement('w:rPr')
        properties.append(run_properties)
    size = OxmlElement('w:sz')
    size.set(qn('w:val'), '2')
    run_properties.append(size)


def _has_body_content(document):
    for child in document._element.body:
        if child.tag == qn('w:sectPr'):
            continue
        if child.tag != qn('w:p'):
            return True
        if child.xpath('.//w:t | .//w:drawing | .//w:pict | .//w:br | .//w:sectPr'):
            return True
    return False


def _page_section(document, width, height, *, reuse_empty):
    if reuse_empty:
        section = document.sections[-1]
    else:
        section = document.add_section(WD_SECTION_START.NEW_PAGE)
        # add_section appends an empty paragraph containing the previous
        # section properties. A normal-height paragraph can create blank pages.
        _page_paragraph(document.paragraphs[-1])

    # Word cannot represent paper larger than 22 inches on either axis.
    # Preserve aspect ratio for oversized drawings; standard PDF paper sizes
    # retain their exact physical dimensions (within Word's twip precision).
    paper_scale = min(1.0, MAX_WORD_PAGE_POINTS / max(width, height))
    section.page_width = Pt(width * paper_scale)
    section.page_height = Pt(height * paper_scale)
    section.orientation = WD_ORIENT.LANDSCAPE if width > height else WD_ORIENT.PORTRAIT
    section.top_margin = section.bottom_margin = Pt(0)
    section.left_margin = section.right_margin = Pt(0)
    section.header_distance = section.footer_distance = Pt(0)
    section.gutter = Pt(0)
    section.different_first_page_header_footer = False
    for name in (
        'header', 'first_page_header', 'even_page_header',
        'footer', 'first_page_footer', 'even_page_footer',
    ):
        part = getattr(section, name)
        part.is_linked_to_previous = False
        # Covers already contain branding. Never overlay the editable body's
        # header, footer, signature, or page numbers on supporting evidence.
        for child in list(part._element):
            part._element.remove(child)
        part._element.append(OxmlElement('w:p'))
    return section


def _anchor_page_image(paragraph, content, width, height, page_number):
    run = paragraph.add_run()
    run.font.size = Pt(1)
    inline = run.add_picture(BytesIO(content), width=width, height=height)._inline
    anchor = OxmlElement('wp:anchor')
    for name, value in {
        'distT': '0', 'distB': '0', 'distL': '0', 'distR': '0',
        'simplePos': '0', 'relativeHeight': '0', 'behindDoc': '0',
        'locked': '0', 'layoutInCell': '1', 'allowOverlap': '1',
    }.items():
        anchor.set(name, value)
    simple_position = OxmlElement('wp:simplePos')
    simple_position.set('x', '0')
    simple_position.set('y', '0')
    anchor.append(simple_position)
    for axis in ('H', 'V'):
        position = OxmlElement(f'wp:position{axis}')
        position.set('relativeFrom', 'page')
        offset = OxmlElement('wp:posOffset')
        offset.text = '0'
        position.append(offset)
        anchor.append(position)
    anchor.append(inline.extent)
    anchor.append(OxmlElement('wp:wrapNone'))
    inline.docPr.set('descr', f'Original PDF page {page_number}')
    anchor.append(inline.docPr)
    frame = inline.find(qn('wp:cNvGraphicFramePr'))
    if frame is not None:
        anchor.append(frame)
    anchor.append(inline.graphic)
    inline.getparent().replace(inline, anchor)


def append_pdf_pages(document, pdf_bytes: bytes, start_page: int = 0) -> int:
    """Append canonical PDF pages, in order, without duplicating Word branding.

    ``start_page`` is zero based: skip the generated PDF body when keeping the
    editable Word body, or pass zero for an entirely visual document. Empty Word
    documents reuse their first page. Returns the number of pages appended.

    Images retain the PDF's visible crop and rotation at 180 dpi, bounded to a
    3000-pixel longest edge. Only one decoded image is held at a time; compressed
    PNG images remain in the resulting DOCX. No attachment links are followed.
    """
    if isinstance(start_page, bool) or not isinstance(start_page, int) or start_page < 0:
        raise ValueError('start_page must be a non-negative integer.')
    with fitz.open(stream=pdf_bytes, filetype='pdf') as source:
        if source.needs_pass:
            raise ValueError('The supporting PDF requires a password.')
        if start_page > len(source):
            raise ValueError('start_page exceeds the number of PDF pages.')
        appended = 0
        reuse_empty = not _has_body_content(document)
        for page_index in range(start_page, len(source)):
            page = source[page_index]
            width, height = page.rect.width, page.rect.height
            scale = min(PDF_PAGE_DPI / 72, MAX_PAGE_IMAGE_EDGE / max(width, height))
            pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
            content = pixmap.tobytes('png')
            del pixmap
            section = _page_section(document, width, height, reuse_empty=reuse_empty)
            reuse_empty = False
            paragraph = document.add_paragraph()
            _page_paragraph(paragraph)
            _anchor_page_image(
                paragraph, content, section.page_width, section.page_height,
                page_index + 1,
            )
            del content
            appended += 1
        return appended
