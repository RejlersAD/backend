"""Editable requisition documents from saved data; never an approval command."""

from datetime import datetime
from decimal import InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.exceptions import InvalidImageStreamError, UnexpectedEndOfFileError, UnrecognizedImageError
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .approval_integrity import stage_signature_issue
from .procurement_vat import CENT, decimal_amount
from .purchase_order_approval_artwork import approval_image_stream


DOCX_MIME_TYPE = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
LOGO_PATH = Path(__file__).resolve().parent.parent / 'assets' / 'rejlers-pr-po-logo.png'


def _text(value):
    # Imported text can contain XML control characters which Word cannot store.
    text = '' if value is None else str(value)
    return ''.join(char for char in text if char in '\t\n\r' or 32 <= ord(char) <= 0xD7FF
                   or 0xE000 <= ord(char) <= 0xFFFD or 0x10000 <= ord(char) <= 0x10FFFF).strip() or '—'


def _first(*values):
    return next((value for value in values if value is not None and value != ''), None)


def _money(value, currency):
    amount = decimal_amount(value)
    if amount is None:
        return '—'
    try:
        return f'{_text(currency)} {amount.quantize(CENT, rounding=ROUND_HALF_UP):,.2f}'
    except InvalidOperation:
        return '—'


def _timestamp(value):
    if not value:
        return '—'
    try:
        parsed = datetime.fromisoformat(str(value).replace('Z', '+00:00'))
        if parsed.tzinfo is not None:
            return parsed.astimezone(ZoneInfo('Asia/Dubai')).strftime('%d.%m.%Y %H:%M:%S GST')
    except ValueError:
        pass
    return _text(value)


def _icv(record, metadata):
    value = metadata.get('icv')
    if value in (None, ''):
        vendor_id = str(record.get('vendor') or '')
        supplier = str(record.get('preferred_supplier_if_any') or record.get('supplier_name') or '').strip().casefold()

        def matches(row):
            if vendor_id:
                return str(row.get('vendor_id') or row.get('id') or '') == vendor_id
            return bool(supplier and str(row.get('vendor_name') or row.get('name') or '').strip().casefold() == supplier)

        shortlist = record.get('selected_vendors') or []
        selected = next((row for row in shortlist if isinstance(row, dict) and matches(row)), {})
        details = record.get('vendor_details') or {}
        if matches(details):
            selected = {**selected, **details}
            if 'icv_percentage' in details:
                selected.pop('icv_value', None)
        value = _first(selected.get('icv_percentage'), selected.get('icv_value'))
    number = decimal_amount(str(value).rstrip('%').strip())
    return f'{number:g}%' if number is not None and 0 <= number <= 100 else value


def _signature_state(stage):
    """Match the existing preview's identity evidence, without profile lookups."""
    if stage_signature_issue(stage) or stage.get('signature_review_required'):
        return False, 'Signature needs review'
    if str(stage.get('status', '')).lower() != 'approved':
        return False, 'Not recorded'
    if stage.get('external') or stage.get('evidence_document_id'):
        verified = stage.get('signature_verified') is True
        return verified, 'Source evidence recorded' if verified else 'See original source evidence'
    assigned_email = str(stage.get('user_email') or stage.get('approver_email') or '').strip().lower()
    assigned_id = str(stage.get('user_id') or stage.get('approver_id') or '')
    for prefix in ('approved_by', 'decided_by', 'signature_user'):
        email = str(stage.get(f'{prefix}_email') or '').strip().lower()
        identifier = str(stage.get(f'{prefix}_id') or '')
        matches = assigned_email == email if assigned_email and email else bool(assigned_id and assigned_id == identifier)
        if matches:
            return True, 'Recorded; see RADAI evidence'
    return False, 'Signer not verified'


def _approval_rows(record):
    workflow = record.get('approval_workflow_config') or record.get('approval_hierarchy')
    if isinstance(workflow, list) and workflow:
        return [row for row in workflow if isinstance(row, dict)]
    # Legacy fixed fields lack the actor metadata needed to verify an image.
    rows = [dict(role=label, user_name=record.get(f'{prefix}_name_display'),
                 status=record.get(f'{prefix}_approval_status', 'pending'),
                 approved_at=record.get(f'{prefix}_approved_at'))
            for prefix, label in [('pm', 'PM'), ('eng_manager', 'MoE'),
                                  ('manager_projects', 'MoP'), ('vp_op', 'VP')]]
    if record.get('status') == 'converted':
        for row in rows:
            if row['status'] in ('', 'pending', 'in_review', None):
                row['status'] = 'not_recorded'
    return rows


class _CompanyForm:
    """One continuous editable grid, using the existing company PR proportions."""

    # Union of column boundaries from the header, prices, approvals and final row.
    STOPS = (0, 1000, 3333, 3548, 3700, 4878, 5700, 6667, 6829, 7400, 10000)
    FULL = (0, 10000)
    TWO = (0, 6667, 10000)
    THREE = (0, 3333, 6667, 10000)
    PRICE = (0, 4878, 6829, 10000)
    APPROVAL = (0, 1000, 3700, 5700, 7400, 10000)

    def __init__(self, document):
        self.table = document.add_table(rows=0, cols=len(self.STOPS) - 1)
        self.table.autofit = False
        self.table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for index, column in enumerate(self.table.columns):
            column.width = Mm(198 * (self.STOPS[index + 1] - self.STOPS[index]) / 10000)
        props = self.table._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            border = OxmlElement(f'w:{edge}')
            for name, value in {'val': 'single', 'sz': '8' if edge in ('top', 'left', 'bottom', 'right') else '4', 'color': '374151'}.items():
                border.set(qn(f'w:{name}'), value)
            borders.append(border)
        props.append(borders)
        margins = OxmlElement('w:tblCellMar')
        for edge, millimetres in [('top', 1.7), ('bottom', 1.7), ('left', 2.3), ('right', 2.3)]:
            margin = OxmlElement(f'w:{edge}')
            margin.set(qn('w:w'), str(int(Mm(millimetres).twips)))
            margin.set(qn('w:type'), 'dxa')
            margins.append(margin)
        props.append(margins)

    def row(self, boundaries=FULL, height=9, *, repeat=False):
        row = self.table.add_row()
        # Word adds the top/bottom cell margins to its minimum row height.
        row.height, row.height_rule = Mm(max(0, height - 3.4)), WD_ROW_HEIGHT_RULE.AT_LEAST
        if repeat:
            row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        original = row.cells
        cells = []
        for start, end in zip(boundaries, boundaries[1:]):
            left, right = self.STOPS.index(start), self.STOPS.index(end)
            cell = original[left]
            if right - left > 1:
                cell = cell.merge(original[right - 1])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells.append(cell)
        return cells


def _write(cell, value, *, bold=False, align=None, size=None):
    cell.text = ''
    paragraph = cell.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(_text(value) if value != '' else '')
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    return paragraph


def _label(cell, label, value):
    paragraph = _write(cell, '')
    paragraph.add_run(label + ' ').bold = True
    paragraph.add_run(_text(value))
    return paragraph


def _section(cell, label, value):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    heading = _write(cell, label, bold=True)
    heading.paragraph_format.space_after = Mm(5.5)
    cell.add_paragraph(_text(value))


def _role_label(stage, index, level_one_index, labels):
    saved = stage.get('approval_label') or labels.get(str(stage.get('user_id') or stage.get('approver_id') or ''))
    if saved:
        return saved
    role = f"{stage.get('role', '')} {stage.get('stage', '')}".lower()
    if 'procurement' in role:
        return 'L0- PRO'
    if 'general manager' in role or 'ceo' in role:
        return 'CEO'
    if 'engineering' in role:
        return 'L2 MoE'
    if 'manager of projects' in role or 'projects manager' in role:
        return 'L3 MoP'
    if any(value in role for value in ('vice president', 'vp delivery', 'vp operations')):
        return 'L4 VOP/VP'
    if 'level 1' in role:
        return f'L1-{level_one_index or index + 1}'
    return stage.get('role') or f'L{stage.get("level") or index + 1}'


def build_purchase_requisition_docx(record):
    """Render the company PR form as editable Word cells, without remote reads."""
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.left_margin = section.right_margin = Mm(6)
    section.bottom_margin, section.footer_distance = Mm(12), Mm(5)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(9)
    normal.font.color.rgb = RGBColor.from_string('1F2937')
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.15
    document.core_properties.title = f'Purchase Requisition {_text(record.get("pr_number"))}'
    document.core_properties.author = 'RADAI'
    metadata = record.get('price_remarks_data')
    metadata = metadata if isinstance(metadata, dict) else {}
    form = _CompanyForm(document)

    title, logo = form.row(form.TWO, 15.3, repeat=True)
    _write(title, 'Purchase Requisition', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=18)
    paragraph = _write(logo, '', align=WD_ALIGN_PARAGRAPH.CENTER)
    if LOGO_PATH.exists():
        paragraph.add_run().add_picture(str(LOGO_PATH), width=Mm(32))
    else:
        paragraph.add_run('REJLERS')
    issuer, number, issued_date = form.row(form.THREE, 9.2)
    _label(issuer, 'Issued by:', record.get('issued_by_name'))
    _label(number, 'PR No.', record.get('pr_number'))
    date = str(record.get('issued_date') or '')
    try:
        date = datetime.strptime(date[:10], '%Y-%m-%d').strftime('%d.%m.%Y')
    except ValueError:
        pass
    _label(issued_date, 'Date:', date)
    product, supplier = form.row(form.TWO, 15.5)
    _label(product, 'Product/ Service:', record.get('product_service') or record.get('title'))
    _label(supplier, 'Supplier:', record.get('supplier_name'))
    project, icv = form.row(form.TWO, 24.5)
    _label(project, 'Project/Department:', record.get('project_department') or record.get('project'))
    _label(icv, 'ICV:', _icv(record, metadata))
    _section(form.row(height=24.7)[0], '1. Description and Reason for Purchase:', record.get('description_reason'))
    _label(form.row(height=11.3)[0], '2. Preferred Supplier (if any):',
           record.get('preferred_supplier_if_any') or record.get('supplier_name'))

    description, price, remarks = form.row(form.PRICE)
    _write(description, '3. Price', bold=True)
    _write(price, 'Price', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _write(remarks, 'Remarks', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    currency = record.get('currency')
    total = _first(record.get('total_price'), record.get('net_total_excl_vat'))
    items = record.get('items') or metadata.get('price_lines') or []
    items = [row for row in items if isinstance(row, dict)] if isinstance(items, list) else []
    budget = _first(metadata.get('budget_in_aed'), record.get('estimated_budget'))
    budget_currency = 'AED' if metadata.get('budget_in_aed') not in (None, '') else currency
    budget_text = f'Budget → {_money(budget, budget_currency)}' if budget is not None else ''
    if not items:
        items = [{'description': record.get('price_description'), 'total': total}]
    for index, item in enumerate(items):
        description, price, remarks = form.row(form.PRICE, 9.2)
        _write(description, item.get('description') or item.get('name'))
        _write(price, _money(_first(item.get('total'), item.get('total_price'), item.get('amount'), item.get('line_total')),
                             item.get('currency') or currency), align=WD_ALIGN_PARAGRAPH.RIGHT)
        remark = item.get('remarks') or (budget_text or record.get('price_remarks') if index == 0 else '')
        _write(remarks, remark)
    label, amount, equivalent = form.row(form.PRICE, 9.2)
    _write(label, 'Total', bold=True)
    _write(amount, _money(total, currency), bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    # Keep captured AED equivalence; do not introduce a different backend FX rate.
    total_amount = decimal_amount(total)
    aed = total if str(currency).upper() == 'AED' else (
        metadata.get('net_total_aed') if total_amount is not None
        and total_amount == decimal_amount(record.get('net_total_excl_vat')) else None
    )
    _write(equivalent, _money(aed, 'AED'), bold=True)
    _label(form.row()[0], 'Negotiation Remarks:', record.get('price_remarks') or metadata.get('negotiation_remarks'))
    _label(form.row()[0], 'PO Reference:', record.get('po_number_reference'))
    recommendation = form.row(height=40.2)[0]
    _section(recommendation, 'Purchase recommendation', record.get('purchase_recommendation') or record.get('notes'))
    if metadata.get('attachment_reference'):
        reference = recommendation.add_paragraph(_text(metadata['attachment_reference']))
        reference.paragraph_format.space_before = Mm(3)
        reference.runs[0].underline = True

    _write(form.row(height=6.8)[0], 'APPROVALS', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for cell, label in zip(form.row(form.APPROVAL, 6.8), ['', 'Name', 'Signature', 'Status', 'Approval Timestamp']):
        _write(cell, label, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    labels = metadata.get('approval_table_labels') or {}
    level_one_index = 0
    for index, stage in enumerate(_approval_rows(record)):
        if str(stage.get('level')) == '1' or 'level 1' in str(stage.get('role') or '').lower():
            level_one_index += 1
        role, name, signature, status, timestamp = form.row(form.APPROVAL, 10.2)
        _write(role, _role_label(stage, index, level_one_index, labels), bold=True)
        _write(name, stage.get('user_name') or stage.get('approver_name') or stage.get('approver'))
        stage_status = str(stage.get('status') or 'pending').lower()
        _write(status, 'Approved' if stage_status == 'approved' else 'Not recorded' if stage_status == 'not_recorded' else stage_status)
        _write(timestamp, _timestamp(stage.get('approved_at') or stage.get('decided_at')) if stage_status == 'approved' else '—', size=8)
        verified, evidence = _signature_state(stage)
        if stage_status in ('pending', 'in_review') and evidence == 'Not recorded':
            evidence = 'Pending'
        stream = approval_image_stream(stage.get('signature')) if verified else None
        paragraph = _write(signature, evidence, align=WD_ALIGN_PARAGRAPH.CENTER, size=8)
        if stream:
            try:
                paragraph.clear()
                paragraph.add_run().add_picture(stream, width=Mm(22))
            except (UnrecognizedImageError, InvalidImageStreamError, UnexpectedEndOfFileError, ValueError):
                _write(signature, evidence, align=WD_ALIGN_PARAGRAPH.CENTER, size=8)
    label, value = form.row((0, 3548, 10000), 9.2)
    _write(label, 'Final Approval Timestamp', bold=True)
    _write(value, _timestamp(record.get('approved_at') or record.get('vp_op_approved_at') or metadata.get('signed_approval_date')))

    # Status stays visible without adding a non-company section to the form.
    footer = section.footer.paragraphs[0]
    footer.add_run(f'{_text(record.get("form_reference"))}  |  Status: {_text(record.get("status_display") or record.get("status"))}  |  Page ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(6.5)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
