"""Deterministic POC workbook reader. Reads cached values; never evaluates formulas.

read_workbook(path_or_bytes) returns metadata plus ``rows`` suitable for
PortfolioRow construction. Dates/decimals in model fields retain Python types;
provenance/history under ``extra`` and reconciliation are JSON-safe. The current
period is chosen by the dated revenue triplets, never by a hard-coded column.
"""
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import hashlib
from io import BytesIO
from pathlib import Path
import zipfile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

PARSER_VERSION = '2.1'
HEADER_ROW = 7
MAX_BYTES = 25 * 1024 * 1024
MISSING = {'', '-', 'n/a', 'n.a', 'n.a.', 'na', 'nil'}
TEXT_HEADERS = {
    'project_code': 'PROJECT ID', 'subproject_code': 'PROJECT SUB ID',
    'title': 'PROJECT TITLE', 'pm': 'PM', 'pc': 'PC Name',
    'business_unit': 'BU', 'client': 'CLIENT', 'scope_type': 'SCOPE TYPE',
}
VALUE_HEADERS = {
    'contract_value_aed': 'Modified Contract Value (AED)',
    'backlog_without_pt_aed': 'Backlog Without PT', 'backlog_with_pt_aed': 'Backlog With PT',
    'eddr_pct': 'PROGRESS EDDR', 'target_margin_pct': 'Target GM', 'forecast_margin_pct': 'FC GM',
    'ld_exposure_aed': 'LD Exposure Risk (AED)', 'prolongation_cost_aed': 'Total Prolongation cost (AED)',
}
DATE_HEADERS = {'start_date': 'Contractual Start', 'contractual_finish': 'Contractual Finish',
                'forecast_finish': 'Forecast completion'}
REVENUE_HEADERS = ['Cumm RR - To-date (AED)', 'POC / EV - To-date (%)', 'RR For the period (AED)']


def _text(value):
    if value is None:
        return ''
    if isinstance(value, (int, float)) and not isinstance(value, bool) and float(value).is_integer():
        return str(int(value))
    return str(value).strip()


def _normal(value):
    return ' '.join(_text(value).casefold().split())


def _json_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value


def _warning(warnings, sheet, row, col, field, code):
    warnings.append({'sheet': sheet, 'cell': f'{get_column_letter(col)}{row}', 'field': field, 'code': code})


def _number(cell, formula, *, sheet, row, col, field, warnings, percent=False):
    value = cell.value
    if cell.data_type == 'e':
        _warning(warnings, sheet, row, col, field, 'excel_error')
        return None
    if _normal(value) in MISSING:
        if formula.data_type == 'f' and value is None:
            _warning(warnings, sheet, row, col, field, 'missing_formula_cache')
        return None
    try:
        if isinstance(value, (bool, date, datetime)):
            raise ValueError
        # Native Excel percentages are fractions. Explicit text percentages
        # already carry the 0..100 unit and must not be multiplied twice.
        explicit_percent = isinstance(value, str) and value.strip().endswith('%')
        number = Decimal(str(value).strip().rstrip('%').replace(',', ''))
        if not number.is_finite():
            raise ValueError
        if percent and not explicit_percent:
            number *= 100
        # Match persisted precision before range validation: cached Excel
        # arithmetic can produce 100.00000000000022 for an exact 100%.
        number = number.quantize(Decimal('0.00000001'))
        if percent and (number > 100 or (number < 0 and 'margin' not in field)):
            raise ValueError
        if abs(number) >= Decimal('1e12' if percent else '1e20'):
            raise ValueError
        return number
    except (InvalidOperation, TypeError, ValueError):
        _warning(warnings, sheet, row, col, field, 'invalid_percentage' if percent else 'invalid_number')
        return None


def _date_value(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _columns(sheet):
    headers = {}
    for col, cell in enumerate(sheet[HEADER_ROW], 1):
        if cell.value is not None:
            headers.setdefault(_normal(cell.value), []).append(col)
    return headers


def _single(headers, name):
    matches = headers.get(_normal(name), [])
    if len(matches) != 1:
        raise ValueError(f'POC header {name!r} must occur exactly once in row {HEADER_ROW}.')
    return matches[0]


def _forecast(workbook, formulas, rows, warnings):
    name = 'Delivery Rolling Forecast'
    if name not in workbook.sheetnames:
        warnings.append({'sheet': name, 'code': 'optional_sheet_missing'})
        return {'matched': 0, 'unmatched': [], 'review_rows': []}
    sheet, source = workbook[name], formulas[name]
    headers = _columns(sheet)
    mapping = {key: _single(headers, label) for key, label in (
        ('project_code', 'Project ID'), ('subproject_code', 'Project Sub ID'), ('title', 'Project Title'))}
    identities = {(r['project_code'], r['subproject_code']): r for r in rows}
    # Section labels define metric units. Monthly columns move as the workbook
    # grows; joining on identifiers avoids the sheet's auxiliary row blocks.
    anchors = []
    sections = {'revenue (cumulative %)': 'poc_pct', 'revenue per month (aed)': 'revenue_aed',
                'manhours': 'manhours'}
    for n in (3, 4):
        for col, cell in enumerate(sheet[n], 1):
            if _normal(cell.value) in sections:
                anchors.append((col, sections[_normal(cell.value)]))
    anchors.sort()
    if not any(metric == 'revenue_aed' for _, metric in anchors):
        raise ValueError('Delivery Rolling Forecast is missing its revenue-per-month section header.')
    period_columns = []
    for index, (start, metric) in enumerate(anchors):
        end = anchors[index + 1][0] if index + 1 < len(anchors) else sheet.max_column + 1
        for col in range(start, end):
            period = _date_value(sheet.cell(HEADER_ROW, col).value)
            # Later sections (EDDR / overclaim) have their own meaning; the
            # manhours block ends at the first nonempty non-date header.
            if metric == 'manhours' and col > start and sheet.cell(HEADER_ROW, col).value and not period:
                break
            if period:
                period_columns.append((col, metric, period))
    if len({(metric, period) for _, metric, period in period_columns}) != len(period_columns):
        raise ValueError('Delivery Rolling Forecast contains duplicate dated columns for a metric.')
    matched, unmatched, review, seen = 0, [], [], set()
    for number in range(HEADER_ROW + 1, sheet.max_row + 1):
        code = _text(sheet.cell(number, mapping['project_code']).value)
        sub = _text(sheet.cell(number, mapping['subproject_code']).value)
        title = _text(sheet.cell(number, mapping['title']).value)
        if not code and not title:
            continue  # Auxiliary identity-only lookup block, not a forecast row.
        if not code or not sub or not title:
            review.append({'source_row': number, 'project_code': code, 'subproject_code': sub,
                           'title': title, 'reason': 'incomplete_forecast_identity'})
            continue
        identity = (code, sub)
        if identity in seen:
            raise ValueError(f'Duplicate forecast project/subproject identity at row {number}.')
        seen.add(identity)
        target = identities.get(identity)
        if target is None:
            unmatched.append({'source_row': number, 'project_code': code, 'subproject_code': sub, 'title': title})
            continue
        forecast = []
        warning_start = len(warnings)
        for col, metric, period in period_columns:
            value = _number(sheet.cell(number, col), source.cell(number, col), sheet=name, row=number,
                            col=col, field=metric, warnings=warnings, percent=metric.endswith('_pct'))
            forecast.append({'date': period.isoformat(), 'metric': metric, 'value': _json_value(value),
                             'cell': f'{get_column_letter(col)}{number}'})
        target['extra']['rolling_forecast'] = forecast
        periods = {}
        for observation in forecast:
            period = periods.setdefault(observation['date'], {'period': observation['date'], 'source_cells': {}})
            period[observation['metric']] = observation['value']
            period['source_cells'][observation['metric']] = observation['cell']
        target['extra']['forecasts'] = sorted(periods.values(), key=lambda item: item['period'])
        target['extra']['forecast_source_row'] = number
        target['extra']['warnings'].extend(warnings[warning_start:])
        matched += 1
    return {'matched': matched, 'unmatched': unmatched, 'review_rows': review}


def read_workbook(path_or_bytes):
    path = None if isinstance(path_or_bytes, (bytes, bytearray)) else Path(path_or_bytes)
    content = bytes(path_or_bytes) if path is None else path.read_bytes()
    if not content or len(content) > MAX_BYTES:
        raise ValueError('Portfolio workbook is empty or exceeds 25 MB.')
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 150 * 1024 * 1024:
                raise ValueError('Expanded portfolio workbook exceeds 150 MB.')
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=False, keep_links=False)
        formulas = load_workbook(BytesIO(content), data_only=False, read_only=False, keep_links=False)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise ValueError('A valid .xlsx portfolio workbook is required.') from exc
    try:
        if 'POC' not in workbook.sheetnames:
            raise ValueError('Workbook must contain the POC source sheet.')
        for name in ('POC', 'Delivery Rolling Forecast'):
            if name in workbook.sheetnames and (workbook[name].max_row > 10000 or workbook[name].max_column > 500):
                raise ValueError(f'{name} exceeds the supported worksheet dimensions.')
        sheet, source = workbook['POC'], formulas['POC']
        headers = _columns(sheet)
        columns = {key: _single(headers, label) for key, label in
                   {**TEXT_HEADERS, **VALUE_HEADERS, **DATE_HEADERS,
                    'include_without_pt': 'WITHOUT PT', 'include_with_pt': 'WITH PT'}.items()}
        # In this template contract currency is the second of two identically
        # named columns; reject changed shape instead of guessing.
        currency_columns = headers.get(_normal('Original Contract Value (Cont. Curr.)'), [])
        if len(currency_columns) != 2:
            raise ValueError('Expected amount and currency columns for Original Contract Value.')
        columns['currency'] = currency_columns[1]
        groups = []
        for col in headers.get(_normal(REVENUE_HEADERS[0]), []):
            if [_normal(sheet.cell(HEADER_ROW, col + n).value) for n in range(3)] != list(map(_normal, REVENUE_HEADERS)):
                raise ValueError(f'Invalid revenue triplet at {get_column_letter(col)}7.')
            period = _date_value(sheet.cell(5, col).value)
            if period is None:
                raise ValueError(f'Missing reporting date at {get_column_letter(col)}5.')
            groups.append((period, col))
        if not groups or len({period for period, _ in groups}) != len(groups):
            raise ValueError('Revenue history requires distinct dated column groups.')
        groups.sort()
        reporting_date, current_col = groups[-1]
        overclaim = headers.get(_normal('Overclaim (POC & EDDR)'), [])
        current_overclaim = [col for col in overclaim if _date_value(sheet.cell(4, col).value) is None]
        if len(current_overclaim) != 1:
            raise ValueError('Expected one current Overclaim (POC & EDDR) column.')
        columns['overclaim_aed'] = current_overclaim[0]
        columns.update(recognized_revenue_aed=current_col, poc_pct=current_col + 1, period_revenue_aed=current_col + 2)
        rows, warnings, seen = [], [], set()
        resource_section = False
        for number in range(HEADER_ROW + 1, sheet.max_row + 1):
            code = _text(sheet.cell(number, columns['project_code']).value)
            sub = _text(sheet.cell(number, columns['subproject_code']).value)
            if not code and not sub:
                if _normal(sheet.cell(number, current_col).value) == 'booked mhr (period)':
                    for _, col in groups:
                        if [_normal(sheet.cell(number, col + offset).value) for offset in range(3)] != [
                                'booked mhr (period)', _normal(REVENUE_HEADERS[0]), _normal(REVENUE_HEADERS[2])]:
                            raise ValueError(f'Inconsistent resource revenue headers at row {number}.')
                    resource_section = True
                    continue
                title = _normal(sheet.cell(number, columns['title']).value)
                summary_caption = title.rstrip(':').strip() in {'total', 'subtotal', 'grand total'}
                project_descriptors = ('pm', 'pc', 'business_unit', 'client', 'scope_type',
                                       'include_without_pt', 'include_with_pt')
                if summary_caption and not any(
                    _text(sheet.cell(number, columns[key]).value)
                    or source.cell(number, columns[key]).data_type == 'f'
                    for key in project_descriptors
                ):
                    continue
                monetary_columns = {col for key, col in columns.items() if key.endswith('_aed')}
                monetary_columns.update(col + offset for _, col in groups
                                        for offset in ((1, 2) if resource_section else (0, 2)))
                has_financial_data = any(
                    _text(sheet.cell(number, col).value) or source.cell(number, col).data_type == 'f'
                    for col in monetary_columns
                )
                if title or source.cell(number, columns['title']).data_type == 'f' or has_financial_data:
                    raise ValueError(f'POC row {number} requires both project and subproject identifiers.')
                continue
            if not code or not sub:
                raise ValueError(f'POC row {number} requires both project and subproject identifiers.')
            if any(sheet.cell(number, columns[key]).data_type == 'e' for key in ('project_code', 'subproject_code')):
                raise ValueError(f'POC row {number} contains an invalid project identifier.')
            if (code, sub) in seen:
                raise ValueError(f'Duplicate POC project/subproject identity at row {number}: {code}/{sub}.')
            seen.add((code, sub))
            row = {key: _text(sheet.cell(number, columns[key]).value) for key in TEXT_HEADERS}
            if not row['title']:
                raise ValueError(f'POC row {number} has no project title.')
            for key, limit in {'project_code': 128, 'subproject_code': 128, 'pm': 128, 'pc': 128,
                               'business_unit': 128, 'client': 255, 'scope_type': 128}.items():
                if len(row[key]) > limit:
                    raise ValueError(f'POC row {number} exceeds the {key} length limit.')
            start_warning = len(warnings)
            row_columns = {**columns}
            if resource_section:
                row_columns['recognized_revenue_aed'] = current_col + 1
            for key in [*VALUE_HEADERS, 'overclaim_aed', 'recognized_revenue_aed', 'period_revenue_aed', 'poc_pct']:
                if resource_section and key in ('poc_pct', 'eddr_pct', 'overclaim_aed'):
                    row[key] = None
                    row_columns.pop(key, None)
                    continue
                col = row_columns[key]
                row[key] = _number(sheet.cell(number, col), source.cell(number, col), sheet='POC', row=number,
                                   col=col, field=key, warnings=warnings, percent=key.endswith('_pct'))
            for key in DATE_HEADERS:
                col = columns[key]
                cell = sheet.cell(number, col)
                row[key] = _date_value(cell.value)
                if row[key] is None and (_normal(cell.value) not in MISSING or source.cell(number, col).data_type == 'f'):
                    _warning(warnings, 'POC', number, col, key, 'invalid_or_missing_date')
            if row['start_date'] and row['contractual_finish'] and row['contractual_finish'] < row['start_date']:
                _warning(warnings, 'POC', number, columns['contractual_finish'], 'contractual_finish', 'finish_before_start')
            for key in ('include_without_pt', 'include_with_pt'):
                value = _normal(sheet.cell(number, columns[key]).value)
                row[key] = {'yes': True, 'no': False}.get(value)
                if row[key] is None:
                    _warning(warnings, 'POC', number, columns[key], key, 'unknown_inclusion_flag')
            currency = _text(sheet.cell(number, columns['currency']).value).upper()
            row['currency'] = {'EURO': 'EUR'}.get(currency, currency)
            if row['currency'] not in {'AED', 'USD', 'EUR', 'GBP', 'SAR', 'QAR', 'KWD', 'BHD', 'OMR'}:
                _warning(warnings, 'POC', number, columns['currency'], 'currency', 'unknown_currency')
                row['currency'] = ''
            history = []
            for period, col in groups:
                values = {}
                metrics = ('booked_manhours', 'recognized_revenue_aed', 'period_revenue_aed') if resource_section else (
                    'recognized_revenue_aed', 'poc_pct', 'period_revenue_aed')
                for offset, key in enumerate(metrics):
                    values[key] = _json_value(_number(sheet.cell(number, col + offset), source.cell(number, col + offset),
                        sheet='POC', row=number, col=col + offset, field=key, warnings=warnings, percent=key.endswith('_pct')))
                history.append({'date': period.isoformat(), **values, 'cell': f'{get_column_letter(col)}{number}'})
            row['source_row'] = number
            row['extra'] = {'sheet': 'POC', 'history': history, 'warnings': warnings[start_warning:],
                            'source_cells': {key: f'{get_column_letter(col)}{number}' for key, col in row_columns.items()},
                            'raw_values': {key: _json_value(sheet.cell(number, col).value) for key, col in row_columns.items()},
                            'rolling_forecast': [], 'forecasts': []}
            if resource_section:
                row['extra']['section'] = 'resource_deputation'
                row['extra']['booked_manhours'] = history[-1].get('booked_manhours')
                row['extra']['sold_rate_per_hour'] = _json_value(sheet.cell(number, columns['eddr_pct']).value)
                row['extra']['previous_overclaim_raw'] = _json_value(sheet.cell(number, columns['overclaim_aed']).value)
            rows.append(row)
        if not rows:
            raise ValueError('POC sheet contains no project rows.')
        forecast = _forecast(workbook, formulas, rows, warnings)
        reconciliation = {'project_count': len({row['project_code'] for row in rows}),
            'subproject_count': len(rows), 'forecast': forecast,
            'scope_counts': dict(Counter(row['scope_type'] for row in rows)),
            'included_without_pt_count': sum(row['include_without_pt'] is True for row in rows),
            'included_with_pt_count': sum(row['include_with_pt'] is True for row in rows),
            'reporting_periods': [period.isoformat() for period, _ in groups]}
        reconciliation['source_summary_cells'] = [
            {'field': key, 'cell': f'{get_column_letter(columns[key])}{number}',
             'value': _json_value(sheet.cell(number, columns[key]).value),
             'formula': source.cell(number, columns[key]).value if source.cell(number, columns[key]).data_type == 'f' else None}
            for key in ('contract_value_aed', 'recognized_revenue_aed', 'period_revenue_aed',
                        'backlog_without_pt_aed', 'backlog_with_pt_aed', 'overclaim_aed')
            for number in (5, 6)
        ]
        from .supplemental import read_supplemental
        reconciliation.update(read_supplemental(workbook, formulas, rows, reporting_date, warnings))
        if path is not None and path.read_bytes() != content:
            raise ValueError('Workbook changed while it was being read.')
        return {'sha256': hashlib.sha256(content).hexdigest(), 'parser_version': PARSER_VERSION,
                'file_name': path.name if path is not None else 'portfolio.xlsx', 'reporting_date': reporting_date,
                'row_count': len(rows), 'warnings': warnings, 'reconciliation': reconciliation, 'rows': rows}
    finally:
        workbook.close()
        formulas.close()
