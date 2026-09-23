from datetime import date
from decimal import Decimal
from io import BytesIO
import json
from unittest import TestCase
from xml.etree import ElementTree
import zipfile

from openpyxl.utils import get_column_letter

from apps.portfolio.supplemental import EXECUTIVE_HEADERS, REPORT_METRICS
from apps.portfolio.workbook import read_workbook
from .test_workbook import workbook_bytes


def cached_values(content, sheet_name, values):
    """Provide saved Excel formula caches without evaluating any formula."""
    namespace = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(BytesIO(content)) as source:
        workbook = ElementTree.fromstring(source.read('xl/workbook.xml'))
        sheets = workbook.find('m:sheets', namespace)
        index = next(i for i, node in enumerate(sheets, 1) if node.attrib['name'] == sheet_name)
        target = f'xl/worksheets/sheet{index}.xml'
        document = ElementTree.fromstring(source.read(target))
        for node in document.findall('.//m:c', namespace):
            if node.attrib['r'] not in values:
                continue
            value = node.find('m:v', namespace)
            if value is None:
                value = ElementTree.SubElement(node, '{' + namespace['m'] + '}v')
            value.text = str(values[node.attrib['r']])
        output = BytesIO()
        with zipfile.ZipFile(output, 'w') as destination:
            for info in source.infolist():
                destination.writestr(info, ElementTree.tostring(document) if info.filename == target else source.read(info.filename))
    return output.getvalue()


def add_executive_columns(workbook, sheet, columns):
    for field, label in EXECUTIVE_HEADERS.items():
        col = sheet.max_column + 1
        sheet.cell(7, col, label)
        columns[label] = col
    sheet.cell(8, columns[EXECUTIVE_HEADERS['current_forecast_aed']], 150)
    sheet.cell(8, columns[EXECUTIVE_HEADERS['pm_forecast_aed']], 175)
    sheet.cell(8, columns[EXECUTIVE_HEADERS['forecast_variance_aed']], -25)
    sheet.cell(8, columns[EXECUTIVE_HEADERS['direct_cost_aed']], 0)
    sheet.cell(8, columns[EXECUTIVE_HEADERS['actual_manhours']], 12.5)


def add_invoice(workbook, sheet, columns):
    invoice = workbook.create_sheet('INVOICING. STATUS')
    labels = ['PROJECT ID', 'PROJECT SUB ID', 'PROJECT TITLE', 'Core / Subitem', 'I.C',
              'SUBITEM CONTRACT VALUE', 'SUBITEM CONTRACT VALUE', 'TOTAL INVOICED',
              'BALANCE', '% INVOICE (CONTACT VS INVOICED)', 'Revenue Recognized till July, 2026',
              'Invoicing Variance AED (May POC - Current Invoicing)',
              'Invoicing Variance % (May POC - Current Invoicing)',
              'INVOICING STATUS REMARKS', 'POC Recognized CURRENT']
    for col, label in enumerate(labels, 1):
        invoice.cell(7, col, label)
    invoice.cell(5, 7, 'SUB-ITEM INVOICING')
    invoice.cell(6, 15, date(2026, 9, 18))
    for row, values in [(8, ['unmatched', 'foreign', 'Unmatched invoice', 'Core', 'Y', 9999, 9999, 9999]),
                        (12, ['00042', '00042-1', 'Sample project', 'Sub-items', 'Y', 9999,
                              100, 120, -20, 1.2, 130, -10, -0.1, 'Sample note'])]:
        for col, value in enumerate(values, 1):
            invoice.cell(row, col, value)
    return invoice


def add_report(workbook, sheet, columns):
    report = workbook.create_sheet('Report')
    for col, kind, key, label in [(2, 'BUSINESS UNIT', 'DE', 'Detail Engineering'),
                                   (10, 'CLIENT', 'Client', 'Client'), (18, 'PM', 'PM1', 'PM1')]:
        report.cell(1, col, 'RAD REVENUE PERFORMANCE SUMMARY BY ' + kind)
        report.cell(2, col, 'REVENUE')
        for offset, (field, title) in zip([1, 2, 4, 5, 6], REPORT_METRICS.items()):
            report.cell(2, col + offset, title)
            value = {'actual_revenue_aed': 10, 'gross_margin_pct': -0.5,
                     'current_forecast_aed': 15, 'pm_forecast_aed': 20, 'forecast_variance_aed': -5}[field]
            report.cell(3, col + offset, value)
            report.cell(4, col + offset, value)
        if kind == 'BUSINESS UNIT':
            report.cell(3, col - 1, key)
        report.cell(3, col, label)
        report.cell(4, col, 'TOTAL')
        report.cell(5, col, 'DIRECT COSTS')
        report.cell(6, col, label)
        report.cell(6, col + 1, 15)
        report.cell(7, col, 'TOTAL')
        report.cell(7, col + 1, 15)
    report['Z1'] = date(2026, 9, 18)
    report['Z2'] = 'BACKLOG (AED)'
    report['Z3'] = 500
    report['Z4'] = 500
    report['AE2'] = 'Risk'
    report['AD3'] = 'PM1'
    report['AE3'] = 50
    report['AD4'] = 'TOTAL'
    report['AE4'] = 50
    return report


def add_capacity(workbook, sheet, columns):
    capacity = workbook.create_sheet('Rolling Forecast Graph')
    for row, label in [(10, 'Backlog type'), (11, 'Delivery Backlog Mhrs'),
                       (13, 'Current Capacity - "Delivery"'),
                       (15, 'Current capacity considering Avg mhrs lost in absences')]:
        capacity.cell(row, 4, label)
    # Deliberately reorder months and leave a gap between labelled rows.
    for col, period in [(5, date(2026, 10, 1)), (6, date(2026, 9, 1))]:
        capacity.cell(10, col, period)
        capacity.cell(11, col, 343.5)
        capacity.cell(13, col, 396)
        capacity.cell(15, col, 356.4)
    return capacity


def add_pm_kpi(workbook, sheet, columns):
    kpi = workbook.create_sheet('Graphs')
    kpi['B3'] = 'PM'
    kpi['A4'], kpi['B4'] = 'PM1', 'First Manager'
    kpi['A5'], kpi['B5'] = 'PM2', 'Second Manager'
    kpi['B6'] = 'Overall'
    kpi['X1'] = 'KPI - JUN - 2026'
    for col, label in enumerate(['PM', 'Revenue', 'Invoicing', 'GM Delivered', 'GM%', 'CPI', 'Overall'], 24):
        kpi.cell(3, col, label)
    for col, value in [(25, .35), (26, .1), (28, .4), (29, .15)]:
        kpi.cell(2, col, value)
    for row, name in [(4, 'Second Manager'), (5, 'First Manager')]:
        for col, value in enumerate([name, 0.5, 2.4, -0.25, 0.7, 1, 0.9], 24):
            kpi.cell(row, col, value)
    return kpi


class SupplementalWorkbookTests(TestCase):
    def test_current_forecasts_are_separate_from_cumulative_revenue(self):
        parsed = read_workbook(workbook_bytes(mutate=add_executive_columns))
        self.assertEqual(parsed['parser_version'], '2.1')
        row = parsed['rows'][0]
        facts = row['extra']['executive']
        self.assertEqual(facts['period'], '2026-09-01')
        self.assertEqual(Decimal(facts['current_forecast_aed']), 150)
        self.assertEqual(Decimal(facts['pm_forecast_aed']), 175)
        self.assertEqual(Decimal(facts['forecast_variance_aed']), -25)
        self.assertEqual(Decimal(facts['direct_cost_aed']), 0)
        self.assertEqual(Decimal(facts['actual_manhours']), Decimal('12.5'))
        self.assertIsNone(facts['delay_days'])
        self.assertEqual(row['recognized_revenue_aed'], 0)
        json.dumps(row['extra'])
        json.dumps(parsed['reconciliation'])

    def test_supplemental_formula_cache_is_unknown_and_never_executed(self):
        def mutate(workbook, sheet, columns):
            add_executive_columns(workbook, sheet, columns)
            sheet.cell(8, columns[EXECUTIVE_HEADERS['current_forecast_aed']], '=999999')
            sheet.cell(8, columns[EXECUTIVE_HEADERS['pm_forecast_aed']], '#REF!')
        parsed = read_workbook(workbook_bytes(mutate=mutate))
        facts = parsed['rows'][0]['extra']['executive']
        self.assertIsNone(facts['current_forecast_aed'])
        self.assertIsNone(facts['pm_forecast_aed'])
        self.assertEqual(facts['source_formulas']['current_forecast_aed'], '=999999')
        self.assertTrue({'missing_formula_cache', 'excel_error'} <= {item['code'] for item in parsed['warnings']})

    def test_invoice_identity_join_and_subitem_totals_do_not_double_count_core(self):
        parsed = read_workbook(workbook_bytes(mutate=add_invoice))
        facts = parsed['rows'][0]['extra']['invoicing']
        self.assertEqual(facts['source_row'], 12)
        self.assertEqual(Decimal(facts['contract_value_aed']), 100)
        self.assertEqual(Decimal(facts['invoiced_aed']), 120)
        self.assertEqual(Decimal(facts['invoice_pct']), 120)
        self.assertEqual(Decimal(facts['variance_pct']), -10)
        self.assertEqual(facts['comparison_date'], '2026-07-31')
        self.assertEqual(facts['comparison_date_basis'], 'header')
        self.assertTrue(facts['included'])
        self.assertEqual(parsed['reconciliation']['invoicing']['matched'], 1)
        self.assertEqual(len(parsed['reconciliation']['invoicing']['unmatched']), 1)

    def test_invoice_formula_reference_identifies_real_comparison_period(self):
        def mutate(workbook, sheet, columns):
            invoice = add_invoice(workbook, sheet, columns)
            col = get_column_letter(columns['Cumm RR - To-date (AED)'])
            invoice['K12'] = f'=_xlfn.XLOOKUP(B12,POC!$B$8:$B$100,POC!${col}$8:${col}$100,"-",FALSE)'
        content = cached_values(workbook_bytes(mutate=mutate), 'INVOICING. STATUS', {'K12': 130})
        parsed = read_workbook(content)
        facts = parsed['rows'][0]['extra']['invoicing']
        self.assertEqual(facts['comparison_date'], '2026-09-18')
        self.assertEqual(facts['comparison_date_basis'], 'formula_reference')
        self.assertEqual(Decimal(facts['comparison_revenue_aed']), 130)
        self.assertIn('source_header_period_conflict', {item['code'] for item in parsed['warnings']})

    def test_invoice_unknown_formulas_and_blank_amounts_do_not_become_zero(self):
        def mutate(workbook, sheet, columns):
            invoice = add_invoice(workbook, sheet, columns)
            invoice['K12'] = '=500'
            invoice['H12'] = None
            invoice['I12'] = '#N/A'
            invoice['H20'] = 50
        parsed = read_workbook(workbook_bytes(mutate=mutate))
        facts = parsed['rows'][0]['extra']['invoicing']
        self.assertIsNone(facts['invoiced_aed'])
        self.assertIsNone(facts['balance_aed'])
        self.assertIsNone(facts['comparison_date'])
        self.assertIsNone(facts['comparison_revenue_aed'])
        self.assertEqual(parsed['reconciliation']['invoicing']['review_rows'][0]['source_row'], 20)

    def test_duplicate_invoice_identities_are_rejected(self):
        def mutate(workbook, sheet, columns):
            invoice = add_invoice(workbook, sheet, columns)
            invoice['A15'], invoice['B15'], invoice['C15'] = '00042', '00042-1', 'Duplicate'
        with self.assertRaisesRegex(ValueError, 'Duplicate invoicing'):
            read_workbook(workbook_bytes(mutate=mutate))

    def test_report_summary_keeps_missing_detail_coverage_and_source_disagreements(self):
        parsed = read_workbook(workbook_bytes(mutate=add_report))
        summary = parsed['reconciliation']['executive_summary']
        self.assertEqual(Decimal(summary['totals']['actual_revenue_aed']), 10)
        self.assertEqual(Decimal(summary['totals']['direct_cost_aed']), 15)
        self.assertEqual(Decimal(summary['totals']['gross_margin_pct']), -50)
        self.assertEqual(Decimal(summary['totals']['backlog_aed']), 500)
        self.assertEqual(summary['business_units'][0]['key'], 'DE')
        self.assertEqual(summary['detail_reconciliation']['actual_revenue_aed']['difference'], '10.00000000')
        self.assertEqual(summary['detail_reconciliation']['backlog_aed']['missing_count'], 1)
        self.assertIsNone(summary['detail_reconciliation']['backlog_aed']['known_total'])
        self.assertIn('summary_detail_mismatch', {item['code'] for item in parsed['warnings']})

    def test_capacity_uses_dates_and_manhours_not_headcount(self):
        def mutate(workbook, sheet, columns):
            capacity = add_capacity(workbook, sheet, columns)
            capacity['E15'] = '=396*.9'
            report = add_report(workbook, sheet, columns)
            report['B30'], report['C30'], report['F30'] = 'Engineering FTE (including partners)', 10.5, 1000
            report['B31'], report['C31'], report['F31'] = 'Local Employees', 2, 200
        parsed = read_workbook(workbook_bytes(mutate=mutate))
        facts = parsed['reconciliation']['capacity']
        self.assertEqual(facts['unit'], 'manhours')
        self.assertEqual([item['period'] for item in facts['periods']], ['2026-09-01', '2026-10-01'])
        self.assertEqual(Decimal(facts['periods'][0]['adjusted_capacity_manhours']), Decimal('356.4'))
        self.assertIsNone(facts['periods'][1]['adjusted_capacity_manhours'])
        self.assertEqual([item['unit'] for item in facts['staffing']], ['fte', 'unspecified'])
        self.assertIsNone(facts['staffing'][0]['period'])
        self.assertFalse(any(item.get('is_total') for item in facts['staffing']))
        self.assertNotIn('cost_aed', facts['staffing'][0])
        self.assertEqual(Decimal(facts['staffing'][0]['planning_amount_aed']), 1000)

    def test_staffing_saved_total_is_verified_without_recalculating_or_inventing_units(self):
        def mutate(workbook, sheet, columns):
            report = add_report(workbook, sheet, columns)
            report['B30'], report['C30'], report['F30'] = 'Engineering FTE (including partners)', 10.5, 1000
            report['B31'], report['C31'], report['F31'] = 'Local Employees', 2, 200
            # One blank separating row. The deliberately distinct saved value
            # proves we preserve the cache instead of recalculating 10.5 + 2.
            report['C33'], report['F33'] = '=SUM($C$30:C32)', '=SUM(F30:F32)'
        content = cached_values(workbook_bytes(mutate=mutate), 'Report', {'C33': 279, 'F33': 5000})
        staffing = read_workbook(content)['reconciliation']['capacity']['staffing']
        self.assertEqual(len(staffing), 3)
        total = staffing[-1]
        self.assertEqual(total['label'], 'Reported planning total')
        self.assertEqual(Decimal(total['value']), 279)
        self.assertEqual(total['unit'], 'mixed_source_units')
        self.assertTrue(total['is_total'])
        self.assertIsNone(total['period'])
        self.assertEqual(total['source_cells']['value'], 'C33')
        self.assertEqual(Decimal(total['planning_amount_aed']), 5000)
        self.assertNotIn('cost_aed', total)

    def test_staffing_total_with_missing_cache_stays_unknown(self):
        def mutate(workbook, sheet, columns):
            report = add_report(workbook, sheet, columns)
            report['B30'], report['C30'] = 'Engineering FTE', 10
            report['C32'] = '=SUM(C30:C31)'
        parsed = read_workbook(workbook_bytes(mutate=mutate))
        total = parsed['reconciliation']['capacity']['staffing'][-1]
        self.assertTrue(total['is_total'])
        self.assertIsNone(total['value'])
        self.assertIn('missing_formula_cache', {item['code'] for item in parsed['warnings']})

    def test_unverified_staffing_totals_are_not_invented(self):
        for formula in ('=SUM(D30:D32)', '=SUM(C29:C32)', '=SUM(C30:C31)', '=SUM(C30:C33)', 279):
            with self.subTest(formula=formula):
                def mutate(workbook, sheet, columns):
                    report = add_report(workbook, sheet, columns)
                    report['B30'], report['C30'] = 'Engineering FTE', 10
                    report['B31'], report['C31'] = 'Local Employees', 2
                    report['C33'] = formula
                content = workbook_bytes(mutate=mutate)
                if isinstance(formula, str):
                    content = cached_values(content, 'Report', {'C33': 279})
                parsed = read_workbook(content)
                self.assertFalse(any(item.get('is_total') for item in parsed['reconciliation']['capacity']['staffing']))
                self.assertIn('unverified_staffing_total', {item['code'] for item in parsed['warnings']})

    def test_duplicate_capacity_months_abort(self):
        def mutate(workbook, sheet, columns):
            capacity = add_capacity(workbook, sheet, columns)
            capacity['E10'] = capacity['F10'].value
        with self.assertRaisesRegex(ValueError, 'distinct dated capacity'):
            read_workbook(workbook_bytes(mutate=mutate))

    def test_pm_kpi_joins_names_preserves_ratio_units_and_stale_period(self):
        def mutate(workbook, sheet, columns):
            kpi = add_pm_kpi(workbook, sheet, columns)
            kpi['Z4'] = '240%'
            kpi['AC5'] = '=MIN(1,2)'
        parsed = read_workbook(workbook_bytes(mutate=mutate))
        kpi = parsed['reconciliation']['pm_kpi']
        self.assertEqual(kpi['period'], '2026-06-01')
        self.assertEqual(kpi['entries'][0]['pm'], 'PM2')
        self.assertEqual(Decimal(kpi['entries'][0]['invoicing_ratio']), Decimal('2.4'))
        self.assertEqual(Decimal(kpi['entries'][0]['gm_delivered_pct']), -25)
        self.assertIsNone(kpi['entries'][1]['cpi_ratio'])
        self.assertIn('kpi_label_period_mismatch', {item['code'] for item in parsed['warnings']})

    def test_risk_history_requires_aed_unit_and_preserves_observation_date(self):
        def mutate(workbook, sheet, columns):
            risk = workbook.create_sheet('KPI Trends')
            for col, title in enumerate(['Date', 'EDDR % vs Revenue', 'LD Exposure Risk', 'Prolongation cost Risk'], 1):
                risk.cell(2, col, title)
            risk['B1'] = 'AED'
            risk['A3'], risk['B3'], risk['C3'], risk['D3'] = date(2025, 6, 30), 40, 0, None
        parsed = read_workbook(workbook_bytes(mutate=mutate))
        history = parsed['reconciliation']['risk_history']
        self.assertEqual(history[0]['date'], '2025-06-30')
        self.assertEqual(Decimal(history[0]['poc_risk_aed']), 40)
        self.assertEqual(Decimal(history[0]['ld_exposure_aed']), 0)
        self.assertIsNone(history[0]['prolongation_cost_aed'])

    def test_weekly_poc_risk_keeps_dated_totals_and_manager_values(self):
        def mutate(workbook, sheet, columns):
            risk = workbook.create_sheet('POC RISK')
            risk['B1'] = 'POC Risk :'
            risk['C2'], risk['D2'] = date(2026, 9, 18), date(2026, 9, 11)
            risk['B3'], risk['B4'] = 'PM1', 'TOTAL'
            risk['C3'], risk['C4'], risk['D3'], risk['D4'] = 40, 40, None, 0
        parsed = read_workbook(workbook_bytes(mutate=mutate))
        history = parsed['reconciliation']['poc_risk_history']
        self.assertEqual([item['date'] for item in history], ['2026-09-11', '2026-09-18'])
        self.assertIsNone(history[0]['by_pm'][0]['poc_risk_aed'])
        self.assertEqual(Decimal(history[0]['poc_risk_aed']), 0)
        self.assertEqual(Decimal(history[1]['poc_risk_aed']), 40)
        self.assertEqual(history[1]['by_pm'][0]['pm'], 'PM1')

    def test_duplicate_weekly_poc_risk_dates_abort(self):
        def mutate(workbook, sheet, columns):
            risk = workbook.create_sheet('POC RISK')
            risk['B1'] = 'POC Risk :'
            risk['C2'], risk['D2'] = date(2026, 9, 18), date(2026, 9, 18)
        with self.assertRaisesRegex(ValueError, 'distinct dated risk'):
            read_workbook(workbook_bytes(mutate=mutate))
