from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from unittest import TestCase
from unittest.mock import patch

from openpyxl import Workbook

from apps.portfolio.workbook import DATE_HEADERS, REVENUE_HEADERS, TEXT_HEADERS, VALUE_HEADERS, read_workbook


def workbook_bytes(*, reporting_date=date(2026, 9, 18), rows=None, mutate=None):
    """Synthetic source schema; no customer data or real workbook committed."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'POC'
    labels = [*TEXT_HEADERS.values(), *VALUE_HEADERS.values(), *DATE_HEADERS.values(),
              'WITHOUT PT', 'WITH PT', 'Original Contract Value (Cont. Curr.)',
              'Original Contract Value (Cont. Curr.)', 'Overclaim (POC & EDDR)', *REVENUE_HEADERS]
    columns = {}
    for col, label in enumerate(labels, 1):
        sheet.cell(7, col, label)
        columns[label] = col
    sheet.cell(5, columns[REVENUE_HEADERS[0]], reporting_date)
    default = {'PROJECT ID': '00042', 'PROJECT SUB ID': '00042-1', 'PROJECT TITLE': 'Sample project',
               'PM': 'PM1', 'PC Name': 'Coordinator', 'BU': 'DE', 'CLIENT': 'Client', 'SCOPE TYPE': 'Base',
               'WITHOUT PT': 'Yes', 'WITH PT': 'Yes', 'Original Contract Value (Cont. Curr.)': 'AED',
               'Modified Contract Value (AED)': Decimal('123.45'), 'Cumm RR - To-date (AED)': 0,
               'POC / EV - To-date (%)': 0.25, 'RR For the period (AED)': 0,
               'Contractual Start': date(2026, 1, 1), 'Contractual Finish': date(2026, 12, 31),
               'Forecast completion': date(2026, 12, 31)}
    for number, overrides in enumerate(rows or [{}], 8):
        for label, value in {**default, **overrides}.items():
            if value is not None:
                sheet.cell(number, columns[label], value)
    if mutate:
        mutate(workbook, sheet, columns)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


class WorkbookReaderTests(TestCase):
    def test_decimal_identity_units_and_missing_are_preserved(self):
        parsed = read_workbook(workbook_bytes(rows=[{'PROJECT ID': '\t00042 ', 'Target GM': -0.15}]))
        row = parsed['rows'][0]
        self.assertEqual(parsed['reporting_date'], date(2026, 9, 18))
        self.assertEqual(row['project_code'], '00042')
        self.assertEqual(row['contract_value_aed'], Decimal('123.45000000'))
        self.assertEqual(row['recognized_revenue_aed'], 0)
        self.assertIsNone(row['backlog_without_pt_aed'])
        self.assertEqual(row['poc_pct'], 25)
        self.assertEqual(row['target_margin_pct'], -15)
        self.assertTrue(row['include_without_pt'])

    def test_formula_without_cache_and_invalid_values_become_unknown(self):
        parsed = read_workbook(workbook_bytes(rows=[{
            'Modified Contract Value (AED)': '=1+2', 'PROGRESS EDDR': '#REF!',
            'FC GM': 'not a number', 'POC / EV - To-date (%)': '12.5%', 'WITHOUT PT': 'maybe',
        }]))
        row = parsed['rows'][0]
        self.assertIsNone(row['contract_value_aed'])
        self.assertIsNone(row['eddr_pct'])
        self.assertIsNone(row['forecast_margin_pct'])
        self.assertEqual(row['poc_pct'], Decimal('12.5'))
        self.assertIsNone(row['include_without_pt'])
        codes = {item['code'] for item in parsed['warnings']}
        self.assertTrue({'missing_formula_cache', 'excel_error', 'invalid_percentage', 'unknown_inclusion_flag'} <= codes)

    def test_percentage_roundoff_is_valid_but_actual_overflow_is_unknown(self):
        parsed = read_workbook(workbook_bytes(rows=[{'POC / EV - To-date (%)': 1.0000000000000022,
                                                   'PROGRESS EDDR': 1.2}]))
        self.assertEqual(parsed['rows'][0]['poc_pct'], 100)
        self.assertIsNone(parsed['rows'][0]['eddr_pct'])

    def test_duplicate_normalized_identity_and_missing_keys_abort(self):
        with self.assertRaisesRegex(ValueError, 'Duplicate POC'):
            read_workbook(workbook_bytes(rows=[{}, {'PROJECT ID': '\t00042'}]))
        with self.assertRaisesRegex(ValueError, 'both project and subproject'):
            read_workbook(workbook_bytes(rows=[{'PROJECT SUB ID': None}]))

    def test_data_row_with_both_identifiers_missing_is_rejected(self):
        with self.assertRaisesRegex(ValueError, 'POC row 9 requires both'):
            read_workbook(workbook_bytes(rows=[{}, {'PROJECT ID': None, 'PROJECT SUB ID': None}]))

    def test_untitled_financial_rows_cannot_be_silently_dropped(self):
        for value in (0, 123, '=SUM(A1:A2)'):
            with self.subTest(value=value):
                def financial_row(workbook, sheet, columns):
                    sheet.cell(9, columns['Modified Contract Value (AED)'], value)
                with self.assertRaisesRegex(ValueError, 'POC row 9 requires both'):
                    read_workbook(workbook_bytes(mutate=financial_row))

    def test_blank_and_explicit_summary_rows_are_not_projects(self):
        def summary(workbook, sheet, columns):
            # Row 9 is blank. A caption alone is insufficient to exempt a row
            # carrying project descriptors, but this standalone total is valid.
            sheet.cell(10, columns['PROJECT TITLE'], 'TOTAL:')
            sheet.cell(10, columns['Modified Contract Value (AED)'], '=SUM(A1:A2)')
        self.assertEqual(read_workbook(workbook_bytes(mutate=summary))['row_count'], 1)
        with self.assertRaisesRegex(ValueError, 'both project and subproject'):
            read_workbook(workbook_bytes(rows=[{'PROJECT ID': None, 'PROJECT SUB ID': None,
                                               'PROJECT TITLE': 'TOTAL'}]))

    def test_required_header_and_period_date_are_validated(self):
        def bad_header(workbook, sheet, columns):
            sheet.cell(7, columns['CLIENT'], 'Something else')
        with self.assertRaisesRegex(ValueError, 'CLIENT'):
            read_workbook(workbook_bytes(mutate=bad_header))
        def bad_date(workbook, sheet, columns):
            sheet.cell(5, columns[REVENUE_HEADERS[0]], 'September sometime')
        with self.assertRaisesRegex(ValueError, 'Missing reporting date'):
            read_workbook(workbook_bytes(mutate=bad_date))

    def test_source_file_must_remain_stable_while_reading(self):
        with patch('apps.portfolio.workbook.Path.read_bytes', side_effect=[workbook_bytes(), b'changed']):
            with self.assertRaisesRegex(ValueError, 'changed while'):
                read_workbook('source.xlsx')

    def test_latest_dated_triplet_selected_instead_of_last_column(self):
        def add_old_period(workbook, sheet, columns):
            col = sheet.max_column + 2
            for offset, label in enumerate(REVENUE_HEADERS):
                sheet.cell(7, col + offset, label)
                sheet.cell(8, col + offset, 0)
            sheet.cell(5, col, date(2025, 12, 31))
        parsed = read_workbook(workbook_bytes(mutate=add_old_period))
        self.assertEqual(parsed['reporting_date'], date(2026, 9, 18))
        self.assertEqual(parsed['rows'][0]['poc_pct'], 25)
        self.assertEqual(len(parsed['rows'][0]['extra']['history']), 2)

    def test_forecast_joins_by_identity_not_row_position(self):
        def forecast(workbook, sheet, columns):
            forecast = workbook.create_sheet('Delivery Rolling Forecast')
            for col, label in enumerate(['Project ID', 'Project Sub ID', 'Project Title', date(2026, 10, 1)], 1):
                forecast.cell(7, col, label)
            forecast.cell(3, 4, 'REVENUE PER MONTH (AED)')
            for n, values in [(8, ['99', '99', 'Unmatched', 999]), (12, ['00042', '00042-1', 'Sample project', 10])]:
                for col, value in enumerate(values, 1):
                    forecast.cell(n, col, value)
        parsed = read_workbook(workbook_bytes(mutate=forecast))
        row = parsed['rows'][0]
        self.assertEqual(row['extra']['forecast_source_row'], 12)
        self.assertEqual(row['extra']['forecasts'][0]['revenue_aed'], '10.00000000')
        self.assertEqual(parsed['reconciliation']['forecast']['matched'], 1)
        self.assertEqual(parsed['reconciliation']['forecast']['unmatched'][0]['project_code'], '99')

    def test_resource_secondary_header_changes_metric_units(self):
        def resource_section(workbook, sheet, columns):
            current = columns[REVENUE_HEADERS[0]]
            for offset, label in enumerate(['BOOKED MHR (PERIOD)', REVENUE_HEADERS[0], REVENUE_HEADERS[2]]):
                sheet.cell(9, current + offset, label)
            for cell in list(sheet[8]):
                if cell.value is not None:
                    sheet.cell(10, cell.column, cell.value)
            sheet.cell(10, columns['PROJECT ID'], '43')
            sheet.cell(10, columns['PROJECT SUB ID'], '43')
            sheet.cell(10, current, 100)
            sheet.cell(10, current + 1, 25000)
            sheet.cell(10, columns['PROGRESS EDDR'], 125)
            sheet.cell(10, columns['Overclaim (POC & EDDR)'], -50)
        parsed = read_workbook(workbook_bytes(mutate=resource_section))
        row = parsed['rows'][1]
        self.assertEqual(row['recognized_revenue_aed'], 25000)
        self.assertIsNone(row['poc_pct'])
        self.assertIsNone(row['eddr_pct'])
        self.assertIsNone(row['overclaim_aed'])
        self.assertEqual(row['extra']['booked_manhours'], '100.00000000')
        self.assertEqual(row['extra']['executive']['actual_manhours'], '100.00000000')
