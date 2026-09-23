"""Cached workbook facts used by the executive portfolio.

All amounts are AED decimal strings or null. Percent fields use percentage
points; ``*_ratio`` fields retain ratios (including values greater than one).
No formula is evaluated. Formula references are inspected only to identify a
dated source column or document provenance. Global facts are reconciliation
data, and must never be exposed as totals for a restricted project scope.
"""
import calendar
from datetime import date
from decimal import Decimal
import re

from openpyxl.utils import column_index_from_string, get_column_letter

from .workbook import _date_value, _json_value, _normal, _number, _text, _warning


EXECUTIVE_HEADERS = {
    'current_forecast_aed': 'Fcst this period (AED)',
    'pm_forecast_aed': 'PM FC Month (AED)',
    'forecast_variance_aed': 'Variance this period (Frst - PM)',
    'direct_cost_aed': 'Total', 'actual_manhours': 'Man-Hrs',
    'labor_cost_aed': 'Man-Hrs Cost', 'other_cost_aed': 'Others',
    'head_office_cost_aed': 'HO', 'resource_cost_aed': 'Resource',
    'delay_days': 'No of days delay', 'ld_min_pct': 'LD clause Min limit % of Contract Value',
    'ld_max_pct': 'LD Max cap',
    'extension_resources_per_day': 'Avg PMT Resources required / day of extension (From PM)',
    'extension_cost_per_hour_aed': 'Avg Mhr costs (AED) - Incl G&A',
}


def _headers(sheet, row):
    result = {}
    for cell in sheet[row]:
        if _normal(cell.value):
            result.setdefault(_normal(cell.value), []).append(cell.column)
    return result


def _column(headers, label, sheet, *, required=True):
    matches = headers.get(_normal(label), [])
    if len(matches) > 1 or required and not matches:
        raise ValueError(f'{sheet} requires one {label!r} header.')
    return matches[0] if matches else None


def _find(sheet, text, *, maximum_row=None):
    return [cell for row in sheet.iter_rows(max_row=maximum_row) for cell in row
            if _normal(cell.value) == _normal(text)]


def _one_label(sheet, text):
    matches = _find(sheet, text)
    if len(matches) != 1:
        raise ValueError(f'{sheet.title} requires one {text!r} label.')
    return matches[0]


def _numeric(sheet, formulas, row, col, field, warnings, *, percentage=False):
    if col is None:
        return None
    cell = sheet.cell(row, col)
    value = _number(cell, formulas.cell(row, col), sheet=sheet.title, row=row,
                    col=col, field=field, warnings=warnings)
    # Scores/invoicing percentages can exceed 100, and invoicing variances can
    # be negative. The POC completion-percentage bounds do not apply to them.
    if value is not None and percentage and not (
        isinstance(cell.value, str) and cell.value.strip().endswith('%')
    ):
        value *= 100
    if value is not None and field.endswith('_ratio') and isinstance(cell.value, str) and cell.value.strip().endswith('%'):
        value /= 100
    return _json_value(value)


def _facts(sheet, formulas, row, mapping, warnings):
    values = {key: _numeric(sheet, formulas, row, col, key, warnings,
                           percentage=key.endswith('_pct')) for key, col in mapping.items()}
    values['source_cells'] = {key: f'{get_column_letter(col)}{row}' for key, col in mapping.items() if col}
    values['source_formulas'] = {
        key: formulas.cell(row, col).value for key, col in mapping.items()
        if col and formulas.cell(row, col).data_type == 'f'
    }
    return values


def _optional_sheet(workbook, name, warnings):
    if name not in workbook.sheetnames:
        warnings.append({'sheet': name, 'code': 'optional_sheet_missing'})
        return None
    sheet = workbook[name]
    if sheet.max_row > 10000 or sheet.max_column > 500:
        raise ValueError(f'{name} exceeds the supported worksheet dimensions.')
    return sheet


def _executive_rows(workbook, formulas, rows, reporting_date, warnings):
    sheet, source = workbook['POC'], formulas['POC']
    headers = _headers(sheet, 7)
    mapping = {key: _column(headers, label, sheet.title, required=False)
               for key, label in EXECUTIVE_HEADERS.items()}
    text_columns = {key: _column(headers, label, sheet.title, required=False) for key, label in {
        'ld_possible': 'Possibility of LD', 'ld_frequency': 'Frequency',
        'remarks': 'Remarks', 'award_status': 'PROJECT AWARD STATUS',
        'priority': 'PRIORITY', 'eddr_status': 'EDDR',
    }.items()}
    date_columns = {key: _column(headers, label, sheet.title, required=False) for key, label in {
        'actual_completion': 'Actual Completion Date',
        'handover_date': 'Delivery to Sales Handover Date',
    }.items()}
    for key, col in mapping.items():
        if col is None:
            warnings.append({'sheet': 'POC', 'field': key, 'code': 'optional_header_missing'})
    for row in rows:
        start = len(warnings)
        number = row['source_row']
        facts = _facts(sheet, source, number, mapping, warnings)
        facts['period'] = reporting_date.replace(day=1).isoformat()
        for key, col in text_columns.items():
            if col:
                cell = sheet.cell(number, col)
                value = _text(cell.value)
                if cell.data_type == 'e' or source.cell(number, col).data_type == 'f' and cell.value is None:
                    _warning(warnings, 'POC', number, col, key, 'invalid_or_missing_text')
                    value = ''
                facts[key] = {'yes': True, 'no': False}.get(_normal(value)) if key == 'ld_possible' else value
                facts['source_cells'][key] = cell.coordinate
            else:
                facts[key] = None if key == 'ld_possible' else ''
        for key, col in date_columns.items():
            facts[key] = None
            if col:
                cell = sheet.cell(number, col)
                facts[key] = _json_value(_date_value(cell.value))
                facts['source_cells'][key] = cell.coordinate
                if facts[key] is None and (cell.value is not None or source.cell(number, col).data_type == 'f'):
                    _warning(warnings, 'POC', number, col, key, 'invalid_or_missing_date')
        if row['extra'].get('section') == 'resource_deputation':
            # EE is empty in the deputation section; the current dated CW
            # column is explicitly labelled BOOKED MHR (PERIOD).
            facts['actual_manhours'] = row['extra'].get('booked_manhours')
            history = row['extra'].get('history', [])
            if history:
                facts['source_cells']['actual_manhours'] = history[-1]['cell']
        row['extra']['executive'] = facts
        row['extra']['warnings'].extend(warnings[start:])


def _month_in_text(value):
    text = _normal(value)
    for month in range(1, 13):
        match = re.search(r'\b' + calendar.month_abbr[month].lower() + r'[a-z]*[\s,\-]+(20\d{2})\b', text)
        if match:
            return date(int(match.group(1)), month, 1)
    return None


def _invoice_comparison_date(formula, poc, header_date):
    # Only recognize XLOOKUP's explicit POC return-column range. Looking at
    # arbitrary referenced columns could accidentally select its lookup key.
    if formula.data_type == 'f' and 'XLOOKUP(' in formula.value.upper():
        ranges = re.findall(r"(?:'POC'|POC)!\$?([A-Z]+)\$?\d+:\$?([A-Z]+)\$?\d+", formula.value, re.I)
        if ranges:
            first, last = ranges[-1]
            if first.upper() == last.upper():
                col = column_index_from_string(first.upper())
                if _normal(poc.cell(7, col).value) == _normal('Cumm RR - To-date (AED)'):
                    value = _date_value(poc.cell(5, col).value)
                    if value:
                        return value, 'formula_reference'
        return None, 'unresolved_formula_reference'
    if formula.data_type == 'f':
        return None, 'unresolved_formula_reference'
    if header_date:
        return header_date.replace(day=calendar.monthrange(header_date.year, header_date.month)[1]), 'header'
    return None, 'unknown'


def _invoicing(workbook, formulas, rows, warnings):
    name = 'INVOICING. STATUS'
    sheet = _optional_sheet(workbook, name, warnings)
    result = {'matched': 0, 'unmatched': [], 'review_rows': [], 'missing': len(rows)}
    if sheet is None:
        return result
    source = formulas[name]
    headers = _headers(sheet, 7)
    identities = {(row['project_code'], row['subproject_code']): row for row in rows}
    identity_columns = {key: _column(headers, label, name) for key, label in {
        'project_code': 'PROJECT ID', 'subproject_code': 'PROJECT SUB ID', 'title': 'PROJECT TITLE',
    }.items()}
    # The sheet repeats core/subitem contract columns. The invoicing section
    # anchor explicitly identifies the additive SUB-ITEM block.
    anchor = _one_label(sheet, 'SUB-ITEM INVOICING')
    if _normal(sheet.cell(7, anchor.column).value) != _normal('SUBITEM CONTRACT VALUE'):
        raise ValueError('INVOICING. STATUS sub-item invoicing contract header is invalid.')
    mapping = {'contract_value_aed': anchor.column}
    mapping.update({key: _column(headers, label, name) for key, label in {
        'invoiced_aed': 'TOTAL INVOICED', 'balance_aed': 'BALANCE',
        'invoice_pct': '% INVOICE (CONTACT VS INVOICED)',
    }.items()})
    for key, prefix in {'comparison_revenue_aed': 'revenue recognized',
                         'variance_aed': 'invoicing variance aed',
                         'variance_pct': 'invoicing variance %'}.items():
        matches = [col for label, cols in headers.items() if label.startswith(prefix) for col in cols]
        if len(matches) != 1:
            raise ValueError(f'{name} requires one {prefix!r} header.')
        mapping[key] = matches[0]
    text_columns = {key: _column(headers, label, name) for key, label in {
        'core_subitem': 'Core / Subitem', 'included': 'I.C', 'remarks': 'INVOICING STATUS REMARKS',
    }.items()}
    current = _column(headers, 'POC Recognized CURRENT', name)
    reporting_date = _date_value(sheet.cell(6, current).value)
    if reporting_date is None:
        _warning(warnings, name, 6, current, 'reporting_date', 'invalid_or_missing_date')
    result['reporting_date'] = _json_value(reporting_date)
    comparison_header = sheet.cell(7, mapping['comparison_revenue_aed'])
    header_month = _month_in_text(comparison_header.value)
    result['comparison_header'] = _text(comparison_header.value)
    seen, period_conflicts = set(), set()
    for number in range(8, sheet.max_row + 1):
        identity = tuple(_text(sheet.cell(number, identity_columns[key]).value)
                         for key in ('project_code', 'subproject_code'))
        title = _text(sheet.cell(number, identity_columns['title']).value)
        record = {'source_row': number, 'project_code': identity[0], 'subproject_code': identity[1], 'title': title}
        if not any(identity):
            if title or any(sheet.cell(number, col).value is not None or source.cell(number, col).data_type == 'f'
                            for col in mapping.values()):
                result['review_rows'].append({**record, 'reason': 'missing_invoice_identity'})
            continue
        if not all(identity) or any(sheet.cell(number, identity_columns[key]).data_type == 'e'
                                    for key in ('project_code', 'subproject_code')):
            result['review_rows'].append({**record, 'reason': 'incomplete_invoice_identity'})
            continue
        if identity in seen:
            raise ValueError(f'Duplicate invoicing project/subproject identity at row {number}.')
        seen.add(identity)
        if identity not in identities:
            result['unmatched'].append(record)
            continue
        row = identities[identity]
        start = len(warnings)
        facts = _facts(sheet, source, number, mapping, warnings)
        comparison_date, basis = _invoice_comparison_date(
            source.cell(number, mapping['comparison_revenue_aed']), workbook['POC'], header_month)
        facts.update(sheet=name, source_row=number, reporting_date=_json_value(reporting_date),
                     comparison_date=_json_value(comparison_date), comparison_date_basis=basis)
        if comparison_date and header_month and comparison_date.replace(day=1) != header_month:
            period_conflicts.add(comparison_date.isoformat())
        if comparison_date is None:
            _warning(warnings, name, number, mapping['comparison_revenue_aed'],
                     'comparison_date', 'unknown_comparison_period')
        for key, col in text_columns.items():
            cell = sheet.cell(number, col)
            facts['source_cells'][key] = cell.coordinate
            value = _text(cell.value)
            if key == 'included':
                facts[key] = {'y': True, 'yes': True, 'n': False, 'no': False}.get(_normal(value))
            else:
                facts[key] = '' if cell.data_type == 'e' else value
        row['extra']['invoicing'] = facts
        row['extra']['warnings'].extend(warnings[start:])
        result['matched'] += 1
    result['missing'] = len(rows) - result['matched']
    if period_conflicts:
        warnings.append({'sheet': name, 'cell': comparison_header.coordinate,
                         'field': 'comparison_date', 'code': 'source_header_period_conflict',
                         'formula_periods': sorted(period_conflicts)})
    if result['unmatched'] or result['review_rows']:
        warnings.append({'sheet': name, 'code': 'invoice_identity_review_required',
                         'unmatched_count': len(result['unmatched']), 'review_count': len(result['review_rows'])})
    return result


def _capacity(workbook, formulas, warnings):
    name = 'Rolling Forecast Graph'
    sheet = _optional_sheet(workbook, name, warnings)
    result = {'sheet': name, 'unit': 'manhours', 'periods': [], 'staffing': [], 'warnings': []}
    if sheet is not None:
        source = formulas[name]
        labels = {key: _one_label(sheet, label) for key, label in {
            'demand_manhours': 'Delivery Backlog Mhrs',
            'gross_capacity_manhours': 'Current Capacity - "Delivery"',
            'adjusted_capacity_manhours': 'Current capacity considering Avg mhrs lost in absences',
        }.items()}
        anchor = _one_label(sheet, 'Backlog type')
        if any(cell.column != anchor.column for cell in labels.values()):
            raise ValueError('Rolling Forecast Graph capacity row labels are not aligned.')
        periods = []
        for col in range(anchor.column + 1, sheet.max_column + 1):
            period = _date_value(sheet.cell(anchor.row, col).value)
            if period:
                periods.append((period, col))
        if not periods or len({period for period, _ in periods}) != len(periods):
            raise ValueError('Rolling Forecast Graph requires distinct dated capacity columns.')
        for period, col in sorted(periods):
            facts = {'period': period.isoformat(), 'source_cells': {}, 'source_formulas': {}}
            for key, label in labels.items():
                value = _numeric(sheet, source, label.row, col, key, warnings)
                if value is not None and Decimal(value) < 0:
                    _warning(warnings, name, label.row, col, key, 'negative_capacity_or_demand')
                    value = None
                facts[key] = value
                facts['source_cells'][key] = sheet.cell(label.row, col).coordinate
                if source.cell(label.row, col).data_type == 'f':
                    facts['source_formulas'][key] = source.cell(label.row, col).value
            result['periods'].append(facts)
    # This separate staffing list has one explicit FTE row and otherwise
    # unspecified quantities. It is not an employee roster or a dated capacity.
    if 'Report' in workbook.sheetnames:
        sheet, source = workbook['Report'], formulas['Report']
        anchors = [cell for row in sheet for cell in row if _normal(cell.value).startswith('engineering fte')]
        if len(anchors) > 1:
            raise ValueError('Report contains duplicate engineering FTE summaries.')
        if anchors:
            anchor = anchors[0]
            first_blank = None
            for number in range(anchor.row, min(anchor.row + 20, sheet.max_row) + 1):
                label = _text(sheet.cell(number, anchor.column).value)
                if not label:
                    first_blank = number
                    break
                facts = _facts(sheet, source, number,
                               {'value': anchor.column + 1, 'planning_amount_aed': anchor.column + 4}, warnings)
                facts.update(label=label, unit='fte' if 'fte' in _normal(label) else 'unspecified', period=None)
                result['staffing'].append(facts)
            if first_blank is not None:
                quantity_col = anchor.column + 1
                letter = get_column_letter(quantity_col)
                for number in range(first_blank, min(anchor.row + 20, sheet.max_row) + 1):
                    if _text(sheet.cell(number, anchor.column).value):
                        break
                    cell, formula = sheet.cell(number, quantity_col), source.cell(number, quantity_col)
                    if cell.value is None and formula.data_type != 'f':
                        continue
                    # A blank label alone is not a total. Accept only the
                    # saved SUM of this exact staffing quantity column from
                    # the anchor through the row immediately above the total.
                    expected = (rf'=\s*\+?\s*SUM\(\s*\$?{letter}\$?{anchor.row}'
                                rf'\s*:\s*\$?{letter}\$?{number - 1}\s*\)\s*')
                    if formula.data_type == 'f' and re.fullmatch(expected, formula.value, re.I):
                        facts = _facts(sheet, source, number,
                                       {'value': quantity_col, 'planning_amount_aed': anchor.column + 4}, warnings)
                        facts.update(label='Reported planning total', unit='mixed_source_units', period=None, is_total=True)
                        result['staffing'].append(facts)
                    else:
                        _warning(warnings, 'Report', number, quantity_col, 'staffing_total', 'unverified_staffing_total')
                    break
            result['warnings'].append('Staffing quantities have no source date and mixed or unspecified units; do not total as headcount.')
            result['warnings'].append('Staffing planning amounts are saved AED figures with an unspecified monetary basis; they are not verified costs.')
    return result


def _pm_kpi(workbook, formulas, reporting_date, warnings):
    name = 'Graphs'
    sheet = _optional_sheet(workbook, name, warnings)
    result = {'sheet': name, 'period': None, 'period_label': '', 'entries': [], 'weights': {}, 'warnings': []}
    if sheet is None:
        return result
    source = formulas[name]
    anchors = [cell for row in sheet for cell in row if _normal(cell.value).startswith('kpi -')]
    if len(anchors) != 1:
        raise ValueError('Graphs requires one dated KPI section label.')
    anchor = anchors[0]
    result['period_label'] = _text(anchor.value)
    period = _month_in_text(anchor.value)
    result['period'] = _json_value(period)
    if period is None or period != reporting_date.replace(day=1):
        issue = {'sheet': name, 'cell': anchor.coordinate, 'field': 'pm_kpi', 'code': 'kpi_label_period_mismatch'}
        warnings.append(issue)
        result['warnings'].append(issue)
    header_row = anchor.row + 2
    expected = ['PM', 'Revenue', 'Invoicing', 'GM Delivered', 'GM%', 'CPI', 'Overall']
    if [_normal(sheet.cell(header_row, anchor.column + offset).value) for offset in range(7)] != list(map(_normal, expected)):
        raise ValueError('Graphs KPI headers are invalid.')
    # Join the KPI names to the source's named PM-code table. The two tables
    # may be reordered independently; row offsets are never identity joins.
    pm_headers = [cell for cell in sheet[header_row] if _normal(cell.value) == 'pm' and cell.column < anchor.column]
    if not pm_headers:
        raise ValueError('Graphs KPI section has no PM identity table.')
    identity_col = pm_headers[0].column
    identities = {}
    for number in range(header_row + 1, sheet.max_row + 1):
        pm_name = _text(sheet.cell(number, identity_col).value)
        if _normal(pm_name) in {'overall', 'total'}:
            break
        code = _text(sheet.cell(number, identity_col - 1).value)
        if pm_name and code:
            if pm_name in identities:
                raise ValueError('Graphs has duplicate PM names in its identity table.')
            identities[pm_name] = code
    metrics = ['revenue_ratio', 'invoicing_ratio', 'gm_delivered_pct', 'gm_score_ratio', 'cpi_ratio', 'overall_ratio']
    mapping = {key: anchor.column + offset for offset, key in enumerate(metrics, 1)}
    seen = set()
    for number in range(header_row + 1, sheet.max_row + 1):
        pm_name = _text(sheet.cell(number, anchor.column).value)
        if not pm_name or _normal(pm_name) in {'overall', 'total'}:
            break
        if pm_name in seen:
            raise ValueError('Graphs has duplicate PM KPI rows.')
        seen.add(pm_name)
        code = identities.get(pm_name)
        if not code:
            warnings.append({'sheet': name, 'cell': sheet.cell(number, anchor.column).coordinate,
                             'field': 'pm_kpi', 'code': 'unmatched_pm_identity'})
            continue
        facts = _facts(sheet, source, number, mapping, warnings)
        facts.update(pm=code, name=pm_name)
        result['entries'].append(facts)
    result['weights'] = {key: _numeric(sheet, source, anchor.row + 1, anchor.column + offset,
                                     key + '_ratio', warnings) for key, offset in [('revenue', 1), ('invoicing', 2), ('gm', 4), ('cpi', 5)]}
    result['warnings'].append('Workbook CPI is a capped revenue/direct-cost score, not an earned-value cost performance index.')
    if any('Report!' in formula for item in result['entries'] for formula in item['source_formulas'].values()):
        result['warnings'].append('KPI formulas reference the current Report while other inputs may be static; retain the source KPI period label.')
    return result


REPORT_METRICS = {'actual_revenue_aed': 'ACTUAL as on Cutoff', 'gross_margin_pct': 'GM %',
                  'current_forecast_aed': 'CURRENT FORECAST for the month',
                  'pm_forecast_aed': 'PM FORECAST for the month', 'forecast_variance_aed': 'VAR.'}


def _report_block(sheet, source, label, warnings, *, business=False):
    anchor = _one_label(sheet, label)
    header_row = anchor.row + 1
    width = 7
    headers = {}
    for col in range(anchor.column + 1, anchor.column + width):
        headers.setdefault(_normal(sheet.cell(header_row, col).value), []).append(col)
    mapping = {key: _column(headers, name, sheet.title) for key, name in REPORT_METRICS.items()}
    entries, total, end = [], {}, None
    for number in range(header_row + 1, sheet.max_row + 1):
        title = _text(sheet.cell(number, anchor.column).value)
        if not title:
            raise ValueError(f'Report {label!r} has no total before a blank label.')
        facts = _facts(sheet, source, number, mapping, warnings)
        if _normal(title) == 'total':
            total, end = facts, number
            break
        key = _text(sheet.cell(number, anchor.column - 1).value) if business else title
        if not key or any(item['key'] == key for item in entries):
            raise ValueError(f'Report {label!r} has missing or duplicate identity labels.')
        entries.append({'key': key, 'label': title, 'source_row': number, **facts})
    if end is None:
        raise ValueError(f'Report {label!r} has no total.')
    # Direct costs use the same named keys but a separately located block.
    if _normal(sheet.cell(end + 1, anchor.column).value) == 'direct costs':
        by_label = {entry['label']: entry for entry in entries}
        for number in range(end + 2, sheet.max_row + 1):
            title = _text(sheet.cell(number, anchor.column).value)
            if not title:
                break
            target = total if _normal(title) == 'total' else by_label.get(title)
            if target is not None:
                col = mapping['actual_revenue_aed']
                target['direct_cost_aed'] = _numeric(sheet, source, number, col, 'direct_cost_aed', warnings)
                target['source_cells']['direct_cost_aed'] = sheet.cell(number, col).coordinate
            if _normal(title) == 'total':
                break
    return entries, total


def _executive_summary(workbook, formulas, reporting_date, warnings):
    name = 'Report'
    sheet = _optional_sheet(workbook, name, warnings)
    result = {'sheet': name, 'reporting_date': reporting_date.isoformat(), 'totals': {},
              'business_units': [], 'clients': [], 'pms': [], 'source_cells': {}}
    if sheet is None:
        return result
    source = formulas[name]
    for key, title in [('business_units', 'BUSINESS UNIT'), ('clients', 'CLIENT'), ('pms', 'PM')]:
        entries, total = _report_block(sheet, source, 'RAD REVENUE PERFORMANCE SUMMARY BY ' + title,
                                       warnings, business=key == 'business_units')
        result[key] = entries
        if key == 'business_units':
            result['totals'] = {field: value for field, value in total.items() if field not in {'source_cells', 'source_formulas'}}
            result['source_cells'].update(total['source_cells'])
    # The PM backlog column sits in the same table as PM revenue; verify its
    # label and date rather than borrowing an unrelated summary cell.
    anchor = _one_label(sheet, 'RAD REVENUE PERFORMANCE SUMMARY BY PM')
    backlog_col = anchor.column + 8
    if _normal(sheet.cell(anchor.row + 1, backlog_col).value) != 'backlog (aed)':
        raise ValueError('Report PM backlog header is invalid.')
    report_date = _date_value(sheet.cell(anchor.row, backlog_col).value)
    result['reporting_date'] = _json_value(report_date)
    if report_date != reporting_date:
        _warning(warnings, name, anchor.row, backlog_col, 'reporting_date', 'report_period_mismatch')
    for entry in result['pms']:
        number = entry['source_row']
        entry['backlog_aed'] = _numeric(sheet, source, number, backlog_col, 'backlog_aed', warnings)
        entry['source_cells']['backlog_aed'] = sheet.cell(number, backlog_col).coordinate
    total_row = anchor.row + 2 + len(result['pms'])
    result['totals']['backlog_aed'] = _numeric(sheet, source, total_row, backlog_col, 'backlog_aed', warnings)
    result['source_cells']['backlog_aed'] = sheet.cell(total_row, backlog_col).coordinate
    risk = _one_label(sheet, 'Risk')
    by_key = {entry['key']: entry for entry in result['pms']}
    for number in range(risk.row + 1, sheet.max_row + 1):
        key = _text(sheet.cell(number, risk.column - 1).value)
        if not key:
            break
        value = _numeric(sheet, source, number, risk.column, 'poc_risk_aed', warnings)
        if _normal(key) == 'total':
            result['totals']['poc_risk_aed'] = value
            result['source_cells']['poc_risk_aed'] = sheet.cell(number, risk.column).coordinate
            break
        if key in by_key:
            by_key[key]['poc_risk_aed'] = value
            by_key[key]['source_cells']['poc_risk_aed'] = sheet.cell(number, risk.column).coordinate
    return result


def _risk_history(workbook, formulas, warnings):
    name = 'KPI Trends'
    sheet = _optional_sheet(workbook, name, warnings)
    if sheet is None:
        return []
    source = formulas[name]
    headers = _headers(sheet, 2)
    dates = _column(headers, 'Date', name)
    mapping = {key: _column(headers, label, name) for key, label in {
        'poc_risk_aed': 'EDDR % vs Revenue', 'ld_exposure_aed': 'LD Exposure Risk',
        'prolongation_cost_aed': 'Prolongation cost Risk',
    }.items()}
    # The misleading EDDR % label denotes an AED exposure in this section.
    if _normal(sheet.cell(1, mapping['poc_risk_aed']).value) != 'aed':
        raise ValueError('KPI Trends must explicitly identify risk values as AED.')
    history, seen = [], set()
    for number in range(3, sheet.max_row + 1):
        cell = sheet.cell(number, dates)
        period = _date_value(cell.value)
        if period is None:
            if cell.value is not None or any(sheet.cell(number, col).value is not None for col in mapping.values()):
                _warning(warnings, name, number, dates, 'risk_history', 'invalid_or_missing_date')
            continue
        if period in seen:
            raise ValueError('KPI Trends contains duplicate risk observation dates.')
        seen.add(period)
        history.append({'date': period.isoformat(), **_facts(sheet, source, number, mapping, warnings)})
    return sorted(history, key=lambda item: item['date'])


def _poc_risk_history(workbook, formulas, warnings):
    name = 'POC RISK'
    sheet = _optional_sheet(workbook, name, warnings)
    if sheet is None:
        return []
    source = formulas[name]
    anchor = _one_label(sheet, 'POC Risk :')
    periods = [(period, cell.column) for cell in sheet[anchor.row + 1]
               if (period := _date_value(cell.value))]
    if not periods or len({period for period, _ in periods}) != len(periods):
        raise ValueError('POC RISK requires distinct dated risk columns.')
    pm_rows, seen, total_row = [], set(), None
    for number in range(anchor.row + 2, sheet.max_row + 1):
        pm = _text(sheet.cell(number, anchor.column).value)
        if _normal(pm) == 'total':
            total_row = number
            break
        if not pm or pm in seen:
            raise ValueError('POC RISK requires distinct PM labels before its total.')
        seen.add(pm)
        pm_rows.append((pm, number))
    if total_row is None:
        raise ValueError('POC RISK requires a total row.')
    return [{
        'date': period.isoformat(),
        'poc_risk_aed': _numeric(sheet, source, total_row, col, 'poc_risk_aed', warnings),
        'source_cells': {'poc_risk_aed': f'{name}!{get_column_letter(col)}{total_row}'},
        'by_pm': [{'pm': pm, 'poc_risk_aed': _numeric(sheet, source, number, col, 'poc_risk_aed', warnings),
                   'source_cells': {'poc_risk_aed': f'{name}!{get_column_letter(col)}{number}'}}
                  for pm, number in pm_rows],
    } for period, col in sorted(periods)]


def _summary_reconciliation(workbook, formulas, rows, summary, warnings):
    """Keep cached totals and detail coverage separate; never invent zeros."""
    sheet, source = workbook['POC'], formulas['POC']
    headers = _headers(sheet, 7)
    for key, label in [('ld_exposure_aed', 'LD Exposure Risk (AED)'),
                       ('prolongation_cost_aed', 'Total Prolongation cost (AED)')]:
        col = _column(headers, label, 'POC')
        summary['totals'][key] = _numeric(sheet, source, 6, col, key, warnings)
        summary['source_cells'][key] = f'POC!{get_column_letter(col)}6'
    mapping = {'actual_revenue_aed': 'period_revenue_aed', 'backlog_aed': 'backlog_without_pt_aed',
               'poc_risk_aed': 'overclaim_aed', 'ld_exposure_aed': 'ld_exposure_aed',
               'prolongation_cost_aed': 'prolongation_cost_aed'}
    checks = {}
    for key in [*mapping, 'current_forecast_aed', 'pm_forecast_aed', 'forecast_variance_aed', 'direct_cost_aed']:
        values = [row.get(mapping[key]) if key in mapping else row['extra']['executive'].get(key) for row in rows]
        known = [Decimal(str(value)) for value in values if value is not None]
        total = sum(known, Decimal(0)) if known else None
        reported = summary['totals'].get(key)
        difference = Decimal(reported) - total if reported is not None and total is not None else None
        checks[key] = {'known_total': _json_value(total), 'known_count': len(known),
                       'missing_count': len(values) - len(known), 'reported_total': reported,
                       'difference': _json_value(difference)}
        if difference is not None and abs(difference) > Decimal('0.01'):
            warnings.append({'sheet': 'Report', 'field': key, 'code': 'summary_detail_mismatch',
                             'difference': str(difference)})
    summary['detail_reconciliation'] = checks
    return summary


def read_supplemental(workbook, formulas, rows, reporting_date, warnings):
    """Populate row extras; return global summaries with explicit source units."""
    _executive_rows(workbook, formulas, rows, reporting_date, warnings)
    return {
        'invoicing': _invoicing(workbook, formulas, rows, warnings),
        'executive_summary': _summary_reconciliation(workbook, formulas, rows,
            _executive_summary(workbook, formulas, reporting_date, warnings), warnings),
        'pm_kpi': _pm_kpi(workbook, formulas, reporting_date, warnings),
        'capacity': _capacity(workbook, formulas, warnings),
        'risk_history': _risk_history(workbook, formulas, warnings),
        'poc_risk_history': _poc_risk_history(workbook, formulas, warnings),
    }
