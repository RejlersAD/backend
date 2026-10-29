from datetime import datetime, timezone
from io import BytesIO
from unittest.mock import patch
from uuid import uuid4

from django.test import SimpleTestCase, TestCase
from rest_framework.exceptions import PermissionDenied, ValidationError

from apps.sales.letter_extraction import _eml_fields, extract_letter_fields
from apps.sales.models import SalesEmailIntake
from apps.sales.tests.access_fixtures import grant_sales_actions
from apps.sales.tests import tests_letter_regeneration as letter_regression
from apps.sales.workflow import prepare_letter
from apps.sales.workspace_models import OpportunityWorkspace, OpportunityWorkspaceUpload


class ContactParsingTests(SimpleTestCase):
    def test_explicit_client_reference_is_extracted_for_outgoing_letter(self):
        fields = extract_letter_fields('Contact Name: Aisha Ali\nClient Reference: ADG-47\nEmail: aisha@client.test')
        self.assertEqual(fields['client_reference'], 'ADG-47')

    def test_internal_forward_uses_original_client_body_date_and_message_id(self):
        data = (b'From: Internal Sales <sales@rejlers.ae>\nDate: Fri, 16 Oct 2026 10:00:00 +0000\n'
                b'Message-ID: <forward@rejlers.ae>\n\nPlease prepare a response.\n'
                b'Regards,\nInternal Sales\nProposal Engineer\nRejlers\nsales@rejlers.ae\n\n'
                b'From: SAP Ariba <notification@ariba.com>\n'
                b'Sent: Wednesday, July 1, 2026 11:47 AM\n'
                b'To: Sales <sales@rejlers.ae>\nSubject: Client invitation\n\n'
                b'Example Oil Company has invited you to participate in this event.\n'
                b'Contact Name: Client Buyer\nPosition: Procurement Manager\n'
                b'Company: Example Oil Company\nEmail: buyer@client.test\nAddress: Client Tower\n\n'
                b'Example Oil Company sourcing site, Event Doc123, Message ID: MSG123456\n')
        fields = _eml_fields(BytesIO(data))
        self.assertEqual(fields['recipient_name'], 'Client Buyer')
        self.assertEqual(fields['recipient_email'], 'buyer@client.test')
        self.assertEqual(fields['recipient_title'], 'Procurement Manager')
        self.assertEqual(fields['recipient_company'], 'Example Oil Company')
        self.assertEqual(fields['your_ref'], 'SAP Ariba email notifications dated: 01,07,2026 ,Message ID :MSG123456')

    def test_vertical_contact_labels_from_readable_documents(self):
        fields = extract_letter_fields(
            'Contact Person:\nLayla Ahmed\nPosition:\nProcurement Manager\n'
            'Company Name:\nClient Engineering LLC\nEmail:\nlayla@client.test\n'
            'Address:\nTower 7, Floor 3\nP.O. Box 456\nAbu Dhabi, UAE\n'
            'Telephone: +971 2 1234567')
        self.assertEqual(fields['recipient_name'], 'Layla Ahmed')
        self.assertEqual(fields['recipient_title'], 'Procurement Manager')
        self.assertEqual(fields['recipient_company'], 'Client Engineering LLC')
        self.assertEqual(fields['recipient_email'], 'layla@client.test')
        self.assertEqual(fields['recipient_address'], 'Tower 7, Floor 3, P.O. Box 456, Abu Dhabi, UAE')

    def test_empty_vertical_label_does_not_capture_the_next_field(self):
        fields = extract_letter_fields(
            'Contact Name:\nPosition:\nProcurement Manager\nCompany:\nClient LLC\nEmail:\nbuyer@client.test')
        self.assertNotIn('recipient_name', fields)
        self.assertEqual(fields['recipient_title'], 'Procurement Manager')

    def test_word_table_with_formatted_runs_extracts_complete_contact(self):
        from docx import Document
        from apps.sales.analysis_content import extract_analysis_chunks
        document = Document()
        table = document.add_table(rows=0, cols=2)
        for label, values in [
            ('Contact Name', ['Layla ', 'Ahmed']),
            ('Position', ['Procurement ', 'Manager']),
            ('Company', ['Client Engineering LLC']),
            ('Email', ['layla', '@client.test']),
            ('Address', ['Abu Dhabi, UAE']),
        ]:
            cells = table.add_row().cells
            cells[0].text = label
            for value in values:
                cells[1].paragraphs[0].add_run(value).bold = True
        stream = BytesIO()
        document.save(stream)
        data = stream.getvalue()
        chunks, _ = extract_analysis_chunks(BytesIO(data), 'client.docx', len(data))
        fields = extract_letter_fields('\n'.join(text for _, text in chunks))
        self.assertEqual(fields['recipient_name'], 'Layla Ahmed')
        self.assertEqual(fields['recipient_title'], 'Procurement Manager')
        self.assertEqual(fields['recipient_company'], 'Client Engineering LLC')
        self.assertEqual(fields['recipient_email'], 'layla@client.test')
        self.assertEqual(fields['recipient_address'], 'Abu Dhabi, UAE')

    def test_html_table_contact_and_script_are_passively_projected(self):
        fields = extract_letter_fields('<table><tr><td>Contact Name</td><td>Table Buyer</td></tr>'
            '<tr><td>Email</td><td>buyer@client.test</td></tr></table>'
            '<script>Contact Name: Fake Person</script>')
        self.assertEqual(fields['recipient_name'], 'Table Buyer')
        self.assertEqual(fields['recipient_email'], 'buyer@client.test')

    def test_indented_document_labels(self):
        fields = extract_letter_fields('   Contact Person: Fatima Hassan\n  Designation: Buyer\n'
                                      '  Company Name: Client Ltd\n E-mail Address: fatima@client.test')
        self.assertEqual(fields['recipient_name'], 'Fatima Hassan')
        self.assertEqual(fields['recipient_email'], 'fatima@client.test')

    def test_outlook_signature_without_closing(self):
        fields = extract_letter_fields('Please submit a quotation.\n\nOmar Ali\n'
            'Purchasing Manager | Client Engineering LLC\nomar@client.test\nPO Box 123',
            sender_name='Omar Ali', sender_email='omar@client.test')
        self.assertEqual(fields['recipient_title'], 'Purchasing Manager')
        self.assertEqual(fields['recipient_company'], 'Client Engineering LLC')

    def test_contact_prose_is_not_limited_to_ariba(self):
        fields = extract_letter_fields('For clarification please contact Sara Ahmed at sara@client.test.',
                                      sender_name='Notifications', sender_email='noreply@portal.test')
        self.assertEqual(fields['recipient_name'], 'Sara Ahmed')
        self.assertEqual(fields['recipient_email'], 'sara@client.test')

    def test_contact_after_twenty_thousand_characters(self):
        fields = extract_letter_fields('Scope details.\n' * 2500 + '\nContact Name: Late Buyer\nEmail: late@client.test')
        self.assertEqual(fields['recipient_name'], 'Late Buyer')

    def test_large_email_and_empty_plain_part_html_fallback(self):
        data = (b'From: Buyer Name <buyer@client.test>\nMIME-Version: 1.0\n'
                b'Content-Type: multipart/alternative; boundary=parts\n\n'
                b'--parts\nContent-Type: text/plain\n\n\n'
                b'--parts\nContent-Type: text/html\n\n<p>Contact Name: HTML Buyer</p>'
                b'<p>Company: Client Ltd</p>' + b' ' * 1100000 + b'\n--parts--\n')
        self.assertEqual(_eml_fields(BytesIO(data))['recipient_name'], 'HTML Buyer')

    def test_ariba_event_contact_overrides_system_sender_and_supplier_login(self):
        fields = _eml_fields(BytesIO(
            b'From: Noor Buyer <system@ariba.com>\n'
            b'To: Supplier Director <supplier@example.test>\n'
            b'Date: Mon, 17 Aug 2026 04:38:30 -0700\n'
            b'Message-ID: <transport-123@ariba.com>\n\n'
            b'Example Oil Company has invited you to participate in the following event: EOI.\n\n'
            b'Use the following username to log in: supplier@example.test.\n\n'
            b'If you have questions about this event, contact Noor Buyer via telephone at  '
            b'or via e-mail at Noor.Buyer@client.test.\n\nThank You,\nExample Oil Company\n\n'
            b'Example Oil Company sourcing site, Event Doc123: EOI, Realm: example, '
            b'Message ID: MSG123456, https://example.ariba.com/event\n'))
        self.assertEqual(fields['recipient_name'], 'Noor Buyer')
        self.assertEqual(fields['recipient_email'], 'Noor.Buyer@client.test')
        self.assertEqual(fields['recipient_company'], 'Example Oil Company')
        self.assertEqual(fields['your_ref'], 'SAP Ariba email notifications dated: 17,08,2026 ,Message ID :MSG123456')
        self.assertNotIn('recipient_title', fields)
        self.assertNotIn('recipient_address', fields)

    def test_ariba_contact_labels_and_original_email_reference(self):
        fields = _eml_fields(BytesIO(
            b'From: SAP Ariba <notification@ariba.com>\r\n'
            b'To: Internal Sales <sales@example.test>\r\n'
            b'Date: Wed, 14 Oct 2026 23:30:00 +0000\r\n'
            b'Message-ID: <invitation-123@ariba.com>\r\n'
            b'Content-Type: text/plain; charset=utf-8\r\n\r\n'
            b'Contact Name: Aisha Ali\nPosition: Procurement Manager\n'
            b'Company: Client Engineering LLC\nEmail: aisha@client.test\n'
            b'Address: Tower 2, Abu Dhabi\n'))
        self.assertEqual(fields, {
            'recipient_name': 'Aisha Ali', 'recipient_title': 'Procurement Manager',
            'recipient_company': 'Client Engineering LLC', 'recipient_email': 'aisha@client.test',
            'recipient_address': 'Tower 2, Abu Dhabi',
            'your_ref': 'SAP Ariba email notifications dated: 14,10,2026 ,Message ID :<invitation-123@ariba.com>',
        })

    def test_signature_beats_automation_sender(self):
        fields = extract_letter_fields(
            'Dear Sales Team,\nPlease quote.\nBest regards,\nAisha Ali\n'
            'Procurement Manager\nClient Engineering LLC\naisha@client.test\n'
            'P.O. Box 123, Abu Dhabi', sender_name='Notifications', sender_email='robot@ariba.com')
        self.assertEqual(fields['recipient_name'], 'Aisha Ali')
        self.assertEqual(fields['recipient_email'], 'aisha@client.test')
        self.assertEqual(fields['recipient_company'], 'Client Engineering LLC')
        self.assertIn('Box 123', fields['recipient_address'])

    def test_from_fallback_never_uses_to_or_salutation(self):
        fields = _eml_fields(BytesIO(b'From: Omar Ali <omar@client.test>\n'
            b'To: Sales Director <sales@example.test>\n\nDear Sales Director,\nPlease quote.'))
        self.assertEqual(fields['recipient_name'], 'Omar Ali')
        self.assertEqual(fields['recipient_email'], 'omar@client.test')
        self.assertNotIn('recipient_title', fields)
        self.assertNotIn('your_ref', fields)

    def test_forwarded_signature_is_not_current_contact(self):
        fields = extract_letter_fields('Please quote.\n\nFrom: Older Person <old@client.test>\n'
            'Date: Tue, 1 Sep 2026 10:00:00 +0000\n\nRegards,\nOlder Person\n'
            'Director\nOld Company', sender_name='Current Person', sender_email='current@client.test')
        self.assertNotEqual(fields.get('recipient_name'), 'Older Person')

    def test_missing_values_and_invalid_date_are_not_invented(self):
        fields = extract_letter_fields('Date: unavailable\n\nPlease respond.')
        self.assertEqual(fields, {})

    def test_html_mime_preserves_header_message_id(self):
        fields = _eml_fields(BytesIO(b'From: SAP Ariba <robot@ariba.com>\n'
            b'Date: Wed, 14 Oct 2026 10:00:00 +0000\nMessage-ID: <html@ariba.com>\n'
            b'Content-Type: text/html; charset=utf-8\n\n'
            b'<p>Contact Name: Aisha Ali</p><p>Email: aisha@client.test</p>'))
        self.assertEqual(fields['recipient_name'], 'Aisha Ali')
        self.assertIn('<html@ariba.com>', fields['your_ref'])

    def test_outlook_sent_time_can_supply_reference_date(self):
        import base64
        import json
        from pathlib import Path
        import struct
        import olefile
        from apps.sales.document_classification_msg import extract_msg_text

        fixtures = json.loads(Path(__file__).with_name('fixtures').joinpath('classification-messages.json').read_text('utf-8'))
        stream = BytesIO(base64.b64decode(fixtures['plain']))
        sent = datetime(2026, 10, 14, 9, tzinfo=timezone.utc)
        ticks = int((sent - datetime(1601, 1, 1, tzinfo=timezone.utc)).total_seconds() * 10_000_000)
        with olefile.OleFileIO(stream, write_mode=True) as archive:
            props = bytearray(archive.openstream('__properties_version1.0').read())
            self.assertGreaterEqual(len(props), 48)
            offset = next((offset for offset in range(32, len(props), 16)
                           if struct.unpack_from('<I', props, offset)[0] == 0x00390040), len(props) - 16)
            struct.pack_into('<IIQ', props, offset, 0x00390040, 0, ticks)
            archive.write_stream('__properties_version1.0', bytes(props))
            for path in archive.listdir():
                if len(path) == 1 and path[0].startswith('__substg1.0_007D'):
                    transport = archive.openstream(path).read()
                    transport = transport.replace(b'Date:', b'Xate:').replace('Date:'.encode('utf-16-le'), 'Xate:'.encode('utf-16-le'))
                    archive.write_stream(path, transport)
        payload = stream.getvalue()
        text, error = extract_msg_text(BytesIO(payload), len(payload))
        self.assertEqual(error, '')
        fields = extract_letter_fields(text)
        self.assertIn('14,10,2026', fields['your_ref'])


class LetterExtractionWorkflowTests(TestCase):
    """Reuse the existing synthetic Sales owner and letter setup."""

    def setUp(self):
        letter_regression.LetterRegenerationStorageTests.setUp(self)
        grant_sales_actions(self.actor, 'sales_email_intake')
        self.deal = self.letter.opportunity
        self.intake = SalesEmailIntake.objects.create(
            opportunity=self.deal, source_message_id='intake-id', internet_message_id='<mail@ariba.com>',
            sender_name='SAP Ariba', sender_email='notification@ariba.com', subject='Invitation',
            received_at=datetime(2026, 10, 15, 10, tzinfo=timezone.utc),
            sent_at=datetime(2026, 10, 14, 9, tzinfo=timezone.utc),
            body_preview='Contact Name: Aisha Ali\nPosition: Buyer\nCompany: Client Ltd\n'
                         'Email: aisha@client.test\nAddress: Abu Dhabi',
        )

    def prepare(self, kind='eoi', custom_data=None):
        with patch('apps.sales.letter_pdf.generate_letter_pdf', return_value=b'%PDF-test'), \
                patch('apps.sales.letter_docx.generate_letter_docx', return_value=b'PK-test'):
            return prepare_letter(self.deal, self.actor, kind, custom_data)

    def test_letter_preparation_context_uses_client_reference_and_requires_authentication(self):
        from rest_framework.test import APIClient
        url = f'/api/v1/sales/deals/{self.deal.pk}/letter-preparation-context/'
        api = APIClient()
        self.assertIn(api.get(url).status_code, (401, 403))
        self.deal.client_reference = 'CLIENT-REF-47'
        self.deal.save(update_fields=['client_reference'])
        api.force_authenticate(self.actor)
        response = api.get(url)
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['deal_code'], self.deal.deal_code)
        self.assertEqual(response.data['letterhead']['your_ref'], 'CLIENT-REF-47')
        self.assertEqual(response.data['letterhead']['recipient_name'], 'Aisha Ali')
        self.assertIn('confidential_date', response.data['missing_fields'])

    def test_letter_preparation_context_prefills_detected_email_reference(self):
        from rest_framework.test import APIClient
        api = APIClient()
        api.force_authenticate(self.actor)
        response = api.get(f'/api/v1/sales/deals/{self.deal.pk}/letter-preparation-context/')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(
            response.data['letterhead']['your_ref'],
            'SAP Ariba email notifications dated: 14,10,2026 ,Message ID :<mail@ariba.com>',
        )
        self.assertNotIn('your_ref', response.data['missing_fields'])

    def test_intake_fields_are_snapshotted_and_edits_win(self):
        letter = self.prepare(custom_data={'letterhead': {'recipient_title': 'Reviewed title'}})
        head = letter.custom_data['letterhead']
        self.assertEqual(head['recipient_name'], 'Aisha Ali')
        self.assertEqual(head['recipient_title'], 'Reviewed title')
        self.assertIn('14,10,2026', head['your_ref'])
        self.assertIn('<mail@ariba.com>', head['your_ref'])
        self.assertEqual(letter.custom_data['letter_extraction']['source'], {'intake_id': str(self.intake.pk)})
        self.deal.client.refresh_from_db()
        self.assertEqual(self.deal.client.company_name, 'Letter Storage')

    def upload(self):
        workspace = OpportunityWorkspace.objects.create(opportunity=self.deal)
        return OpportunityWorkspaceUpload.objects.create(workspace=workspace, request_id=uuid4(),
            actor=self.actor, folder_key='tender', name='invitation.eml', size=100, sha256='a' * 64,
            provider='radai', status='ready')

    def test_new_source_updates_automatic_fields_without_replacing_reviewed_wording(self):
        from apps.sales.workflow import refresh_letter_files
        letter = self.prepare()
        letter.body = 'Reviewed letter body remains exactly as written.'
        letter.subject = 'Reviewed subject'
        letter.custom_data['letterhead']['recipient_title'] = 'Reviewed role'
        with patch('apps.sales.letter_pdf.generate_letter_pdf', return_value=b'%PDF-test'), \
                patch('apps.sales.letter_docx.generate_letter_docx', return_value=b'PK-test'):
            refresh_letter_files(letter, self.actor)
        previous_revision = letter.custom_data['letterhead']['rev']
        self.intake.body_preview = ('Contact Name: New Buyer\nPosition: New automatic role\n'
            'Company: New Client Ltd\nEmail: new@client.test\nAddress: New Client Tower')
        self.intake.internet_message_id = '<updated@ariba.com>'
        self.intake.save()
        with patch('apps.sales.letter_pdf.generate_letter_pdf', return_value=b'%PDF-new') as pdf, \
                patch('apps.sales.letter_docx.generate_letter_docx', return_value=b'PK-new'):
            updated = prepare_letter(self.deal, self.actor, 'eoi')
            repeated = prepare_letter(self.deal, self.actor, 'eoi')
        head = repeated.custom_data['letterhead']
        self.assertEqual(head['recipient_name'], 'New Buyer')
        self.assertEqual(head['recipient_email'], 'new@client.test')
        self.assertEqual(head['recipient_company'], 'New Client Ltd')
        self.assertEqual(head['recipient_address'], 'New Client Tower')
        self.assertEqual(head['recipient_title'], 'Reviewed role')
        self.assertIn('<updated@ariba.com>', head['your_ref'])
        self.assertEqual(repeated.body, 'Reviewed letter body remains exactly as written.')
        self.assertEqual(repeated.subject, 'Reviewed subject')
        self.assertEqual(updated.pk, letter.pk)
        self.assertEqual(repeated.pk, letter.pk)
        self.assertEqual(head['rev'], previous_revision + 1)
        self.assertEqual(pdf.call_count, 1)

    def test_current_document_version_is_extracted_with_visible_provenance(self):
        from apps.sales.workspace_models import OpportunityDocument
        from apps.sales.letter_extraction import extract_opportunity_letter_fields
        old = self.upload()
        document = OpportunityDocument.objects.create(id=uuid4(), workspace=old.workspace, folder_key='tender',
            name=old.name, root_upload=old, head_upload=old)
        old.document = document
        old.save()
        current = OpportunityWorkspaceUpload.objects.create(
            workspace=old.workspace, document=document, request_id=uuid4(), actor=self.actor,
            folder_key='tender', name=old.name, size=120, sha256='b' * 64, provider='radai',
            status='ready', version_number=2, previous_upload=old)
        document.head_upload = current
        document.save()
        with patch('apps.sales.private_attachments.download_private_file', return_value=(
                BytesIO(b'From: Current Buyer <current@client.test>\n\nPlease quote.'), current.name)) as download:
            fields = extract_opportunity_letter_fields(self.deal, self.actor)
        self.assertEqual(fields['recipient_name'], 'Current Buyer')
        self.assertEqual(fields['_source']['name'], current.name)
        self.assertEqual(fields['_source']['version'], 2)
        self.assertEqual(fields['_source']['document_id'], str(document.pk))
        self.assertEqual(download.call_args.args[-1], 'radai-' + str(current.pk))

    def test_new_uploaded_version_replaces_previous_version_contact_in_same_draft(self):
        from apps.sales.workspace_models import OpportunityDocument
        old = self.upload()
        document = OpportunityDocument.objects.create(id=uuid4(), workspace=old.workspace, folder_key='tender',
            name=old.name, root_upload=old, head_upload=old)
        old.document = document
        old.save()
        old_mail = (b'From: Old Buyer <old@client.test>\nDate: Tue, 13 Oct 2026 10:00:00 +0000\n'
                    b'Message-ID: <old@client.test>\n\nPlease quote.')
        with patch('apps.sales.private_attachments.download_private_file',
                   side_effect=lambda *args, **kwargs: (BytesIO(old_mail), old.name)):
            first = self.prepare()
        self.assertEqual(first.custom_data['letterhead']['recipient_name'], 'Old Buyer')
        current = OpportunityWorkspaceUpload.objects.create(
            workspace=old.workspace, document=document, request_id=uuid4(), actor=self.actor,
            folder_key='tender', name=old.name, size=120, sha256='b' * 64, provider='radai',
            status='ready', version_number=2, previous_upload=old)
        document.head_upload = current
        document.save()
        new_mail = (b'From: New Buyer <new@client.test>\nDate: Fri, 16 Oct 2026 10:00:00 +0000\n'
                    b'Message-ID: <new@client.test>\n\nPlease quote.')
        with patch('apps.sales.private_attachments.download_private_file',
                   side_effect=lambda *args, **kwargs: (BytesIO(new_mail), current.name)):
            updated = self.prepare()
        head = updated.custom_data['letterhead']
        self.assertEqual(updated.pk, first.pk)
        self.assertEqual(head['recipient_name'], 'New Buyer')
        self.assertEqual(head['recipient_email'], 'new@client.test')
        self.assertIn('16,10,2026', head['your_ref'])
        self.assertNotIn('13,10,2026', head['your_ref'])
        self.assertEqual(updated.custom_data['letter_extraction']['source']['version'], 2)

    def test_new_source_clears_obsolete_automatic_optional_fields(self):
        first = self.prepare()
        self.intake.body_preview = 'Contact Name: Aisha Ali\nCompany: Client Ltd\nEmail: aisha@client.test'
        self.intake.save()
        with patch('apps.sales.letter_pdf.generate_letter_pdf', return_value=b'%PDF-test') as pdf, \
                patch('apps.sales.letter_docx.generate_letter_docx', return_value=b'PK-test'):
            updated = prepare_letter(self.deal, self.actor, 'eoi')
            repeated = prepare_letter(self.deal, self.actor, 'eoi')
        self.assertEqual(updated.pk, first.pk)
        self.assertEqual(repeated.custom_data['letterhead']['recipient_title'], '')
        self.assertEqual(repeated.custom_data['letterhead']['recipient_address'], '')
        self.assertEqual(pdf.call_count, 1)

    def test_updated_source_fields_appear_in_real_pdf_and_word_output(self):
        import fitz
        from docx import Document
        from apps.sales.letter_pdf import generate_letter_pdf
        from apps.sales.letter_docx import generate_letter_docx
        self.prepare()
        self.intake.body_preview = ('Contact Name: Current Client Buyer\nPosition: Procurement Manager\n'
            'Company: Current Client Ltd\nEmail: current@client.test\nAddress: Current Client Tower')
        self.intake.internet_message_id = '<current-source@ariba.com>'
        self.intake.save()
        letter = self.prepare()
        with fitz.open(stream=generate_letter_pdf(letter), filetype='pdf') as pdf:
            pdf_text = '\n'.join(page.get_text() for page in pdf)
        word = Document(BytesIO(generate_letter_docx(letter)))
        word_text = '\n'.join(
            [paragraph.text for paragraph in word.paragraphs]
            + [cell.text for table in word.tables for row in table.rows for cell in row.cells])
        for text in (pdf_text, word_text):
            self.assertIn('Current Client Buyer', text)
            self.assertIn('current@client.test', text)
            self.assertIn('Current Client Ltd', text)
            self.assertIn('Current Client Tower', text)
            self.assertNotIn('Aisha Ali', text)

    def test_newest_incoming_email_is_preferred_over_an_older_upload(self):
        from apps.sales.letter_extraction import extract_opportunity_letter_fields
        old = self.upload()
        current = OpportunityWorkspaceUpload.objects.create(
            workspace=old.workspace, request_id=uuid4(), actor=self.actor, folder_key='tender',
            name='latest-invitation.eml', size=120, sha256='b' * 64, provider='radai', status='ready')
        with patch('apps.sales.private_attachments.download_private_file', return_value=(
                BytesIO(b'From: Current Buyer <current@client.test>\n\nPlease quote.'), current.name)) as download:
            fields = extract_opportunity_letter_fields(self.deal, self.actor)
        self.assertEqual(fields['_source']['upload_id'], str(current.pk))
        self.assertEqual(download.call_count, 1)

    def test_manual_contact_identity_is_not_mixed_with_a_different_source_person(self):
        from apps.sales.workflow import refresh_letter_files
        letter = self.prepare()
        letter.custom_data['letterhead']['recipient_name'] = 'Reviewed Buyer'
        letter.custom_data['letterhead']['recipient_email'] = 'reviewed@client.test'
        with patch('apps.sales.letter_pdf.generate_letter_pdf', return_value=b'%PDF-test'), \
                patch('apps.sales.letter_docx.generate_letter_docx', return_value=b'PK-test'):
            refresh_letter_files(letter, self.actor)
        self.intake.body_preview = ('Contact Name: Other Buyer\nPosition: Different Role\nCompany: Other Company\n'
                                    'Email: other@client.test\nAddress: Other Address')
        self.intake.save()
        with patch('apps.sales.letter_pdf.generate_letter_pdf') as pdf:
            reopened = prepare_letter(self.deal, self.actor, 'eoi')
        self.assertEqual(reopened.custom_data['letterhead']['recipient_name'], 'Reviewed Buyer')
        self.assertEqual(reopened.custom_data['letterhead']['recipient_email'], 'reviewed@client.test')
        self.assertNotEqual(reopened.custom_data['letterhead']['recipient_company'], 'Other Company')
        self.assertIn('recipient_name', reopened.custom_data['letter_extraction']['retained_fields'])
        pdf.assert_not_called()
    def test_uploaded_email_is_used_and_stream_closed(self):
        self.upload()
        stream = BytesIO(b'From: Uploaded Person <upload@client.test>\n'
            b'Date: Tue, 13 Oct 2026 10:00:00 +0000\nMessage-ID: <uploaded@client.test>\n\nPlease quote.')
        with patch('apps.sales.private_attachments.download_private_file', return_value=(stream, 'invitation.eml')):
            letter = self.prepare()
        self.assertEqual(letter.custom_data['letterhead']['recipient_name'], 'Uploaded Person')
        self.assertIn('13,10,2026', letter.custom_data['letterhead']['your_ref'])
        self.assertTrue(stream.closed)

    def test_email_after_twenty_other_documents_is_still_selected(self):
        upload = self.upload()
        OpportunityWorkspaceUpload.objects.filter(pk=upload.pk).update(
            created_at=datetime(2026, 12, 1, tzinfo=timezone.utc))
        OpportunityWorkspaceUpload.objects.bulk_create([
            OpportunityWorkspaceUpload(workspace=upload.workspace, request_id=uuid4(), actor=self.actor,
                folder_key='tender', name=f'scope-{index}.pdf', size=100, sha256='b' * 64,
                provider='radai', status='ready') for index in range(25)])
        stream = BytesIO(b'From: Uploaded Person <upload@client.test>\n\nPlease quote.')
        with patch('apps.sales.private_attachments.download_private_file', return_value=(stream, upload.name)) as download:
            letter = self.prepare()
        self.assertEqual(letter.custom_data['letterhead']['recipient_email'], 'upload@client.test')
        self.assertEqual(download.call_count, 1)

    def test_reopening_a_legacy_blank_draft_applies_extraction_and_preserves_reviewed_body(self):
        from django.core.files.base import ContentFile
        self.letter.pdf_file.save('legacy.pdf', ContentFile(b'%PDF-old'), save=False)
        self.letter.docx_file.save('legacy.docx', ContentFile(b'PK-old'), save=False)
        self.letter.body = 'Keep this reviewed body exactly.'
        self.letter.subject = 'Reviewed subject'
        self.letter.custom_data = {'letterhead': {'recipient_name': '  ', 'recipient_email': '',
            'recipient_company': 'Reviewed company', 'recipient_title': '', 'recipient_address': '',
            'your_ref': '', 'rev': 2, 'our_ref': 'KEEP-REF'}}
        self.letter.save()
        with patch('apps.sales.letter_pdf.generate_letter_pdf', return_value=b'%PDF-new') as pdf, \
                patch('apps.sales.letter_docx.generate_letter_docx', return_value=b'PK-new'):
            letter = prepare_letter(self.deal, self.actor, 'eoi')
            repeated = prepare_letter(self.deal, self.actor, 'eoi')
        self.assertEqual(letter.pk, self.letter.pk)
        self.assertEqual(repeated.pk, self.letter.pk)
        self.assertEqual(letter.body, 'Keep this reviewed body exactly.')
        self.assertEqual(letter.subject, 'Reviewed subject')
        head = repeated.custom_data['letterhead']
        self.assertEqual(head['recipient_name'], 'Aisha Ali')
        self.assertEqual(head['recipient_email'], 'aisha@client.test')
        self.assertEqual(head['recipient_title'], 'Buyer')
        self.assertEqual(head['recipient_address'], 'Abu Dhabi')
        self.assertEqual(head['recipient_company'], 'Reviewed company')
        self.assertEqual(head['our_ref'], 'KEEP-REF')
        self.assertEqual(head['rev'], 3)
        self.assertIn('<mail@ariba.com>', head['your_ref'])
        self.assertEqual(pdf.call_count, 1)

    def test_intentionally_cleared_recipient_field_is_not_refilled_on_reopen(self):
        from apps.sales.workflow import refresh_letter_files
        letter = self.prepare()
        letter.custom_data['letterhead']['recipient_title'] = ''
        with patch('apps.sales.letter_pdf.generate_letter_pdf', return_value=b'%PDF-test'), \
                patch('apps.sales.letter_docx.generate_letter_docx', return_value=b'PK-test'):
            refresh_letter_files(letter, self.actor)
            reopened = prepare_letter(self.deal, self.actor, 'eoi')
        self.assertEqual(reopened.custom_data['letterhead']['recipient_title'], '')
        self.assertIn('recipient_title', reopened.custom_data['letterhead_manual_fields'])

    def test_prepare_api_enriches_reused_draft_from_uploaded_email_once(self):
        from django.core.files.base import ContentFile
        from rest_framework.test import APIClient
        upload = self.upload()
        self.letter.pdf_file.save('legacy-api.pdf', ContentFile(b'%PDF-old'), save=False)
        self.letter.docx_file.save('legacy-api.docx', ContentFile(b'PK-old'), save=False)
        self.letter.body = 'Reviewed body must survive extraction.'
        self.letter.custom_data = {'letterhead': {'recipient_name': '', 'recipient_email': '',
            'recipient_company': self.deal.client.company_name, 'your_ref': ''}}
        self.letter.save()
        data = (b'From: Buyer Name <system@ariba.com>\nDate: Mon, 17 Aug 2026 04:38:30 -0700\n\n'
            b'Example Oil Company has invited you to participate in the following event: EOI.\n'
            b'Contact Buyer Name via e-mail at buyer@client.test.\n'
            b'Example Oil Company sourcing site, Event Doc123, Message ID: MSG123456\n')
        api = APIClient()
        api.force_authenticate(self.actor)
        url = f'/api/v1/sales/deals/{self.deal.pk}/prepare-letter/'
        with patch('apps.sales.private_attachments.download_private_file', side_effect=lambda *args: (BytesIO(data), upload.name)), \
                patch('apps.sales.letter_pdf.generate_letter_pdf', return_value=b'%PDF-new') as pdf, \
                patch('apps.sales.letter_docx.generate_letter_docx', return_value=b'PK-new'), \
                patch('apps.sales.views.DealViewSet._attach_letter_files'):
            first = api.post(url, {'letter_type': 'eoi'}, format='json')
            second = api.post(url, {'letter_type': 'eoi'}, format='json')
        self.assertEqual(first.status_code, 200, first.data)
        self.assertEqual(second.status_code, 200, second.data)
        self.assertEqual(first.data['id'], second.data['id'])
        self.assertEqual(first.data['body'], 'Reviewed body must survive extraction.')
        self.assertEqual(first.data['custom_data']['letterhead']['recipient_name'], 'Buyer Name')
        self.assertEqual(first.data['custom_data']['letterhead']['recipient_company'], 'Example Oil Company')
        self.assertEqual(first.data['custom_data']['letterhead']['recipient_email'], 'buyer@client.test')
        self.assertIn('MSG123456', first.data['custom_data']['letterhead']['your_ref'])
        self.assertEqual(pdf.call_count, 1)

    def test_prepare_uses_ariba_contact_and_company_instead_of_saved_client(self):
        self.upload()
        stream = BytesIO(b'From: Noor Buyer <system@ariba.com>\n'
            b'Date: Mon, 17 Aug 2026 04:38:30 -0700\nMessage-ID: <transport@ariba.com>\n\n'
            b'Example Oil Company has invited you to participate in the following event: EOI.\n'
            b'Use the following username to log in: supplier@example.test.\n'
            b'If you have questions about this event, contact Noor Buyer via telephone at '
            b'or via e-mail at Noor.Buyer@client.test.\n'
            b'Example Oil Company sourcing site, Event Doc123: EOI, Message ID: MSG123456\n')
        with patch('apps.sales.private_attachments.download_private_file', return_value=(stream, 'invitation.eml')):
            letter = self.prepare()
        head = letter.custom_data['letterhead']
        self.assertEqual(head['recipient_name'], 'Noor Buyer')
        self.assertEqual(head['recipient_email'], 'Noor.Buyer@client.test')
        self.assertEqual(head['recipient_company'], 'Example Oil Company')
        self.assertIn('MSG123456', head['your_ref'])
        self.assertTrue(letter.custom_data['letter_extraction']['source']['upload_id'])

    def test_denied_file_falls_back_to_intake(self):
        self.upload()
        with patch('apps.sales.private_attachments.download_private_file', side_effect=PermissionDenied):
            letter = self.prepare()
        self.assertEqual(letter.custom_data['letterhead']['recipient_name'], 'Aisha Ali')

    def test_email_read_denial_does_not_extract_intake_body(self):
        self.letter.custom_data = {}
        self.letter.save()
        with patch('apps.sales.letter_extraction.module_action_allowed', return_value=False):
            with self.assertRaises(ValidationError):
                self.prepare()

    def test_missing_recipient_api_blocks_all_types_without_render_or_attachment(self):
        from apps.sales.models import SalesLetter, SalesLetterTemplate
        from rest_framework.test import APIClient
        self.intake.body_preview = 'Please quote.'
        self.intake.save()
        self.letter.custom_data = {}
        self.letter.save()
        api = APIClient()
        api.force_authenticate(self.actor)
        for kind in ('eoi', 'regret_expertise', 'regret_manpower'):
            self.deal.stage, self.deal.bid_decision = ('proposal', 'bid') if kind == 'eoi' else ('no_bid', 'no_bid')
            self.deal.save()
            SalesLetterTemplate.objects.update_or_create(letter_type=kind,
                defaults={'subject_template': 'Subject', 'body_template': 'Body'})
            count = SalesLetter.objects.count()
            with patch('apps.sales.letter_pdf.generate_letter_pdf') as pdf, \
                    patch('apps.sales.letter_docx.generate_letter_docx') as docx, \
                    patch('apps.sales.views.DealViewSet._attach_letter_files') as attach:
                response = api.post(f'/api/v1/sales/deals/{self.deal.pk}/prepare-letter/',
                                    {'letter_type': kind}, format='json')
            self.assertEqual(response.status_code, 400, response.data)
            self.assertEqual(str(response.data['code']), 'recipient_empty')
            self.assertIn('Recipient is empty', str(response.data['detail']))
            self.assertEqual(SalesLetter.objects.count(), count)
            pdf.assert_not_called()
            docx.assert_not_called()
            attach.assert_not_called()

    def test_notification_source_is_skipped_for_later_named_contact(self):
        self.upload()
        with patch('apps.sales.private_attachments.download_private_file', return_value=(
                BytesIO(b'From: Notifications <noreply@portal.test>\n\nPlease quote.'), 'invitation.eml')):
            letter = self.prepare()
        self.assertEqual(letter.custom_data['letterhead']['recipient_name'], 'Aisha Ali')

    def test_sharepoint_upload_uses_authorized_workspace_download(self):
        upload = self.upload()
        upload.provider = 'sharepoint'
        upload.result = {'id': 'remote-file-id'}
        upload.save()
        with patch('apps.sales.opportunity_workspace.download_workspace_file', return_value=(
                BytesIO(b'From: Remote Buyer <remote@client.test>\n\nPlease quote.'), upload.name)) as download:
            letter = self.prepare()
        self.assertEqual(letter.custom_data['letterhead']['recipient_name'], 'Remote Buyer')
        download.assert_called_once_with(self.deal, self.actor, 'tender', 'remote-file-id')

    def test_readable_pdf_docx_and_html_uploads(self):
        import fitz
        from docx import Document
        upload = self.upload()
        text = 'Contact Name: Document Buyer\nPosition: Procurement Manager\nCompany: Example Ltd\nEmail: buyer@example.test\nAddress: PO Box 123'
        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((72, 72), text)
        pdf_data = pdf.tobytes()
        pdf.close()
        docx = Document()
        for line in text.splitlines():
            docx.add_paragraph(line)
        buffer = BytesIO()
        docx.save(buffer)
        for suffix, data in (('.pdf', pdf_data), ('.docx', buffer.getvalue()),
                             ('.html', ('<p>' + text.replace('\n', '</p><p>') + '</p>').encode())):
            with self.subTest(suffix=suffix):
                upload.name, upload.size = 'invitation' + suffix, len(data)
                upload.save()
                with patch('apps.sales.private_attachments.download_private_file', return_value=(BytesIO(data), upload.name)):
                    letter = self.prepare()
                self.assertEqual(letter.custom_data['letterhead']['recipient_name'], 'Document Buyer')
                self.assertEqual(letter.custom_data['letterhead']['recipient_email'], 'buyer@example.test')

    def test_vertical_upload_fields_populate_all_letter_types_and_reopen_without_render(self):
        import fitz
        from docx import Document
        from apps.sales.models import SalesLetterTemplate
        text = ('Contact Name:\nDocument Buyer\nPosition:\nProcurement Manager\n'
                'Company:\nExample Ltd\nEmail:\nbuyer@example.test\n'
                'Address:\nTower 7\nAbu Dhabi, UAE')
        pdf = fitz.open()
        pdf.new_page().insert_text((72, 72), text)
        pdf_data = pdf.tobytes()
        pdf.close()
        document = Document()
        for line in text.splitlines():
            document.add_paragraph(line)
        buffer = BytesIO()
        document.save(buffer)
        fixtures = [
            ('.pdf', pdf_data), ('.docx', buffer.getvalue()), ('.txt', text.encode()),
            ('.html', ('<p>' + text.replace('\n', '</p><p>') + '</p>').encode()),
            ('.eml', b'From: Notifications <noreply@portal.test>\n\n' + text.encode()),
        ]
        upload = self.upload()
        for kind in ('eoi', 'regret_expertise', 'regret_manpower'):
            self.deal.stage, self.deal.bid_decision = ('proposal', 'bid') if kind == 'eoi' else ('no_bid', 'no_bid')
            self.deal.save()
            SalesLetterTemplate.objects.update_or_create(letter_type=kind,
                defaults={'subject_template': 'Subject', 'body_template': 'Dear {{ client_contact_name }}'})
            for suffix, data in fixtures:
                with self.subTest(kind=kind, suffix=suffix):
                    self.deal.letters.all().delete()
                    upload.name, upload.size = 'client' + suffix, len(data)
                    upload.save()
                    with patch('apps.sales.private_attachments.download_private_file',
                               side_effect=lambda *args: (BytesIO(data), upload.name)), \
                            patch('apps.sales.letter_pdf.generate_letter_pdf', return_value=b'%PDF-test') as render, \
                            patch('apps.sales.letter_docx.generate_letter_docx', return_value=b'PK-test'):
                        letter = prepare_letter(self.deal, self.actor, kind)
                        repeated = prepare_letter(self.deal, self.actor, kind)
                    head = repeated.custom_data['letterhead']
                    self.assertEqual(head['recipient_name'], 'Document Buyer')
                    self.assertEqual(head['recipient_title'], 'Procurement Manager')
                    self.assertEqual(head['recipient_company'], 'Example Ltd')
                    self.assertEqual(head['recipient_email'], 'buyer@example.test')
                    self.assertEqual(head['recipient_address'], 'Tower 7, Abu Dhabi, UAE')
                    self.assertEqual(letter.pk, repeated.pk)
                    self.assertEqual(render.call_count, 1)

    def test_name_only_contact_is_allowed_without_inventing_optional_fields(self):
        self.intake.body_preview = 'Contact Name: Named Buyer'
        self.intake.save()
        letter = self.prepare()
        self.assertEqual(letter.custom_data['letterhead']['recipient_name'], 'Named Buyer')

    def test_regeneration_cannot_clear_required_recipient(self):
        from apps.sales.workflow import refresh_letter_files
        letter = self.prepare()
        letter.custom_data['letterhead']['recipient_name'] = ' '
        with patch('apps.sales.letter_pdf.generate_letter_pdf') as pdf:
            with self.assertRaises(ValidationError):
                refresh_letter_files(letter, self.actor)
        pdf.assert_not_called()

    def test_on_demand_file_creation_also_requires_recipient(self):
        from apps.sales.letter_pdf import save_letter_pdf
        from apps.sales.letter_docx import save_letter_docx
        self.letter.custom_data = {}
        with patch('apps.sales.letter_pdf.generate_letter_pdf') as pdf, \
                patch('apps.sales.letter_docx.generate_letter_docx') as docx:
            for save in (save_letter_pdf, save_letter_docx):
                with self.assertRaises(ValidationError):
                    save(self.letter)
        pdf.assert_not_called()
        docx.assert_not_called()

    def test_both_regret_types_use_same_extraction(self):
        from apps.sales.models import SalesLetterTemplate
        self.deal.stage, self.deal.bid_decision = 'no_bid', 'no_bid'
        self.deal.save()
        for kind in ('regret_expertise', 'regret_manpower'):
            SalesLetterTemplate.objects.update_or_create(letter_type=kind,
                defaults={'subject_template': 'Regret', 'body_template': 'Dear {{ client_contact_name }}'})
            letter = self.prepare(kind)
            self.assertEqual(letter.body, 'Dear Aisha Ali')

    def test_both_regret_types_enrich_existing_blank_drafts(self):
        from django.core.files.base import ContentFile
        from apps.sales.models import SalesLetter, SalesLetterTemplate
        self.deal.stage, self.deal.bid_decision = 'no_bid', 'no_bid'
        self.deal.save()
        for kind in ('regret_expertise', 'regret_manpower'):
            with self.subTest(kind=kind):
                template, _ = SalesLetterTemplate.objects.update_or_create(letter_type=kind,
                    defaults={'subject_template': 'Regret', 'body_template': 'Default body'})
                legacy = SalesLetter.objects.create(opportunity=self.deal, letter_type=kind, template=template,
                    subject='Reviewed regret', body='Preserve reviewed regret wording.', generated_by=self.actor,
                    custom_data={'letterhead': {'recipient_name': '', 'recipient_email': '', 'your_ref': ''}})
                legacy.pdf_file.save(f'{kind}-old.pdf', ContentFile(b'%PDF-old'), save=False)
                legacy.docx_file.save(f'{kind}-old.docx', ContentFile(b'PK-old'), save=False)
                legacy.save()
                letter = self.prepare(kind)
                repeated = self.prepare(kind)
                self.assertEqual(letter.pk, legacy.pk)
                self.assertEqual(repeated.pk, legacy.pk)
                self.assertEqual(letter.body, 'Preserve reviewed regret wording.')
                self.assertEqual(letter.custom_data['letterhead']['recipient_email'], 'aisha@client.test')
                self.assertEqual(letter.custom_data['letterhead']['recipient_name'], 'Aisha Ali')
