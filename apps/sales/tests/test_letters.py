"""Sales letters: standard letterhead context, PDF/DOCX output, Correspondence attach."""
from datetime import datetime, timedelta
from importlib import import_module
from tempfile import TemporaryDirectory
from unittest import skipUnless
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.utils import timezone

from apps.sales.letter_attachments import attach_letter_files
from apps.sales.letter_context import (
    build_letterhead, format_date_ordinal, letter_salutation, normalize_body_text, split_body_paragraphs,
)
from apps.sales.letter_docx import generate_letter_docx, get_letter_docx_bytes
from apps.sales.models import (
    Client, Contact, Deal, OpportunityAuditEvent, SalesEmailIntake, SalesLetter,
    SalesLetterTemplate,
)
from apps.sales.opportunity_workspace import WorkspaceAPIError
from apps.sales.tests.access_fixtures import grant_sales_actions
from apps.sales.workflow import prepare_letter, refresh_letter_files

try:
    import weasyprint  # noqa: F401
    WEASYPRINT_AVAILABLE = True
except ImportError:
    WEASYPRINT_AVAILABLE = False

try:
    import docx  # noqa: F401
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


def _weasyprint_available():
    return WEASYPRINT_AVAILABLE


def _docx_available():
    return DOCX_AVAILABLE


def seed_test_letter_templates():
    templates = import_module('apps.sales.migrations.0028_update_letter_templates').TEMPLATES
    for letter_type, content in templates.items():
        SalesLetterTemplate.objects.update_or_create(
            letter_type=letter_type, defaults={**content, 'is_active': True},
        )


class LetterContextTests(TestCase):
    def setUp(self):
        seed_test_letter_templates()
        User = get_user_model()
        self.actor = User.objects.create_user(username='letter-owner', email='owner@example.test')
        grant_sales_actions(self.actor, 'sales', 'sales_opportunities', 'sales_clients')
        self.client_record = Client.objects.create(
            client_code='L-CLIENT', company_name='Abu Dhabi National Oil Company',
            industry_type='oil_gas', address='Corniche Road', city='Abu Dhabi',
            country='United Arab Emirates', account_manager=self.actor,
        )
        self.contact = Contact.objects.create(
            client=self.client_record, first_name='Ridhesh', last_name='Rajan',
            job_title='Senior Contract Specialist', email='ridhesh.rajan@adnoc.ae',
        )
        self.opportunity = Deal.objects.create(
            deal_code='Q-101883', deal_name='EOI for Conceptual Facility and Infrastructure Design Study',
            client=self.client_record, client_contact=self.contact, owner=self.actor,
            stage='proposal', bid_decision='bid', client_reference='TENDER-2026-001',
        )

    def test_format_date_ordinal_uses_company_style(self):
        self.assertEqual(format_date_ordinal(datetime(2026, 6, 10)), '10th June 2026')

    def test_normalize_body_text_strips_editor_html(self):
        html = '<p>Dear Sir,</p><p>First <strong>paragraph</strong>.</p><p>Second paragraph.</p>'
        self.assertEqual(
            normalize_body_text(html),
            'Dear Sir,\nFirst paragraph.\nSecond paragraph.',
        )
        self.assertEqual(
            split_body_paragraphs(html),
            ['First paragraph.', 'Second paragraph.'],
        )

    def test_salutation_with_body_in_same_plain_text_paragraph_is_rendered_once(self):
        self.assertEqual(
            split_body_paragraphs('Dear Sir,\nProject details.'),
            ['Project details.'],
        )
        self.assertEqual(letter_salutation('Dear Sir / Madam,\nProject details.'), 'Dear Sir / Madam,')
        self.assertEqual(split_body_paragraphs('Dear Sir / Madam,\nProject details.'), ['Project details.'])

    def test_closing_with_signature_in_same_paragraph_is_preserved(self):
        self.assertEqual(
            split_body_paragraphs('Project details.\n\nSincerely yours,\nSender'),
            ['Project details.', 'Sincerely yours,\nSender'],
        )

    def test_letterhead_builds_standard_fields_from_records(self):
        received = timezone.now() - timedelta(days=5)
        SalesEmailIntake.objects.create(
            source_message_id='MSG65153313',
            subject='SAP Ariba Email Notifications',
            sender_name='SAP Ariba Email Notifications',
            sender_email='notification@ariba.com',
            received_at=received,
            opportunity=self.opportunity,
        )
        head = build_letterhead(self.opportunity)
        self.assertEqual(head['response_label'], 'EOI Response')
        self.assertEqual(head['response_code'], 'Q-101883')
        self.assertTrue(head['confidential'].endswith('/ Rev 0'))
        self.assertEqual(head['sender_name'], 'Mr. Jarmo Suominen')
        self.assertEqual(head['focal_contact'], 'Ms. Anam Abbas')
        self.assertEqual(head['email'], 'uae.sales@rejlers.ae')
        self.assertEqual(head['adnoc_unified_code'], '0010001089')
        self.assertEqual(head['recipient_name'], 'Ridhesh Rajan')
        self.assertEqual(head['recipient_title'], 'Senior Contract Specialist')
        self.assertEqual(head['recipient_company'], 'Abu Dhabi National Oil Company')
        self.assertIn('Abu Dhabi', head['recipient_address'])
        self.assertEqual(head['recipient_email'], 'ridhesh.rajan@adnoc.ae')
        self.assertEqual(head['your_ref'], 'TENDER-2026-001')
        self.assertEqual(
            head['our_ref'],
            'LR-RAD-Abu Dhabi National Oil Company-001 / Q-101883',
        )
        self.assertEqual(head['signature_name'], 'Jarmo Suominen')
        self.assertIn('CEO, Rejlers Abu Dhabi', head['signature_titles'])

    def test_letterhead_falls_back_to_client_reference_without_intake(self):
        head = build_letterhead(self.opportunity)
        self.assertEqual(head['your_ref'], 'TENDER-2026-001')

    def test_letterhead_keeps_edited_values_from_custom_data(self):
        head = build_letterhead(self.opportunity, {
            'letterhead': {'recipient_name': 'Edited Name', 'rev': 3},
        })
        self.assertEqual(head['recipient_name'], 'Edited Name')
        self.assertEqual(head['rev'], 3)
        self.assertTrue(head['confidential'].endswith('/ Rev 3'))
        # Untouched fields are still computed from the records.
        self.assertEqual(head['recipient_company'], 'Abu Dhabi National Oil Company')

    def test_letterhead_sequence_increments_per_letter(self):
        head = build_letterhead(self.opportunity)
        self.assertIn('-001 /', head['our_ref'])
        SalesLetter.objects.create(
            opportunity=self.opportunity, letter_type='eoi',
            template=SalesLetterTemplate.objects.get(letter_type='eoi'),
            subject='First', body='Body',
        )
        head = build_letterhead(self.opportunity)
        self.assertIn('-002 /', head['our_ref'])


class LetterGenerationTests(TestCase):
    def setUp(self):
        seed_test_letter_templates()
        User = get_user_model()
        self.actor = User.objects.create_user(username='letter-owner', email='owner@example.test')
        grant_sales_actions(self.actor, 'sales', 'sales_opportunities', 'sales_clients')
        self.client_record = Client.objects.create(
            client_code='L-CLIENT', company_name='Abu Dhabi National Oil Company',
            industry_type='oil_gas', account_manager=self.actor,
        )
        self.contact = Contact.objects.create(
            client=self.client_record, first_name='Sample', last_name='Recipient',
            email='recipient@example.test',
        )
        self.opportunity = Deal.objects.create(
            deal_code='Q-101883', deal_name='Conceptual Facility and Infrastructure Design Study',
            client=self.client_record, client_contact=self.contact, owner=self.actor,
            stage='proposal', bid_decision='bid',
        )

    def prepare(self, letter_type):
        return prepare_letter(self.opportunity, self.actor, letter_type,
                              {'letterhead': {'recipient_name': 'Sample Recipient'}})

    def test_prepare_letter_persists_letterhead_and_audits(self):
        letter = self.prepare('eoi')
        self.assertEqual(letter.subject, 'Conceptual Facility and Infrastructure Design Study')
        self.assertIn('Thank you for your interest', letter.body)
        head = letter.custom_data['letterhead']
        self.assertEqual(head['response_code'], 'Q-101883')
        self.assertEqual(head['rev'], 0)
        event = OpportunityAuditEvent.objects.get(event_type='letter_generated')
        self.assertEqual(event.data['letter_type'], 'eoi')

    def test_prepare_letter_regret_requires_no_bid(self):
        self.opportunity.stage = 'no_bid'
        self.opportunity.bid_decision = 'no_bid'
        self.opportunity.save(update_fields=['stage', 'bid_decision'])
        letter = self.prepare('regret_expertise')
        self.assertIn('area of expertise', letter.body)
        head = letter.custom_data['letterhead']
        self.assertEqual(head['confidential'].split(' / ')[-1], 'Rev 0')

    def test_refresh_letter_files_bumps_revision(self):
        letter = self.prepare('eoi')
        letter.body += '\n\nReviewed clarification.'
        refresh_letter_files(letter, self.actor)
        letter.refresh_from_db()
        self.assertEqual(letter.custom_data['letterhead']['rev'], 1)
        self.assertTrue(letter.custom_data['letterhead']['confidential'].endswith('/ Rev 1'))
        self.assertTrue(
            OpportunityAuditEvent.objects.filter(event_type='letter_pdf_regenerated').exists()
        )

    @skipUnless(_docx_available(), 'python-docx not installed')
    def test_docx_generation_produces_word_document(self):
        from io import BytesIO
        from docx import Document
        letter = self.prepare('eoi')
        docx_bytes = generate_letter_docx(letter)
        self.assertTrue(docx_bytes.startswith(b'PK'))
        self.assertGreater(len(docx_bytes), 2000)
        again = get_letter_docx_bytes(letter)
        self.assertTrue(again.startswith(b'PK'))
        letter.body = 'Dear Sir / Madam,\n\nUpdated draft wording.'
        updated = Document(BytesIO(get_letter_docx_bytes(letter)))
        self.assertIn('Updated draft wording.', [p.text for p in updated.paragraphs])

    @skipUnless(_weasyprint_available(), 'weasyprint not installed')
    def test_pdf_generation_renders_pdf_document(self):
        from apps.sales.letter_pdf import generate_letter_pdf
        import pymupdf
        letter = self.prepare('eoi')
        letter.body = 'Dear Sir / Madam,\n\nProject details.\n\nSincerely yours,'
        pdf_bytes = generate_letter_pdf(letter)
        self.assertTrue(pdf_bytes.startswith(b'%PDF'))
        self.assertGreater(len(pdf_bytes), 1000)
        document = pymupdf.open(stream=pdf_bytes, filetype='pdf')
        page = document[0]
        self.assertEqual(len(page.search_for('Dear Sir / Madam,')), 1)
        self.assertEqual(len(page.search_for('Sincerely yours,')), 1)
        self.assertLess(page.search_for('Subject:')[0].x0, page.search_for(letter.subject)[0].x0)
        self.assertGreater(page.search_for('Rejlers International Engineering Solutions AB')[-1].y0, 750)
        self.assertEqual(len(page.search_for('1 (1)')), 1)
        confidential = page.search_for('Confidential')[0]
        sender = page.search_for('From:')[0]
        self.assertLess(sender.y0 - confidential.y1, 32)
        recipient_email = page.search_for(letter.custom_data['letterhead']['recipient_email'])[0]
        your_reference = page.search_for('Your Ref. No.:')[0]
        self.assertLess(your_reference.y0 - recipient_email.y1, 45)
        signature = page.search_for('CEO, Rejlers Abu Dhabi')[0]
        footer_company = page.search_for('Rejlers International Engineering Solutions AB')[-1]
        self.assertLess(signature.y1 + 5, footer_company.y0)

    @skipUnless(_docx_available(), 'python-docx not installed')
    def test_docx_uses_one_greeting_and_aligned_company_footer(self):
        from io import BytesIO
        from docx import Document
        letter = self.prepare('eoi')
        letter.body = 'Dear Sir / Madam,\n\nProject details.'
        document = Document(BytesIO(generate_letter_docx(letter)))
        self.assertEqual([p.text for p in document.paragraphs].count('Dear Sir / Madam,'), 1)
        self.assertEqual(document.tables[0].cell(1, 0).text, 'Confidential')
        self.assertEqual(document.tables[-1].cell(0, 0).text, 'Subject:')
        self.assertEqual(document.sections[0].footer.paragraphs[1].text,
                         'Rejlers International Engineering Solutions AB')

    @skipUnless(_weasyprint_available(), 'weasyprint not installed')
    def test_draft_preview_uses_current_layout_without_changing_saved_file(self):
        import pymupdf
        from apps.sales.letter_pdf import get_letter_pdf_bytes
        letter = self.prepare('eoi')
        original_file = letter.pdf_file.name
        letter.body = 'Dear Sir / Madam,\n\nUpdated draft wording.'
        pdf_bytes = get_letter_pdf_bytes(letter)
        document = pymupdf.open(stream=pdf_bytes, filetype='pdf')
        self.assertEqual(len(document[0].search_for('Updated draft wording.')), 1)
        self.assertEqual(len(document[0].search_for('Dear Sir / Madam,')), 1)
        self.assertEqual(letter.pdf_file.name, original_file)
        letter.status = 'sent'
        issued = pymupdf.open(stream=get_letter_pdf_bytes(letter), filetype='pdf')
        self.assertEqual(len(issued[0].search_for('Updated draft wording.')), 0)


@override_settings(SALES_WORKSPACE_ENABLED=False)
class LetterAttachTests(TestCase):
    def setUp(self):
        seed_test_letter_templates()
        User = get_user_model()
        self.actor = User.objects.create_user(username='letter-owner', email='owner@example.test')
        grant_sales_actions(self.actor, 'sales', 'sales_opportunities', 'sales_clients')
        self.client_record = Client.objects.create(
            client_code='L-CLIENT', company_name='Abu Dhabi National Oil Company',
            industry_type='oil_gas', account_manager=self.actor,
        )
        self.contact = Contact.objects.create(
            client=self.client_record, first_name='Sample', last_name='Recipient',
            email='recipient@example.test',
        )
        self.opportunity = Deal.objects.create(
            deal_code='Q-101883', deal_name='Conceptual Facility Study',
            client=self.client_record, client_contact=self.contact, owner=self.actor,
            stage='proposal', bid_decision='bid',
        )
        self.private_root = TemporaryDirectory(prefix='radai-letter-attach-tests-')
        self.addCleanup(self.private_root.cleanup)
        override = override_settings(SALES_ATTACHMENT_ROOT=self.private_root.name)
        override.enable()
        self.addCleanup(override.disable)

    def prepare(self, letter_type):
        return prepare_letter(
            self.opportunity, self.actor, letter_type,
            {'letterhead': {'recipient_name': self.contact.full_name}},
        )

    @skipUnless(_docx_available(), 'python-docx not installed')
    def test_attach_uploads_pdf_and_docx_to_correspondence(self):
        letter = self.prepare('eoi')
        attachments = attach_letter_files(letter, self.actor)
        self.assertEqual(attachments['folder'], 'correspondence')
        self.assertEqual(attachments['version'], 1)
        kinds = {f['kind']: f for f in attachments['files']}
        self.assertEqual(set(kinds), {'pdf', 'docx'})
        for item in kinds.values():
            self.assertEqual(item['status'], 'ready')
            self.assertEqual(item['provider'], 'radai')
        event = OpportunityAuditEvent.objects.get(event_type='letter_attached')
        self.assertEqual(event.data['folder'], 'correspondence')
        self.assertEqual(len(event.data['files']), 2)

        with patch('apps.sales.letter_attachments.upload_private_file') as upload:
            repeated = attach_letter_files(letter, self.actor)
        self.assertEqual(repeated['version'], 1)
        upload.assert_not_called()

        letter.refresh_from_db()
        letter.custom_data['attachments'].pop('fingerprint')
        letter.save(update_fields=['custom_data'])
        with patch('apps.sales.letter_attachments.upload_private_file') as upload:
            legacy = attach_letter_files(letter, self.actor)
        self.assertEqual(legacy['version'], 1)
        upload.assert_not_called()

        letter.refresh_from_db()
        letter.body += '\n\nReviewed clarification.'
        refresh_letter_files(letter, self.actor)
        letter.refresh_from_db()
        attachments = attach_letter_files(letter, self.actor)
        self.assertEqual(attachments['version'], 2)
        self.assertEqual(
            {f['name'] for f in attachments['files']},
            {'Q-101883-eoi.pdf', 'Q-101883-eoi.docx'},
        )

    @skipUnless(_docx_available(), 'python-docx not installed')
    def test_attach_mirrors_manual_upload_single_call_per_file(self):
        """Letter attach makes one service call per file, exactly like the
        manual workspace upload endpoint — no cross-provider fallback. A
        failing SharePoint upload surfaces as a failed attach, identical to a
        failed manual upload."""
        letter = self.prepare('eoi')

        with patch('apps.sales.letter_attachments.workspace_config', return_value=object()), \
             patch('apps.sales.letter_attachments.upload_workspace_file',
                   side_effect=WorkspaceAPIError('remote_busy')):
            attachments = attach_letter_files(letter, self.actor)

        self.assertEqual(len(attachments['files']), 2)
        for item in attachments['files']:
            self.assertEqual(item['status'], 'failed')
            self.assertEqual(item['provider'], 'sharepoint')
            self.assertTrue(item.get('error'))
        event = OpportunityAuditEvent.objects.get(event_type='letter_attached')
        self.assertEqual(event.data['storage_provider'], 'sharepoint')
        self.assertEqual(len(event.data['failed_files']), 2)

    @skipUnless(_docx_available(), 'python-docx not installed')
    def test_prepare_and_refresh_via_api_attach_each_time(self):
        from rest_framework.test import APIClient

        api = APIClient()
        api.force_authenticate(self.actor)
        url = f'/api/v1/sales/deals/{self.opportunity.pk}/prepare-letter/'
        response = api.post(url, {
            'letter_type': 'eoi',
            'custom_data': {'letterhead': {'recipient_name': self.contact.full_name}},
        }, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        letter_id = response.data['id']
        self.assertEqual(response.data['attachments']['version'], 1)
        self.assertEqual(
            {f['kind'] for f in response.data['attachments']['files']},
            {'pdf', 'docx'},
        )
        repeated = api.post(url, {
            'letter_type': 'eoi',
            'custom_data': {'letterhead': {'recipient_name': self.contact.full_name}},
        }, format='json')
        self.assertEqual(repeated.status_code, 200, repeated.data)
        self.assertEqual(repeated.data['id'], letter_id)
        self.assertEqual(repeated.data['attachments']['version'], 1)
        self.assertEqual(SalesLetter.objects.count(), 1)

        detail = f'/api/v1/sales/deals/{self.opportunity.pk}/letters/{letter_id}/'
        unchanged = api.post(detail + 'regenerate-pdf/', {
            'subject': response.data['subject'],
            'body': response.data['body'],
            'custom_data': response.data['custom_data'],
        }, format='json')
        self.assertEqual(unchanged.status_code, 200, unchanged.data)
        self.assertEqual(unchanged.data['custom_data']['letterhead']['rev'], 0)
        self.assertEqual(unchanged.data['attachments']['version'], 1)
        regen = api.post(detail + 'regenerate-pdf/', {
            'subject': 'Edited subject',
            'custom_data': response.data['custom_data'],
        }, format='json')
        self.assertEqual(regen.status_code, 200, regen.data)
        self.assertEqual(regen.data['custom_data']['letterhead']['rev'], 1)
        self.assertEqual(regen.data['attachments']['version'], 2)
        retry = api.post(detail + 'regenerate-pdf/', {
            'subject': 'Edited subject',
            'custom_data': response.data['custom_data'],
        }, format='json')
        self.assertEqual(retry.status_code, 200, retry.data)
        self.assertEqual(retry.data['custom_data']['letterhead']['rev'], 1)
        self.assertEqual(retry.data['attachments']['version'], 2)

        docx = api.get(detail + 'docx/')
        self.assertEqual(docx.status_code, 200)
        self.assertTrue(
            b''.join(docx.streaming_content).startswith(b'PK')
        )

        upload = OpportunityAuditEvent.objects.filter(event_type='workspace_upload_started')
        self.assertGreaterEqual(upload.count(), 4)
