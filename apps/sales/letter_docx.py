"""DOCX generation for sales letters (python-docx).

Mirrors the standard Rejlers letterhead produced by the PDF renderer so the
downloadable Word copy can be edited offline without losing the layout.
"""

import logging
from io import BytesIO

from django.core.files.base import ContentFile
from django.utils import timezone
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from .letter_context import build_letter_context, split_body_paragraphs
from .letter_defaults import get_logo_path

logger = logging.getLogger(__name__)

DOCX_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


def _strip_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        element = OxmlElement(f'w:{edge}')
        element.set(qn('w:val'), 'none')
        element.set(qn('w:sz'), '0')
        borders.append(element)
    tbl_pr.append(borders)


def _set_cell_width(cell, width_mm):
    cell.width = Mm(width_mm)


def _set_table_widths(table, widths_mm):
    for index, width in enumerate(widths_mm):
        table.columns[index].width = Mm(width)
        for row in table.rows:
            _set_cell_width(row.cells[index], width)


def _add_run(paragraph, text, *, bold=False):
    run = paragraph.add_run(text)
    run.bold = bold
    return run


def _meta_row(table, row_idx, label, lines):
    label_cell = table.cell(row_idx, 0)
    value_cell = table.cell(row_idx, 1)
    label_cell.text = ''
    value_cell.text = ''
    _add_run(label_cell.paragraphs[0], label)
    first = True
    for line in lines:
        paragraph = value_cell.paragraphs[0] if first else value_cell.add_paragraph()
        _add_run(paragraph, line)
        first = False
    return label_cell, value_cell


def _add_bottom_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'B7B7B7')
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _add_page_field(paragraph, instruction):
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), instruction)
    paragraph._p.append(field)


def generate_letter_docx(letter, custom_data=None):
    """Generate DOCX bytes for a letter."""
    merged = dict(letter.custom_data or {})
    merged.update(custom_data or {})
    context = build_letter_context(letter.opportunity, letter=letter, custom_data=merged)
    lh = context['letterhead']

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.bottom_margin = Mm(14), Mm(16)
    section.left_margin, section.right_margin = Mm(18), Mm(18)
    section.footer_distance = Mm(9)

    normal = document.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(9)
    footer_style = document.styles['Footer']
    footer_style.font.name = 'Arial'
    footer_style.font.size = Pt(7)

    # Use fixed columns so both header values align with the letter body.
    header_table = document.add_table(rows=2, cols=3)
    header_table.autofit = False
    _strip_table_borders(header_table)
    _set_table_widths(header_table, (41, 95, 38))
    for row_idx, (label, value) in enumerate(((lh['response_label'], lh['response_code']),
                                              ('Confidential', lh['confidential_date']))):
        label_cell, value_cell = header_table.cell(row_idx, 0), header_table.cell(row_idx, 1)
        label_cell.text = value_cell.text = ''
        label_p, value_p = label_cell.paragraphs[0], value_cell.paragraphs[0]
        label_p.paragraph_format.space_after = value_p.paragraph_format.space_after = Pt(0)
        if row_idx:
            label_p.paragraph_format.space_before = value_p.paragraph_format.space_before = Pt(8)
        _add_run(label_p, label, bold=row_idx == 0)
        _add_run(value_p, value)
    right = header_table.cell(0, 2).merge(header_table.cell(1, 2))
    right.text = ''
    logo_path = context.get('logo_path') or get_logo_path()
    if logo_path:
        right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        right.paragraphs[0].add_run().add_picture(str(logo_path), height=Mm(7))
    else:
        run = _add_run(right.paragraphs[0], 'REJLERS', bold=True)
        run.font.size = Pt(14)

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(15)

    # Meta block - compact spacing, left aligned at same margin.
    recipient_lines = [lh['recipient_name']]
    if lh.get('recipient_title'):
        recipient_lines.append(lh['recipient_title'])
    recipient_lines.append(lh['recipient_company'])
    if lh.get('recipient_address'):
        recipient_lines.append(lh['recipient_address'])
    if lh.get('recipient_email'):
        recipient_lines.append(f"E-mail: {lh['recipient_email']}")

    meta = document.add_table(rows=9, cols=2)
    meta.autofit = False
    _strip_table_borders(meta)
    _set_table_widths(meta, (41, 133))
    rows = [
        ('From:', [lh['sender_name'], lh['sender_company']]),
        ('Bid Focal Contact', [lh['focal_contact']]),
        ('Contact Number:', [lh['contact_number']]),
        ('Fax:', [lh['fax']]),
        ('Email:', [f"E: {lh['email']}"]),
        ('To:', recipient_lines),
        ('Your Ref. No.:', [lh['your_ref'] or '—']),
        ('Our Ref. No.:', [lh['our_ref']]),
        ('Our ADNOC Unified\nCode:', [lh['adnoc_unified_code']]),
    ]
    for idx, (label, lines) in enumerate(rows):
        _meta_row(meta, idx, label, lines)
        for cell in [meta.cell(idx, 0), meta.cell(idx, 1)]:
            for paragraph_idx, paragraph in enumerate(cell.paragraphs):
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt({1: 12, 5: 18, 6: 31, 7: 12, 8: 12}.get(idx, 0)) if paragraph_idx == 0 else Pt(0)
                paragraph.paragraph_format.line_spacing = Pt(11)

    subject_table = document.add_table(rows=1, cols=2)
    subject_table.autofit = False
    _strip_table_borders(subject_table)
    _set_table_widths(subject_table, (41, 133))
    subject_label, subject_value = subject_table.cell(0, 0), subject_table.cell(0, 1)
    subject_label.text = subject_value.text = ''
    subject_label.paragraphs[0].paragraph_format.space_before = Pt(12)
    subject_value.paragraphs[0].paragraph_format.space_before = Pt(12)
    _add_run(subject_label.paragraphs[0], 'Subject:')
    _add_run(subject_value.paragraphs[0], context['subject'], bold=True)
    if context.get('subject_suffix'):
        suffix_p = subject_value.add_paragraph()
        suffix_p.paragraph_format.space_before = Pt(4)
        _add_run(suffix_p, context['subject_suffix'], bold=True)

    salutation = document.add_paragraph(context['salutation'])
    salutation.paragraph_format.space_before = Pt(12)
    salutation.paragraph_format.space_after = Pt(8)

    # Body.
    for para in split_body_paragraphs(letter.body):
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(9)
        p.paragraph_format.line_spacing = Pt(11.5)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _add_run(p, para)

    # Signature block.
    closing_p = document.add_paragraph('Sincerely yours,')
    closing_p.paragraph_format.space_before = Pt(13)
    closing_p.paragraph_format.space_after = Pt(24)
    closing_p.paragraph_format.keep_with_next = True
    signature_p = document.add_paragraph()
    signature_p.paragraph_format.keep_with_next = bool(lh['signature_titles'])
    _add_run(signature_p, lh['signature_name'])
    for index, title in enumerate(lh['signature_titles']):
        title_p = document.add_paragraph()
        title_p.paragraph_format.space_after = Pt(0)
        title_p.paragraph_format.space_before = Pt(0)
        title_p.paragraph_format.keep_with_next = index < len(lh['signature_titles']) - 1
        _add_run(title_p, title)

    # Footer - add to document footer section for proper bottom positioning
    footer = section.footer
    footer.is_linked_to_previous = False
    footer.paragraphs[0].clear()
    
    # Footer separator line
    rule_p = footer.paragraphs[0]
    rule_p.style = footer_style
    _add_bottom_border(rule_p)
    rule_p.paragraph_format.space_after = Pt(2)
    rule_p.paragraph_format.space_before = Pt(0)
    rule_p.paragraph_format.left_indent = Mm(-2)
    rule_p.paragraph_format.right_indent = Mm(-13)
    
    footer_lines = [lh['footer_company'], lh['footer_address'], lh['footer_contact']]
    for line in footer_lines:
        p = footer.add_paragraph()
        p.style = footer_style
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.left_indent = Mm(-2)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = _add_run(p, line)
        run.font.size = Pt(7)
        if line == footer_lines[-1]:
            p.paragraph_format.tab_stops.add_tab_stop(Mm(171), WD_TAB_ALIGNMENT.RIGHT)
            p.add_run('\t')
            _add_page_field(p, 'PAGE')
            p.add_run(' (')
            _add_page_field(p, 'NUMPAGES')
            p.add_run(')')

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def save_letter_docx(letter, custom_data=None):
    """Generate and save the DOCX to the letter's docx_file field."""
    from .letter_extraction import require_recipient
    require_recipient({**(letter.custom_data or {}), **(custom_data or {})}.get('letterhead') or {})
    docx_bytes = generate_letter_docx(letter, custom_data)
    filename = f"{letter.opportunity.deal_code}-{letter.letter_type}-{letter.generated_at.strftime('%Y%m%d')}.docx"
    letter.docx_file.save(filename, ContentFile(docx_bytes), save=False)
    letter.docx_generated_at = timezone.now()
    letter.save(update_fields=['docx_file', 'docx_generated_at', 'updated_at'])
    logger.info(f"DOCX generated for letter {letter.id}: {filename}")
    return letter.docx_file.name


def regenerate_letter_docx(letter, custom_data=None):
    old_name = letter.docx_file.name if letter.docx_file else None
    storage = letter.docx_file.storage
    new_name = save_letter_docx(letter, custom_data)
    if old_name and old_name != new_name:
        try:
            storage.delete(old_name)
        except Exception as exc:
            logger.warning("Unable to remove previous DOCX for letter %s: %s", letter.id, exc)
    return new_name


def get_letter_docx_bytes(letter):
    """Get DOCX bytes (generates on demand if missing)."""
    if letter.status == 'draft':
        return generate_letter_docx(letter)
    if not letter.docx_file:
        save_letter_docx(letter)
    if letter.docx_file:
        letter.docx_file.open('rb')
        docx_bytes = letter.docx_file.read()
        letter.docx_file.close()
        return docx_bytes
    return generate_letter_docx(letter)
