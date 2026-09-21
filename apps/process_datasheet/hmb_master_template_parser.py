"""
HMB Master Template Parser
==========================
Soft-coded parser for the shared HMB master workbook layout used as a
reference template for multi-case data collection.

Phase 1 scope:
  - Parse workbook structure (sheet, stream columns, section row ranges)
  - Extract row-schema metadata (property + unit)
  - Return a compact preview payload for frontend review
"""

from __future__ import annotations

from io import BytesIO
from typing import Dict, List, Any
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries


HMB_MASTER_TEMPLATE_CONFIG: Dict[str, Any] = {
    # Master layout assumptions (kept soft-coded for future variants)
    'sheet_name': 'Master',
    'stream_header_row': 1,
    'stream_description_row': 2,
    'stream_start_col': 4,  # Column D
    'property_col': 2,      # Column B
    'unit_col': 3,          # Column C
    'max_stream_preview': 12,
    'max_record_preview': 120,
    'sections': [
        {'key': 'general', 'label': 'General', 'start_row': 6, 'end_row': 10},
        {'key': 'vapour', 'label': 'Vapour', 'start_row': 12, 'end_row': 21},
        {'key': 'light_liquid', 'label': 'Light Liquid', 'start_row': 23, 'end_row': 31},
        {'key': 'heavy_liquid', 'label': 'Heavy Liquid', 'start_row': 33, 'end_row': 41},
        {'key': 'composition', 'label': 'Composition', 'start_row': 43, 'end_row': 78},
    ],
}

HMB_MASTER_TEMPLATE_ANALYSIS_VERSION = '1.2'

HMB_FIXED_CASE_LABELS = [
    f'CASE {group} ({number}{variant})'
    for group in ('A', 'B', 'C')
    for number in (1, 2)
    for variant in ('a', 'b')
]


HMB_CASE_PARSER_CONFIG: Dict[str, Any] = {
    'preferred_sheet_names': ['Overall', 'Master'],
    'stream_header_row': 4,
    'stream_description_row': 0,
    'stream_start_col': 3,
    'property_col': 1,
    'unit_col': 2,
    'sections': HMB_MASTER_TEMPLATE_CONFIG['sections'],
    'section_markers': {
        'conditions': ['conditions'],
        'vapour': ['vapour'],
        'light_liquid': ['light liquid'],
        'heavy_liquid': ['heavy liquid'],
        'composition': ['composition'],
    },
}

_EXCEPTIONS_SAMPLE_LIMIT = 120


def _color_value(color) -> str:
    if color is None:
        return ''
    if getattr(color, 'type', None) == 'rgb':
        return _to_text(color.rgb).lstrip('#')
    return ''


def _json_safe_cell_value(value: Any):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if hasattr(value, 'isoformat'):
        return value.isoformat()
    return str(value)


def _serialize_master_layout(ws) -> Dict[str, Any]:
    max_row = ws.max_row
    max_col = ws.max_column
    if ws.print_area:
        for print_range in str(ws.print_area).split(','):
            cell_range = print_range.rsplit('!', 1)[-1].replace("'", '')
            _, _, print_max_col, print_max_row = range_boundaries(cell_range)
            max_row = max(max_row, print_max_row)
            max_col = max(max_col, print_max_col)

    merged_anchors = {}
    merged_ranges = []
    for merged_range in ws.merged_cells.ranges:
        min_col, min_row, range_max_col, range_max_row = merged_range.bounds
        if min_row > max_row or min_col > max_col:
            continue
        merged_anchors[(min_row, min_col)] = {
            'row_span': range_max_row - min_row + 1,
            'col_span': range_max_col - min_col + 1,
        }
        merged_ranges.append(str(merged_range))

    cells = []
    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            cell = ws.cell(row, col)
            if cell.coordinate in ws.merged_cells and (row, col) not in merged_anchors:
                continue
            merge = merged_anchors.get((row, col), {})
            alignment = cell.alignment
            fill = cell.fill
            font = cell.font
            border = cell.border
            cells.append({
                'row': row,
                'column': col,
                'coordinate': cell.coordinate,
                'value': _json_safe_cell_value(cell.value),
                'row_span': merge.get('row_span', 1),
                'col_span': merge.get('col_span', 1),
                'style': {
                    'style_id': cell.style_id,
                    'fill': _color_value(fill.fgColor) if fill.fill_type else '',
                    'font_color': _color_value(font.color),
                    'bold': bool(font.bold),
                    'italic': bool(font.italic),
                    'horizontal': alignment.horizontal or '',
                    'vertical': alignment.vertical or '',
                    'wrap_text': bool(alignment.wrap_text),
                    'number_format': cell.number_format or '',
                    'border_top': border.top.style or '',
                    'border_right': border.right.style or '',
                    'border_bottom': border.bottom.style or '',
                    'border_left': border.left.style or '',
                },
            })

    return {
        'max_row': max_row,
        'max_col': max_col,
        'cells': cells,
        'merged_ranges': merged_ranges,
        'column_widths': {
            get_column_letter(col): ws.column_dimensions[get_column_letter(col)].width or 13
            for col in range(1, max_col + 1)
        },
        'row_heights': {
            str(row): ws.row_dimensions[row].height or 15
            for row in range(1, max_row + 1)
        },
    }


def _to_text(value: Any) -> str:
    if value is None:
        return ''
    text = str(value).strip()
    return text


def _is_formula_or_ref_error(value: Any) -> bool:
    text = _to_text(value)
    if not text:
        return False
    upper = text.upper()
    return text.startswith('=') or '#REF!' in upper


def _is_valid_stream_header(value: Any) -> bool:
    text = _to_text(value)
    if not text:
        return False
    return not _is_formula_or_ref_error(text)


def _pick_case_sheet(wb, preferred_names: List[str]):
    for name in preferred_names:
        if name in wb.sheetnames:
            return wb[name]
    # fallback: widest sheet typically contains stream matrix
    return max(wb.worksheets, key=lambda ws: ws.max_column)


def _is_stream_like(value: str) -> bool:
    txt = _to_text(value)
    if not txt:
        return False
    txt = txt.replace('-', '').replace('.', '').replace('_', '').replace('/', '')
    return txt.isdigit()


def _detect_case_layout(ws, cfg: Dict[str, Any]) -> Dict[str, int]:
    # Find a header row that contains Property/Unit plus many stream ids.
    for r in range(1, min(15, ws.max_row) + 1):
        c1 = _to_text(ws.cell(r, 1).value).lower()
        c2 = _to_text(ws.cell(r, 2).value).lower()
        if c1 == 'property' and c2 == 'unit':
            stream_start = 3
            stream_like_count = 0
            for col in range(stream_start, min(ws.max_column, 500) + 1):
                if _is_stream_like(ws.cell(r, col).value):
                    stream_like_count += 1
            if stream_like_count >= 3:
                return {
                    'stream_header_row': r,
                    'property_col': 1,
                    'unit_col': 2,
                    'stream_start_col': stream_start,
                    'stream_description_row': 0,
                }

    # Fallback to configured defaults
    return {
        'stream_header_row': cfg['stream_header_row'],
        'property_col': cfg['property_col'],
        'unit_col': cfg['unit_col'],
        'stream_start_col': cfg['stream_start_col'],
        'stream_description_row': cfg['stream_description_row'],
    }


def _section_from_label(text: str, marker_cfg: Dict[str, List[str]], current: Dict[str, str]) -> Dict[str, str]:
    lowered = _to_text(text).lower()
    if not lowered:
        return current
    for key, aliases in marker_cfg.items():
        for alias in aliases:
            if alias in lowered:
                return {'key': key, 'label': key.replace('_', ' ').title()}
    return current


def _normalise_text(value: Any) -> str:
    return re.sub(r'\s+', ' ', _to_text(value)).strip()


def _normalise_key(value: Any) -> str:
    return _normalise_text(value).lower().rstrip(':')


def _normalise_unit(value: Any) -> str:
    key = _normalise_key(value)
    key = key.replace('lbm/', 'lb/').replace('lbm ', 'lb ')
    key = {
        'btu/hr/ft/f': 'btu/hr-ft-f',
        'btu/lb/f': 'btu/lb-f',
        'lb/hr': 'lb/hr',
    }.get(key, key)
    if key in {'', '-', 'none', '<none>'}:
        return '<none>'
    return key


def _normalise_component(value: Any) -> str:
    return _normalise_text(value).lower()


def _to_float(value: Any):
    if isinstance(value, (int, float)):
        return float(value)
    txt = _to_text(value)
    if not txt or txt == '---':
        return None
    try:
        return float(txt)
    except Exception:
        return None


def _find_property_layout(ws, max_scan_rows: int = 120, max_scan_cols: int = 20):
    for r in range(1, min(ws.max_row, max_scan_rows) + 1):
        vals = [_normalise_key(ws.cell(r, c).value) for c in range(1, min(ws.max_column, max_scan_cols) + 1)]
        if 'property' in vals and 'unit' in vals:
            p_col = vals.index('property') + 1
            u_col = vals.index('unit') + 1
            return {'header_row': r, 'property_col': p_col, 'unit_col': u_col, 'stream_start_col': max(p_col, u_col) + 1}
    return {
        'header_row': HMB_CASE_PARSER_CONFIG['stream_header_row'],
        'property_col': HMB_CASE_PARSER_CONFIG['property_col'],
        'unit_col': HMB_CASE_PARSER_CONFIG['unit_col'],
        'stream_start_col': HMB_CASE_PARSER_CONFIG['stream_start_col'],
    }


def _extract_stream_columns(ws, layout):
    columns = []
    header_row = layout['header_row']
    empty_streak = 0
    for c in range(layout['stream_start_col'], ws.max_column + 1):
        stream_id = _normalise_text(ws.cell(header_row, c).value)
        if _is_valid_stream_header(stream_id):
            columns.append({'column_index': c, 'stream_id': stream_id})
            empty_streak = 0
            continue
        if stream_id == '':
            empty_streak += 1
            if columns and empty_streak >= 8:
                break
    return columns


def _parse_phase_property_sheet(ws):
    layout = _find_property_layout(ws)
    streams = _extract_stream_columns(ws, layout)
    rows = {}
    for r in range(layout['header_row'] + 1, ws.max_row + 1):
        raw_prop = _normalise_text(ws.cell(r, layout['property_col']).value)
        if not raw_prop:
            continue
        if raw_prop.isupper() and len(raw_prop) <= 40:
            continue
        prop_key = _normalise_key(raw_prop)
        unit = _normalise_text(ws.cell(r, layout['unit_col']).value)
        values = {}
        for stream in streams:
            v = ws.cell(r, stream['column_index']).value
            if v in (None, '', '---'):
                continue
            if _is_formula_or_ref_error(v):
                continue
            values[stream['stream_id']] = v
        rows.setdefault(prop_key, []).append({
            'raw_property': raw_prop,
            'unit': unit,
            'values': values,
        })
    return {
        'layout': layout,
        'stream_ids': [s['stream_id'] for s in streams],
        'rows': rows,
    }


def _extract_case_name(wb, source_filename: str = '') -> str:
    if 'Info' in wb.sheetnames:
        ws = wb['Info']
        case_name = _normalise_text(ws.cell(3, 2).value)
        if not case_name:
            for r in range(1, min(ws.max_row, 40) + 1):
                key = _normalise_key(ws.cell(r, 1).value)
                if key == 'case name':
                    case_name = _normalise_text(ws.cell(r, 2).value)
                    break
        if case_name.lower().endswith('.hsc'):
            case_name = case_name[:-4]
        if case_name:
            return _canonical_case_name(case_name)
    return _canonical_case_name(source_filename)


def _canonical_case_name(value: Any) -> str:
    text = _normalise_text(value)
    stem = re.sub(r'\.(xlsx|xlsm|hsc)$', '', text, flags=re.IGNORECASE)
    compact = re.sub(r'[^a-z0-9]+', '', stem.lower())

    explicit = re.search(r'case\s*([abc])\s*\(?\s*([12])\s*([ab])\s*\)?', stem, re.IGNORECASE)
    if explicit:
        return f'CASE {explicit.group(1).upper()} ({explicit.group(2)}{explicit.group(3).lower()})'

    legacy = re.search(r'case\s*[_-]?\s*([12])\s*([ab])', stem, re.IGNORECASE)
    if legacy:
        return f'CASE A ({legacy.group(1)}{legacy.group(2).lower()})'

    for label in HMB_FIXED_CASE_LABELS:
        if compact == re.sub(r'[^a-z0-9]+', '', label.lower()):
            return label
    return stem or 'Case'


def _is_normalized_summary_workbook(wb) -> bool:
    return {'Conditions', 'Properties', 'Mole Fraction'}.issubset(set(wb.sheetnames))


def _split_summary_header(value: Any):
    text = _normalise_text(value)
    if not text or '|' not in text:
        return None
    property_part, phase_part = (part.strip() for part in text.rsplit('|', 1))
    match = re.match(r'^(.*?)\s*:?\s*\(([^()]*)\)\s*$', property_part)
    if match:
        property_name = match.group(1).strip()
        unit = match.group(2).strip().replace('-', '/')
    else:
        property_name = property_part.rstrip(':').strip()
        unit = ''
    section_key = {
        'overall': 'general',
        'vapour phase': 'vapour',
        'vapor phase': 'vapour',
        'liquid phase': 'light_liquid',
        'aqueous phase': 'heavy_liquid',
    }.get(_normalise_key(phase_part))
    if not property_name or not section_key:
        return None
    return property_name, unit, section_key


def _parse_normalized_summary_sheet(ws):
    sections = {
        key: {'layout': {'header_row': 1}, 'stream_ids': [], 'rows': {}}
        for key in ('general', 'vapour', 'light_liquid', 'heavy_liquid')
    }
    stream_ids = []
    seen_streams = set()
    for row in range(2, ws.max_row + 1):
        stream_id = _normalise_text(ws.cell(row, 1).value)
        if stream_id and stream_id not in seen_streams:
            stream_ids.append(stream_id)
            seen_streams.add(stream_id)

    for col in range(2, ws.max_column + 1):
        parsed_header = _split_summary_header(ws.cell(1, col).value)
        if not parsed_header:
            continue
        property_name, unit, section_key = parsed_header
        values = {}
        for row in range(2, ws.max_row + 1):
            stream_id = _normalise_text(ws.cell(row, 1).value)
            value = ws.cell(row, col).value
            if not stream_id or value in (None, '', '---') or _is_formula_or_ref_error(value):
                continue
            values[stream_id] = value
        prop_key = _normalise_key(property_name)
        sections[section_key]['rows'].setdefault(prop_key, []).append({
            'raw_property': property_name,
            'unit': unit,
            'values': values,
        })

    for section in sections.values():
        section['stream_ids'] = list(stream_ids)
    return sections, stream_ids


def _parse_normalized_composition(ws, allowed_streams=None, allowed_components=None):
    allowed_streams = set(allowed_streams or [])
    allowed_components = set(_normalise_component(c) for c in (allowed_components or []))
    components = [
        (_normalise_component(ws.cell(1, col).value), col)
        for col in range(2, ws.max_column + 1)
        if _normalise_text(ws.cell(1, col).value)
    ]
    out = {}
    for row in range(2, ws.max_row + 1):
        stream_id = _normalise_text(ws.cell(row, 1).value)
        if not stream_id or (allowed_streams and stream_id not in allowed_streams):
            continue
        for component, col in components:
            if allowed_components and component not in allowed_components:
                continue
            value = _to_float(ws.cell(row, col).value)
            out[(stream_id, component)] = 0.0 if value is None else value
    return out


def _parse_master_case_workbook(wb, source_filename: str, template_profile_payload: Dict[str, Any]):
    ws = wb['Master']
    case_name = _canonical_case_name(ws.cell(1, 1).value or source_filename)
    template_sections = template_profile_payload.get('sections', [])
    template_streams = template_profile_payload.get('stream_columns', [])
    source_columns = {
        _normalise_text(ws.cell(1, col).value): col
        for col in range(4, ws.max_column + 1)
        if _normalise_text(ws.cell(1, col).value)
    }
    records = []
    for section in template_sections:
        section_label = _normalise_text(section.get('label'))
        section_key = _normalise_key(section_label).replace(' ', '_')
        for prop in section.get('properties', []) or []:
            row = int(prop.get('row') or 0)
            if not row:
                continue
            for stream in template_streams:
                stream_id = _normalise_text(stream.get('stream_id'))
                col = source_columns.get(stream_id)
                if not col:
                    continue
                value = ws.cell(row, col).value
                if value in (None, '', '---') or _is_formula_or_ref_error(value):
                    continue
                records.append({
                    'case_name': case_name,
                    'source_filename': source_filename,
                    'sheet_name': 'Master',
                    'section_key': section_key,
                    'section_label': section_label,
                    'row_index': row,
                    'property_name': _normalise_text(prop.get('property')),
                    'unit': _normalise_text(prop.get('unit')),
                    'stream_id': stream_id,
                    'source_stream_id': stream_id,
                    'stream_description': _normalise_text(stream.get('description')),
                    'value_text': str(value),
                })
    matched = {stream_id: stream_id for stream_id in source_columns if stream_id in {
        _normalise_text(stream.get('stream_id')) for stream in template_streams
    }}
    return {
        'case_name': case_name,
        'sheet_name': 'Master',
        'detected_format': 'master_case',
        'stream_count': len(matched),
        'record_count': len(records),
        'records': records,
        'exceptions': {},
        'stream_mapping': {
            'matched': matched,
            'source_stream_count': len(source_columns),
            'mapped_source_stream_count': len(matched),
            'ignored_source_stream_count': len(source_columns) - len(matched),
            'ignored_source_streams': [],
            'template_stream_count': len(template_streams),
        },
    }


def _resolve_stream_mapping(template_streams, source_stream_ids):
    source_ids = [_normalise_text(s) for s in source_stream_ids if _normalise_text(s)]
    source_set = set(source_ids)
    mapping = {}
    unmatched = []
    duplicates = []
    used = set()

    for stream in template_streams:
        template_id = _normalise_text(stream.get('stream_id'))
        if not template_id:
            continue
        matched = None

        if template_id in source_set:
            matched = template_id
        else:
            desc = _normalise_text(stream.get('description'))
            if desc:
                hits = []
                for sid in source_ids:
                    if re.search(rf'(^|\W){re.escape(sid)}(\W|$)', desc):
                        hits.append(sid)
                if len(hits) == 1:
                    matched = hits[0]

        if not matched:
            unmatched.append(template_id)
            continue
        if matched in used:
            duplicates.append({'template_stream': template_id, 'source_stream': matched})
            continue
        used.add(matched)
        mapping[template_id] = matched
    return mapping, unmatched, duplicates


def _property_candidates(section_key: str, property_name: str):
    p = _normalise_key(property_name)
    aliases = {
        'vapour fraction': ['vapour / phase fraction', 'phase fraction', 'vapour fraction'],
        'standard ideal liquid volume flow': ['standard ideal liquid volume flow', 'std ideal liq vol flow', 'std ideal liquid vol flow'],
        'std gas flow': ['std gas flow', 'std. gas flow'],
        'actual volume flow': ['actual volume flow', 'act. volume flow', 'act. liq. flow', 'act. gas flow'],
        'mass heat capacity': ['mass heat capacity'],
        'cp/cv (gamma)': ['cp/cv (gamma)', 'cp/cv', 'cp/(cp - r)', 'gamma'],
        'compressibility': ['compressibility', 'z factor', 'cp/(cp - r)'],
        'mass density': ['mass density', 'liq. mass density (std. cond)'],
        'thermal conductivity': ['thermal conductivity'],
        'viscosity': ['viscosity', 'kinematic viscosity'],
    }
    cands = [p]
    cands.extend(aliases.get(p, []))
    if section_key == 'general' and p == 'vapour fraction':
        cands.append('vapour / phase fraction')
    # punctuation-normalized variants to match keys like "std. gas flow"
    normalized = []
    for c in cands:
        normalized.append(c)
        normalized.append(re.sub(r'[^a-z0-9]+', ' ', c).strip())
    cands = normalized
    return list(dict.fromkeys(cands))


def _unit_compatible(template_unit: str, source_unit: str) -> bool:
    t = _normalise_unit(template_unit)
    s = _normalise_unit(source_unit)
    if t == s:
        return True
    if t == '<none>' and s in {'<none>', ''}:
        return True
    return False


def _try_convert_unit(value: Any, source_unit: str, template_unit: str):
    num = _to_float(value)
    if num is None:
        return None
    s = _normalise_unit(source_unit)
    t = _normalise_unit(template_unit)

    # Known conversion used by HYSYS exports in sample streams.
    if s == 'acfm' and t == 'ft3/hr':
        return num * 60.0
    if s == 'barrel/day' and t == 'ft3/hr':
        return num * 5.614583333333333 / 24.0
    return None


def _best_source_row(source_sheet_data: Dict[str, Any], section_key: str, prop_name: str, source_stream_id: str, template_unit: str):
    rows = source_sheet_data.get('rows', {})
    first_found = None
    first_convertible = None
    for cand in _property_candidates(section_key, prop_name):
        row_entries = rows.get(cand)
        if not row_entries:
            continue
        if isinstance(row_entries, dict):
            row_entries = [row_entries]
        for row_data in row_entries:
            if source_stream_id not in row_data.get('values', {}):
                continue
            if first_found is None:
                first_found = row_data
            if _unit_compatible(template_unit, row_data.get('unit', '')):
                return row_data, None, None
            converted = _try_convert_unit(
                row_data.get('values', {}).get(source_stream_id),
                row_data.get('unit', ''),
                template_unit,
            )
            if converted is not None and first_convertible is None:
                first_convertible = (row_data, converted)
    if first_convertible is not None:
        row_data, converted = first_convertible
        return row_data, converted, None
    if first_found is not None:
        return None, None, first_found
    return None, None, None


def _parse_composition_all(ws, allowed_streams=None, allowed_components=None):
    allowed_streams = set(allowed_streams or [])
    allowed_components = set(_normalise_component(c) for c in (allowed_components or []))
    headers = [
        _normalise_key(v)
        for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    ]
    try:
        idx_stream = headers.index('stream')
        idx_phase = headers.index('phase')
        idx_component = headers.index('component')
        idx_mole_frac = headers.index('mole fraction')
    except ValueError:
        return {}

    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        stream_id = _normalise_text(row[idx_stream] if idx_stream < len(row) else '')
        if allowed_streams and stream_id not in allowed_streams:
            continue
        phase = _normalise_key(row[idx_phase] if idx_phase < len(row) else '')
        component = _normalise_component(row[idx_component] if idx_component < len(row) else '')
        if allowed_components and component not in allowed_components:
            continue
        if not stream_id or not component:
            continue
        if phase != 'overall':
            continue
        mole_frac = _to_float(row[idx_mole_frac] if idx_mole_frac < len(row) else None)
        out[(stream_id, component)] = 0.0 if mole_frac is None else mole_frac
    return out


def _parse_composition_overall(ws, allowed_streams=None, allowed_components=None):
    allowed_streams = set(allowed_streams or [])
    allowed_components = set(_normalise_component(c) for c in (allowed_components or []))

    layout = _find_property_layout(ws)
    stream_columns = _extract_stream_columns(ws, layout)
    stream_ids = [s['stream_id'] for s in stream_columns]
    if allowed_streams:
        stream_ids = [sid for sid in stream_ids if sid in allowed_streams]

    start_row = None
    end_row = None
    for r in range(layout['header_row'] + 1, ws.max_row + 1):
        token = _normalise_key(ws.cell(r, layout['property_col']).value)
        if token == 'mole fraction':
            start_row = r + 1
            continue
        if start_row and token in {'mass flow', 'mass fraction', 'liquid volume flow', 'liquid volume fraction', 'molar flow'}:
            end_row = r - 1
            break

    if start_row is None:
        return {}
    if end_row is None:
        end_row = ws.max_row

    col_by_stream = {s['stream_id']: s['column_index'] for s in stream_columns if s['stream_id'] in stream_ids}
    out = {}
    for r in range(start_row, end_row + 1):
        component_raw = _normalise_text(ws.cell(r, layout['property_col']).value)
        if not component_raw:
            continue
        comp_key = _normalise_component(component_raw)
        if allowed_components and comp_key not in allowed_components:
            continue
        for sid, col in col_by_stream.items():
            value = _to_float(ws.cell(r, col).value)
            out[(sid, comp_key)] = 0.0 if value is None else value
    return out


def _push_exception(exceptions: Dict[str, Any], key: str, item: Dict[str, Any]):
    count_key = f'{key}_count'
    exceptions[count_key] = int(exceptions.get(count_key, 0)) + 1
    bucket = exceptions.get(key)
    if not isinstance(bucket, list):
        bucket = []
        exceptions[key] = bucket
    if len(bucket) < _EXCEPTIONS_SAMPLE_LIMIT:
        bucket.append(item)


def analyze_hmb_master_template(workbook_path: str) -> Dict[str, Any]:
    with open(workbook_path, 'rb') as workbook_file:
        workbook_bytes = workbook_file.read()
    wb = load_workbook(BytesIO(workbook_bytes), data_only=False)
    values_wb = load_workbook(BytesIO(workbook_bytes), data_only=True)
    cfg = HMB_MASTER_TEMPLATE_CONFIG

    sheet_name = cfg['sheet_name'] if cfg['sheet_name'] in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]
    values_ws = values_wb[sheet_name]

    header_row = cfg['stream_header_row']
    desc_row = cfg['stream_description_row']
    start_col = cfg['stream_start_col']

    stream_columns: List[Dict[str, Any]] = []
    for col in range(start_col, ws.max_column + 1):
        stream_id = _to_text(ws.cell(header_row, col).value)
        if not _is_valid_stream_header(stream_id):
            continue
        stream_columns.append({
            'column_index': col,
            'column_letter': ws.cell(header_row, col).column_letter,
            'stream_id': stream_id,
            'description': _to_text(ws.cell(desc_row, col).value),
        })

    section_summaries: List[Dict[str, Any]] = []
    normalized_preview: List[Dict[str, Any]] = []
    total_property_rows = 0

    preview_streams = stream_columns[: cfg['max_stream_preview']]

    for section in cfg['sections']:
        properties = []
        for row in range(section['start_row'], section['end_row'] + 1):
            prop = _to_text(ws.cell(row, cfg['property_col']).value)
            if not prop:
                continue

            unit = _to_text(ws.cell(row, cfg['unit_col']).value)
            properties.append({'row': row, 'property': prop, 'unit': unit})
            total_property_rows += 1

            # Compact long-format preview for frontend review
            for s in preview_streams:
                value = ws.cell(row, s['column_index']).value
                if value in (None, ''):
                    continue
                if len(normalized_preview) >= cfg['max_record_preview']:
                    continue
                normalized_preview.append({
                    'section': section['label'],
                    'row': row,
                    'property': prop,
                    'unit': unit,
                    'stream_id': s['stream_id'],
                    'value': value,
                })

        section_summaries.append({
            'key': section['key'],
            'label': section['label'],
            'start_row': section['start_row'],
            'end_row': section['end_row'],
            'property_count': len(properties),
            'properties': properties,
        })

    case_title = _to_text(ws.cell(1, 1).value)
    named_ranges = []
    for defined_name in wb.defined_names.values():
        named_ranges.append({
            'name': getattr(defined_name, 'name', ''),
            'ref': getattr(defined_name, 'attr_text', ''),
        })

    warnings = []
    if not stream_columns:
        warnings.append('No stream columns detected in row 1 from column D onward.')

    baseline_case_name = _canonical_case_name(case_title)
    baseline_records = []
    for section in section_summaries:
        section_key = section['key']
        section_label = section['label']
        for prop in section['properties']:
            row = prop['row']
            for stream in stream_columns:
                value = values_ws.cell(row, stream['column_index']).value
                if value in (None, '', '---'):
                    continue
                baseline_records.append({
                    'case_name': baseline_case_name,
                    'source_filename': '',
                    'sheet_name': sheet_name,
                    'section_key': section_key,
                    'section_label': section_label,
                    'row_index': row,
                    'property_name': prop['property'],
                    'unit': prop['unit'],
                    'stream_id': stream['stream_id'],
                    'source_stream_id': stream['stream_id'],
                    'stream_description': stream['description'],
                    'value_text': str(value),
                })

    result = {
        'template_meta': {
            'sheet_name': sheet_name,
            'sheet_count': len(wb.sheetnames),
            'case_title': case_title,
            'max_row': ws.max_row,
            'max_col': ws.max_column,
        },
        'summary': {
            'stream_count': len(stream_columns),
            'section_count': len(section_summaries),
            'property_row_count': total_property_rows,
            'preview_stream_count': len(preview_streams),
            'preview_record_count': len(normalized_preview),
        },
        'stream_columns': stream_columns,
        'sections': section_summaries,
        'template_layout': _serialize_master_layout(ws),
        'baseline': {
            'case_name': baseline_case_name,
            'stream_count': len(stream_columns),
            'record_count': len(baseline_records),
            'records': baseline_records,
        },
        'normalized_preview': normalized_preview,
        'named_ranges': named_ranges,
        'warnings': warnings,
    }
    values_wb.close()
    wb.close()
    return result


def parse_hmb_case_workbook(workbook_path: str, source_filename: str = '', template_profile_payload: Dict[str, Any] | None = None) -> Dict[str, Any]:
    """Parse one HMB case workbook into normalized long-format records."""
    # Random cell access dominates this parser; normal mode is significantly
    # faster than read_only mode for this access pattern.
    with open(workbook_path, 'rb') as workbook_file:
        workbook_bytes = workbook_file.read()
    wb = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=False)
    case_name = _extract_case_name(wb, source_filename=source_filename)

    if 'Master' in wb.sheetnames and (template_profile_payload or {}).get('stream_columns'):
        parsed_master = _parse_master_case_workbook(wb, source_filename, template_profile_payload or {})
        wb.close()
        return parsed_master

    normalized_summary = _is_normalized_summary_workbook(wb)
    sheet_map = {
        'general': 'Overall',
        'vapour': 'Vapour Phase',
        'light_liquid': 'Liquid Phase',
        'heavy_liquid': 'Aqueous Phase',
    }

    if normalized_summary:
        phase_data = {
            key: {'layout': {'header_row': 1}, 'stream_ids': [], 'rows': {}}
            for key in sheet_map
        }
        source_stream_ids = []
        for summary_sheet in ('Conditions', 'Properties'):
            parsed_sections, parsed_streams = _parse_normalized_summary_sheet(wb[summary_sheet])
            if not source_stream_ids:
                source_stream_ids = parsed_streams
            for key, parsed_section in parsed_sections.items():
                phase_data[key]['stream_ids'] = parsed_section['stream_ids']
                for property_key, entries in parsed_section['rows'].items():
                    phase_data[key]['rows'].setdefault(property_key, []).extend(entries)
        sheet_map = {key: 'Conditions / Properties' for key in sheet_map}
    else:
        phase_data = {}
        for key, sheet_name in sheet_map.items():
            if sheet_name in wb.sheetnames:
                phase_data[key] = _parse_phase_property_sheet(wb[sheet_name])
            else:
                phase_data[key] = {'stream_ids': [], 'rows': {}, 'layout': None}
        source_stream_ids = phase_data['general']['stream_ids'] or phase_data['vapour']['stream_ids'] or phase_data['light_liquid']['stream_ids']

    template_sections = (template_profile_payload or {}).get('sections', [])
    template_streams = (template_profile_payload or {}).get('stream_columns', [])

    allowed_components = []
    for sec in template_sections:
        sec_key = _normalise_key(sec.get('label')).replace(' ', '_')
        if sec_key != 'composition':
            continue
        for p in (sec.get('properties') or []):
            prop_name = _normalise_text(p.get('property'))
            if prop_name:
                allowed_components.append(prop_name)

    comp_all = {}
    if normalized_summary:
        comp_all = _parse_normalized_composition(
            wb['Mole Fraction'],
            allowed_streams=source_stream_ids,
            allowed_components=allowed_components,
        )
    elif 'Composition (Overall)' in wb.sheetnames:
        comp_all = _parse_composition_overall(
            wb['Composition (Overall)'],
            allowed_streams=source_stream_ids,
            allowed_components=allowed_components,
        )
    if not comp_all and 'Composition (All)' in wb.sheetnames:
        comp_all = _parse_composition_all(
            wb['Composition (All)'],
            allowed_streams=source_stream_ids,
            allowed_components=allowed_components,
        )

    exceptions = {
        'unmatched_streams': [],
        'duplicate_stream_matches': [],
        'unit_mismatches': [],
        'unmapped_properties': [],
        'zero_filled_components': [],
        'composition_sum_warnings': [],
        'unmatched_streams_count': 0,
        'duplicate_stream_matches_count': 0,
        'unit_mismatches_count': 0,
        'unmapped_properties_count': 0,
        'zero_filled_components_count': 0,
        'composition_sum_warnings_count': 0,
    }

    records: List[Dict[str, Any]] = []

    if not template_sections or not template_streams:
        # Backward-compatible fallback when no template profile payload is supplied.
        ws = _pick_case_sheet(wb, HMB_CASE_PARSER_CONFIG['preferred_sheet_names'])
        layout = _detect_case_layout(ws, HMB_CASE_PARSER_CONFIG)
        stream_columns = _extract_stream_columns(ws, {
            'header_row': layout['stream_header_row'],
            'property_col': layout['property_col'],
            'unit_col': layout['unit_col'],
            'stream_start_col': layout['stream_start_col'],
        })
        current_section = {'key': 'general', 'label': 'General'}
        marker_cfg = HMB_CASE_PARSER_CONFIG.get('section_markers', {})
        for row in range(layout['stream_header_row'] + 1, ws.max_row + 1):
            property_name = _to_text(ws.cell(row, layout['property_col']).value)
            if not property_name:
                continue
            if property_name.isupper() and len(property_name) <= 32:
                current_section = _section_from_label(property_name, marker_cfg, current_section)
                continue
            current_section = _section_from_label(property_name, marker_cfg, current_section)
            unit = _to_text(ws.cell(row, layout['unit_col']).value)
            for stream in stream_columns:
                value = ws.cell(row, stream['column_index']).value
                if value in (None, '', '---'):
                    continue
                if _is_formula_or_ref_error(value):
                    continue
                records.append({
                    'case_name': case_name,
                    'source_filename': source_filename,
                    'sheet_name': ws.title,
                    'section_key': current_section['key'],
                    'section_label': current_section['label'],
                    'row_index': row,
                    'property_name': property_name,
                    'unit': unit,
                    'stream_id': stream['stream_id'],
                    'source_stream_id': stream['stream_id'],
                    'stream_description': '',
                    'value_text': str(value),
                })
        wb.close()
        return {
            'case_name': case_name,
            'sheet_name': ws.title,
            'stream_count': len(stream_columns),
            'record_count': len(records),
            'records': records,
            'exceptions': exceptions,
        }

    stream_map, unmatched, duplicates = _resolve_stream_mapping(template_streams, source_stream_ids)
    template_by_id = {
        _normalise_text(stream.get('stream_id')): stream
        for stream in template_streams
        if _normalise_text(stream.get('stream_id'))
    }
    template_id_by_source = {source_id: template_id for template_id, source_id in stream_map.items()}
    mapped_source_stream_ids = [sid for sid in source_stream_ids if sid in template_id_by_source]
    ignored_source_stream_ids = [sid for sid in source_stream_ids if sid not in template_id_by_source]
    for item in unmatched:
        _push_exception(exceptions, 'unmatched_streams', item)
    for item in duplicates:
        _push_exception(exceptions, 'duplicate_stream_matches', item)

    for section in template_sections:
        section_label = _normalise_text(section.get('label'))
        section_key = _normalise_key(section_label).replace(' ', '_')
        properties = section.get('properties', []) or []

        for source_stream_id in mapped_source_stream_ids:
            template_stream_id = template_id_by_source.get(source_stream_id, source_stream_id)
            template_stream = template_by_id.get(template_stream_id, {})

            comp_sum = 0.0
            comp_count = 0

            for prop in properties:
                prop_name = _normalise_text(prop.get('property'))
                prop_unit = _normalise_text(prop.get('unit'))
                row_idx = prop.get('row') or 0
                if not prop_name:
                    continue

                value = None
                source_unit = ''
                source_sheet = ''

                if section_key == 'composition':
                    source_sheet = 'Mole Fraction' if normalized_summary else 'Composition (All)'
                    comp_key = _normalise_component(prop_name)
                    comp_val = comp_all.get((source_stream_id, comp_key))
                    if comp_val is None:
                        value = 0.0
                        _push_exception(exceptions, 'zero_filled_components', {
                            'template_stream': template_stream_id,
                            'source_stream': source_stream_id,
                            'component': prop_name,
                        })
                    else:
                        value = comp_val
                    if _normalise_key(prop_name) != 'total':
                        comp_sum += float(value)
                        comp_count += 1
                else:
                    source_sheet_data = phase_data.get(section_key, {'rows': {}})
                    source_sheet = sheet_map.get(section_key, '')
                    found, converted_value, first_incompatible = _best_source_row(
                        source_sheet_data,
                        section_key,
                        prop_name,
                        source_stream_id,
                        prop_unit,
                    )
                    if found:
                        source_unit = _normalise_text(found.get('unit'))
                        value = converted_value
                        if value is None:
                            value = found['values'].get(source_stream_id)
                    else:
                        if first_incompatible is not None:
                            _push_exception(exceptions, 'unit_mismatches', {
                                'template_stream': template_stream_id,
                                'source_stream': source_stream_id,
                                'section': section_label,
                                'property': prop_name,
                                'template_unit': prop_unit or '<none>',
                                'source_unit': _normalise_text(first_incompatible.get('unit')) or '<none>',
                            })
                        else:
                            _push_exception(exceptions, 'unmapped_properties', {
                                'template_stream': template_stream_id,
                                'source_stream': source_stream_id,
                                'section': section_label,
                                'property': prop_name,
                                'source_sheet': source_sheet,
                            })

                if value in (None, '', '---'):
                    continue
                records.append({
                    'case_name': case_name,
                    'source_filename': source_filename,
                    'sheet_name': source_sheet,
                    'section_key': section_key,
                    'section_label': section_label,
                    'row_index': int(row_idx) if row_idx else 0,
                    'property_name': prop_name,
                    'unit': prop_unit,
                    'stream_id': template_stream_id,
                    'source_stream_id': source_stream_id,
                    'stream_description': _normalise_text(template_stream.get('description')),
                    'value_text': str(value),
                })

            if section_key == 'composition' and comp_count > 0:
                if abs(comp_sum - 1.0) > 0.005:
                    _push_exception(exceptions, 'composition_sum_warnings', {
                        'template_stream': template_stream_id,
                        'source_stream': source_stream_id,
                        'mole_fraction_sum': round(comp_sum, 6),
                    })

    wb.close()
    return {
        'case_name': case_name,
        'sheet_name': 'Conditions' if normalized_summary else 'Overall',
        'detected_format': 'normalized_summary' if normalized_summary else 'phase_workbook',
        'stream_count': len(mapped_source_stream_ids),
        'record_count': len(records),
        'records': records,
        'exceptions': exceptions,
        'stream_mapping': {
            'matched': stream_map,
            'source_stream_count': len(source_stream_ids),
            'mapped_source_stream_count': len(mapped_source_stream_ids),
            'ignored_source_stream_count': len(ignored_source_stream_ids),
            'ignored_source_streams': ignored_source_stream_ids[:120],
            'template_stream_count': len(template_streams),
        },
    }
