from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from rest_framework.test import APIRequestFactory, force_authenticate

from apps.project_organizer.models import Project
from apps.process_datasheet.hmb_extractor_view import (
    list_hmb_master_templates_view,
    analyze_hmb_master_template_view,
    execute_hmb_case_preview_view,
    preview_hmb_case_files_view,
    retrieve_hmb_master_template_view,
)
from apps.process_datasheet.models import HMBCaseImportBatch, HMBMasterTemplateProfile, HMBSourceUpload


class HMBTemplateAccessTests(TestCase):
    def setUp(self):
        users = get_user_model().objects
        self.owner = users.create_user(
            username='hmb-owner', email='hmb-owner@example.test', password='test'
        )
        self.profile_creator = users.create_user(
            username='hmb-profile-creator',
            email='hmb-profile-creator@example.test',
            password='test',
        )
        self.outsider = users.create_user(
            username='hmb-outsider', email='hmb-outsider@example.test', password='test'
        )
        self.project = Project.objects.create(name='HMB Project', created_by=self.owner)
        self.profile = HMBMasterTemplateProfile.objects.create(
            source_filename='master.xlsx',
            file_sha256='a' * 64,
            project=self.project,
            created_by=self.profile_creator,
        )
        self.factory = APIRequestFactory()

    def test_project_owner_can_list_profile_created_by_another_user(self):
        request = self.factory.get(
            '/api/v1/process-datasheet/datasheets/hmb-master-templates/',
            {'project_id': str(self.project.project_id)},
        )
        force_authenticate(request, user=self.owner)

        response = list_hmb_master_templates_view(request)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['count'], 1)
        self.assertEqual(response.data['results'][0]['id'], str(self.profile.id))

    def test_project_preserves_linked_owned_global_master(self):
        master = HMBMasterTemplateProfile.objects.create(
            source_filename='Master.xlsx', file_sha256='d' * 64, created_by=self.owner,
        )
        HMBCaseImportBatch.objects.create(
            project=self.project, template_profile=master, imported_by=self.owner,
            source_file_count=1, total_records=5135,
        )
        HMBMasterTemplateProfile.objects.create(
            source_filename='Private.xlsx', file_sha256='e' * 64, created_by=self.outsider,
        )
        request = self.factory.get('/templates/', {'project_id': str(self.project.pk)})
        force_authenticate(request, user=self.owner)
        response = list_hmb_master_templates_view(request)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['preferred_template_profile_id'], str(master.pk))
        self.assertEqual(response.data['results'][0]['id'], str(master.pk))
        self.assertEqual(response.data['count'], 2)
        master.refresh_from_db()
        self.assertIsNone(master.project_id)

    def test_project_owner_can_retrieve_profile_created_by_another_user(self):
        request = self.factory.get('/api/v1/process-datasheet/datasheets/hmb-master-templates/')
        force_authenticate(request, user=self.owner)

        response = retrieve_hmb_master_template_view(request, self.profile.id)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['template_profile']['id'], str(self.profile.id))

    def test_case_configuration_is_isolated_per_project(self):
        from apps.process_datasheet.hmb_extractor_view import _build_hmb_stream_comparison
        self.profile.analysis_payload = {'stream_columns': [{'stream_id': 'S-001'}], 'sections': []}
        self.profile.save()
        second = Project.objects.create(name='Second project', created_by=self.owner)
        for project, slots in [(self.project, ['Design', 'Turndown']), (second, ['Winter', 'Summer'])]:
            HMBSourceUpload.objects.create(project=project, uploaded_by=self.owner, upload_kind='output_template',
                storage_key=f'private/{project.pk}.xlsx', original_filename='Final.xlsx', file_sha256='f' * 64,
                metadata={'case_slots': slots})
        result = _build_hmb_stream_comparison(self.project, self.profile, 'S-001')
        self.assertEqual(result['case_names'], ['Design', 'Turndown'])
        self.assertNotIn('Winter', result['case_names'])
        self.assertNotIn('CASE A (1a)', result['case_names'])

    def test_outsider_cannot_list_or_retrieve_project_profile(self):
        list_request = self.factory.get(
            '/api/v1/process-datasheet/datasheets/hmb-master-templates/',
            {'project_id': str(self.project.project_id)},
        )
        force_authenticate(list_request, user=self.outsider)
        detail_request = self.factory.get('/api/v1/process-datasheet/datasheets/hmb-master-templates/')
        force_authenticate(detail_request, user=self.outsider)

        list_response = list_hmb_master_templates_view(list_request)
        detail_response = retrieve_hmb_master_template_view(detail_request, self.profile.id)

        self.assertEqual(list_response.status_code, 403)
        self.assertEqual(detail_response.status_code, 403)


class HMBUploadStorageTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username='hmb-uploader', email='hmb-uploader@example.test', password='test'
        )
        self.project = Project.objects.create(name='Stored HMB Project', created_by=self.user)
        self.factory = APIRequestFactory()

    @staticmethod
    def storage_result(key):
        return {
            'stored': True,
            'key': key,
            'sha256': 'b' * 64,
            'size': 128,
            'content_type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        }

    @patch('apps.process_datasheet.hmb_extractor_view.store_hmb_source')
    @patch('apps.process_datasheet.hmb_extractor_view.analyze_hmb_master_template')
    def test_master_template_creates_private_source_audit(self, analyze, store):
        analyze.return_value = {
            'template_meta': {'sheet_name': 'HMB'},
            'summary': {'stream_count': 1, 'section_count': 1, 'property_row_count': 1},
            'stream_columns': [],
            'sections': [],
            'template_layout': {},
            'baseline': {},
            'named_ranges': [],
            'warnings': [],
            'normalized_preview': [],
        }
        store.return_value = self.storage_result('media/hmb/master-templates/project/master.xlsx')
        upload = SimpleUploadedFile('master.xlsx', b'workbook-content')
        request = self.factory.post(
            '/api/v1/process-datasheet/datasheets/analyze-hmb-master-template/',
            {'master_template_file': upload, 'project_id': str(self.project.project_id)},
            format='multipart',
        )
        force_authenticate(request, user=self.user)

        response = analyze_hmb_master_template_view(request)

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.data['source_stored'])
        source = HMBSourceUpload.objects.get()
        self.assertEqual(source.upload_kind, HMBSourceUpload.KIND_MASTER_TEMPLATE)
        self.assertEqual(source.project, self.project)
        self.assertEqual(str(source.template_profile_id), response.data['template_profile_id'])

    @patch('apps.process_datasheet.hmb_extractor_view.store_hmb_source')
    def test_rejected_master_template_never_reaches_storage(self, store):
        upload = SimpleUploadedFile('master.txt', b'not-an-excel-workbook')
        request = self.factory.post(
            '/api/v1/process-datasheet/datasheets/analyze-hmb-master-template/',
            {'master_template_file': upload, 'project_id': str(self.project.project_id)},
            format='multipart',
        )
        force_authenticate(request, user=self.user)

        response = analyze_hmb_master_template_view(request)

        self.assertEqual(response.status_code, 400)
        store.assert_not_called()
        self.assertFalse(HMBSourceUpload.objects.exists())

    @patch('apps.process_datasheet.hmb_extractor_view.store_hmb_source')
    @patch('apps.process_datasheet.hmb_extractor_view.parse_hmb_case_workbook')
    def test_case_preview_source_is_linked_when_import_executes(self, parse, store):
        profile = HMBMasterTemplateProfile.objects.create(
            source_filename='master.xlsx',
            file_sha256='c' * 64,
            project=self.project,
            created_by=self.user,
        )
        parse.return_value = {
            'case_name': 'CASE A (1a)',
            'detected_format': 'phase_workbook',
            'stream_count': 1,
            'record_count': 1,
            'records': [{
                'source_filename': 'case.xlsx',
                'stream_id': '1001',
                'source_stream_id': '1001',
                'section_key': 'general',
                'section_label': 'General',
                'property_name': 'Temperature',
                'unit': 'C',
                'value_text': '25',
                'row_index': 1,
            }],
        }
        store.return_value = self.storage_result('media/hmb/case-files/project/case.xlsx')
        upload = SimpleUploadedFile('case.xlsx', b'case-workbook-content')
        preview_request = self.factory.post(
            '/api/v1/process-datasheet/datasheets/preview-hmb-cases/',
            {
                'case_files': upload,
                'project_id': str(self.project.project_id),
                'template_profile_id': str(profile.id),
            },
            format='multipart',
        )
        force_authenticate(preview_request, user=self.user)

        preview_response = preview_hmb_case_files_view(preview_request)

        self.assertEqual(preview_response.status_code, 200)
        source = HMBSourceUpload.objects.get(upload_kind=HMBSourceUpload.KIND_CASE_FILE)
        self.assertTrue(preview_response.data['files'][0]['source_stored'])
        self.assertEqual(preview_response.data['files'][0]['source_upload_id'], str(source.id))
        self.assertEqual(source.status, HMBSourceUpload.STATUS_ANALYZED)

        execute_request = self.factory.post(
            '/api/v1/process-datasheet/datasheets/execute-hmb-cases/',
            {
                'preview_token': preview_response.data['preview_token'],
                'case_assignments': {'case.xlsx': 'CASE A (1a)'},
            },
            format='json',
        )
        force_authenticate(execute_request, user=self.user)
        execute_response = execute_hmb_case_preview_view(execute_request)

        self.assertEqual(execute_response.status_code, 200)
        source.refresh_from_db()
        self.assertEqual(source.status, HMBSourceUpload.STATUS_IMPORTED)
        self.assertIsNotNone(source.imported_at)
        self.assertEqual(source.import_batch, HMBCaseImportBatch.objects.get())
        retry = self.factory.post('/execute/', {'preview_token': preview_response.data['preview_token']}, format='json')
        force_authenticate(retry, user=self.user)
        retried = execute_hmb_case_preview_view(retry)
        self.assertEqual(retried.status_code, 200)
        self.assertTrue(retried.data['already_imported'])
        self.assertEqual(HMBCaseImportBatch.objects.count(), 1)
        upload = SimpleUploadedFile('case.xlsx', b'case-workbook-content')
        store.return_value = self.storage_result('media/hmb/case-files/project/retry.xlsx')
        preview_request = self.factory.post('/preview/', {
            'case_files': upload, 'project_id': str(self.project.pk), 'template_profile_id': str(profile.pk),
        }, format='multipart')
        force_authenticate(preview_request, user=self.user)
        replacement_preview = preview_hmb_case_files_view(preview_request)
        self.assertEqual(replacement_preview.status_code, 200)
        denied = self.factory.post('/execute/', {'preview_token': replacement_preview.data['preview_token']}, format='json')
        force_authenticate(denied, user=self.user)
        self.assertEqual(execute_hmb_case_preview_view(denied).status_code, 409)
        self.assertEqual(HMBCaseImportBatch.objects.count(), 1)
        from django.core.cache import cache
        from apps.process_datasheet.hmb_extractor_view import HMB_IMPORT_PREVIEW_CONFIG
        cache_key = HMB_IMPORT_PREVIEW_CONFIG['cache_prefix'] + replacement_preview.data['preview_token']
        cached = cache.get(cache_key)
        cached['files'][0]['exceptions'] = {'unit_mismatches_count': 1}
        cache.set(cache_key, cached)
        blocked = self.factory.post('/execute/', {'preview_token': replacement_preview.data['preview_token'], 'replace_existing': True}, format='json')
        force_authenticate(blocked, user=self.user)
        self.assertEqual(execute_hmb_case_preview_view(blocked).status_code, 400)
        self.assertEqual(HMBCaseImportBatch.objects.count(), 1)

    def test_output_template_round_trip_and_project_authorization(self):
        import io
        import tempfile
        from openpyxl import Workbook, load_workbook
        from django.test import override_settings
        from apps.process_datasheet.hmb_output_view import hmb_output_template_view, hmb_final_export_view
        from apps.process_datasheet.models import HMBCaseRecord

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Overall'
        for coordinate, value in {'A4': 'Phase', 'B4': 'Property', 'C4': 'Unit', 'D4': 'CASE A (1a)', 'A5': 'Overall', 'B5': 'Temperature', 'C5': 'F'}.items():
            sheet[coordinate] = value
        content = io.BytesIO()
        workbook.save(content)
        profile = HMBMasterTemplateProfile.objects.create(
            source_filename='Master.xlsx', file_sha256='f' * 64, created_by=self.user,
            analysis_payload={'stream_columns': [{'stream_id': '1001'}], 'sections': [
                {'key': 'general', 'label': 'General', 'properties': [{'property': 'Temperature', 'unit': 'F', 'row': 6}]}]},
        )
        batch = HMBCaseImportBatch.objects.create(project=self.project, template_profile=profile, imported_by=self.user)
        HMBCaseRecord.objects.create(batch=batch, project=self.project, template_profile=profile,
            source_filename='case.xlsx', case_name='CASE A (1a)', stream_id='1001', section_key='general',
            section_label='General', property_name='Temperature', unit='F', value_text='0')
        with tempfile.TemporaryDirectory() as directory, override_settings(BASE_DIR=directory, USE_S3=False, DEBUG=True):
            request = self.factory.post('/output-template/', {'output_template_file': SimpleUploadedFile('Final.xlsx', content.getvalue())}, format='multipart')
            force_authenticate(request, user=self.user)
            saved = hmb_output_template_view(request, self.project.pk)
            self.assertEqual(saved.status_code, 201)
            request = self.factory.post('/export/', {'template_profile_id': str(profile.pk),
                'output_template_id': saved.data['output_template']['id'], 'stream_ids': ['1001']}, format='json')
            force_authenticate(request, user=self.user)
            result = hmb_final_export_view(request, self.project.pk)
            self.assertEqual(result.status_code, 200)
            exported = load_workbook(io.BytesIO(result.content))
            self.assertEqual(exported['1001']['D5'].value, 0)
            exported.close()
            stranger = get_user_model().objects.create_user(username='stranger', email='stranger@example.test')
            request = self.factory.get('/output-template/')
            force_authenticate(request, user=stranger)
            self.assertEqual(hmb_output_template_view(request, self.project.pk).status_code, 403)