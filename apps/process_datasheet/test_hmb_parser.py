import tempfile
import os
from pathlib import Path
from unittest import TestCase, skipUnless

from openpyxl import Workbook, load_workbook
from io import BytesIO
from .services.hmb_export import build_final_workbook

from .hmb_master_template_parser import (
    _canonical_case_name,
    analyze_hmb_master_template,
    parse_hmb_case_workbook,
    parse_hmb_csv_file,
    _property_candidates,
)


class HMBWorkbookParserTests(TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.master_path = self.root / 'Master.xlsx'
        self._write_master(self.master_path, populated=False)
        self.template = analyze_hmb_master_template(str(self.master_path))

    def tearDown(self):
        self.temp_dir.cleanup()

    @staticmethod
    def _write_master(path, populated):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Master'
        sheet['A1'] = 'CASE A (1a) - baseline'
        sheet['C1'] = 'Stream'
        sheet['D1'] = '1048'
        sheet['C2'] = 'Description'
        sheet['D2'] = 'Test stream'
        sheet['A3'] = 'Phase'
        sheet['B3'] = 'Property'
        sheet['C4'] = 'Unit'
        sheet['A43'] = 'Composition'
        sheet['B6'] = 'Temperature'
        sheet['C6'] = 'F'
        sheet['B43'] = 'H2S'
        sheet['C43'] = 'mol frac.'
        if populated:
            sheet['D6'] = 100.5
            sheet['D43'] = 0.01
        workbook.save(path)
        workbook.close()

    def test_canonical_case_names(self):
        self.assertEqual(_canonical_case_name('Case_1B_Stream_Summary.xlsx'), 'CASE A (1b)')
        self.assertEqual(_canonical_case_name('CASE C (2a)'), 'CASE C (2a)')
        self.assertEqual(_canonical_case_name('CASE-A-1a-SUM-MAX OIL-GAS-REV D'), 'CASE A (1a)')
        self.assertEqual(_canonical_case_name('Relief scenario.xlsx'), 'Relief scenario')

    def test_project_master_detects_shifted_headers_and_custom_sections(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Project Bravo'
        sheet['D6'] = 'Stream'
        sheet['E6'] = 'BR-001'
        sheet['F6'] = 'BR-002'
        sheet['D7'] = 'Description'
        sheet['B8'] = 'Phase'
        sheet['C8'] = 'Property'
        sheet['D8'] = 'Unit'
        sheet['B10'] = 'Overall'
        sheet['C10'] = 'Temperature'
        sheet['D10'] = 'C'
        sheet['B13'] = 'Solid'
        sheet['C13'] = 'Mass Flow'
        sheet['D13'] = 'kg/h'
        path = self.root / 'ProjectBravo.xlsx'
        workbook.save(path)
        workbook.close()
        result = analyze_hmb_master_template(str(path))
        self.assertEqual([stream['stream_id'] for stream in result['stream_columns']], ['BR-001', 'BR-002'])
        self.assertEqual([section['key'] for section in result['sections']], ['general', 'solid'])
        self.assertEqual(result['sections'][1]['properties'][0]['row'], 13)
        self.assertEqual(result['template_meta']['layout']['header_row'], 8)
        source = load_workbook(BytesIO(path.read_bytes()))
        source.active.insert_rows(10, 3)
        source.active['E13'] = 25
        source.active['F13'] = 30
        source.active['E16'] = 100
        source.active['F16'] = 0
        source_path = self.root / 'Summer operating.xlsx'
        source.save(source_path)
        source.close()
        parsed = parse_hmb_case_workbook(str(source_path), source_path.name, result)
        self.assertEqual(parsed['case_name'], 'Summer operating')
        self.assertEqual(parsed['record_count'], 4)
        temperature = next(record for record in parsed['records'] if record['stream_id'] == 'BR-001' and record['property_name'] == 'Temperature')
        self.assertEqual(temperature['value_text'], '25')
        self.assertEqual(temperature['source_metadata']['cell'], 'E13')
        self.assertEqual(temperature['row_index'], 10)

    def test_csv_long_layout_maps_to_template(self):
        path = self.root / 'Case_A_1b.csv'
        path.write_text(
            'Case,Stream ID,Phase,Parameter,UOM,Result\n'
            'CASE A (1b),1048,Overall,Temperature,F,64.61\n'
            'CASE A (1b),1048,Composition,H2S,mol frac.,0.0012\n',
            encoding='utf-8',
        )

        parsed = parse_hmb_csv_file(str(path), path.name, self.template)

        self.assertEqual(parsed['detected_format'], 'csv_long')
        self.assertEqual(parsed['case_name'], 'CASE A (1b)')
        self.assertEqual(parsed['record_count'], 2)
        self.assertEqual({row['property_name'] for row in parsed['records']}, {'Temperature', 'H2S'})

    def test_csv_wide_layout_maps_stream_columns(self):
        path = self.root / 'CASE B (1a).csv'
        path.write_text(
            'Section,Property,Unit,1048\n'
            'General,Temperature,F,71.2\n'
            'Composition,H2S,mol frac.,0.002\n',
            encoding='utf-8',
        )

        parsed = parse_hmb_csv_file(str(path), path.name, self.template)

        self.assertEqual(parsed['detected_format'], 'csv_wide')
        self.assertEqual(parsed['case_name'], 'CASE B (1a)')
        self.assertEqual(parsed['stream_count'], 1)
        self.assertEqual(parsed['record_count'], 2)

    def test_normalized_summary_maps_to_master(self):
        path = self.root / 'Case_1B_Stream_Summary.xlsx'
        workbook = Workbook()
        conditions = workbook.active
        conditions.title = 'Conditions'
        conditions.append(['Stream', 'Temperature: (F) | Overall'])
        conditions.append(['1048', 64.61])
        properties = workbook.create_sheet('Properties')
        properties.append(['Stream', 'Molecular Weight | Overall'])
        properties.append(['1048', 18.06])
        composition = workbook.create_sheet('Mole Fraction')
        composition.append(['Stream', 'H2S'])
        composition.append(['1048', 0.0012])
        workbook.save(path)
        workbook.close()

        parsed = parse_hmb_case_workbook(str(path), path.name, self.template)

        self.assertEqual(parsed['case_name'], 'CASE A (1b)')
        self.assertEqual(parsed['detected_format'], 'normalized_summary')
        self.assertEqual(parsed['stream_count'], 1)
        self.assertEqual(parsed['stream_mapping']['mapped_source_stream_count'], 1)
        values = {(row['section_key'], row['property_name']): row['value_text'] for row in parsed['records']}
        self.assertEqual(values[('general', 'Temperature')], '64.61')
        self.assertEqual(values[('composition', 'H2S')], '0.0012')
        temperature = next(record for record in parsed['records'] if record['property_name'] == 'Temperature')
        self.assertEqual(temperature['source_metadata']['cell'], 'B2')
        self.assertEqual(temperature['source_metadata']['sheet'], 'Conditions')
        component = next(record for record in parsed['records'] if record['property_name'] == 'H2S')
        self.assertEqual(component['source_metadata']['cell'], 'B2')
        self.assertEqual(component['source_metadata']['sheet'], 'Mole Fraction')

    def test_missing_composition_is_not_measured_zero(self):
        path = self.root / 'Case_1B_Stream_Summary.xlsx'
        workbook = Workbook()
        conditions = workbook.active
        conditions.title = 'Conditions'
        conditions.append(['Stream', 'Temperature: (F) | Overall'])
        conditions.append(['1048', 0])
        properties = workbook.create_sheet('Properties')
        properties.append(['Stream', 'Molecular Weight | Overall'])
        properties.append(['1048', 18])
        composition = workbook.create_sheet('Mole Fraction')
        composition.append(['Stream', 'H2S'])
        composition.append(['1048', None])
        workbook.save(path)
        workbook.close()
        parsed = parse_hmb_case_workbook(str(path), path.name, self.template)
        self.assertEqual(len(parsed['records']), 1)
        self.assertEqual(parsed['records'][0]['value_text'], '0')
        self.assertEqual(parsed['exceptions']['missing_components_count'], 1)
        self.assertEqual(parsed['exceptions']['zero_filled_components_count'], 0)

    def test_output_template_is_not_a_master(self):
        workbook = Workbook()
        workbook.active['D4'] = 'CASE A (1a)'
        path = self.root / 'Final_Template.xlsx'
        workbook.save(path)
        workbook.close()
        with self.assertRaisesRegex(ValueError, 'final comparison template'):
            analyze_hmb_master_template(str(path))

    def test_distinct_engineering_quantities_are_not_aliases(self):
        self.assertNotIn('cp/(cp - r)', _property_candidates('vapour', 'Compressibility'))
        self.assertNotIn('kinematic viscosity', _property_candidates('vapour', 'Viscosity'))
        self.assertNotIn('liq. mass density (std. cond)', _property_candidates('light_liquid', 'Mass Density'))

    def test_populated_master_is_case_data(self):
        path = self.root / 'Case 1A Template.xlsx'
        self._write_master(path, populated=True)

        parsed = parse_hmb_case_workbook(str(path), path.name, self.template)

        self.assertEqual(parsed['case_name'], 'CASE A (1a)')
        self.assertEqual(parsed['detected_format'], 'master_case')
        self.assertEqual(parsed['stream_count'], 1)
        self.assertEqual(parsed['record_count'], 2)
        self.assertEqual({row['value_text'] for row in parsed['records']}, {'100.5', '0.01'})

    def test_final_export_preserves_layout_cases_and_blanks(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Overall'
        sheet.merge_cells('A1:B3')
        sheet['A4'] = 'Phase'
        sheet['B4'] = 'Property'
        sheet['C4'] = 'Unit'
        sheet['D4'] = 'CASE A (1a)'
        sheet['E4'] = 'CASE A (1b)'
        sheet['A5'] = 'Overall'
        sheet['B5'] = 'Temperature'
        sheet['C5'] = 'F'
        sheet['D5'] = '=HLOOKUP(A1,#REF!,6,FALSE)'
        sheet['E5'] = '=HLOOKUP(A1,#REF!,6,FALSE)'
        sheet.column_dimensions['B'].width = 35
        sheet.print_area = 'A1:E79'
        content = BytesIO()
        workbook.save(content)
        comparisons = [{
            'stream': {'stream_id': stream, 'description': 'Test'},
            'case_names': ['CASE A (1a)', 'CASE A (1b)'],
            'rows': [{'section_key': 'general', 'property': 'Temperature', 'unit': 'F',
                      'values': {'CASE A (1a)': '0', 'CASE A (1b)': ''}}],
        } for stream in ['1001', '1048']]
        result = load_workbook(BytesIO(build_final_workbook(content.getvalue(), comparisons)))
        for stream in ['1001', '1048']:
            exported = result[stream]
            self.assertEqual(exported['D5'].value, 0)
            self.assertIsNone(exported['E5'].value)
            self.assertIn('A1:B3', [str(merged) for merged in exported.merged_cells.ranges])
            self.assertEqual(exported.column_dimensions['B'].width, 35)
            self.assertIn('$A$1:$E$79', str(exported.print_area))
            self.assertFalse(any(cell.data_type == 'f' for row in exported for cell in row))
        self.assertIn('Validation', result.sheetnames)
        self.assertIn('Sources', result.sheetnames)
        result.close()

    @skipUnless(os.environ.get('HMB_SAMPLE_DIR'), 'Set HMB_SAMPLE_DIR to validate the supplied workbooks.')
    def test_supplied_workbooks_and_all_stream_export(self):
        root = Path(os.environ['HMB_SAMPLE_DIR'])
        master = analyze_hmb_master_template(str(root / 'Template' / 'Master.xlsx'))
        self.assertEqual(len(master['stream_columns']), 86)
        parsed_cases = []
        for filename in ['Case_1B_Stream_Summary.xlsx', 'Case_1A_Stream_Report.xlsx']:
            parsed = parse_hmb_case_workbook(str(root / filename), filename, master)
            parsed_cases.append(parsed)
            source = load_workbook(BytesIO((root / filename).read_bytes()), data_only=True)
            traced = 0
            for record in parsed['records']:
                provenance = record.get('source_metadata', {})
                if not provenance:
                    continue
                actual = source[provenance['sheet']][provenance['cell']].value
                self.assertIsNotNone(actual, (filename, provenance))
                self.assertEqual(str(actual), provenance['value'])
                if provenance['status'] == 'present':
                    self.assertAlmostEqual(float(actual), float(record['value_text']), places=8)
                traced += 1
            self.assertEqual(traced, len(parsed['records']), filename)
            source.close()
            self.assertEqual(parsed['stream_count'], 86)
            self.assertEqual(parsed['exceptions']['zero_filled_components_count'], 0)
        case_b = parsed_cases[0]
        samples = {(record['stream_id'], record['section_key'], record['property_name']): record['value_text'] for record in case_b['records']}
        self.assertEqual(float(samples['1001', 'general', 'Temperature']), 155.7)
        self.assertEqual(float(samples['1048', 'general', 'Temperature']), 64.61)
        self.assertEqual(float(samples['1048', 'composition', 'H2S']), 0.0006)
        from .hmb_master_template_parser import hmb_property_identity
        comparisons = []
        for stream in master['stream_columns']:
            rows = {}
            for parsed in parsed_cases:
                for record in parsed['records']:
                    if record['stream_id'] != stream['stream_id']:
                        continue
                    identity = hmb_property_identity(record['section_key'], record['property_name'], record['unit'])
                    row = rows.setdefault(identity, {'section_key': record['section_key'], 'property': record['property_name'], 'unit': record['unit'], 'values': {}, 'sources': {}})
                    row['values'][parsed['case_name']] = record['value_text']
                    row['sources'][parsed['case_name']] = record['source_metadata']
            comparisons.append({'stream': stream, 'case_names': [parsed['case_name'] for parsed in parsed_cases], 'rows': list(rows.values())})
        template_bytes = (root / 'Template' / 'Final_Template.xlsx').read_bytes()
        from .services.hmb_export import inspect_output_template
        layout = inspect_output_template(template_bytes)
        exported = load_workbook(BytesIO(build_final_workbook(template_bytes, comparisons)))
        self.assertEqual(len(exported.sheetnames), 88)
        for comparison in comparisons:
            sheet = exported[str(comparison['stream']['stream_id'])]
            by_identity = {hmb_property_identity(row['section_key'], row['property'], row['unit']): row for row in comparison['rows']}
            for prop in layout['properties']:
                for case, column in layout['cases'].items():
                    expected = by_identity.get(tuple(prop['identity']), {}).get('values', {}).get(case)
                    actual = sheet.cell(prop['row'], column).value
                    if expected is None:
                        self.assertIsNone(actual)
                    else:
                        self.assertAlmostEqual(float(actual), float(expected), places=7)
            self.assertFalse(any(cell.data_type == 'f' and '#REF!' in cell.value for row in sheet for cell in row))
        exported.close()

    def test_dynamic_final_layout_and_custom_case_columns(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Project Bravo Deliverable'
        sheet['D3'] = 'Stream'
        sheet['D4'] = 'Description'
        sheet['B7'] = 'Phase'
        sheet['C7'] = 'Parameter'
        sheet['D7'] = 'UOM'
        sheet['E7'] = 'Summer operating'
        sheet['F7'] = 'Winter operating'
        sheet['B9'] = 'Overall'
        sheet['C9'] = 'Temperature'
        sheet['D9'] = 'C'
        sheet['B15'] = 'Solid'
        sheet['C15'] = 'Mass Flow'
        sheet['D15'] = 'kg/h'
        content = BytesIO()
        workbook.save(content)
        comparisons = [{'stream': {'stream_id': 'BR-001'},
            'case_names': ['Summer operating', 'Winter operating', 'Turndown'],
            'rows': [
                {'section_key': 'general', 'property': 'Temperature', 'unit': 'C', 'values': {'Summer operating': '35', 'Turndown': '10'}},
                {'section_key': 'solid', 'property': 'Mass Flow', 'unit': 'kg/h', 'values': {'Winter operating': '0'}},
            ]}]
        result = load_workbook(BytesIO(build_final_workbook(content.getvalue(), comparisons)))
        output = result['BR-001']
        self.assertEqual(output['E9'].value, 35)
        self.assertEqual(output['F15'].value, 0)
        self.assertEqual(output['G9'].value, 10)
        self.assertEqual(output['G7'].value, 'Turndown')
        self.assertEqual(output['E3'].value, 'BR-001')
        self.assertEqual(output.freeze_panes, 'E8')
        self.assertIsNone(output['F9'].value)
        result.close()
