from copy import copy
from io import BytesIO
import math
import re

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter, range_boundaries

from ..hmb_master_template_parser import canonical_hmb_case_name, hmb_property_identity
from .hmb_layout import detect_template_layout, template_sections


def inspect_output_template(content):
    workbook = load_workbook(BytesIO(content), data_only=False, keep_links=False)
    try:
        layout = detect_template_layout(workbook, 'output')
        sheet = workbook[layout['sheet_name']]
        cases = {}
        for column in range(layout['stream_start_col'], sheet.max_column + 1):
            label = str(sheet.cell(layout['header_row'], column).value or '').strip()
            if label:
                case = canonical_hmb_case_name(label)
                if case in cases:
                    raise ValueError(f'Duplicate final-template case slot: {case}')
                cases[case] = column
        if not cases:
            raise ValueError('Final template requires case labels next to its Property and Unit headers.')
        properties = []
        identities = set()
        for section in template_sections(sheet, layout):
            for prop in section['properties']:
                identity = hmb_property_identity(section['key'], prop['property'], prop['unit'])
                if identity in identities:
                    raise ValueError(f'Duplicate final-template property at row {prop["row"]}.')
                identities.add(identity)
                properties.append({**prop, 'identity': identity})
        if not properties:
            raise ValueError('Final template contains no property rows.')
        return {**layout, 'sheet': sheet.title, 'cases': cases, 'properties': properties}
    finally:
        workbook.close()


def build_final_workbook(content, comparisons):
    if not comparisons:
        raise ValueError('Select at least one stream for export.')
    layout = inspect_output_template(content)
    workbook = load_workbook(BytesIO(content), data_only=False, keep_links=False)
    original = workbook[layout['sheet']]
    validation = workbook.create_sheet('Validation')
    validation.append(['Stream', 'Case', 'Phase', 'Property', 'Unit', 'Status'])
    provenance = workbook.create_sheet('Sources')
    provenance.append(['Stream', 'Case', 'Phase', 'Property', 'Source file', 'Source stream', 'Status', 'Sheet', 'Cell', 'Original value', 'Original unit'])
    for comparison in comparisons:
        if comparison.get('conflicts'):
            raise ValueError('Resolve conflicting imported values before exporting.')
        stream = comparison['stream']
        stream_id = str(stream['stream_id'])
        sheet = workbook.copy_worksheet(original)
        name = re.sub(r'[\[\]:*?/\\]', '_', stream_id)[:31] or 'Stream'
        sheet.title = name
        sheet.print_area = str(original.print_area).rsplit('!', 1)[-1] if original.print_area else None
        sheet.freeze_panes = f'{get_column_letter(layout["stream_start_col"])}{layout["header_row"] + 1}'
        cases = dict(layout['cases'])
        for case in comparison['case_names']:
            if case not in cases:
                column = max(sheet.max_column, max(cases.values())) + 1
                previous = max(cases.values())
                cases[case] = column
                for row in range(1, original.max_row + 1):
                    sheet.cell(row, column)._style = copy(sheet.cell(row, previous)._style)
                sheet.column_dimensions[get_column_letter(column)].width = sheet.column_dimensions[get_column_letter(previous)].width
                print_last_row = original.max_row
                if original.print_area:
                    print_range = str(original.print_area).rsplit('!', 1)[-1]
                    print_last_row = max(print_last_row, range_boundaries(print_range)[3])
                sheet.print_area = f'A1:{get_column_letter(column)}{print_last_row}'
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == 'f' and '#REF!' in cell.value:
                    cell.value = None
        for merged in original.merged_cells.ranges:
            if merged.min_row == 1 and merged.min_col == 1 and merged.max_row < layout['header_row'] and merged.max_col < layout['stream_start_col']:
                sheet['A1'] = stream_id
                sheet['A1'].data_type = 's'
        for case, column in cases.items():
            for row, value in [(layout['stream_label_row'], stream_id),
                               (layout['stream_description_row'], stream.get('description', '')),
                               (layout['pfd_row'], stream.get('pfd_no', '')),
                               (layout['header_row'], case)]:
                if row:
                    sheet.cell(row, column, value).data_type = 's'
        rows_by_identity = {
            hmb_property_identity(row['section_key'], row['property'], row['unit']): row
            for row in comparison['rows']
        }
        for property_row in layout['properties']:
            identity = tuple(property_row['identity'])
            mapped = rows_by_identity.get(identity, {})
            for case, column in cases.items():
                value = mapped.get('values', {}).get(case, '')
                cell = sheet.cell(property_row['row'], column)
                cell.value = None
                if value not in ('', None):
                    try:
                        numeric = float(value)
                        cell.value = numeric if math.isfinite(numeric) else str(value)
                    except (ValueError, TypeError):
                        cell.value = str(value)
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
                else:
                    cell.comment = Comment('No validated source value. Blank is not zero.', 'RADAI')
                    validation.append([stream_id, case, identity[0], property_row['property'], property_row['unit'], 'missing'])
                source = mapped.get('sources', {}).get(case, {})
                if value not in ('', None):
                    provenance.append([stream_id, case, identity[0], property_row['property'], source.get('filename', ''), source.get('stream_id', ''), source.get('status', 'legacy'), source.get('sheet', ''), source.get('cell', ''), source.get('value', ''), source.get('unit', '')])
    workbook.remove(original)
    for name in list(workbook.defined_names):
        definition = workbook.defined_names[name]
        if '#REF!' in str(definition.attr_text) or '[' in str(definition.attr_text):
            del workbook.defined_names[name]
    workbook.move_sheet(validation, offset=len(workbook.sheetnames) - 1)
    workbook.move_sheet(provenance, offset=len(workbook.sheetnames) - 1)
    for sheet in (validation, provenance):
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()