"""Bounded, deterministic document extraction for reviewed replica records.

Nothing here evaluates formulas, follows links, performs OCR, calls an AI service,
or updates projects. Suggestions preserve source values rather than interpreting
dates, quantities, or completion percentages.
"""

from __future__ import annotations

import base64
import csv
import io
import re
import zipfile
import zlib
from datetime import date, datetime, time
from pathlib import PureWindowsPath
from xml.parsers import expat


MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_ZIP_ENTRIES = 2000
MAX_ZIP_UNCOMPRESSED_BYTES = 60 * 1024 * 1024
MAX_ZIP_MEMBER_BYTES = 20 * 1024 * 1024
MAX_PDF_PAGES = 50
MAX_PDF_OBJECTS = 20_000
MAX_PDF_STREAM_BYTES = 4 * 1024 * 1024
MAX_PDF_TOTAL_STREAM_BYTES = 20 * 1024 * 1024
MAX_SHEETS = 20
MAX_TABLES = 20
MAX_ROWS = 2000
MAX_COLUMNS = 100
MAX_SECTIONS = 2000
MAX_SECTION_CHARS = 12_000
MAX_TOTAL_CHARS = 300_000
MAX_SUGGESTIONS = 500
MAX_VALUE_CHARS = 1000
MAX_EVIDENCE_CHARS = 2000

SUPPORTED_EXTENSIONS = frozenset({'.pdf', '.xlsx', '.docx', '.txt', '.csv'})

# A source label is required. These are recognition rules, not field mappings.
_LABELS = frozenset({
    'project code', 'project number', 'project no', 'project id',
    'project name', 'project title',
    'date', 'start date', 'end date', 'due date', 'target date',
    'planned start', 'planned finish', 'actual start', 'actual finish',
    'completion date', 'milestone', 'milestone name', 'deliverable',
    'deliverable name', 'document number', 'document no', 'document title',
    'status', 'quantity', 'qty', 'unit', 'progress', 'progress %',
    'completion', 'completion %', 'percent complete', 'percentage complete',
})
_EXPLICIT_PAIR = re.compile(r'^\s*([^:=\n]{1,80})\s*[:=]\s*(\S.*?)\s*$')


def _label_key(value):
    return re.sub(r'\s+', ' ', str(value).strip().lower().replace('.', '')
                  .replace('(', '').replace(')', '')).strip()


def _known_label(value):
    return _label_key(value) in _LABELS


def _as_text(value):
    if value is None:
        return ''
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value).strip()


class _Output:
    def __init__(self):
        self.sections = []
        self.suggestions = []
        self.warnings = []
        self.chars = 0
        self._seen = set()

    def warn(self, message):
        if message not in self.warnings:
            self.warnings.append(message)

    @property
    def full(self):
        return len(self.sections) >= MAX_SECTIONS or self.chars >= MAX_TOTAL_CHARS

    def section(self, location, text):
        text = text.strip()
        if not text:
            return ''
        if self.full:
            self.warn('Output limit reached; remaining source content was not extracted.')
            return ''
        available = min(MAX_SECTION_CHARS, MAX_TOTAL_CHARS - self.chars)
        retained = text[:available]
        if len(retained) < len(text):
            self.warn('Some source text was truncated by the extraction character limits.')
        self.sections.append({'location': location, 'text': retained})
        self.chars += len(retained)
        return retained

    def suggest(self, label, value, location, evidence):
        label, value = _as_text(label), _as_text(value)
        if not value or not _known_label(label):
            return
        if len(value) > MAX_VALUE_CHARS or len(evidence) > MAX_EVIDENCE_CHARS:
            self.warn('Long values or evidence rows were omitted from suggestions; review the source text.')
            return
        key = (label, value, location)
        if key in self._seen:
            return
        if len(self.suggestions) >= MAX_SUGGESTIONS:
            self.warn('Suggestion limit reached; additional source values were not suggested.')
            return
        self._seen.add(key)
        self.suggestions.append({
            'label': label, 'value': value, 'location': location, 'evidence': evidence,
        })

    def labeled_lines(self, text, location):
        for line in text.splitlines():
            match = _EXPLICIT_PAIR.match(line)
            if match:
                self.suggest(match[1], match[2], location, line.strip())

    def result(self):
        return {
            'sections': self.sections,
            'suggestions': self.suggestions,
            'warnings': self.warnings,
        }


def _read_bounded(file_obj):
    """Read from the beginning where possible, restoring the caller's position."""
    previous = None
    try:
        previous = file_obj.tell()
        file_obj.seek(0)
    except (AttributeError, OSError, io.UnsupportedOperation):
        previous = None
    try:
        chunks = []
        remaining = MAX_FILE_BYTES + 1
        while remaining:
            chunk = file_obj.read(min(64 * 1024, remaining))
            if not chunk:
                break
            if not isinstance(chunk, bytes):
                raise ValueError('Extraction requires a binary file.')
            chunks.append(chunk)
            remaining -= len(chunk)
        data = b''.join(chunks)
    except (AttributeError, OSError, TypeError) as exc:
        raise ValueError('The source file could not be read.') from exc
    finally:
        if previous is not None:
            file_obj.seek(previous)
    if len(data) > MAX_FILE_BYTES:
        raise ValueError(f'The source file exceeds the {MAX_FILE_BYTES // (1024 * 1024)} MB extraction limit.')
    if not data:
        raise ValueError('The source file is empty.')
    return data


def _validate_office_archive(data, extension, output):
    """Validate actual inflated sizes and XML before an Office parser sees them."""
    required = 'xl/workbook.xml' if extension == '.xlsx' else 'word/document.xml'

    def reject_xml(*_args):
        raise ValueError('Office XML with document types or entities is not supported.')

    def inspect_element(name, attrs):
        local_name = name.rsplit('}', 1)[-1].rsplit(':', 1)[-1]
        content_type = attrs.get('ContentType', '').lower()
        if 'macroenabled' in content_type or 'vbaproject' in content_type:
            raise ValueError('Macro-enabled Office content is not supported.')
        if local_name == 'Relationship' and attrs.get('TargetMode', '').lower() == 'external':
            if attrs.get('Type', '').rstrip('/').endswith('/hyperlink'):
                output.warn('External hyperlinks were not followed.')
            else:
                raise ValueError('Office documents with external data or attached templates are not supported.')

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                raise ValueError('The Office archive contains too many entries.')
            names = [entry.filename for entry in entries]
            if len(set(names)) != len(names):
                raise ValueError('The Office archive contains duplicate entries.')
            if required not in names or '[Content_Types].xml' not in names:
                raise ValueError(f'The file is not a valid {extension[1:].upper()} document.')
            if sum(entry.file_size for entry in entries) > MAX_ZIP_UNCOMPRESSED_BYTES:
                raise ValueError('The Office archive exceeds the uncompressed size limit.')
            total = 0
            for entry in entries:
                name = entry.filename.lower().replace('\\', '/')
                if name.startswith('/') or '..' in name.split('/'):
                    raise ValueError('The Office archive contains an unsafe entry path.')
                if entry.flag_bits & 1:
                    raise ValueError('Encrypted Office archives are not supported.')
                if entry.compress_type not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED):
                    raise ValueError('The Office archive uses an unsupported compression method.')
                if entry.file_size > MAX_ZIP_MEMBER_BYTES:
                    raise ValueError('An Office archive entry exceeds the uncompressed size limit.')
                if any(marker in name for marker in ('vbaproject', '/activex/', '/embeddings/', '/externallinks/')):
                    raise ValueError('Office macros, active objects, embedded files, and external data links are not supported.')
                if entry.is_dir():
                    continue
                parser = None
                if name.endswith(('.xml', '.rels')):
                    parser = expat.ParserCreate(namespace_separator='}')
                    parser.StartDoctypeDeclHandler = reject_xml
                    parser.EntityDeclHandler = reject_xml
                    parser.ExternalEntityRefHandler = reject_xml
                    parser.StartElementHandler = inspect_element
                member_total = 0
                with archive.open(entry) as stream:
                    while True:
                        block = stream.read(64 * 1024)
                        if not block:
                            break
                        member_total += len(block)
                        total += len(block)
                        if member_total > MAX_ZIP_MEMBER_BYTES or total > MAX_ZIP_UNCOMPRESSED_BYTES:
                            raise ValueError('The Office archive exceeds the uncompressed size limit.')
                        if parser:
                            parser.Parse(block, False)
                    if parser:
                        parser.Parse(b'', True)
    except (zipfile.BadZipFile, RuntimeError, expat.ExpatError, EOFError, zlib.error) as exc:
        raise ValueError('The Office archive is corrupt or contains invalid XML.') from exc


def _table_row(output, values, locations, row_location, headers, *, allow_table_suggestions=True):
    """Extract explicit pairs or a recognized table without guessing its meaning."""
    values = [_as_text(value) for value in values]
    while values and not values[-1]:
        values.pop()
    if not any(values):
        return None
    header_count = sum(_known_label(value) for value in values if value)
    if allow_table_suggestions and header_count >= 2 and all(not value or len(value) <= 80 for value in values):
        output.section(row_location, ' | '.join(values))
        return list(values)
    if headers:
        evidence = ' | '.join(
            f'{headers[index] if index < len(headers) and headers[index] else locations[index]}: {value}'
            for index, value in enumerate(values) if value
        )
    else:
        evidence = ' | '.join(values)
    retained = output.section(row_location, evidence)
    # A truncated source row must never create a partial/inaccurate suggestion.
    if retained != evidence:
        return headers
    for index, value in enumerate(values):
        if not value:
            continue
        output.labeled_lines(value, locations[index])
        if headers and index < len(headers):
            output.suggest(headers[index], value, locations[index], evidence)
    nonempty = [(index, value) for index, value in enumerate(values) if value]
    if allow_table_suggestions and not headers and len(nonempty) == 2 and _known_label(nonempty[0][1]):
        index, value = nonempty[1]
        output.suggest(nonempty[0][1], value, locations[index], evidence)
    return headers


def _extract_xlsx(data, output):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        if len(workbook.worksheets) > MAX_SHEETS:
            output.warn(f'Only the first {MAX_SHEETS} worksheets were inspected.')
        for sheet in workbook.worksheets[:MAX_SHEETS]:
            if output.full:
                output.warn('Output limit reached; remaining source content was not extracted.')
                break
            if sheet.sheet_state != 'visible':
                output.warn('Hidden worksheets were omitted.')
                continue
            # Ignore potentially incorrect dimensions; iterate only a bounded area.
            if sheet.max_row and sheet.max_row > MAX_ROWS:
                output.warn(f'Only the first {MAX_ROWS} rows of each worksheet were inspected.')
            if sheet.max_column and sheet.max_column > MAX_COLUMNS:
                output.warn(f'Only the first {MAX_COLUMNS} columns of each worksheet were inspected.')
            sheet.reset_dimensions()
            headers = None
            for number, row in enumerate(sheet.iter_rows(max_row=MAX_ROWS + 1, max_col=MAX_COLUMNS + 1), 1):
                if number > MAX_ROWS:
                    if any(cell.value is not None for cell in row):
                        output.warn(f'Only the first {MAX_ROWS} rows of each worksheet were inspected.')
                    break
                if row[MAX_COLUMNS].value is not None:
                    output.warn(f'Only the first {MAX_COLUMNS} columns of each worksheet were inspected.')
                values = []
                for cell in row[:MAX_COLUMNS]:
                    if cell.data_type == 'f':
                        output.warn('Spreadsheet formulas were omitted; calculated or cached values were not used.')
                        values.append('')
                    elif cell.data_type == 'e':
                        output.warn('Spreadsheet error cells were omitted.')
                        values.append('')
                    else:
                        values.append(cell.value)
                locations = [f'Sheet "{sheet.title}", cell {get_column_letter(index + 1)}{number}'
                             for index in range(len(values))]
                headers = _table_row(output, values, locations, f'Sheet "{sheet.title}", row {number}', headers)
                if output.full:
                    output.warn('Output limit reached; remaining source content was not extracted.')
                    return
    finally:
        workbook.close()


def _extract_docx(data, output):
    from docx import Document
    from docx.table import Table, _Cell
    from docx.text.paragraph import Paragraph

    document = Document(io.BytesIO(data))
    output.warn('Word extraction includes body paragraphs and tables; headers, footers, text boxes, and images were not inspected.')
    paragraph_number = table_number = 0
    for element in document.element.body:
        if output.full:
            output.warn('Output limit reached; remaining source content was not extracted.')
            break
        if element.tag.endswith('}p'):
            paragraph_number += 1
            if paragraph_number > MAX_ROWS:
                output.warn(f'Only the first {MAX_ROWS} body paragraphs were inspected.')
                break
            text = Paragraph(element, document).text
            location = f'Paragraph {paragraph_number}'
            retained = output.section(location, text)
            if retained == text.strip():
                output.labeled_lines(retained, location)
        elif element.tag.endswith('}tbl'):
            table_number += 1
            if table_number > MAX_TABLES:
                output.warn(f'Only the first {MAX_TABLES} Word tables were inspected.')
                break
            headers = None
            table = Table(element, document)
            merged = any(node.tag.endswith(('}gridSpan', '}vMerge', '}gridBefore', '}gridAfter'))
                         for node in table._tbl.iter())
            if merged:
                output.warn('Merged or offset Word tables were retained as text; table-column suggestions were omitted.')
            for row_number, row in enumerate(table.rows, 1):
                if row_number > MAX_ROWS:
                    output.warn(f'Only the first {MAX_ROWS} rows of each Word table were inspected.')
                    break
                # row.cells expands the entire table on each call (quadratic).
                # Read the physical cells directly and avoid assuming alignment
                # when a table contains merged/offset cells.
                cells = row._tr.tc_lst
                if len(cells) > MAX_COLUMNS:
                    output.warn(f'Only the first {MAX_COLUMNS} cells of each Word table row were inspected.')
                values = [_Cell(cell, table).text for cell in cells[:MAX_COLUMNS]]
                locations = [f'Table {table_number}, row {row_number}, cell {index + 1}'
                             for index in range(len(values))]
                headers = _table_row(output, values, locations, f'Table {table_number}, row {row_number}',
                                     headers, allow_table_suggestions=not merged)
                if output.full:
                    output.warn('Output limit reached; remaining source content was not extracted.')
                    return


def _validate_pdf(reader):
    """Reject active content and preflight ordinary non-image decoded streams.

    PyPDF2 remains an in-process parser: these limits are not a CPU/memory sandbox
    for every possible malformed PDF object graph.
    """
    from PyPDF2.generic import ArrayObject, DictionaryObject, IndirectObject, StreamObject

    pending = [reader.trailer]
    seen = set()
    total_decoded = 0
    forbidden = {'/JS', '/JavaScript', '/Launch', '/EmbeddedFiles', '/EF', '/XFA', '/RichMedia', '/AA', '/OpenAction'}
    actions = {'/JavaScript', '/Launch', '/SubmitForm', '/ImportData', '/GoToR', '/GoToE'}
    while pending:
        item = pending.pop()
        marker = ('ref', item.idnum, item.generation) if isinstance(item, IndirectObject) else ('obj', id(item))
        if marker in seen:
            continue
        seen.add(marker)
        if len(seen) > MAX_PDF_OBJECTS:
            raise ValueError('The PDF exceeds the object complexity limit.')
        if isinstance(item, IndirectObject):
            pending.append(item.get_object())
        elif isinstance(item, DictionaryObject):
            if forbidden.intersection(item.keys()) or str(item.get('/S', '')) in actions or str(item.get('/Type', '')) == '/EmbeddedFile':
                raise ValueError('PDF actions, scripts, forms with active content, and embedded files are not supported.')
            if isinstance(item, StreamObject) and str(item.get('/Subtype', '')) != '/Image':
                decoded = item._data
                filters = item.get('/Filter', [])
                if not isinstance(filters, (list, ArrayObject)):
                    filters = [filters]
                for filter_name in filters:
                    if str(filter_name) in ('/FlateDecode', '/Fl'):
                        decompressor = zlib.decompressobj()
                        decoded = decompressor.decompress(decoded, MAX_PDF_STREAM_BYTES + 1)
                        if len(decoded) > MAX_PDF_STREAM_BYTES or decompressor.unconsumed_tail:
                            raise ValueError('A PDF stream exceeds the decompressed size limit.')
                        decoded += decompressor.flush(MAX_PDF_STREAM_BYTES - len(decoded) + 1)
                        if not decompressor.eof:
                            raise ValueError('The PDF contains an incomplete compressed stream.')
                    elif str(filter_name) in ('/ASCII85Decode', '/A85'):
                        encoded = decoded.strip()
                        if encoded.startswith(b'<~'):
                            decoded = base64.a85decode(encoded, adobe=True)
                        else:
                            decoded = base64.a85decode(encoded.removesuffix(b'~>'))
                    elif str(filter_name) in ('/ASCIIHexDecode', '/AHx'):
                        encoded = re.sub(rb'\s', b'', decoded).removesuffix(b'>')
                        if len(encoded) % 2:
                            encoded += b'0'
                        decoded = bytes.fromhex(encoded.decode('ascii'))
                    else:
                        raise ValueError('The PDF uses an unsupported text-stream compression method.')
                    if len(decoded) > MAX_PDF_STREAM_BYTES:
                        raise ValueError('A PDF stream exceeds the decompressed size limit.')
                if len(decoded) > MAX_PDF_STREAM_BYTES:
                    raise ValueError('A PDF stream exceeds the decompressed size limit.')
                total_decoded += len(decoded)
                if total_decoded > MAX_PDF_TOTAL_STREAM_BYTES:
                    raise ValueError('The PDF exceeds the total decompressed stream size limit.')
            pending.extend(value for value in item.values()
                           if isinstance(value, (IndirectObject, DictionaryObject, ArrayObject)))
        elif isinstance(item, ArrayObject):
            pending.extend(item)


def _extract_pdf(data, output):
    from PyPDF2 import PdfReader

    reader = PdfReader(io.BytesIO(data), strict=True)
    if reader.is_encrypted:
        raise ValueError('Encrypted PDFs are not supported.')
    _validate_pdf(reader)
    page_count = len(reader.pages)
    if page_count > MAX_PDF_PAGES:
        output.warn(f'Only the first {MAX_PDF_PAGES} PDF pages were inspected.')
    for index in range(min(page_count, MAX_PDF_PAGES)):
        text = reader.pages[index].extract_text() or ''
        location = f'Page {index + 1}'
        if not text.strip():
            output.warn(f'{location} has no extractable text; OCR was not performed.')
            continue
        retained = output.section(location, text)
        if retained == text.strip():
            output.labeled_lines(retained, location)
        if output.full:
            output.warn('Output limit reached; remaining source content was not extracted.')
            break
    if not output.sections:
        raise ValueError('No extractable text was found in the inspected PDF pages. Scanned or image-only files require OCR.')


def _extract_text(data, extension, output):
    try:
        text = data.decode('utf-8-sig')
    except UnicodeDecodeError as exc:
        raise ValueError('Text and CSV files must use UTF-8 encoding.') from exc
    if '\x00' in text:
        raise ValueError('The file contains binary content instead of UTF-8 text.')
    if extension == '.txt':
        for number, line in enumerate(io.StringIO(text), 1):
            if number > MAX_ROWS:
                output.warn(f'Only the first {MAX_ROWS} text lines were inspected.')
                break
            location = f'Line {number}'
            retained = output.section(location, line)
            if retained == line.strip():
                output.labeled_lines(retained, location)
            if output.full:
                output.warn('Output limit reached; remaining source content was not extracted.')
                break
    else:
        headers = None
        for number, row in enumerate(csv.reader(io.StringIO(text, newline=''), strict=True), 1):
            if number > MAX_ROWS:
                output.warn(f'Only the first {MAX_ROWS} CSV rows were inspected.')
                break
            if len(row) > MAX_COLUMNS:
                output.warn(f'Only the first {MAX_COLUMNS} columns of each CSV row were inspected.')
            row = row[:MAX_COLUMNS]
            locations = [f'CSV row {number}, column {index + 1}' for index in range(len(row))]
            headers = _table_row(output, row, locations, f'CSV row {number}', headers)
            if output.full:
                output.warn('Output limit reached; remaining source content was not extracted.')
                break


def extract_content(file_obj, filename):
    """Return JSON-safe sections, evidence-backed suggestions, and limitations.

    Oversized/unsafe/unreadable files raise ValueError. Bounded partial results
    carry explicit warnings. The caller must retain the source file/version and
    require review before using any suggestion as project information.
    """
    extension = PureWindowsPath(str(filename)).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError('Unsupported file type. Extraction supports PDF, XLSX, DOCX, TXT, and CSV.')
    data = _read_bounded(file_obj)
    output = _Output()
    try:
        if extension in ('.xlsx', '.docx'):
            _validate_office_archive(data, extension, output)
        if extension == '.xlsx':
            _extract_xlsx(data, output)
        elif extension == '.docx':
            _extract_docx(data, output)
        elif extension == '.pdf':
            _extract_pdf(data, output)
        else:
            _extract_text(data, extension, output)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f'The {extension[1:].upper()} file could not be extracted; it may be corrupt or unsupported.') from exc
    if not output.sections:
        raise ValueError('No readable content was found in the inspected source content.')
    return output.result()
