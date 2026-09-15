"""Pure extraction tests; run without a database or Django settings."""

import io
import json
import unittest
import zipfile
from datetime import date
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook
from PyPDF2 import PdfWriter
from PyPDF2.generic import DictionaryObject, NameObject, DecodedStreamObject

from apps.file_replica import extraction


def make_xlsx(rows, *, extra_sheet=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Deliverables'
    for row in rows:
        sheet.append(row)
    if extra_sheet:
        workbook.create_sheet('Other').append(['Project Code', 'UNINSPECTED'])
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def make_docx():
    document = Document()
    document.add_paragraph('Project Code: 5900738')
    document.add_paragraph('Project Name: Grid Power Integration Project')
    table = document.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ['Milestone', 'Due Date', 'Status']):
        cell.text = text
    for cell, text in zip(table.add_row().cells, ['Design review', '2026-12-01', 'Pending']):
        cell.text = text
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def make_pdf(texts, *, compressed=False, active=False):
    writer = PdfWriter()
    for text in texts:
        writer.add_blank_page(width=400, height=400)
        page = writer.pages[-1]
        if text:
            font = DictionaryObject({
                NameObject('/Type'): NameObject('/Font'),
                NameObject('/Subtype'): NameObject('/Type1'),
                NameObject('/BaseFont'): NameObject('/Helvetica'),
            })
            page[NameObject('/Resources')] = DictionaryObject({
                NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)}),
            })
            stream = DecodedStreamObject()
            safe = text.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)')
            stream.set_data(f'BT /F1 12 Tf 30 300 Td ({safe}) Tj ET'.encode('ascii'))
            if compressed:
                stream = stream.flate_encode()
            page[NameObject('/Contents')] = writer._add_object(stream)
    if active:
        writer.add_js('app.alert("ignored");')
    result = io.BytesIO()
    writer.write(result)
    return result.getvalue()


def replace_zip_entry(data, name, content):
    result = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(result, 'w', zipfile.ZIP_DEFLATED) as destination:
        for info in source.infolist():
            if info.filename != name:
                destination.writestr(info.filename, source.read(info.filename))
        destination.writestr(name, content)
    return result.getvalue()


class ExtractionTests(unittest.TestCase):
    def extract(self, data, filename):
        result = extraction.extract_content(io.BytesIO(data), filename)
        json.dumps(result)  # Every supported result is safe for a JSONField/API.
        return result

    def test_explicit_text_labels_preserve_values_and_locations(self):
        data = b'Project Code: 5900738\nProject Name: Grid Power\nProgress: 42%\nSome prose mentions 85% progress.\nDue Date: TBD\n'
        result = self.extract(data, 'notes.TXT')
        suggestions = {item['label']: item for item in result['suggestions']}
        self.assertEqual(suggestions['Project Code']['value'], '5900738')
        self.assertEqual(suggestions['Project Code']['location'], 'Line 1')
        self.assertEqual(suggestions['Progress']['value'], '42%')
        self.assertEqual(suggestions['Due Date']['value'], 'TBD')
        self.assertEqual(len(suggestions), 4)
        self.assertEqual(suggestions['Progress']['evidence'], 'Progress: 42%')

    def test_csv_table_keeps_row_context_and_cell_locations(self):
        data = b'WBS,Deliverable,Quantity,Status\n1.1,"Pumps, duty",2,Approved\n'
        result = self.extract(data, 'register.csv')
        quantity = next(item for item in result['suggestions'] if item['label'] == 'Quantity')
        self.assertEqual(quantity['value'], '2')
        self.assertEqual(quantity['location'], 'CSV row 2, column 3')
        self.assertIn('WBS: 1.1', quantity['evidence'])
        self.assertIn('Deliverable: Pumps, duty', quantity['evidence'])

    def test_xlsx_explicit_pairs_tables_and_dates(self):
        data = make_xlsx([
            ['Project Code', '5900738'], [],
            ['Milestone', 'Due Date', 'Status'],
            ['Design review', date(2026, 12, 1), 'Pending'],
        ])
        result = self.extract(data, 'register.xlsx')
        code = next(item for item in result['suggestions'] if item['label'] == 'Project Code')
        due = next(item for item in result['suggestions'] if item['label'] == 'Due Date')
        self.assertEqual(code['location'], 'Sheet "Deliverables", cell B1')
        self.assertEqual(due['value'], '2026-12-01T00:00:00')
        self.assertIn('Milestone: Design review', due['evidence'])
        self.assertEqual(due['location'], 'Sheet "Deliverables", cell B4')

    def test_xlsx_omits_formulas_instead_of_guessing_completion(self):
        result = self.extract(make_xlsx([
            ['Deliverable', 'Progress'], ['Design', '=1/2'],
        ]), 'register.xlsx')
        self.assertFalse(any(item['label'] == 'Progress' for item in result['suggestions']))
        self.assertTrue(any('formulas were omitted' in warning for warning in result['warnings']))
        self.assertNotIn('=1/2', json.dumps(result))

    def test_docx_body_and_tables_have_source_locations(self):
        result = self.extract(make_docx(), 'project.docx')
        code = next(item for item in result['suggestions'] if item['label'] == 'Project Code')
        status = next(item for item in result['suggestions'] if item['label'] == 'Status')
        self.assertEqual(code['location'], 'Paragraph 1')
        self.assertEqual(status['location'], 'Table 1, row 2, cell 3')
        self.assertIn('Milestone: Design review', status['evidence'])
        self.assertTrue(any('headers, footers' in warning for warning in result['warnings']))

    def test_merged_word_tables_do_not_guess_column_alignment(self):
        document = Document(io.BytesIO(make_docx()))
        document.tables[0].rows[1].cells[0].merge(document.tables[0].rows[1].cells[1])
        stream = io.BytesIO()
        document.save(stream)
        result = self.extract(stream.getvalue(), 'merged.docx')
        self.assertFalse(any(item['label'] == 'Status' for item in result['suggestions']))
        self.assertIn('Design review', json.dumps(result['sections']))
        self.assertTrue(any('Merged or offset' in item for item in result['warnings']))

    def test_encrypted_pdf_has_a_readable_failure(self):
        writer = PdfWriter()
        writer.add_blank_page(width=400, height=400)
        writer.encrypt('private')
        stream = io.BytesIO()
        writer.write(stream)
        with self.assertRaisesRegex(ValueError, 'Encrypted PDFs'):
            self.extract(stream.getvalue(), 'encrypted.pdf')

    def test_pdf_text_and_compressed_streams(self):
        for compressed in (False, True):
            with self.subTest(compressed=compressed):
                result = self.extract(make_pdf(['Project Code: 5900738'], compressed=compressed), 'project.pdf')
                self.assertEqual(result['suggestions'][0]['value'], '5900738')
                self.assertEqual(result['suggestions'][0]['location'], 'Page 1')

    def test_empty_pdf_fails_instead_of_claiming_ocr_success(self):
        with self.assertRaisesRegex(ValueError, 'require OCR'):
            self.extract(make_pdf(['']), 'scan.pdf')

    def test_mixed_pdf_warns_about_unreadable_pages(self):
        result = self.extract(make_pdf(['', 'Status: Issued']), 'mixed.pdf')
        self.assertEqual(result['suggestions'][0]['location'], 'Page 2')
        self.assertTrue(any('Page 1 has no extractable text' in item for item in result['warnings']))

    def test_pdf_page_limit_is_visible(self):
        with patch.object(extraction, 'MAX_PDF_PAGES', 1):
            result = self.extract(make_pdf(['Status: Draft', 'Status: Approved']), 'long.pdf')
        self.assertEqual([item['value'] for item in result['suggestions']], ['Draft'])
        self.assertIn('Only the first 1 PDF pages were inspected.', result['warnings'])

    def test_pdf_decompression_limit(self):
        data = make_pdf(['Project Name: ' + 'A' * 1000], compressed=True)
        with patch.object(extraction, 'MAX_PDF_STREAM_BYTES', 100):
            with self.assertRaisesRegex(ValueError, 'decompressed size limit'):
                self.extract(data, 'large.pdf')

    def test_pdf_active_content_is_rejected(self):
        with self.assertRaisesRegex(ValueError, 'scripts'):
            self.extract(make_pdf(['Status: Draft'], active=True), 'active.pdf')

    def test_row_sheet_and_column_limits_are_visible(self):
        data = make_xlsx([
            ['Deliverable', 'Status', 'Quantity'],
            ['Design', 'Draft', 4], ['Construction', 'Approved', 10],
        ], extra_sheet=True)
        with patch.object(extraction, 'MAX_ROWS', 2), patch.object(extraction, 'MAX_SHEETS', 1), patch.object(extraction, 'MAX_COLUMNS', 2):
            result = self.extract(data, 'large.xlsx')
        self.assertEqual({item['value'] for item in result['suggestions']}, {'Design', 'Draft'})
        for word in ('rows', 'worksheets', 'columns'):
            self.assertTrue(any(word in warning for warning in result['warnings']))

    def test_partial_text_never_becomes_a_partial_suggestion(self):
        with patch.object(extraction, 'MAX_SECTION_CHARS', 20):
            result = self.extract(b'Project Name: A very long project description', 'notes.txt')
        self.assertEqual(len(result['sections'][0]['text']), 20)
        self.assertEqual(result['suggestions'], [])
        self.assertTrue(any('truncated' in item for item in result['warnings']))

    def test_output_and_suggestion_limits(self):
        with patch.object(extraction, 'MAX_TOTAL_CHARS', 30), patch.object(extraction, 'MAX_SUGGESTIONS', 1):
            result = self.extract(b'Status: Draft\nQuantity: 4\nProject Code: 5900738\n', 'notes.txt')
        self.assertLessEqual(sum(len(item['text']) for item in result['sections']), 30)
        self.assertEqual(len(result['suggestions']), 1)
        self.assertTrue(any('Suggestion limit' in item for item in result['warnings']))

    def test_size_limit_and_stream_position(self):
        stream = io.BytesIO(b'Status: Draft')
        stream.seek(3)
        result = extraction.extract_content(stream, 'notes.txt')
        self.assertEqual(stream.tell(), 3)
        self.assertEqual(result['suggestions'][0]['value'], 'Draft')
        with patch.object(extraction, 'MAX_FILE_BYTES', 3):
            with self.assertRaisesRegex(ValueError, 'extraction limit'):
                self.extract(b'abcd', 'notes.txt')

    def test_office_uncompressed_limit(self):
        with patch.object(extraction, 'MAX_ZIP_UNCOMPRESSED_BYTES', 100):
            with self.assertRaisesRegex(ValueError, 'uncompressed size limit'):
                self.extract(make_xlsx([['Project Code', '5900738']]), 'large.xlsx')

    def test_disguised_macro_archive_is_rejected(self):
        data = replace_zip_entry(make_xlsx([['Status', 'Draft']]), 'xl/vbaProject.bin', b'macro')
        with self.assertRaisesRegex(ValueError, 'macros'):
            self.extract(data, 'disguised.xlsx')

    def test_office_xml_entities_are_rejected(self):
        malicious = b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY x "boom">]><x>&x;</x>'
        data = replace_zip_entry(make_docx(), 'customXml/item99.xml', malicious)
        with self.assertRaisesRegex(ValueError, 'document types or entities'):
            self.extract(data, 'entities.docx')

    def test_external_office_data_is_rejected(self):
        relationship = b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="r1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/attachedTemplate" Target="https://example.invalid/template" TargetMode="External"/></Relationships>'
        data = replace_zip_entry(make_docx(), 'word/_rels/settings.xml.rels', relationship)
        with self.assertRaisesRegex(ValueError, 'external data or attached templates'):
            self.extract(data, 'external.docx')

    def test_corrupt_unsupported_empty_and_non_utf8_fail_readably(self):
        cases = [
            (b'bad pdf', 'bad.pdf', 'could not be extracted'),
            (b'bad zip', 'bad.xlsx', 'corrupt'),
            (b'bad zip', 'bad.docx', 'corrupt'),
            (b'anything', 'macro.xlsm', 'Unsupported file type'),
            (b'', 'empty.txt', 'empty'),
            (b'\xff\xfe', 'legacy.txt', 'UTF-8'),
            (b'abc\x00', 'binary.txt', 'binary content'),
            (b'"unterminated', 'broken.csv', 'could not be extracted'),
        ]
        for data, filename, message in cases:
            with self.subTest(filename=filename):
                with self.assertRaisesRegex(ValueError, message):
                    self.extract(data, filename)


if __name__ == '__main__':
    unittest.main()
