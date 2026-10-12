from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from django.test import SimpleTestCase
from openpyxl import load_workbook

from .excel_export import export_project_to_xlsx
from .models import ValveMTOProject, ValveMTORow
from .views import ValveMTOProjectViewSet, ValveMTOLegendViewSet


class ValveMTODataTests(SimpleTestCase):
    def test_size_nb_is_derived_only_from_explicit_quoted_line_prefix(self):
        from .specification_lookup import size_nb_from_line_number, enrich_specifications
        for line, expected in [
            ('4"-PL-DC3N-8198-P', '4"'),
            ('1/2"-CD-AC3N-8052', '1/2"'),
            ('1-1/2"-PL-DC3N-8198-P', '1-1/2"'),
            ('1 1/2\u2033\u2013PL\u2013DC3N\u20138198-P', '1 1/2"'),
            ('604-LFG-3-AC2GA0-2012', ''),
            ('DN100-PL-DC3N-8198-P', ''),
        ]:
            self.assertEqual(size_nb_from_line_number(line), expected)
        row = enrich_specifications({'line_number': '4"-PL-DC3N-8198-P', 'size_1': '2"'}, {})
        self.assertEqual(row['size_1'], '2"')

    def test_server_export_fills_blank_size_from_line_number(self):
        row = ValveMTORow(line_number='4"-PL-DC3N-8198-P', size_primary='', area='Field')
        rows = MagicMock()
        rows.all.return_value = [row]
        wb = load_workbook(BytesIO(export_project_to_xlsx(SimpleNamespace(rows=rows))))
        exported = dict(zip((c.value for c in wb['FIELD'][1]), (c.value for c in wb['FIELD'][2])))
        self.assertEqual(exported['SIZE (NB)'], '4"')

    def test_explicit_specifications_fill_only_exact_class_and_source_area_mappings(self):
        from .specification_lookup import load_specification_lookups, enrich_specifications
        legend = SimpleNamespace(definition={'fields': [
            {'key': 'spec', 'lookup': {'AC3N': {'pms_class': 'CARBON STEEL', 'rating': '150', 'facing': 'RF'}}},
            {'key': 'area', 'lookup': {'2"-CD-AC3N-8052': 'FIELD'}},
        ]})
        with patch('apps.valve_mto.models.ValveMTOLegend.objects.filter', return_value=[legend]) as query:
            lookups = load_specification_lookups(10)
        self.assertEqual(query.call_args.kwargs['created_by_id'], 10)
        row = enrich_specifications({'line_number': '2"-CD-AC3N-8052'}, lookups)
        self.assertEqual((row['pms_class'], row['rating'], row['facing'], row['tab']),
                         ('CARBON STEEL', '150', 'RF', 'field'))
        self.assertEqual(row['size_1'], '2"')
        unsupported = enrich_specifications({'line_number': '2"-CD-ACSN-8052'}, lookups)
        self.assertNotIn('pms_class', unsupported)
        self.assertNotIn('area', unsupported)

    def test_installation_lookup_never_replaces_the_drawing_title(self):
        from .specification_lookup import enrich_specifications
        title = 'PIPING AND INSTRUMENTATION DIAGRAM AA PRODUCTION PLATFORM MPFM'
        row = enrich_specifications(
            {'area': title, 'pid_number': 'PJ6-EXD-AAP-BQDA-0001'},
            {'area': {'PJ6-EXD-AAP-BQDA-0001': 'FIELD'}},
        )
        self.assertEqual(row['area'], title)
        self.assertEqual(row['tab'], 'field')

    def test_long_drawing_title_round_trips_and_exports_using_independent_tab(self):
        from .serializers import ValveMTORowSerializer
        title = 'PIPING AND INSTRUMENTATION DIAGRAM AA PRODUCTION PLATFORM MPFM'
        data = ValveMTORowSerializer(data={'area': title, 'tab': 'island'}, partial=True)
        self.assertTrue(data.is_valid(), data.errors)
        row = ValveMTORow(**data.validated_data)
        rows = MagicMock()
        rows.all.return_value = [row]
        wb = load_workbook(BytesIO(export_project_to_xlsx(SimpleNamespace(rows=rows))))
        exported = dict(zip((c.value for c in wb['ISLAND'][1]), (c.value for c in wb['ISLAND'][2])))
        self.assertEqual(exported['Area'], title)
        self.assertEqual(wb['FIELD'].max_row, 1)

    def test_conflicting_active_project_lookup_is_explicit_not_last_writer_wins(self):
        from .specification_lookup import load_specification_lookups
        legends = [SimpleNamespace(definition={'fields': [{'key': 'spec', 'lookup': {'AC3N': value}}]})
                   for value in ('CARBON STEEL', 'STAINLESS STEEL')]
        with patch('apps.valve_mto.models.ValveMTOLegend.objects.filter', return_value=legends):
            with self.assertRaisesRegex(ValueError, 'Conflicting active legend'):
                load_specification_lookups(10)

    def test_conflicting_row_specification_is_preserved_and_flagged(self):
        from .specification_lookup import enrich_specifications
        row = enrich_specifications({'piping_class': 'AC3N', 'rating': '300'}, {'rating': {'AC3N': '150'}})
        self.assertEqual(row['rating'], '300')
        self.assertEqual(row['confidence'], 'LOW')
        self.assertIn('CONFLICTS WITH ACTIVE PMS', row['remarks'])

    def test_workbook_actions_have_explicit_module_permissions(self):
        from apps.rbac.action_policy import operation_action
        for action_name, expected in [('workbook_preview', 'read'), ('workbook_import', 'update')]:
            options = getattr(ValveMTOLegendViewSet, action_name).kwargs
            view = ValveMTOLegendViewSet(permission_action=options['permission_action'])
            self.assertEqual(operation_action(SimpleNamespace(method='POST'), view), expected)

    def test_bulk_save_and_serializer_preserve_extraction_columns(self):
        project = ValveMTOProject(project_name='Regression fixture')
        created = []
        queryset = MagicMock()
        queryset.__iter__.side_effect = lambda: iter(created)
        fields = {
            'tag_number': 'HV-101', 'valve_type': 'GATE', 'pms_class': 'CARBON STEEL',
            'piping_class': 'AC2GA0', 'pid_number': 'PJ6-EXD-AAP-BQDA-0001',
            'line_number': '604-LFG-3-AC2GA0-2012', 'rating': '150', 'facing': 'RF',
            'unit': 'NOS', 'confidence': 'HIGH', 'location_hint': 'D-7',
            'area': 'Field', 'qty_field': 1, 'qty_island': 0, 'qty_observed': 1,
        }
        view = ValveMTOProjectViewSet()
        locked = MagicMock()
        locked.get.return_value = project
        with patch.object(view, 'get_object', return_value=project), \
                patch.object(ValveMTOProject.objects, 'select_for_update', return_value=locked) as lock, \
                patch.object(type(project.rows), 'all', return_value=queryset), \
                patch.object(ValveMTORow.objects, 'bulk_create', side_effect=lambda rows: created.extend(rows)), \
                patch.object(project, 'save'), \
                patch('apps.valve_mto.views.transaction.atomic'):
            response = view.save_rows(SimpleNamespace(data={'rows': [fields]}))
        self.assertEqual(response.status_code, 200)
        lock.assert_called_once_with()
        locked.get.assert_called_once_with(pk=project.pk)
        self.assertEqual(response.data['row_count'], 1)
        for key, value in fields.items():
            self.assertEqual(response.data['rows'][0][key], value, key)

    def test_server_export_keeps_drawing_line_and_material_fields_distinct(self):
        row = ValveMTORow(
            piping_class='AC2GA0', pms_class='CARBON STEEL',
            pid_number='PJ6-EXD-AAP-BQDA-0001', line_number='604-LFG-3-AC2GA0-2012',
            facing='RF', confidence='HIGH', location_hint='D-7', area='Field', unit='NOS',
            size_primary='2"', size_secondary_2='1"', line_list_ref='Legacy line list',
        )
        rows = MagicMock()
        rows.all.return_value = [row]
        workbook = load_workbook(BytesIO(export_project_to_xlsx(SimpleNamespace(rows=rows))))
        headers = [cell.value for cell in workbook['FIELD'][1]]
        values = [cell.value for cell in workbook['FIELD'][2]]
        exported = dict(zip(headers, values))
        self.assertEqual(exported['P&ID Number'], row.pid_number)
        self.assertEqual(exported['Line Number'], row.line_number)
        self.assertEqual(exported['Piping Class'], 'AC2GA0')
        self.assertEqual(exported['PMS Class'], 'CARBON STEEL')
        self.assertEqual(exported['Facing'], 'RF')
        self.assertEqual(exported['Confidence'], 'HIGH')
        self.assertEqual(exported['Location'], 'D-7')
        self.assertEqual(exported['Unit'], 'NOS')
        self.assertEqual(exported['SIZE (NB)'], '2"')
        for sheet in workbook:
            labels = [cell.value for cell in sheet[1]]
            self.assertIn('SIZE (NB)', labels)
            for removed in ('Size 1', 'Size 2', 'Line List Ref'):
                self.assertNotIn(removed, labels)
        self.assertEqual(row.size_secondary_2, '1"')
        self.assertEqual(row.line_list_ref, 'Legacy line list')

    def test_invalid_observed_quantities_are_rejected_before_replacing_existing_rows(self):
        project = ValveMTOProject(project_name='Existing review data')
        view = ValveMTOProjectViewSet()
        for value in (-1, 1.5, 'invalid'):
            with self.subTest(value=value), \
                    patch.object(view, 'get_object', return_value=project), \
                    patch.object(type(project.rows), 'all') as existing:
                response = view.save_rows(SimpleNamespace(data={'rows': [{'qty_observed': value}]}))
            self.assertEqual(response.status_code, 400)
            self.assertIn('observed quantity', response.data['error'])
            existing.assert_not_called()

    def test_unassigned_local_matches_are_exported_with_observed_count(self):
        row = ValveMTORow(valve_type='BALL VALVE', qty_observed=1, confidence='LOW', location_hint='Page 1')
        rows = MagicMock()
        rows.all.return_value = [row]
        workbook = load_workbook(BytesIO(export_project_to_xlsx(SimpleNamespace(rows=rows))))
        sheet = workbook['UNASSIGNED REVIEW']
        exported = dict(zip((cell.value for cell in sheet[1]), (cell.value for cell in sheet[2])))
        self.assertEqual(exported['Observed Qty (Review)'], 1)
        self.assertEqual(exported['Type'], 'BALL VALVE')
