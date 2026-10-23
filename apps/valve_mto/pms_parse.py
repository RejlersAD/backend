"""
Valve MTO — PMS workbook parsing (Stage 2).

Accepts two shapes (§4 Stage 2):
  1. The standard PMS template — header row with columns
     PIPING CLASS*, RATING*, FACING*, PIPING MATERIAL*, SERVICE,
     CORROSION ALLOWANCE, VALVE BODY MATERIAL (optional).
  2. The ADNOC Offshore legend workbook's 'Piping Material Codes' sheet —
     header auto-map (PIPING CLASS / ASME B16.5 FLANGE RATING /
     PIPING MATERIAL / SERVICE), header not necessarily on row 1.

parse_pms_workbook(bytes, filename) →
    {classes: {CODE: {rating, facing, material, service, corrosion_allowance,
                      valve_body_material}}, source, sheet}
or {error: ...}.
"""
from __future__ import annotations

import io
import logging

import openpyxl

logger = logging.getLogger(__name__)

# Header auto-map: canonical field → accepted header aliases (lowercase).
HEADER_ALIASES = {
    'code': ['piping class', 'piping material code', 'class', 'class code', 'pipe class'],
    'rating': ['rating', 'flange rating', 'asme b16.5 flange rating', 'pressure rating'],
    'facing': ['facing', 'face', 'face type', 'flange facing'],
    'material': ['piping material', 'material', 'piping material class'],
    'service': ['service', 'services'],
    'corrosion_allowance': ['corrosion allowance', 'ca'],
    'valve_body_material': ['valve body material', 'body material'],
}
REQUIRED_FIELDS = ('code', 'rating', 'material')
TEMPLATE_REQUIRED = ('code', 'rating', 'facing', 'material')  # template marks FACING* required

MAX_SCAN_ROWS = 30  # header can start anywhere in the first N rows


def _norm(cell) -> str:
    return str(cell or '').strip().lower()


def _map_header(row) -> dict:
    """Map column index → canonical field for one candidate header row.
    Returns {} when fewer than the required fields are present."""
    mapping = {}
    for idx, cell in enumerate(row):
        text = _norm(cell).rstrip('*').strip()
        if not text:
            continue
        for field, aliases in HEADER_ALIASES.items():
            if field not in mapping.values() and text in aliases:
                mapping[idx] = field
                break
    return mapping if all(f in mapping.values() for f in REQUIRED_FIELDS) else {}


def _read_sheet(ws) -> dict | None:
    rows = list(ws.iter_rows(values_only=True))
    for start in range(min(MAX_SCAN_ROWS, len(rows))):
        mapping = _map_header(rows[start])
        if not mapping:
            continue
        classes = {}
        for row in rows[start + 1:]:
            code = ''
            entry = {}
            for idx, field in mapping.items():
                if idx >= len(row):
                    continue
                value = str(row[idx]).strip() if row[idx] is not None else ''
                if field == 'code':
                    code = value
                else:
                    entry[field] = value
            if code:
                classes[code] = entry
        if classes:
            return classes
    return None


def parse_pms_workbook(data: bytes, filename: str = '') -> dict:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        return {'error': f'Could not open workbook: {exc}'}

    # Preferred: the ADNOC-style dedicated sheet; else every sheet in order.
    sheet_order = []
    for name in wb.sheetnames:
        if 'material code' in name.lower() or 'piping class' in name.lower():
            sheet_order.insert(0, name)
        else:
            sheet_order.append(name)

    for name in sheet_order:
        classes = _read_sheet(wb[name])
        if classes:
            return {
                'classes': classes,
                'source': filename,
                'sheet': name,
                'class_count': len(classes),
            }
    return {'error': 'No Piping Class table found (expected the PMS template '
                     'or a "Piping Material Codes" sheet).'}


def coverage_check(pms_classes: dict, drawing_classes: list) -> dict:
    """§4 Stage 2 coverage: which Stage-1 classes the PMS covers."""
    covered = [c for c in drawing_classes if c in pms_classes]
    uncovered = [c for c in drawing_classes if c not in pms_classes]
    return {
        'covered': len(covered),
        'total': len(drawing_classes),
        'covered_classes': covered,
        'uncovered_classes': uncovered,
        'blocked': bool(drawing_classes) and not covered,
    }
