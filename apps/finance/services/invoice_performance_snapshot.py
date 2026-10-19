"""Reproducible daily performance aggregates from the audited Finance workbook."""
import hashlib
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .invoice_performance import aggregate_invoice_days
from .workbook_layout import WORKBOOK_MAX_COLUMN, validated_column_indices
from .workbook_summary_snapshot import HEADERS, _cell_currency, _cell_kind, _text


def _invoice_date(cell):
    value = cell.value
    if cell.data_type == 'e':
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for pattern in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%d.%m.%Y', '%d-%b-%Y'):
            try:
                return datetime.strptime(value.strip(), pattern).date()
            except ValueError:
                continue
    return None


def generate_invoice_performance_snapshot(path, *, sheet='External Invoice ', first_row=6,
                                          last_row=4409, header_row=5, snapshot_at=None):
    """Return an aggregate-only artifact; do not import, update or expose invoice rows."""
    path = Path(path)
    captured = snapshot_at or datetime.now(timezone.utc)
    if captured.tzinfo is None or captured.utcoffset() is None:
        raise ValueError('snapshot_at must include a timezone.')
    with path.open('rb') as source:
        digest = hashlib.file_digest(source, 'sha256').hexdigest()
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet]
        header = next(worksheet.iter_rows(min_row=header_row, max_row=header_row, max_col=WORKBOOK_MAX_COLUMN))
        indices = validated_column_indices(
            header, {**HEADERS, 'B': 'Invoice Date'}, header_row=header_row, casefold=True,
        )
        if last_row > worksheet.max_row:
            raise ValueError('The last invoice row is outside the worksheet.')
        counters = {'currency_conflict_count': 0, 'receipt_currency_mismatch_count': 0,
                    'invalid_invoice_date_count': 0}

        def source_rows():
            for row_number, physical in enumerate(
                    worksheet.iter_rows(min_row=first_row, max_col=WORKBOOK_MAX_COLUMN), first_row):
                row = tuple(physical[index] for index in indices)
                invoice_id = _text(row[0].value)
                footer = invoice_id.casefold().rstrip(':') in {'total', 'grand total', 'subtotal', 'sub total'}
                if row_number > last_row:
                    if invoice_id and not footer:
                        raise ValueError(f'An invoice row exists after the specified last row: {row_number}.')
                    continue
                if not invoice_id or footer or row[0].data_type == 'e':
                    raise ValueError(f'Row {row_number} is not an invoice row.')
                currency, currency_status = _cell_currency(row[11], row[30])
                receipt_currency, _ = _cell_currency(row[26], row[30])
                if currency_status != 'recorded':
                    counters['currency_conflict_count'] += 1
                invoiced = Decimal(str(row[11].value)) if _cell_kind(row[11]) == 'numeric_count' else None
                received = Decimal(str(row[26].value)) if _cell_kind(row[26]) == 'numeric_count' else None
                if received is not None and receipt_currency != currency:
                    counters['receipt_currency_mismatch_count'] += 1
                    received = None
                invoice_date = _invoice_date(row[1])
                if invoice_date is None and row[1].value is not None:
                    counters['invalid_invoice_date_count'] += 1
                status = _text(row[17].value).casefold().replace('_', ' ').strip()
                payment_status = 'cancelled' if status in {'cancelled', 'canceled'} else 'credit_note' if status in {'credit note', 'creditnote'} else status
                yield {'currency': currency or 'UNSPECIFIED', 'category': 'external',
                       'invoice_date': invoice_date, 'payment_status': payment_status,
                       'invoice_amount': invoiced, 'actual_payment_received': received}

        days = aggregate_invoice_days(source_rows())
    finally:
        workbook.close()
    return {'schema_version': '1.0',
            'source': {'file_name': path.name, 'sheet': sheet, 'first_row': first_row, 'last_row': last_row,
                       'snapshot_at': captured.isoformat(), 'sha256': digest, 'scope': 'full_external_workbook',
                       'invoice_date_column': 'B', 'amount_column': get_column_letter(indices[11] + 1),
                       'receipt_column': get_column_letter(indices[26] + 1),
                       'currency_column': get_column_letter(indices[30] + 1),
                       'currency_method': 'strict_agreement_or_single_source',
                       'tax_basis': 'recorded_invoice_amount'},
            'coverage': counters, 'days': days}
