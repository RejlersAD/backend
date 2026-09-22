"""Import cached source facts without recalculating or changing customer invoices."""
import hashlib
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from django.db import transaction
from openpyxl import load_workbook

from apps.finance.receivables_source_models import ReceivablesSourceRow, ReceivablesSourceSnapshot
from .workbook_summary import validate_snapshot, validate_source_summary
from .workbook_summary_snapshot import (
    FOOTER_LABELS, _cell_currency, _cell_kind, _text, generate_workbook_snapshot,
)


SOURCE_HEADERS = {'A': 'Invoice #', 'B': 'Invoice Date', 'C': 'Invoice Sent',
                  'E': 'COMPANY', 'F': 'COMPANY', 'G': 'RAD Project #', 'H': 'Project Name',
                  'L': 'Invoice Amount', 'M': 'Inv Amt. (AED)', 'N': 'Due Date',
                  'O': 'Payment terms', 'P': 'PM', 'R': 'Payment Status',
                  'Y': 'Balance to be received', 'Z': 'Payment Date',
                  'AA': 'Actual Payment Received', 'AC': 'Remarks', 'AD': 'Project ID', 'AE': 'Inv. CUR'}
MONEY_QUANTUM = Decimal('0.00000001')


def _money(cell):
    if _cell_kind(cell) != 'numeric_count':
        return None
    value = Decimal(str(cell.value))
    if abs(value) >= Decimal('1e20'):
        raise ValueError('A monetary source value exceeds the supported precision.')
    value = value.quantize(MONEY_QUANTUM, rounding=ROUND_HALF_UP)
    if abs(value) >= Decimal('1e20'):
        raise ValueError('A monetary source value exceeds the supported precision.')
    return value


def _date(cell):
    if cell.data_type == 'e':
        return None
    value = cell.value
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for pattern in ('%Y-%m-%d', '%d/%m/%Y', '%d/%m/%y', '%d-%m-%Y', '%d.%m.%Y',
                        '%d-%b-%Y', '%d.%b.%y', '%d.%B.%y'):
            try:
                return datetime.strptime(value.strip(), pattern).date()
            except ValueError:
                pass
    return None


def _status(value):
    status = ' '.join(_text(value).casefold().replace('_', ' ').split())
    # Match complete labels; "Paid (partial)" is not a paid invoice.
    aliases = {'paid (partial)': 'partial', 'partially paid': 'partial',
               'canceled': 'cancelled', 'credit note': 'credit_note', 'creditnote': 'credit_note'}
    return aliases.get(status, status) if len(status) <= 64 else 'unknown'


def _reconciliation(rows):
    overdue = [row for row in rows if row['payment_status'] == 'overdue']
    amounts = [row['invoice_amount_aed'] for row in overdue]
    known = [value for value in amounts if value is not None]
    ids = Counter(row['invoice_number'] for row in rows)
    return {
        'row_count': len(rows), 'payment_status_counts': dict(Counter(row['payment_status'] for row in rows)),
        'original_currency_counts': dict(Counter(row['currency'] or 'UNSPECIFIED' for row in rows)),
        'duplicate_invoice_number_count': sum(count > 1 for count in ids.values()),
        'overdue': {'count': len(overdue), 'amount_aed': str(sum(known, Decimal('0')).quantize(
            Decimal('0.01'), rounding=ROUND_HALF_UP)) if not overdue or known else None,
                    'missing_amount_count': len(amounts) - len(known)},
    }


def read_receivables_source(path, *, sheet='External Invoice ', first_row=6, last_row=4409,
                            header_row=5, original_filename=None):
    """Read a complete external sheet; None bounds detect its contiguous invoice rows."""
    path = Path(path)
    if sheet.strip().casefold() != 'external invoice':
        raise ValueError('Only the External Invoice source sheet is supported.')
    if header_row < 1 or first_row != header_row + 1 or (last_row is not None and last_row < first_row):
        raise ValueError('The invoice range must immediately follow the header row.')
    file_name = Path(str(original_filename or path.name).replace('\\', '/')).name
    if not file_name or len(file_name) > ReceivablesSourceSnapshot._meta.get_field('file_name').max_length:
        raise ValueError('The source filename is empty or too long.')
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = []
    try:
        candidates = [name for name in workbook.sheetnames if name.strip().casefold() == 'external invoice']
        if not candidates:
            raise ValueError('The requested source sheet was not found.')
        if len(candidates) != 1:
            raise ValueError('The workbook has more than one External Invoice source sheet.')
        sheet = candidates[0]
        worksheet = workbook[sheet]
        if last_row is not None and last_row > worksheet.max_row:
            raise ValueError('The last invoice row is outside the worksheet.')
        from openpyxl.utils import column_index_from_string
        header = next(worksheet.iter_rows(min_row=header_row, max_row=header_row, max_col=31))
        for column, expected in SOURCE_HEADERS.items():
            if _text(header[column_index_from_string(column) - 1].value).casefold() != expected.casefold():
                raise ValueError(f'Unexpected source header at {column}{header_row}.')
        detected_end = False
        for number, cells in enumerate(worksheet.iter_rows(min_row=first_row, max_col=31), first_row):
            invoice_number = _text(cells[0].value)
            footer = invoice_number.casefold().rstrip(':') in FOOTER_LABELS
            if last_row is None:
                if not invoice_number or footer:
                    detected_end = True
                    continue
                if detected_end:
                    raise ValueError(f'An invoice row exists after a blank or footer row: {number}.')
            elif number > last_row:
                if invoice_number and not footer:
                    raise ValueError(f'An invoice row exists after the specified last row: {number}.')
                continue
            if not invoice_number or footer or cells[0].data_type == 'e':
                raise ValueError(f'Row {number} is not an invoice row.')
            currency, currency_status = _cell_currency(cells[11], cells[30])
            balance_currency, balance_currency_status = _cell_currency(cells[24], cells[30])
            receipt_currency, receipt_currency_status = _cell_currency(cells[26], cells[30])
            row = {
                'row_number': number, 'invoice_number': invoice_number,
                'company': _text(cells[5].value), 'account': _text(cells[4].value),
                'pm': _text(cells[15].value), 'project_name': _text(cells[7].value),
                'rad_project_no': _text(cells[6].value), 'project_id': _text(cells[29].value),
                'payment_terms': _text(cells[14].value), 'remarks': _text(cells[28].value),
                'invoice_date': _date(cells[1]), 'invoice_sent_date': _date(cells[2]),
                'due_date': _date(cells[13]), 'payment_date': _date(cells[25]),
                'payment_status': _status(cells[17].value), 'raw_payment_status': _text(cells[17].value),
                'currency': currency or '', 'currency_status': currency_status,
                'balance_currency': balance_currency or '', 'balance_currency_status': balance_currency_status,
                'actual_payment_currency': receipt_currency or '',
                'actual_payment_currency_status': receipt_currency_status,
                'invoice_amount': _money(cells[11]), 'invoice_amount_aed': _money(cells[12]),
                'balance_to_be_received': _money(cells[24]), 'actual_payment_received': _money(cells[26]),
            }
            for name, value in row.items():
                field = ReceivablesSourceRow._meta.get_field(name)
                if isinstance(value, str) and field.max_length and len(value) > field.max_length:
                    raise ValueError(f'Row {number} exceeds the supported {name} length.')
            rows.append(row)
    finally:
        workbook.close()
    if hashlib.sha256(path.read_bytes()).hexdigest() != digest:
        raise ValueError('The source workbook changed while it was being read.')
    if not rows:
        raise ValueError('The source worksheet contains no invoice rows.')
    if last_row is None:
        last_row = rows[-1]['row_number']
    metadata = {'sha256': digest, 'file_name': file_name, 'sheet_name': sheet,
                'header_row': header_row, 'first_row': first_row, 'last_row': last_row,
                'row_count': len(rows)}
    return metadata, rows, _reconciliation(rows)


def import_receivables_source(path, *, sheet='External Invoice ', first_row=6, last_row=4409,
                               header_row=5, dry_run=False, original_filename=None):
    """Stage and activate one immutable source version; never save operational invoices."""
    metadata, rows, reconciliation = read_receivables_source(
        path, sheet=sheet, first_row=first_row, last_row=last_row, header_row=header_row,
        original_filename=original_filename)
    # Validate the complete aggregate before publishing either the rows or totals.
    # Both readers must have consumed the identical file, including cached cells.
    summary = generate_workbook_snapshot(
        path, sheet=metadata['sheet_name'], last_row=metadata['last_row'], header_row=header_row)
    summary['source']['file_name'] = metadata['file_name']
    validate_snapshot(summary)
    if (summary['source']['sha256'] != metadata['sha256']
            or hashlib.sha256(Path(path).read_bytes()).hexdigest() != metadata['sha256']
            or summary['invoice_count'] != metadata['row_count']):
        raise ValueError('The source workbook changed while its summary was being read.')
    reconciliation['workbook_summary'] = summary
    result = {**metadata, 'reconciliation': reconciliation, 'dry_run': dry_run,
              'snapshot_id': None, 'created': False, 'activated': False}
    if dry_run:
        return result
    identity = {key: metadata[key] for key in ('sha256', 'sheet_name', 'header_row', 'first_row', 'last_row')}
    with transaction.atomic():
        # Serialize switches when versions exist; the conditional unique constraint
        # also prevents two first imports from publishing simultaneous active sources.
        list(ReceivablesSourceSnapshot.objects.select_for_update().values_list('pk', flat=True))
        snapshot = ReceivablesSourceSnapshot.objects.filter(**identity).first()
        created = snapshot is None
        if created:
            snapshot = ReceivablesSourceSnapshot.objects.create(**metadata, reconciliation=reconciliation)
            # Workbook facts are independent of operational invoice identities.
            # Matching numbers cannot establish a reliable link in a legacy
            # register that may contain duplicate IDs or invoice numbers.
            records = [ReceivablesSourceRow(snapshot=snapshot, updated_at=snapshot.imported_at, **row)
                       for row in rows]
            ReceivablesSourceRow.objects.bulk_create(records, batch_size=250)
        elif snapshot.row_count != len(rows) or snapshot.rows.count() != len(rows):
            raise ValueError('The existing source version has an inconsistent row count.')
        else:
            existing = snapshot.reconciliation.get('workbook_summary')
            if existing is None:
                # Older versions contain the same immutable facts but predate the
                # stored aggregate. Preserve their original import provenance.
                summary['source']['file_name'] = snapshot.file_name
                reconciliation = {**snapshot.reconciliation, 'workbook_summary': summary}
                ReceivablesSourceSnapshot.objects.filter(pk=snapshot.pk).update(reconciliation=reconciliation)
                snapshot.reconciliation = reconciliation
            else:
                validate_source_summary(existing, snapshot)
                reconciliation = snapshot.reconciliation
            result.update(file_name=snapshot.file_name, reconciliation=reconciliation)
        activated = not snapshot.is_active
        if activated:
            ReceivablesSourceSnapshot.objects.filter(is_active=True).update(is_active=False)
            ReceivablesSourceSnapshot.objects.filter(pk=snapshot.pk).update(is_active=True)
        result.update(snapshot_id=snapshot.pk, created=created, activated=activated)
    return result


def get_active_receivables_source():
    """Return a queryset pinned to one source version, or None before first import."""
    snapshot = ReceivablesSourceSnapshot.objects.filter(is_active=True).first()
    if snapshot is None:
        return None
    queryset = ReceivablesSourceRow.objects.filter(snapshot_id=snapshot.pk)
    queryset._receivables_snapshot = snapshot
    return queryset
