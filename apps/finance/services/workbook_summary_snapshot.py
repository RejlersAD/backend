"""Reproducible aggregate-only snapshot generation; never imports invoices."""
import hashlib
from collections import Counter
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from .workbook_summary import AMOUNT_FIELDS, COVERAGE_FIELDS, STATUS_GROUPS, validate_snapshot


HEADERS = {'A': 'Invoice #', 'G': 'RAD Project #', 'L': 'Invoice Amount',
           'M': 'Inv Amt. (AED)', 'R': 'Payment Status', 'AA': 'Actual Payment Received'}
AMOUNT_INDICES = {'invoice_amount': 11, 'invoice_amount_aed': 12, 'actual_payment_received': 26}
FOOTER_LABELS = {'total', 'grand total', 'subtotal', 'sub total'}


def _text(value):
    if value is None:
        return ''
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        number = Decimal(str(value))
        if number.is_finite() and number == number.to_integral_value():
            return str(int(number))
    return str(value).strip()


def _cell_kind(cell):
    value = cell.value
    if cell.data_type == 'e':
        return 'error_count'
    if value is None or (isinstance(value, str) and not value.strip()):
        return 'blank_count'
    if (not isinstance(value, bool) and isinstance(value, (int, float, Decimal))
            and Decimal(str(value)).is_finite()):
        return 'numeric_count'
    return 'text_count'


def generate_workbook_snapshot(workbook_path, *, last_row, sheet='External Invoice ',
                               header_row=5, snapshot_at=None):
    """Sum cached numeric cells within a verified, contiguous invoice row range.

    Explicit bounds plus invoice-ID checks prevent the worksheet's later total
    formulas from entering the sum. Duplicate invoice IDs remain separate rows.
    Text amounts, errors, and blank cells are reported, never coerced into money.
    """
    path = Path(workbook_path)
    first_row = header_row + 1
    if header_row < 1 or last_row < first_row:
        raise ValueError('The invoice range must follow the header row.')
    captured = snapshot_at or datetime.now(timezone.utc)
    if captured.tzinfo is None or captured.utcoffset() is None:
        raise ValueError('snapshot_at must include a timezone.')
    with path.open('rb') as source:
        digest = hashlib.file_digest(source, 'sha256').hexdigest()
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in workbook.sheetnames:
            raise ValueError('The requested invoice sheet was not found.')
        worksheet = workbook[sheet]
        if last_row > worksheet.max_row:
            raise ValueError('The last invoice row is outside the worksheet.')
        headers = next(worksheet.iter_rows(min_row=header_row, max_row=header_row, max_col=27))
        for column, expected in HEADERS.items():
            if _text(headers[column_index_from_string(column) - 1].value) != expected:
                raise ValueError(f'Unexpected header at {column}{header_row}.')
        totals = {key: Decimal('0') for key in AMOUNT_FIELDS}
        coverage = {key: dict.fromkeys(COVERAGE_FIELDS, 0) for key in AMOUNT_FIELDS}
        statuses = Counter()
        other_statuses = Counter()
        projects = set()
        project_excluded_rows = 0
        count = 0
        for row_number, row in enumerate(worksheet.iter_rows(min_row=first_row, max_col=27), first_row):
            invoice_id = _text(row[0].value)
            footer = invoice_id.casefold().rstrip(':') in FOOTER_LABELS
            if row_number > last_row:
                if invoice_id and not footer:
                    raise ValueError(f'An invoice row exists after the specified last row: {row_number}.')
                continue
            if not invoice_id or footer or row[0].data_type == 'e':
                raise ValueError(f'Row {row_number} is not an invoice row; exclude footers from the range.')
            count += 1
            for key, index in AMOUNT_INDICES.items():
                cell = row[index]
                kind = _cell_kind(cell)
                coverage[key][kind] += 1
                if kind == 'numeric_count':
                    totals[key] += Decimal(str(cell.value))
            project = _text(row[6].value)
            if not project or project.casefold() == 'n/a' or row[6].data_type == 'e':
                project_excluded_rows += 1
            else:
                projects.add(project)
            status = _text(row[17].value).casefold()
            if status in {'paid', 'cancelled', 'pending', 'new'}:
                statuses[status] += 1
            else:
                statuses['other'] += 1
                label = (status.upper() if row[17].data_type == 'e'
                         else status.title() if status else 'Not recorded')
                other_statuses[label] += 1
    finally:
        workbook.close()
    return validate_snapshot({
        'schema_version': '1.0', 'status': 'available',
        'source': {'file_name': path.name, 'sheet': sheet, 'first_row': first_row,
                   'last_row': last_row, 'snapshot_at': captured.isoformat(),
                   'sha256': digest, 'scope': 'full_workbook'},
        'invoice_count': count,
        'totals': {**{key: str(value.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP))
                     for key, value in totals.items()}, 'project_count': len(projects)},
        'coverage': coverage,
        'payment_status': [{'id': key, 'label': label, 'count': statuses[key]}
                           for key, label in STATUS_GROUPS],
        'other_statuses': [{'label': label, 'count': value}
                           for label, value in sorted(other_statuses.items())],
        'project_excluded_rows': project_excluded_rows,
        'currency_basis': 'mixed_original',
    })
