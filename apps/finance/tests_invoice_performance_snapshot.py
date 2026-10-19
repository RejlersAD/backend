"""Offline workbook layout changes preserve Finance's daily aggregate facts."""
from datetime import date, datetime, timezone
from pathlib import Path
from tempfile import TemporaryDirectory

from django.test import SimpleTestCase
from openpyxl import Workbook, load_workbook

from apps.finance.services.invoice_performance_snapshot import generate_invoice_performance_snapshot
from apps.finance.services.workbook_summary_snapshot import HEADERS


class InvoicePerformanceSnapshotLayoutTests(SimpleTestCase):
    def setUp(self):
        directory = TemporaryDirectory(prefix='finance-performance-layout-')
        self.addCleanup(directory.cleanup)
        self.legacy = Path(directory.name) / 'original.xlsx'
        self.shifted = Path(directory.name) / 'business-unit.xlsx'
        self.captured = datetime(2026, 10, 19, 8, tzinfo=timezone.utc)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'External Invoice '
        for column, label in {**HEADERS, 'B': 'Invoice Date'}.items():
            sheet[f'{column}5'] = label
        for number, values in enumerate([
            {'A': '1', 'B': date(2026, 10, 1), 'G': 'P1', 'L': 100, 'M': 100,
             'R': 'Pending', 'AA': 40, 'AE': 'AED'},
            {'A': '2', 'B': date(2026, 10, 2), 'G': 'P2', 'L': 200, 'M': 734,
             'R': 'Paid', 'AA': 75, 'AE': 'USD'},
            {'A': '3', 'B': date(2026, 10, 3), 'G': 'P2', 'L': 300, 'M': 300,
             'R': 'Cancelled', 'AA': 0, 'AE': 'AED'},
        ], 6):
            for column, value in values.items():
                sheet[f'{column}{number}'] = value
        # This receipt conflicts with the invoice currency; mapping must retain
        # the cell format and existing unknown-receipt treatment.
        sheet['AA7'].number_format = '[$EUR] #,##0.00'
        workbook.save(self.legacy)
        sheet.insert_cols(8)
        sheet['H5'] = 'Business Unit'
        for number in range(6, 9):
            sheet[f'H{number}'] = 'Synthetic business unit'
        workbook.save(self.shifted)
        workbook.close()

    def generate(self, path):
        return generate_invoice_performance_snapshot(path, last_row=8, snapshot_at=self.captured)

    def test_business_unit_insertion_preserves_daily_facts_and_reports_physical_columns(self):
        original = self.generate(self.legacy)
        shifted = self.generate(self.shifted)
        self.assertEqual(shifted['days'], original['days'])
        self.assertEqual(shifted['coverage'], original['coverage'])
        self.assertEqual(shifted['coverage']['receipt_currency_mismatch_count'], 1)
        for key, old, new in [('amount_column', 'L', 'M'), ('receipt_column', 'AA', 'AB'),
                              ('currency_column', 'AE', 'AF')]:
            self.assertEqual(original['source'][key], old)
            self.assertEqual(shifted['source'][key], new)
        self.assertNotEqual(shifted['source']['sha256'], original['source']['sha256'])

    def test_shifted_required_header_must_still_match_before_aggregation(self):
        workbook = load_workbook(self.shifted)
        workbook.active['S5'] = 'Unexpected status field'
        workbook.save(self.shifted)
        workbook.close()
        with self.assertRaisesMessage(ValueError, 'Unexpected header at S5.'):
            self.generate(self.shifted)
