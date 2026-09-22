"""Source imports preserve Finance's cells and never rewrite the operational register."""
import io
import json
from datetime import date
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from django.core.management import call_command
from django.db import IntegrityError, connection, transaction
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from openpyxl import Workbook, load_workbook

from apps.finance.receivables_source_models import ReceivablesSourceRow, ReceivablesSourceSnapshot
from apps.finance.services.receivables_source import (
    SOURCE_HEADERS, get_active_receivables_source, import_receivables_source, read_receivables_source,
)
from apps.finance.services.workbook_summary import build_workbook_summary
from apps.finance.services.receivables_dashboard import build_receivables_dashboard
from apps.invoice_tracker.models import CustomerInvoice


class ReceivablesSourceTests(TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)

    def workbook(self, rows, *, name='finance.xlsx'):
        path = Path(self.directory.name) / name
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'External Invoice '
        for column, header in SOURCE_HEADERS.items():
            sheet[f'{column}5'] = header
        for number, row in enumerate(rows, 6):
            values = {'A': f'INV-{number}', 'B': date(2026, 9, 1), 'E': 'Account reference',
                      'F': 'Customer', 'N': date(2026, 9, 20), 'R': 'Pending', **row}
            for column, value in values.items():
                if column != 'formats':
                    sheet[f'{column}{number}'] = value
            for column, number_format in row.get('formats', {}).items():
                sheet[f'{column}{number}'].number_format = number_format
        workbook.save(path)
        workbook.close()
        return path

    def load(self, path, count, **kwargs):
        return import_receivables_source(path, last_row=5 + count, **kwargs)

    def test_source_preserves_monetary_cells_status_and_independent_currencies(self):
        path = self.workbook([
            {'A': 'USD', 'L': 13650, 'M': 50095.5, 'R': 'Overdue',
             'formats': {'L': '[$USD] #,##0.00', 'M': '[$AED] #,##0.00'}},
            {'A': 'PARTIAL', 'L': 1000, 'M': 1000, 'R': 'Paid (partial)', 'Y': 12.34567891,
             'AA': 300, 'formats': {'L': '[$AED] #,##0.00', 'Y': '[$AED] #,##0.00',
                                    'AA': '[$EUR] #,##0.00'}},
            {'A': 'NEW', 'L': 40, 'M': 80, 'R': 'New', 'Y': '#VALUE!', 'AA': None},
            {'A': 'REJECTED', 'L': 60, 'M': '#VALUE!', 'R': 'Rejected', 'N': None},
            {'A': 'CONFLICT', 'L': 10, 'M': 40, 'AE': 'EUR', 'formats': {'L': '[$USD] #,##0.00'}},
        ])
        metadata, rows, report = read_receivables_source(path, last_row=10)
        self.assertEqual(metadata['row_count'], 5)
        self.assertEqual(rows[0]['invoice_amount'], Decimal('13650'))
        self.assertEqual(rows[0]['invoice_amount_aed'], Decimal('50095.5'))
        self.assertEqual(rows[0]['currency'], 'USD')
        self.assertEqual(rows[0]['balance_currency'], '')
        self.assertIsNone(rows[0]['balance_to_be_received'])
        self.assertEqual(rows[1]['payment_status'], 'partial')
        self.assertEqual(rows[1]['raw_payment_status'], 'Paid (partial)')
        self.assertEqual(rows[1]['balance_to_be_received'], Decimal('12.34567891'))
        self.assertEqual(rows[1]['balance_currency'], 'AED')
        self.assertEqual(rows[1]['actual_payment_currency'], 'EUR')
        self.assertEqual(rows[2]['payment_status'], 'new')
        self.assertIsNone(rows[2]['balance_to_be_received'])
        self.assertEqual(rows[2]['currency_status'], 'not_recorded')
        self.assertEqual(rows[3]['payment_status'], 'rejected')
        self.assertIsNone(rows[3]['invoice_amount_aed'])
        self.assertIsNone(rows[3]['due_date'])
        self.assertEqual(rows[4]['currency_status'], 'conflict')
        self.assertEqual(rows[4]['currency'], '')
        self.assertEqual(report['overdue'], {'count': 1, 'amount_aed': '50095.50', 'missing_amount_count': 0})

    def test_dry_run_command_reconciles_without_database_queries(self):
        path = self.workbook([{'L': 100, 'M': 367, 'R': 'Overdue'},
                              {'L': 20, 'M': None, 'R': 'Overdue'}])
        output = io.StringIO()
        with CaptureQueriesContext(connection) as queries:
            call_command('import_receivables_source', str(path), last_row=7, dry_run=True, stdout=output)
        self.assertEqual(len(queries), 0)
        report = json.loads(output.getvalue())
        self.assertTrue(report['dry_run'])
        self.assertFalse(report['activated'])
        self.assertEqual(report['reconciliation']['overdue'],
                         {'count': 2, 'amount_aed': '367.00', 'missing_amount_count': 1})
        self.assertEqual(ReceivablesSourceSnapshot.objects.count(), 0)

    def test_import_never_queries_links_or_changes_operational_invoices(self):
        for number in ('UNIQUE', 'DUPLICATE'):
            invoice = CustomerInvoice(invoice_number=number, payment_status='paid', currency='AED',
                                      invoice_amount=999, invoice_amount_aed=999,
                                      balance_to_be_received=0, actual_payment_received=999)
            invoice.save(_skip_recompute=True)
        before = list(CustomerInvoice.objects.order_by('pk').values())
        path = self.workbook([{'A': 'UNIQUE', 'L': 10, 'M': 36.7, 'R': 'Overdue'},
                              {'A': 'DUPLICATE', 'L': 20, 'M': 20},
                              {'A': 'DUPLICATE', 'L': 30, 'M': 30}])
        with CaptureQueriesContext(connection) as queries, patch.object(
            CustomerInvoice, 'save', side_effect=AssertionError('Operational invoices must not be saved')
        ):
            report = self.load(path, 3)
        self.assertFalse(any('invoice_tracker_customerinvoice' in query['sql'] for query in queries))
        self.assertTrue(report['created'])
        self.assertTrue(report['activated'])
        self.assertEqual(report['reconciliation']['duplicate_invoice_number_count'], 1)
        self.assertEqual(list(CustomerInvoice.objects.order_by('pk').values()), before)
        rows = list(get_active_receivables_source())
        self.assertEqual([row.register_invoice_id for row in rows], [None, None, None])
        self.assertEqual(rows[0].payment_status, 'overdue')
        self.assertEqual(rows[0].invoice_amount_aed, Decimal('36.7'))
        self.assertEqual([row.row_number for row in rows], [6, 7, 8])

    def test_same_source_is_idempotent_and_activation_retains_previous_versions(self):
        self.assertIsNone(get_active_receivables_source())
        first_path = self.workbook([{'L': 10, 'M': 10}])
        first = self.load(first_path, 1)
        pinned = get_active_receivables_source()
        repeated = self.load(first_path, 1)
        self.assertEqual(repeated['snapshot_id'], first['snapshot_id'])
        self.assertFalse(repeated['created'])
        self.assertFalse(repeated['activated'])
        second_path = self.workbook([{'L': 20, 'M': 20}], name='replacement.xlsx')
        second = self.load(second_path, 1)
        self.assertNotEqual(first['snapshot_id'], second['snapshot_id'])
        self.assertEqual(ReceivablesSourceSnapshot.objects.count(), 2)
        self.assertEqual(ReceivablesSourceRow.objects.count(), 2)
        self.assertEqual(ReceivablesSourceSnapshot.objects.filter(is_active=True).count(), 1)
        self.assertEqual(get_active_receivables_source()._receivables_snapshot.pk, second['snapshot_id'])
        self.assertEqual(pinned.get().invoice_amount, Decimal('10'))
        self.assertEqual(get_active_receivables_source().get().invoice_amount, Decimal('20'))
        restored = self.load(first_path, 1)
        self.assertFalse(restored['created'])
        self.assertTrue(restored['activated'])
        self.assertEqual(ReceivablesSourceRow.objects.count(), 2)

    def test_failed_import_keeps_previous_active_snapshot_and_no_partial_rows(self):
        first = self.load(self.workbook([{'L': 10, 'M': 10}]), 1)
        replacement = self.workbook([{'L': 20, 'M': 20}], name='replacement.xlsx')
        with patch.object(ReceivablesSourceRow.objects, 'bulk_create', side_effect=RuntimeError('Write failed')):
            with self.assertRaisesMessage(RuntimeError, 'Write failed'):
                self.load(replacement, 1)
        self.assertEqual(ReceivablesSourceSnapshot.objects.count(), 1)
        self.assertEqual(ReceivablesSourceRow.objects.count(), 1)
        self.assertEqual(get_active_receivables_source()._receivables_snapshot.pk, first['snapshot_id'])

    def test_invalid_bounds_or_headers_cannot_publish_incomplete_source(self):
        path = self.workbook([{'L': 10, 'M': 10}, {'L': 20, 'M': 20}])
        with self.assertRaisesMessage(ValueError, 'after the specified last row'):
            self.load(path, 1)
        with self.assertRaisesMessage(ValueError, 'Only the External Invoice'):
            self.load(path, 2, sheet='Internal Invoice')
        with self.assertRaisesMessage(ValueError, 'immediately follow'):
            self.load(path, 2, first_row=7)
        from openpyxl import load_workbook
        workbook = load_workbook(path)
        workbook.active['M5'] = 'Wrong amount'
        workbook.save(path)
        workbook.close()
        with self.assertRaisesMessage(ValueError, 'Unexpected source header at M5'):
            self.load(path, 2)
        self.assertEqual(ReceivablesSourceSnapshot.objects.count(), 0)

    def test_database_prevents_two_active_snapshots(self):
        first = self.load(self.workbook([{'L': 10, 'M': 10}]), 1)
        with self.assertRaises(IntegrityError), transaction.atomic():
            ReceivablesSourceSnapshot.objects.create(sha256='a' * 64, file_name='other.xlsx',
                                                     sheet_name='External Invoice ', last_row=6,
                                                     row_count=1, is_active=True)
        self.assertEqual(get_active_receivables_source()._receivables_snapshot.pk, first['snapshot_id'])

    def test_upload_detects_sheet_and_bounds_and_publishes_matching_summary(self):
        path = self.workbook([{'L': 10, 'M': 36.7, 'R': 'Overdue'},
                              {'L': 20, 'M': 80, 'R': 'Paid'}])
        workbook = load_workbook(path)
        workbook.active.title = ' external invoice '
        workbook.active['A8'] = 'Grand Total:'
        workbook.active['M8'] = 99999
        workbook.active['M12'] = 'Trailing note'
        workbook.save(path)
        workbook.close()
        report = import_receivables_source(path, last_row=None, original_filename='Uploaded Finance.xlsx')
        snapshot = ReceivablesSourceSnapshot.objects.get(pk=report['snapshot_id'])
        summary = snapshot.reconciliation['workbook_summary']
        self.assertEqual(snapshot.file_name, 'Uploaded Finance.xlsx')
        self.assertEqual(snapshot.sheet_name, ' external invoice ')
        self.assertEqual((snapshot.first_row, snapshot.last_row, snapshot.row_count), (6, 7, 2))
        self.assertEqual(summary['source']['sha256'], snapshot.sha256)
        self.assertEqual(summary['source']['file_name'], snapshot.file_name)
        self.assertEqual(summary['source']['sheet'], snapshot.sheet_name)
        self.assertEqual(summary['invoice_count'], 2)
        self.assertEqual(summary['totals']['invoice_amount_aed'], '116.70')
        self.assertEqual(snapshot.rows.count(), 2)

    def test_auto_bounds_reject_empty_gaps_footers_followed_by_invoices_and_ambiguous_sheets(self):
        for rows in ([], [{'A': None}, {'L': 20}], [{'A': 'Total'}, {'L': 20}]):
            with self.subTest(rows=rows):
                with self.assertRaises(ValueError):
                    import_receivables_source(self.workbook(rows), last_row=None)
        path = self.workbook([{'L': 10, 'M': 10}])
        workbook = load_workbook(path)
        workbook.create_sheet('External Invoice')
        workbook.save(path)
        workbook.close()
        with self.assertRaisesMessage(ValueError, 'more than one External Invoice'):
            import_receivables_source(path, last_row=None)
        self.assertFalse(ReceivablesSourceSnapshot.objects.exists())

    def test_summary_validation_failure_never_publishes_rows(self):
        first = self.load(self.workbook([{'L': 10, 'M': 10}]), 1)
        replacement = self.workbook([{'L': 20, 'M': 20}], name='replacement.xlsx')
        with patch('apps.finance.services.receivables_source.generate_workbook_snapshot',
                   side_effect=ValueError('Invalid aggregate')):
            with self.assertRaisesMessage(ValueError, 'Invalid aggregate'):
                import_receivables_source(replacement, last_row=None)
        self.assertEqual(ReceivablesSourceSnapshot.objects.count(), 1)
        self.assertEqual(ReceivablesSourceRow.objects.count(), 1)
        self.assertEqual(get_active_receivables_source()._receivables_snapshot.pk, first['snapshot_id'])

    def test_oversized_source_amount_is_a_validation_error_before_publication(self):
        for amount in (1e20, -1e20):
            with self.subTest(amount=amount):
                with self.assertRaisesMessage(ValueError, 'exceeds the supported precision'):
                    import_receivables_source(self.workbook([{'L': amount, 'M': 1}]), last_row=None)
        self.assertFalse(ReceivablesSourceSnapshot.objects.exists())

    def test_same_file_backfills_legacy_summary_and_keeps_original_provenance(self):
        path = self.workbook([{'L': 10, 'M': 36.7}])
        first = import_receivables_source(path, last_row=None, original_filename='First upload.xlsx')
        snapshot = ReceivablesSourceSnapshot.objects.get(pk=first['snapshot_id'])
        snapshot.reconciliation.pop('workbook_summary')
        snapshot.save(update_fields=['reconciliation'])
        before_rows = list(snapshot.rows.values())
        repeated = import_receivables_source(path, last_row=None, original_filename='Renamed upload.xlsx')
        snapshot.refresh_from_db()
        self.assertFalse(repeated['created'])
        self.assertFalse(repeated['activated'])
        self.assertEqual(repeated['snapshot_id'], first['snapshot_id'])
        self.assertEqual(repeated['file_name'], 'First upload.xlsx')
        self.assertEqual(snapshot.reconciliation['workbook_summary']['source']['file_name'], snapshot.file_name)
        self.assertEqual(snapshot.reconciliation['workbook_summary']['source']['sha256'], snapshot.sha256)
        self.assertEqual(list(snapshot.rows.values()), before_rows)
        repeated_again = import_receivables_source(path, last_row=None)
        self.assertEqual(repeated_again['reconciliation'], snapshot.reconciliation)

    def test_summary_reader_uses_active_or_explicitly_pinned_source_and_no_static_totals_before_upload(self):
        with patch('apps.finance.services.workbook_summary.module_action_allowed', return_value=True):
            self.assertEqual(build_workbook_summary('reader')['status'], 'unavailable')
            first = self.load(self.workbook([{'L': 10, 'M': 36.7}]), 1)
            pinned = ReceivablesSourceSnapshot.objects.get(pk=first['snapshot_id'])
            self.load(self.workbook([{'L': 20, 'M': 80}], name='replacement.xlsx'), 1)
            self.assertEqual(build_workbook_summary('reader')['totals']['invoice_amount_aed'], '80.00')
            with CaptureQueriesContext(connection) as queries, patch.object(Path, 'open') as opened:
                summary = build_workbook_summary('reader', source_snapshot=pinned)
            self.assertEqual(len(queries), 0)
            opened.assert_not_called()
            self.assertEqual(summary['source']['sha256'], pinned.sha256)
            self.assertEqual(summary['totals']['invoice_amount_aed'], '36.70')
            self.assertEqual(build_workbook_summary('reader', source_snapshot=None)['status'], 'unavailable')

    def test_dashboard_pins_full_summary_to_its_selected_rows_even_after_another_upload(self):
        self.load(self.workbook([{'L': 10, 'M': 36.7, 'F': 'Acme', 'AE': 'USD'},
                                 {'L': 20, 'M': 80, 'F': 'Other', 'AE': 'AED'}]), 2)
        pinned = get_active_receivables_source()
        self.load(self.workbook([{'L': 30, 'M': 100}], name='replacement.xlsx'), 1)
        with patch('apps.finance.services.receivables_dashboard.get_active_receivables_source', return_value=pinned), \
                patch('apps.finance.services.receivables_dashboard.module_action_allowed',
                      side_effect=lambda user, module, action: module == 'finance_outgoing'), \
                patch('apps.finance.services.receivables_dashboard.build_invoice_performance', return_value={}), \
                patch('apps.finance.services.workbook_summary.module_action_allowed', return_value=True):
            all_rows = build_receivables_dashboard('reader', as_of=date(2026, 9, 22))
            filtered = build_receivables_dashboard('reader', currency='USD', company='Acme',
                                                   as_of=date(2026, 9, 22))
        self.assertEqual(all_rows['workbook_summary'], filtered['workbook_summary'])
        self.assertEqual(filtered['sources']['receivables']['invoice_count'], 1)
        self.assertEqual(filtered['workbook_summary']['invoice_count'], 2)
        self.assertEqual(filtered['workbook_summary']['totals']['invoice_amount_aed'], '116.70')
        self.assertEqual(filtered['sources']['receivables']['sha256'], pinned._receivables_snapshot.sha256)
        self.assertEqual(filtered['workbook_summary']['source']['sha256'], pinned._receivables_snapshot.sha256)
