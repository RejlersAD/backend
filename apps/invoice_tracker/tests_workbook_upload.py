"""Workbook uploads publish one coherent reporting source without rewriting invoices."""
import hashlib
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch

from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook
from rest_framework.test import APIRequestFactory, force_authenticate

from apps.finance import tests_command_center as finance_helpers
from apps.finance import tests_receivables_source as workbook_helpers
from apps.finance.command_center_views import FinanceCustomerInvoiceRegisterView, FinanceReceivablesDashboardView
from apps.finance.receivables_source_models import ReceivablesSourceRow, ReceivablesSourceSnapshot
from apps.invoice_tracker.models import CustomerInvoice
from apps.invoice_tracker.views import CustomerInvoiceViewSet
from apps.rbac.models import Permission, RolePermission


UPLOAD_URL = '/api/v1/invoice-tracker/invoices/import-excel/'
XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


class InvoiceWorkbookUploadTests(TestCase):
    workbook = workbook_helpers.ReceivablesSourceTests.workbook
    grant = finance_helpers.FinanceCommandCenterTests.grant

    def setUp(self):
        finance_helpers.FinanceCommandCenterTests.setUp(self)
        self.directory = TemporaryDirectory(prefix='invoice-upload-tests-')
        self.addCleanup(self.directory.cleanup)
        self.factory = APIRequestFactory()
        self.grant('finance_overview', 'finance_outgoing')
        for permission in Permission.objects.filter(module__code='finance_outgoing', action='create', is_active=True):
            RolePermission.objects.get_or_create(role=self.role, permission=permission)
        invoice = CustomerInvoice(invoice_number='OPERATIONAL-UNCHANGED', payment_status='paid', currency='AED',
                                  invoice_amount=Decimal('999'), invoice_amount_aed=Decimal('999'),
                                  balance_to_be_received=Decimal('0'), actual_payment_received=Decimal('999'))
        invoice.save(_skip_recompute=True)
        self.operational_before = list(CustomerInvoice.objects.order_by('pk').values())

    def upload_bytes(self, name, content, **params):
        request = self.factory.post(UPLOAD_URL, {
            'file': SimpleUploadedFile(name, content, content_type=XLSX), **params,
        }, format='multipart')
        force_authenticate(request, self.user)
        with CaptureQueriesContext(connection) as queries:
            response = CustomerInvoiceViewSet.as_view({'post': 'import_excel'})(request)
        if params.get('mode') != 'operational':
            self.assertEqual([row['sql'] for row in queries.captured_queries
                              if CustomerInvoice._meta.db_table.lower() in row['sql'].lower()], [])
            self.assertEqual(list(CustomerInvoice.objects.order_by('pk').values()), self.operational_before)
        return response

    def upload(self, path, **params):
        return self.upload_bytes(path.name, path.read_bytes(), **params)

    def get_report(self, view, url):
        request = self.factory.get(url)
        force_authenticate(request, self.user)
        response = view.as_view()(request)
        self.assertEqual(response.status_code, 200, response.data)
        return response.data

    def assert_reports_match(self, upload, amount, count):
        source = upload['source']
        register = self.get_report(FinanceCustomerInvoiceRegisterView, '/api/v1/finance/dashboard/customer-invoices/')
        dashboard = self.get_report(FinanceReceivablesDashboardView, '/api/v1/finance/dashboard/receivables/')
        summary = dashboard['workbook_summary']
        self.assertEqual(register['source']['snapshot_id'], source['snapshot_id'])
        self.assertEqual(register['source']['sha256'], source['sha256'])
        self.assertEqual(dashboard['sources']['receivables']['snapshot_id'], source['snapshot_id'])
        self.assertEqual(dashboard['sources']['receivables']['sha256'], source['sha256'])
        self.assertEqual(summary['source']['sha256'], source['sha256'])
        self.assertEqual(summary['source']['file_name'], source['file_name'])
        self.assertEqual(summary['invoice_count'], count)
        self.assertEqual(register['pagination']['count'], count)
        self.assertEqual(register['totals']['amount_home']['amount'], amount)
        self.assertEqual(dashboard['kpis']['overdue']['amount'], amount)
        self.assertEqual(summary['totals']['invoice_amount_aed'], amount)

    def test_default_upload_publishes_workbook_and_updates_both_reporting_surfaces(self):
        path = self.workbook([
            {'A': 'REPEATED', 'L': 100, 'M': 100, 'R': 'Overdue', 'AE': 'AED'},
            {'A': 'REPEATED', 'L': -25, 'M': -25, 'R': 'Overdue', 'AE': 'AED'},
        ], name='Finance September.xlsx')
        response = self.upload(path)
        self.assertEqual(response.status_code, 200, response.data)
        data = response.data
        self.assertEqual(data['mode'], 'workbook')
        self.assertEqual(data['rows_published'], 2)
        self.assertEqual(data['source']['sha256'], hashlib.sha256(path.read_bytes()).hexdigest())
        self.assertEqual(data['source']['file_name'], path.name)
        self.assertEqual(data['source']['row_count'], 2)
        self.assertEqual({key: data[key] for key in ('rows_created', 'rows_updated', 'rows_skipped', 'rows_seen', 'sheets_processed')},
                         {'rows_created': 2, 'rows_updated': 0, 'rows_skipped': 0, 'rows_seen': 2, 'sheets_processed': 1})
        self.assertEqual(data['errors'], [])
        self.assertEqual(data['warnings'], [])
        self.assertEqual(data['reconciliation']['duplicate_invoice_number_count'], 1)
        self.assertEqual(response['Cache-Control'], 'private, no-store')
        self.assert_reports_match(data, '75.00', 2)

        replacement = self.workbook([
            {'A': 'REPLACEMENT', 'L': 250, 'M': 250, 'R': 'Overdue', 'AE': 'AED'},
        ], name='Finance October.xlsx')
        updated = self.upload(replacement)
        self.assertEqual(updated.status_code, 200, updated.data)
        self.assertNotEqual(updated.data['source']['sha256'], data['source']['sha256'])
        self.assertNotEqual(updated.data['source']['snapshot_id'], data['source']['snapshot_id'])
        self.assert_reports_match(updated.data, '250.00', 1)
        self.assertEqual(ReceivablesSourceSnapshot.objects.count(), 2)
        self.assertEqual(ReceivablesSourceSnapshot.objects.filter(is_active=True).count(), 1)
        self.assertEqual(ReceivablesSourceRow.objects.filter(snapshot_id=data['source']['snapshot_id']).count(), 2)

    def test_upload_preserves_source_statuses_cells_and_duplicate_rows(self):
        path = self.workbook([
            {'A': 'DUPLICATE', 'L': 100, 'M': 120, 'R': 'Paid (partial)', 'Y': 7.125,
             'AA': 3, 'AE': 'AED'},
            {'A': 'DUPLICATE', 'L': -10, 'M': -12, 'R': 'New', 'AA': None, 'AE': 'AED'},
            {'A': 'ERROR', 'L': 10, 'M': '#VALUE!', 'R': 'Rejected', 'AA': None, 'AE': 'AED'},
        ])
        response = self.upload(path, mode='workbook')
        self.assertEqual(response.status_code, 200, response.data)
        rows = list(ReceivablesSourceRow.objects.order_by('row_number'))
        self.assertEqual([row.row_number for row in rows], [6, 7, 8])
        self.assertEqual([row.invoice_number for row in rows], ['DUPLICATE', 'DUPLICATE', 'ERROR'])
        self.assertEqual([row.raw_payment_status for row in rows], ['Paid (partial)', 'New', 'Rejected'])
        self.assertEqual([row.payment_status for row in rows], ['partial', 'new', 'rejected'])
        self.assertEqual(rows[0].balance_to_be_received, Decimal('7.125'))
        self.assertEqual(rows[0].actual_payment_received, Decimal('3'))
        self.assertEqual(rows[1].invoice_amount_aed, Decimal('-12'))
        self.assertIsNone(rows[1].actual_payment_received)
        self.assertIsNone(rows[2].invoice_amount_aed)
        self.assertTrue(all(row.register_invoice_id is None for row in rows))

    def test_repeated_upload_is_idempotent_and_does_not_duplicate_source_rows(self):
        path = self.workbook([{'L': 20, 'M': 20, 'R': 'Overdue', 'AE': 'AED'}])
        first, second = self.upload(path), self.upload(path)
        self.assertEqual(first.status_code, 200, first.data)
        self.assertEqual(second.status_code, 200, second.data)
        self.assertEqual(first.data['source'], second.data['source'])
        self.assertEqual(second.data['rows_created'], 0)
        self.assertEqual(ReceivablesSourceSnapshot.objects.count(), 1)
        self.assertEqual(ReceivablesSourceRow.objects.count(), 1)
        self.assert_reports_match(second.data, '20.00', 1)

    def test_invalid_workbooks_leave_previous_source_active_without_invoice_writes(self):
        accepted = self.upload(self.workbook([{'L': 15, 'M': 15, 'R': 'Overdue', 'AE': 'AED'}]))
        self.assertEqual(accepted.status_code, 200, accepted.data)
        active_id = accepted.data['source']['snapshot_id']
        no_external = self.workbook([{'L': 100, 'M': 100}], name='missing-external.xlsx')
        workbook = load_workbook(no_external)
        workbook.active.title = 'Internal Invoice'
        workbook.save(no_external)
        workbook.close()
        invalid_header = self.workbook([{'L': 100, 'M': 100}], name='invalid-header.xlsx')
        workbook = load_workbook(invalid_header)
        workbook.active['M5'] = 'Unexpected amount column'
        workbook.save(invalid_header)
        workbook.close()
        after_footer = self.workbook([
            {'A': 'FIRST', 'L': 1, 'M': 1}, {'A': 'Total', 'M': 1}, {'A': 'AFTER-TOTAL', 'L': 2, 'M': 2},
        ], name='after-footer.xlsx')
        invalid = [('malformed.xlsx', b'not an Excel ZIP archive'),
                   *[(path.name, path.read_bytes()) for path in (no_external, invalid_header, after_footer)]]
        for name, content in invalid:
            with self.subTest(workbook=name):
                response = self.upload_bytes(name, content)
                self.assertEqual(response.status_code, 400, response.data)
                self.assertEqual(ReceivablesSourceSnapshot.objects.count(), 1)
                self.assertEqual(ReceivablesSourceSnapshot.objects.get(is_active=True).pk, active_id)
                self.assertEqual(ReceivablesSourceRow.objects.count(), 1)
        self.assert_reports_match(accepted.data, '15.00', 1)

    def test_read_only_user_cannot_publish_a_workbook(self):
        RolePermission.objects.filter(role=self.role, permission__module__code='finance_outgoing',
                                      permission__action='create').delete()
        cache.clear()
        response = self.upload(self.workbook([{'L': 100, 'M': 100}]))
        self.assertEqual(response.status_code, 403, response.data)
        self.assertEqual(ReceivablesSourceSnapshot.objects.count(), 0)
        self.assertEqual(ReceivablesSourceRow.objects.count(), 0)

    def test_explicit_operational_mode_uses_legacy_importer_without_publishing_snapshot(self):
        result = {'sheets_processed': 1, 'rows_seen': 2, 'rows_created': 1, 'rows_updated': 1,
                  'rows_skipped': 0, 'errors': [], 'warnings': []}
        with patch('apps.invoice_tracker.views.import_workbook',
                   return_value=SimpleNamespace(as_dict=lambda: result)) as legacy:
            response = self.upload(self.workbook([{'L': 100, 'M': 100}]),
                                   mode='operational', sheets='External Invoice ')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['mode'], 'operational')
        legacy.assert_called_once()
        self.assertEqual(legacy.call_args.kwargs['user'], self.user)
        self.assertEqual(legacy.call_args.kwargs['sheet_names'], ['External Invoice'])
        self.assertFalse(Path(legacy.call_args.args[0]).exists())
        for key, value in result.items():
            self.assertEqual(response.data[key], value)
        self.assertEqual(ReceivablesSourceSnapshot.objects.count(), 0)
