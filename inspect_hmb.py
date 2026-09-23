import inspect,re
from openpyxl import load_workbook
from apps.process_datasheet.hmb_master_template_parser import parse_hmb_case_workbook,_pick_case_sheet,_detect_case_layout,HMB_CASE_PARSER_CONFIG
paths=[r"C:\Users\Mohammed.Agra\OneDrive - Rejlers AB\Desktop\RADAI\Documents\Process\HMB Extractor\Case_A.xlsx",r"C:\Users\Mohammed.Agra\OneDrive - Rejlers AB\Desktop\RADAI\Documents\Process\HMB Extractor\Case_1A_Stream_Report.xlsx",r"C:\Users\Mohammed.Agra\OneDrive - Rejlers AB\Desktop\RADAI\Documents\Process\HMB Extractor\Case 1A Template.xlsx"]
def invoke(fn,obj):
 a=[]
 for p in inspect.signature(fn).parameters.values():
  if p.name in ('ws','worksheet','sheet','wb','workbook'): a.append(obj)
  elif p.name in ('config','parser_config'): a.append(HMB_CASE_PARSER_CONFIG)
  elif p.default is inspect.Parameter.empty: raise TypeError(p.name)
 return fn(*a)
def get(o,k,d=None): return o.get(k,d) if isinstance(o,dict) else getattr(o,k,d)
def out(x): return x if isinstance(x,dict) else vars(x) if hasattr(x,'__dict__') else str(x)
for path in paths:
 print('\n'+'='*100+'\nFILE:',path); wb=load_workbook(path,data_only=False,read_only=False); print('SHEET NAMES:',wb.sheetnames); ws=invoke(_pick_case_sheet,wb); print('SELECTED SHEET:',ws.title); l=invoke(_detect_case_layout,ws); print('LAYOUT:',l)
 for k in ('header_row','property_row','unit_row','start_col','stream_start_row','data_start_row'): print(' ',k,':',get(l,k))
 hr,pr,sc=get(l,'header_row'),get(l,'property_row'),get(l,'start_col') or 1
 ids=[]
 if hr:
  for c in range(sc,ws.max_column+1):
   x=ws.cell(hr,c).value
   if x is not None and str(x).strip() not in ids: ids.append(str(x).strip())
 print('FIRST 20 STREAM IDS:',ids[:20]); props=[]
 if pr:
  for r in range(pr,ws.max_row+1):
   vals=[ws.cell(r,c).value for c in range(1,max(1,sc))]; name=next((x for x in vals if x is not None and str(x).strip()),None)
   if name is not None: props.append((r,str(name).strip(),ws.cell(r,sc-1).value if sc>1 else None))
 print('FIRST 20 PROPERTIES (row, name, unit):',props[:20]); markers=[]
 for r in range(1,ws.max_row+1):
  labels=[]
  for c in range(1,min(ws.max_column,max(1,sc))+1):
   x=ws.cell(r,c).value
   if isinstance(x,str) and re.fullmatch(r'[A-Z][A-Z0-9 _/()&.\-:]{2,}',x.strip()): labels.append(x.strip())
  if labels: markers.append((r,labels))
 print('SECTION MARKER ROWS (uppercase labels):',markers); records=parse_hmb_case_workbook(path); print('RECORD COUNT:',len(records)); print('10-RECORD SAMPLE:')
 for x in records[:10]: print(out(x))
 wb.close()
